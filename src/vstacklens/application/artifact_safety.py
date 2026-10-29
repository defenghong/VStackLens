from __future__ import annotations

import json
import io
import hashlib
from html import unescape
import re
import zipfile
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from docx import Document


SENSITIVE_PATTERNS: tuple[tuple[re.Pattern[str], str], ...] = (
    (re.compile(r"(?i)\btraceback\b"), "包含底层 traceback"),
    (re.compile(r"(?i)\bTimeoutError\b"), "包含底层 TimeoutError"),
    (re.compile(r"(?i)\bsite-packages\b"), "包含本机 Python 路径"),
    (re.compile(r"(?i)[A-Z]:\\Users\\[^\s\"'<>|]*"), "包含本机用户目录路径"),
    (re.compile(r"D:\\ExampleProject(?:\\|$)[^\s\"'<>|]*"), "包含本机源码路径"),
    (re.compile(r"(?i)src\\+vstacklens"), "包含本机源码相对路径"),
)

CREDENTIAL_ASSIGNMENT = re.compile(r'''(?i)["']?\b(?P<field>password|passwd|pwd|api[_-]?key|access[_-]?token|token|cookie|session|secret|authorization)\b["']?\s*[:=]\s*(?:"([^"]*)"|'([^']*)'|(\[[^\]\r\n]*\]|\{[^}\r\n]*\})|([^,\s;<>"'{}\[\]]+))''')
BEARER_CREDENTIAL = re.compile(r"(?i)\bbearer\s+[a-z0-9._\-]+")
REDACTED_VALUES = {"", "none", "null", "[redacted]", "[redacted_credential]", "[已隐藏]", "[敏感字段已隐藏]"}
PRIVATE_KEY = re.compile(r"-----BEGIN (?:[A-Z0-9]+ )*PRIVATE KEY-----", re.IGNORECASE)
CREDENTIAL_URL = re.compile(r"https?://[^\s/:@]+:[^\s/@]+@", re.IGNORECASE)
CREDENTIAL_XML = re.compile(r"(?is)<(?:[\w.-]+:)?(?:password|passwd|pwd|api[_-]?key|access[_-]?token|token|cookie|session|secret|authorization)\b[^>]*>\s*([^<]+)")
PDF_INFO_TEXT_LIMIT = 1024 * 1024

CUSTOMER_FORBIDDEN_PATTERNS: tuple[tuple[re.Pattern[str], str], ...] = (
    (re.compile(r"(?i)\bP4\b"), "客户产物包含 P4"),
    (re.compile(r"(?i)P4\s*级"), "客户产物包含 P4级"),
    (re.compile(r"期望值"), "客户产物包含期望值"),
    (re.compile(r"期望状态"), "客户产物包含期望状态"),
    (re.compile(r"(?i)\bPDF\s*/\s*PPT\b"), "客户产物包含 PDF/PPT"),
    (re.compile(r"(?i)\bPDF\b"), "客户产物包含 PDF"),
    (re.compile(r"(?i)\bPPT\b"), "客户产物包含 PPT"),
    (re.compile(r"(?i)PDF\s*导出"), "客户产物包含 PDF导出"),
    (re.compile(r"(?i)PPT\s*导出"), "客户产物包含 PPT导出"),
)


@dataclass(frozen=True, slots=True)
class ArtifactScanFinding:
    path: Path
    reason: str


class UnsafeArtifactError(RuntimeError):
    def __init__(self, findings: list[ArtifactScanFinding]) -> None:
        self.findings = findings
        super().__init__("交付已阻止：产物含敏感内容或无法完成有效扫描。" + "；".join(f"{item.path.name}: {item.reason}" for item in findings))


def sanitize_text(value: Any, *secrets: str) -> str:
    text = str(value or "")
    for secret in secrets:
        if secret:
            text = re.sub(re.escape(secret), "[已隐藏]", text, flags=re.IGNORECASE)
    text = re.sub(r"(?i)authorization\s*:\s*bearer\s+[^,\s;]+", "[敏感字段已隐藏]", text)
    text = re.sub(r"(?i)\bauthorization[\"']?\s*[:=]\s*[\"']?(?:basic|bearer)\s+[^\s,\"';]+", "[敏感字段已隐藏]", text)
    text = re.sub(r"(?i)(https?://)[^\s/:]+:[^\s@]+@", r"\1[敏感字段已隐藏]@", text)
    text = re.sub(r"(?i)\b(password|passwd|pwd|api[_-]?key|access[_-]?token|token|authorization|cookie|session|secret)\b\s*[:=]\s*[^,\s;]+", "[敏感字段已隐藏]", text)
    text = re.sub(r"(?i)bearer\s+[a-z0-9._\-]+", "[敏感字段已隐藏]", text)
    text = re.sub(r"(?i)\b(password|passwd|pwd|token|authorization|cookie|session|secret)\b", "[敏感字段已隐藏]", text)
    text = re.sub(r"(?i)[A-Z]:\\Users\\[^\s;\"'<>|]+", "[本机路径已隐藏]", text)
    text = re.sub(r"(?i)[A-Z]:\\[^\s;\"'<>|]*(?:site-packages|codex_extracted|ExampleProject)[^\s;\"'<>|]*", "[本机路径已隐藏]", text)
    text = re.sub(r"(?i)\bsrc\\+vstacklens[^\s;\"'<>|]*", "[本机路径已隐藏]", text)
    text = re.sub(r"(?i)\btraceback\b", "异常摘要", text)
    text = re.sub(r"(?i)\bTimeoutError\b", "连接超时", text)
    return text[:4000]


def scan_artifacts(paths: list[Path], *, engineering_paths: set[Path] | None = None, secrets: tuple[str, ...] = (), customer_copy: bool = True) -> list[ArtifactScanFinding]:
    engineering_paths = {path.resolve() for path in engineering_paths or set()}
    findings: list[ArtifactScanFinding] = []
    for path in paths:
        if not path:
            continue
        if not path.is_file():
            findings.append(ArtifactScanFinding(path, "无法扫描：文件不存在或不是文件"))
            continue
        try:
            text = unescape(_artifact_text(path))
            is_customer = customer_copy and path.resolve() not in engineering_paths
            customer_text = unescape(_artifact_text(path, customer_wording=True)) if is_customer else ""
        except Exception as exc:
            findings.append(ArtifactScanFinding(path, f"无法扫描：{type(exc).__name__}"))
            continue
        resolved = path.resolve()
        matched_reasons: set[str] = set()
        is_customer = customer_copy and resolved not in engineering_paths
        # Credential checks are shared by every artifact role. Customer-only
        # checks concern wording and diagnostic disclosure, not secret strength.
        for pattern, reason in SENSITIVE_PATTERNS if is_customer else ():
            if pattern.search(text):
                matched_reasons.add(reason)
        if BEARER_CREDENTIAL.search(text):
            matched_reasons.add("包含 Bearer token")
        for match in CREDENTIAL_ASSIGNMENT.finditer(text):
            value = next((part for part in match.groups()[1:] if part is not None), "").strip()
            if len(value) >= 2 and (value[0], value[-1]) in {("[", "]"), ("{", "}")} and not value[1:-1].strip():
                continue
            if value.casefold() not in REDACTED_VALUES:
                matched_reasons.add("包含凭据或会话字段")
                if match.group("field").casefold() == "authorization":
                    matched_reasons.add("包含 Authorization 敏感字段")
        if PRIVATE_KEY.search(text):
            matched_reasons.add("包含私钥")
        if CREDENTIAL_URL.search(text):
            matched_reasons.add("包含 URL 凭据")
        if any(match.group(1).strip().casefold() not in REDACTED_VALUES for match in CREDENTIAL_XML.finditer(text)):
            matched_reasons.add("包含 XML 凭据或会话值")
        if any(secret and secret.casefold() in text.casefold() for secret in secrets):
            matched_reasons.add("包含本次会话凭据")
        for pattern, reason in CUSTOMER_FORBIDDEN_PATTERNS if is_customer else ():
            if pattern.search(customer_text):
                matched_reasons.add(reason)
        findings.extend(ArtifactScanFinding(path, reason) for reason in sorted(matched_reasons))
    return findings


def _artifact_text(path: Path, *, customer_wording: bool = False) -> str:
    return _artifact_bytes_text(path.read_bytes(), path.suffix.lower(), customer_wording=customer_wording)


@lru_cache(maxsize=1)
def _default_docx_thumbnail_sha() -> str:
    import docx
    template = Path(docx.__file__).parent / "templates" / "default.docx"
    with zipfile.ZipFile(template) as archive:
        return hashlib.sha256(archive.read("docProps/thumbnail.jpeg")).hexdigest()


def _pdf_info_projection(document: Any) -> dict[str, Any] | None:
    """Decode trailer.Info only; document.metadata omits arbitrary Info keys."""
    import pymupdf

    api = pymupdf.mupdf
    pdf = api.PdfDocument(document.this)
    root = api.pdf_dict_gets(api.pdf_trailer(pdf), "Info")
    node_count = 0
    text_chars = 0
    reference_stack: set[int] = set()

    def text(value: str) -> str:
        nonlocal text_chars
        text_chars += len(value)
        if text_chars > PDF_INFO_TEXT_LIMIT:
            raise ValueError("PDF Info text budget exceeded")
        return value

    def visit(value: Any, depth: int) -> Any:
        nonlocal node_count
        node_count += 1
        if depth > 8 or node_count > 512:
            raise ValueError("PDF Info traversal budget exceeded")
        if not value.m_internal:
            return None
        if api.pdf_is_indirect(value):
            xref = api.pdf_to_num(value)
            if xref in reference_stack or not 0 < xref < document.xref_length():
                raise ValueError("PDF Info cyclic or invalid reference")
            if document.xref_is_stream(xref):
                raise ValueError("PDF Info stream cannot be scanned")
            reference_stack.add(xref)
            try:
                resolved = api.pdf_load_object(pdf, xref)
                if not resolved.m_internal:
                    raise ValueError("PDF Info unresolved reference")
                return visit(resolved, depth + 1)
            finally:
                reference_stack.remove(xref)
        if api.pdf_is_null(value):
            return None
        if api.pdf_is_dict(value):
            return {text(api.pdf_to_name(api.pdf_dict_get_key(value, index))): visit(api.pdf_dict_get_val(value, index), depth + 1)
                    for index in range(api.pdf_dict_len(value))}
        if api.pdf_is_array(value):
            return [visit(api.pdf_array_get(value, index), depth + 1) for index in range(api.pdf_array_len(value))]
        if api.pdf_is_string(value):
            return text(api.pdf_to_text_string(value))
        if api.pdf_is_name(value):
            return text(api.pdf_to_name(value))
        if api.pdf_is_bool(value):
            return bool(api.pdf_to_bool(value))
        if api.pdf_is_int(value):
            return api.pdf_to_int64(value)
        if api.pdf_is_real(value):
            return api.pdf_to_real(value)
        raise ValueError("PDF Info value type cannot be scanned")

    info = visit(root, 0)
    if info is not None and not isinstance(info, dict):
        raise ValueError("PDF Info must be a dictionary")
    return info


def _artifact_bytes_text(content: bytes, suffix: str, depth: int = 0, *, customer_wording: bool = False) -> str:
    if depth > 8:
        raise ValueError("nested artifact depth limit reached")
    if suffix in {".docx", ".zip"}:
        chunks = []
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if archive.testzip() is not None:
                raise ValueError("archive integrity check failed")
            for name in archive.namelist():
                if name.endswith("/"):
                    continue
                if not customer_wording:
                    chunks.append(name)
                member_suffix = ".rels" if name.rsplit("/", 1)[-1] == ".rels" else Path(name).suffix.lower()
                if suffix == ".docx" and name == "docProps/thumbnail.jpeg":
                    if hashlib.sha256(archive.read(name)).hexdigest() == _default_docx_thumbnail_sha():
                        continue
                    raise ValueError("DOCX thumbnail cannot be scanned")
                if suffix == ".docx" and name.startswith(("word/media/", "word/embeddings/")) and not name.endswith("/"):
                    from vstacklens.resources import package_resource_path
                    logo = package_resource_path("assets", "icons", "vstacklens.png")
                    if name.startswith("word/media/") and logo.is_file() and hashlib.sha256(archive.read(name)).digest() == hashlib.sha256(logo.read_bytes()).digest():
                        continue
                    raise ValueError("embedded media cannot be scanned")
                if member_suffix in {".xml", ".rels"}:
                    raw = archive.read(name)
                    root = ElementTree.fromstring(raw)
                    chunks.extend(["".join(root.itertext()), raw.decode("utf-8")])
                elif member_suffix in {".html", ".json", ".ndjson", ".txt", ".md", ".js", ".css", ".svg", ".docx", ".pdf", ".zip"}:
                    chunks.append(_artifact_bytes_text(archive.read(name), member_suffix, depth + 1, customer_wording=customer_wording))
                else:
                    raise ValueError("archive member format cannot be scanned")
        if not chunks:
            raise ValueError("archive has no scannable content")
        return "\n".join(chunks)
    if suffix == ".pdf":
        import pymupdf
        with pymupdf.open(stream=content, filetype="pdf") as document:
            if document.needs_pass:
                raise ValueError("encrypted PDF cannot be scanned")
            texts = [page.get_text() for page in document]
            if any(not text.strip() and page.get_images() for text, page in zip(texts, document)):
                raise ValueError("PDF image-only page cannot be scanned")
            if not any(text.strip() for text in texts):
                raise ValueError("PDF has no readable text; scan is unavailable")
            # These two values are computed by the parser, rather than stored
            # document metadata. In particular format='PDF 1.7' is not wording
            # supplied by the report. Keep all actual metadata and attachments.
            metadata = {key: value for key, value in document.metadata.items() if key not in {"format", "encryption"}}
            texts.append(json.dumps(metadata, ensure_ascii=False))
            # Preserve actual custom keys and decoded string values, including
            # indirect/nested fields. Limits and unsupported representations
            # raise instead of silently accepting an incompletely scanned PDF.
            texts.append(json.dumps(_pdf_info_projection(document), ensure_ascii=False))
            xmp = document.get_xml_metadata()
            if xmp:
                root = ElementTree.fromstring(xmp)
                if customer_wording:
                    for node in root.iter():
                        if node.tag == "{http://purl.org/dc/elements/1.1/}format" and (node.text or "").strip().lower() == "application/pdf":
                            node.text = ""  # Standard MIME descriptor, not report wording; raw XMP is still security-scanned.
                texts.append(" ".join(root.itertext()))
                texts.extend(value for node in root.iter() for value in node.attrib.values())
                if not customer_wording:
                    texts.append(xmp)
            for name in document.embfile_names():
                texts.extend([name, json.dumps(document.embfile_info(name), ensure_ascii=False)])
                texts.append(_artifact_bytes_text(document.embfile_get(name), Path(name).suffix.lower(), depth + 1, customer_wording=customer_wording))
            return "\n".join(texts)
    if suffix == ".json":
        return json.dumps(json.loads(content.decode("utf-8")), ensure_ascii=False)
    if suffix == ".ndjson":
        return "\n".join(json.dumps(json.loads(line), ensure_ascii=False) for line in content.decode("utf-8").splitlines() if line.strip())
    if suffix in {".md", ".html", ".txt", ".log", ".xml", ".rels", ".js", ".css", ".svg"}:
        return content.decode("utf-8")
    raise ValueError("artifact format cannot be scanned")


def ensure_artifacts_safe(paths: list[Path], *, secrets: tuple[str, ...] = ()) -> None:
    findings = scan_artifacts(paths, secrets=secrets, customer_copy=False)
    if findings:
        raise UnsafeArtifactError(findings)


def _docx_table_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)


def _zip_text(path: Path) -> str:
    return _artifact_bytes_text(path.read_bytes(), ".zip")


def summarize_scan_findings(findings: list[ArtifactScanFinding]) -> list[dict[str, str]]:
    return [{"file": finding.path.name, "reason": finding.reason} for finding in findings]


def json_summary(value: Any) -> str:
    return sanitize_text(json.dumps(value, ensure_ascii=False, indent=2))
