"""Batch-three Word visual primitives on the batch-four semantic document."""
from io import BytesIO
import re
from copy import deepcopy
import pymupdf
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt,Cm
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.text.run import Run
from vstacklens.reports.document_word import WordDocumentWriter,_font,_bookmark,_field
from vstacklens.reports.document_model import COLORS,NAVY,INK,NEUTRAL
from vstacklens.reports.document_visual import format_reference

SURFACE='E9EEF3'

class VisualWordWriter(WordDocumentWriter):
    def __init__(self,engine,report,model):
        super().__init__(engine,report,model)
        # Native TOC entries must not inherit Normal's extra paragraph gap.
        # Text remains 10.5pt at the existing 1.15 line spacing.
        for level in (1,2,3):
            name=f'TOC {level}'
            style=self.document.styles[name] if name in self.document.styles else self.document.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH,builtin=True)
            style.base_style=self.document.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(10.5)
            style.paragraph_format.line_spacing=1.15
            style.paragraph_format.space_before=Pt(0);style.paragraph_format.space_after=Pt(0)
            style.paragraph_format.left_indent=Pt((level-1)*21)

    def linked_text(self,p,text,size=10.5,bold=False,color=INK):
        letters={a['letter'] for a in self.model['appendices']};cursor=0
        for match in re.finditer(r'附录\s+([A-Z]+)',str(text)):
            if match.group(1) not in letters:continue
            _font(p.add_run(str(text)[cursor:match.start()]),size,bold,color)
            link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'d_appendix_'+match.group(1))
            run=OxmlElement('w:r');link.append(run);p._p.append(link)
            native=Run(run,p);native.text=match.group();_font(native,size,bold,NAVY)
            cursor=match.end()
        _font(p.add_run(str(text)[cursor:]),size,bold,color)

    def paragraph(self,text,*,bold=False,size=10.5,color=INK,keep=False):
        if text is None or text=='':return None
        text=format_reference(text,'word')
        p=super().paragraph(text,bold=bold,size=size,color=color,keep=keep)
        if re.search(r'附录\s+[A-Z]+',text):
            for child in list(p._p):
                if child.tag!=qn('w:pPr'):p._p.remove(child)
            self.linked_text(p,text,size,bold,color)
        return p
    def _cell(self,cell,fill=SURFACE,padding=120):
        self.engine._shade_cell(cell,fill)
        pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
        for side in ('left','right','top','bottom'):
            n=OxmlElement('w:'+side);n.set(qn('w:w'),str(padding));n.set(qn('w:type'),'dxa');m.append(n)
        pr.append(m)

    def cp(self,cell,text='',size=10.5,bold=False,color=INK,first=False,after=4):
        p=cell.paragraphs[0] if first else cell.add_paragraph()
        p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=1.15
        self.linked_text(p,format_reference(str(text),'word'),size,bold,color);return p

    def frame(self,columns=1):
        self._before_table();t=self.document.add_table(rows=1,cols=columns);t.autofit=False
        self.fixed_width(t,[int(self.width*20/columns)]*columns)
        caption=OxmlElement('w:tblCaption');caption.set(qn('w:val'),'视觉卡片布局');t._tbl.tblPr.append(caption)
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        return t

    def metric_cards(self,items,*,cover=False):
        if not items:return
        t=self.frame(len(items))
        for c,item in zip(t.rows[0].cells,items):
            self._cell(c,padding=140);self.cp(c,item['label'],color=NEUTRAL,first=True)
            self.cp(c,item['value'],size=22 if cover else 17,bold=True)
            if item['detail']:self.cp(c,item['detail'],size=10,color=NEUTRAL)
        return t

    def story(self,text):
        if not text:return
        t=self.frame();c=t.cell(0,0);self._cell(c,'D5E3EF')
        self.cp(c,'检查解读',bold=True,color=NAVY,first=True);self.cp(c,text)

    def bar(self,cell,value,maximum,*,threshold=None,width=None):
        """Native editable geometry; values are supplied by the common model."""
        width=width or self.width-24
        if maximum<=0:return
        stops=sorted({0.,float(maximum),max(0.,min(float(value),maximum)),*([float(threshold)] if threshold is not None and 0<threshold<maximum else [])})
        segments=[(a,b) for a,b in zip(stops,stops[1:]) if b>a]
        if not segments:return
        t=cell.add_table(rows=1,cols=len(segments));t.autofit=False
        self.fixed_width(t,[max(1,int((b-a)/maximum*width*20)) for a,b in segments])
        cap=OxmlElement('w:tblCaption');cap.set(qn('w:val'),'图形几何');t._tbl.tblPr.append(cap)
        for c,(a,b) in zip(t.rows[0].cells,segments):
            self._cell(c,'506B84' if a<value and threshold is not None and value>threshold else '8CA9BF' if a<value else 'F1F5F9',padding=0)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(10.5)
            if threshold is not None and b==threshold:
                borders=OxmlElement('w:tcBorders');line=OxmlElement('w:right')
                for k,v in [('val','single'),('sz','8'),('color',NAVY.lstrip('#'))]:line.set(qn('w:'+k),v)
                borders.append(line);c._tc.get_or_add_tcPr().append(borders)
        # Word requires a trailing paragraph after a nested table. It has no
        # text; keep its caret from adding a blank body-text line per bar.
        trailing=cell.paragraphs[-1]
        if not trailing.text:
            trailing.paragraph_format.space_after=Pt(0);trailing.paragraph_format.line_spacing=Pt(1)

    def graph_cells(self,nodes):
        if not nodes:return
        t=self.frame(len(nodes));width=self.width/len(nodes)-24
        for cell,g in zip(t.rows[0].cells,nodes):
            self._cell(cell);self.cp(cell,g['title'],size=11.5,bold=True,color=NAVY,first=True)
            if g.get('equal_text'):self.cp(cell,g['equal_text']);continue
            for r in g['series']:
                p=self.cp(cell,r['name']+(' · '+r['detail'] if r.get('detail') else ''),size=10.5)
                _font(p.add_run(' · '+r['text']),10.5,True)
                self.bar(cell,r['value'],g['axis_max'],threshold=g['threshold'],width=width)
            self.cp(cell,'  ·  '.join(t['text'] for t in g['ticks']),size=10,color=NEUTRAL)
            if g.get('more'):self.cp(cell,format_reference(g['more'],'word'),size=10,color=NEUTRAL)

    def cover(self):
        d=self.model;v=d['visual'];t=self.frame();t.rows[0].height=Cm(8.6);t.rows[0].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        c=t.cell(0,0);self._cell(c,NAVY.lstrip('#'),padding=260)
        self.cp(c,'客户交付报告',size=11,bold=True,color='#BBD4E9',first=True)
        if v['cover_label']:self.cp(c,v['cover_label'],color='#FFFFFF')
        self.cp(c,v['display_title'],size=24,bold=True,color='#FFFFFF')
        for label,value in [('客户名称',d['customer']),('巡检日期',d['date']),('平台版本',v['platform'])]:
            if value:self.cp(c,label+'：'+value,color='#FFFFFF')
        if d['opening']['statistics']:
            self.cp(c,'执行摘要',bold=True,color='#FFFFFF');self.cp(c,d['opening']['statistics'],color='#FFFFFF')
        self.metric_cards(v['opening_kpis'],cover=True)
        t=self.frame();c=t.cell(0,0);self._cell(c,'EDF3F8',padding=160)
        self.cp(c,'总体结论：'+d['opening']['conclusion'],size=13,bold=True,color=NAVY,first=True)
        if d['opening']['priorities']:
            self.cp(c,'需优先处理的问题',bold=True)
            for r in d['opening']['priorities']:
                p=c.add_paragraph();p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.15
                _font(p.add_run(r['level']+'  '),10.5,True,COLORS[r['level']]);_font(p.add_run(r['title']+f' · {r["objects"]} 个受影响对象'))
        if d['opening']['more']:self.cp(c,d['opening']['more'])
        self.cp(c,d['opening']['statistics'])
        self.page_break()

    def labeled(self,label,text):
        p=self.paragraph('',keep=True)
        if p is None:p=self.document.add_paragraph()
        p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.15
        # Leave a full character after the longest label; at 52.5pt its
        # trailing tab could jump to Word's next default stop instead.
        p.paragraph_format.left_indent=Pt(63);p.paragraph_format.first_line_indent=Pt(-63)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Pt(63));p.paragraph_format.keep_with_next=True
        _font(p.add_run(label+'：'),bold=True);_font(p.add_run('\t'+text));return p

    def reference(self,text,letter=''):
        text=format_reference(text,'word')
        if not text:return
        p=self.paragraph(text,color=NEUTRAL)
        if letter:
            anchor='d_appendix_'+letter
            for node in list(p._p):
                if node.tag!=qn('w:pPr'):p._p.remove(node)
            link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),anchor)
            run=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),NAVY.lstrip('#'));rp.append(color)
            run.append(rp);node=OxmlElement('w:t');node.text=text;run.append(node);link.append(run);p._p.append(link)

    def problem(self,p):
        start=len(self.document.paragraphs)
        h=self.heading(p['title'],2,'problem-'+p['id']);borders=OxmlElement('w:pBdr');left=OxmlElement('w:left')
        for k,v in [('val','single'),('sz','12'),('space','5'),('color',COLORS[p['level']].lstrip('#'))]:left.set(qn('w:'+k),v)
        borders.append(left);h._p.get_or_add_pPr().append(borders)
        self.paragraph('等级：'+p['level'],color=COLORS[p['level']],keep=True)
        if p.get('capacity_note'):self.paragraph(p['capacity_note'],keep=True)
        for f in p['fields']:self.labeled(f['label'],f['text'])
        for group in p['render_groups']:
            if group['title']:self.paragraph(group['title'],bold=True,keep=True)
            for i,step in enumerate(group['steps']):self.paragraph(f'{i+1}. {step}',keep=True)
        if p['visual']['meta']:self.paragraph(' · '.join(x['label']+'：'+x['value'] for x in p['visual']['meta']),color=NEUTRAL,size=10,keep=True)
        # End the keep chain at the review metadata, before the table. Keep
        # wrapped field/step paragraphs intact as part of the same preamble.
        preamble=self.document.paragraphs[start:]
        for index,paragraph in enumerate(preamble):
            paragraph.paragraph_format.keep_together=True
            paragraph.paragraph_format.keep_with_next=index<len(preamble)-1
        representative=p['representative_table']
        if representative:self.table({**representative,'_problem_layout':True},keep_after=bool(p['reference']))
        self.reference(p['reference'],p['appendix'])
        for note in p['object_notes']:
            if note.startswith('NSX Manager 的资源预留'):self.paragraph(note)

    def cluster_cards(self,node):
        # Portrait Word uses equal-width pairs; a four-column portrait row can
        # exceed a page solely because long host names wrap in tiny cells.
        rows=node['rows']
        for offset in range(0,len(rows),2):
            pair=rows[offset:offset+2];t=self.frame(len(pair))
            for cell,row in zip(t.rows[0].cells,pair):
                c=self.model['visual']['clusters'][row['key']];self._cell(cell)
                self.cp(cell,c['name'],size=11.5,bold=True,color=NAVY,first=True);self.cp(cell,c['summary'],size=10,color=NEUTRAL)
                for f in c['features']+c['fields']:self.cp(cell,f['label']+'：'+f['value'])
                for g in c['graphs']:
                    self.cp(cell,g['title']+' · '+g['text'],size=10.5,bold=True);self.bar(cell,g['value'],100,width=self.width/len(pair)-24)
                if c['host_names']:self.cp(cell,'主机：'+c['host_names'],size=10,color=NEUTRAL)
                if c['cpu_models']:self.cp(cell,'CPU 型号：'+c['cpu_models'],size=10,color=NEUTRAL)

    def top_grid(self,nodes):
        for offset in range(0,len(nodes),2):
            pair=nodes[offset:offset+2];t=self.frame(len(pair));width=self.width/len(pair)-24
            for cell,node in zip(t.rows[0].cells,pair):
                self._cell(cell,padding=80);self.cp(cell,node['title'],size=11.5,bold=True,color=NAVY,first=True,after=2)
                for note in node['notes']:self.cp(cell,note,size=10,color=NEUTRAL,after=2)
                for row in self.model['visual']['top'][node['id']]:
                    p=self.cp(cell,str(row['rank'])+'. '+row['name'],bold=True,after=2)
                    _font(p.add_run(' · '+row['display']),10.5,True)
                    if row['meta_text']:self.cp(cell,row['meta_text'],size=10,color=NEUTRAL,after=2)
                    self.bar(cell,row.get('bar_pct',0),100,width=width)

    def scope(self):
        d=self.model;v=d['visual']
        if not v['has_execution_stats']:return
        self.metric_cards([{'label':a,'value':str(d['statistics'][k]),'detail':''} for a,k in [('检查总数','total'),('通过','passed'),('未通过','failed')]])
        self.paragraph('通过比例 '+v['scope']['pass_percent_text'],color=NEUTRAL)
        t=self.frame();self._cell(t.cell(0,0),'FFFFFF');self.bar(t.cell(0,0),v['scope']['pass_percent'],100)
        t=self.frame();c=t.cell(0,0);self._cell(c,'EDF3F8');self.cp(c,'未通过检查结果按等级分布',bold=True,first=True)
        for i,r in enumerate(v['scope']['level_distribution'] or []):
            self.cp(c,f'{r["level"]}  {r["count"]} 条 · {r["share_pct_text"]}',bold=True,color=COLORS[r['level']])
        rows=[r for r in v['scope']['level_distribution'] or [] if r['count']>0]
        if rows:
            strip=c.add_table(rows=1,cols=len(rows));strip.autofit=False
            self.fixed_width(strip,[max(1,int((self.width-24)*20*r['count']/v['scope']['problem_checks'])) for r in rows])
            caption=OxmlElement('w:tblCaption');caption.set(qn('w:val'),'图形几何');strip._tbl.tblPr.append(caption)
            for cell,row in zip(strip.rows[0].cells,rows):
                self._cell(cell,COLORS[row['level']].lstrip('#'),padding=0)
                cell.paragraphs[0].paragraph_format.space_after=Pt(0);cell.paragraphs[0].paragraph_format.line_spacing=Pt(10.5)

    def vm_state_card(self):
        c=self.model['visual'].get('vm_state_card')
        if not c:return
        t=self.frame();cell=t.cell(0,0);self._cell(cell)
        self.cp(cell,c['title'],bold=True,color=NAVY,first=True)
        self.cp(cell,' · '.join(s['label']+' '+str(s['value']) for s in c['segments']))
        rows=[s for s in c['segments'] if s['value']>0]
        if c['total']>0 and rows:
            strip=cell.add_table(rows=1,cols=len(rows));strip.autofit=False
            self.fixed_width(strip,[max(1,int((self.width-24)*20*s['value']/c['total'])) for s in rows])
            caption=OxmlElement('w:tblCaption');caption.set(qn('w:val'),'图形几何');strip._tbl.tblPr.append(caption)
            for target,s in zip(strip.rows[0].cells,rows):
                self._cell(target,s['color'].lstrip('#'),padding=0)
                target.paragraphs[0].paragraph_format.space_after=Pt(0);target.paragraphs[0].paragraph_format.line_spacing=Pt(10.5)

    def env_group(self,g):
        v=self.model['visual'];heading=self.heading(g['title'],2,g['id'])
        if g['id']=='resource-top':
            heading.paragraph_format.space_before=Pt(2);heading.paragraph_format.space_after=Pt(4)
        if g['id']=='scale-clusters':
            self.metric_cards(v['asset_kpis']);self.vm_state_card();self.metric_cards(v['snapshots'])
            if v['narratives'].get('clusters'):self.story(v['narratives']['clusters'])
        if g['id']=='host-resources':
            self.metric_cards(v['host_kpis'])
            if v['narratives'].get('hosts'):self.story(v['narratives']['hosts'])
            self.graph_cells([v[k] for k in ('allocation_cpu','allocation_memory') if k in v])
            self.graph_cells([n for n in g['blocks'] if n['kind']=='chart' and n['id']!='memory_allocation_pct'])
        if g['id']=='storage':self.metric_cards(v['vsan_cards'])
        if g['id']=='resource-top':self.top_grid(g['blocks']);return
        if g['id'] in ('network','security-certificates'):
            self.story(v['group_summaries'][g['id']]['references']['word']['text'])
        for node in g['blocks']:
            if node['id'] in ('scale','snapshot-overview') or (g['id']=='host-resources' and node['kind']=='chart'):continue
            if node['kind']=='cards':self.cluster_cards(node)
            elif node['kind']=='table':self.table({**node,'notes':[format_reference(x,'word') for x in node.get('notes',[])]})
            elif node['kind']=='paragraph' and g['id'] not in ('network','security-certificates'):
                if node['id'] in ('environment-identity','capacity-note','judgment-boundary'):
                    self.paragraph(format_reference(node['text'],'word'),color=NEUTRAL)
                else:self.story(format_reference(node['text'],'word'))
        if g['id']=='scale-clusters':
            if v['cluster_coverage']:
                t=self.frame(len(v['cluster_coverage']))
                for cell,c in zip(t.rows[0].cells,v['cluster_coverage']):
                    self._cell(cell);self.cp(cell,c['title']+' · '+f'{c["value"]}/{c["total"]}',bold=True,color=NAVY,first=True)
                    self.bar(cell,c['value'],c['total'],width=self.width/len(v['cluster_coverage'])-24)
                    if c['text']:self.cp(cell,c['text'],size=10,color=NEUTRAL)
        if g['id']=='storage':
            for chart in v['vsan_charts']:
                t=self.frame();c=t.cell(0,0);self._cell(c)
                self.cp(c,chart['title'],bold=True,color=NAVY,first=True)
                self.cp(c,chart['text'])
                self.cp(c,' · '.join(r['label']+' '+r['display'] for r in chart['segments']))
                if chart['total']>0:self.bar(c,chart['segments'][0]['value'],chart['total'])

    def render(self,path):
        d=self.model;self.cover()
        toc_section=self.document.add_section(WD_SECTION_START.CONTINUOUS)
        columns=toc_section._sectPr.find(qn('w:cols'))
        if columns is None:columns=OxmlElement('w:cols');toc_section._sectPr.append(columns)
        columns.set(qn('w:num'),'1')
        p=self.document.add_paragraph();self._apply_break(p);_font(p.add_run('目录'),16,True,NAVY);p.paragraph_format.keep_with_next=True
        _field(self.document.add_paragraph(),' TOC \\o "1-3" \\h \\z \\u ')
        body_section=self.document.add_section(WD_SECTION_START.CONTINUOUS)
        columns=body_section._sectPr.find(qn('w:cols'))
        if columns is None:columns=OxmlElement('w:cols');body_section._sectPr.append(columns)
        columns.set(qn('w:num'),'1')
        for i,s in enumerate(d['sections'],1):
            self.heading(f'{i:02d}  '+s['title'],1,s['id'])
            subtitle=d['visual']['section_subtitles'].get(s['kind'],'')
            if subtitle:self.paragraph(subtitle,color=NEUTRAL,keep=True)
            if s['kind']=='overview':
                self.scope();self.table({'id':'overview','title':s['title'],'weights':[1,5,1],'headers':['等级','问题类别','受影响对象'],
                  'rows':[{'id':'ov'+str(j),'key':'ov'+str(j),'cells':[r['level'],r['title'],str(r['objects'])]} for j,r in enumerate(s['rows'])]})
            elif s['kind']=='problems':
                for p in s['items']:self.problem(p)
            elif s['kind']=='checklist':
                for j in range(0,len(s['items']),2):self.paragraph('   '.join('✓ '+r['label'] for r in s['items'][j:j+2]),keep=j<4)
            elif s['kind']=='environment':
                for g in s['groups']:self.env_group(g)
        if d['appendices']:
            self.page_break()
            for table in d['appendices']:self.table(table,appendix=True)
        self.document.save(path);self.engine.last_document_trace=self.trace
