from __future__ import annotations

from pathlib import Path

from vstacklens.reports.report_model import ReportData


class PptxReportEngine:
    def render(self, report: ReportData, output_path: Path) -> None:
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
        except ImportError as exc:
            raise RuntimeError("PPTX export requires python-pptx. Install or bundle python-pptx for PPTX output.") from exc

        prs = Presentation()
        self._title_slide(prs, report, Inches, Pt)
        self._environment_slide(prs, report, Inches, Pt)
        self._risk_slide(prs, report, Inches, Pt)
        self._top_risks_slide(prs, report, Inches, Pt)
        self._roadmap_slide(prs, report, Inches, Pt)
        self._summary_slide(prs, report, Inches, Pt)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_path)

    def _title_slide(self, prs, report: ReportData, Inches, Pt) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = report.customer_info.customer_name
        slide.placeholders[1].text = f"{report.customer_info.project_name}\n健康评分：{report.health_score.score}"

    def _environment_slide(self, prs, report: ReportData, Inches, Pt) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "环境概览"
        text = "\n".join(f"{k}: {v}" for k, v in report.asset_inventory.get("summary", {}).items())
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.5), Inches(4.2))
        box.text_frame.text = text

    def _risk_slide(self, prs, report: ReportData, Inches, Pt) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "风险分布"
        text = f"P1: {report.risk_summary.P1}\nP2: {report.risk_summary.P2}\nP3: {report.risk_summary.P3}\n优化建议: {report.risk_summary.P4}"
        slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(4.0), Inches(3.0)).text_frame.text = text

    def _top_risks_slide(self, prs, report: ReportData, Inches, Pt) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "TOP10 风险"
        text = "\n".join(f"{item.risk_level} {item.title} - {item.object_name}" for item in report.findings[:10])
        slide.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(9.0), Inches(5.2)).text_frame.text = text or "无风险"

    def _roadmap_slide(self, prs, report: ReportData, Inches, Pt) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "整改路线图"
        slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.5), Inches(3.0)).text_frame.text = "立即整改：P1/P2\n30天整改：P3\n长期优化：容量与配置优化建议"

    def _summary_slide(self, prs, report: ReportData, Inches, Pt) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "总结"
        slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.5), Inches(3.0)).text_frame.text = report.health_score.explanation
