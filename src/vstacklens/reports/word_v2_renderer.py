from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from vstacklens.reports.word_v2_view_model import (
    ConclusionView,
    OverviewView,
    RiskItemView,
    VsanView,
    WordV2ViewModel,
)
from vstacklens.reports.word_v2_templates import DetailTable, TableColumn


PAGE_WIDTH_DXA = 9650
HEADER_BLUE = "17365D"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F7FA"
BORDER_GRAY = "D9E1E8"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(90, 90, 90)


class WordV2Renderer:
    """Render the fixed stage-one Word V2 customer report structure."""

    def render(self, model: WordV2ViewModel, output_path: Path) -> Path:
        if model.validation_errors:
            raise ValueError("Word V2 data validation failed: " + "; ".join(model.validation_errors))
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        self._document = document
        self._setup_document(document, model)
        self._render_cover(document, model)
        document.add_page_break()
        self._render_toc(document, model)
        document.add_page_break()
        self._render_overview(document, model.overview)
        self._render_result_analysis(document, model)
        self._render_risk_details(document, model.risk_items)
        self._render_vsan(document, model.vsan)
        document.add_page_break()
        self._render_conclusion(document, model.conclusion)
        document.save(output_path)
        return output_path

    def _setup_document(self, document: Document, model: WordV2ViewModel) -> None:
        section = document.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(16)
        section.bottom_margin = Mm(16)
        section.left_margin = Mm(16)
        section.right_margin = Mm(16)
        section.header_distance = Mm(8)
        section.footer_distance = Mm(8)
        section.different_first_page_header_footer = True

        styles = document.styles
        self._set_style(styles["Normal"], 10.5, False)
        styles["Normal"].paragraph_format.line_spacing = 1.12
        styles["Normal"].paragraph_format.space_after = Pt(4)
        self._set_style(styles["Title"], 26, True)
        self._remove_style_borders(styles["Title"])
        for name, size in (("Heading 1", 17), ("Heading 2", 13.5)):
            self._set_style(styles[name], size, True)
            styles[name].paragraph_format.keep_with_next = True
            styles[name].paragraph_format.keep_together = True
            styles[name].paragraph_format.space_before = Pt(10)
            styles[name].paragraph_format.space_after = Pt(4)

        customer = model.cover.customer_name
        header = section.header.paragraphs[0]
        header.paragraph_format.tab_stops.add_tab_stop(Mm(178), WD_TAB_ALIGNMENT.RIGHT)
        left = header.add_run("VStackLens 虚拟化平台健康巡检报告")
        self._format_run(left, 8.5, False, GRAY)
        right = header.add_run(f"\t{customer}")
        self._format_run(right, 8.5, False, GRAY)
        section.first_page_header.paragraphs[0].text = ""

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        prefix = footer.add_run(f"保密资料 | 巡检日期：{model.cover.inspection_date} | 第 ")
        self._format_run(prefix, 8, False, GRAY)
        self._append_field(footer, "PAGE")
        suffix = footer.add_run(" 页")
        self._format_run(suffix, 8, False, GRAY)
        section.first_page_footer.paragraphs[0].text = ""

    def _render_cover(self, document: Document, model: WordV2ViewModel) -> None:
        for _ in range(4):
            document.add_paragraph()
        brand = document.add_paragraph()
        brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._format_run(brand.add_run("VStackLens"), 16, True, RGBColor(47, 117, 181))

        title = document.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._format_run(title.add_run(model.cover.report_title), 26, True, BLACK)

        metadata = [
            ["客户名称", model.cover.customer_name],
            ["巡检日期", model.cover.inspection_date],
            ["报告版本", model.cover.report_version],
            ["生成时间", model.cover.generated_at],
        ]
        self._add_table(document, ["项目", "内容"], metadata, widths=[2200, 7450], header=False)

    def _render_toc(self, document: Document, model: WordV2ViewModel) -> None:
        self._heading(document, "目录", 1)
        for section in model.toc_sections:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Mm(5)
            paragraph.paragraph_format.space_after = Pt(4)
            self._format_run(paragraph.add_run(section), 11, False, BLACK)

    def _render_overview(self, document: Document, view: OverviewView) -> None:
        self._heading(document, "1. 巡检概况", 1)
        self._heading(document, "1.1 巡检范围", 2)
        scope_rows = [
            [
                item["serial_no"], item["vcenter_name"], item["management_address"],
                item["esxi_host_count"], item["datastore_count"], item["vm_count"],
            ]
            for item in view.scope_rows
        ]
        self._add_table(document, ["序号", "vCenter名称", "管理地址", "ESXi主机", "Datastore", "虚拟机"], scope_rows)

        self._heading(document, "1.2 平台资源规模", 2)
        self._add_table(document, ["指标", "数值", "指标", "数值"], view.resource_scale)

        self._heading(document, "1.3 风险等级说明", 2)
        self._add_table(document, ["等级", "定义"], [[item.level, item.definition] for item in view.risk_definitions])

        self._heading(document, "1.4 整体巡检评价", 2)
        for index, statement in enumerate(view.evaluation_statements, start=1):
            self._numbered_paragraph(document, index, statement)

    def _render_result_analysis(self, document: Document, model: WordV2ViewModel) -> None:
        self._heading(document, "2. 巡检结果总体分析", 1)
        self._heading(document, "2.1 风险问题统计", 2)
        rows = [[row.serial_no, row.level, row.title, row.object_type, row.affected_object_count] for row in model.result_analysis.risk_rows]
        if rows:
            self._add_table(document, ["序号", "风险等级", "风险项", "对象类型", "问题对象数量"], rows)
        else:
            self._paragraph(document, "本次巡检未识别需要列入风险详情的异常。")
        if model.result_analysis.risk_distribution:
            self._heading(document, "2.2 风险分布", 2)
            rows = [[item.get("serial_no"), item.get("environment"), item.get("risk_count")] for item in model.result_analysis.risk_distribution]
            self._add_table(document, ["序号", "环境", "风险数量"], rows)

    def _render_risk_details(self, document: Document, risks: list[RiskItemView]) -> None:
        self._heading(document, "3. 风险项详细说明及优化建议", 1)
        if not risks:
            self._paragraph(document, "本次巡检未识别需要列入风险详情的异常。")
            return
        for risk in risks:
            self._heading(document, f"{risk.display_id} {risk.level} {risk.title}", 2)
            self._labeled_paragraph(document, "问题", risk.problem)
            self._labeled_paragraph(document, "建议", risk.recommendation)
            self._add_detail_table(risk.detail_table)

    def _render_vsan(self, document: Document, view: VsanView) -> None:
        self._heading(document, "4. vSAN 专项巡检", 1)
        self._heading(document, "4.1 vSAN 基本状态", 2)
        if not view.applicable:
            self._paragraph(document, view.conclusion)
            return
        self._add_table(document, ["项目", "结果"], view.basic_rows)
        self._heading(document, "4.2 磁盘组与物理磁盘", 2)
        disk_table = next((table for table in view.detail_tables if table.template_key == "vsan_disk_details"), None)
        if disk_table:
            self._add_detail_table(disk_table)
        else:
            self._paragraph(document, "本次未采集到物理磁盘明细，磁盘健康状态未确认。")
        self._heading(document, "4.3 对象健康", 2)
        self._paragraph(document, view.object_health_text)
        self._heading(document, "4.4 Resync", 2)
        self._paragraph(document, view.resync_text)
        self._heading(document, "4.5 容量", 2)
        self._add_table(document, ["项目", "结果"], view.capacity_rows)
        self._heading(document, "4.6 vSAN VMkernel 网络", 2)
        vmkernel_table = next((table for table in view.detail_tables if table.template_key == "vsan_vmk_network"), None)
        if vmkernel_table:
            self._add_detail_table(vmkernel_table)
        else:
            self._paragraph(document, "本次未采集到 vSAN VMkernel 网络明细，网络信息未确认。")

    def _render_conclusion(self, document: Document, view: ConclusionView) -> None:
        self._heading(document, "6. 巡检结论与建议", 1)
        self._heading(document, "6.1 vCenter / vSphere 环境", 2)
        self._paragraph(document, f"经本次巡检，当前 vCenter 及 vSphere 环境整体运行{view.vsphere_status}，具体检查结果如下：")
        if view.vsphere_items:
            for index, item in enumerate(view.vsphere_items, start=1):
                self._numbered_paragraph(document, index, item)
        else:
            self._numbered_paragraph(document, 1, "当前未发现需要优先处理的明显异常。")

        self._heading(document, "6.2 vSAN 环境", 2)
        self._paragraph(document, view.vsan_items[0] if view.vsan_items else "本次未形成 vSAN 专项结论。")

    def _add_detail_table(self, table: DetailTable) -> None:
        if not table.rows or not table.columns:
            return
        headers = ["序号", *[column.label for column in table.columns]]
        rows = []
        for index, row in enumerate(table.rows, start=1):
            rows.append([index, *[row.get(column.key, "") for column in table.columns]])
        self._add_table(self._document, headers, rows)

    def _add_table(self, document: Document, headers: list[str], rows: list[list[Any]], *, widths: list[int] | None = None, header: bool = True) -> Any:
        if not rows:
            return None
        table = document.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        normalized = self._normalize_widths(widths, len(headers))
        self._set_table_width(table, normalized)
        self._set_borders(table)
        header_row = table.rows[0]
        self._repeat_header(header_row)
        self._keep_row(header_row)
        for index, value in enumerate(headers):
            self._set_cell(header_row.cells[index], value, bold=True, center=True)
            self._set_cell_width(header_row.cells[index], normalized[index])
            if header:
                self._shade(header_row.cells[index], HEADER_BLUE)
                self._color_cell(header_row.cells[index], RGBColor(255, 255, 255))
        for row_index, values in enumerate(rows):
            row = table.add_row()
            self._keep_row(row)
            for index, value in enumerate(values):
                self._set_cell(row.cells[index], value, center=index == 0)
                self._set_cell_width(row.cells[index], normalized[index])
                if row_index % 2:
                    self._shade(row.cells[index], LIGHT_GRAY)
        document.add_paragraph().paragraph_format.space_after = Pt(1)
        return table

    def _heading(self, document: Document, text: str, level: int) -> None:
        paragraph = document.add_paragraph(text, style=f"Heading {level}")
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True

    def _paragraph(self, document: Document, text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        self._format_run(paragraph.add_run(self._safe_text(text)), 10.5, False, BLACK)

    def _labeled_paragraph(self, document: Document, label: str, text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        self._format_run(paragraph.add_run(f"{label}："), 10.5, True, BLACK)
        self._format_run(paragraph.add_run(self._safe_text(text).rstrip("。；;，,") + "。"), 10.5, False, BLACK)

    def _numbered_paragraph(self, document: Document, index: int, text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Mm(5)
        paragraph.paragraph_format.first_line_indent = Mm(-5)
        paragraph.paragraph_format.keep_together = True
        self._format_run(paragraph.add_run(f"{index}. {self._safe_text(text)}"), 10.5, False, BLACK)

    def _set_style(self, style: Any, size: float, bold: bool) -> None:
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = BLACK

    def _remove_style_borders(self, style: Any) -> None:
        p_pr = style._element.pPr
        if p_pr is None:
            return
        borders = p_pr.find(qn("w:pBdr"))
        if borders is not None:
            p_pr.remove(borders)

    def _format_run(self, run: Any, size: float, bold: bool, color: RGBColor) -> None:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(size)
        run.bold = bold
        run.font.color.rgb = color

    def _set_cell(self, cell: Any, value: Any, *, bold: bool = False, center: bool = False) -> None:
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        self._format_run(paragraph.add_run(self._safe_text(value)), 9.2 if not bold else 9.5, bold, BLACK)
        self._set_cell_margins(cell)

    def _set_cell_margins(self, cell: Any) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        margins = tc_pr.first_child_found_in("w:tcMar")
        if margins is None:
            margins = OxmlElement("w:tcMar")
            tc_pr.append(margins)
        for name in ("top", "start", "bottom", "end"):
            node = margins.find(qn(f"w:{name}"))
            if node is None:
                node = OxmlElement(f"w:{name}")
                margins.append(node)
            node.set(qn("w:w"), "80")
            node.set(qn("w:type"), "dxa")

    def _normalize_widths(self, widths: list[int] | None, count: int) -> list[int]:
        if widths and len(widths) == count:
            return widths
        base = PAGE_WIDTH_DXA // count
        result = [base for _ in range(count)]
        result[-1] += PAGE_WIDTH_DXA - sum(result)
        return result

    def _set_table_width(self, table: Any, widths: list[int]) -> None:
        tbl_pr = table._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl_pr.append(layout)
        tbl_w = OxmlElement("w:tblW")
        tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
        tbl_w.set(qn("w:type"), "dxa")
        tbl_pr.append(tbl_w)
        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for width in widths:
            column = OxmlElement("w:gridCol")
            column.set(qn("w:w"), str(width))
            grid.append(column)

    def _set_cell_width(self, cell: Any, width: int) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = tc_pr.first_child_found_in("w:tcW")
        if tc_w is None:
            tc_w = OxmlElement("w:tcW")
            tc_pr.append(tc_w)
        tc_w.set(qn("w:w"), str(width))
        tc_w.set(qn("w:type"), "dxa")

    def _set_borders(self, table: Any) -> None:
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), BORDER_GRAY)
            borders.append(element)
        table._tbl.tblPr.append(borders)

    def _shade(self, cell: Any, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        tc_pr.append(shading)

    def _color_cell(self, cell: Any, color: RGBColor) -> None:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = color

    def _repeat_header(self, row: Any) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)

    def _keep_row(self, row: Any) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)

    def _append_field(self, paragraph: Any, instruction: str) -> None:
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = instruction
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend((begin, instr, separate, end))

    def _safe_text(self, value: Any) -> str:
        if value is None:
            return "未采集"
        text = str(value).replace("\u200b", "")
        for forbidden in ("VSL-", "Rule ID", "INTERNAL_ONLY", "SUPPRESSED", "NEEDS_REVIEW", "Collector"):
            text = text.replace(forbidden, "")
        return " ".join(text.replace("\r", " ").replace("\n", " ").split()).strip() or "未采集"


__all__ = ["WordV2Renderer"]
