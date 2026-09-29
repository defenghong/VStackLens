from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips

from vstacklens.reports.report_model import FindingItem, ReportData


LEVEL_ORDER = {"P1": 0, "P2": 1, "P3": 2, "P4": 3}
LEVEL_STYLE = {
    "P1": {"fill": "C00000", "font": RGBColor(255, 255, 255), "text": RGBColor(192, 0, 0)},
    "P2": {"fill": "7030A0", "font": RGBColor(255, 255, 255), "text": RGBColor(112, 48, 160)},
    "P3": {"fill": "FFC000", "font": RGBColor(0, 0, 0), "text": RGBColor(191, 143, 0)},
    "P4": {"fill": "70AD47", "font": RGBColor(0, 0, 0), "text": RGBColor(84, 130, 53)},
}
TABLE_TOTAL_WIDTH_DXA = 9525
LEVEL_DEFINITION_WIDTHS_DXA = [1134, 8391]
OBJECT_OVERVIEW_WIDTHS_DXA = [1587, 1587, 1587, 1588, 1588, 1588]
SUMMARY_WIDTHS_DXA = [907, 1814, 6804]
RESOURCE_CHANGE_WIDTHS_DXA = [2268, 1134, 1134, 1134, 3855]
VCENTER_WIDTHS_DXA = [2199, 2200, 2563, 2563]
GENERIC_5_COL_WIDTHS_DXA = [1905, 1905, 1905, 1905, 1905]
ISSUE_SUMMARY_WIDTHS_DXA = [950, 1400, 2400, 950, 3825]
CLUSTER_FEATURE_WIDTHS_DXA = [4762, 4763]
NETWORK_LINK_WIDTHS_DXA = [1500, 900, 1400, 1150, 1850, 2725]
VMOTION_WIDTHS_DXA = [2382, 2381, 2381, 2381]
GENERIC_3_COL_WIDTHS_DXA = [3175, 3175, 3175]
FALLBACK_DETAIL_WIDTHS_DXA = [2500, 3300, 3725]
ISO_DETAIL_WIDTHS_DXA = [3175, 6350]
PORTGROUP_SECURITY_WIDTHS_DXA = [1531, 2268, 1531, 1531, 2664]
LOCAL_DATASTORE_VM_WIDTHS_DXA = [2500, 2200, 2200, 2625]
VSAN_DETAIL_WIDTHS_DXA = [2800, 1900, 2200, 2625]
CROSS_CLUSTER_DATASTORE_WIDTHS_DXA = [2600, 1600, 3725, 1600]
HOST_OVERCOMMIT_WIDTHS_DXA = [2600, 2300, 2300, 2325]
GUEST_OS_WIDTHS_DXA = [2500, 3512, 3513]
LONG_TEXT_HEADER_KEYWORDS = (
    "对象",
    "当前状态",
    "说明",
    "建议",
    "问题描述",
    "关注点",
    "备注",
    "告警内容",
    "UserAgent",
    "趋势判断",
    "ISO镜像",
    "所在位置",
    "vSwitch或端口组",
    "关联集群",
    "实际操作系统",
    "配置操作系统",
    "影响说明",
    "判断与建议",
)
LONG_TOKEN_PATTERN = re.compile(r"[A-Za-z0-9][A-Za-z0-9._:@-]{15,}")
WORD_BREAK = "\u200b"
TABLE_DATA_FONT_SIZE = 8.5
MAX_SUMMARY_ROWS = 20

RISK_CATEGORY_ORDER = ["vCenter", "集群", "主机", "网络", "存储", "虚拟机"]
TYPE_LABELS = {
    "vCenter": "vCenter",
    "ClusterComputeResource": "集群",
    "HostSystem": "主机",
    "Datastore": "存储",
    "VirtualMachine": "虚拟机",
}
FORBIDDEN_REPLACEMENTS = {
    "立即整改": "建议优先关注",
    "紧急整改": "建议优先关注",
    "必须整改": "建议评估处理",
    "严重威胁": "可能影响",
    "授权服务": "授权状态",
    "不可判定项": "数据缺口",
    "无不可判定项": "暂无数据缺口",
    "更新域": "刷新域",
    "更新目录": "刷新目录",
    "报告生成说明": "说明",
    "规则覆盖清单": "检查清单",
    "巡检范围": "检查对象",
    "风险统计": "检查结果汇总",
    "主要关注方向": "建议关注事项",
    "整改路线图": "后续处理建议",
}
CUSTOMER_WORD_REPLACEMENTS = {
    "VM 配置资源限制 Limit": "VM 配置资源限制",
    "虚拟机配置了资源 Limit": "虚拟机配置了资源限制",
    "CPU/Memory limit": "CPU/内存资源限制",
    "CPU/Memory Limit": "CPU/内存资源限制",
    "CPU/内存 Limit": "CPU/内存资源限制",
    "CPU 和内存 Limit": "CPU 和内存资源限制",
    "CPU 或内存 Limit": "CPU 或内存资源限制",
    "CPU Limit": "CPU 资源限制",
    "CPU limit": "CPU 资源限制",
    "Memory Limit": "内存资源限制",
    "memory limit": "内存资源限制",
    "期望值": "建议状态",
    "当前值": "当前状态",
    "期望：": "建议：",
    "失败": "未通过",
}
CUSTOMER_WORD_REGEX_REPLACEMENTS = (
    (re.compile(r"(?i)\bCPU\s*/\s*Memory\s+limits?\b"), "CPU/内存资源限制"),
    (re.compile(r"(?i)\bCPU\s+limits?\b"), "CPU 资源限制"),
    (re.compile(r"(?i)\bmemory\s+limits?\b"), "内存资源限制"),
    (re.compile(r"(?i)\bresource\s+limits?\b"), "资源限制"),
    (re.compile(r"(?i)\blimits?\b"), "资源限制"),
)
COMMON_MOJIBAKE_REPLACEMENTS = {
    "鏈寚瀹氬鎴?": "未指定客户",
    "鏈寚瀹氱珯鐐?": "未指定站点",
    "鏈噰闆?": "未采集",
    "VMware 铏氭嫙鍖栧仴搴疯瘎浼?": "VMware 虚拟化健康评估",
    "VStackLens VMware 铏氭嫙鍖栧仴搴疯瘎浼版姤鍛?": "vSphere 环境健康巡检报告",
}
TOOLS_CUSTOMER_RECOMMENDATION = (
    "安装 VMware Tools / open-vm-tools 后，vCenter 可以更准确获取客户机状态、IP、心跳和运行信息；"
    "支持优雅关机/重启；改善驱动、性能和运维管理能力；安装或修复后需复查 Tools 运行状态"
)


@dataclass(slots=True)
class RiskGroup:
    level: str
    category: str
    title: str
    findings: list[FindingItem]
    rule_id: str = ""


class DocxReportEngine:
    """Render the customer-facing Word V2 report."""

    def render(self, report: ReportData, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        self._setup_document(document)
        risk_groups = self._aggregate_risks(report)
        customer_risk_groups = self._customer_visible_risk_groups(risk_groups)

        self._cover(document, report)
        self._table_of_contents(document)
        self._executive_summary(document, report, customer_risk_groups)
        self._asset_overview(document, report)
        self._risk_recommendations(document, report, customer_risk_groups)
        vsan_status = (report.vsan_summary or {}).get("status")
        vsan_visible = bool(vsan_status and vsan_status != "not_applicable")
        if vsan_visible:
            self._vsan_section(document, report)
        self._follow_up(document, customer_risk_groups, section_number=5 if vsan_visible else 4)
        self._appendix(document, report, section_number=6 if vsan_visible else 5)

        document.save(output_path)
        return output_path

    def _customer_visible_risk_groups(self, risk_groups: list[RiskGroup]) -> list[RiskGroup]:
        base_groups = [group for group in risk_groups if group.level in {"P1", "P2", "P3"}]
        syslog_hosts = {
            self._object_key(finding.object_name)
            for group in base_groups
            if self._is_remote_syslog_group(group)
            for finding in group.findings
        }
        visible_groups: list[RiskGroup] = []

        for group in base_groups:
            findings = [finding for finding in group.findings if not self._is_default_vm_finding(finding)]
            if not findings:
                continue

            candidate = RiskGroup(
                level=self._risk_level_for_findings(findings, group.level),
                category=group.category,
                title=self._clean(group.title),
                findings=findings,
                rule_id=group.rule_id,
            )
            if self._is_tools_group(candidate) or self._is_syslog_group(candidate):
                continue

            if self._is_log_core_dump_group(candidate) and syslog_hosts:
                findings = [finding for finding in findings if self._object_key(finding.object_name) not in syslog_hosts]
                if not findings:
                    continue
                candidate = RiskGroup(
                    level=self._risk_level_for_findings(findings, candidate.level),
                    category=candidate.category,
                    title=candidate.title,
                    findings=findings,
                    rule_id=candidate.rule_id,
                )
            visible_groups.append(candidate)

        return sorted(visible_groups, key=self._risk_group_sort_key)

    def _risk_level_for_findings(self, findings: list[FindingItem], default: str = "P3") -> str:
        levels = [finding.risk_level.upper() for finding in findings if finding.risk_level.upper() in LEVEL_ORDER]
        return min(levels, key=lambda item: LEVEL_ORDER[item]) if levels else default

    def _object_key(self, value: Any) -> str:
        return self._clean(value).casefold()

    def _is_default_vm_finding(self, finding: FindingItem) -> bool:
        if finding.object_type != "VirtualMachine":
            return False
        return self._is_default_vm_name(finding.object_name)

    def _is_default_vm_name(self, value: Any) -> bool:
        text = self._clean(value)
        if not text:
            return False
        candidates = [text, *re.split(r"[\\/]", text)]
        for candidate in candidates:
            name = candidate.strip().casefold()
            if name.startswith("vcls") or name.startswith("vmware vcls"):
                return True
        return False

    def _is_remote_syslog_group(self, group: RiskGroup) -> bool:
        text = self._group_text(group)
        lowered = text.casefold()
        return "syslog" in lowered and not self._is_log_core_dump_text(text)

    def _is_syslog_group(self, group: RiskGroup) -> bool:
        lowered = self._group_text(group).casefold()
        return "syslog" in lowered or "远程日志" in lowered

    def _is_log_core_dump_group(self, group: RiskGroup) -> bool:
        return self._is_log_core_dump_text(self._group_text(group))

    def _is_log_core_dump_text(self, text: str) -> bool:
        lowered = text.casefold()
        return any(keyword in lowered for keyword in ("core dump", "coredump")) or any(
            keyword in text for keyword in ("核心转储", "转储配置不完整", "日志与核心转储")
        )

    def _is_tools_group(self, group: RiskGroup) -> bool:
        return any(self._is_tools_finding(finding) for finding in group.findings)

    def _group_text(self, group: RiskGroup) -> str:
        parts: list[str] = [group.title]
        for finding in group.findings:
            parts.extend(
                getattr(finding, name, "")
                for name in (
                    "rule_id",
                    "rule_name",
                    "title",
                    "object_name",
                    "current_value",
                    "expected_value",
                    "current_value_zh",
                    "expected_value_zh",
                    "business_impact",
                    "technical_impact",
                    "consequence",
                    "evidence_summary_zh",
                    "observed_detail_zh",
                    "expected_detail_zh",
                    "remediation",
                    "recommended_action_zh",
                    "fault_detail_zh",
                    "explanation",
                    "explanation_zh",
                    "affected_components_zh",
                )
            )
            parts.extend(finding.remediation_steps or [])
        return " ".join(self._clean(part) for part in parts)

    def _merged_tools_group(self, findings: list[FindingItem]) -> RiskGroup | None:
        if not findings:
            return None
        by_vm: dict[str, FindingItem] = {}
        for finding in findings:
            key = self._object_key(finding.object_name)
            existing = by_vm.get(key)
            if existing is None or self._tools_priority(finding) < self._tools_priority(existing):
                by_vm[key] = finding
        merged = sorted(by_vm.values(), key=lambda finding: self._clean(finding.object_name))
        has_uninstalled = any(self._tools_priority(finding) == 0 for finding in merged)
        title = "虚拟机未安装 VMware Tools" if has_uninstalled else "虚拟机 VMware Tools 需安装或修复"
        return RiskGroup(
            level=self._risk_level_for_findings(merged, "P3"),
            category="虚拟机",
            title=title,
            findings=merged,
            rule_id="VMWARE-TOOLS",
        )

    def _tools_priority(self, finding: FindingItem) -> int:
        text = self._clean(
            " ".join(
                str(value)
                for value in (
                    finding.rule_id,
                    finding.rule_name,
                    finding.title,
                    finding.current_value,
                    finding.evidence_summary_zh,
                    finding.remediation,
                    finding.recommended_action_zh,
                )
                if value is not None
            )
        ).casefold()
        if any(keyword in text for keyword in ("未安装", "not installed", "install state")):
            return 0
        if any(keyword in text for keyword in ("未运行", "not running", "notrunning")):
            return 1
        if any(keyword in text for keyword in ("版本过旧", "版本较旧", "outdated", "upgrade")):
            return 2
        return 3

    def _setup_document(self, document: Document) -> None:
        section = document.sections[0]
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        styles = document.styles
        normal = styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        normal.font.size = Pt(10.5)
        for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
            style = styles[name]
            style.font.name = "Microsoft YaHei"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            style.font.color.rgb = RGBColor(31, 78, 121)
        for name in ("TOC 1", "TOC 2", "TOC 3"):
            if name in styles:
                style = styles[name]
                style.font.name = "Microsoft YaHei"
                style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        settings = document.settings.element
        if settings.find(qn("w:updateFields")) is None:
            update_fields = OxmlElement("w:updateFields")
            update_fields.set(qn("w:val"), "true")
            settings.append(update_fields)

    def _cover(self, document: Document, report: ReportData) -> None:
        for _ in range(4):
            document.add_paragraph()
        customer = self._clean(report.customer_info.customer_name or "未指定客户")
        vcenter = self._clean(report.customer_info.vcenter or "未指定 vCenter")
        inspection_date = self._date_label(report.report_info.generated_at)

        customer_paragraph = document.add_paragraph()
        customer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        customer_run = customer_paragraph.add_run(customer)
        self._format_run(customer_run, size=18, bold=True, color=RGBColor(31, 78, 121))

        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_before = Pt(18)
        title_run = title.add_run("vSphere 环境健康巡检报告")
        self._format_run(title_run, size=28, bold=True, color=RGBColor(31, 78, 121))

        date = document.add_paragraph()
        date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date.paragraph_format.space_before = Pt(24)
        date_run = date.add_run(f"巡检日期：{inspection_date}")
        self._format_run(date_run, size=13, bold=False, color=RGBColor(80, 80, 80))

        target = document.add_paragraph()
        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        target.paragraph_format.space_before = Pt(8)
        target_run = target.add_run(f"检查对象：vCenter {vcenter}")
        self._format_run(target_run, size=12, bold=False, color=RGBColor(80, 80, 80))

        for _ in range(4):
            document.add_paragraph()

        summary_title = document.add_paragraph()
        summary_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        summary_title_run = summary_title.add_run("摘要")
        self._format_run(summary_title_run, size=14, bold=True, color=RGBColor(31, 78, 121))

        summary = document.add_paragraph()
        summary.paragraph_format.first_line_indent = Cm(0.75)
        summary.paragraph_format.line_spacing = 1.25
        summary_run = summary.add_run(f"基于当前巡检数据，对 vCenter 环境 {vcenter} 的平台配置、运行状态、容量、可用性和建议关注事项进行汇总。")
        self._format_run(summary_run, size=12)
        document.add_page_break()

    def _table_of_contents(self, document: Document) -> None:
        title = document.add_paragraph()
        title.style = "Title"
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title_run = title.add_run("目录")
        self._format_run(title_run, size=16, bold=True, color=RGBColor(31, 78, 121))

        toc = document.add_paragraph()
        if "TOC 1" in document.styles:
            toc.style = "TOC 1"
        tabs = toc.paragraph_format.tab_stops
        tabs.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = toc.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = 'TOC \\o "1-3" \\h \\z \\u'
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_sep)
        placeholder = toc.add_run("打开文档后可在 Word 中刷新此目录。")
        self._format_run(placeholder, size=10.5, color=RGBColor(90, 90, 90))
        toc.add_run()._r.append(fld_end)
        document.add_page_break()

    def _executive_summary(self, document: Document, report: ReportData, risk_groups: list[RiskGroup]) -> None:
        self._add_heading(document, "1. 执行摘要", 1)
        self._add_paragraph(document, f"环境健康状态：{report.health_score.label}。{report.health_score.explanation}")
        self._add_heading(document, "1.1 巡检报告概述", 2)
        self._add_paragraph(
            document,
            (
                "本报告记录 vSphere 虚拟化平台当前健康状况。检查内容覆盖平台配置、运行状态、容量使用、"
                "可用性、安全配置、证书与授权等方面，目的是帮助客户了解当前环境状态和建议关注事项，"
                "并为后续运维优化提供参考。"
            ),
        )

        self._add_heading(document, "1.2 评估等级说明", 2)
        self._add_level_definition_table(
            document,
            [
                ["P1", "建议优先关注，通常涉及可用性、安全、容量或关键配置风险"],
                ["P2", "建议纳入近期优化计划"],
                ["P3", "建议结合日常运维持续优化"],
            ],
        )

        self._add_heading(document, "1.3 检查对象概览", 2)
        summary = self._asset_summary(report)
        self._add_table(
            document,
            ["vCenter", "数据中心", "集群", "ESXi 主机", "共享/业务 Datastore", "虚拟机"],
            [
                [
                    summary.get("vCenter", 0),
                    len(self._unique_datacenters(report)),
                    summary.get("ClusterComputeResource", 0),
                    summary.get("HostSystem", 0),
                    summary.get("Datastore", 0),
                    summary.get("VirtualMachine", 0),
                ]
            ],
            widths=OBJECT_OVERVIEW_WIDTHS_DXA,
            align_center_columns={0, 1, 2, 3, 4, 5},
        )

        self._add_heading(document, "1.4 重点问题", 2)
        counts = self._risk_counts(report, risk_groups)
        detail_counts = self._finding_counts(risk_groups)
        total = sum(counts.values())
        detail_total = sum(detail_counts.values())
        self._add_paragraph(
            document,
            (
                f"本次 vSphere 环境健康检查共识别 {total} 类建议关注问题，涉及 {detail_total} 条对象明细。"
                f"按问题类型统计：P1 {counts['P1']} 类、P2 {counts['P2']} 类、P3 {counts['P3']} 类。"
                f"按影响对象明细统计：P1 {detail_counts['P1']} 条、P2 {detail_counts['P2']} 条、P3 {detail_counts['P3']} 条。"
                "以下按问题类型汇总展示，Word 正文仅保留代表性对象明细，完整对象清单建议在 HTML 报告中查看。"
            ),
        )
        rows = self._summary_rows(risk_groups)
        self._add_heading(document, "建议关注问题分类", 3)
        self._add_issue_summary_table(document, rows[:MAX_SUMMARY_ROWS])
        if len(rows) > MAX_SUMMARY_ROWS:
            self._add_paragraph(document, "其余建议关注事项请参见后续风险与建议章节。")

        self._add_heading(document, "1.5 巡检结论", 2)
        self._resource_change_summary(document, report)

    def _asset_overview(self, document: Document, report: ReportData) -> None:
        self._add_heading(document, "3. VMware 巡检", 1)
        self._add_heading(document, "3.1 vCenter 与集群", 2)
        vcenter_rows = []
        for item in self._asset_details(report, "vCenter"):
            props = item.get("properties") or {}
            vcenter_rows.append(
                [
                    item.get("object_name", ""),
                    self._fmt(props.get("version")),
                    self._fmt(props.get("build")),
                    "已连接" if props.get("connected") is not False else "未连接",
                ]
            )
        if vcenter_rows:
            self._add_table(document, ["名称", "版本", "Build", "连接状态"], vcenter_rows, widths=VCENTER_WIDTHS_DXA)
        clusters = self._asset_details(report, "ClusterComputeResource")
        if clusters:
            self._add_table(
                document,
                ["集群", "主机数", "HA", "DRS", "说明"],
                [
                    [
                        item.get("object_name", ""),
                        self._fmt((item.get("properties") or {}).get("host_count")),
                        self._enabled_label((item.get("properties") or {}).get("ha_enabled")),
                        self._enabled_label((item.get("properties") or {}).get("drs_enabled")),
                        self._cluster_attention(item.get("properties") or {}),
                    ]
                    for item in clusters
                ],
                widths=GENERIC_5_COL_WIDTHS_DXA,
            )
        if not vcenter_rows and not clusters:
            self._add_paragraph(document, "本次未采集到可展示的 vCenter 与集群数据。")

        hosts = self._asset_details(report, "HostSystem")
        self._add_heading(document, "3.2 ESXi 主机", 2)
        if hosts:
            rows = []
            for item in hosts:
                props = item.get("properties") or {}
                rows.append(
                    [
                        item.get("object_name", ""),
                        self._cluster_from_location(item.get("asset_location") or props.get("asset_location")),
                        self._host_state_label(props.get("connection_state")),
                        self._fmt_percent(props.get("cpu_usage_avg")),
                        self._fmt_percent(props.get("memory_usage_avg")),
                    ]
                )
            self._add_table(
                document,
                ["主机", "所属集群", "连接状态", "CPU 使用率", "内存使用率"],
                rows,
                widths=GENERIC_5_COL_WIDTHS_DXA,
            )
        else:
            self._add_paragraph(document, "本次未采集到 ESXi 主机资源数据。")

        stores = self._asset_details(report, "Datastore")
        self._add_heading(document, "3.3 存储使用情况", 2)
        if stores:
            rows = []
            for item in sorted(stores, key=lambda row: self._safe_float_sort_key((row.get("properties") or {}).get("used_percent")), reverse=True):
                props = item.get("properties") or {}
                rows.append(
                    [
                        item.get("object_name", ""),
                        self._fmt(props.get("datastore_filesystem_type")),
                        self._fmt_percent(props.get("used_percent")),
                        self._fmt_percent(props.get("free_percent")),
                        "可访问" if props.get("accessible") is not False else "不可访问",
                    ]
                )
            self._add_table(
                document,
                ["Datastore", "类型", "已用", "剩余", "状态"],
                rows,
                widths=GENERIC_5_COL_WIDTHS_DXA,
            )
        else:
            self._add_paragraph(document, "本次未采集到 Datastore 容量数据。")

        self._add_heading(document, "3.4 版本、证书与授权状态", 2)
        version_rows = self._version_certificate_license_rows(report)
        self._add_table(
            document,
            ["对象", "版本/产品", "证书状态", "授权状态", "说明"],
            version_rows or [["vCenter", "未采集", "未采集", "未采集", "建议定期复核"]],
            widths=GENERIC_5_COL_WIDTHS_DXA,
        )

    def _risk_recommendations(self, document: Document, report: ReportData, risk_groups: list[RiskGroup]) -> None:
        self._add_heading(document, "2. 问题与整改建议", 1)

        by_level: dict[str, list[RiskGroup]] = defaultdict(list)
        for group in risk_groups:
            by_level[group.level].append(group)

        for level_index, level in enumerate(("P1", "P2", "P3"), start=1):
            section_number = f"2.{level_index}"
            self._add_heading(document, f"{section_number} {level}", 2)
            level_groups: dict[str, list[RiskGroup]] = defaultdict(list)
            for group in by_level.get(level, []):
                level_groups[group.category].append(group)
            if not level_groups:
                self._add_paragraph(document, f"本次未识别 {level} 等级需要列入报告的建议关注事项。")
                continue
            category_index = 1
            for category in RISK_CATEGORY_ORDER:
                items = level_groups.get(category)
                if not items:
                    continue
                category_number = f"{section_number}.{category_index}"
                self._add_heading(document, f"{category_number} {category}", 3)
                for item_index, group in enumerate(sorted(items, key=self._risk_group_sort_key), start=1):
                    self._add_heading(document, f"{category_number}.{item_index} {group.title}", 4)
                    self._add_risk_sentence(document, group)
                    table = self._detail_table(report, group)
                    if table:
                        headers, rows = table
                        self._add_table(document, headers, rows, widths=self._detail_widths(headers))
                category_index += 1

    def _follow_up(self, document: Document, risk_groups: list[RiskGroup], *, section_number: int = 5) -> None:
        self._add_heading(document, f"{section_number}. 后续处理建议", 1)
        priority = [group for group in risk_groups if group.level in {"P1", "P2"}]
        planned = [group for group in risk_groups if group.level == "P3"]

        self._add_heading(document, f"{section_number}.1 建议优先确认", 2)
        self._add_bullets(
            document,
            self._suggestions(priority, "建议优先确认", "结合影响范围和变更窗口安排后续动作。"),
            empty_text="当前未识别 P1/P2 等级建议关注事项，建议保持定期复核。",
        )

        self._add_heading(document, f"{section_number}.2 建议计划优化", 2)
        self._add_bullets(
            document,
            self._suggestions(planned, "建议计划优化", "纳入近期运维优化计划。"),
            empty_text="当前未识别 P3 等级建议关注事项，建议结合日常运维持续观察。",
        )

    def _vsan_section(self, document: Document, report: ReportData) -> None:
        summary = report.vsan_summary or {}
        if summary.get("status") == "not_applicable":
            return
        self._add_heading(document, "4. vSAN 专项巡检", 1)
        status = str(summary.get("status") or "unknown")
        status_label = {"collected": "已采集", "unavailable": "部分采集", "unknown": "未确认", "not_applicable": "不适用"}.get(status, status)
        self._add_paragraph(document, f"vSAN 采集状态：{status_label}；架构：{summary.get('architecture') or '未采集'}；集群：{', '.join(summary.get('clusters') or []) or '未采集'}。")
        self._add_paragraph(document, "判定说明：采集成功且异常计数为 0，表示本次采集时点未发现该项异常；存在异常时说明影响和建议；未采集或未确认的数据不按正常处理。")
        self._add_paragraph(document, self._vsan_overall_judgement(summary))

        self._add_heading(document, "4.1 集群健康", 2)
        self._add_paragraph(document, self._vsan_category_line(summary, "VSAN-CLUSTER-HEALTH"))
        self._add_paragraph(document, self._vsan_health_judgement(summary))
        self._add_table(document, ["项目", "结果"], [["总体 Health", status_label], ["健康异常", self._fmt(summary.get("health_issue_count"))], ["对象异常", self._fmt(summary.get("object_issue_count"))]], widths=CLUSTER_FEATURE_WIDTHS_DXA)

        self._add_heading(document, "4.2 磁盘与磁盘组", 2)
        self._add_paragraph(document, self._vsan_category_line(summary, "VSAN-DISK"))
        self._add_table(document, ["主机数", "磁盘组数", "缓存盘数", "容量盘数", "状态"], [[self._fmt(summary.get("host_count")), self._fmt(summary.get("disk_group_count")), self._fmt(summary.get("cache_disk_count")), self._fmt(summary.get("capacity_disk_count")), status_label]], widths=GENERIC_5_COL_WIDTHS_DXA)
        self._add_paragraph(document, self._vsan_disk_judgement(summary))
        self._add_vsan_disk_details(document, summary)

        self._add_heading(document, "4.3 对象健康", 2)
        self._add_paragraph(document, self._vsan_category_line(summary, "VSAN-OBJECT-HEALTH"))
        self._add_paragraph(document, self._vsan_object_judgement(summary))
        self._add_vsan_object_chart(document, summary)
        if summary.get("object_issues"):
            self._add_table(document, ["对象/组件", "Health", "说明"], [[item.get("component", "未采集"), item.get("status", "未知"), item.get("summary", "")] for item in summary["object_issues"]], widths=GENERIC_3_COL_WIDTHS_DXA)

        self._add_heading(document, "4.4 Resync", 2)
        self._add_paragraph(document, self._vsan_category_line(summary, "VSAN-RESYNC"))
        self._resync_status_table(document, summary)
        self._add_paragraph(document, self._vsan_resync_judgement(summary))

        self._add_heading(document, "4.5 vSAN VMkernel 网络", 2)
        self._add_paragraph(document, self._vsan_category_line(summary, "VSAN-NETWORK"))
        network = summary.get("network") or {}
        self._add_paragraph(document, self._vsan_network_judgement(summary))
        self._add_vsan_vmk_details(document, summary)

        self._add_heading(document, "4.6 Capacity", 2)
        self._add_paragraph(document, self._vsan_category_line(summary, "VSAN-CAPACITY"))
        capacity = summary.get("capacity") or {}
        self._add_paragraph(document, self._vsan_capacity_judgement(summary))
        self._add_vsan_capacity_chart(document, summary)

        if summary.get("storage_policy_summary"):
            self._add_heading(document, "4.7 Storage Policy", 2)
            self._add_paragraph(document, self._vsan_category_line(summary, "VSAN-POLICY"))
            self._add_paragraph(document, f"Storage Policy 不合规对象：{self._fmt(summary.get('policy_noncompliant_count'))}。未采集时显示 Unknown，不将其判定为通过。")

    def _vsan_category_line(self, summary: dict[str, Any], category_id: str) -> str:
        item = next((item for item in summary.get("report_categories") or [] if item.get("category_id") == category_id), None)
        if not item:
            return "检查等级：未确认；本次结论：未确认。"
        return f"检查等级：{item.get('risk_level') or '未确认'}；本次结论：{item.get('status') or '未确认'}。"

    def _vsan_overall_judgement(self, summary: dict[str, Any]) -> str:
        issue_values = [summary.get("health_issue_count"), summary.get("disk_issue_count"), summary.get("object_issue_count")]
        issue_count = sum(int(value or 0) for value in issue_values if value is not None)
        resync_object_count = summary.get("resync_object_count")
        resync_bytes = summary.get("resync_bytes")
        disk_rows = self._vsan_disk_rows(summary)
        capacity = summary.get("capacity") or {}
        network = summary.get("network") or {}
        vmkernels = network.get("vmkernels") or []
        if issue_count or any(row[4] == "异常" for row in disk_rows) or capacity.get("status") in {"attention", "high"}:
            return "判断：本次 vSAN 检查发现需要关注的健康、磁盘、对象或容量状态，具体影响和建议见下方检查项。"
        if (
            summary.get("status") != "collected"
            or any(value is None for value in issue_values)
            or summary.get("object_count") is None
            or summary.get("vmdk_count") is None
            or resync_object_count is None
            or resync_bytes is None
            or capacity.get("status") == "unknown"
            or capacity.get("status") is None
            or not disk_rows
            or any(row[4] == "未确认" for row in disk_rows)
            or not vmkernels
            or self._vsan_missing_vmk_labels(vmkernels)
        ):
            return "判断：本次 vSAN 部分检查数据未确认，已采集结果见下文；当前暂不判定整体状态正常。"
        if resync_object_count == 0 and resync_bytes == 0:
            return "判断：本次采集未发现 vSAN 集群健康、物理磁盘、对象健康或容量水位异常；当前无待重同步数据，整体状态正常。"
        return "判断：本次采集未发现已确认的健康异常，但当前存在重同步活动，建议持续观察对象可用性和完成进度。"

    def _vsan_health_judgement(self, summary: dict[str, Any]) -> str:
        count = summary.get("health_issue_count")
        if count is None:
            return "判断：集群健康结果未确认，无法据此判断集群整体正常。"
        if count == 0:
            return "判断：本次 vSAN 集群健康检查未发现异常项，当前采集时点总体 Health 正常。"
        return f"判断：发现 {count} 项 vSAN 集群健康异常，可能影响数据服务或故障恢复能力，建议结合异常明细和 vCenter 事件进一步确认。"

    def _vsan_disk_judgement(self, summary: dict[str, Any]) -> str:
        if not summary.get("disk_topology"):
            return "判断：本次未采集到磁盘组拓扑，物理磁盘状态无法确认。"
        rows = self._vsan_disk_rows(summary)
        issue_count = summary.get("disk_issue_count")
        if summary.get("physical_disks") is None or any(row[4] == "未确认" for row in rows):
            return "判断：磁盘组和物理磁盘拓扑已采集，但物理磁盘健康字段未完整返回，盘级健康状态标记为未确认。"
        if issue_count is None:
            return "判断：物理磁盘状态已列出，但磁盘健康检查结果未完整返回，盘级结论未确认。"
        if issue_count == 0:
            return f"判断：已检查 {len(rows)} 块物理磁盘，未发现磁盘健康异常；详细盘位和状态见下表。"
        return f"判断：发现 {issue_count} 项物理磁盘健康异常，可能影响磁盘组冗余和数据可用性，建议优先核对异常设备和所在主机。"

    def _vsan_object_judgement(self, summary: dict[str, Any]) -> str:
        object_count = summary.get("object_count")
        vmdk_count = summary.get("vmdk_count")
        issue_count = summary.get("object_issue_count")
        if object_count is None or vmdk_count is None or issue_count is None:
            return "判断：对象或 VMDK 健康数据未完整采集，当前不能将其判定为正常。"
        if issue_count == 0:
            return f"判断：本次采集识别 {object_count} 个 vSAN 对象和 {vmdk_count} 个 VMDK，未发现对象不可访问或异常健康状态。VMDK 图示为已识别数量。"
        return f"判断：发现 {issue_count} 个对象健康异常，可能影响虚拟机磁盘访问或数据冗余，建议结合对象明细和受影响虚拟机进行处理。"

    def _vsan_resync_judgement(self, summary: dict[str, Any]) -> str:
        objects = summary.get("resync_object_count")
        bytes_left = summary.get("resync_bytes")
        if objects is None or bytes_left is None:
            return "判断：Resync 状态未完整采集，当前不对重同步活动作正常判断。"
        if objects == 0 and bytes_left == 0:
            return "判断：当前没有待重同步对象和待同步数据，预计完成时间与计划重同步按空闲状态显示为 0。"
        return f"判断：当前有 {self._fmt(objects)} 个对象或 {self._resync_bytes_label(bytes_left)} 待重同步，可能影响恢复窗口和存储性能，建议持续观察完成进度。"

    def _vsan_network_judgement(self, summary: dict[str, Any]) -> str:
        network = summary.get("network") or {}
        rows = [item for item in network.get("vmkernels") or [] if isinstance(item, dict)]
        if not rows:
            return "判断：未采集到 vSAN VMkernel 明细，网络状态无法确认。"
        missing = self._vsan_missing_vmk_labels(rows)
        if missing:
            return f"判断：已采集 {len(rows)} 条 vSAN VMkernel 明细，但其中 {missing} 条缺少 IP 或 Network Label，相关字段标记为未记录。"
        return f"判断：已采集 {len(rows)} 条 vSAN VMkernel 明细，主机、IP、子网、Network Label 和 MTU 均已返回。"

    def _vsan_missing_vmk_labels(self, rows: list[dict[str, Any]]) -> int:
        missing_values = {"", "未记录", "未解析"}
        return sum(
            1
            for item in rows
            if not item.get("ip_address")
            or str(item.get("network_label") or item.get("portgroup") or "").strip() in missing_values
        )

    def _vsan_capacity_judgement(self, summary: dict[str, Any]) -> str:
        capacity = summary.get("capacity") or {}
        used = capacity.get("used_percent")
        status = capacity.get("status")
        if used is None:
            return "判断：vSAN 容量使用率未采集，当前不能判断容量水位。"
        if status == "normal":
            return f"判断：当前容量使用率为 {self._fmt_percent(used)}，低于 75% 关注阈值，整体容量水位正常。"
        if status == "attention":
            return f"判断：当前容量使用率为 {self._fmt_percent(used)}，已达到关注区间，建议纳入容量扩展和对象增长观察。"
        return f"判断：当前容量使用率为 {self._fmt_percent(used)}，已达到高容量水位，建议优先评估扩容和数据增长趋势。"

    def _add_vsan_object_chart(self, document: Document, summary: dict[str, Any]) -> None:
        object_count = summary.get("object_count")
        vmdk_count = summary.get("vmdk_count")
        issue_count = summary.get("object_issue_count")
        rows = []
        if object_count is None or issue_count is None:
            rows.append(("vSAN 对象", [("未确认", 1, "9EADBC")], "对象 未采集"))
        else:
            normal_count = max(int(object_count) - int(issue_count or 0), 0)
            segments = []
            if normal_count:
                segments.append(("正常", normal_count, "4EA72E"))
            if issue_count:
                segments.append(("异常", int(issue_count), "C00000"))
            rows.append(("vSAN 对象", segments or [("未确认", 1, "9EADBC")], f"对象 {int(object_count)}；异常 {int(issue_count or 0)}"))
        if vmdk_count is None:
            rows.append(("VMDK", [("未确认", 1, "9EADBC")], "VMDK 未采集"))
        else:
            rows.append(("VMDK", [("已识别", int(vmdk_count), "4EA72E")], f"VMDK {int(vmdk_count)}"))
        self._add_vsan_bar_chart(document, "vSAN 对象与 VMDK 健康", rows, ["正常：未发现对象异常", "已识别：VMDK 数量", "未确认：数据不足"])

    def _add_vsan_capacity_chart(self, document: Document, summary: dict[str, Any]) -> None:
        capacity = summary.get("capacity") or {}
        total = capacity.get("total_gb")
        used = capacity.get("used_gb")
        free = capacity.get("free_gb")
        if total is None or used is None or free is None:
            used_percent = capacity.get("used_percent")
            if used_percent is None:
                rows = [("容量", [("未确认", 1, "9EADBC")], "容量 未采集")]
            else:
                rows = [("容量", [("已用", float(used_percent), "2F75B5"), ("可用", max(100 - float(used_percent), 0), "9EADBC")], f"使用率 {self._fmt_percent(used_percent)}")]
        else:
            rows = [("容量", [("已用", float(used), "2F75B5"), ("可用", float(free), "9EADBC")], f"总容量 {float(total):,.1f} GB；已用 {float(used):,.1f} GB；可用 {float(free):,.1f} GB；使用率 {self._fmt_percent(capacity.get('used_percent'))}")]
        self._add_vsan_bar_chart(document, "vSAN 容量使用", rows, ["已用", "可用"])

    def _add_vsan_bar_chart(self, document: Document, title: str, rows: list[tuple[str, list[tuple[str, float, str]], str]], legend: list[str]) -> None:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        self._format_run(paragraph.add_run(title), size=11.5, color=RGBColor(31, 78, 121))
        for label, segments, total_label in rows:
            bar_width = TABLE_TOTAL_WIDTH_DXA - 1800
            row_total = sum(max(float(value), 0.0) for _, value, _ in segments) or 1.0
            segment_widths = [max(1, int(round(bar_width * max(float(value), 0.0) / row_total))) for _, value, _ in segments]
            segment_widths[max(range(len(segment_widths)), key=segment_widths.__getitem__)] += bar_width - sum(segment_widths)
            widths = [1800, *segment_widths]
            table = document.add_table(rows=1, cols=len(widths))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            self._set_table_fixed_width(table, widths)
            self._keep_row_together(table.rows[0])
            self._set_cell_text(table.rows[0].cells[0], label, align_center=True)
            self._set_cell_width(table.rows[0].cells[0], widths[0])
            for index, ((name, value, color), width) in enumerate(zip(segments, segment_widths), start=1):
                cell = table.rows[0].cells[index]
                self._set_cell_text(cell, f"{name} {self._fmt(value)}" if width >= 1100 else "", align_center=True)
                self._set_cell_width(cell, width)
                self._shade_cell(cell, color)
                self._color_cell_text(cell, RGBColor(255, 255, 255))
            self._add_paragraph(document, total_label)
        self._add_paragraph(document, "图示说明：" + "；".join(legend) + "。")

    def _add_vsan_disk_details(self, document: Document, summary: dict[str, Any]) -> None:
        rows = self._vsan_disk_rows(summary)
        if not rows:
            self._add_paragraph(document, "本次未采集到磁盘组和物理磁盘明细，盘级状态未确认。")
            return
        table = self._add_table(document, ["主机", "磁盘组", "磁盘角色", "设备名称", "健康状态"], rows, widths=[1900, 1450, 1450, 3150, 1575])
        self._merge_table_runs(table, rows, (0, 1))
        if summary.get("disk_issues"):
            self._add_table(document, ["主机", "设备", "状态", "说明"], [[item.get("host", "未记录"), item.get("disk", "未记录"), item.get("status", "异常"), item.get("summary", "需要核查")] for item in summary["disk_issues"]], widths=[1900, 3150, 1450, 3025])

    def _vsan_disk_rows(self, summary: dict[str, Any]) -> list[list[str]]:
        rows: list[list[str]] = []
        physical_index: dict[tuple[str, str], dict[str, Any]] = {}
        for disk in summary.get("physical_disks") or []:
            host = str(disk.get("host") or "")
            for identity in (disk.get("name"), disk.get("scsi_device"), disk.get("uuid")):
                if identity:
                    physical_index[(host, str(identity).casefold())] = disk
        for host_item in sorted(summary.get("disk_topology") or [], key=lambda item: str(item.get("host") or "")):
            host = str(host_item.get("host") or "未记录")
            for group_index, group in enumerate(host_item.get("disk_groups") or [], start=1):
                group_name = f"磁盘组 {group_index}"
                devices = [("缓存盘", group.get("cache_disk")), *(('容量盘', item) for item in group.get("capacity_disks") or [])]
                for role, device in devices:
                    if not device:
                        continue
                    disk = physical_index.get((host, str(device).casefold()))
                    values = [str((disk or {}).get(key) or "").casefold() for key in ("summary_health", "operational_health", "capacity_health")]
                    bad = {"red", "yellow", "error", "failed", "unhealthy", "offline", "absent", "lost", "degraded"}
                    health = "异常" if any(value in bad for value in values) else "正常" if disk and any(value in {"green", "healthy", "ok", "normal"} for value in values) else "未确认"
                    rows.append([host, group_name, role, str(device), health])
        return rows

    def _add_vsan_vmk_details(self, document: Document, summary: dict[str, Any]) -> None:
        source_rows = [item for item in (summary.get("network") or {}).get("vmkernels") or [] if isinstance(item, dict)]
        rows = [[
            item.get("host_name") or item.get("host") or "未记录",
            item.get("device") or "未记录",
            item.get("ip_address") or item.get("ip") or "未记录",
            item.get("subnet_mask") or item.get("subnet") or "未记录",
            item.get("network_label") or item.get("portgroup") or "未记录",
            item.get("mtu") if item.get("mtu") is not None else "未记录",
        ] for item in source_rows]
        if not rows:
            self._add_paragraph(document, "本次未采集到 vSAN VMkernel 网络明细。")
            return
        rows.sort(key=lambda row: (str(row[0]), str(row[1])))
        table = self._add_table(document, ["主机", "设备", "IP", "子网", "Network Label", "MTU"], rows, widths=[1650, 900, 1800, 1600, 2700, 875])
        self._merge_table_runs(table, rows, (0,))

    def _merge_table_runs(self, table: Any, rows: list[list[Any]], columns: tuple[int, ...]) -> None:
        for column in columns:
            start = 0
            while start < len(rows):
                end = start
                while end + 1 < len(rows) and rows[end + 1][column] == rows[start][column] and (column == 0 or rows[end + 1][0] == rows[start][0]):
                    end += 1
                if end > start:
                    merged = table.cell(start + 1, column).merge(table.cell(end + 1, column))
                    self._set_cell_text(merged, rows[start][column], align_center=True)
                start = end + 1

    def _resync_status_table(self, document: Document, summary: dict[str, Any]) -> None:
        object_count = self._safe_float(summary.get("resync_object_count"))
        remaining_bytes = self._safe_float(summary.get("resync_bytes"))
        idle = object_count == 0 and remaining_bytes == 0
        has_pending = (object_count is not None and object_count > 0) or (remaining_bytes is not None and remaining_bytes > 0)
        status = "当前无待同步任务" if idle else "当前有待同步任务" if has_pending else "状态未确认"
        rows = [
            ("正在重同步对象", "未采集" if object_count is None else f"{self._fmt(summary.get('resync_object_count'))} 个", idle),
            ("待同步数据", self._resync_bytes_label(summary.get("resync_bytes")), idle),
            ("预计完成时间", "0 秒" if idle else "未采集", False),
            ("计划重同步", "0 项" if idle else "未采集", False),
        ]
        table = self._new_table(document, ["Resync 状态", status], widths=[6500, 3025], repeat_header=False)
        for index, cell in enumerate(table.rows[0].cells):
            self._shade_cell(cell, "F3F6F9")
            for run in cell.paragraphs[0].runs:
                run.font.color.rgb = RGBColor(40, 115, 72) if idle and index == 1 else RGBColor(38, 55, 70)
        for label, value, collected_zero in rows:
            row = table.add_row()
            self._keep_row_together(row)
            self._set_cell_text(row.cells[0], label)
            self._set_cell_text(row.cells[1], value)
            row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for cell, width in zip(row.cells, table._vstacklens_widths_dxa):
                self._set_cell_width(cell, width)
            if collected_zero:
                for run in row.cells[1].paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(40, 115, 72)
        document.add_paragraph()

    def _resync_bytes_label(self, value: Any) -> str:
        number = self._safe_float(value)
        if number is None:
            return "未采集"
        if number == 0:
            return "0 B"
        for unit, divisor in (("TB", 1024 ** 4), ("GB", 1024 ** 3), ("MB", 1024 ** 2), ("KB", 1024)):
            if number >= divisor:
                return f"{number / divisor:.1f} {unit}"
        return f"{number:g} B"

    def _appendix(self, document: Document, report: ReportData, *, section_number: int = 6) -> None:
        self._add_heading(document, f"{section_number}. 环境资产与对象清单", 1)
        self._add_heading(document, f"{section_number}.1 资产清单", 2)
        has_any = False
        vcenters = self._asset_details(report, "vCenter")
        if vcenters:
            has_any = True
            self._add_paragraph(document, "vCenter")
            self._add_table(
                document,
                ["名称", "版本", "Build", "连接状态"],
                [
                    [
                        item.get("object_name", ""),
                        self._fmt((item.get("properties") or {}).get("version")),
                        self._fmt((item.get("properties") or {}).get("build")),
                        "已连接" if (item.get("properties") or {}).get("connected") is not False else "未连接",
                    ]
                    for item in vcenters
                ],
                widths=VCENTER_WIDTHS_DXA,
            )

        clusters = self._asset_details(report, "ClusterComputeResource")
        if clusters:
            has_any = True
            self._add_paragraph(document, "集群")
            self._add_table(
                document,
                ["集群", "主机数", "HA", "DRS", "说明"],
                [
                    [
                        item.get("object_name", ""),
                        self._fmt((item.get("properties") or {}).get("host_count")),
                        self._enabled_label((item.get("properties") or {}).get("ha_enabled")),
                        self._enabled_label((item.get("properties") or {}).get("drs_enabled")),
                        self._cluster_attention(item.get("properties") or {}),
                    ]
                    for item in clusters
                ],
                widths=GENERIC_5_COL_WIDTHS_DXA,
            )

        hosts = self._asset_details(report, "HostSystem")
        if hosts:
            has_any = True
            self._add_paragraph(document, "ESXi 主机")
            self._add_table(
                document,
                ["主机", "所属集群", "连接状态", "CPU 使用率", "内存使用率"],
                [
                    [
                        item.get("object_name", ""),
                        self._cluster_from_location(item.get("asset_location") or (item.get("properties") or {}).get("asset_location")),
                        self._host_state_label((item.get("properties") or {}).get("connection_state")),
                        self._fmt_percent((item.get("properties") or {}).get("cpu_usage_avg")),
                        self._fmt_percent((item.get("properties") or {}).get("memory_usage_avg")),
                    ]
                    for item in hosts
                ],
                widths=GENERIC_5_COL_WIDTHS_DXA,
            )

        stores = self._asset_details(report, "Datastore")
        if stores:
            has_any = True
            self._add_paragraph(document, "Datastore")
            self._add_table(
                document,
                ["Datastore", "类型", "已用", "剩余", "状态"],
                [
                    [
                        item.get("object_name", ""),
                        self._fmt((item.get("properties") or {}).get("datastore_filesystem_type")),
                        self._fmt_percent((item.get("properties") or {}).get("used_percent")),
                        self._fmt_percent((item.get("properties") or {}).get("free_percent")),
                        "可访问" if (item.get("properties") or {}).get("accessible") is not False else "不可访问",
                    ]
                    for item in stores
                ],
                widths=GENERIC_5_COL_WIDTHS_DXA,
            )

        vms = self._asset_details(report, "VirtualMachine")
        if vms:
            has_any = True
            self._add_paragraph(document, "虚拟机")
            self._add_table(
                document,
                ["虚拟机", "所在位置", "电源状态", "vCPU"],
                [
                    [
                        item.get("object_name", ""),
                        self._short_location(item.get("asset_location") or (item.get("properties") or {}).get("asset_location")),
                        self._power_state_label((item.get("properties") or {}).get("power_state")),
                        self._fmt((item.get("properties") or {}).get("vcpu_count")),
                    ]
                    for item in vms
                ],
                widths=[2600, 3400, 1900, 1625],
            )

        if not has_any:
            self._add_paragraph(document, "本次未采集到可展示的资产清单数据。")

    def _aggregate_risks(self, report: ReportData) -> list[RiskGroup]:
        grouped: dict[tuple[str, str], list[FindingItem]] = defaultdict(list)
        for finding in report.findings:
            level = (finding.risk_level or "").upper()
            if level not in LEVEL_ORDER:
                continue
            category = self._classify_finding(finding)
            title = self._risk_title(finding)
            grouped[(finding.rule_id, level)].append(finding)

        result: list[RiskGroup] = []
        for (rule_id, level), findings in grouped.items():
            category = self._classify_finding(findings[0])
            title = self._risk_title(findings[0])
            level = min((finding.risk_level.upper() for finding in findings if finding.risk_level.upper() in LEVEL_ORDER), key=lambda item: LEVEL_ORDER[item])
            result.append(RiskGroup(level=level, category=category, title=title, findings=findings, rule_id=rule_id))
        return sorted(result, key=self._risk_group_sort_key)

    def _risk_group_sort_key(self, group: RiskGroup) -> tuple[int, int, str]:
        category_order = RISK_CATEGORY_ORDER.index(group.category) if group.category in RISK_CATEGORY_ORDER else len(RISK_CATEGORY_ORDER)
        return (LEVEL_ORDER.get(group.level, 99), category_order, group.title)

    def _risk_title(self, finding: FindingItem) -> str:
        return self._clean(finding.title or finding.rule_name or finding.rule_id or "建议关注事项")

    def _classify_finding(self, finding: FindingItem) -> str:
        title_rule_text = " ".join(self._clean(value) for value in (finding.title, finding.rule_name, finding.object_name))
        text = " ".join(
            self._clean(value)
            for value in (finding.title, finding.rule_name, finding.object_type, finding.object_name, finding.business_impact, finding.technical_impact)
        )
        lowered = text.lower()
        title_rule_lowered = title_rule_text.lower()
        object_type = finding.object_type
        network_hit = any(keyword.lower() in title_rule_lowered for keyword in ("网络", "网卡", "vmotion", "vmkernel", "链路", "端口组", "vswitch", "uplink", "mtu", "vmnic", "pnic"))
        log_audit_hit = any(keyword.lower() in title_rule_lowered for keyword in ("syslog", "日志", "审计", "ntp", "core dump", "coredump"))
        customer_confirm_hit = any(keyword.lower() in title_rule_lowered for keyword in ("orphaned", "inaccessible", "孤立", "不可访问", "需确认", "客户确认"))
        if object_type == "vCenter":
            return "vCenter"
        if object_type == "ClusterComputeResource":
            return "网络" if network_hit else "集群"
        if object_type == "HostSystem":
            return "网络" if network_hit else "主机"
        if object_type == "Datastore":
            return "存储"
        if object_type == "VirtualMachine":
            return "虚拟机"
        if customer_confirm_hit or log_audit_hit:
            return "主机"
        if network_hit:
            return "网络"
        if any(keyword in text for keyword in ("虚拟机", "VMware Tools", "快照", "ISO")):
            return "虚拟机"
        if any(keyword in text for keyword in ("ESXi", "主机", "SSH", "Shell", "Lockdown", "电源策略")):
            return "主机"
        if any(keyword in text for keyword in ("Datastore", "存储", "容量", "路径")):
            return "存储"
        return "集群"

    def _summary_rows(self, risk_groups: list[RiskGroup]) -> list[list[str]]:
        if not risk_groups:
            return [["-", "-", "本次未识别需要列入风险详情的建议关注事项。", "0", "保持定期复核"]]
        rows: list[list[str]] = []
        for group in sorted(risk_groups, key=self._risk_group_sort_key):
            recommendation = self._recommendation(group.findings[0]) if group.findings else "结合实际业务影响评估处理"
            rows.append([group.level, group.category, group.title, str(len(group.findings)), recommendation])
        return rows

    def _risk_counts(self, report: ReportData, risk_groups: list[RiskGroup]) -> dict[str, int]:
        counts = {level: 0 for level in ("P1", "P2", "P3")}
        for group in risk_groups:
            if group.level in counts:
                counts[group.level] += 1
        return counts

    def _finding_counts(self, risk_groups: list[RiskGroup]) -> dict[str, int]:
        counts = {level: 0 for level in ("P1", "P2", "P3")}
        for group in risk_groups:
            if group.level in counts:
                counts[group.level] += len(group.findings)
        return counts

    def _resource_change_summary(self, document: Document, report: ReportData) -> None:
        summary = self._asset_summary(report)
        rows = [
            ["vCenter", "未提供", self._fmt(summary.get("vCenter", 0)), "本次基线", "建议后续结合历史巡检形成趋势"],
            ["集群", "未提供", self._fmt(summary.get("ClusterComputeResource", 0)), "本次基线", "建议关注集群容量与可用性配置变化"],
            ["ESXi 主机", "未提供", self._fmt(summary.get("HostSystem", 0)), "本次基线", "建议结合新增或下线主机变更记录确认"],
            ["Datastore", "未提供", self._fmt(summary.get("Datastore", 0)), "本次基线", "建议持续观察容量增长和共享存储使用情况"],
            ["虚拟机", "未提供", self._fmt(summary.get("VirtualMachine", 0)), "本次基线", "建议结合业务上线和退役计划确认变化"],
        ]
        self._add_table(
            document,
            ["指标", "上一轮", "本次", "变化", "趋势判断"],
            rows,
            widths=RESOURCE_CHANGE_WIDTHS_DXA,
            align_center_columns={1, 2, 3},
        )

    def _detail_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]] | None:
        title = group.title
        rule_text = " ".join(self._clean(f"{item.rule_id} {item.rule_name} {item.title}") for item in group.findings)
        lowered = rule_text.lower()
        if "VSL-CL-001" in rule_text or "集群 HA 未启用" in title or "cluster ha disabled" in lowered:
            return self._cluster_feature_table(report, group, "ha_enabled", "HA 当前状态")
        if "VSL-CL-003" in rule_text or "集群 DRS 未启用" in title or "cluster drs disabled" in lowered:
            return self._cluster_feature_table(report, group, "drs_enabled", "DRS 当前状态")
        if "告警" in title and group.category == "vCenter":
            return self._vcenter_alarm_table(group)
        if "会话" in title:
            return self._idle_session_table(group)
        if ("ssh" in lowered or "shell" in lowered) and group.category == "主机":
            return self._host_ssh_table(report, group)
        if "syslog" in lowered:
            return self._host_syslog_table(report, group)
        if "VSL-NET-015" in rule_text:
            return self._portgroup_security_table(report, group)
        if "VSL-HOST-024" in rule_text:
            return self._hardware_health_table(report, group)
        if "VSL-VM-023" in rule_text:
            return self._local_datastore_vm_table(report, group)
        if "VSL-VM-024" in rule_text:
            return self._guest_os_table(report, group)
        if "VSL-DS-019" in rule_text:
            return self._cross_cluster_datastore_table(group)
        if "VSL-HOST-026" in rule_text:
            return self._host_overcommit_table(group)
        if "电源策略" in title:
            return self._host_power_table(report, group)
        if "vmotion" in lowered:
            return self._vmotion_table(report, group)
        if "链路" in title or "物理网卡" in title:
            return self._network_link_table(report, group)
        if "快照" in title:
            return self._snapshot_table(report, group)
        if "ISO" in title.upper():
            return self._iso_table(group)
        if "Tools" in title or "VMware Tools" in title:
            return self._tools_table(report, group)
        if "VSL-DS-020" in rule_text or "vsan" in lowered:
            return self._vsan_detail_table(group)
        return self._fallback_detail_table(group)

    def _cluster_feature_table(self, report: ReportData, group: RiskGroup, prop_name: str, status_header: str) -> tuple[list[str], list[list[Any]]]:
        rows = []
        for finding in group.findings:
            props = self._asset_props(report, "ClusterComputeResource", finding.object_name)
            status = props.get(prop_name) if prop_name in props else finding.current_value
            rows.append([finding.object_name, self._enabled_label(status)])
        return ["集群", status_header], rows

    def _vcenter_alarm_table(self, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        for finding in group.findings:
            observed = self._observed(finding)
            alarms = observed.get("active_red_alarms") or observed.get("alarms") or []
            for alarm in alarms:
                if isinstance(alarm, dict):
                    rows.append(
                        [
                            alarm.get("entity_name") or finding.object_name,
                            alarm.get("alarm_name") or alarm.get("name") or group.title,
                            alarm.get("acknowledged") if alarm.get("acknowledged") is not None else "未确认",
                            self._time_label(alarm.get("time") or alarm.get("created_time") or finding.collected_at),
                            alarm.get("status") or self._current_display(finding),
                        ]
                    )
            if not alarms:
                rows.append([finding.object_name, group.title, "未确认", self._time_label(finding.collected_at), self._current_display(finding)])
        return ["对象", "告警内容", "确认状态", "时间", "状态"], rows

    def _idle_session_table(self, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        for finding in group.findings:
            observed = self._observed(finding)
            sessions = observed.get("idle_sessions") or observed.get("sessions") or []
            for session in sessions:
                if isinstance(session, dict):
                    rows.append(
                        [
                            session.get("username") or finding.object_name,
                            session.get("ip_address") or session.get("ip") or "未采集",
                            session.get("user_agent") or session.get("userAgent") or "未采集",
                            self._time_label(session.get("login_time") or session.get("created_time")),
                            self._time_label(session.get("last_active_time") or session.get("last_activity")),
                        ]
                    )
            if not sessions:
                rows.append([finding.object_name, "未采集", "未采集", self._time_label(finding.collected_at), "未采集"])
        return ["用户名", "IP地址", "UserAgent", "登录时间", "上次活动时间"], rows

    def _host_ssh_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows = []
        for finding in group.findings:
            props = self._asset_props(report, "HostSystem", finding.object_name)
            observed = self._observed(finding)
            running = props.get("ssh_running")
            if running is None:
                running = props.get("shell_running")
            if running is None:
                running = observed.get("service_running") if "service_running" in observed else observed.get("running")
            if running is None:
                running = finding.current_value
            rows.append(
                [
                    finding.object_name,
                    self._host_cluster(report, finding.object_name),
                    "已启用" if self._truthy(running) else self._current_display(finding),
                ]
            )
        return ["ESXi 主机", "所属集群", "当前状态"], rows

    def _host_syslog_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows = []
        for finding in group.findings:
            props = self._asset_props(report, "HostSystem", finding.object_name)
            configured = props.get("syslog_configured")
            status = "已配置" if configured is True else "未配置"
            rows.append([finding.object_name, self._host_cluster(report, finding.object_name), status])
        return ["ESXi 主机", "所属集群", "远程 Syslog 状态"], rows

    def _host_power_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows = []
        for finding in group.findings:
            props = self._asset_props(report, "HostSystem", finding.object_name)
            observed = self._observed(finding)
            policy = props.get("host_power_policy") or props.get("power_policy") or observed.get("host_power_policy") or finding.current_value
            rows.append([finding.object_name, self._host_cluster(report, finding.object_name), self._fmt(policy)])
        return ["ESXi 主机", "所属集群", "电源策略"], rows

    def _portgroup_security_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        for finding in group.findings:
            observed = self._observed(finding)
            props = self._asset_props(report, "HostSystem", finding.object_name)
            issues = observed.get("portgroup_security_issues") or props.get("portgroup_security_issues") or []
            added = False
            if issues:
                for issue in issues:
                    if not isinstance(issue, dict):
                        continue
                    switch = self._fmt(issue.get("switch"))
                    portgroup = self._fmt(issue.get("portgroup"))
                    target = switch if switch == portgroup else f"{switch} / {portgroup}"
                    rows.append(
                        [
                            observed.get("host_name") or finding.object_name,
                            target,
                            issue.get("policy") or "未采集",
                            issue.get("current_value") or "已启用",
                            f"建议调整为{issue.get('recommended_value') or '禁用'}",
                        ]
                    )
                    added = True
            if not added:
                rows.append(
                    [
                        finding.object_name,
                        "未采集",
                        "安全策略",
                        self._fallback_status_value(finding),
                        self._recommendation(finding),
                    ]
                )
        return ["ESXi 主机", "vSwitch或端口组", "策略项", "当前状态", "建议"], rows

    def _hardware_health_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        for finding in group.findings:
            observed = self._observed(finding)
            props = self._asset_props(report, "HostSystem", finding.object_name)
            issues = observed.get("host_hardware_health_issues") or props.get("host_hardware_health_issues") or []
            if issues:
                for issue in issues:
                    if isinstance(issue, dict):
                        rows.append(
                            [
                                observed.get("host_name") or finding.object_name,
                                issue.get("component") or "未采集",
                                self._health_status_label(issue.get("status")),
                            ]
                        )
            else:
                rows.append([finding.object_name, "未采集", self._current_display(finding)])
        return ["ESXi 主机", "组件", "健康状态"], rows

    def _local_datastore_vm_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        for finding in group.findings:
            observed = self._observed(finding)
            datastores = observed.get("local_datastore_names") or observed.get("datastore_names") or []
            rows.append(
                [
                    finding.object_name,
                    self._join_display(datastores),
                    self._vm_location(report, finding),
                    "位于本地 Datastore，主机维护或故障时迁移与恢复能力可能受限",
                ]
            )
        return ["虚拟机", "Datastore", "所属主机或集群", "影响说明"], rows

    def _guest_os_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        for finding in group.findings:
            observed = self._observed(finding)
            props = self._asset_props(report, "VirtualMachine", finding.object_name)
            rows.append(
                [
                    finding.object_name,
                    observed.get("guest_os_actual") or props.get("guest_os_actual") or "未采集",
                    observed.get("guest_os_configured") or props.get("guest_os_configured") or "未采集",
                ]
            )
        return ["虚拟机", "实际操作系统", "配置操作系统"], rows

    def _cross_cluster_datastore_table(self, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        for finding in group.findings:
            observed = self._observed(finding)
            rows.append(
                [
                    finding.object_name,
                    self._fmt(observed.get("datastore_cluster_count")),
                    self._join_display(observed.get("datastore_cluster_names") or []),
                    self._fmt(observed.get("datastore_host_count")),
                ]
            )
        return ["Datastore", "关联集群数", "关联集群", "主机数"], rows

    def _host_overcommit_table(self, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        for finding in group.findings:
            observed = self._observed(finding)
            rows.append(
                [
                    finding.object_name,
                    self._ratio_label(observed.get("host_vcpu_to_pcpu_ratio")),
                    self._ratio_percent_label(observed.get("host_memory_allocation_ratio")),
                    self._fallback_status_value(finding),
                ]
            )
        return ["ESXi 主机", "vCPU:pCPU", "内存分配比例", "当前状态"], rows

    def _vmotion_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        for finding in group.findings:
            observed = self._observed(finding)
            missing = observed.get("missing_vmotion_hosts") or []
            if missing:
                for host in missing:
                    rows.append([host, self._host_cluster(report, str(host)), "未配置可用 vMotion VMkernel", self._recommendation(finding)])
            else:
                rows.append([finding.object_name, self._cluster_from_location(finding.object_path), self._current_display(finding), self._recommendation(finding)])
        return ["ESXi 主机", "所属集群", "当前状态", "建议"], rows

    def _network_link_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        issue_labels = {
            "link_down": "已绑定上行，但链路断开",
            "speed_zero": "实际速率为 0",
            "speed_below_1gbps": "实际速率低于 1Gbps",
            "speed_mismatch": "固定配置速率与实际速率不一致",
        }
        for finding in group.findings:
            observed = self._observed(finding)
            host_name = observed.get("host_name") or finding.object_name
            host_props = self._asset_props(report, "HostSystem", str(host_name))
            aggregate_switches = observed.get("affected_switches") or host_props.get("affected_switches") or []
            physical_nics = observed.get("physical_nic_details")
            details = physical_nics if isinstance(physical_nics, list) else observed.get("link_speed_detail") or []
            if details:
                reported = 0
                for item in details:
                    if isinstance(item, dict):
                        issue_codes = item.get("issue_codes") or []
                        switches = item.get("assigned_switches") or []
                        if switches:
                            binding = "、".join(str(value) for value in switches if value)
                        elif item.get("is_uplink") is False:
                            binding = "未绑定"
                        elif aggregate_switches:
                            aggregate = "、".join(str(value) for value in aggregate_switches if value)
                            binding = f"{aggregate}（主机级映射，未逐卡保存）"
                        else:
                            binding = "未记录逐卡绑定"
                        link_state = item.get("link_state") or item.get("status")
                        if link_state == "down":
                            link_label = "down"
                            actual_label = "无链路"
                        elif link_state == "up":
                            link_label = "up"
                            actual_label = self._speed_label(item.get("actual_speed_mb", item.get("speed_mb")))
                        else:
                            link_label = "未采集"
                            actual_label = self._speed_label(item.get("actual_speed_mb", item.get("speed_mb")))
                        if item.get("autonegotiation") is True:
                            configured_label = "自动协商"
                        elif item.get("configured_speed_mb") is not None:
                            configured_label = self._speed_label(item.get("configured_speed_mb"))
                            configured_duplex = item.get("configured_duplex")
                            if configured_duplex is not None:
                                configured_label += " 全双工" if configured_duplex else " 半双工"
                        elif item.get("autonegotiation") is False:
                            configured_label = "配置速率未采集"
                        else:
                            configured_label = "旧版证据未记录"
                        reasons = [issue_labels[code] for code in issue_codes if code in issue_labels]
                        if not reasons and item.get("status") == "degraded":
                            reasons.append("物理网卡链路降速")
                        if not reasons and item.get("status") == "down":
                            reasons.append("旧版采集记录显示链路 down")
                        is_uplink = item.get("is_uplink")
                        if isinstance(physical_nics, list) and is_uplink is False:
                            judgment = "未绑定到 vSwitch/vDS/opaque switch；该网卡不按上行链路判故障。"
                        elif reasons:
                            judgment = "；".join(reasons)
                            recommendation = self._recommendation(finding)
                            if recommendation:
                                judgment += f"；{recommendation}"
                        elif isinstance(physical_nics, list):
                            judgment = "已采集；链路正常，不计为异常。"
                        else:
                            judgment = f"采集异常明细。{self._recommendation(finding)}"
                        rows.append(
                            [
                                observed.get("host_name") or finding.object_name,
                                item.get("device") or "未采集",
                                binding,
                                link_label,
                                f"实际 {actual_label} / 配置 {configured_label}",
                                judgment,
                            ]
                        )
                        if issue_codes:
                            reported += 1
                if isinstance(physical_nics, list) and not reported:
                    rows.append(
                        [
                            observed.get("host_name") or finding.object_name,
                            "逐卡异常明细为空",
                            "已采集到网卡清单",
                            "规则结果与明细不一致",
                            "未形成链路异常证据",
                            "请复核采集结果和规则输入；不将缺少的明细伪装成具体故障。",
                        ]
                    )
            else:
                affected = observed.get("affected_nics") or finding.affected_components or []
                if len(aggregate_switches) == 1:
                    binding = f"{aggregate_switches[0]}（主机级映射，未逐卡保存）"
                elif aggregate_switches:
                    binding = f"{'、'.join(str(value) for value in aggregate_switches)}（主机级映射，未逐卡保存）"
                else:
                    binding = "旧版证据未记录"
                for nic in affected or ["未采集"]:
                    rows.append(
                        [
                            finding.object_name,
                            nic,
                            binding,
                            self._current_display(finding),
                            "旧版证据未记录",
                            self._recommendation(finding),
                        ]
                    )
        return ["ESXi 主机", "网卡", "绑定交换机", "链路状态", "实际/配置速率", "判断与建议"], rows

    def _snapshot_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows = []
        for finding in group.findings:
            observed = self._observed(finding)
            snapshots = observed.get("snapshots") or observed.get("snapshot_detail") or []
            if not snapshots:
                asset_props = self._asset_props(report, "VirtualMachine", finding.object_name)
                snapshots = asset_props.get("snapshots") or asset_props.get("snapshot_detail") or []
            if snapshots:
                for snapshot in snapshots:
                    if isinstance(snapshot, dict):
                        rows.append(
                            [
                                finding.object_name,
                                self._snapshot_name_label(snapshot),
                                self._snapshot_time_display(snapshot),
                            ]
                        )
            else:
                rows.append([finding.object_name, "虚拟机存在快照", "未采集"])
        return ["所在虚拟机", "快照名称", "创建时间/存在时长"], rows

    def _snapshot_time_display(self, snapshot: dict[str, Any]) -> str:
        created = snapshot.get("create_time") or snapshot.get("created_at") or snapshot.get("creation_time")
        if created:
            return str(created)
        age = snapshot.get("age_days") or snapshot.get("snapshot_age_days")
        return f"约 {age} 天" if age not in (None, "") else "未采集"

    def _iso_table(self, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows = []
        for finding in group.findings:
            observed = self._observed(finding)
            iso = observed.get("iso_path") or observed.get("mounted_iso") or finding.current_value
            rows.append([finding.object_name, self._fmt(iso)])
        return ["虚拟机名称", "ISO镜像"], rows

    def _tools_table(self, report: ReportData, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows = []
        for finding in group.findings:
            props = self._asset_props(report, "VirtualMachine", finding.object_name)
            finding_status = self._tools_status_from_finding(finding)
            rows.append(
                [
                    finding.object_name,
                    finding_status if self._tools_priority(finding) < 3 else self._tools_label(props) if props else finding_status,
                    self._power_state_label(props.get("power_state")) if props else "未采集",
                ]
            )
        return ["虚拟机名称", "VMware Tools 状态", "运行状态"], rows

    def _vsan_detail_table(self, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        rows: list[list[Any]] = []
        for finding in group.findings:
            observed = self._observed(finding)
            used_percent = observed.get("vsan_used_percent") if observed else finding.current_value
            rows.append(
                [
                    observed.get("datastore_name") or finding.object_name,
                    self._fmt_percent(used_percent),
                    self._gb_size_label(observed.get("vsan_free_gb")),
                    "容量水位偏高" if self._is_high_percent(used_percent, 80) else self._fallback_status_value(finding),
                ]
            )
        return ["Datastore", "已用率", "剩余容量", "状态"], rows

    def _fallback_detail_table(self, group: RiskGroup) -> tuple[list[str], list[list[Any]]]:
        return ["对象", "当前状态", "建议"], [
            [finding.object_name, self._fallback_status_value(finding), self._recommendation(finding)] for finding in group.findings
        ]

    def _detail_widths(self, headers: list[str]) -> list[int]:
        presets = {
            ("对象", "告警内容", "确认状态", "时间", "状态"): GENERIC_5_COL_WIDTHS_DXA,
            ("用户名", "IP地址", "UserAgent", "登录时间", "上次活动时间"): GENERIC_5_COL_WIDTHS_DXA,
            ("集群", "HA 当前状态"): CLUSTER_FEATURE_WIDTHS_DXA,
            ("集群", "DRS 当前状态"): CLUSTER_FEATURE_WIDTHS_DXA,
            ("ESXi 主机", "所属集群", "远程 Syslog 状态"): GENERIC_3_COL_WIDTHS_DXA,
            ("ESXi 主机", "所属集群", "电源策略"): GENERIC_3_COL_WIDTHS_DXA,
            ("ESXi 主机", "所属集群", "当前状态"): GENERIC_3_COL_WIDTHS_DXA,
            ("ESXi 主机", "所属集群", "当前状态", "建议"): VMOTION_WIDTHS_DXA,
            ("ESXi 主机", "网卡", "绑定交换机", "链路状态", "实际/配置速率", "判断与建议"): NETWORK_LINK_WIDTHS_DXA,
            ("所在虚拟机", "快照名称", "创建时间/存在时长"): GENERIC_3_COL_WIDTHS_DXA,
            ("虚拟机名称", "ISO镜像"): ISO_DETAIL_WIDTHS_DXA,
            ("虚拟机名称", "VMware Tools 状态", "运行状态"): GENERIC_3_COL_WIDTHS_DXA,
            ("ESXi 主机", "vSwitch或端口组", "策略项", "当前状态", "建议"): PORTGROUP_SECURITY_WIDTHS_DXA,
            ("虚拟机", "Datastore", "所属主机或集群", "影响说明"): LOCAL_DATASTORE_VM_WIDTHS_DXA,
            ("ESXi 主机", "组件", "健康状态"): GENERIC_3_COL_WIDTHS_DXA,
            ("Datastore", "已用率", "剩余容量", "状态"): VSAN_DETAIL_WIDTHS_DXA,
            ("虚拟机", "实际操作系统", "配置操作系统"): GUEST_OS_WIDTHS_DXA,
            ("Datastore", "关联集群数", "关联集群", "主机数"): CROSS_CLUSTER_DATASTORE_WIDTHS_DXA,
            ("ESXi 主机", "vCPU:pCPU", "内存分配比例", "当前状态"): HOST_OVERCOMMIT_WIDTHS_DXA,
            ("对象", "当前状态", "建议"): FALLBACK_DETAIL_WIDTHS_DXA,
        }
        return presets.get(tuple(headers), self._semantic_widths_dxa(headers))

    def _add_issue_summary_table(self, document: Document, rows: list[list[str]]) -> None:
        headers = ["等级", "分类", "问题描述", "影响数量", "建议动作"]
        table = self._new_table(document, headers, widths=ISSUE_SUMMARY_WIDTHS_DXA, repeat_header=True)
        widths = getattr(table, "_vstacklens_widths_dxa", ISSUE_SUMMARY_WIDTHS_DXA)
        for row in rows:
            table_row = table.add_row()
            self._keep_row_together(table_row)
            cells = table_row.cells
            for idx, header in enumerate(headers):
                value = row[idx] if idx < len(row) else ""
                self._set_cell_text(cells[idx], value, align_center=not self._is_long_text_column(header))
                self._set_cell_width(cells[idx], widths[idx])
            level = row[0] if row else ""
            if level in LEVEL_STYLE:
                self._shade_cell(cells[0], LEVEL_STYLE[level]["fill"])
                self._color_cell_text(cells[0], LEVEL_STYLE[level]["font"], bold=True)
        document.add_paragraph()

    def _add_level_definition_table(self, document: Document, rows: list[list[str]]) -> None:
        table = self._new_table(document, ["等级", "说明"], widths=LEVEL_DEFINITION_WIDTHS_DXA, repeat_header=False)
        widths = getattr(table, "_vstacklens_widths_dxa", LEVEL_DEFINITION_WIDTHS_DXA)
        for level, description in rows:
            table_row = table.add_row()
            self._keep_row_together(table_row)
            cells = table_row.cells
            self._set_cell_text(cells[0], level, align_center=True)
            self._set_cell_text(cells[1], description)
            self._set_cell_width(cells[0], widths[0])
            self._set_cell_width(cells[1], widths[1])
            style = LEVEL_STYLE.get(level, LEVEL_STYLE["P4"])
            self._shade_cell(cells[0], style["fill"])
            self._color_cell_text(cells[0], style["font"], bold=True)
        document.add_paragraph()

    def _add_table(
        self,
        document: Document,
        headers: list[str],
        rows: list[list[Any]],
        widths: list[float] | None = None,
        align_center_columns: set[int] | None = None,
    ) -> Any:
        if not rows:
            return None
        align_center_columns = align_center_columns or set()
        table = self._new_table(document, headers, widths=widths, repeat_header=True)
        normalized_widths = getattr(table, "_vstacklens_widths_dxa", self._semantic_widths_dxa(headers))
        for row in rows:
            table_row = table.add_row()
            self._keep_row_together(table_row)
            cells = table_row.cells
            for idx, header in enumerate(headers):
                value = row[idx] if idx < len(row) else ""
                align_center = idx in align_center_columns
                self._set_cell_text(cells[idx], value, align_center=align_center)
                self._set_cell_width(cells[idx], normalized_widths[idx])
        document.add_paragraph()
        return table

    def _new_table(self, document: Document, headers: list[str], widths: list[float] | None, repeat_header: bool) -> Any:
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        normalized_widths = self._normalize_widths_dxa(widths, headers)
        table._vstacklens_widths_dxa = normalized_widths
        self._set_table_autofit_width(table, normalized_widths)
        self._keep_row_together(table.rows[0])
        if repeat_header:
            self._repeat_header(table.rows[0])
        for idx, header in enumerate(headers):
            cell = table.rows[0].cells[idx]
            self._set_cell_text(cell, header, bold=True, align_center=True)
            self._shade_cell(cell, "D9EAF7")
            self._set_cell_width(cell, normalized_widths[idx])
        return table

    def _set_table_fixed_width(self, table: Any, widths: list[int]) -> None:
        self._set_table_width(table, widths, layout="fixed")

    def _set_table_autofit_width(self, table: Any, widths: list[int]) -> None:
        self._set_table_width(table, widths, layout="autofit")

    def _set_table_width(self, table: Any, widths: list[int], *, layout: str) -> None:
        tbl_pr = table._tbl.tblPr
        if tbl_pr is None:
            tbl_pr = OxmlElement("w:tblPr")
            table._tbl.insert(0, tbl_pr)
        tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
        if tbl_layout is None:
            tbl_layout = OxmlElement("w:tblLayout")
            tbl_pr.append(tbl_layout)
        tbl_layout.set(qn("w:type"), layout)

        tbl_w = tbl_pr.first_child_found_in("w:tblW")
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(TABLE_TOTAL_WIDTH_DXA))
        tbl_w.set(qn("w:type"), "dxa")

        tbl_grid = table._tbl.tblGrid
        if tbl_grid is None:
            tbl_grid = OxmlElement("w:tblGrid")
            table._tbl.insert(1, tbl_grid)
        for grid_col in list(tbl_grid.findall(qn("w:gridCol"))):
            tbl_grid.remove(grid_col)
        for width in widths:
            grid_col = OxmlElement("w:gridCol")
            grid_col.set(qn("w:w"), str(int(width)))
            tbl_grid.append(grid_col)

    def _set_borders(self, table: Any) -> None:
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tbl_pr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = borders.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                borders.append(node)
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), "4")
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), "D9E1E8")

    def _normalize_widths_dxa(self, widths: list[float] | None, headers: list[str]) -> list[int]:
        if not headers:
            return []
        if not widths or len(widths) != len(headers):
            return self._semantic_widths_dxa(headers)
        numeric_widths: list[float] = []
        try:
            numeric_widths = [float(width) for width in widths]
        except (TypeError, ValueError):
            return self._semantic_widths_dxa(headers)
        if any(width <= 0 for width in numeric_widths):
            return self._semantic_widths_dxa(headers)
        if max(numeric_widths) <= 20 and sum(numeric_widths) <= 50:
            raw = [int(round(width * 1440)) for width in numeric_widths]
        else:
            raw = [int(round(width)) for width in numeric_widths]
        total = sum(raw)
        if total <= 0:
            return self._semantic_widths_dxa(headers)
        normalized = [max(1, int(round(width * TABLE_TOTAL_WIDTH_DXA / total))) for width in raw]
        diff = TABLE_TOTAL_WIDTH_DXA - sum(normalized)
        if diff:
            target = max(range(len(normalized)), key=lambda idx: normalized[idx])
            normalized[target] += diff
        return normalized

    def _semantic_widths_dxa(self, headers: list[str]) -> list[int]:
        if not headers:
            return []
        if len(headers) == 1:
            return [TABLE_TOTAL_WIDTH_DXA]
        long_indexes = [idx for idx, header in enumerate(headers) if self._is_long_text_column(header)]
        if not long_indexes:
            base = TABLE_TOTAL_WIDTH_DXA // len(headers)
            widths = [base] * len(headers)
            widths[-1] += TABLE_TOTAL_WIDTH_DXA - sum(widths)
            return widths
        if len(long_indexes) == 1:
            widths = [0] * len(headers)
            long_idx = long_indexes[0]
            widths[long_idx] = min(6804, max(6000, TABLE_TOTAL_WIDTH_DXA - (len(headers) - 1) * 900))
            remaining = TABLE_TOTAL_WIDTH_DXA - widths[long_idx]
            short_indexes = [idx for idx in range(len(headers)) if idx != long_idx]
            base = remaining // len(short_indexes)
            for idx in short_indexes:
                widths[idx] = base
            widths[short_indexes[-1]] += remaining - base * len(short_indexes)
            return widths
        base = TABLE_TOTAL_WIDTH_DXA // len(headers)
        widths = [base] * len(headers)
        widths[-1] += TABLE_TOTAL_WIDTH_DXA - sum(widths)
        return widths

    def _is_long_text_column(self, header: str) -> bool:
        return any(keyword in header for keyword in LONG_TEXT_HEADER_KEYWORDS)

    def _add_heading(self, document: Document, text: str, level: int) -> None:
        paragraph = document.add_heading(self._clean(text), level=level)
        for run in paragraph.runs:
            self._format_run(run, bold=True)

    def _add_paragraph(self, document: Document, text: Any) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(self._clean(text))
        self._format_run(run, size=10.5)

    def _add_bullets(self, document: Document, items: list[str], empty_text: str) -> None:
        rows = items or [empty_text]
        for item in rows:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(self._clean(item))
            self._format_run(run, size=10.5)

    def _add_risk_sentence(self, document: Document, group: RiskGroup) -> None:
        finding = group.findings[0]
        observed = self._sentence_observation(group)
        impact = self._sentence_impact(finding)
        recommendation = self._recommendation(finding)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        level_run = paragraph.add_run(f"{group.level}级。")
        self._format_run(level_run, size=10.5, bold=True, color=LEVEL_STYLE.get(group.level, LEVEL_STYLE["P4"])["text"])
        body = f"检测到{observed}，可能导致{impact}。建议{recommendation}。"
        body_run = paragraph.add_run(self._clean(body))
        self._format_run(body_run, size=10.5)

    def _set_cell_text(self, cell: Any, text: Any, bold: bool = False, align_center: bool = False) -> None:
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if align_center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        value = self._clean(text)
        if not bold:
            value = self._wrap_long_tokens(value)
        run = paragraph.add_run(value)
        self._format_run(run, size=9 if bold else TABLE_DATA_FONT_SIZE, bold=bold)

    def _format_run(self, run: Any, size: float | None = None, bold: bool = False, color: RGBColor | None = None) -> None:
        run.bold = bold
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        if size is not None:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color

    def _set_cell_width(self, cell: Any, width_dxa: int) -> None:
        cell.width = Twips(int(width_dxa))
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = tc_pr.first_child_found_in("w:tcW")
        if tc_w is None:
            tc_w = OxmlElement("w:tcW")
            tc_pr.append(tc_w)
        tc_w.set(qn("w:w"), str(int(width_dxa)))
        tc_w.set(qn("w:type"), "dxa")

    def _shade_cell(self, cell: Any, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.first_child_found_in("w:shd")
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), fill)

    def _color_cell_text(self, cell: Any, color: RGBColor, bold: bool = False) -> None:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = color
                run.bold = bold

    def _repeat_header(self, row: Any) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)

    def _keep_row_together(self, row: Any) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.first_child_found_in("w:cantSplit") is None:
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)

    def _wrap_long_tokens(self, text: str) -> str:
        return LONG_TOKEN_PATTERN.sub(lambda match: self._wrap_token(match.group(0)), text)

    def _wrap_token(self, token: str) -> str:
        # Preserve object identifiers exactly so customers can copy and search them.
        return token

    def _clean(self, value: Any) -> str:
        text = "" if value is None else str(value)
        for old, new in COMMON_MOJIBAKE_REPLACEMENTS.items():
            text = text.replace(old, new)
        for old, new in FORBIDDEN_REPLACEMENTS.items():
            text = text.replace(old, new)
        for old, new in CUSTOMER_WORD_REPLACEMENTS.items():
            text = text.replace(old, new)
        for pattern, replacement in CUSTOMER_WORD_REGEX_REPLACEMENTS:
            text = pattern.sub(replacement, text)
        text = text.replace("\r", " ").replace("\n", " ")
        return " ".join(text.split()).strip()

    def _fmt(self, value: Any, default: str = "未采集") -> str:
        if value is None or value == "":
            return default
        if isinstance(value, bool):
            return "是" if value else "否"
        if isinstance(value, float):
            return f"{value:.2f}".rstrip("0").rstrip(".")
        return self._clean(value)

    def _fmt_percent(self, value: Any) -> str:
        if value is None or value == "":
            return "未采集"
        try:
            return f"{float(value):.1f}%"
        except (TypeError, ValueError):
            return self._fmt(value)

    def _join_display(self, values: Any) -> str:
        if values is None or values == "":
            return "未采集"
        if isinstance(values, (list, tuple, set)):
            parts = [self._fmt(value, default="") for value in values if value not in (None, "")]
            return "、".join(part for part in parts if part) or "未采集"
        return self._fmt(values)

    def _health_status_label(self, value: Any) -> str:
        mapping = {"red": "红色", "yellow": "黄色", "green": "绿色", "unknown": "未知"}
        return mapping.get(str(value or "").strip().casefold(), self._fmt(value))

    def _ratio_label(self, value: Any) -> str:
        if value is None or value == "":
            return "未采集"
        try:
            return f"{float(value):.2f}:1"
        except (TypeError, ValueError):
            return self._fmt(value)

    def _ratio_percent_label(self, value: Any) -> str:
        if value is None or value == "":
            return "未采集"
        try:
            return f"{float(value) * 100:.1f}%"
        except (TypeError, ValueError):
            return self._fmt(value)

    def _safe_float(self, value: Any) -> float | None:
        if value is None or value == "":
            return None
        try:
            return float(value)
        except (TypeError, ValueError):
            return None

    def _safe_float_sort_key(self, value: Any) -> float:
        number = self._safe_float(value)
        return number if number is not None else -1.0

    def _is_high_percent(self, value: Any, threshold: float) -> bool:
        number = self._safe_float(value)
        return number is not None and number >= threshold

    def _truthy(self, value: Any) -> bool:
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return value > 0
        return str(value).strip().lower() in {"true", "yes", "enabled", "running", "1", "是", "已启用", "运行"}

    def _asset_inventory(self, report: ReportData) -> dict[str, Any]:
        return report.asset_inventory or {}

    def _asset_summary(self, report: ReportData) -> dict[str, int]:
        return (self._asset_inventory(report).get("summary") or {}) if isinstance(report.asset_inventory, dict) else {}

    def _asset_details(self, report: ReportData, key: str) -> list[dict[str, Any]]:
        details = self._asset_inventory(report).get("details") or {}
        return list(details.get(key) or [])

    def _asset_props(self, report: ReportData, key: str, name: str) -> dict[str, Any]:
        for item in self._asset_details(report, key):
            if item.get("object_name") == name:
                return item.get("properties") or {}
        return {}

    def _unique_datacenters(self, report: ReportData) -> list[str]:
        values: set[str] = set()
        details = self._asset_inventory(report).get("details") or {}
        for items in details.values():
            for item in items or []:
                path = item.get("asset_location") or item.get("object_path") or (item.get("properties") or {}).get("asset_location") or ""
                parts = [part.strip() for part in str(path).split("/") if part.strip()]
                if len(parts) >= 2:
                    values.add(parts[1])
        return sorted(values)

    def _cluster_from_location(self, value: Any) -> str:
        text = self._clean(value)
        if not text:
            return "未采集"
        parts = [part.strip() for part in text.split("/") if part.strip()]
        if len(parts) >= 3:
            candidate = parts[2]
            if candidate.startswith("挂载集群 "):
                return candidate.replace("挂载集群 ", "", 1)
            if candidate.startswith("主机 "):
                return "未采集"
            return candidate
        return "未采集"

    def _short_location(self, value: Any) -> str:
        text = self._clean(value)
        if not text:
            return "未采集"
        parts = [part.strip() for part in text.split("/") if part.strip()]
        if len(parts) >= 3:
            return " / ".join(parts[-2:])
        return text

    def _host_cluster(self, report: ReportData, host_name: str) -> str:
        props = self._asset_props(report, "HostSystem", host_name)
        return self._cluster_from_location(props.get("asset_location")) if props else "未采集"

    def _vm_location(self, report: ReportData, finding: FindingItem) -> str:
        props = self._asset_props(report, "VirtualMachine", finding.object_name)
        location = props.get("asset_location") or finding.object_path
        cluster = self._cluster_from_location(location)
        if cluster != "未采集":
            return cluster
        return self._short_location(location)

    def _host_state_label(self, value: Any) -> str:
        mapping = {"connected": "正常连接", "disconnected": "连接断开", "notResponding": "无响应"}
        return mapping.get(str(value), self._fmt(value))

    def _power_state_label(self, value: Any) -> str:
        mapping = {"poweredOn": "运行", "poweredOff": "关机", "suspended": "挂起"}
        return mapping.get(str(value), self._fmt(value))

    def _enabled_label(self, value: Any) -> str:
        if value is True:
            return "已启用"
        if value is False:
            return "未启用"
        return self._fmt(value)

    def _host_attention(self, props: dict[str, Any]) -> str:
        items = []
        if props.get("ssh_running"):
            items.append("SSH/Shell 入口需定期复核")
        if props.get("syslog_configured") is False:
            items.append("日志外送配置需关注")
        if props.get("physical_nic_link_issue_count"):
            items.append("网络链路状态需关注")
        if props.get("storage_path_dead_count"):
            items.append("存储路径状态需关注")
        return "；".join(items) if items else "状态正常"

    def _storage_attention(self, props: dict[str, Any]) -> str:
        used = self._safe_float(props.get("used_percent"))
        if props.get("accessible") is False:
            return "访问状态需关注"
        if used is None:
            return "未采集"
        if used >= 80:
            return "建议持续观察容量增长"
        if props.get("thin_provisioning_overcommit_ratio"):
            return "Thin Provisioning 使用需定期复核"
        return "状态正常"

    def _cluster_attention(self, props: dict[str, Any]) -> str:
        items = []
        if props.get("ha_enabled") is False:
            items.append("HA 未启用")
        if props.get("drs_enabled") is False:
            items.append("DRS 未启用")
        missing = props.get("cluster_vmotion_missing_host_count")
        if missing:
            items.append(f"{missing} 台主机缺少 vMotion 网络")
        return "；".join(items) if items else "状态正常"

    def _tools_label(self, props: dict[str, Any]) -> str:
        if props.get("vmware_tools_installed") is False:
            return "未安装"
        if props.get("tools_outdated"):
            return "版本较旧"
        if props.get("tools_running") is False:
            return "未运行"
        if props:
            return "正常"
        return "未采集"

    def _tools_status_from_finding(self, finding: FindingItem) -> str:
        priority = self._tools_priority(finding)
        if priority == 0:
            return "未安装"
        if priority == 1:
            return "未运行"
        if priority == 2:
            return "版本较旧"
        return self._fmt(finding.current_value)

    def _version_certificate_license_rows(self, report: ReportData) -> list[list[str]]:
        rows: list[list[str]] = []
        for item in self._asset_details(report, "vCenter"):
            props = item.get("properties") or {}
            rows.append(
                [
                    f"vCenter {item.get('object_name', '')}",
                    f"vCenter {self._fmt(props.get('version'))} / Build {self._fmt(props.get('build'))}",
                    self._certificate_label(props.get("certificate_days_remaining")),
                    self._license_label(props.get("license_days_remaining")),
                    self._status_summary(props.get("certificate_days_remaining"), props.get("license_days_remaining")),
                ]
            )
        for item in self._asset_details(report, "HostSystem"):
            props = item.get("properties") or {}
            rows.append(
                [
                    f"ESXi {item.get('object_name', '')}",
                    self._fmt(props.get("host_license_name"), "未采集产品授权信息"),
                    self._certificate_label(props.get("host_certificate_days_remaining")),
                    self._license_label(props.get("host_license_expiration_days")),
                    self._status_summary(props.get("host_certificate_days_remaining"), props.get("host_license_expiration_days")),
                ]
            )
        return rows

    def _certificate_label(self, value: Any) -> str:
        if value is None or value == "":
            return "未采集"
        try:
            return f"证书剩余 {int(float(value))} 天"
        except (TypeError, ValueError):
            return self._fmt(value)

    def _license_label(self, value: Any) -> str:
        if value is None or value == "":
            return "未采集"
        try:
            days = int(float(value))
        except (TypeError, ValueError):
            return self._fmt(value)
        if days >= 999999:
            return "永久授权"
        if days < 0:
            return "已过期"
        return f"剩余 {days} 天"

    def _status_summary(self, certificate_days: Any, license_days: Any) -> str:
        cert_ok = self._days_ok(certificate_days, 90)
        license_ok = self._days_ok(license_days, 30)
        if cert_ok and license_ok:
            return "状态正常"
        return "建议定期复核"

    def _days_ok(self, value: Any, threshold: int) -> bool:
        try:
            return int(float(value)) >= threshold
        except (TypeError, ValueError):
            return False

    def _has_vsan(self, report: ReportData) -> bool:
        for item in self._asset_details(report, "ClusterComputeResource") + self._asset_details(report, "Datastore"):
            text = f"{item.get('object_name', '')} {item.get('object_path', '')} {item.get('asset_location', '')}"
            props = item.get("properties") or {}
            if "vsan" in text.lower() or str(props.get("datastore_filesystem_type", "")).lower() == "vsan":
                return True
        return False

    def _vsan_rows(self, report: ReportData) -> list[list[str]]:
        rows: list[list[str]] = []
        for item in self._asset_details(report, "ClusterComputeResource"):
            if "vsan" in str(item.get("object_name", "")).lower():
                props = item.get("properties") or {}
                rows.append([item.get("object_name", ""), "已识别", self._cluster_attention(props)])
        for item in self._asset_details(report, "Datastore"):
            props = item.get("properties") or {}
            if str(props.get("datastore_filesystem_type", "")).lower() == "vsan" or "vsan" in str(item.get("object_name", "")).lower():
                rows.append([item.get("object_name", ""), "可访问" if props.get("accessible") is not False else "不可访问", self._storage_attention(props)])
        return rows

    def _observed(self, finding: FindingItem) -> dict[str, Any]:
        if isinstance(finding.observed_detail, dict) and finding.observed_detail:
            return finding.observed_detail
        raw = finding.raw_evidence or {}
        observed = raw.get("observed_detail") if isinstance(raw, dict) else None
        return observed if isinstance(observed, dict) else {}

    def _current_display(self, finding: FindingItem, default: str = "未采集") -> str:
        return self._clean(finding.current_value_zh) or self._fmt(finding.current_value, default=default)

    def _status_value(self, finding: FindingItem) -> str:
        if finding.evidence_summary_zh:
            return self._clean(finding.evidence_summary_zh)
        if finding.current_value_zh:
            return self._clean(finding.current_value_zh)
        if finding.current_value is not None:
            return self._fmt(finding.current_value)
        if finding.observed_detail_zh and finding.observed_detail_zh != "未列出":
            return self._clean(finding.observed_detail_zh)
        observed = self._observed(finding)
        if observed:
            return self._clean(", ".join(f"{key}: {self._fmt(value)}" for key, value in list(observed.items())[:3]))
        return "需关注"

    def _fallback_status_value(self, finding: FindingItem) -> str:
        status = self._status_value(finding)
        if len(status) <= 48:
            return status
        current = self._clean(finding.current_value_zh) or self._fmt(finding.current_value, default="")
        parts = []
        if current:
            parts.append(f"当前：{self._short_status_fragment(current)}")
        expected = self._clean(finding.expected_value_zh) or self._fmt(finding.expected_value, default="")
        if expected:
            parts.append(f"建议：{self._short_status_fragment(expected)}")
        return "；".join(parts) if parts else self._short_status_fragment(status, max_len=48)

    def _short_status_fragment(self, value: Any, max_len: int = 36) -> str:
        text = self._clean(value)
        if len(text) <= max_len:
            return text
        for separator in ("；", ";", "，", ",", "。"):
            head = text.split(separator, 1)[0].strip()
            if head and len(head) < len(text):
                text = head
                break
        if len(text) > max_len:
            return f"{text[: max_len - 1]}…"
        return text

    def _sentence_observation(self, group: RiskGroup) -> str:
        finding = group.findings[0]
        if len(group.findings) > 1:
            return f"{len(group.findings)} 个对象存在“{group.title}”"
        if finding.evidence_summary_zh:
            text = self._clean(finding.evidence_summary_zh)
            if text.endswith("。"):
                text = text[:-1]
            return text
        return f"{finding.object_name} 存在“{group.title}”"

    def _sentence_impact(self, finding: FindingItem) -> str:
        for value in (finding.business_impact, finding.consequence, finding.technical_impact):
            text = self._clean(value)
            if text:
                return text.rstrip("。；;，,")
        return "环境运行稳定性或后续运维效率受到影响"

    def _recommendation(self, finding: FindingItem) -> str:
        if self._is_tools_finding(finding):
            return TOOLS_CUSTOMER_RECOMMENDATION
        for value in (finding.recommended_action_zh, finding.remediation):
            text = self._clean(value)
            if text:
                return text.rstrip("。；;，,")
        return "结合实际业务影响评估处理"

    def _is_tools_finding(self, finding: FindingItem) -> bool:
        rule_id = self._clean(finding.rule_id).upper()
        if rule_id in {"VSL-VM-002", "VSL-VM-003", "VSL-VM-014"}:
            return True
        text = self._clean(
            " ".join(
                self._clean(getattr(finding, name, ""))
                for name in (
                    "rule_id",
                    "rule_name",
                    "title",
                    "current_value",
                    "expected_value",
                    "current_value_zh",
                    "expected_value_zh",
                    "evidence_summary_zh",
                    "observed_detail_zh",
                    "expected_detail_zh",
                    "remediation",
                    "recommended_action_zh",
                )
            )
        ).casefold()
        if "vmware tools" not in text and "tools" not in text:
            return False
        return any(
            keyword in text
            for keyword in (
                "未安装",
                "not installed",
                "install state",
                "未运行",
                "not running",
                "notrunning",
                "版本过旧",
                "版本较旧",
                "outdated",
                "upgrade",
                "tools_running",
                "tools_outdated",
                "vmware_tools_installed",
            )
        )

    def _suggestions(self, groups: list[RiskGroup], prefix: str, suffix: str) -> list[str]:
        result = []
        for group in sorted(groups, key=self._risk_group_sort_key)[:5]:
            result.append(f"{prefix}“{group.title}”，{suffix}")
        return result

    def _date_label(self, value: str | None) -> str:
        parsed = self._parse_datetime(value)
        return parsed.strftime("%Y-%m-%d") if parsed else self._clean(value or "未记录")

    def _time_label(self, value: Any) -> str:
        parsed = self._parse_datetime(str(value)) if value else None
        return parsed.strftime("%Y-%m-%d %H:%M") if parsed else self._fmt(value)

    def _snapshot_size_label(self, snapshot: dict[str, Any]) -> str:
        for key in ("size_gb", "snapshot_size_gb", "used_gb", "capacity_gb"):
            if key in snapshot and snapshot.get(key) not in (None, ""):
                return self._gb_size_label(snapshot.get(key))
        for key in ("size_mb", "snapshot_size_mb", "used_mb", "capacity_mb"):
            if key in snapshot and snapshot.get(key) not in (None, ""):
                return self._size_label(snapshot.get(key))
        for key in ("size_bytes", "snapshot_size_bytes", "used_bytes", "capacity_bytes"):
            if key in snapshot and snapshot.get(key) not in (None, ""):
                return self._bytes_size_label(snapshot.get(key))
        for key in ("size", "capacity"):
            if key in snapshot and snapshot.get(key) not in (None, ""):
                return self._ambiguous_size_label(snapshot.get(key), infer_bytes=key == "capacity")
        return "未采集"

    def _snapshot_name_label(self, snapshot: dict[str, Any]) -> str:
        name = snapshot.get("name") or snapshot.get("snapshot_name")
        return self._fmt(name, default="虚拟机存在快照")

    def _snapshot_size_display(self, snapshot: dict[str, Any]) -> str:
        size = self._snapshot_size_label(snapshot)
        return "" if size == "未采集" else size

    def _snapshot_retention_label(self, created_at: Any, report: ReportData) -> str:
        if created_at is None or created_at == "":
            return "未识别"
        parsed = self._parse_datetime(str(created_at))
        if not parsed:
            return self._time_label(created_at) or self._fmt(created_at)
        reference = self._parse_datetime(report.report_info.generated_at) if report.report_info else None
        if reference is None:
            reference = datetime.now(parsed.tzinfo)
        if parsed.tzinfo is not None and reference.tzinfo is None:
            reference = reference.replace(tzinfo=parsed.tzinfo)
        if parsed.tzinfo is None and reference.tzinfo is not None:
            parsed = parsed.replace(tzinfo=reference.tzinfo)
        try:
            days = max(0, (reference - parsed).days)
        except TypeError:
            return self._time_label(created_at) or self._fmt(created_at)
        return f"约 {days} 天"

    def _parse_datetime(self, value: str | None) -> datetime | None:
        if not value:
            return None
        text = value.replace("Z", "+00:00")
        try:
            return datetime.fromisoformat(text)
        except ValueError:
            return None

    def _speed_label(self, value: Any) -> str:
        if value is None or value == "":
            return "未协商"
        try:
            return f"{int(float(value))} Mbps"
        except (TypeError, ValueError):
            return self._fmt(value)

    def _size_label(self, value: Any) -> str:
        if value is None or value == "":
            return "未采集"
        try:
            number = float(value)
        except (TypeError, ValueError):
            return self._fmt(value)
        if number >= 1024:
            return f"{number / 1024:.1f} GB"
        return f"{number:.1f} MB"

    def _gb_size_label(self, value: Any) -> str:
        if value is None or value == "":
            return "未采集"
        try:
            number = float(value)
        except (TypeError, ValueError):
            return self._fmt(value)
        return f"{number:.1f} GB"

    def _bytes_size_label(self, value: Any) -> str:
        if value is None or value == "":
            return "未采集"
        try:
            number = float(value)
        except (TypeError, ValueError):
            return self._fmt(value)
        mb = number / (1024 * 1024)
        if mb >= 1024:
            return f"{mb / 1024:.1f} GB"
        return f"{mb:.1f} MB"

    def _ambiguous_size_label(self, value: Any, infer_bytes: bool = False) -> str:
        text = self._clean(value)
        if re.search(r"(?i)\b(?:kb|mb|gb|tb|bytes?|b)\b", text):
            return text
        try:
            number = float(value)
        except (TypeError, ValueError):
            return text or "未采集"
        if infer_bytes and number >= 10_000_000:
            return self._bytes_size_label(number)
        return self._size_label(number)
