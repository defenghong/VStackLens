from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor, Twips

from vstacklens.reports.docx_report import (
    LEVEL_ORDER,
    LEVEL_STYLE,
    TABLE_TOTAL_WIDTH_DXA,
    DocxReportEngine,
    RiskGroup,
)
from vstacklens.reports.report_model import FindingItem, ReportData


PORTRAIT_CONTENT_DXA = 9650
LANDSCAPE_CONTENT_DXA = 14500
HEADER_BLUE = "17365D"
ACCENT_BLUE = "2F75B5"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F7FA"
BORDER_GRAY = "D9E1E8"
TEXT_GRAY = RGBColor(80, 80, 80)
BLACK = RGBColor(0, 0, 0)
REPORT_SECTIONS = (
    "1. 巡检报告概览",
    "2. 环境健康状态",
    "3. 检查结果汇总",
    "4. 重点问题详情与整改建议",
    "5. VMware 平台巡检",
    "6. vSAN 巡检",
    "7. 历史对比与整改情况",
    "8. 环境资产与对象清单",
    "9. 附录",
)


class DocxReportV2Engine(DocxReportEngine):
    """Render the staged Word V2 presentation layer.

    The class intentionally reuses the legacy engine only for existing finding
    grouping, customer-language normalization, and detail extraction. It owns
    the document shell so the report can be migrated section by section.
    """

    def render(self, report: ReportData, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        self._setup_v2_document(document, report)
        self._document = document
        risk_groups = self._display_risk_groups(report)

        self._cover_v2(document, report)
        self._summary_page(document, report, risk_groups)
        self._version_and_toc(document, report)
        self._chapter_overview(document, report)
        self._chapter_health(document, report, risk_groups)
        self._chapter_results(document, report, risk_groups)
        self._chapter_problem_details(document, report, risk_groups)
        self._chapter_vmware(document, report)
        self._chapter_vsan(document, report)
        self._chapter_history(document, report)
        self._chapter_assets(document, report)
        self._chapter_appendix(document, report)

        document.save(output_path)
        return output_path

    def _setup_v2_document(self, document: Document, report: ReportData) -> None:
        section = document.sections[0]
        self._configure_section(section, landscape=False)
        section.different_first_page_header_footer = True
        self._configure_styles(document)
        self._configure_header_footer(section, report)
        section.first_page_header.paragraphs[0].text = ""
        section.first_page_footer.paragraphs[0].text = ""

        settings = document.settings.element
        if settings.find(qn("w:updateFields")) is None:
            update_fields = OxmlElement("w:updateFields")
            update_fields.set(qn("w:val"), "true")
            settings.append(update_fields)

    def _configure_section(self, section: Any, *, landscape: bool) -> None:
        if landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width = Mm(297)
            section.page_height = Mm(210)
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width = Mm(210)
            section.page_height = Mm(297)
        section.top_margin = Mm(18)
        section.bottom_margin = Mm(18)
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)
        section.header_distance = Mm(8)
        section.footer_distance = Mm(8)

    def _configure_styles(self, document: Document) -> None:
        styles = document.styles
        self._set_style_font(styles["Normal"], 10.5, bold=False, color=BLACK)
        styles["Normal"].paragraph_format.space_after = Pt(5)
        styles["Normal"].paragraph_format.line_spacing = 1.15
        self._set_style_font(styles["Title"], 26, bold=True, color=BLACK)
        styles["Title"].paragraph_format.space_after = Pt(8)
        self._remove_style_borders(styles["Title"])
        for name, size in (("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 11.5)):
            style = styles[name]
            self._set_style_font(style, size, bold=True, color=BLACK)
            style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
            style.paragraph_format.space_after = Pt(5)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True
        for name in ("TOC 1", "TOC 2", "TOC 3"):
            if name in styles:
                self._set_style_font(styles[name], 10.5, bold=False, color=BLACK)

    def _set_style_font(self, style: Any, size: float, *, bold: bool, color: RGBColor) -> None:
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color

    def _configure_header_footer(self, section: Any, report: ReportData) -> None:
        customer = self._clean(report.customer_info.customer_name or "未指定客户")
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header.paragraph_format.space_after = Pt(0)
        header.paragraph_format.tab_stops.add_tab_stop(Cm(16.8), WD_TAB_ALIGNMENT.RIGHT)
        left = header.add_run("VStackLens 虚拟化平台健康巡检报告")
        self._format_run(left, size=8.5, color=TEXT_GRAY)
        right = header.add_run(f"\t{customer}")
        self._format_run(right, size=8.5, color=TEXT_GRAY)

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(0)
        footer_run = footer.add_run(
            f"保密资料 | 巡检日期：{self._date_label(report.report_info.generated_at)} | 第 "
        )
        self._format_run(footer_run, size=8, color=TEXT_GRAY)
        self._append_field(footer, "PAGE")
        suffix = footer.add_run(" 页")
        self._format_run(suffix, size=8, color=TEXT_GRAY)

    def _append_field(self, paragraph: Any, instruction: str) -> None:
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = instruction
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        separate.append(text)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend((begin, instr, separate, end))

    def _cover_v2(self, document: Document, report: ReportData) -> None:
        for _ in range(3):
            document.add_paragraph()
        brand = document.add_paragraph()
        brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
        brand_run = brand.add_run("VStackLens")
        self._format_run(brand_run, size=16, bold=True, color=RGBColor(47, 117, 181))

        title = document.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run("虚拟化平台健康巡检报告")
        self._format_run(title_run, size=27, bold=True, color=BLACK)

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(28)
        subtitle_run = subtitle.add_run(self._clean(report.customer_info.project_name or "VMware 虚拟化健康评估"))
        self._format_run(subtitle_run, size=12, color=TEXT_GRAY)

        metadata = (
            ("客户名称", report.customer_info.customer_name),
            ("检查站点", report.customer_info.site_name),
            ("巡检日期", self._date_label(report.report_info.generated_at)),
            ("报告版本", report.report_info.report_version or "V1.0"),
            ("生成时间", self._date_time_label(report.report_info.generated_at)),
        )
        table = document.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        self._set_table_fixed_width(table, [1900, 7600], PORTRAIT_CONTENT_DXA)
        self._set_table_borders(table, BORDER_GRAY)
        for label, value in metadata:
            row = table.add_row()
            self._keep_row_together(row)
            self._set_cell_text_v2(row.cells[0], label, bold=True, align_center=True)
            self._set_cell_text_v2(row.cells[1], value or "未指定")
            self._set_cell_width(row.cells[0], 1900)
            self._set_cell_width(row.cells[1], 7600)
            self._shade_cell(row.cells[0], LIGHT_BLUE)
        document.add_paragraph()
        notice = document.add_paragraph()
        notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
        notice_run = notice.add_run("本报告基于巡检时间点采集的数据生成")
        self._format_run(notice_run, size=9, color=TEXT_GRAY)
        document.add_page_break()

    def _summary_page(self, document: Document, report: ReportData, groups: list[RiskGroup]) -> None:
        self._add_heading_v2(document, "首页摘要", 1, page_break=False)
        status = self._overall_status(report)
        status_table = self._add_table_v2(
            document,
            ["整体健康状态", "当前判断", "说明"],
            [[status, self._health_status_short(status), self._health_explanation(report)]],
            widths=[1900, 1900, 5850],
            align_center_columns={0, 1},
        )
        self._shade_cell(status_table.rows[1].cells[0], self._status_fill(status))
        self._color_cell_text(status_table.rows[1].cells[0], self._status_font(status), bold=True)

        counts = self._risk_counts_v2(report, groups)
        affected = self._affected_counts_v2(report)
        self._add_heading_v2(document, "风险统计", 2, page_break=False)
        self._add_table_v2(
            document,
            ["P1", "P2", "P3", "问题类别总数", "影响对象明细"],
            [[counts["P1"], counts["P2"], counts["P3"], sum(counts.values()), sum(affected.values())]],
            widths=[1500, 1500, 1500, 2500, 2650],
            align_center_columns={0, 1, 2, 3, 4},
        )

        self._add_heading_v2(document, "环境规模", 2, page_break=False)
        scale = self._environment_scale(report)
        self._add_table_v2(
            document,
            ["vCenter", "数据中心", "Cluster", "ESXi 主机", "虚拟机", "Datastore", "vSAN 集群"],
            [[scale[key] for key in ("vcenter", "datacenter", "cluster", "host", "vm", "datastore", "vsan_cluster")]],
            widths=[1200, 1350, 1350, 1450, 1450, 1450, 1450],
            align_center_columns=set(range(7)),
        )

        self._add_heading_v2(document, "重点问题摘要", 2, page_break=False)
        top_rows = self._summary_rows_v2(report, groups)[:5]
        if top_rows:
            self._add_issue_summary_v2(document, top_rows)
        else:
            self._add_paragraph_v2(document, "本次巡检未发现 P1/P2/P3 级风险项。")

        document.add_page_break()

    def _version_and_toc(self, document: Document, report: ReportData) -> None:
        self._add_heading_v2(document, "版本记录与目录", 1, page_break=False)
        version = report.report_info.report_version or "V1.0"
        self._add_table_v2(
            document,
            ["版本", "日期", "说明"],
            [[version, self._date_label(report.report_info.generated_at), "Word V2 第一阶段展示层"]],
            widths=[1800, 2400, 5850],
        )
        self._add_heading_v2(document, "目录", 2, page_break=False)
        sections = [item for item in REPORT_SECTIONS if item != "7. 历史对比与整改情况" or self._history_visible(report)]
        rows = [[item, ""] for item in sections]
        self._add_table_v2(document, ["章节", "页码"], rows, widths=[7600, 2450])
        self._add_paragraph_v2(document, "目录章节采用固定结构生成；页码以最终分页结果为准。")
        document.add_page_break()

    def _chapter_overview(self, document: Document, report: ReportData) -> None:
        self._add_heading_v2(document, "1. 巡检报告概览", 1)
        self._add_heading_v2(document, "1.1 巡检说明", 2)
        self._add_paragraph_v2(document, "本报告用于评估 VMware 虚拟化环境的配置、运行、容量、可用性和最佳实践状态，识别需要关注的问题并提供整改建议。")
        self._add_heading_v2(document, "1.2 巡检范围", 2)
        scale = self._environment_scale(report)
        scope_rows = [[label, scale[key]] for key, label in (
            ("vcenter", "vCenter"), ("datacenter", "数据中心"), ("cluster", "Cluster"),
            ("host", "ESXi 主机"), ("datastore", "Datastore"), ("vm", "虚拟机"), ("vsan_cluster", "vSAN 集群"),
        ) if scale[key] or key != "vsan_cluster"]
        self._add_table_v2(document, ["资源类型", "覆盖数量"], scope_rows, widths=[6500, 3350], align_center_columns={1})

        self._add_heading_v2(document, "1.3 风险等级说明", 2)
        self._add_table_v2(
            document,
            ["等级", "定义"],
            [["P1", "需要重点关注的问题，建议优先确认影响范围和处理窗口。"],
             ["P2", "非关键问题，但需要关注或进一步调查。"],
             ["P3", "优化建议或持续关注事项。"]],
            widths=[1500, 8350],
        )

        self._add_heading_v2(document, "1.4 环境架构概览", 2)
        self._architecture_chart(document, report)
        self._add_heading_v2(document, "1.5 报告说明", 2)
        self._add_paragraph_v2(document, "本报告反映巡检时间点的环境状态。部分检查结果仍需结合业务需求、厂商建议和维护窗口进一步确认。")

    def _chapter_health(self, document: Document, report: ReportData, groups: list[RiskGroup]) -> None:
        self._add_heading_v2(document, "2. 环境健康状态", 1)
        self._add_heading_v2(document, "2.1 整体健康状态", 2)
        self._add_paragraph_v2(document, f"当前环境健康状态为 {self._overall_status(report)}。{self._health_explanation(report)}")
        self._add_heading_v2(document, "2.2 风险等级统计", 2)
        counts = self._risk_counts_v2(report, groups)
        self._risk_chart(document, counts)
        resource_rows = self._resource_rows(report)
        if resource_rows:
            self._add_heading_v2(document, "2.3 资源概览", 2)
            self._resource_chart(document, resource_rows)
        self._add_heading_v2(document, "2.4 重点问题", 2)
        top_rows = self._summary_rows_v2(report, groups)[:5]
        if top_rows:
            self._add_issue_summary_v2(document, top_rows)
        else:
            self._add_paragraph_v2(document, "本次巡检未发现 P1/P2/P3 级风险项。")
        self._add_heading_v2(document, "2.5 整改优先级", 2)
        priorities = []
        if counts["P1"] or counts["P2"]:
            priorities.append("第一阶段：优先确认 P1/P2 问题的影响对象、根因和维护窗口。")
        if counts["P3"]:
            priorities.append("第二阶段：将 P3 问题纳入近期运维优化计划。")
        priorities.append("整改后建议重新巡检，确认问题状态和影响范围发生变化。")
        self._add_bullets_v2(document, priorities)

    def _chapter_results(self, document: Document, report: ReportData, groups: list[RiskGroup]) -> None:
        self._add_heading_v2(document, "3. 检查结果汇总", 1)
        self._add_heading_v2(document, "3.1 问题类别汇总", 2)
        rows = self._summary_rows_v2(report, groups)
        if rows:
            self._add_issue_summary_v2(document, rows)
        else:
            self._add_paragraph_v2(document, "本次巡检未发现 P1/P2/P3 级风险项。")
        self._add_heading_v2(document, "3.2 已验证正常摘要", 2)
        normal = self._normal_summary(report)
        if normal:
            self._add_bullets_v2(document, normal)
        else:
            self._add_paragraph_v2(document, "本次没有可单独展示的正常项摘要。")

    def _chapter_problem_details(self, document: Document, report: ReportData, groups: list[RiskGroup]) -> None:
        self._add_heading_v2(document, "4. 重点问题详情与整改建议", 1)
        category_map = {
            str(item.get("category_id") or item.get("rule_id") or ""): item
            for item in (report.report_presentation or {}).get("problem_categories", [])
        }
        by_level = {level: [group for group in groups if group.level == level] for level in ("P1", "P2", "P3")}
        for level_index, level in enumerate(("P1", "P2", "P3"), start=1):
            section = f"4.{level_index}"
            self._add_heading_v2(document, f"{section} {level}", 2)
            items = sorted(by_level[level], key=self._risk_group_sort_key)
            if not items:
                self._add_paragraph_v2(document, f"本次未发现 {level} 问题类别。")
                continue
            for item_index, group in enumerate(items, start=1):
                self._add_heading_v2(document, f"{section}.{item_index} {group.title}", 3)
                category = category_map.get(group.rule_id, {})
                self._add_table_v2(
                    document,
                    ["优先级", "影响对象", "整改方式", "整改复杂度"],
                    [[group.level, category.get("affected_object_count", len(group.findings)), category.get("remediation_mode") or "需结合场景评估", category.get("remediation_effort") or "中"]],
                    widths=[1800, 2500, 3000, 2650],
                    align_center_columns={0, 1, 2, 3},
                )
                self._problem_text(document, report, group, category)
                detail = self._detail_table(report, group)
                if detail:
                    headers, rows = detail
                    self._add_heading_v2(document, "详细问题明细", 3, page_break=False)
                    self._add_table_v2(document, headers, rows, widths=self._detail_widths(headers), landscape=self._wide_headers(headers))
                else:
                    self._add_heading_v2(document, "影响范围", 3, page_break=False)
                    rows = [[finding.object_name, finding.object_type, finding.object_path or "未采集"] for finding in group.findings]
                    self._add_table_v2(document, ["对象名称", "对象类型", "所在位置"], rows, widths=[3000, 2200, 4750])

    def _chapter_vmware(self, document: Document, report: ReportData) -> None:
        self._add_heading_v2(document, "5. VMware 平台巡检", 1)
        modules = (report.report_presentation or {}).get("vmware_modules") or []
        rows = []
        for module in modules:
            rows.append([
                module.get("label") or module.get("module_id") or "未命名模块",
                len(module.get("problem_categories") or []),
                len(module.get("normal_checks") or []),
                len(module.get("unconfirmed_checks") or []),
            ])
        if rows:
            self._add_table_v2(document, ["模块", "问题类别", "已验证正常", "未确认"], rows, widths=[4300, 1800, 1900, 1950], align_center_columns={1, 2, 3})
        else:
            self._add_paragraph_v2(document, "本次未生成 VMware 模块汇总数据。")
        self._add_paragraph_v2(document, "各模块的详细检查对象和问题整改内容见第四章；工程明细将在后续版本继续迁移到对应技术章节。")

    def _chapter_vsan(self, document: Document, report: ReportData) -> None:
        summary = report.vsan_summary or {}
        if not summary or summary.get("applicable") is False or summary.get("status") == "not_applicable":
            return
        self._add_heading_v2(document, "6. vSAN 巡检", 1)
        view = (report.report_presentation or {}).get("vsan") or {}
        self._add_heading_v2(document, "6.1 集群健康", 2)
        self._add_paragraph_v2(document, view.get("health_text") or "本次未能获取 vSAN 集群健康状态，建议进一步确认。")
        self._add_heading_v2(document, "6.2 磁盘与磁盘组", 2)
        self._add_table_v2(document, ["主机数", "磁盘组数", "缓存盘数", "容量盘数"], [[self._fmt(summary.get("host_count")), self._fmt(summary.get("disk_group_count")), self._fmt(summary.get("cache_disk_count")), self._fmt(summary.get("capacity_disk_count"))]], widths=[2400, 2400, 2400, 2350], align_center_columns={0, 1, 2, 3})
        self._add_heading_v2(document, "6.3 Object / VMDK", 2)
        self._add_paragraph_v2(document, view.get("object_text") or "本次未能获取 vSAN 对象状态，建议进一步确认。")
        self._add_heading_v2(document, "6.4 Resync", 2)
        self._add_paragraph_v2(document, view.get("resync_text") or "本次未能获取 vSAN 重同步状态，建议进一步确认。")
        self._add_heading_v2(document, "6.5 vSAN Network", 2)
        self._add_paragraph_v2(document, view.get("network_text") or "本次未能获取完整 vSAN 网络配置，建议进一步确认。")
        self._add_heading_v2(document, "6.6 Capacity", 2)
        capacity = view.get("capacity") or summary.get("capacity") or {}
        self._add_paragraph_v2(document, f"使用率：{self._fmt_percent(capacity.get('used_percent'))}；总容量：{self._fmt(capacity.get('total_gb'))} GB；已用：{self._fmt(capacity.get('used_gb'))} GB；可用：{self._fmt(capacity.get('free_gb'))} GB。")
        self._add_heading_v2(document, "6.7 Storage Policy", 2)
        self._add_paragraph_v2(document, view.get("policy_text") or "本次未能获取完整 Storage Policy 合规状态。")
        policy = summary.get("storage_policy_summary") or {}
        categories = policy.get("policy_categories") or []
        if categories:
            self._add_table_v2(
                document,
                ["策略名称", "策略 UUID", "已检查", "合规", "不合规", "未确认", "不适用"],
                [[item.get("policy_name") or "未命名策略", item.get("policy_uuid") or "未记录", self._fmt(item.get("checked_count")), self._fmt(item.get("compliant_count")), self._fmt(item.get("noncompliant_count")), self._fmt(item.get("unknown_count")), self._fmt(item.get("not_applicable_count"))] for item in categories],
                widths=[3000, 3300, 1400, 1400, 1400, 1400, 1400],
                landscape=True,
                align_center_columns={2, 3, 4, 5, 6},
            )

    def _chapter_history(self, document: Document, report: ReportData) -> None:
        if not self._history_visible(report):
            return
        self._add_heading_v2(document, "7. 历史对比与整改情况", 1)
        comparison = (report.appendix or {}).get("history_comparison") or {}
        summary = comparison.get("summary") or {}
        self._add_paragraph_v2(document, f"本次对比基于有效历史巡检记录，新增 {summary.get('new', 0)} 项，持续 {summary.get('existing', 0)} 项，已解决 {summary.get('resolved', 0)} 项。")
        rows = []
        for key, label in (("new_findings", "新增"), ("existing_findings", "持续"), ("resolved_findings", "已解决")):
            for item in comparison.get(key, []):
                rows.append([label, item.get("risk_level", ""), item.get("title", ""), item.get("object_name", "")])
        if rows:
            self._add_table_v2(document, ["变化", "等级", "问题类别", "对象"], rows, widths=[1500, 1500, 3900, 3050])

    def _chapter_assets(self, document: Document, report: ReportData) -> None:
        self._add_heading_v2(document, "8. 环境资产与对象清单", 1)
        summary = self._environment_scale(report)
        self._add_heading_v2(document, "8.1 资产规模", 2)
        self._add_table_v2(document, ["对象类型", "数量"], [[label, summary[key]] for key, label in (("vcenter", "vCenter"), ("datacenter", "数据中心"), ("cluster", "Cluster"), ("host", "ESXi 主机"), ("vm", "虚拟机"), ("datastore", "Datastore"), ("vsan_cluster", "vSAN 集群"))], widths=[6500, 3350], align_center_columns={1})
        self._add_heading_v2(document, "8.2 vCenter 与集群", 2)
        rows = []
        for item in self._asset_details(report, "vCenter"):
            props = item.get("properties") or {}
            rows.append([item.get("object_name", ""), self._fmt(props.get("version")), self._fmt(props.get("build")), "已连接" if props.get("connected") is not False else "未连接"])
        for item in self._asset_details(report, "ClusterComputeResource"):
            props = item.get("properties") or {}
            rows.append([item.get("object_name", ""), f"主机 {self._fmt(props.get('host_count'))}", f"HA {self._enabled_label(props.get('ha_enabled'))}", f"DRS {self._enabled_label(props.get('drs_enabled'))}"])
        self._add_table_v2(document, ["对象", "版本或规模", "Build 或 HA", "连接或 DRS"], rows or [["未采集", "未采集", "未采集", "未采集"]], widths=[3300, 2200, 2200, 2250])
        self._add_heading_v2(document, "8.3 ESXi 主机与 Datastore", 2)
        rows = []
        for item in self._asset_details(report, "HostSystem"):
            props = item.get("properties") or {}
            rows.append([item.get("object_name", ""), "ESXi 主机", self._host_state_label(props.get("connection_state")), self._fmt_percent(props.get("cpu_usage_avg")), self._fmt_percent(props.get("memory_usage_avg"))])
        for item in self._asset_details(report, "Datastore"):
            props = item.get("properties") or {}
            rows.append([item.get("object_name", ""), "Datastore", "可访问" if props.get("accessible") is not False else "不可访问", self._fmt_percent(props.get("used_percent")), self._fmt(props.get("datastore_filesystem_type"))])
        self._add_table_v2(document, ["对象", "类型", "状态", "使用率或 CPU", "内存或格式"], rows or [["未采集", "未采集", "未采集", "未采集", "未采集"]], widths=[3000, 1700, 2200, 1600, 1450])

    def _chapter_appendix(self, document: Document, report: ReportData) -> None:
        self._add_heading_v2(document, "9. 附录", 1)
        self._add_heading_v2(document, "9.1 检查范围", 2)
        checklist = (report.appendix or {}).get("rule_checklist") or []
        self._add_paragraph_v2(document, f"本次纳入检查清单 {len(checklist)} 项；正文仅展示客户可读的检查结果和整改建议。")
        self._add_heading_v2(document, "9.2 未确认与未适用项", 2)
        unavailable = (report.appendix or {}).get("unavailable_rules") or []
        if unavailable:
            self._add_table_v2(document, ["检查项", "对象", "状态", "说明"], [[item.get("title") or "未命名检查", item.get("object_name") or "未记录", "未确认", item.get("reason") or "需进一步确认"] for item in unavailable], widths=[3000, 2600, 1500, 2850])
        else:
            self._add_paragraph_v2(document, "本次未记录需要单独列出的未确认项。")
        self._add_heading_v2(document, "9.3 工程数据说明", 2)
        self._add_paragraph_v2(document, "完整对象明细、原始证据和历史数据保留在同一次巡检生成的 HTML 报告及数据包中。Word 正文不展示内部规则标识和采集器字段。")

    def _problem_text(self, document: Document, report: ReportData, group: RiskGroup, category: dict[str, Any]) -> None:
        finding = group.findings[0]
        description = category.get("description") or category.get("description_zh") or finding.explanation_zh or finding.explanation or group.title
        impact = category.get("potential_impact") or category.get("impact_scope") or finding.business_impact or finding.technical_impact or "可能影响当前虚拟化环境的运行稳定性或后续运维效率。"
        remediation = category.get("remediation") or category.get("remediation_zh") or finding.recommended_action_zh or finding.remediation or "结合受影响对象和技术证据复核后，按维护窗口安排处理。"
        self._labeled_paragraph(document, "问题说明", description)
        self._labeled_paragraph(document, "影响范围", f"涉及 {len(group.findings)} 条对象明细，具体对象见下方明细表。")
        self._labeled_paragraph(document, "可能影响", impact)
        self._labeled_paragraph(document, "整改建议", remediation)
        mode = category.get("remediation_mode") or ("维护窗口" if finding.maintenance_window_required else "需结合场景评估")
        self._labeled_paragraph(document, "整改方式", mode)

    def _labeled_paragraph(self, document: Document, label: str, text: Any) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_together = True
        label_run = paragraph.add_run(f"{label}：")
        self._format_run(label_run, size=10.5, bold=True, color=BLACK)
        body_run = paragraph.add_run(self._clean(text).rstrip("。；;，,") + "。")
        self._format_run(body_run, size=10.5, color=BLACK)

    def _add_heading_v2(self, document: Document, text: str, level: int, *, page_break: bool = True) -> None:
        heading = self._clean(text)
        if text.startswith("1.2 ") and "巡检范围" in text:
            heading = heading.replace("检查对象", "巡检范围")
        paragraph = document.add_paragraph(heading, style=f"Heading {level}")
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        if level == 1 and page_break:
            paragraph.paragraph_format.page_break_before = True

    def _add_paragraph_v2(self, document: Document, text: Any) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run(self._clean(text))
        self._format_run(run, size=10.5, color=BLACK)

    def _add_bullets_v2(self, document: Document, items: list[str]) -> None:
        for item in items:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run(self._clean(item))
            self._format_run(run, size=10.5, color=BLACK)

    def _add_issue_summary_v2(self, document: Document, rows: list[list[Any]]) -> None:
        table = self._add_table_v2(
            document,
            ["等级", "分类", "问题描述", "影响数量", "建议动作"],
            rows,
            widths=[1300, 1850, 3450, 1300, 2050],
            align_center_columns={0, 3},
        )
        for row in table.rows[1:]:
            level = self._clean(row.cells[0].text)
            style = LEVEL_STYLE.get(level)
            if style:
                self._shade_cell(row.cells[0], style["fill"])
                self._color_cell_text(row.cells[0], style["font"], bold=True)
                self._shade_cell(row.cells[2], {"P1": "F4CCCC", "P2": "FCE4D6", "P3": "FFF2CC"}.get(level, LIGHT_GRAY))
                self._color_cell_text(row.cells[2], style["text"], bold=True)

    def _add_table_v2(
        self,
        document: Document,
        headers: list[str],
        rows: list[list[Any]],
        *,
        widths: list[float] | None = None,
        align_center_columns: set[int] | None = None,
        landscape: bool = False,
    ) -> Any:
        if landscape:
            self._start_landscape(document)
        total = LANDSCAPE_CONTENT_DXA if landscape else PORTRAIT_CONTENT_DXA
        normalized = self._normalize_widths_for_total(widths, headers, total)
        table = document.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        self._set_table_fixed_width(table, normalized, total)
        self._set_table_borders(table, BORDER_GRAY)
        self._repeat_header(table.rows[0])
        self._keep_row_together(table.rows[0])
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            self._set_cell_text_v2(cell, header, bold=True, align_center=True)
            self._set_cell_width(cell, normalized[index])
            self._shade_cell(cell, HEADER_BLUE)
            self._color_cell_text(cell, RGBColor(255, 255, 255), bold=True)
        align_center_columns = align_center_columns or set()
        for row_index, values in enumerate(rows):
            table_row = table.add_row()
            self._keep_row_together(table_row)
            for index, header in enumerate(headers):
                cell = table_row.cells[index]
                value = values[index] if index < len(values) else ""
                centered = index in align_center_columns or (index != 0 and not self._is_long_text_column(header))
                self._set_cell_text_v2(cell, value, align_center=centered)
                self._set_cell_width(cell, normalized[index])
                if row_index % 2 == 1:
                    self._shade_cell(cell, LIGHT_GRAY)
        document.add_paragraph().paragraph_format.space_after = Pt(1)
        if landscape:
            self._end_landscape(document)
        return table

    def _set_table_fixed_width(self, table: Any, widths: list[int], total: int) -> None:
        tbl_pr = table._tbl.tblPr
        layout = tbl_pr.first_child_found_in("w:tblLayout")
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
        tbl_w = tbl_pr.first_child_found_in("w:tblW")
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(total))
        tbl_w.set(qn("w:type"), "dxa")
        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for width in widths:
            col = OxmlElement("w:gridCol")
            col.set(qn("w:w"), str(int(width)))
            grid.append(col)

    def _set_table_borders(self, table: Any, color: str) -> None:
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.first_child_found_in("w:tblBorders")
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tbl_pr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = qn(f"w:{edge}")
            element = borders.find(tag)
            if element is None:
                element = OxmlElement(f"w:{edge}")
                borders.append(element)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), color)

    def _set_cell_text_v2(self, cell: Any, value: Any, *, bold: bool = False, align_center: bool = False) -> None:
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if align_center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        run = paragraph.add_run(self._clean(value))
        self._format_run(run, size=9 if bold else 9.2, bold=bold, color=BLACK)
        self._set_cell_margins(cell, top=80, start=100, bottom=80, end=100)

    def _set_cell_margins(self, cell: Any, *, top: int, start: int, bottom: int, end: int) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        margins = tc_pr.first_child_found_in("w:tcMar")
        if margins is None:
            margins = OxmlElement("w:tcMar")
            tc_pr.append(margins)
        for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
            node = margins.find(qn(f"w:{name}"))
            if node is None:
                node = OxmlElement(f"w:{name}")
                margins.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def _normalize_widths_for_total(self, widths: list[float] | None, headers: list[str], total: int) -> list[int]:
        if not headers:
            return []
        if not widths or len(widths) != len(headers):
            raw = [1] * len(headers)
        else:
            raw = [float(item) for item in widths]
        if sum(raw) <= 0:
            raw = [1] * len(headers)
        result = [max(1, int(round(item * total / sum(raw)))) for item in raw]
        result[-1] += total - sum(result)
        return result

    def _start_landscape(self, document: Document) -> None:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        self._configure_section(section, landscape=True)
        section.different_first_page_header_footer = False
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True

    def _end_landscape(self, document: Document) -> None:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        self._configure_section(section, landscape=False)
        section.different_first_page_header_footer = False
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True

    def _architecture_chart(self, document: Document, report: ReportData) -> None:
        scale = self._environment_scale(report)
        headers = ["vCenter", ">", "数据中心", ">", "Cluster", ">", "ESXi 主机", ">", "虚拟机"]
        values = [scale["vcenter"], ">", scale["datacenter"], ">", scale["cluster"], ">", scale["host"], ">", scale["vm"]]
        self._add_table_v2(document, headers, [values], widths=[1400, 450, 1450, 450, 1500, 450, 1550, 450, 3300], align_center_columns={0, 1, 2, 3, 4, 5, 6, 7, 8})
        self._add_paragraph_v2(document, f"存储关联对象：Datastore {scale['datastore']} 个；vSAN 集群 {scale['vsan_cluster']} 个。")

    def _risk_chart(self, document: Document, counts: dict[str, int]) -> None:
        maximum = max(counts.values(), default=0)
        rows = []
        for level in ("P1", "P2", "P3"):
            value = counts[level]
            bar = "=" * (max(1, round(value / maximum * 16)) if maximum else 0)
            rows.append([level, value, bar or "-"])
        table = self._add_table_v2(document, ["等级", "问题类别数", "分布"], rows, widths=[1600, 2300, 6050], align_center_columns={0, 1})
        for row in table.rows[1:]:
            level = self._clean(row.cells[0].text)
            style = LEVEL_STYLE.get(level)
            if style:
                self._shade_cell(row.cells[0], style["fill"])
                self._color_cell_text(row.cells[0], style["font"], bold=True)

    def _resource_chart(self, document: Document, rows: list[list[Any]]) -> None:
        self._add_table_v2(document, ["资源项", "指标", "使用率", "概览"], rows, widths=[2600, 2700, 1800, 2850], align_center_columns={2})

    def _wide_headers(self, headers: list[str]) -> bool:
        """Use landscape for dense tables before reducing the readable font size."""
        if len(headers) >= 7:
            return True
        long_columns = sum(1 for header in headers if self._is_long_text_column(header))
        return len(headers) >= 6 and long_columns >= 2

    def _resource_rows(self, report: ReportData) -> list[list[Any]]:
        rows: list[list[Any]] = []
        hosts = self._asset_details(report, "HostSystem")
        cpu = [self._safe_float((item.get("properties") or {}).get("cpu_usage_avg")) for item in hosts]
        cpu = [value for value in cpu if value is not None]
        memory = [self._safe_float((item.get("properties") or {}).get("memory_usage_avg")) for item in hosts]
        memory = [value for value in memory if value is not None]
        if cpu:
            average = sum(cpu) / len(cpu)
            rows.append(["ESXi CPU", "主机平均使用率", f"{average:.1f}%", "=" * max(1, round(average / 100 * 16))])
        if memory:
            average = sum(memory) / len(memory)
            rows.append(["ESXi 内存", "主机平均使用率", f"{average:.1f}%", "=" * max(1, round(average / 100 * 16))])
        capacity = ((report.vsan_summary or {}).get("capacity") or {})
        used = self._safe_float(capacity.get("used_percent"))
        if used is not None:
            rows.append(["vSAN 容量", "集群使用率", f"{used:.1f}%", "=" * max(1, round(used / 100 * 16))])
        return rows

    def _environment_scale(self, report: ReportData) -> dict[str, int]:
        presentation = report.report_presentation or {}
        scale = presentation.get("environment_scale") or {}
        summary = self._asset_summary(report)
        return {
            "vcenter": int(scale.get("vcenter", summary.get("vCenter", 0)) or 0),
            "datacenter": int(scale.get("datacenter", len(self._unique_datacenters(report))) or 0),
            "cluster": int(scale.get("cluster", summary.get("ClusterComputeResource", 0)) or 0),
            "host": int(scale.get("host", summary.get("HostSystem", 0)) or 0),
            "vm": int(scale.get("vm", summary.get("VirtualMachine", 0)) or 0),
            "datastore": int(scale.get("datastore", summary.get("Datastore", 0)) or 0),
            "vsan_cluster": int(scale.get("vsan_cluster", 0) or 0),
        }

    def _display_risk_groups(self, report: ReportData) -> list[RiskGroup]:
        groups = self._customer_visible_risk_groups(self._aggregate_risks(report))
        if groups or not report.problem_categories:
            return groups

        # Real customer payloads may contain presentation categories and
        # compact findings without rule/category identifiers. Build a Word-
        # only view from the already-approved customer-facing category data.
        result: list[RiskGroup] = []
        for category in report.problem_categories:
            level = str(category.get("priority") or category.get("risk_level") or "")
            if level not in {"P1", "P2", "P3"}:
                continue
            objects = category.get("affected_objects") or []
            if not objects:
                objects = [{"object_name": name, "object_type": "", "object_path": ""} for name in category.get("sample_objects", [])]
            if not objects:
                objects = [{"object_name": "未记录", "object_type": "", "object_path": ""}]
            findings = [
                FindingItem(
                    risk_level=level,
                    rule_id=str(category.get("category_id") or ""),
                    rule_name=str(category.get("title") or category.get("category_id") or ""),
                    title=str(category.get("title") or category.get("category_id") or ""),
                    object_type=str(item.get("object_type") or ""),
                    object_name=str(item.get("object_name") or "未记录"),
                    object_path=str(item.get("object_path") or ""),
                    status="open",
                    business_impact=str(category.get("potential_impact") or category.get("impact_scope") or ""),
                    remediation=str(category.get("remediation") or ""),
                    evidence_summary_zh="；".join(str(value) for value in category.get("issue_summaries", []) if value),
                    explanation_zh=str(category.get("description") or category.get("summary") or ""),
                    recommended_action_zh=str(category.get("remediation") or ""),
                    evidence={"items": category.get("evidence_items") or []},
                )
                for item in objects
                if isinstance(item, dict)
            ]
            if findings:
                result.append(RiskGroup(
                    level=level,
                    category=self._classify_finding(findings[0]),
                    title=self._clean(category.get("title") or category.get("category_id") or "建议关注事项"),
                    findings=findings,
                    rule_id=str(category.get("category_id") or ""),
                    metadata=dict(category),
                ))
        return sorted(result, key=self._risk_group_sort_key)

    def _summary_rows_v2(self, report: ReportData, groups: list[RiskGroup]) -> list[list[str]]:
        rows = self._summary_rows(groups)
        category_by_title = {self._clean(group.title): group.category for group in groups}
        for row in rows:
            if len(row) >= 3 and not self._clean(row[1]):
                row[1] = category_by_title.get(self._clean(row[2]), "综合")
        return rows

    def _remove_style_borders(self, style: Any) -> None:
        p_pr = style._element.pPr
        if p_pr is None:
            return
        borders = p_pr.find(qn("w:pBdr"))
        if borders is not None:
            p_pr.remove(borders)

    def _risk_counts_v2(self, report: ReportData, groups: list[RiskGroup]) -> dict[str, int]:
        source = (report.report_presentation or {}).get("risk_summary") or {}
        return {level: int(source.get(level, getattr(report.risk_summary, level, 0)) or 0) for level in ("P1", "P2", "P3")}

    def _affected_counts_v2(self, report: ReportData) -> dict[str, int]:
        source = (report.report_presentation or {}).get("affected_object_summary") or report.affected_object_summary or {}
        return {level: int(source.get(level, 0) or 0) for level in ("P1", "P2", "P3")}

    def _normal_summary(self, report: ReportData) -> list[str]:
        normal = []
        tree = (report.report_presentation or {}).get("normal_check_tree", [])
        if isinstance(tree, dict):
            tree = [
                {"module_id": key, "label": key, "checks": value if isinstance(value, list) else []}
                for key, value in tree.items()
            ]
        for item in tree:
            if not isinstance(item, dict):
                continue
            checks = item.get("checks") or []
            if checks:
                normal.append(f"{item.get('label') or item.get('module_id')}: 已验证 {len(checks)} 类检查项正常。")
        return normal[:8]

    def _overall_status(self, report: ReportData) -> str:
        return self._clean((report.report_presentation or {}).get("overall_status") or report.overall_status or report.health_score.label or "评估受限")

    def _health_status_short(self, value: str) -> str:
        text = value.casefold()
        if any(token in text for token in ("风险", "异常", "需处理", "problem", "risk")):
            return "需要处理"
        if any(token in text for token in ("关注", "limited", "受限")):
            return "需要关注"
        return "状态正常"

    def _health_explanation(self, report: ReportData) -> str:
        return self._clean(report.health_score.explanation or "请结合本次检查结果和影响对象继续确认。")

    def _status_fill(self, value: str) -> str:
        short = self._health_status_short(value)
        return {"需要处理": "F4CCCC", "需要关注": "FCE4D6", "状态正常": "E2F0D9"}.get(short, LIGHT_BLUE)

    def _status_font(self, value: str) -> RGBColor:
        short = self._health_status_short(value)
        return {"需要处理": RGBColor(156, 0, 6), "需要关注": RGBColor(156, 87, 0), "状态正常": RGBColor(0, 97, 0)}.get(short, BLACK)

    def _history_visible(self, report: ReportData) -> bool:
        return bool(report.history_visible or ((report.appendix or {}).get("history_comparison") or {}).get("state") == "ready")

    def _date_time_label(self, value: Any) -> str:
        text = self._clean(value)
        return text.replace("T", " ").split("+", 1)[0] if text else "未采集"


__all__ = ["DocxReportV2Engine"]
