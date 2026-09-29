from __future__ import annotations

import html
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


LOG_CHECK_REPORT_TITLE = "VStackLens 日志粗排查报告"
MODEL_ASSISTED_DIAGNOSIS_TITLE = "VStackLens 模型辅助日志诊断报告"
MODEL_ASSISTED_TRIAGE_TITLE = "VStackLens 模型辅助粗排查报告"

REPORT_CSS = """
:root {
  color-scheme: light;
  --bg: #f5f7fb;
  --card: #ffffff;
  --text: #1f2937;
  --muted: #64748b;
  --line: #dbe3ef;
  --blue: #1d4ed8;
  --red: #c00000;
  --orange: #d97706;
  --green: #15803d;
}
* { box-sizing: border-box; }
body {
  margin: 0;
  background: var(--bg);
  color: var(--text);
  font-family: "Microsoft YaHei", "Segoe UI", Arial, sans-serif;
  line-height: 1.65;
}
.page { max-width: 1180px; margin: 0 auto; padding: 28px; }
.hero {
  background: linear-gradient(135deg, #123a6f, #2563eb);
  color: white;
  padding: 28px 32px;
  border-radius: 8px;
}
.hero h1 { margin: 0 0 12px; font-size: 28px; }
.hero p { margin: 4px 0; color: #dbeafe; }
.grid { display: grid; grid-template-columns: repeat(4, minmax(0, 1fr)); gap: 12px; margin: 18px 0; }
.metric, section {
  background: var(--card);
  border: 1px solid var(--line);
  border-radius: 8px;
  padding: 18px;
}
.metric strong { display: block; font-size: 24px; color: var(--blue); }
.metric span { color: var(--muted); }
section { margin: 16px 0; }
h2 { margin: 0 0 14px; font-size: 20px; }
h3 { margin: 18px 0 8px; font-size: 16px; }
table { width: 100%; border-collapse: collapse; margin: 10px 0; table-layout: fixed; }
th, td { border: 1px solid var(--line); padding: 8px 10px; vertical-align: top; word-break: break-word; }
th { background: #eaf0f8; text-align: left; }
.badge { display: inline-block; padding: 2px 8px; border-radius: 999px; color: white; font-size: 12px; }
.high { background: var(--red); }
.medium { background: var(--orange); }
.low { background: var(--green); }
.muted { color: var(--muted); }
.diagnosis-grid { grid-template-columns: repeat(3, minmax(0, 1fr)); }
.diagnosis-card strong { font-size: 14px; color: var(--muted); }
.diagnosis-card p { margin: 6px 0 0; }
.evidence-block { background: #f8fafc; border: 1px solid var(--line); border-radius: 6px; padding: 10px 12px; margin: 8px 0; }
.evidence-block p { margin: 4px 0; }
.evidence-excerpt { font-family: Consolas, "Courier New", monospace; white-space: pre-wrap; word-break: break-word; }
.numbered-list p { margin: 6px 0; }
.chain-table { margin: 0; }
.chain-table th:first-child, .chain-table td:first-child { width: 96px; }
@media print {
  body { background: white; }
  .page { max-width: none; padding: 0; }
  .hero, .metric, section { border-radius: 0; }
}
"""


def render_log_analysis_html(payload: dict[str, Any], report_dir: Path) -> Path:
    report_dir.mkdir(parents=True, exist_ok=True)
    (report_dir / "assets").mkdir(exist_ok=True)
    (report_dir / "assets" / "log_analysis.css").write_text(REPORT_CSS.strip(), encoding="utf-8")
    index_path = report_dir / "index.html"
    index_path.write_text(_html_document(payload), encoding="utf-8")
    return index_path


def render_log_analysis_docx(payload: dict[str, Any], docx_path: Path) -> Path:
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _configure_document(doc)
    if _model_assisted_diagnosis(payload) and _snapshot_diagnosis(payload):
        return _render_snapshot_diagnosis_docx(payload, docx_path, doc)
    if _model_assisted_diagnosis(payload) and _problem_diagnosis(payload):
        return _render_problem_diagnosis_docx(payload, docx_path, doc)
    if _model_assisted_triage(payload) and _problem_diagnosis(payload):
        return _render_model_triage_docx(payload, docx_path, doc)
    return _render_log_check_docx(payload, docx_path, doc)


def _html_document(payload: dict[str, Any]) -> str:
    if _model_assisted_diagnosis(payload) and _snapshot_diagnosis(payload):
        return _snapshot_html_document(payload)
    if _model_assisted_diagnosis(payload) and _problem_diagnosis(payload):
        return _problem_html_document(payload)
    if _model_assisted_triage(payload) and _problem_diagnosis(payload):
        return _model_triage_html_document(payload)
    return _log_check_html_document(payload)


def _render_log_check_docx(payload: dict[str, Any], docx_path: Path, doc: Document) -> Path:
    meta = payload.get("metadata", {}) if isinstance(payload.get("metadata"), dict) else {}
    summary = payload.get("summary", {}) if isinstance(payload.get("summary"), dict) else {}
    findings = payload.get("findings", []) if isinstance(payload.get("findings"), list) else []
    files = payload.get("files", []) if isinstance(payload.get("files"), list) else []
    diagnosis = payload.get("diagnosis", {}) if isinstance(payload.get("diagnosis"), dict) else {}
    engine = diagnosis.get("diagnosis_engine") if isinstance(diagnosis.get("diagnosis_engine"), dict) else {}
    customer = str(meta.get("customer_name") or "未指定客户")
    generated_at = str(meta.get("generated_at") or "")
    bundle_name = str(meta.get("support_bundle_name") or "")
    problem = str(meta.get("problem_description") or diagnosis.get("customer_question") or "未记录")
    stats = _log_check_stats(summary, diagnosis, findings)
    high_clues, general_clues, observations = _log_check_clue_groups(payload)
    recommendations = _log_check_recommendations(payload)
    missing = _log_check_missing_materials(payload)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(customer)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(18)
    run.font.bold = True

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_paragraph.add_run(LOG_CHECK_REPORT_TITLE)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

    for label, value in (
        ("排查时间", generated_at),
        ("日志包", bundle_name),
        ("客户问题", problem),
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"{label}：").bold = True
        paragraph.add_run(value or "未记录")

    doc.add_page_break()

    doc.add_heading("1. 排查概览", level=1)
    doc.add_heading("1.1 客户问题描述", level=2)
    doc.add_paragraph(problem)
    doc.add_heading("1.2 日志包处理结果", level=2)
    overview = doc.add_table(rows=1, cols=4)
    _style_table(overview)
    _fill_header(overview.rows[0].cells, ["日志包文件", "识别日志文件", "关键线索", "建议核对项"])
    row = overview.add_row().cells
    row[0].text = str(stats["total_files"])
    row[1].text = str(stats["log_files"])
    row[2].text = str(stats["key_evidence"])
    row[3].text = str(len(recommendations))
    doc.add_heading("1.3 当前排查范围", level=2)
    doc.add_paragraph("本次仅基于规则提取日志线索，供后续排查参考。日志分析过程不连接 vCenter，不要求输入 vCenter 用户名或密码，也不写入虚拟化健康巡检结果。")
    doc.add_paragraph("当前模式不进行深度原因分析。建议结合故障时间、对象名称、vCenter 任务事件进一步核对。")

    doc.add_heading("2. 相关日志线索", level=1)
    doc.add_heading("2.1 高相关线索", level=2)
    _add_log_check_clue_table_docx(doc, high_clues, "未提取到高相关日志线索。")
    doc.add_heading("2.2 一般相关线索", level=2)
    _add_log_check_clue_table_docx(doc, general_clues, "未提取到一般相关日志线索。")
    if observations:
        doc.add_heading("2.3 其他观察项", level=2)
        table = doc.add_table(rows=1, cols=3)
        _style_table(table)
        _fill_header(table.rows[0].cells, ["分类", "观察项", "建议"])
        for item in observations[:12]:
            cells = table.add_row().cells
            cells[0].text = str(item.get("category") or "日志")
            cells[1].text = str(item.get("title") or item.get("description") or "日志观察项")
            cells[2].text = str(item.get("recommendation") or "建议结合业务影响进一步核查。")

    doc.add_heading("3. 初步排查方向", level=1)
    doc.add_heading("3.1 建议优先核对", level=2)
    for index, item in enumerate(recommendations, start=1):
        doc.add_paragraph(f"3.1.{index} {item}")
    doc.add_heading("3.2 需要补充的信息", level=2)
    for index, item in enumerate(missing, start=1):
        doc.add_paragraph(f"3.2.{index} {item}")

    doc.add_heading("4. 附录", level=1)
    doc.add_heading("4.1 日志文件统计", level=2)
    file_table = doc.add_table(rows=1, cols=3)
    _style_table(file_table)
    _fill_header(file_table.rows[0].cells, ["文件路径", "大小", "识别类型"])
    for item in files[:80]:
        if not isinstance(item, dict):
            continue
        cells = file_table.add_row().cells
        cells[0].text = str(item.get("path") or "")
        cells[1].text = str(item.get("size_label") or "")
        cells[2].text = "日志文件" if item.get("is_log") else "其他文件"
    if len(files) > 80:
        doc.add_paragraph("文件数量较多，Word 报告仅展示前 80 个文件。")
    doc.add_heading("4.2 处理说明", level=2)
    notes = [
        "本报告展示规则筛选后的日志线索，不上传完整 support bundle，不展示完整原始日志正文。",
        "建议结合故障时间、对象名称、vCenter 任务事件进一步核对。",
    ]
    fallback_reason = str(engine.get("fallback_reason") or "").strip()
    if fallback_reason:
        notes.append("云端辅助分析未启用成功，本报告仅提供日志粗排查线索。")
    for index, item in enumerate(notes, start=1):
        doc.add_paragraph(f"4.2.{index} {item}")

    _set_document_font(doc)
    doc.save(docx_path)
    return docx_path


def _log_check_html_document(payload: dict[str, Any]) -> str:
    meta = payload.get("metadata", {}) if isinstance(payload.get("metadata"), dict) else {}
    summary = payload.get("summary", {}) if isinstance(payload.get("summary"), dict) else {}
    findings = payload.get("findings", []) if isinstance(payload.get("findings"), list) else []
    files = payload.get("files", []) if isinstance(payload.get("files"), list) else []
    diagnosis = payload.get("diagnosis", {}) if isinstance(payload.get("diagnosis"), dict) else {}
    engine = diagnosis.get("diagnosis_engine") if isinstance(diagnosis.get("diagnosis_engine"), dict) else {}
    problem = str(meta.get("problem_description") or diagnosis.get("customer_question") or "未记录")
    stats = _log_check_stats(summary, diagnosis, findings)
    high_clues, general_clues, observations = _log_check_clue_groups(payload)
    recommendations = _log_check_recommendations(payload)
    missing = _log_check_missing_materials(payload)

    high_html = _log_check_clue_table_html(high_clues, "未提取到高相关日志线索。")
    general_html = _log_check_clue_table_html(general_clues, "未提取到一般相关日志线索。")
    observations_html = ""
    if observations:
        rows = "\n".join(
            "<tr>"
            f"<td>{_esc(item.get('category') if isinstance(item, dict) else '日志')}</td>"
            f"<td>{_esc((item.get('title') or item.get('description')) if isinstance(item, dict) else item)}</td>"
            f"<td>{_esc(item.get('recommendation') if isinstance(item, dict) else '建议结合业务影响进一步核查。')}</td>"
            "</tr>"
            for item in observations[:12]
        )
        observations_html = f"""
      <h3>2.3 其他观察项</h3>
      <table><thead><tr><th>分类</th><th>观察项</th><th>建议</th></tr></thead><tbody>{rows}</tbody></table>"""
    recommendation_html = "".join(f"<p><strong>3.1.{index}</strong> {_esc(item)}</p>" for index, item in enumerate(recommendations, start=1))
    missing_html = "".join(f"<p><strong>3.2.{index}</strong> {_esc(item)}</p>" for index, item in enumerate(missing, start=1))
    file_rows = "\n".join(
        "<tr>"
        f"<td>{_esc(item.get('path'))}</td>"
        f"<td>{_esc(item.get('size_label'))}</td>"
        f"<td>{'日志文件' if item.get('is_log') else '其他文件'}</td>"
        "</tr>"
        for item in files[:120]
        if isinstance(item, dict)
    ) or "<tr><td colspan=\"3\">未读取到文件清单。</td></tr>"
    overflow_note = "<p class=\"muted\">文件数量较多，HTML 报告仅展示前 120 个文件。</p>" if len(files) > 120 else ""
    processing_notes = [
        "本报告展示规则筛选后的日志线索，不上传完整 support bundle，不展示完整原始日志正文。",
        "建议结合故障时间、对象名称、vCenter 任务事件进一步核对。",
    ]
    fallback_reason = str(engine.get("fallback_reason") or "").strip()
    if fallback_reason:
        processing_notes.append("云端辅助分析未启用成功，本报告仅提供日志粗排查线索。")
    processing_html = "".join(f"<p><strong>4.2.{index}</strong> {_esc(item)}</p>" for index, item in enumerate(processing_notes, start=1))

    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{_esc(LOG_CHECK_REPORT_TITLE)}</title>
  <link rel="stylesheet" href="assets/log_analysis.css">
</head>
<body>
  <main class="page">
    <div class="hero">
      <h1>{_esc(LOG_CHECK_REPORT_TITLE)}</h1>
      <p>客户名称：{_esc(meta.get("customer_name") or "未指定客户")}</p>
      <p>日志包：{_esc(meta.get("support_bundle_name") or "未记录")}</p>
      <p>排查时间：{_esc(meta.get("generated_at") or "")}</p>
      <p>客户问题：{_esc(problem)}</p>
    </div>
    <div class="grid">
      <div class="metric"><strong>{stats["total_files"]}</strong><span>日志包文件</span></div>
      <div class="metric"><strong>{stats["log_files"]}</strong><span>识别日志文件</span></div>
      <div class="metric"><strong>{stats["key_evidence"]}</strong><span>关键线索</span></div>
      <div class="metric"><strong>{len(recommendations)}</strong><span>建议核对项</span></div>
    </div>
    <section>
      <h2>1. 排查概览</h2>
      <h3>1.1 客户问题描述</h3>
      <p>{_esc(problem)}</p>
      <h3>1.2 日志包处理结果</h3>
      <table><thead><tr><th>日志包文件</th><th>识别日志文件</th><th>关键线索</th><th>建议核对项</th></tr></thead><tbody>
        <tr><td>{stats["total_files"]}</td><td>{stats["log_files"]}</td><td>{stats["key_evidence"]}</td><td>{len(recommendations)}</td></tr>
      </tbody></table>
      <h3>1.3 当前排查范围</h3>
      <p>本次仅基于规则提取日志线索，供后续排查参考。日志分析过程不连接 vCenter，不要求输入 vCenter 用户名或密码，也不写入虚拟化健康巡检结果。</p>
      <p>当前模式不进行深度原因分析。建议结合故障时间、对象名称、vCenter 任务事件进一步核对。</p>
    </section>
    <section>
      <h2>2. 相关日志线索</h2>
      <h3>2.1 高相关线索</h3>
      {high_html}
      <h3>2.2 一般相关线索</h3>
      {general_html}
      {observations_html}
    </section>
    <section>
      <h2>3. 初步排查方向</h2>
      <h3>3.1 建议优先核对</h3>
      <div class="numbered-list">{recommendation_html}</div>
      <h3>3.2 需要补充的信息</h3>
      <div class="numbered-list">{missing_html}</div>
    </section>
    <section>
      <h2>4. 附录</h2>
      <h3>4.1 日志文件统计</h3>
      <table><thead><tr><th>文件路径</th><th>大小</th><th>识别类型</th></tr></thead><tbody>{file_rows}</tbody></table>
      {overflow_note}
      <h3>4.2 处理说明</h3>
      <div class="numbered-list">{processing_html}</div>
    </section>
  </main>
</body>
</html>
"""


def _log_check_stats(summary: dict[str, Any], diagnosis: dict[str, Any], findings: list[dict[str, Any]]) -> dict[str, int]:
    evidence_chain = diagnosis.get("evidence_chain") if isinstance(diagnosis.get("evidence_chain"), list) else []
    evidence_sections = diagnosis.get("evidence_sections") if isinstance(diagnosis.get("evidence_sections"), list) else []
    section_excerpt_count = 0
    evidence_files: set[str] = set()
    for item in evidence_chain:
        if isinstance(item, dict):
            file_path = str(item.get("file") or "").strip()
            if file_path and "未提取" not in file_path and "未明确" not in file_path:
                evidence_files.add(file_path)
    for section in evidence_sections:
        if isinstance(section, dict) and isinstance(section.get("excerpts"), list):
            section_excerpt_count += len(section["excerpts"])
            section_file = str(section.get("log_file") or "").strip()
            if section_file and "未提取" not in section_file and "未明确" not in section_file:
                evidence_files.add(section_file)
            for excerpt in section["excerpts"]:
                if isinstance(excerpt, dict):
                    file_path = str(excerpt.get("file") or "").strip()
                    if file_path and "未提取" not in file_path and "未明确" not in file_path:
                        evidence_files.add(file_path)
    key_evidence = _safe_int(diagnosis.get("key_evidence_count")) or len(evidence_chain) or section_excerpt_count or _safe_int(summary.get("evidence_count"))
    return {
        "total_files": _safe_int(summary.get("total_files")),
        "log_files": _safe_int(summary.get("log_files") or summary.get("log_file_count")) or len(evidence_files),
        "key_evidence": key_evidence,
        "findings": len(findings),
    }


def _safe_int(value: Any) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _log_check_clue_groups(payload: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    diagnosis = payload.get("diagnosis", {}) if isinstance(payload.get("diagnosis"), dict) else {}
    findings = payload.get("findings", []) if isinstance(payload.get("findings"), list) else []
    rows: list[dict[str, Any]] = []
    for item in diagnosis.get("evidence_chain") or []:
        if not isinstance(item, dict):
            continue
        rows.append(
            {
                "file": item.get("file"),
                "location": item.get("location"),
                "message": item.get("message") or item.get("event"),
                "relevance": item.get("relevance") or item.get("strength") or item.get("level") or "与客户问题相关",
                "strength": item.get("strength") or item.get("level") or "提示",
            }
        )
    if not rows:
        for finding in findings:
            if not isinstance(finding, dict):
                continue
            for evidence in finding.get("evidence") or []:
                if not isinstance(evidence, dict):
                    continue
                rows.append(
                    {
                        "file": evidence.get("file"),
                        "location": evidence.get("location"),
                        "message": evidence.get("message"),
                        "relevance": finding.get("title") or finding.get("description"),
                        "strength": finding.get("severity") or "提示",
                    }
                )
    high: list[dict[str, Any]] = []
    general: list[dict[str, Any]] = []
    for row in rows:
        strength = str(row.get("strength") or "")
        if any(token in strength for token in ("高", "中", "strong", "直接")):
            high.append(row)
        else:
            general.append(row)
    if not high and rows:
        high = rows[: min(8, len(rows))]
        general = rows[len(high):]
    observations = [
        item
        for item in findings
        if isinstance(item, dict) and not item.get("problem_related")
    ][:8]
    return high[:20], general[:20], observations


def _add_log_check_clue_table_docx(doc: Document, clues: list[dict[str, Any]], empty_text: str) -> None:
    if not clues:
        doc.add_paragraph(empty_text)
        return
    table = doc.add_table(rows=1, cols=4)
    _style_table(table)
    _fill_header(table.rows[0].cells, ["日志文件", "位置", "线索摘录", "关联说明"])
    for item in clues[:12]:
        cells = table.add_row().cells
        cells[0].text = str(item.get("file") or "")
        cells[1].text = str(item.get("location") or "")
        cells[2].text = str(item.get("message") or "")
        cells[3].text = str(item.get("relevance") or "")
    if len(clues) > 12:
        doc.add_paragraph("更多日志线索建议结合原始日志继续查看。")


def _log_check_clue_table_html(clues: list[dict[str, Any]], empty_text: str) -> str:
    if not clues:
        return f"<p class=\"muted\">{_esc(empty_text)}</p>"
    rows = "\n".join(
        "<tr>"
        f"<td>{_esc(item.get('file'))}</td>"
        f"<td>{_esc(item.get('location'))}</td>"
        f"<td>{_esc(item.get('message'))}</td>"
        f"<td>{_esc(item.get('relevance'))}</td>"
        "</tr>"
        for item in clues[:20]
    )
    overflow = "<p class=\"muted\">更多日志线索建议结合原始日志继续查看。</p>" if len(clues) > 20 else ""
    return (
        "<table><thead><tr><th>日志文件</th><th>位置</th><th>线索摘录</th><th>关联说明</th></tr></thead>"
        f"<tbody>{rows}</tbody></table>{overflow}"
    )


def _log_check_recommendations(payload: dict[str, Any]) -> list[str]:
    diagnosis = payload.get("diagnosis", {}) if isinstance(payload.get("diagnosis"), dict) else {}
    values = _string_list(diagnosis.get("recommended_next_steps") or diagnosis.get("handling_recommendations"))
    if not values:
        values = _string_list(payload.get("recommendations"))
    if not values:
        values = [
            "建议结合故障发生时间窗口复核相关 vCenter Task/Event 记录。",
            "建议根据问题对象名称在原始日志中继续缩小检索范围。",
            "建议保留原始 support bundle，避免后续导出导致时间窗口偏移。",
        ]
    return values[:8]


def _log_check_missing_materials(payload: dict[str, Any]) -> list[str]:
    diagnosis = payload.get("diagnosis", {}) if isinstance(payload.get("diagnosis"), dict) else {}
    values = _string_list(diagnosis.get("missing_materials") or diagnosis.get("supplemental_info"))
    if not values:
        values = [
            "故障发生的准确时间点和持续时间。",
            "受影响对象名称、业务影响范围和现场截图。",
            "相关 vCenter Task/Event 记录或变更记录。",
        ]
    return values[:10]


def _string_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, tuple | set):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, str) and value.strip():
        return [value.strip()]
    return []


def _snapshot_diagnosis(payload: dict[str, Any]) -> dict[str, Any] | None:
    diagnosis = payload.get("diagnosis")
    if isinstance(diagnosis, dict) and diagnosis.get("scenario") == "snapshot_consolidation_failure":
        return diagnosis
    return None


def _problem_diagnosis(payload: dict[str, Any]) -> dict[str, Any] | None:
    diagnosis = payload.get("diagnosis")
    if isinstance(diagnosis, dict) and diagnosis.get("scenario"):
        return diagnosis
    return None


def _model_used(payload: dict[str, Any]) -> bool:
    diagnosis = payload.get("diagnosis")
    if not isinstance(diagnosis, dict):
        return False
    engine = diagnosis.get("diagnosis_engine")
    return isinstance(engine, dict) and engine.get("model_used") is True


def _model_quality(payload: dict[str, Any]) -> str:
    diagnosis = payload.get("diagnosis")
    if not isinstance(diagnosis, dict):
        return ""
    engine = diagnosis.get("diagnosis_engine")
    if not isinstance(engine, dict) or engine.get("model_used") is not True:
        return ""
    return str(engine.get("model_quality") or "").strip().lower()


def _model_assisted_diagnosis(payload: dict[str, Any]) -> bool:
    return _model_quality(payload) == "accepted"


def _model_assisted_triage(payload: dict[str, Any]) -> bool:
    return _model_quality(payload) == "downgraded"


def _list_values(items: Any, fallback: str = "当前日志包未提取到对应信息。") -> list[str]:
    values = [str(item).strip() for item in items if str(item).strip()] if isinstance(items, list) else []
    return values or [fallback]


def _join_values(items: Any, fallback: str = "未明确") -> str:
    values = _list_values(items, fallback="")
    values = [item for item in values if item]
    return " / ".join(values) if values else fallback


def _diagnosis_current_judgement(diagnosis: dict[str, Any]) -> str:
    return str(
        diagnosis.get("current_judgement")
        or diagnosis.get("accurate_judgement")
        or diagnosis.get("conclusion")
        or "当前证据不足以单独确认完整原因链。"
    )


def _diagnosis_root_cause_text(diagnosis: dict[str, Any]) -> str:
    return str(diagnosis.get("fully_determine_text") or "否，当前 support bundle 尚不能单独证明最终根因。")


def _log_analysis_independence_text() -> str:
    return (
        "本报告基于客户导入的 VMware support bundle 日志包生成。日志分析过程不连接 vCenter，"
        "不要求输入 vCenter 用户名或密码，也不写入虚拟化健康巡检结果。"
    )


def _key_basis_items(diagnosis: dict[str, Any]) -> list[str]:
    items: list[str] = []
    for section in diagnosis.get("evidence_sections") or []:
        if not isinstance(section, dict):
            continue
        title = str(section.get("title") or "").strip()
        judgement = str(section.get("judgement") or "").strip()
        if title and judgement:
            items.append(f"{title}：{judgement}")
        elif judgement:
            items.append(judgement)
        if len(items) >= 4:
            break
    if not items:
        timeline = _timeline_items(diagnosis)
        for item in timeline[:4]:
            event = str(item.get("event") or "").strip()
            relevance = str(item.get("relevance") or "").strip()
            if event and relevance:
                items.append(f"{event}：{relevance}")
            elif event:
                items.append(event)
    return items or ["当前日志包尚未提取到足以闭环的关键依据，需要结合补充材料继续确认。"]


def _numbered_items_html(prefix: str, items: Any, fallback: str = "当前日志包未提取到对应信息。") -> str:
    values = _list_values(items, fallback=fallback)
    rows = "".join(f"<p><strong>{prefix}.{index}</strong> {_esc(item)}</p>" for index, item in enumerate(values, start=1))
    return f"<div class=\"numbered-list\">{rows}</div>"


def _add_numbered_docx_items(doc: Document, prefix: str, items: Any, fallback: str = "当前日志包未提取到对应信息。") -> None:
    for index, item in enumerate(_list_values(items, fallback=fallback), start=1):
        doc.add_paragraph(f"{prefix}.{index} {item}")


def _evidence_judgement_items(diagnosis: dict[str, Any]) -> list[str]:
    values: list[str] = []
    for section in diagnosis.get("evidence_sections") or []:
        if not isinstance(section, dict):
            continue
        title = str(section.get("title") or "日志证据").strip()
        judgement = str(section.get("judgement") or "").strip()
        if judgement:
            values.append(f"{title}：{judgement}")
    return values


def _diagnosis_problem_text(diagnosis: dict[str, Any], meta: dict[str, Any]) -> str:
    return str(diagnosis.get("customer_question") or meta.get("problem_description") or "未记录")


def _diagnosis_domains_text(affected: dict[str, Any], fallback: str = "未明确") -> str:
    return _join_values(affected.get("domains"), fallback=fallback)


def _add_evidence_chain_docx(doc: Document, evidence_chain: Any, max_rows: int = 20) -> None:
    rows = [item for item in evidence_chain if isinstance(item, dict)] if isinstance(evidence_chain, list) else []
    if not rows:
        doc.add_paragraph("当前日志包未提取到足以直接支撑该现象的关键日志摘录。")
        return
    table = doc.add_table(rows=1, cols=5)
    _style_table(table)
    _fill_header(table.rows[0].cells, ["证据强度", "日志文件", "位置", "日志摘录", "关联说明"])
    for item in rows[:max_rows]:
        cells = table.add_row().cells
        cells[0].text = str(item.get("strength") or item.get("level") or "")
        cells[1].text = str(item.get("file") or "")
        cells[2].text = str(item.get("location") or "")
        cells[3].text = str(item.get("message") or "")
        cells[4].text = str(item.get("relevance") or "")
    if len(rows) > max_rows:
        doc.add_paragraph("更多相关证据建议在 HTML 日志分析报告中查看。")


def _evidence_chain_table_html(evidence_chain: Any, max_rows: int = 30) -> str:
    rows = [item for item in evidence_chain if isinstance(item, dict)] if isinstance(evidence_chain, list) else []
    body = "\n".join(
        "<tr>"
        f"<td>{_esc(item.get('strength') or item.get('level') or '')}</td>"
        f"<td>{_esc(item.get('file') or '')}</td>"
        f"<td>{_esc(item.get('location') or '')}</td>"
        f"<td>{_esc(item.get('message') or '')}</td>"
        f"<td>{_esc(item.get('relevance') or '')}</td>"
        "</tr>"
        for item in rows[:max_rows]
    )
    if not body:
        body = "<tr><td colspan=\"5\">当前日志包未提取到足以直接支撑该现象的关键日志摘录。</td></tr>"
    overflow = "<p class=\"muted\">更多相关证据建议结合原始日志继续查看。</p>" if len(rows) > max_rows else ""
    return (
        "<table><thead><tr><th>证据强度</th><th>日志文件</th><th>位置</th><th>日志摘录</th><th>关联说明</th></tr></thead>"
        f"<tbody>{body}</tbody></table>{overflow}"
    )


def _add_evidence_explanations_docx(doc: Document, diagnosis: dict[str, Any], prefix: str = "3.3") -> None:
    _add_numbered_docx_items(
        doc,
        prefix,
        _evidence_judgement_items(diagnosis),
        fallback="当前日志包尚未形成完整证据链，需要结合补充材料继续确认。",
    )


def _evidence_explanations_html(diagnosis: dict[str, Any], prefix: str = "3.3") -> str:
    return _numbered_items_html(
        prefix,
        _evidence_judgement_items(diagnosis),
        fallback="当前日志包尚未形成完整证据链，需要结合补充材料继续确认。",
    )


def _snapshot_evidence_chain(diagnosis: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for section in diagnosis.get("evidence_sections") or []:
        if not isinstance(section, dict):
            continue
        title = str(section.get("title") or "日志证据")
        for excerpt in section.get("excerpts") or []:
            if not isinstance(excerpt, dict):
                continue
            rows.append(
                {
                    "strength": excerpt.get("level") or "提示",
                    "file": excerpt.get("file") or section.get("log_file") or "",
                    "location": excerpt.get("location") or "",
                    "message": excerpt.get("message") or "",
                    "relevance": title,
                }
            )
    return rows


def _triage_evidence_chain(diagnosis: dict[str, Any]) -> list[dict[str, Any]]:
    if diagnosis.get("scenario") == "snapshot_consolidation_failure":
        rows = _snapshot_evidence_chain(diagnosis)
        if rows:
            return rows
    rows = [item for item in diagnosis.get("evidence_chain") or [] if isinstance(item, dict)]
    if rows:
        return rows
    for section in diagnosis.get("evidence_sections") or []:
        if not isinstance(section, dict):
            continue
        title = str(section.get("title") or "日志证据")
        for excerpt in section.get("excerpts") or []:
            if not isinstance(excerpt, dict):
                continue
            rows.append(
                {
                    "strength": excerpt.get("strength") or excerpt.get("level") or section.get("level") or "提示",
                    "file": excerpt.get("file") or section.get("log_file") or "",
                    "location": excerpt.get("location") or "",
                    "message": excerpt.get("message") or "",
                    "relevance": title,
                }
            )
    return rows


def _triage_boundary_items(diagnosis: dict[str, Any]) -> list[str]:
    engine = diagnosis.get("diagnosis_engine") if isinstance(diagnosis.get("diagnosis_engine"), dict) else {}
    items: list[str] = ["当前证据不足，不能仅凭当前 support bundle 确认最终根因。"]
    reason = str(engine.get("model_quality_reason") or "").strip()
    if reason:
        items.append(f"模型输出质量门禁结论：{reason}")
    evidence_quality = str(diagnosis.get("model_evidence_quality") or "").strip()
    if evidence_quality:
        items.append(f"当前模型证据质量评估：{evidence_quality}")
    risk = str(diagnosis.get("model_overclaiming_risk") or "").strip()
    if risk:
        items.append(f"过度判断风险：{risk}")
    for item in _string_list(diagnosis.get("evidence_reasoning")):
        if item not in items:
            items.append(item)
    return items


def _render_model_triage_docx(payload: dict[str, Any], docx_path: Path, doc: Document) -> Path:
    meta = payload.get("metadata", {}) if isinstance(payload.get("metadata"), dict) else {}
    diagnosis = _problem_diagnosis(payload) or {}
    engine = diagnosis.get("diagnosis_engine") if isinstance(diagnosis.get("diagnosis_engine"), dict) else {}
    affected = diagnosis.get("affected_objects", {}) if isinstance(diagnosis.get("affected_objects"), dict) else {}
    report_title = str(meta.get("report_title") or MODEL_ASSISTED_TRIAGE_TITLE)
    diagnosis_title = str(diagnosis.get("title") or diagnosis.get("diagnosis_title") or "客户问题日志粗排查")
    customer = str(meta.get("customer_name") or "未指定客户")
    generated_at = str(meta.get("generated_at") or "")
    bundle_name = str(meta.get("support_bundle_name") or "")
    problem = _diagnosis_problem_text(diagnosis, meta)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(customer)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(18)
    run.font.bold = True

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_paragraph.add_run(report_title)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

    for label, value in (
        ("排查类型", diagnosis_title),
        ("分析时间", generated_at),
        ("日志包", bundle_name),
        ("客户问题", problem),
        ("问题对象", affected.get("primary_object") or affected.get("vm_name") or "当前问题描述未明确提供"),
        ("证据边界", "当前证据不足，不能仅凭当前 support bundle 确认最终根因。"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"{label}：").bold = True
        paragraph.add_run(str(value or "未记录"))

    doc.add_page_break()
    doc.add_heading("1. 粗排查结论", level=1)
    doc.add_heading("1.1 当前判断", level=2)
    doc.add_paragraph(_diagnosis_current_judgement(diagnosis))
    doc.add_heading("1.2 证据质量与边界", level=2)
    boundary_table = doc.add_table(rows=1, cols=2)
    _style_table(boundary_table)
    _fill_header(boundary_table.rows[0].cells, ["项目", "内容"])
    _add_docx_row(boundary_table, "处理方式", "模型输出已按证据质量降级为粗排查，不作为完整根因诊断。")
    _add_docx_row(boundary_table, "证据质量", str(diagnosis.get("model_evidence_quality") or "insufficient"))
    _add_docx_row(boundary_table, "质量门禁", str(engine.get("model_quality_reason") or "证据不足，按粗排查报告处理"))
    _add_docx_row(boundary_table, "判断边界", "当前证据不足，不能仅凭当前 support bundle 确认最终根因。")
    doc.add_heading("1.3 主要排查方向", level=2)
    _add_numbered_docx_items(doc, "1.3", diagnosis.get("likely_direction"), fallback="建议结合客户问题和补充材料继续缩小排查方向。")

    doc.add_heading("2. 相关证据线索", level=1)
    doc.add_heading("2.1 关键日志摘录", level=2)
    _add_evidence_chain_docx(doc, _triage_evidence_chain(diagnosis))
    doc.add_heading("2.2 证据不足原因", level=2)
    _add_numbered_docx_items(doc, "2.2", _triage_boundary_items(diagnosis))

    doc.add_heading("3. 当前不能确认", level=1)
    _add_numbered_docx_items(doc, "3", diagnosis.get("what_cannot_be_confirmed") or diagnosis.get("cannot_confirm_items") or [])

    doc.add_heading("4. 需要客户补充的信息", level=1)
    _add_numbered_docx_items(doc, "4", diagnosis.get("missing_materials") or diagnosis.get("supplemental_info") or [])

    doc.add_heading("5. 建议处理方向", level=1)
    _add_numbered_docx_items(doc, "5", diagnosis.get("recommended_next_steps") or diagnosis.get("handling_recommendations") or [])

    observations = diagnosis.get("other_observations") if isinstance(diagnosis.get("other_observations"), list) else []
    if observations:
        doc.add_heading("6. 附录：其他观察项", level=1)
        table = doc.add_table(rows=1, cols=4)
        _style_table(table)
        _fill_header(table.rows[0].cells, ["等级", "分类", "观察项", "建议"])
        for item in observations:
            if not isinstance(item, dict):
                continue
            cells = table.add_row().cells
            cells[0].text = str(item.get("severity") or "提示")
            cells[1].text = str(item.get("category") or "日志")
            cells[2].text = str(item.get("title") or "")
            cells[3].text = str(item.get("recommendation") or "建议结合业务影响进一步核查。")

    _set_document_font(doc)
    doc.save(docx_path)
    return docx_path


def _model_triage_html_document(payload: dict[str, Any]) -> str:
    meta = payload.get("metadata", {}) if isinstance(payload.get("metadata"), dict) else {}
    diagnosis = _problem_diagnosis(payload) or {}
    engine = diagnosis.get("diagnosis_engine") if isinstance(diagnosis.get("diagnosis_engine"), dict) else {}
    affected = diagnosis.get("affected_objects", {}) if isinstance(diagnosis.get("affected_objects"), dict) else {}
    report_title = str(meta.get("report_title") or MODEL_ASSISTED_TRIAGE_TITLE)
    diagnosis_title = str(diagnosis.get("title") or diagnosis.get("diagnosis_title") or "客户问题日志粗排查")
    problem = _diagnosis_problem_text(diagnosis, meta)
    missing = _string_list(diagnosis.get("missing_materials") or diagnosis.get("supplemental_info"))
    missing_text = "；".join(missing[:3]) if missing else "当前证据不足，不能仅凭当前 support bundle 确认最终根因。"
    directions_html = _numbered_items_html("1.3", diagnosis.get("likely_direction"), fallback="建议结合客户问题和补充材料继续缩小排查方向。")
    evidence_html = _evidence_chain_table_html(_triage_evidence_chain(diagnosis))
    boundary_html = _numbered_items_html("2.2", _triage_boundary_items(diagnosis))
    cannot_html = _numbered_items_html("3", diagnosis.get("what_cannot_be_confirmed") or diagnosis.get("cannot_confirm_items") or [])
    supplemental_html = _numbered_items_html("4", diagnosis.get("missing_materials") or diagnosis.get("supplemental_info") or [])
    recommendation_html = _numbered_items_html("5", diagnosis.get("recommended_next_steps") or diagnosis.get("handling_recommendations") or [])
    observations = diagnosis.get("other_observations") if isinstance(diagnosis.get("other_observations"), list) else []
    observations_html = ""
    if observations:
        rows = "\n".join(
            "<tr>"
            f"<td>{_esc(item.get('severity') if isinstance(item, dict) else '')}</td>"
            f"<td>{_esc(item.get('category') if isinstance(item, dict) else '')}</td>"
            f"<td>{_esc(item.get('title') if isinstance(item, dict) else '')}</td>"
            f"<td>{_esc(item.get('recommendation') if isinstance(item, dict) else '')}</td>"
            "</tr>"
            for item in observations
            if isinstance(item, dict)
        )
        observations_html = f"""
    <section>
      <h2>6. 附录：其他观察项</h2>
      <table><thead><tr><th>等级</th><th>分类</th><th>观察项</th><th>建议</th></tr></thead><tbody>{rows}</tbody></table>
    </section>"""

    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{_esc(report_title)}</title>
  <link rel="stylesheet" href="assets/log_analysis.css">
</head>
<body>
  <main class="page">
    <div class="hero">
      <h1>{_esc(report_title)}</h1>
      <p>排查类型：{_esc(diagnosis_title)}</p>
      <p>客户名称：{_esc(meta.get("customer_name") or "未指定客户")}</p>
      <p>日志包：{_esc(meta.get("support_bundle_name") or "未记录")}</p>
      <p>分析时间：{_esc(meta.get("generated_at") or "")}</p>
    </div>
    <div class="grid diagnosis-grid">
      <div class="metric diagnosis-card"><strong>当前判断</strong><p>{_esc(_diagnosis_current_judgement(diagnosis))}</p></div>
      <div class="metric diagnosis-card"><strong>证据边界</strong><p>当前证据不足，不能仅凭当前 support bundle 确认最终根因。</p></div>
      <div class="metric diagnosis-card"><strong>客户问题</strong><p>{_esc(problem)}</p></div>
      <div class="metric diagnosis-card"><strong>问题对象</strong><p>{_esc(affected.get("primary_object") or affected.get("vm_name") or "当前问题描述未明确提供")}</p></div>
      <div class="metric diagnosis-card"><strong>证据质量</strong><p>{_esc(diagnosis.get("model_evidence_quality") or "insufficient")}</p></div>
      <div class="metric diagnosis-card"><strong>还缺哪些材料</strong><p>{_esc(missing_text)}</p></div>
    </div>
    <section>
      <h2>1. 粗排查结论</h2>
      <h3>1.1 当前判断</h3>
      <p>{_esc(_diagnosis_current_judgement(diagnosis))}</p>
      <h3>1.2 证据质量与边界</h3>
      <table><tbody>
        <tr><th>处理方式</th><td>模型输出已按证据质量降级为粗排查，不作为完整根因诊断。</td></tr>
        <tr><th>证据质量</th><td>{_esc(diagnosis.get("model_evidence_quality") or "insufficient")}</td></tr>
        <tr><th>质量门禁</th><td>{_esc(engine.get("model_quality_reason") or "证据不足，按粗排查报告处理")}</td></tr>
        <tr><th>判断边界</th><td>当前证据不足，不能仅凭当前 support bundle 确认最终根因。</td></tr>
      </tbody></table>
      <h3>1.3 主要排查方向</h3>
      {directions_html}
    </section>
    <section>
      <h2>2. 相关证据线索</h2>
      <h3>2.1 关键日志摘录</h3>
      {evidence_html}
      <h3>2.2 证据不足原因</h3>
      {boundary_html}
    </section>
    <section>
      <h2>3. 当前不能确认</h2>
      {cannot_html}
    </section>
    <section>
      <h2>4. 需要客户补充的信息</h2>
      {supplemental_html}
    </section>
    <section>
      <h2>5. 建议处理方向</h2>
      {recommendation_html}
    </section>{observations_html}
  </main>
</body>
</html>
"""


def _render_problem_diagnosis_docx(payload: dict[str, Any], docx_path: Path, doc: Document) -> Path:
    meta = payload.get("metadata", {})
    diagnosis = _problem_diagnosis(payload) or {}
    affected = diagnosis.get("affected_objects", {}) if isinstance(diagnosis.get("affected_objects"), dict) else {}
    report_title = str(meta.get("report_title") or "VMware 日志分析报告")
    diagnosis_title = str(diagnosis.get("title") or diagnosis.get("diagnosis_title") or "客户问题驱动日志诊断报告")
    customer = str(meta.get("customer_name") or "未指定客户")
    generated_at = str(meta.get("generated_at") or "")
    bundle_name = str(meta.get("support_bundle_name") or "")

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(customer)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(18)
    run.font.bold = True

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_paragraph.add_run(report_title)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

    for label, value in (
        ("诊断类型", diagnosis_title),
        ("分析时间", generated_at),
        ("日志包", bundle_name),
        ("客户问题", diagnosis.get("customer_question") or meta.get("problem_description") or ""),
        ("问题对象", affected.get("primary_object") or "当前问题描述未明确提供"),
        ("是否可以完全定位", diagnosis.get("fully_determine_text") or "否，当前日志包无法完全确认最终根因。"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"{label}：").bold = True
        paragraph.add_run(str(value or "未记录"))

    doc.add_page_break()
    doc.add_heading("1. 诊断结论", level=1)
    doc.add_heading("1.1 当前判断", level=2)
    doc.add_paragraph(_diagnosis_current_judgement(diagnosis))
    doc.add_heading("1.2 是否能够确认根因", level=2)
    doc.add_paragraph(_diagnosis_root_cause_text(diagnosis))
    doc.add_heading("1.3 最可能排查方向", level=2)
    doc.add_paragraph(_join_values(diagnosis.get("likely_direction"), fallback="当前日志包尚未形成明确排查方向，需要结合补充材料继续确认。"))
    doc.add_heading("1.4 关键依据摘要", level=2)
    _add_numbered_docx_items(doc, "1.4", _key_basis_items(diagnosis))

    doc.add_heading("2. 问题概况", level=1)
    doc.add_paragraph(_log_analysis_independence_text())
    doc.add_heading("2.1 客户问题描述", level=2)
    doc.add_paragraph(_diagnosis_problem_text(diagnosis, meta))
    doc.add_heading("2.2 影响对象", level=2)
    table = doc.add_table(rows=1, cols=2)
    _style_table(table)
    _fill_header(table.rows[0].cells, ["项目", "内容"])
    _add_docx_row(table, "问题对象", str(affected.get("primary_object") or "当前问题描述未明确提供"))
    _add_docx_row(table, "问题动作", " / ".join(str(item) for item in affected.get("actions", []) if str(item).strip()) or "未明确")
    _add_docx_row(table, "问题症状", " / ".join(str(item) for item in affected.get("symptoms", []) if str(item).strip()) or "未明确")
    _add_docx_row(table, "问题时间", " / ".join(str(item) for item in affected.get("time_hints", []) if str(item).strip()) or "未提供")
    doc.add_heading("2.3 识别到的问题领域", level=2)
    doc.add_paragraph(_diagnosis_domains_text(affected, fallback="未知"))

    doc.add_heading("3. 关键日志证据", level=1)
    doc.add_heading("3.1 时间线摘要", level=2)
    _add_timeline_docx(doc, diagnosis, include_heading=False)
    doc.add_heading("3.2 关键日志摘录", level=2)
    _add_evidence_chain_docx(doc, diagnosis.get("evidence_chain"))
    doc.add_heading("3.3 证据说明", level=2)
    _add_evidence_explanations_docx(doc, diagnosis)

    doc.add_heading("4. 当前能够确认", level=1)
    _add_numbered_docx_items(doc, "4", diagnosis.get("what_can_be_confirmed") or diagnosis.get("confirmed_items") or [])

    doc.add_heading("5. 当前不能确认", level=1)
    _add_numbered_docx_items(doc, "5", diagnosis.get("what_cannot_be_confirmed") or diagnosis.get("cannot_confirm_items") or [])

    doc.add_heading("6. 需要客户补充的信息", level=1)
    _add_numbered_docx_items(doc, "6", diagnosis.get("missing_materials") or diagnosis.get("supplemental_info") or [])

    doc.add_heading("7. 建议处理方向", level=1)
    _add_numbered_docx_items(doc, "7", diagnosis.get("recommended_next_steps") or diagnosis.get("handling_recommendations") or [])

    observations = diagnosis.get("other_observations") if isinstance(diagnosis.get("other_observations"), list) else []
    if observations:
        doc.add_heading("8. 附录：其他观察项", level=1)
        table = doc.add_table(rows=1, cols=4)
        _style_table(table)
        _fill_header(table.rows[0].cells, ["等级", "分类", "观察项", "建议"])
        for item in observations:
            if not isinstance(item, dict):
                continue
            cells = table.add_row().cells
            cells[0].text = str(item.get("severity") or "提示")
            cells[1].text = str(item.get("category") or "日志")
            cells[2].text = str(item.get("title") or "")
            cells[3].text = str(item.get("recommendation") or "建议结合业务影响进一步核查。")

    _set_document_font(doc)
    doc.save(docx_path)
    return docx_path


def _problem_html_document(payload: dict[str, Any]) -> str:
    meta = payload.get("metadata", {})
    diagnosis = _problem_diagnosis(payload) or {}
    affected = diagnosis.get("affected_objects", {}) if isinstance(diagnosis.get("affected_objects"), dict) else {}
    report_title = meta.get("report_title") or "VMware 日志分析报告"
    diagnosis_title = diagnosis.get("title") or diagnosis.get("diagnosis_title") or "客户问题驱动日志诊断报告"
    missing = diagnosis.get("missing_materials") if isinstance(diagnosis.get("missing_materials"), list) else []
    missing_text = "；".join(str(item) for item in missing[:3]) if missing else "当前日志包无法完全确认最终根因。"
    domains_text = " / ".join(str(item) for item in affected.get("domains", []) if str(item).strip()) or "未知"
    actions_text = " / ".join(str(item) for item in affected.get("actions", []) if str(item).strip()) or "未明确"
    symptoms_text = " / ".join(str(item) for item in affected.get("symptoms", []) if str(item).strip()) or "未明确"
    time_text = " / ".join(str(item) for item in affected.get("time_hints", []) if str(item).strip()) or "未提供"

    timeline_html = _timeline_html(diagnosis, include_heading=False)
    evidence_html = _evidence_chain_table_html(diagnosis.get("evidence_chain"))
    basis_html = _numbered_items_html("1.4", _key_basis_items(diagnosis))
    confirmed_html = _numbered_items_html("4", diagnosis.get("what_can_be_confirmed") or diagnosis.get("confirmed_items") or [])
    cannot_html = _numbered_items_html("5", diagnosis.get("what_cannot_be_confirmed") or diagnosis.get("cannot_confirm_items") or [])
    supplemental_html = _numbered_items_html("6", diagnosis.get("missing_materials") or diagnosis.get("supplemental_info") or [])
    recommendation_html = _numbered_items_html("7", diagnosis.get("recommended_next_steps") or diagnosis.get("handling_recommendations") or [])
    evidence_explanation_html = _evidence_explanations_html(diagnosis)
    observations = diagnosis.get("other_observations") if isinstance(diagnosis.get("other_observations"), list) else []
    observations_html = ""
    if observations:
        rows = "\n".join(
            "<tr>"
            f"<td>{_esc(item.get('severity') if isinstance(item, dict) else '')}</td>"
            f"<td>{_esc(item.get('category') if isinstance(item, dict) else '')}</td>"
            f"<td>{_esc(item.get('title') if isinstance(item, dict) else '')}</td>"
            f"<td>{_esc(item.get('recommendation') if isinstance(item, dict) else '')}</td>"
            "</tr>"
            for item in observations
            if isinstance(item, dict)
        )
        observations_html = f"""
    <section>
      <h2>8. 附录：其他观察项</h2>
      <table><thead><tr><th>等级</th><th>分类</th><th>观察项</th><th>建议</th></tr></thead><tbody>{rows}</tbody></table>
    </section>"""

    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{_esc(report_title)}</title>
  <link rel="stylesheet" href="assets/log_analysis.css">
</head>
<body>
  <main class="page">
    <div class="hero">
      <h1>{_esc(report_title)}</h1>
      <p>诊断类型：{_esc(diagnosis_title)}</p>
      <p>客户名称：{_esc(meta.get("customer_name") or "未指定客户")}</p>
      <p>日志包：{_esc(meta.get("support_bundle_name") or "未记录")}</p>
      <p>分析时间：{_esc(meta.get("generated_at") or "")}</p>
    </div>
    <div class="grid diagnosis-grid">
      <div class="metric diagnosis-card"><strong>当前判断</strong><p>{_esc(_diagnosis_current_judgement(diagnosis))}</p></div>
      <div class="metric diagnosis-card"><strong>是否可以完全定位</strong><p>{_esc(diagnosis.get("fully_determine_text") or "否，当前日志包无法完全确认最终根因。")}</p></div>
      <div class="metric diagnosis-card"><strong>客户问题</strong><p>{_esc(diagnosis.get("customer_question") or meta.get("problem_description") or "未记录")}</p></div>
      <div class="metric diagnosis-card"><strong>问题对象</strong><p>{_esc(affected.get("primary_object") or "当前问题描述未明确提供")}</p></div>
      <div class="metric diagnosis-card"><strong>问题领域</strong><p>{_esc(domains_text)}</p></div>
      <div class="metric diagnosis-card"><strong>还缺哪些材料</strong><p>{_esc(missing_text)}</p></div>
    </div>
    <section>
      <h2>1. 诊断结论</h2>
      <h3>1.1 当前判断</h3>
      <p>{_esc(_diagnosis_current_judgement(diagnosis))}</p>
      <h3>1.2 是否能够确认根因</h3>
      <p>{_esc(_diagnosis_root_cause_text(diagnosis))}</p>
      <h3>1.3 最可能排查方向</h3>
      <p>{_esc(_join_values(diagnosis.get("likely_direction"), fallback="当前日志包尚未形成明确排查方向，需要结合补充材料继续确认。"))}</p>
      <h3>1.4 关键依据摘要</h3>
      {basis_html}
    </section>
    <section>
      <h2>2. 问题概况</h2>
      <p>{_esc(_log_analysis_independence_text())}</p>
      <h3>2.1 客户问题描述</h3>
      <p>{_esc(_diagnosis_problem_text(diagnosis, meta))}</p>
      <h3>2.2 影响对象</h3>
      <table><tbody>
        <tr><th>问题对象</th><td>{_esc(affected.get("primary_object") or "当前问题描述未明确提供")}</td></tr>
        <tr><th>问题动作</th><td>{_esc(actions_text)}</td></tr>
        <tr><th>问题症状</th><td>{_esc(symptoms_text)}</td></tr>
        <tr><th>问题时间</th><td>{_esc(time_text)}</td></tr>
      </tbody></table>
      <h3>2.3 识别到的问题领域</h3>
      <p>{_esc(domains_text)}</p>
    </section>
    <section>
      <h2>3. 关键日志证据</h2>
      <h3>3.1 时间线摘要</h3>
      {timeline_html}
      <h3>3.2 关键日志摘录</h3>
      {evidence_html}
      <h3>3.3 证据说明</h3>
      {evidence_explanation_html}
    </section>
    <section>
      <h2>4. 当前能够确认</h2>
      {confirmed_html}
    </section>
    <section>
      <h2>5. 当前不能确认</h2>
      {cannot_html}
    </section>
    <section>
      <h2>6. 需要客户补充的信息</h2>
      {supplemental_html}
    </section>
    <section>
      <h2>7. 建议处理方向</h2>
      {recommendation_html}
    </section>{observations_html}
  </main>
</body>
</html>
"""


def _render_snapshot_diagnosis_docx(payload: dict[str, Any], docx_path: Path, doc: Document) -> Path:
    meta = payload.get("metadata", {})
    diagnosis = _snapshot_diagnosis(payload) or {}
    affected = diagnosis.get("affected_objects", {}) if isinstance(diagnosis.get("affected_objects"), dict) else {}
    title = str(meta.get("report_title") or diagnosis.get("title") or MODEL_ASSISTED_DIAGNOSIS_TITLE)
    customer = str(meta.get("customer_name") or "未指定客户")
    generated_at = str(meta.get("generated_at") or "")
    bundle_name = str(meta.get("support_bundle_name") or "")

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(customer)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(18)
    run.font.bold = True

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_paragraph.add_run(title)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

    for label, value in (
        ("分析时间", generated_at),
        ("日志包", bundle_name),
        ("客户问题", diagnosis.get("customer_question") or meta.get("problem_description") or ""),
        ("问题虚拟机", affected.get("vm_name") or "当前日志包未明确识别"),
        ("是否可以完全定位", diagnosis.get("fully_determine_text") or "否，当前日志包无法完全确认现场最终状态。"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"{label}：").bold = True
        paragraph.add_run(str(value or "未记录"))

    doc.add_page_break()
    doc.add_heading("1. 诊断结论", level=1)
    doc.add_heading("1.1 当前判断", level=2)
    doc.add_paragraph(_diagnosis_current_judgement(diagnosis))
    doc.add_heading("1.2 是否能够确认根因", level=2)
    doc.add_paragraph(_diagnosis_root_cause_text(diagnosis))
    doc.add_heading("1.3 最可能排查方向", level=2)
    doc.add_paragraph(_join_values(diagnosis.get("likely_direction"), fallback="整合任务中断 / VM 状态变化 / 备份任务占用 / 文件锁 / VMDK 描述符链状态"))
    doc.add_heading("1.4 关键依据摘要", level=2)
    _add_numbered_docx_items(doc, "1.4", _key_basis_items(diagnosis))

    doc.add_heading("2. 问题概况", level=1)
    doc.add_paragraph(_log_analysis_independence_text())
    doc.add_heading("2.1 客户问题描述", level=2)
    doc.add_paragraph(_diagnosis_problem_text(diagnosis, meta))
    doc.add_heading("2.2 影响对象", level=2)
    table = doc.add_table(rows=1, cols=2)
    _style_table(table)
    _fill_header(table.rows[0].cells, ["项目", "内容"])
    _add_docx_row(table, "问题虚拟机", str(affected.get("vm_name") or "当前日志包未明确识别"))
    _add_docx_row(table, "虚拟机配置文件", str(affected.get("vmx_path") or "当前日志包未明确识别"))
    snapshot_disks = affected.get("snapshot_disks") if isinstance(affected.get("snapshot_disks"), list) else []
    disk_text = "\n".join(f"{item.get('controller')}: {item.get('file')}" for item in snapshot_disks if isinstance(item, dict)) or "当前日志包未提取到仍挂载的快照磁盘。"
    _add_docx_row(table, "涉及 vmdk / 快照链文件", disk_text)
    chain_files = affected.get("vmdk_chain_files") if isinstance(affected.get("vmdk_chain_files"), list) else []
    chain_rows = _vmdk_chain_rows(chain_files)
    doc.add_paragraph("VMDK 描述符链：")
    if chain_rows:
        chain_table = doc.add_table(rows=1, cols=2)
        _style_table(chain_table)
        _fill_header(chain_table.rows[0].cells, ["层级", "描述符文件"])
        for level, filename in chain_rows:
            cells = chain_table.add_row().cells
            cells[0].text = level
            cells[1].text = filename
    else:
        doc.add_paragraph("当前日志包未提取到完整描述符链。")
    doc.add_heading("2.3 识别到的问题领域", level=2)
    doc.add_paragraph(_diagnosis_domains_text(affected, fallback="快照 / 存储 / 虚拟机"))

    doc.add_heading("3. 关键日志证据", level=1)
    doc.add_heading("3.1 时间线摘要", level=2)
    _add_timeline_docx(doc, diagnosis, include_heading=False)
    doc.add_heading("3.2 关键日志摘录", level=2)
    _add_evidence_chain_docx(doc, _snapshot_evidence_chain(diagnosis))
    doc.add_heading("3.3 证据说明", level=2)
    _add_evidence_explanations_docx(doc, diagnosis)

    doc.add_heading("4. 当前能够确认", level=1)
    _add_numbered_docx_items(doc, "4", diagnosis.get("confirmed_items") or ["当前日志包包含快照整合相关线索，但仍需结合现场信息确认。"])

    doc.add_heading("5. 当前不能确认", level=1)
    _add_numbered_docx_items(doc, "5", diagnosis.get("cannot_confirm_items") or [])

    doc.add_heading("6. 需要客户补充的信息", level=1)
    _add_numbered_docx_items(doc, "6", diagnosis.get("supplemental_info") or [])

    doc.add_heading("7. 建议处理方向", level=1)
    _add_numbered_docx_items(doc, "7", diagnosis.get("handling_recommendations") or [])

    observations = diagnosis.get("other_observations") if isinstance(diagnosis.get("other_observations"), list) else []
    if observations:
        doc.add_heading("8. 附录：其他观察项", level=1)
        table = doc.add_table(rows=1, cols=4)
        _style_table(table)
        _fill_header(table.rows[0].cells, ["等级", "分类", "观察项", "建议"])
        for item in observations:
            if not isinstance(item, dict):
                continue
            cells = table.add_row().cells
            cells[0].text = str(item.get("severity") or "提示")
            cells[1].text = str(item.get("category") or "日志")
            cells[2].text = str(item.get("title") or "")
            cells[3].text = str(item.get("recommendation") or "建议结合业务影响进一步核查。")

    _set_document_font(doc)
    doc.save(docx_path)
    return docx_path


def _snapshot_html_document(payload: dict[str, Any]) -> str:
    meta = payload.get("metadata", {})
    diagnosis = _snapshot_diagnosis(payload) or {}
    affected = diagnosis.get("affected_objects", {}) if isinstance(diagnosis.get("affected_objects"), dict) else {}
    title = meta.get("report_title") or diagnosis.get("title") or MODEL_ASSISTED_DIAGNOSIS_TITLE
    missing = diagnosis.get("missing_materials") if isinstance(diagnosis.get("missing_materials"), list) else []
    missing_text = "；".join(str(item) for item in missing[:3]) if missing else "当前日志包无法完全确认现场最终状态。"
    snapshot_disks = affected.get("snapshot_disks") if isinstance(affected.get("snapshot_disks"), list) else []
    chain_files = affected.get("vmdk_chain_files") if isinstance(affected.get("vmdk_chain_files"), list) else []

    timeline_html = _timeline_html(diagnosis, include_heading=False)
    evidence_html = _evidence_chain_table_html(_snapshot_evidence_chain(diagnosis))
    evidence_explanation_html = _evidence_explanations_html(diagnosis)
    basis_html = _numbered_items_html("1.4", _key_basis_items(diagnosis))
    confirmed_html = _numbered_items_html("4", diagnosis.get("confirmed_items") or [])
    cannot_html = _numbered_items_html("5", diagnosis.get("cannot_confirm_items") or [])
    supplemental_html = _numbered_items_html("6", diagnosis.get("supplemental_info") or [])
    recommendation_html = _numbered_items_html("7", diagnosis.get("handling_recommendations") or [])
    observations = diagnosis.get("other_observations") if isinstance(diagnosis.get("other_observations"), list) else []
    observations_html = ""
    if observations:
        rows = "\n".join(
            "<tr>"
            f"<td>{_esc(item.get('severity') if isinstance(item, dict) else '')}</td>"
            f"<td>{_esc(item.get('category') if isinstance(item, dict) else '')}</td>"
            f"<td>{_esc(item.get('title') if isinstance(item, dict) else '')}</td>"
            f"<td>{_esc(item.get('recommendation') if isinstance(item, dict) else '')}</td>"
            "</tr>"
            for item in observations
            if isinstance(item, dict)
        )
        observations_html = f"""
    <section>
      <h2>8. 附录：其他观察项</h2>
      <table><thead><tr><th>等级</th><th>分类</th><th>观察项</th><th>建议</th></tr></thead><tbody>{rows}</tbody></table>
    </section>"""

    disk_rows = "\n".join(
        f"<tr><td>{_esc(item.get('controller'))}</td><td>{_esc(item.get('file'))}</td></tr>"
        for item in snapshot_disks
        if isinstance(item, dict)
    ) or "<tr><td colspan=\"2\">当前日志包未提取到仍挂载的快照磁盘。</td></tr>"
    chain_rows = _vmdk_chain_rows(chain_files)
    chain_html = (
        "<table class=\"chain-table\"><thead><tr><th>层级</th><th>描述符文件</th></tr></thead><tbody>"
        + "".join(f"<tr><td>{_esc(level)}</td><td>{_esc(filename)}</td></tr>" for level, filename in chain_rows)
        + "</tbody></table>"
        if chain_rows
        else "当前日志包未提取到完整描述符链。"
    )

    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{_esc(title)}</title>
  <link rel="stylesheet" href="assets/log_analysis.css">
</head>
<body>
  <main class="page">
    <div class="hero">
      <h1>{_esc(title)}</h1>
      <p>客户名称：{_esc(meta.get("customer_name") or "未指定客户")}</p>
      <p>日志包：{_esc(meta.get("support_bundle_name") or "未记录")}</p>
      <p>分析时间：{_esc(meta.get("generated_at") or "")}</p>
    </div>
    <div class="grid diagnosis-grid">
      <div class="metric diagnosis-card"><strong>当前判断</strong><p>{_esc(_diagnosis_current_judgement(diagnosis))}</p></div>
      <div class="metric diagnosis-card"><strong>是否可以完全定位</strong><p>{_esc(diagnosis.get("fully_determine_text") or "否，当前日志包无法完全确认现场最终状态。")}</p></div>
      <div class="metric diagnosis-card"><strong>客户问题</strong><p>{_esc(diagnosis.get("customer_question") or meta.get("problem_description") or "未记录")}</p></div>
      <div class="metric diagnosis-card"><strong>影响对象</strong><p>{_esc(affected.get("vm_name") or "当前日志包未明确识别")}</p></div>
      <div class="metric diagnosis-card"><strong>关键证据数量</strong><p>{int(diagnosis.get("key_evidence_count", 0) or 0)}</p></div>
      <div class="metric diagnosis-card"><strong>还缺哪些材料</strong><p>{_esc(missing_text)}</p></div>
    </div>
    <section>
      <h2>1. 诊断结论</h2>
      <h3>1.1 当前判断</h3>
      <p>{_esc(_diagnosis_current_judgement(diagnosis))}</p>
      <h3>1.2 是否能够确认根因</h3>
      <p>{_esc(_diagnosis_root_cause_text(diagnosis))}</p>
      <h3>1.3 最可能排查方向</h3>
      <p>{_esc(_join_values(diagnosis.get("likely_direction"), fallback="整合任务中断 / VM 状态变化 / 备份任务占用 / 文件锁 / VMDK 描述符链状态"))}</p>
      <h3>1.4 关键依据摘要</h3>
      {basis_html}
    </section>
    <section>
      <h2>2. 问题概况</h2>
      <p>{_esc(_log_analysis_independence_text())}</p>
      <h3>2.1 客户问题描述</h3>
      <p>{_esc(_diagnosis_problem_text(diagnosis, meta))}</p>
      <h3>2.2 影响对象</h3>
      <table><tbody>
        <tr><th>问题虚拟机</th><td>{_esc(affected.get("vm_name") or "当前日志包未明确识别")}</td></tr>
        <tr><th>虚拟机配置文件</th><td>{_esc(affected.get("vmx_path") or "当前日志包未明确识别")}</td></tr>
        <tr><th>涉及 vmdk / 快照链文件</th><td><table><thead><tr><th>磁盘</th><th>文件</th></tr></thead><tbody>{disk_rows}</tbody></table></td></tr>
        <tr><th>VMDK 描述符链</th><td>{chain_html}</td></tr>
      </tbody></table>
      <h3>2.3 识别到的问题领域</h3>
      <p>{_esc(_diagnosis_domains_text(affected, fallback="快照 / 存储 / 虚拟机"))}</p>
    </section>
    <section>
      <h2>3. 关键日志证据</h2>
      <h3>3.1 时间线摘要</h3>
      {timeline_html}
      <h3>3.2 关键日志摘录</h3>
      {evidence_html}
      <h3>3.3 证据说明</h3>
      {evidence_explanation_html}
    </section>
    <section>
      <h2>4. 当前能够确认</h2>
      {confirmed_html}
    </section>
    <section>
      <h2>5. 当前不能确认</h2>
      {cannot_html}
    </section>
    <section>
      <h2>6. 需要客户补充的信息</h2>
      {supplemental_html}
    </section>
    <section>
      <h2>7. 建议处理方向</h2>
      {recommendation_html}
    </section>{observations_html}
  </main>
</body>
</html>
"""


def _snapshot_evidence_section_html(index: int, section: dict[str, Any]) -> str:
    excerpts = section.get("excerpts") if isinstance(section.get("excerpts"), list) else []
    excerpt_html = "\n".join(
        f"<p class=\"evidence-excerpt\">{_esc(item.get('message') if isinstance(item, dict) else item)}</p>"
        for item in excerpts
    ) or "<p class=\"muted\">当前日志包未提取到可展示的日志摘录。</p>"
    return f"""
      <div class="evidence-block">
        <h3>3.{index} {_esc(section.get("title") or "日志证据")}</h3>
        <p><strong>日志文件：</strong>{_esc(section.get("log_file") or "当前日志包未提取到对应文件")}</p>
        <p><strong>日志摘录：</strong></p>
        {excerpt_html}
        <p><strong>判断：</strong>{_esc(section.get("judgement") or "建议结合现场信息继续确认。")}</p>
      </div>
"""


def _ordered_list_html(items: Any) -> str:
    values = [str(item) for item in items if str(item).strip()] if isinstance(items, list) else []
    if not values:
        values = ["当前日志包未提取到对应信息。"]
    return "<ol>" + "".join(f"<li>{_esc(item)}</li>" for item in values) + "</ol>"


def _timeline_items(diagnosis: dict[str, Any]) -> list[dict[str, Any]]:
    timeline = diagnosis.get("timeline") if isinstance(diagnosis.get("timeline"), list) else []
    return [item for item in timeline if isinstance(item, dict)]


def _add_timeline_docx(doc: Document, diagnosis: dict[str, Any], include_heading: bool = True) -> None:
    timeline = _timeline_items(diagnosis)
    if not timeline:
        if include_heading:
            doc.add_heading("问题时间线", level=2)
        doc.add_paragraph("当前日志包未提取到可展示的时间线摘要。")
        return
    if include_heading:
        doc.add_heading("问题时间线", level=2)
    table = doc.add_table(rows=1, cols=5)
    _style_table(table)
    _fill_header(table.rows[0].cells, ["时间", "日志文件", "位置", "事件", "关联说明"])
    for item in timeline[:12]:
        cells = table.add_row().cells
        cells[0].text = str(item.get("time") or "")
        cells[1].text = str(item.get("file") or "")
        cells[2].text = str(item.get("location") or "")
        cells[3].text = str(item.get("event") or "")
        cells[4].text = str(item.get("relevance") or "")
    if len(timeline) > 12:
        doc.add_paragraph("更多时间线事件建议在 HTML 日志分析报告中查看。")


def _timeline_html(diagnosis: dict[str, Any], include_heading: bool = True) -> str:
    timeline = _timeline_items(diagnosis)
    if not timeline:
        heading = "<h3>问题时间线</h3>" if include_heading else ""
        return f"{heading}<p class=\"muted\">当前日志包未提取到可展示的时间线摘要。</p>"
    rows = "\n".join(
        "<tr>"
        f"<td>{_esc(item.get('time'))}</td>"
        f"<td>{_esc(item.get('file'))}</td>"
        f"<td>{_esc(item.get('location'))}</td>"
        f"<td>{_esc(item.get('event'))}</td>"
        f"<td>{_esc(item.get('relevance'))}</td>"
        "</tr>"
        for item in timeline[:20]
    )
    overflow = "<p class=\"muted\">更多时间线事件建议结合原始日志继续查看。</p>" if len(timeline) > 20 else ""
    heading = "<h3>问题时间线</h3>" if include_heading else ""
    return (
        f"{heading}"
        "<table><thead><tr><th>时间</th><th>日志文件</th><th>位置</th><th>事件</th><th>关联说明</th></tr></thead>"
        f"<tbody>{rows}</tbody></table>{overflow}"
    )


def _vmdk_chain_rows(chain_files: list[Any]) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for item in chain_files:
        if isinstance(item, dict):
            raw_value = item.get("file") or item.get("path") or item.get("name")
        else:
            raw_value = item
        filename = _vmdk_descriptor_filename(raw_value)
        if not filename:
            continue
        level = _vmdk_descriptor_level(filename)
        key = (level, filename)
        if key in seen:
            continue
        seen.add(key)
        rows.append(key)
    return rows


def _vmdk_descriptor_filename(value: Any) -> str:
    text = str(value or "").strip().replace("\\", "/")
    return text.rsplit("/", 1)[-1].strip()


def _vmdk_descriptor_level(filename: str) -> str:
    lower = filename.lower()
    stem = filename[:-5] if lower.endswith(".vmdk") else filename
    suffix = stem.rsplit("-", 1)[-1]
    if len(suffix) == 6 and suffix.isdigit():
        return suffix
    return "base"


def _add_docx_row(table, label: str, value: str) -> None:
    cells = table.add_row().cells
    cells[0].text = label
    cells[1].text = value


def _add_numbered_docx_list(doc: Document, items: Any) -> None:
    values = [str(item) for item in items if str(item).strip()] if isinstance(items, list) else []
    if not values:
        values = ["当前日志包未提取到对应信息。"]
    for item in values:
        doc.add_paragraph(item, style="List Number")


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(31, 78, 121)


def _style_table(table) -> None:
    table.style = "Table Grid"
    table.autofit = True
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _fill_header(cells, headers: list[str]) -> None:
    for cell, header in zip(cells, headers, strict=False):
        cell.text = header
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "D9EAF7")
        cell._tc.get_or_add_tcPr().append(shading)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def _set_document_font(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Microsoft YaHei"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _evidence_rows(findings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for finding in findings:
        for evidence in finding.get("evidence", []) or []:
            if isinstance(evidence, dict):
                rows.append(evidence)
    return rows


def _severity_class(value: Any) -> str:
    text = str(value or "")
    if "高" in text:
        return "high"
    if "中" in text:
        return "medium"
    return "low"


def _esc(value: Any) -> str:
    if value is None:
        return ""
    return html.escape(str(value))
