"""Word presentation of the shared customer document, including full appendices."""
from __future__ import annotations

import math
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
import pymupdf

from vstacklens.reports.document_model import get_document, COLORS, NAVY, INK, NEUTRAL
from vstacklens.reports.word_toc import update_word_toc


def _font(run, size=10.5, bold=False, color=INK):
    run.font.name = 'Microsoft YaHei'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor.from_string(color.lstrip('#'))


def _soft(value):
    value = str(value)
    return re.sub(r'([/\\_.-])',lambda m:m.group()+'\u200b',value)


def _bookmark(p, name, number):
    name='d_'+re.sub(r'[^A-Za-z0-9_]','_',name)[:35]
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number))
    p._p.insert(0,start);p._p.append(end)


def _field(paragraph, instruction, cache=''):
    for kind in ('begin','code','separate','text','end'):
        run=OxmlElement('w:r')
        if kind=='code':
            node=OxmlElement('w:instrText');node.set(qn('xml:space'),'preserve');node.text=instruction
        elif kind=='text':
            node=OxmlElement('w:t');node.text=cache
        else:
            node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),kind)
            if kind=='begin':node.set(qn('w:dirty'),'true')
        run.append(node);paragraph._p.append(run)


class WordDocumentWriter:
    def __init__(self, engine, report, model):
        self.engine=engine;self.report=report;self.model=model;self.document=Document()
        engine._setup_document(self.document)
        self.document.core_properties.title=model['title'];self.document.core_properties.author='VStackLens'
        self.document.core_properties.last_modified_by='VStackLens'
        for name in ('Normal','TOC 1','TOC 2','TOC 3'):
            if name in self.document.styles:
                s=self.document.styles[name];s.font.size=Pt(10.5);s.paragraph_format.line_spacing=1.15
        self.width=481.9;self.page_height=739.8;self.remaining=self.page_height
        # LO's Microsoft YaHei auto line spacing includes the East Asian grid.
        # This is a pagination measurement, not a reduced line-spacing setting.
        self.line_height=20.15
        self.serial=0;self.trace=[]
        self.pending_page_break=False
        try:self.metrics=pymupdf.Font(fontfile=r'C:\Windows\Fonts\msyh.ttc')
        except Exception:self.metrics=pymupdf.Font('china-s')
        self.glyph_widths={}
        labels=[row['name'] for s in model['sections'] if s['kind']=='environment'
                for node in s['blocks'] if node['kind']=='chart' for row in node['series']]
        label_width=max(120,min(190,max((self.metrics.text_length(name,fontsize=10.5)*1.06+12 for name in labels),default=120)))
        self.chart_widths=[label_width,self.width-label_width-56,56]
        footer=self.document.sections[0].footer.paragraphs[0]
        footer.clear();_font(footer.add_run((model['date_label']+' · ' if model['date_label'] else '')+'第 '),10.5)
        _field(footer,' PAGE ');_font(footer.add_run(' 页'),10.5)
        header=self.document.sections[0].header.paragraphs[0]
        _font(header.add_run(model['header_identity']),10.5,color=NAVY)

    def lines(self,text,width):
        if not text:return 1
        result=0
        for part in str(text).split('\n'):
            used=0;count=1;available=max(width-12,10)
            for token in part:
                if token=='\u200b':continue
                if token not in self.glyph_widths:self.glyph_widths[token]=self.metrics.text_length(token,fontsize=10.5)
                length=self.glyph_widths[token]
                if used and used+length>available:count+=1;used=0
                if not used and token.isspace():continue
                while length>available:
                    count+=1;length-=available
                used+=length
            result+=count
        return result

    def page_break(self):
        self.pending_page_break=True;self.remaining=self.page_height

    def _apply_break(self,p):
        if self.pending_page_break:
            p.paragraph_format.page_break_before=True
            self.pending_page_break=False

    def _before_table(self):
        if self.pending_page_break or (len(self.document.element.body)>1 and self.document.element.body[-2].tag==qn('w:tbl')):
            p=self.document.add_paragraph();p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=Pt(1);p.paragraph_format.keep_with_next=True
            self._apply_break(p)

    def fixed_width(self,table,widths):
        self.engine._set_table_fixed_width(table,widths)
        # The legacy helper has a full-page preferred width; a nested chart
        # bar must instead use its actual cell width. Main tables respect the
        # declared page margins as well.
        table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'),str(sum(widths)))
        for row in table.rows:
            column=0
            for cell in row._tr.findall(qn('w:tc')):
                properties=cell.find(qn('w:tcPr'));span=properties.find(qn('w:gridSpan')) if properties is not None else None
                count=int(span.get(qn('w:val'))) if span is not None else 1
                preferred=properties.find(qn('w:tcW'))
                if preferred is not None:preferred.set(qn('w:w'),str(sum(widths[column:column+count])))
                column+=count

    def paragraph(self,text,*,bold=False,size=10.5,color=INK,keep=False):
        if text is None or text=='':return None
        p=self.document.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.15
        self._apply_break(p)
        p.paragraph_format.keep_with_next=keep
        _font(p.add_run(str(text)),size,bold,color)
        self.remaining-=self.lines(text,self.width)*size*1.15+4
        return p

    def heading(self,text,level=1,anchor=''):
        p=self.document.add_heading('',level=level)
        self._apply_break(p)
        _font(p.add_run(text),16 if level==1 else 13 if level==2 else 11,True,NAVY)
        p.paragraph_format.space_before=Pt(12 if level==1 else 8);p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.keep_with_next=True
        self.remaining-=40 if level==1 else 30
        self.serial+=1
        if anchor:_bookmark(p,anchor,self.serial)
        return p

    def table(self,node,*,appendix=False,keep_after=False):
        if not node or not node['rows']:return
        total=sum(node['weights']);widths=[self.width*w/total for w in node['weights']]
        rows=node['rows'];index=0;fragment=0
        line_counts=[max(self.lines(value,widths[i]) for i,value in enumerate(row['cells'])) for row in rows]
        # Word/LO can wrap a path or mixed-script identifier one line earlier
        # than font-advance measurement. Reserve that line for wrapped rows
        # when choosing keep-together fragments; do not alter visible spacing.
        heights=[(n+(1 if appendix and n>1 else 0))*self.line_height+6 for n in line_counts]
        suffix=[0.0]*(len(rows)+1)
        for position in range(len(rows)-1,-1,-1):suffix[position]=suffix[position+1]+heights[position]
        # Body short tables are kept whole. Appendix page-sized fragments use
        # actual font advances to retain readable spacing and explicit titles.
        while index<len(rows):
            first=appendix and fragment==0
            title=f'附录 {node["letter"]} {node["title"]}' if first else node['title']+('（续）' if fragment else '')
            if first:
                self.heading(title,1,node.get('anchor',''))
            head_height=self.lines(title,self.width)*(self.line_height*(16/10.5 if first else 1))+6
            head_height+=max(self.lines(value,widths[i]) for i,value in enumerate(node['headers']))*self.line_height+6
            all_height=head_height+suffix[0]
            if appendix:
                if index or (len(rows)<=15 and all_height<=self.page_height-10):
                    selected=rows[index:];used=suffix[index]
                else:
                    # Keep the first titled fragment below half a page. Let
                    # Word paginate the continuing table naturally; its repeat
                    # header already says (续). A guessed current page position
                    # must never inject a blank or prematurely break a page.
                    # Leave room for native word-breaking differences (paths
                    # often wrap before a slash/hyphen). A smaller initial
                    # fragment can share the preceding short table's page;
                    # the continuing table still paginates at full page size.
                    available=max(self.line_height,self.page_height*.42-head_height)
                    selected=[];used=0
                    while index+len(selected)<len(rows):
                        row_height=heights[index+len(selected)]
                        if selected and used+row_height>available:break
                        selected.append(rows[index+len(selected)]);used+=row_height
                        if used>=available:break
            else:
                selected=rows[index:];used=suffix[index]
            self._before_table()
            table=self.document.add_table(rows=0,cols=len(widths));table.style='Table Grid';table.autofit=False
            self.fixed_width(table,[int(w*20) for w in widths])
            oversize=len(selected)==1 and heights[index]>self.page_height-head_height
            caption=title
            if not first:
                self._row(table,[caption]+['']*(len(widths)-1),widths,header=True,merge=True)
            if oversize:
                self._row(table,['续段对象：'+selected[0]['cells'][0]]+['']*(len(widths)-1),widths,header=True,merge=True)
            self._row(table,node['headers'],widths,header=True)
            keep_whole=(len(rows)<=15 and len(selected)==len(rows) or (appendix and index==0)) and len(selected)<=15 and used+head_height<self.page_height-10 and not node.get('_problem_layout')
            for offset,row in enumerate(selected):
                long=heights[index+offset]>self.page_height-head_height
                follows=keep_after or (not appendix and (bool(node.get('notes')) or bool(node.get('reference'))))
                keep=(keep_whole and offset<len(selected)-1) or (follows and index+offset==len(rows)-1)
                self._row(table,row['cells'],widths,key=row['id'],keep=keep,split=long,
                          grade_columns=[i for i,h in enumerate(node['headers']) if h=='等级'])
                self.trace.append({'table':node['id'],'row':row['id'],'key':row['key'],'fragment':fragment})
            self.remaining-=head_height+used
            index+=len(selected);fragment+=1
        notes=node.get('notes') or []
        for index,note in enumerate(notes):self.paragraph(note,keep=not appendix and (index<len(notes)-1 or bool(node.get('reference'))))
        if node.get('reference'):self.paragraph(node['reference'])
        # No printable blank paragraph after a table: it could spill alone to
        # a new page before the next explicit page break. Cell line spacing and
        # visible text sizes are unchanged.

    def _row(self,table,values,widths,*,header=False,merge=False,key='',keep=False,split=False,grade_columns=(),heading=0,anchor=''):
        tr=OxmlElement('w:tr');properties=OxmlElement('w:trPr')
        if header:properties.append(OxmlElement('w:tblHeader'))
        if not split:properties.append(OxmlElement('w:cantSplit'))
        tr.append(properties)
        if merge:values=values[:1];widths=[sum(widths)]
        for i,value in enumerate(values):
            cell=OxmlElement('w:tc');pr=OxmlElement('w:tcPr');w=OxmlElement('w:tcW');w.set(qn('w:w'),str(int(widths[i]*20)));w.set(qn('w:type'),'dxa');pr.append(w)
            if merge:
                span=OxmlElement('w:gridSpan');span.set(qn('w:val'),str(len(table.columns)));pr.append(span)
            margin=OxmlElement('w:tcMar')
            for side,n in [('top',60),('bottom',60),('left',108),('right',108)]:
                node=OxmlElement('w:'+side);node.set(qn('w:w'),str(n));node.set(qn('w:type'),'dxa');margin.append(node)
            pr.append(margin)
            if header:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),NAVY.lstrip('#'));pr.append(fill)
            cell.append(pr);p=OxmlElement('w:p');ppr=OxmlElement('w:pPr')
            if heading:
                style=OxmlElement('w:pStyle');style.set(qn('w:val'),'Heading'+str(heading));ppr.append(style)
            spacing=OxmlElement('w:spacing');spacing.set(qn('w:after'),'0');spacing.set(qn('w:before'),'0');spacing.set(qn('w:line'),'276');spacing.set(qn('w:lineRule'),'auto');ppr.append(spacing)
            wrapping=OxmlElement('w:wordWrap');wrapping.set(qn('w:val'),'off');ppr.append(wrapping)
            if header or keep:ppr.append(OxmlElement('w:keepNext'))
            p.append(ppr)
            if anchor:
                self.serial+=1;bs=OxmlElement('w:bookmarkStart');bs.set(qn('w:id'),str(self.serial));bs.set(qn('w:name'),'d_'+anchor.replace('-','_'));p.append(bs)
            if key and i==0:
                self.serial+=1;bs=OxmlElement('w:bookmarkStart');bs.set(qn('w:id'),str(self.serial));bs.set(qn('w:name'),'r_'+key[:30]);p.append(bs)
            run=OxmlElement('w:r');rp=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
            for attr in ('ascii','hAnsi','eastAsia'):fonts.set(qn('w:'+attr),'Microsoft YaHei')
            rp.append(fonts);size=OxmlElement('w:sz');size.set(qn('w:val'),'32' if heading else '21');rp.append(size)
            text_color=COLORS.get(str(value),INK) if i in grade_columns else INK
            color=OxmlElement('w:color');color.set(qn('w:val'),'FFFFFF' if header else text_color.lstrip('#'));rp.append(color)
            if header:rp.append(OxmlElement('w:b'))
            run.append(rp)
            parts=str(value).split('\n')
            for j,part in enumerate(parts):
                if j:run.append(OxmlElement('w:br'))
                text=OxmlElement('w:t');text.set(qn('xml:space'),'preserve');text.text=part;run.append(text)
            p.append(run)
            if key and i==0:
                end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(self.serial));p.append(end)
            if anchor:
                end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(self.serial));p.append(end)
            cell.append(p);tr.append(cell)
        table._tbl.append(tr)

    def graph(self,node):
        """Scannable native Word geometry; no opaque embedded images."""
        self.paragraph(node['title'],bold=True,size=12,color=NAVY,keep=True)
        container=self.document.add_table(rows=1,cols=1);container.autofit=False
        self.fixed_width(container,[int(self.width*20)])
        container.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        parent=container.rows[0].cells[0]
        margins=OxmlElement('w:tcMar')
        for side in ('left','right','top','bottom'):
            value=OxmlElement('w:'+side);value.set(qn('w:w'),'0');value.set(qn('w:type'),'dxa');margins.append(value)
        parent._tc.get_or_add_tcPr().append(margins)
        parent.paragraphs[0].paragraph_format.line_spacing=Pt(1)
        parent.paragraphs[0].paragraph_format.space_after=Pt(0)
        table=parent.add_table(rows=0,cols=3);table.autofit=False
        widths=self.chart_widths
        self.fixed_width(table,[int(w*20) for w in widths])
        for item in node['series']:
            self._row(table,[item['name'],'',item['text']],widths,keep=True,grade_columns=[0] if item.get('level') else [])
            row=table.rows[-1];cell=row.cells[1]
            # Use the source value and source threshold, not rounded labels.
            stops=sorted(set([0.0,float(item['value']),float(node['axis_max']),*([float(node['threshold'])] if node['threshold'] is not None else [])]))
            stops=[s for s in stops if 0<=s<=node['axis_max']]
            segments=[(a,b) for a,b in zip(stops,stops[1:]) if b>a]
            bar=cell.add_table(rows=1,cols=len(segments));bar.autofit=False
            segment_widths=[max(1,int((b-a)/node['axis_max']*(widths[1]-12)*20)) for a,b in segments]
            self.fixed_width(bar,segment_widths)
            caption=OxmlElement('w:tblCaption');caption.set(qn('w:val'),'图形几何');bar._tbl.tblPr.append(caption)
            for i,(a,b) in enumerate(segments):
                c=bar.rows[0].cells[i];pr=c._tc.get_or_add_tcPr()
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),COLORS.get(item.get('level'),NEUTRAL).lstrip('#') if a<item['value'] else 'F1F5F9');pr.append(fill)
                if b==node['threshold']:
                    borders=OxmlElement('w:tcBorders');line=OxmlElement('w:right');line.set(qn('w:val'),'single');line.set(qn('w:sz'),'8');line.set(qn('w:color'),NAVY.lstrip('#'));borders.append(line);pr.append(borders)
                para=c.paragraphs[0];para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=Pt(11.4)
            cell.paragraphs[0].paragraph_format.line_spacing=Pt(1)
            cell.paragraphs[0].paragraph_format.space_after=Pt(0)
        axis=parent.add_table(rows=1,cols=3);axis.autofit=False
        self.fixed_width(axis,[int(w*20) for w in widths])
        axis.rows[0].cells[0].text=''
        axis.rows[0].cells[2].text=''
        p=axis.rows[0].cells[1].paragraphs[0]
        available=widths[1]-12
        from docx.enum.text import WD_TAB_ALIGNMENT
        ticks=node['ticks'];p.paragraph_format.space_after=Pt(0)
        for i,tick in enumerate(ticks[1:],1):
            alignment=WD_TAB_ALIGNMENT.RIGHT if i==len(ticks)-1 else WD_TAB_ALIGNMENT.CENTER
            p.paragraph_format.tab_stops.add_tab_stop(Pt(available*tick['value']/node['axis_max']),alignment)
        for i,tick in enumerate(ticks):_font(p.add_run(('\t' if i else '')+tick['text']),10.5,color=NEUTRAL)
        self.paragraph(node['more'])

    def cards(self,node):
        """The <=4 cluster mode is a native two-column card grid in Word too."""
        self.paragraph(node['title'],bold=True,size=12,color=NAVY,keep=True)
        table=self.document.add_table(rows=0,cols=2);table.autofit=False
        self.fixed_width(table,[int(self.width*10),int(self.width*10)])
        for offset in range(0,len(node['rows']),2):
            row=table.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for index in range(2):
                if offset+index>=len(node['rows']):continue
                data=node['rows'][offset+index];cell=row.cells[index]
                self.engine._shade_cell(cell,'F1F5F9')
                p=cell.paragraphs[0];_font(p.add_run(data['cells'][0]),11,True,NAVY)
                for label,value in zip(node['headers'][1:],data['cells'][1:]):
                    p=cell.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
                    _font(p.add_run(label+'：'+value))

    def ranked_cards(self,node):
        self.paragraph(node['title'],bold=True,size=12,color=NAVY,keep=True)
        for index,item in enumerate(node['rows']):
            self._before_table()
            table=self.document.add_table(rows=1,cols=1);table.autofit=False;table.style='Table Grid'
            self.fixed_width(table,[int(self.width*20)])
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            cell=table.cell(0,0);p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.15
            wrapping=OxmlElement('w:wordWrap');wrapping.set(qn('w:val'),'off');p._p.get_or_add_pPr().append(wrapping)
            _font(p.add_run(item['name']),10.5,True,NAVY)
            if item['detail']:
                p=cell.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.15
                _font(p.add_run(item['detail']))
            if index==len(node['rows'])-1:
                for note in node.get('notes') or []:
                    p=cell.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.15
                    _font(p.add_run(note))

    def render(self,path):
        d=self.model
        title=self.document.add_paragraph(style='Title');_font(title.add_run(d['title']),24,True,NAVY)
        self.paragraph('总体结论：'+d['opening']['conclusion'],bold=True,size=18)
        self.paragraph(d['opening']['problem_text'],bold=True,size=14)
        if d['opening']['priorities']:
            self.paragraph('需优先处理的问题',bold=True,size=13,color=NAVY)
            for p in d['opening']['priorities']:
                para=self.paragraph(p['title']+f' · {p["objects"]} 个受影响对象')
                _font(para.add_run('  '+p['level']),10.5,True,COLORS[p['level']])
        self.paragraph(d['opening']['more']);self.paragraph(d['opening']['statistics'])
        self.page_break()
        p=self.document.add_paragraph();_font(p.add_run('目录'),16,True,NAVY)
        self._apply_break(p)
        p.paragraph_format.keep_with_next=True
        toc=self.document.add_paragraph();_field(toc,' TOC \\o "1-3" \\h \\z \\u ')
        for section in d['sections']:
            self.heading(section['title'],1,section['id'])
            if section['kind']=='overview':
                rows=section['rows'];t={'id':'overview','title':'问题总览','weights':[1,5,1],'headers':['等级','问题类别','受影响对象'],
                  'rows':[{'id':'ov'+str(i),'key':'ov'+str(i),'cells':[r['level'],r['title'],str(r['objects'])]} for i,r in enumerate(rows)]}
                self.table(t)
                self.graph(section['distribution'])
            elif section['kind']=='problems':
                for item in section['items']:
                    self.heading(item['title'],2,'problem-'+item['id'])
                    self.paragraph(d['labels']['level']+'：'+item['level'],color=COLORS[item['level']],keep=True)
                    self.paragraph(item.get('capacity_note'))
                    for field in item['fields']:self.paragraph(field['label']+'：'+field['text'])
                    for group in item['render_groups']:
                        if group['title']:self.paragraph(group['title'],bold=True,keep=True)
                        for i,step in enumerate(group['steps']):self.paragraph(f'{i+1}. {step}')
                    self.table(item['representative_table'],keep_after=bool(item['reference']));self.paragraph(item['reference'])
                    for note in item['object_notes']:
                        if note.startswith('NSX Manager 的资源预留'):self.paragraph(note)
            elif section['kind']=='checklist':
                items=section['items']
                for i in range(0,len(items),2):
                    self.paragraph('   '.join('✓ '+r['label'] for r in items[i:i+2]),
                                   keep=i<4 and i+2<len(items))
            elif section['kind']=='environment':
                for group in section['groups']:
                    self.heading(group['title'],2,group['id'])
                    for node in group['blocks']:
                        if node.get('presentation')=='ranked-cards':self.ranked_cards(node)
                        elif node['kind']=='table':self.table(node)
                        elif node['kind']=='cards':self.cards(node)
                        elif node['kind']=='paragraph':
                            if node['title']:self.paragraph(node['title'],bold=True,color=NAVY,keep=True)
                            self.paragraph(node['text'])
                        elif node['kind']=='chart':self.graph(node)
        if d['appendices']:
            self.page_break()
            for table in d['appendices']:
                self.table(table,appendix=True)
        self.document.save(path)
        self.engine.last_document_trace=self.trace


def render_document_word(engine,report,output_path):
    from vstacklens.products.reporting import ensure_report_has_valid_data
    from vstacklens.application.artifact_publish import staged_file
    ensure_report_has_valid_data(report)
    output_path=Path(output_path);output_path.parent.mkdir(parents=True,exist_ok=True)
    from vstacklens.reports.document_word_visual import VisualWordWriter
    writer=VisualWordWriter(engine,report,get_document(report))
    with staged_file(output_path) as stage:
        writer.render(stage)
        engine.last_toc_update=update_word_toc(stage)
    return output_path
