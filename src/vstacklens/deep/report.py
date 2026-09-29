from __future__ import annotations

import html
import json
from collections import Counter
from datetime import UTC, datetime
from math import isfinite
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from vstacklens.deep.contracts import DeepCategory, DeepDataset, DeepFinding, DeepReport


CATEGORY_LABELS = {
    DeepCategory.CURRENT_RISK: "当前隐患",
    DeepCategory.HISTORICAL_HEALTH: "历史健康",
    DeepCategory.TREND: "趋势",
}
_MAX_EVENT_SAMPLES_IN_CUSTOMER_REPORT = 3
_EVIDENCE_TABLE_THRESHOLD = 10


def _vlan_summary(value: Any) -> str:
    if not isinstance(value, dict):
        return "VLAN 未显式设置"
    if value.get("vlanId") is not None:
        return f"VLAN ID {value['vlanId']}"
    if value.get("pvlanId") is not None:
        return f"Private VLAN ID {value['pvlanId']}"
    ranges = value.get("ranges") or []
    if ranges:
        rendered = [
            f"{item.get('start', '?')}-{item.get('end', '?')}"
            if isinstance(item, dict)
            else str(item)
            for item in ranges[:8]
        ]
        suffix = f" 等 {len(ranges)} 个范围" if len(ranges) > len(rendered) else ""
        return "Trunk VLAN " + ", ".join(rendered) + suffix
    return str(value.get("spec_type") or "VLAN 设置")


def _teaming_summary(value: Any) -> str:
    if not isinstance(value, dict):
        return "Teaming 未显式设置"
    parts = [str(value.get("policy") or "policy 未设置")]
    active = value.get("active_uplink_count")
    standby = value.get("standby_uplink_count")
    if active is not None or standby is not None:
        parts.append(f"active {active if active is not None else '?'} / standby {standby if standby is not None else '?'}")
    failure = value.get("failure_criteria")
    if isinstance(failure, dict):
        checks = []
        for field, label in (("checkBeacon", "beacon"), ("checkDuplex", "duplex"), ("checkErrorPercent", "error"), ("checkStatus", "status")):
            if field in failure:
                checks.append(f"{label}={'on' if failure[field] else 'off'}")
        speed = failure.get("checkSpeed")
        if speed:
            checks.append(f"speed={speed}" + (f"@{failure['speed']}" if failure.get("speed") is not None else ""))
        if checks:
            parts.append("failure checks " + ", ".join(checks))
    return "; ".join(parts)


def _uses_evidence_table(finding: DeepFinding) -> bool:
    return len(finding.evidence) > _EVIDENCE_TABLE_THRESHOLD or (finding.rule_id == "NET-DEEP-003" and len(finding.evidence) > 1)


def _evidence_summary_lines(item: Any) -> list[str]:
    summary = item.value_summary or {}
    samples = summary.get("event_samples") or []
    sample_count = int(summary.get("event_samples_count", len(samples)) or 0)
    if item.kind != "event" and not samples:
        lines: list[str] = []
        description = str(summary.get("summary") or "").strip()
        if description:
            lines.append(f"摘要：{description}")
        value = summary.get("value")
        if isinstance(value, dict):
            scalar_fields = []
            for key, field_value in value.items():
                if isinstance(field_value, (str, int, float, bool)) or field_value is None:
                    rendered = str(field_value)
                    if len(rendered) > 120:
                        rendered = rendered[:119].rstrip() + "…"
                    scalar_fields.append(f"{key}={rendered}")
            if scalar_fields:
                lines.append("关键数据：" + "；".join(scalar_fields[:5]))
                if len(scalar_fields) > 5:
                    lines.append(f"其余 {len(scalar_fields) - 5} 个标量字段见 Dataset 原始证据。")
            elif value:
                lines.append("结构化字段保留在 Dataset 原始证据中。")
            distributed = value.get("distributed_network_profiles")
            if isinstance(distributed, dict):
                status = str(distributed.get("status") or "unknown")
                lines.append(
                    f"vDS 配置采集：{status}；{distributed.get('switch_count', 0)} 个 Distributed Switch，"
                    f"{distributed.get('portgroup_count', 0)} 个 Distributed Port Group。"
                )
                for profile in (distributed.get("portgroups") or [])[:5]:
                    lines.append(
                        f"{profile.get('dvs_name') or 'DVS'}/{profile.get('name') or 'Port Group'}："
                        f"{_vlan_summary(profile.get('vlan'))}；Teaming={_teaming_summary(profile.get('teaming'))}"
                    )
                omitted_portgroups = max(0, int(distributed.get("portgroup_count", 0) or 0) - min(5, len(distributed.get("portgroups") or [])))
                if omitted_portgroups:
                    lines.append(f"其余 {omitted_portgroups} 个 Distributed Port Group 仍在 Dataset 原始证据中。")
                if distributed.get("truncated"):
                    lines.append("分布式网络配置样本已达到本地保留上限。")
        elif isinstance(value, list):
            values = [str(entry) for entry in value[:5] if isinstance(entry, (str, int, float, bool))]
            if values:
                lines.append("观测值：" + "、".join(values))
                if len(value) > len(values):
                    lines.append(f"其余 {len(value) - len(values)} 项见 Dataset 原始证据。")
            elif value:
                lines.append(f"共 {len(value)} 项结构化观测，详情见 Dataset 原始证据。")
        elif value is not None and (not description or str(value) not in description):
            unit = str(summary.get("unit") or "")
            lines.append(f"观测值：{value}" + (f" {unit}" if unit else ""))
        if item.raw_pointer:
            lines.append(f"Dataset 原始位置：{item.raw_pointer}")
        return lines or ["Evidence 已保存在 Dataset 原始记录中。"]

    lines: list[str] = []
    description = str(summary.get("summary") or "").strip()
    if description:
        lines.append(f"摘要：{description}")
    value = summary.get("value")
    if isinstance(value, dict):
        for key, label in (
            ("event_count", "事件数"),
            ("resolved_event_count", "已解析对象事件数"),
            ("unresolved_object_event_count", "未解析对象事件数"),
            ("cluster_count", "关联簇数"),
        ):
            if value.get(key) is not None:
                lines.append(f"{label}：{value[key]}")
    elif value is not None:
        unit = str(summary.get("unit") or "count")
        lines.append(f"事件数：{value} {unit}")
    event_filter = summary.get("event_filter") or []
    if event_filter:
        lines.append("匹配特征：" + "、".join(str(item) for item in event_filter[:6]))

    if samples:
        visible_samples = samples[:_MAX_EVENT_SAMPLES_IN_CUSTOMER_REPORT]
        lines.append(f"代表事件（显示 {len(visible_samples)} / {sample_count} 条）：")
        for sample in visible_samples:
            if not isinstance(sample, dict):
                lines.append(str(sample))
                continue
            parts = [str(sample[key]) for key in ("timestamp", "event_type") if sample.get(key)]
            location = sample.get("host") or sample.get("entity")
            if location:
                parts.append(f"对象：{location}")
            if sample.get("event_key") is not None:
                parts.append(f"Event Key：{sample['event_key']}")
            message = " ".join(str(sample.get("message") or "").split())
            if len(message) > 200:
                message = message[:199].rstrip() + "…"
            detail = " | ".join(parts)
            lines.append(f"{detail} — {message}" if detail and message else detail or message)
        omitted = sample_count - len(visible_samples)
        if omitted:
            lines.append(f"其余 {omitted} 条样本保留在 Dataset 原始证据中。")
    if item.raw_pointer:
        lines.append(f"Dataset 原始位置：{item.raw_pointer}")
    return lines or ["事件 Evidence 已保存在 Dataset 原始记录中。"]


def _no_analysis_message(dataset: DeepDataset) -> str:
    connection_failure = next(
        (
            item
            for item in dataset.collection_log
            if item.get("action") in {"connection", "session.content"} and item.get("status") not in {"ok", "connected"}
        ),
        None,
    )
    if connection_failure and connection_failure.get("action") == "connection":
        return "本次未能连接 vCenter 并取得环境清单，因此未开展风险评估。本报告不表示环境正常；请恢复只读管理连接后重新巡检。"
    if connection_failure:
        return "本次未能读取 vCenter 环境清单，因此未开展风险评估。本报告不表示环境正常；请检查管理连接后重新巡检。"
    if bool((dataset.manifest.impact or {}).get("degraded")):
        return "本次采集过程已降级，未形成可展示的可信分析结果；数据不足的范围不作正常性判断。"
    return "本次没有形成可展示的可信分析结果；数据不足的范围不作正常性判断。"


def _log_coverage_note(dataset: DeepDataset) -> str | None:
    log_state = next(
        (item for item in reversed(dataset.collection_log) if item.get("action") == "logs.history"),
        None,
    )
    if log_state is None:
        return None

    def safe_count(value: Any) -> int:
        try:
            return max(0, int(value or 0))
        except (TypeError, ValueError):
            return 0

    status = str(log_state.get("status") or "unknown").casefold()
    successful = safe_count(log_state.get("successful_count"))
    partially_successful = safe_count(log_state.get("partially_successful_count"))
    total = safe_count(log_state.get("category_count"))
    truncated = safe_count(log_state.get("truncated_count"))
    missing_privileges = bool(log_state.get("missing_privilege_ids"))
    log_records = [
        record
        for record in dataset.records
        if record.kind == "log" and record.metadata.get("rule_id") == "LOG-DEEP-001"
    ]
    correlated_log_records = sum(
        1
        for record in log_records
        if safe_count(record.metadata.get("related_event_count"))
        or safe_count(record.metadata.get("related_record_count"))
    )
    same_interface_state_links = sum(
        safe_count((record.metadata.get("related_event_evidence_counts") or {}).get("same_object_time_category_interface_state"))
        for record in log_records
    )

    if status == "ok":
        message = "已完成本次接口可返回的日志采集"
    elif status == "partial" or successful:
        message = "日志仅部分采集成功"
    elif status in {"unavailable", "not_requested"} or status.startswith("skipped_") or status in {"budget_exceeded", "cancelled"}:
        message = "本次未取得可用的历史日志"
    else:
        message = "本次历史日志采集状态未能确认"

    if total:
        message += f"（{successful}/{total} 个对象与日志类别组合成功读取或确认无内容"
        if partially_successful:
            message += f"；另有 {partially_successful} 个组合部分读取"
        message += "）"
    if missing_privileges:
        message += "；部分来源读取权限受限"
    if truncated:
        message += f"；{truncated} 项达到读取上限"

    line_count = sum(safe_count((record.value or {}).get("line_count")) for record in log_records)
    segment_count = safe_count(log_state.get("segment_read_count"))
    retained_bytes = safe_count(log_state.get("retained_log_bytes"))
    source_bytes_read = safe_count(log_state.get("source_log_bytes_read", retained_bytes))
    max_total_bytes = safe_count(log_state.get("max_log_total_bytes"))
    primary_bytes = safe_count(log_state.get("retained_primary_log_bytes"))
    fallback_bytes = safe_count(log_state.get("retained_fallback_log_bytes"))
    max_primary_bytes = safe_count(log_state.get("max_log_primary_bytes"))
    max_fallback_bytes = safe_count(log_state.get("max_log_fallback_bytes"))
    if line_count or segment_count:
        message += f"；实际保留 {line_count:,} 行日志，分段读取 {segment_count:,} 次"
    if max_total_bytes:
        message += (
            f"；日志预算：总量 {retained_bytes / (1024 * 1024):.2f}/{max_total_bytes / (1024 * 1024):.2f} MiB，"
            f"vCenter 主来源 {primary_bytes / (1024 * 1024):.2f}/{max_primary_bytes / (1024 * 1024):.2f} MiB，"
            f"备用来源合计 {fallback_bytes / (1024 * 1024):.2f}/{max_fallback_bytes / (1024 * 1024):.2f} MiB"
        )
    requested_log_days = log_state.get("requested_log_days")
    if requested_log_days:
        cutoff = str(log_state.get("log_window_cutoff_utc") or "未知")
        read_bytes = safe_count(log_state.get("source_log_bytes_read"))
        excluded_lines = safe_count(log_state.get("time_window_excluded_line_count"))
        unknown_time_lines = safe_count(log_state.get("time_unknown_line_count"))
        window_status_value = str(log_state.get("time_window_status") or "unknown")
        window_status = {
            "complete": "目标窗口定位完成",
            "complete_reached_file_start": "窗口覆盖到文件起始位置",
            "complete_with_unknown_time": "窗口定位完成但含时间无法判定行",
            "complete_no_recent_rows": "已定位但目标窗口无日志行",
            "complete_no_log_rows": "日志文件为空",
            "partial": "窗口只定位到部分数据",
            "unavailable": "时间定位不可用",
            "not_limited": "未限定天数",
            "not_attempted": "未尝试",
        }.get(window_status_value, "时间范围未能完全确认")
        message += (
            f"；本次日志目标为最近 {requested_log_days} 天（起点 {cutoff}），时间定位状态 {window_status}，"
            f"实际读取约 {read_bytes / (1024 * 1024):.2f} MiB，保留 {retained_bytes / (1024 * 1024):.2f} MiB；"
            f"在已读取的定位页中跳过窗口外 {excluded_lines:,} 行"
        )
        if unknown_time_lines:
            message += f"，保留并标记时间无法判定 {unknown_time_lines:,} 行"
        message += "；更早的未读取行不计入本次日志分析"
    elif source_bytes_read > retained_bytes:
        message += f"；来源实际读取 {source_bytes_read / (1024 * 1024):.2f} MiB，含未保留的定位/去重内容"
    deduplicated_lines = safe_count(log_state.get("deduplicated_source_line_count"))
    if deduplicated_lines:
        message += f"；跨来源按同对象、类别、源行号及相同行内容去重 {deduplicated_lines:,} 行"
    source_labels = {
        "vim.DiagnosticManager.BrowseDiagnosticLog": "vCenter API",
        "vim.DiagnosticManager.QueryDescriptions": "vCenter API 描述查询",
        "PowerCLI.Get-Log": "PowerCLI",
        "vim.DiagnosticManager.DirectESXi": "直连 ESXi",
        "DeepCollectionPolicy.log_budget": "日志预算",
        "DeepCollectionPolicy.resource_guard": "资源保护",
        "DeepCollectionPolicy.time_window": "日志时间范围",
    }
    source_result_labels = {
        "ok": "成功",
        "empty": "无内容",
        "no_logs": "未发现匹配的日志描述符",
        "partial": "部分读取",
        "interface_unavailable": "接口不可用或未返回可用日志",
        "esxcli_host_missing": "PowerCLI 未返回对应主机日志结果",
        "permission_denied": "权限不足",
        "timeout": "超时",
        "total_size_limit_reached": "日志总预算耗尽",
        "file_size_limit_reached": "单文件预算耗尽",
        "file_limit_reached": "文件数上限耗尽",
        "not_attempted": "未尝试",
        "not_needed": "无需回退",
        "time_window_seek_not_supported": "该来源不支持按时间定位",
        "timestamp_unparseable_in_requested_window": "窗口内有日志时间无法判定",
        "time_seek_unavailable": "尾部定位接口不可用",
        "order_unreliable": "日志时间顺序无法验证",
        "budget_limited": "日志预算限制",
    }
    source_statuses: dict[str, Counter[str]] = {}
    source_result_statuses: dict[str, Counter[str]] = {}
    for record in log_records:
        source = str(getattr(record.source, "api", "") or "")
        label = source_labels.get(source)
        status_value = str((record.metadata or {}).get("log_status") or "")
        if label and status_value:
            source_result_statuses.setdefault(label, Counter())[status_value] += 1
    for attempt in log_state.get("attempts") or []:
        source = str(attempt.get("source") or "")
        label = source_labels.get(source)
        status_value = str(attempt.get("status") or "")
        if label and status_value:
            source_statuses.setdefault(label, Counter())[status_value] += 1
            result_counts = attempt.get("result_status_counts") or {}
            if isinstance(result_counts, dict):
                counts_for_source = source_result_statuses.setdefault(label, Counter())
                for result_status, count in result_counts.items():
                    counts_for_source[str(result_status)] = max(counts_for_source[str(result_status)], safe_count(count))
    if source_statuses:
        source_summaries = []
        for label, counts in source_statuses.items():
            statuses = "、".join(f"{status_value}×{count}" if count > 1 else status_value for status_value, count in sorted(counts.items()))
            details = source_result_statuses.get(label)
            if details:
                detail_text = "、".join(
                    f"{source_result_labels.get(status_value, status_value)} {count}"
                    for status_value, count in sorted(details.items())
                )
                source_summaries.append(f"{label} {statuses}（结果：{detail_text}）")
            else:
                source_summaries.append(f"{label} {statuses}")
        message += "；日志来源状态：" + "；".join(source_summaries)

    category_times: dict[str, tuple[datetime, datetime]] = {}
    for record in log_records:
        if str((record.metadata or {}).get("time_source") or "") != "line_timestamp":
            continue
        if not (record.value or {}).get("lines"):
            continue
        try:
            start = datetime.fromisoformat(str(record.window.start).replace("Z", "+00:00")).astimezone(UTC)
            end = datetime.fromisoformat(str(record.window.end).replace("Z", "+00:00")).astimezone(UTC)
        except (AttributeError, TypeError, ValueError):
            continue
        category = str((record.metadata or {}).get("log_category") or "unknown")
        previous = category_times.get(category)
        category_times[category] = (min(start, previous[0]) if previous else start, max(end, previous[1]) if previous else end)
    if category_times:
        coverage_start = min(item[0] for item in category_times.values())
        coverage_end = max(item[1] for item in category_times.values())
        coverage_days = max(0.0, (coverage_end - coverage_start).total_seconds() / 86_400)
        def format_utc(value: datetime) -> str:
            return value.strftime("%Y-%m-%d %H:%M:%S UTC")

        message += f"；日志行时间戳跨度 {format_utc(coverage_start)} 至 {format_utc(coverage_end)}（约 {coverage_days:.1f} 天；这是可见样本跨度，不代表连续或全量覆盖）"
        category_labels = {"vpxd": "vCenter", "ha": "HA", "vmkernel": "VMkernel", "hostd": "hostd", "vpxa": "vpxa", "vobd": "VOBD", "syslog": "Syslog", "vsan": "vSAN"}
        windows = [
            f"{category_labels.get(category, category)} {format_utc(bounds[0])}—{format_utc(bounds[1])}"
            for category, bounds in sorted(category_times.items())
        ]
        message += "；类别时间窗：" + "；".join(windows)

    marker_totals: Counter[str] = Counter()
    for record in log_records:
        for marker, count in ((record.value or {}).get("marker_counts") or {}).items():
            marker_totals[str(marker)] += safe_count(count)
    marker_labels = {
        "error": "错误",
        "timeout": "超时",
        "storage": "存储",
        "network": "网络",
        "performance": "性能",
        "ha": "HA",
    }
    marker_summary = [f"{marker_labels.get(marker, marker)} {count}" for marker, count in sorted(marker_totals.items()) if count]
    if marker_summary:
        message += "；日志关键词标记累计命中：" + "、".join(marker_summary)
        message += "（同一行或不同来源可能重复计数，仅是人工复核线索，不是唯一故障数或根因证明）"
    if correlated_log_records:
        message += f"；{correlated_log_records} 条日志记录存在同对象、时间窗且类型相关的证据关联（共现不代表因果）"
        if same_interface_state_links:
            message += f"；其中 {same_interface_state_links} 条同时核实同一网卡和 up/down 方向"
    else:
        message += "；本次未形成日志与事件/任务/性能的可验证关联，相关范围证据不足，不代表没有历史问题"
    message += "。未读取、未保留或未覆盖的历史范围不作正常性判断。"
    return message


def _log_correlation_lines(dataset: DeepDataset, *, limit: int = 3) -> list[str]:
    record_by_id = {record.record_id: record for record in dataset.records}
    categories = {
        "alarm": "告警",
        "cluster": "集群",
        "compute": "计算",
        "drs": "DRS",
        "ha": "HA",
        "maintenance": "维护",
        "network": "网络",
        "performance": "性能",
        "storage": "存储",
    }
    kinds = {"event": "事件", "task": "任务", "perf": "性能"}
    links: dict[tuple[str, ...], tuple[int, float, str]] = {}

    def distance_label(value: Any) -> tuple[float, str]:
        try:
            seconds = float(value)
        except (TypeError, ValueError):
            return float("inf"), "距离未知"
        if not isfinite(seconds) or seconds < 0:
            return float("inf"), "距离未知"
        if seconds == 0:
            return 0.0, "位于记录时间窗内"
        if seconds < 60:
            label = f"相距 {seconds:.0f} 秒"
        elif seconds < 3600:
            label = f"相距 {seconds / 60:.1f} 分钟"
        else:
            label = f"相距 {seconds / 3600:.1f} 小时"
        return seconds, label

    for log_record in dataset.records:
        if log_record.kind != "log" or log_record.metadata.get("rule_id") != "LOG-DEEP-001":
            continue
        entity = log_record.entity.display_ref or log_record.entity.stable_id
        category = str(log_record.metadata.get("log_category") or "日志")
        category_label = categories.get(category, category)
        log_pointer = log_record.raw_pointer or "logs/log.ndjson"

        for sample in log_record.metadata.get("related_event_samples") or []:
            if not isinstance(sample, dict):
                continue
            event_type = str(sample.get("event_type") or "vSphere Event")
            event_key = sample.get("event_key")
            event_label = f"{event_type}（Event Key {event_key}）" if event_key is not None else event_type
            event_interfaces = ", ".join(str(item) for item in sample.get("interfaces") or [] if item)
            event_state = str(sample.get("state") or "")
            if event_interfaces or event_state:
                state_label = {"down": "down", "up/restored": "up/restored"}.get(event_state, event_state)
                event_label += f"；{event_interfaces} {state_label}".strip("； ")
            evidence_level = str(sample.get("evidence_level") or "")
            evidence_priority = {
                "same_object_time_category_interface_state": 0,
                "same_object_time_category_interface": 1,
                "same_object_time_category": 2,
            }.get(evidence_level, 2)
            if "network" in (sample.get("matching_categories") or []):
                evidence_label = {
                    "same_object_time_category_interface_state": "同对象、同网卡、同方向",
                    "same_object_time_category_interface": "同对象、同网卡，方向未能从日志确认",
                    "same_object_time_category": "同对象、时间窗和网络类型，未核实同网卡或方向",
                }.get(evidence_level, "同对象、时间窗和网络类型")
            else:
                evidence_label = "同对象、时间窗和共享事件类型"
            timestamp = str(sample.get("timestamp") or "时间未知")
            matching = sorted({str(item) for item in sample.get("matching_categories") or [] if item})
            matched_label = "、".join(categories.get(item, item) for item in matching) or "类型未知"
            numeric_distance, distance = distance_label(sample.get("time_distance_seconds"))
            event_location = sample.get("matched_object") or sample.get("object") or sample.get("host") or entity
            object_label = str(event_location or entity)
            event_pointer = str(sample.get("dataset_pointer") or "")
            event_source_api = str(sample.get("event_source_api") or "EventManager.QueryEvents")
            log_timestamp = str(sample.get("log_timestamp") or log_record.window.end)
            log_line_index = sample.get("log_line_index")
            source_line_number = sample.get("source_line_number")
            log_summary = str(log_record.metadata.get("correlation_summary") or "").strip()
            log_evidence = f"{category_label} 日志 {log_timestamp}"
            if source_line_number is not None:
                log_evidence += f"（源日志第 {source_line_number} 行）"
            elif log_line_index is not None:
                log_evidence += f"（保留样本第 {log_line_index} 行）"
            if log_summary:
                log_evidence += f"，摘要 {log_summary}"
            key = ("event", event_type, str(event_key or ""), timestamp, object_label, matched_label)
            event_source_evidence = f"来源 {event_source_api}"
            if event_pointer:
                event_source_evidence += f"；Dataset 位置：{log_pointer} → {event_pointer}"
            else:
                event_ref = f"Event Key {event_key}" if event_key is not None else f"时间 {timestamp}"
                event_source_evidence += f"；{event_ref} 的关联摘要保存在日志 Evidence 中"
            line = (
                f"对象 {entity}；{log_evidence} ↔ 事件 {event_label}；"
                f"事件对象 {object_label}；时间 {timestamp}；{distance}；"
                f"匹配类型 {matched_label}；关联依据：{evidence_label}；{event_source_evidence}。"
            )
            parent_event_pointer = sample.get("parent_dataset_pointer")
            if parent_event_pointer:
                line = line[:-1] + f"；父 Dataset 位置：{parent_event_pointer}。"
            links.setdefault(key, (evidence_priority, numeric_distance, line))

        for sample in log_record.metadata.get("related_record_samples") or []:
            if not isinstance(sample, dict):
                continue
            record_id = str(sample.get("record_id") or "")
            related = record_by_id.get(record_id)
            kind = str(sample.get("kind") or (related.kind if related else "record"))
            kind_label = kinds.get(kind, kind)
            source = str(sample.get("source") or (related.source.api if related else "来源未知"))
            start = str(sample.get("window_start") or (related.window.start if related else ""))
            end = str(sample.get("window_end") or (related.window.end if related else ""))
            matching = sorted({str(item) for item in sample.get("matching_categories") or [] if item})
            matched_label = "、".join(categories.get(item, item) for item in matching) or "类型未知"
            numeric_distance, distance = distance_label(sample.get("time_distance_seconds"))
            related_pointer = related.raw_pointer if related and related.raw_pointer else f"events/{kind}.ndjson#{record_id or 'unknown'}"
            log_timestamp = str(sample.get("log_timestamp") or log_record.window.end)
            log_line_index = sample.get("log_line_index")
            source_line_number = sample.get("source_line_number")
            log_summary = str(log_record.metadata.get("correlation_summary") or "").strip()
            log_evidence = f"{category_label} 日志 {log_timestamp}"
            if source_line_number is not None:
                log_evidence += f"（源日志第 {source_line_number} 行）"
            elif log_line_index is not None:
                log_evidence += f"（保留样本第 {log_line_index} 行）"
            if log_summary:
                log_evidence += f"，摘要 {log_summary}"
            key = (kind, record_id, source, start, end)
            line = (
                f"对象 {entity}；{log_evidence} ↔ {kind_label}；来源 {source}；"
                f"记录时间窗 {start or '未知'} 至 {end or '未知'}；{distance}；"
                f"匹配类型 {matched_label}；Dataset 位置：{log_pointer} → {related_pointer}。"
            )
            links.setdefault(key, (2, numeric_distance, line))

    ordered = sorted(links.values(), key=lambda item: (item[0], item[1], item[2]))
    visible = [line for _, _, line in ordered[: max(0, limit)]]
    if len(ordered) > len(visible):
        visible.append(f"其余 {len(ordered) - len(visible)} 条关联保留在 Dataset 原始证据中。")
    return visible


class DeepHtmlReportBuilder:
    def render(self, report: DeepReport, output_path: Path, dataset: DeepDataset) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(self._html(report, dataset), encoding="utf-8")
        (output_path.parent / "deep_report_payload.json").write_text(
            json.dumps({"report": report.model_dump(mode="json"), "dataset_id": dataset.dataset_id}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return output_path

    def _html(self, report: DeepReport, dataset: DeepDataset) -> str:
        empty_analysis_message = html.escape(_no_analysis_message(dataset))
        log_coverage_note = _log_coverage_note(dataset)
        log_coverage_html = (
            f'<p class="boundary">{html.escape(log_coverage_note)}</p>'
            if log_coverage_note
            else ""
        )
        log_correlation_lines = _log_correlation_lines(dataset)
        log_correlation_html = (
            '<h3>日志时间关联（共现不代表因果）</h3><ul>'
            + "".join(f"<li>{html.escape(line)}</li>" for line in log_correlation_lines)
            + "</ul>"
            if log_correlation_lines
            else ""
        )
        scope_rows = "".join(
            f"<li><strong>{html.escape(item.label)}</strong>：已完成 {html.escape('、'.join(item.completed_capabilities))}，可信 Finding {item.finding_count}，通过检查 {item.pass_count} 项</li>"
            for item in report.analysis_scope
        ) or f"<li>{empty_analysis_message}</li>"
        pass_rows = "".join(
            f"<tr><td>{html.escape(CATEGORY_LABELS.get(item['category'], item['category']))}</td><td>{html.escape(str(item['profile']))}</td><td>{item['passed']}</td></tr>"
            for item in report.pass_summary
        ) or '<tr><td colspan="3">无可展示的通过项聚合</td></tr>'
        trend_summary = report.trend_summary or {}
        trend_block = ""
        ready_windows = []
        for window, summary in (report.trend_summaries or {}).items():
            if summary.get("state") == "READY":
                ready_windows.append(
                    f"<li>过去 {html.escape(window)} 天：对比 {summary.get('dataset_count', 0)} 次 Dataset；风险 {summary.get('finding_count_baseline', 0)} → {summary.get('finding_count_current', 0)}，新增 {summary.get('new_finding_count', 0)}，已解决 {summary.get('resolved_finding_count', 0)}，持续 {summary.get('persistent_finding_count', 0)}，未核实 {summary.get('unverified_finding_count', 0)}（未核实不计作已解决）。</li>"
                )
        if ready_windows:
            trend_block = f"<h3>历史风险趋势</h3><ul>{''.join(ready_windows)}</ul>"
        category_sections = []
        for category in DeepCategory:
            findings = [item for item in report.findings if item.category == category]
            scope = next((item for item in report.analysis_scope if item.category == category), None)
            if not findings and scope is None:
                continue
            cards = "".join(self._finding_html(item) for item in findings) or '<div class="empty">该分类暂无形成可信 Finding。</div>'
            category_sections.append(
                f'<section id="{category.value}" class="l1"><h2>{html.escape(CATEGORY_LABELS[category])}</h2>{cards}</section>'
            )
        return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>VStackLens Deep Inspection</title><style>{self._css()}</style></head><body>
<main class="shell">
<header class="cover"><div class="eyebrow">VStackLens Deep Inspection</div><h1>深度巡检客户报告</h1><p>Dataset：{html.escape(dataset.dataset_id)}</p></header>
<section id="summary" class="l0"><h2>L0 摘要</h2><p class="boundary">{html.escape(report.boundary_statement)}</p>{log_coverage_html}{log_correlation_html}
<h3>本次已完成分析范围</h3><ul>{scope_rows}</ul>
<h3>PASS 聚合</h3><table><thead><tr><th>分类</th><th>能力层</th><th>通过项</th></tr></thead><tbody>{pass_rows}</tbody></table>{trend_block}</section>
{''.join(category_sections)}
<footer>报告默认只展示摘要和可信结论；技术细节与 Evidence 可按 Finding 展开。</footer>
</main></body></html>"""

    def _finding_html(self, finding: DeepFinding) -> str:
        if _uses_evidence_table(finding):
            evidence = self._evidence_table_html(finding.evidence)
        else:
            evidence = "<ol>" + "".join(
                f"<li><code>{html.escape(item.ref)}</code><br><span>来源：{html.escape(item.source.api)}<br>{'<br>'.join(html.escape(line) for line in _evidence_summary_lines(item))}</span></li>"
                for item in finding.evidence
            ) + "</ol>" if finding.evidence else "<p>无</p>"
        details = f"""<article class="finding">
<div class="finding-head"><span class="badge">{html.escape(finding.display_priority)}</span><div><h3>{html.escape(finding.title)}</h3><span class="subtle">{html.escape(finding.confidence_class.value)} · {html.escape(finding.scope.scope_type.value)} · {html.escape(finding.finding_id)}</span></div></div>
<p>{html.escape(finding.finding or finding.fact)}</p>
<details><summary>L3 技术细节与人工复核</summary><div class="detail-grid"><div><strong>事实</strong><p>{html.escape(finding.fact)}</p></div><div><strong>人工复核指引</strong><p>{html.escape(finding.verification_guidance)}</p></div><div><strong>反证条件</strong><p>{html.escape('；'.join(finding.disconfirming_conditions) or '未记录')}</p></div><div><strong>诊断边界</strong><p>{html.escape(finding.diagnosis_boundary or '无')}</p></div></div>
<details><summary>L4 Evidence（{len(finding.evidence)}）</summary>{evidence}</details></details></article>"""
        return details

    @staticmethod
    def _evidence_table_html(items: list[Any]) -> str:
        rows = []
        for item in items:
            summary = "<br>".join(html.escape(line) for line in _evidence_summary_lines(item))
            entity = item.entity.display_ref or item.entity.stable_id
            source_window = f"{item.source.api}<br>{item.window.start} 至 {item.window.end}"
            rows.append(
                f"<tr><td>{html.escape(entity)}</td><td><code>{html.escape(item.ref)}</code></td>"
                f"<td>{html.escape(source_window).replace('&lt;br&gt;', '<br>')}<br>{summary}</td></tr>"
            )
        return (
            '<div class="evidence-table-scroll"><table class="evidence-table">'
            "<thead><tr><th>对象</th><th>Evidence Ref</th><th>来源、时间窗与摘要</th></tr></thead>"
            f"<tbody>{''.join(rows)}</tbody></table></div>"
        )

    @staticmethod
    def _css() -> str:
        return """
:root{--ink:#172033;--muted:#667085;--blue:#173b67;--line:#dbe2ea;--surface:#fff;--bg:#f5f7fa;--accent:#0f766e;--risk:#9f1239}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font-family:"Microsoft YaHei","Segoe UI",Arial,sans-serif;line-height:1.6}.shell{max-width:1120px;margin:0 auto;padding:34px 28px 56px}.cover{padding-bottom:22px;border-bottom:1px solid var(--line);margin-bottom:18px}.eyebrow{font-size:12px;color:var(--muted);letter-spacing:.08em;text-transform:uppercase}h1{margin:4px 0;font-size:30px;color:var(--blue)}h2{margin:0 0 14px;color:var(--blue);font-size:21px}h3{margin:15px 0 7px;font-size:16px;color:var(--blue)}p{margin:6px 0}.l0,.l1{background:var(--surface);border:1px solid var(--line);border-radius:10px;padding:20px;margin:0 0 16px}.boundary{background:#f0fdfa;border:1px solid #99d5cc;padding:11px 13px;color:#134e4a}.subtle{color:var(--muted);font-size:12px}table{width:100%;border-collapse:collapse;margin:8px 0 10px}th,td{border:1px solid var(--line);padding:9px 10px;text-align:left;vertical-align:top;font-size:13px}th{background:#eef3f8;color:var(--blue)}.evidence-table-scroll{overflow-x:auto}.evidence-table td{font-size:12px;padding:7px 8px}.evidence-table code{font-size:10px}.finding{border:1px solid var(--line);border-radius:8px;padding:14px 16px;margin:12px 0;background:#fff}.finding-head{display:flex;gap:11px;align-items:flex-start}.finding h3{margin:0 0 2px}.badge{display:inline-flex;min-width:42px;justify-content:center;border:1px solid #f3b4c2;border-radius:999px;padding:3px 8px;color:var(--risk);font-weight:700}.detail-grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px;margin:8px 0}.detail-grid>div{background:#f8fafc;border:1px solid var(--line);padding:10px;border-radius:6px}.detail-grid strong{font-size:12px;color:var(--muted)}details{margin-top:10px;border-top:1px solid var(--line);padding-top:8px}summary{cursor:pointer;font-weight:700;color:var(--blue)}code{font-size:12px;word-break:break-all}footer{margin-top:18px;color:var(--muted);font-size:12px}@media(max-width:720px){.shell{padding:20px 14px 36px}.detail-grid{grid-template-columns:1fr}h1{font-size:25px}}
"""


class DeepWordReportBuilder:
    def render(self, report: DeepReport, output_path: Path, dataset: DeepDataset) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        self._setup(document)
        title = document.add_paragraph(style="Title")
        title.add_run("VStackLens Deep Inspection")
        subtitle = document.add_paragraph()
        subtitle_run = subtitle.add_run("深度巡检客户报告")
        subtitle_run.bold = True
        subtitle_run.font.size = Pt(18)
        subtitle_run.font.color.rgb = RGBColor(64, 111, 166)
        document.add_paragraph(f"Dataset：{dataset.dataset_id}")
        document.add_heading("摘要", level=1)
        document.add_paragraph(report.boundary_statement)
        log_coverage_note = _log_coverage_note(dataset)
        if log_coverage_note:
            document.add_paragraph(log_coverage_note)
        log_correlation_lines = _log_correlation_lines(dataset)
        if log_correlation_lines:
            document.add_heading("日志时间关联（共现不代表因果）", level=2)
            for line in log_correlation_lines:
                document.add_paragraph(line, style="List Bullet")
        document.add_heading("本次已完成分析范围", level=2)
        if report.analysis_scope:
            for item in report.analysis_scope:
                document.add_paragraph(f"{item.label}：已完成 {', '.join(item.completed_capabilities)}；可信 Finding {item.finding_count} 项；通过检查 {item.pass_count} 项。", style="List Bullet")
        else:
            document.add_paragraph(_no_analysis_message(dataset))
        document.add_heading("PASS 聚合", level=2)
        self._table(document, ["分类", "能力层", "通过项"], [[CATEGORY_LABELS[DeepCategory(item["category"])], item["profile"], item["passed"]] for item in report.pass_summary] or [["无可展示的通过项聚合", "", ""]])
        for category in DeepCategory:
            findings = [item for item in report.findings if item.category == category]
            if not findings:
                continue
            category_heading = document.add_heading(CATEGORY_LABELS[category], level=1)
            category_heading.paragraph_format.keep_with_next = True
            for finding in findings:
                finding_heading = document.add_heading(finding.title, level=2)
                finding_heading.paragraph_format.keep_with_next = True
                document.add_paragraph(finding.finding or finding.fact)
                document.add_paragraph(f"可信度：{finding.confidence_class.value}；Finding ID：{finding.finding_id}")
                document.add_heading("人工复核指引", level=3)
                document.add_paragraph(finding.verification_guidance)
                if finding.diagnosis_boundary:
                    document.add_paragraph(f"诊断边界：{finding.diagnosis_boundary}")
        ready_windows = [summary for summary in (report.trend_summaries or {}).values() if summary.get("state") == "READY"]
        if ready_windows:
            document.add_heading("趋势摘要", level=1)
            for summary in ready_windows:
                document.add_paragraph(
                    f"过去 {summary.get('window_days', 0)} 天：已对比 {summary.get('dataset_count', 0)} 次同环境 Dataset。风险数量从 {summary.get('finding_count_baseline', 0)} 变为 {summary.get('finding_count_current', 0)}，"
                    f"新增 {summary.get('new_finding_count', 0)}，已解决 {summary.get('resolved_finding_count', 0)}，持续 {summary.get('persistent_finding_count', 0)}，未核实 {summary.get('unverified_finding_count', 0)}（未核实不计作已解决）。"
                )
        if report.findings:
            appendix_heading = document.add_heading("Evidence 附录", level=1)
            appendix_heading.paragraph_format.keep_with_next = True
            for finding in report.findings:
                finding_heading = document.add_heading(finding.title, level=2)
                finding_heading.paragraph_format.keep_with_next = True
                if _uses_evidence_table(finding):
                    document.add_paragraph(f"共 {len(finding.evidence)} 条逐对象证据；完整记录见对应 Dataset 引用。")
                    self._append_evidence_table(document, finding.evidence)
                else:
                    for item in finding.evidence:
                        paragraph = document.add_paragraph()
                        paragraph.paragraph_format.keep_together = True
                        paragraph.paragraph_format.keep_with_next = True
                        ref_run = paragraph.add_run("Evidence Ref\n")
                        ref_run.bold = True
                        ref_run.font.color.rgb = RGBColor(23, 59, 103)
                        ref_value = paragraph.add_run(item.ref)
                        ref_value.font.name = "Consolas"
                        ref_value.font.size = Pt(8)
                        details = [
                            f"来源：{item.source.api}",
                            f"窗口：{item.window.start} 至 {item.window.end}",
                            *_evidence_summary_lines(item),
                        ]
                        for index, line in enumerate(details):
                            detail_paragraph = document.add_paragraph(str(line))
                            detail_paragraph.paragraph_format.keep_together = True
                            detail_paragraph.paragraph_format.keep_with_next = index < len(details) - 1
                            detail_paragraph.paragraph_format.space_before = Pt(0)
                            detail_paragraph.paragraph_format.space_after = Pt(0)
        document.save(output_path)
        return output_path

    @staticmethod
    def _append_evidence_table(document: Document, items: list[Any]) -> None:
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.autofit = False
        widths = (Cm(2.9), Cm(5.1), Cm(8.4))
        headers = ("对象", "Evidence Ref", "来源、时间窗与摘要")
        for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers, strict=True)):
            cell.width = widths[index]
            cell.text = header
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor(23, 59, 103)
                run.font.size = Pt(8)
        header_properties = table.rows[0]._tr.get_or_add_trPr()
        repeat_header = OxmlElement("w:tblHeader")
        repeat_header.set(qn("w:val"), "true")
        header_properties.append(repeat_header)
        for item in items:
            cells = table.add_row().cells
            values = (
                item.entity.display_ref or item.entity.stable_id,
                item.ref,
                f"来源：{item.source.api}\n窗口：{item.window.start} 至 {item.window.end}\n" + "\n".join(_evidence_summary_lines(item)),
            )
            row_properties = table.rows[-1]._tr.get_or_add_trPr()
            row_properties.append(OxmlElement("w:cantSplit"))
            for index, (cell, value) in enumerate(zip(cells, values, strict=True)):
                cell.width = widths[index]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                cell.text = str(value)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    for run in paragraph.runs:
                        run.font.size = Pt(7.5 if index == 1 else 8)
                        if index == 1:
                            run.font.name = "Consolas"

    @staticmethod
    def _setup(document: Document) -> None:
        section = document.sections[0]
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        styles = document.styles
        styles["Normal"].font.name = "Microsoft YaHei"
        styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles["Normal"].font.size = Pt(10.5)

    @staticmethod
    def _table(document: Document, headers: list[str], rows: list[list[Any]]) -> None:
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = str(header)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor(23, 59, 103)
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
