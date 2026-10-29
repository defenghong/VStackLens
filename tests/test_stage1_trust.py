from __future__ import annotations

import json
import sqlite3
import ssl
import zipfile
from pathlib import Path

import pytest
from docx import Document

from vstacklens.application.artifact_safety import UnsafeArtifactError, ensure_artifacts_safe, scan_artifacts
from vstacklens.application.artifact_publish import staged_file
from vstacklens.collection.diagnostics import exception_details, exception_category
from vstacklens.collection.pyvmomi_collector import PyVmomiCollector
from vstacklens.deep.analyzer import DeepAnalyzer
from vstacklens.deep.backup import restore_history_backup
from vstacklens.deep.contracts import Capability, CapabilityMatrix, CapabilityStatus, DatasetManifest, DatasetRecord, DeepCategory, DeepDataset, DeepEntity, DeepRuleDefinition, DeepSource, DeepWindow
from vstacklens.deep.history import DeepHistoryStore
from vstacklens.deep.policies import default_threshold_registry
from vstacklens.deep.provenance import policy_snapshot


RULE = DeepRuleDefinition(rule_id="STAGE1-OBJECT", rule_version="1.0", title="对象复检", category=DeepCategory.CURRENT_RISK,
                          dimension_key="state", requires=["inventory.current"], threshold_id="TH-SNAPSHOT-CHAIN-DEPTH",
                          verification_guidance="对原对象复核采集值。")


def sample(index: int, states: dict[str, bool], *, recorded=True) -> DeepDataset:
    stamp = f"2026-01-{index:02d}T00:00:00Z"
    dataset_id = f"stage1-{index}"
    dataset = DeepDataset(manifest=DatasetManifest(dataset_id=dataset_id, created_at_utc=stamp, target={"vcenter_ref": "offline-environment"}),
                          capability=CapabilityMatrix(probed_at_utc=stamp, capabilities=[Capability(id="inventory.current", name="对象清单", status=CapabilityStatus.AVAILABLE)]))
    if recorded:
        dataset.manifest.analysis_policy = policy_snapshot((RULE,), default_threshold_registry())
    for entity_id, finding in states.items():
        dataset.records.append(DatasetRecord(record_id=f"record-{entity_id}", dataset_id=dataset_id, kind="config",
                                            entity=DeepEntity(type="VirtualMachine", stable_id=entity_id, display_ref=entity_id),
                                            collected_at_utc=stamp, source=DeepSource(api="offline.fixture"),
                                            window=DeepWindow(start=stamp, end=stamp, sample_count=1), value=4 if finding else 0,
                                            finding=finding, summary="离线样本", metadata={"rule_id": RULE.rule_id}))
    return dataset


def record(store, dataset):
    analysis = DeepAnalyzer().analyze(dataset, rules=None if dataset.manifest.analysis_policy else (RULE,))
    store.record(dataset, analysis)
    return analysis


def entry(store, finding_id):
    return next(item for item in store.lifecycle_snapshot("offline-environment")["entries"] if item["finding_id"] == finding_id)


@pytest.mark.parametrize("case,reason", [("absent", "OBJECT_ABSENT"), ("not_rechecked", "NOT_RECHECKED"), ("incomplete", "INSUFFICIENT_DATA"), ("permission", "PERMISSION_DENIED"), ("not_applicable", "NOT_RECHECKED")])
def test_original_object_is_required_to_verify_resolution(tmp_path, case, reason):
    store = DeepHistoryStore(tmp_path / "history.db")
    first = record(store, sample(1, {"object-B": True})).findings[0]
    current = sample(2, {"object-A": False})
    if case != "absent":
        original = sample(2, {"object-B": False}).records[0]
        if case == "not_rechecked":
            original.kind = "inventory"
        elif case == "incomplete":
            original.window.completeness = 0.1
        elif case == "permission":
            current.capability.capabilities[0].status = CapabilityStatus.UNAVAILABLE
            current.capability.capabilities[0].detail = {"reason": "permission_denied", "missing_privilege_id": "System.Read"}
        elif case == "not_applicable":
            original.metadata["not_applicable"] = True
        current.records.append(original)
    record(store, current)
    item = entry(store, first.finding_id)
    assert item["status"] == "OPEN"
    assert item["verification_status"] == "UNVERIFIED"
    assert item["verification_reason"] == reason


def test_original_object_can_resolve_while_another_object_still_fails(tmp_path):
    store = DeepHistoryStore(tmp_path / "history.db")
    first = record(store, sample(1, {"object-B": True})).findings[0]
    result = record(store, sample(2, {"object-A": True, "object-B": False}))
    assert result.rule_results[0].status.value == "FINDING"
    resolved = entry(store, first.finding_id)
    assert resolved["verification_status"] == "CONFIRMED_RESOLVED"
    assert resolved["resolution_evidence"][0]["scope"]["scope_id"] == "object-B"
    assert resolved["resolution_evidence"][0]["evidence_refs"]


def test_recorded_policy_is_used_even_when_current_defaults_change(monkeypatch):
    dataset = sample(1, {"object-B": True})
    changed = default_threshold_registry()
    changed.get("TH-SNAPSHOT-CHAIN-DEPTH").value = 99
    monkeypatch.setattr("vstacklens.deep.provenance.default_threshold_registry", lambda: changed)
    result = DeepAnalyzer().analyze(dataset)
    assert result.findings[0].threshold.value == 3
    assert result.diagnostic.analysis_provenance["mode"] == "RECORDED_POLICY"


def test_policy_tampering_is_not_silently_reinterpreted():
    dataset = sample(1, {"object-B": True})
    dataset.manifest.analysis_policy["rules"][0]["rule_version"] = "changed"
    with pytest.raises(ValueError, match="integrity"):
        DeepAnalyzer().analyze(dataset)


def test_legacy_reinterpretation_is_labeled_and_cannot_resolve(tmp_path):
    store = DeepHistoryStore(tmp_path / "history.db")
    first = record(store, sample(1, {"object-B": True}, recorded=False)).findings[0]
    result = record(store, sample(2, {"object-B": False}))
    assert entry(store, first.finding_id)["verification_reason"] == "ORIGINAL_POLICY_MISSING"
    legacy = DeepAnalyzer().analyze(sample(3, {"object-B": True}, recorded=False), rules=(RULE,))
    assert legacy.report.analysis_provenance["mode"] == "LEGACY_REINTERPRETATION"
    assert "不替代历史原始结论" in legacy.report.boundary_statement
    assert result.rule_results[0].status.value == "PASS"


def test_reanalysis_does_not_overwrite_original_or_latest_lifecycle(tmp_path):
    store = DeepHistoryStore(tmp_path / "history.db")
    old = sample(1, {"object-B": True})
    first = record(store, old).findings[0]
    record(store, sample(2, {"object-B": False}))
    changed = default_threshold_registry()
    changed.get("TH-SNAPSHOT-CHAIN-DEPTH").value = 99
    replay = DeepAnalyzer().analyze(old, thresholds=changed)
    store.record(old, replay)
    assert replay.diagnostic.analysis_provenance["mode"] == "EXPLICIT_REINTERPRETATION"
    assert entry(store, first.finding_id)["verification_status"] == "CONFIRMED_RESOLVED"
    old_view = store.lifecycle_snapshot("offline-environment", as_of=old.manifest.created_at_utc)
    assert old_view["entries"][0]["status"] == "OPEN"
    assert "resolved_dataset" not in old_view["entries"][0]
    with sqlite3.connect(store.path) as conn:
        raw = conn.execute("SELECT payload_json FROM deep_finding_observations WHERE dataset_id=?", (old.dataset_id,)).fetchone()[0]
        assert json.loads(raw)["threshold"]["value"] == 3
        assert conn.execute("SELECT status FROM deep_rule_evaluations WHERE dataset_id=?", (old.dataset_id,)).fetchone()[0] == "FINDING"
        assert conn.execute("SELECT COUNT(*) FROM deep_analysis_runs").fetchone()[0] == 3


def test_rule_version_change_cannot_hide_old_risk(tmp_path):
    store = DeepHistoryStore(tmp_path / "history.db")
    first = record(store, sample(1, {"object-B": True})).findings[0]
    current = sample(2, {"object-B": False})
    current.manifest.analysis_policy = policy_snapshot((RULE.model_copy(update={"rule_version": "2.0"}),), default_threshold_registry())
    record(store, current)
    assert entry(store, first.finding_id)["verification_reason"] == "POLICY_CHANGED"


def test_history_queries_exclude_future_dataset_and_observation(tmp_path):
    store = DeepHistoryStore(tmp_path / "history.db")
    for index, value in ((1, 20), (2, 30), (3, 99)):
        dataset = sample(index, {"object-A": False})
        metric = dataset.records[0].model_copy(deep=True)
        metric.entity = DeepEntity(type="Datastore", stable_id="store-A")
        metric.metadata = {"rule_id": "STO-DEEP-004"}
        metric.value = value
        metric.unit = "percent"
        dataset.records.append(metric)
        record(store, dataset)
    records = store.metric_records("offline-environment", metric_id="datastore.used_percent", rule_id="CAP-DEEP-001", as_of="2026-01-02T08:00:00+08:00")
    assert [item.value for item in records] == [20, 30]
    assert store.trend_summary("offline-environment", as_of="2026-01-02T00:00:00Z")["dataset_count"] == 2
    # A later Dataset must also be excluded even if it claims an older metric timestamp.
    with sqlite3.connect(store.path) as conn:
        conn.execute("UPDATE deep_metric_observations SET observed_at_utc='2026-01-01T12:00:00Z' WHERE dataset_id='stage1-3'")
    records = store.metric_records("offline-environment", metric_id="datastore.used_percent", rule_id="CAP-DEEP-001", as_of="2026-01-02T00:00:00Z")
    assert [item.value for item in records] == [20, 30]


def legacy_database(path):
    store = DeepHistoryStore(path)
    record(store, sample(1, {"object-B": True}))
    with sqlite3.connect(path) as conn:
        conn.execute("DROP TABLE deep_analysis_runs")
    return path


def test_migration_consistent_backup_can_restore_wal_data(tmp_path):
    path = legacy_database(tmp_path / "legacy.db")
    with sqlite3.connect(path) as writer:
        writer.execute("PRAGMA journal_mode=WAL")
        writer.execute("UPDATE deep_finding_history SET status='RESOLVED'")
        writer.commit()
        upgraded = DeepHistoryStore(path)
        assert upgraded.migration_backup.is_file()
        restored = restore_history_backup(upgraded.migration_backup, tmp_path / "restored.db")
        with sqlite3.connect(restored) as conn:
            assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            assert conn.execute("SELECT status FROM deep_finding_history").fetchone()[0] == "RESOLVED"
            assert not conn.execute("SELECT name FROM sqlite_master WHERE name='deep_analysis_runs'").fetchone()
        # Raw old records remain, but unsupported old resolutions aren't presented as verified.
        assert upgraded.lifecycle_snapshot("offline-environment")["entries"][0]["verification_status"] == "UNVERIFIED"
        with pytest.raises(FileExistsError):
            restore_history_backup(upgraded.migration_backup, restored)


def test_failed_backup_stops_migration(tmp_path, monkeypatch):
    path = legacy_database(tmp_path / "legacy.db")
    def reject(*_):
        raise RuntimeError("restore validation failed")
    monkeypatch.setattr("vstacklens.deep.history.backup_before_migration", reject)
    with pytest.raises(RuntimeError, match="restore validation"):
        DeepHistoryStore(path)
    with sqlite3.connect(path) as conn:
        assert not conn.execute("SELECT name FROM sqlite_master WHERE name='deep_analysis_runs'").fetchone()
        assert conn.execute("SELECT COUNT(*) FROM deep_finding_observations").fetchone()[0] == 1


@pytest.mark.parametrize("error,category", [(ImportError("missing SDK"), "dependency_missing"), (TimeoutError("RPC timed out"), "timeout"), (ssl.SSLCertVerificationError("certificate verify failed"), "certificate_verification_failed"), (RuntimeError("RPC returned malformed object"), "api_error")])
def test_api_failures_keep_distinct_reason_and_sanitized_details(error, category):
    assert exception_category(error) == category
    details = exception_details(error)
    assert details["exception_repr"] and details["traceback"]
    assert details["category"] == category


def test_partial_standard_api_failures_keep_all_details_without_credentials():
    collector = PyVmomiCollector("offline.invalid", "admin@offline", "dummy-secret")
    for context, error in (("Health", RuntimeError("RPC failed for admin@offline password=dummy-secret")), ("Resync", TimeoutError("RPC timeout"))):
        try:
            raise error
        except Exception as exc:
            collector._record_collection_warning(context, exc)
    details = collector._collection_diagnostics
    assert len(details) == 2
    assert [item["category"] for item in details] == ["api_error", "timeout"]
    assert all(item["exception_repr"] and item["traceback"] for item in details)
    assert "dummy-secret" not in json.dumps(details)
    assert "admin@offline" not in json.dumps(details)


def test_pdf_content_is_actually_scanned(tmp_path):
    import pymupdf
    path = tmp_path / "secret.pdf"
    with pymupdf.open() as document:
        document.new_page().insert_text((30, 30), "api_key=opaque-synthetic-credential")
        document.save(path)
    assert scan_artifacts([path], secrets=("opaque-synthetic-credential",))
    with pytest.raises(UnsafeArtifactError):
        ensure_artifacts_safe([path], secrets=("opaque-synthetic-credential",))


@pytest.mark.parametrize("suffix", ["pdf", "docx", "zip"])
def test_unreadable_artifact_is_blocked_not_passed(tmp_path, suffix):
    path = tmp_path / f"broken.{suffix}"
    path.write_bytes(b"not a valid file")
    assert any("无法扫描" in item.reason for item in scan_artifacts([path]))
    with pytest.raises(UnsafeArtifactError):
        ensure_artifacts_safe([path])


def test_word_headers_footers_and_split_runs_are_scanned(tmp_path):
    path = tmp_path / "header.docx"
    document = Document()
    paragraph = document.sections[0].header.paragraphs[0]
    paragraph.add_run("opaque-synthetic-")
    paragraph.add_run("credential")
    document.sections[0].footer.paragraphs[0].text = "password=hidden-value"
    document.save(path)
    findings = scan_artifacts([path], secrets=("opaque-synthetic-credential",))
    assert any(item.reason == "包含本次会话凭据" for item in findings)
    with pytest.raises(UnsafeArtifactError):
        ensure_artifacts_safe([path])


def test_nested_broken_pdf_is_not_ignored(tmp_path):
    path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("report.pdf", b"broken")
    with pytest.raises(UnsafeArtifactError):
        ensure_artifacts_safe([path])


def test_sensitive_output_never_replaces_existing_delivery(tmp_path):
    path = tmp_path / "original.html"
    path.write_text("previous clean report", encoding="utf-8")
    with pytest.raises(UnsafeArtifactError):
        with staged_file(path) as stage:
            stage.write_text("password=synthetic-only", encoding="utf-8")
    assert path.read_text(encoding="utf-8") == "previous clean report"


def test_service_replaying_old_dataset_never_imports_later_evidence(tmp_path):
    from vstacklens.deep.service import DeepInspectionService
    service = DeepInspectionService()
    datasets = []
    for index, value in ((1, 20), (2, 99)):
        dataset = sample(index, {"object-B": bool(index == 1)})
        metric = dataset.records[0].model_copy(deep=True)
        metric.entity = DeepEntity(type="Datastore", stable_id="store-A")
        metric.metadata = {"rule_id": "STO-DEEP-004"}
        metric.value = value
        metric.unit = "percent"
        dataset.records.append(metric)
        datasets.append(dataset)
        service._analyze_dataset_object(dataset, tmp_path / f"report-{index}", create_offline_package=False)
    old = datasets[0]
    original = old.model_dump_json()
    replay = service._analyze_dataset_object(old, tmp_path / "replayed-old", create_offline_package=False)
    assert old.model_dump_json() == original
    assert replay.analysis.diagnostic.finding_lifecycle_history["entries"][0]["status"] == "OPEN"
    metric_files = list(tmp_path.glob("replayed-old.deep-dataset-*/history/metrics/*.ndjson"))
    assert metric_files
    values = [json.loads(line)["value"] for path in metric_files for line in path.read_text(encoding="utf-8").splitlines()]
    assert 99 not in values
    assert 20 in values


def test_dataset_id_collision_is_rejected_without_touching_original(tmp_path):
    store = DeepHistoryStore(tmp_path / "history.db")
    first = sample(1, {"object-B": True})
    record(store, first)
    changed = first.model_copy(deep=True)
    changed.records[0].finding = False
    with pytest.raises(ValueError, match="original evidence"):
        record(store, changed)
    changed = first.model_copy(deep=True)
    changed.manifest.target = {"vcenter_ref": "another-environment"}
    with pytest.raises(ValueError, match="environment"):
        record(store, changed)
    assert store.lifecycle_snapshot("offline-environment")["entries"][0]["status"] == "OPEN"


def test_same_timestamp_dataset_replay_is_scoped_to_original_dataset(tmp_path):
    store = DeepHistoryStore(tmp_path / "history.db")
    old = sample(1, {"object-B": True})
    later = sample(2, {"object-B": False})
    later.manifest.created_at_utc = old.manifest.created_at_utc
    for dataset, value in ((old, 20), (later, 99)):
        metric = dataset.records[0].model_copy(deep=True)
        metric.entity = DeepEntity(type="Datastore", stable_id="store-A")
        metric.collected_at_utc = old.manifest.created_at_utc
        metric.metadata = {"rule_id": "STO-DEEP-004"}
        metric.value = value
        dataset.records.append(metric)
        record(store, dataset)
    point = old.manifest.created_at_utc
    records = store.metric_records("offline-environment", metric_id="datastore.used_percent", rule_id="CAP-DEEP-001", as_of=point, as_of_dataset_id=old.dataset_id)
    assert [item.value for item in records] == [20]
    assert store.lifecycle_snapshot("offline-environment", as_of=point, as_of_dataset_id=old.dataset_id)["entries"][0]["status"] == "OPEN"
    assert store.trend_summary("offline-environment", as_of=point, as_of_dataset_id=old.dataset_id)["dataset_count"] == 1


def test_different_analysis_engine_is_explicit_reinterpretation():
    from vstacklens.deep.provenance import fingerprint
    dataset = sample(1, {"object-B": True})
    policy = dataset.manifest.analysis_policy
    policy["analyzer_revision"] = "prior-engine"
    policy["fingerprint"] = fingerprint({key: value for key, value in policy.items() if key != "fingerprint"})
    result = DeepAnalyzer().analyze(dataset)
    assert result.diagnostic.analysis_provenance["mode"] == "ENGINE_REINTERPRETATION"
    assert result.diagnostic.analysis_provenance["resolution_eligible"] is False
    assert "不替代历史原始结论" in result.report.boundary_statement


def test_offline_matrix_uses_recorded_rule_definition():
    from vstacklens.deep.offline_package import DeepOfflineEvidencePackager
    dataset = sample(1, {"object-B": True})
    historical_rule = RULE.model_copy(update={"title": "原规则名称", "verification_guidance": "原复核方法"})
    dataset.manifest.analysis_policy = policy_snapshot((historical_rule,), default_threshold_registry())
    result = DeepAnalyzer().analyze(dataset)
    matrix = DeepOfflineEvidencePackager._rule_matrix(dataset, result)
    assert "原复核方法" in json.dumps(matrix, ensure_ascii=False)
    assert "原规则名称" in json.dumps(matrix, ensure_ascii=False)


def test_historical_signal_disappearance_does_not_confirm_repair(tmp_path):
    store = DeepHistoryStore(tmp_path / "history.db")
    historical = RULE.model_copy(update={"category": DeepCategory.HISTORICAL_HEALTH})
    datasets = [sample(1, {"object-B": True}), sample(2, {"object-B": False})]
    for dataset in datasets:
        dataset.manifest.analysis_policy = policy_snapshot((historical,), default_threshold_registry())
        result = record(store, dataset)
    item = store.lifecycle_snapshot("offline-environment")["entries"][0]
    assert result.rule_results[0].status.value == "PASS"
    assert item["verification_status"] == "UNVERIFIED"
    assert item["verification_reason"] == "HISTORICAL_NOT_REPEATED"


def test_unscannable_embedded_word_media_blocks_delivery(tmp_path):
    path = tmp_path / "opaque.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", "<document>Readable normal text</document>")
        archive.writestr("word/media/unknown.png", b"opaque image content")
    with pytest.raises(UnsafeArtifactError):
        ensure_artifacts_safe([path])


def test_exception_redaction_covers_basic_auth_and_url_credentials():
    from vstacklens.application.connection_probe import redact_sensitive
    error = RuntimeError("Authorization: Basic YWRtaW46c3ludGhldGlj https://demo:encoded%40value@offline.invalid/sdk")
    details = exception_details(error)
    public_detail = redact_sensitive(error)
    for text in (json.dumps(details), public_detail):
        assert "YWRtaW46c3ludGhldGlj" not in text
        assert "encoded%40value" not in text
