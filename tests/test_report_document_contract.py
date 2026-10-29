from copy import deepcopy
from pathlib import Path
import json

import pytest

from vstacklens.reports.document_model import COLORS, get_document, select_representatives, _entry, _table
from vstacklens.reports.document_pdf import pdf_payload
from vstacklens.reports.document_html import html_page, JS, CSS
from vstacklens.reports.customer_format import capacity_text, local_time_text


@pytest.fixture
def small_report():
    from test_pdf_report import _report
    return _report()


def test_document_is_read_only_and_renderers_receive_independent_copies(small_report):
    report, context=small_report
    inputs=deepcopy(report.caliber_inputs);assets=deepcopy(report.asset_inventory);original=deepcopy(report.model_dump())
    one=get_document(report,context);saved=deepcopy(one)
    one['counts']['objects']=-1;one['appendices'].clear()
    two=get_document(report,context)
    assert two==saved
    assert report.caliber_inputs==inputs and report.asset_inventory==assets
    assert report.model_dump()==original
    assert two['title']==report.report_info.report_title
    assert [s['id'] for s in two['sections']]==[s for s in ('overview','problems','passed','environment') if s in {x['id'] for x in two['sections']}]


def test_pdf_is_concise_but_full_appendices_and_references_are_not_deleted(small_report):
    report,context=small_report;doc=get_document(report,context);before=deepcopy(doc)
    payload=pdf_payload(doc)
    assert 'appendices' not in payload and 'full_navigation' not in payload
    assert doc==before and doc['appendices'][0]['letter']=='A'
    assert [a['letter'] for a in doc['appendices']]==[chr(65+i) for i in range(len(doc['appendices']))]
    for section in doc['sections']:
        if section['kind']=='problems':
            for item in section['items']:
                appendix=next(a for a in doc['appendices'] if a['letter']==item['appendix'])
                assert len(appendix['rows'])==item['objects']
                assert len(item['representative_keys'])<=5


@pytest.mark.parametrize('count',[5,6])
def test_representatives_are_stable_worst_age_first_without_mutating_input(count):
    objects=[{'object_type':'VirtualMachine','object_key':f'vm{x}','name':f'对象{count-x}'} for x in range(count)]
    assets={('VirtualMachine',f'vm{x}'):{'properties':{'snapshot_age_days_max':x+1}} for x in range(count)}
    original=deepcopy(objects)
    a,_=select_representatives('VM-SNAPSHOT',objects,assets)
    b,_=select_representatives('VM-SNAPSHOT',list(reversed(objects)),assets)
    assert a==b and a[0]['object_key']==f'vm{count-1}' and len(a)==min(5,count)
    assert objects==original


def test_reservations_precede_limits_and_names_determine_order():
    objects=[{'object_type':'VirtualMachine','object_key':s,'name':s} for s in ['z-limit','b-reserve','a-reserve']]
    assets={('VirtualMachine',o['object_key']):{'properties':{'cpu_reservation_mhz':100 if 'reserve' in o['name'] else 0}} for o in objects}
    result,_=select_representatives('VM-RESOURCE-CONFIG',objects,assets)
    assert [x['name'] for x in result]==['a-reserve','b-reserve','z-limit']


def test_real_zero_is_not_an_empty_column_and_table_namespace_prevents_row_collision():
    values=[_entry('host1','host',[0,'—',''],category='环境明细')]
    one=_table('host-table','主机',['值','缺字段','空值'],values)
    two=_table('security-table','安全',['值','缺字段','空值'],values)
    assert one['headers']==['值'] and one['rows'][0]['cells']==['0']
    assert one['rows'][0]['id']!=two['rows'][0]['id']
    assert _table('empty','empty',['值'],[]) is None


def test_colors_and_rounding_match_user_pinned_values():
    assert COLORS=={'P1':'#B3261E','P2':'#A15C07','P3':'#3F6688'}
    assert capacity_text(931.25)=='931.3 GB'
    assert local_time_text('2031-05-25')=='2031-05-25'


def test_projection_reuse_is_scoped_and_reports_do_not_share_mutable_rows(small_report):
    from vstacklens.reports.customer_caliber import approved_result_projection,document_projection_scope
    report,context=small_report
    a=approved_result_projection(report,context)
    with document_projection_scope(report,context):
        b=approved_result_projection(report,context);c=approved_result_projection(report,context)
        assert b is c and b==a
    after=approved_result_projection(report,context)
    assert after is not b and after==a
    if after['object_results']:
        after['object_results'][0]['result_status']='changed-outside-scope'
        assert approved_result_projection(report,context)==a


def test_document_cache_is_invalidated_when_a_recorded_asset_changes(small_report):
    report,context=small_report
    before=get_document(report,context)
    report.asset_inventory['details']['Datastore'][0]['object_name']='新数据存储名称'
    after=get_document(report,context)
    assert before!=after
    assert any('新数据存储名称' in str(a) for a in after['appendices'])


def test_invalid_date_and_missing_customer_do_not_leave_empty_labels(small_report):
    report,context=small_report;report.report_info.generated_at='';report.customer_info.customer_name='未指定客户'
    doc=get_document(report,context)
    assert doc['customer']=='' and doc['date_label']=='' and not doc['header_identity'].startswith('·')


def test_new_html_payload_still_drives_existing_report_center(small_report,tmp_path):
    from vstacklens.reports.html_package import HtmlReportPackageBuilder
    from vstacklens.application.inspection_service import InspectionService
    report,context=small_report
    index,_=HtmlReportPackageBuilder().render(report,context,tmp_path/'report')
    service=object.__new__(InspectionService)
    item=service._report_center_item(index)
    assert item.title==report.report_info.report_title and item.customer_name==report.customer_info.customer_name
    assert item.asset_summary['VM']==1 and service._report_center_item_visible(item)
    assert item.customer_caliber['affected_total']==get_document(report,context)['counts']['objects']


def test_internal_task_caption_is_not_a_customer_report_title(small_report):
    report,context=small_report
    report.report_info.report_title='VC vSAN 标准巡检性能核对'
    original=deepcopy(report.model_dump());doc=get_document(report,context)
    assert doc['title']=='VStackLens 虚拟化巡检报告'
    assert '性能核对' not in doc['header_identity']
    assert report.model_dump()==original
    report.report_info.report_title='客户自定义评估报告'
    assert get_document(report,context)['title']=='客户自定义评估报告'


def test_p4_metadata_is_omitted_without_changing_any_saved_risk(small_report,monkeypatch):
    from dataclasses import replace
    from vstacklens.reports import document_model
    report,context=small_report;report.risk_summary.P4=9;before=deepcopy(report.model_dump())
    real=document_model.get_customer_problem_view
    def captured(*args,**kwargs):
        view=real(*args,**kwargs);caliber=deepcopy(view.caliber)
        caliber['counts']['P4']=0;caliber['affected_counts']['P4']=9
        return replace(view,caliber=caliber)
    monkeypatch.setattr(document_model,'get_customer_problem_view',captured)
    doc=get_document(report,context);meta=doc['metadata']['report_context']['customer_caliber']
    assert set(meta['counts'])<= {'P1','P2','P3'} and set(meta['affected_counts'])<= {'P1','P2','P3'}
    assert report.model_dump()==before and report.risk_summary.P4==9


def test_supplemental_snapshot_keeps_both_name_and_time_aliases(small_report):
    from vstacklens.reports.customer_format import local_time_text
    report,context=small_report
    report.asset_inventory['details']['VirtualMachine'][0]['properties']['snapshots']=[
        {'snapshot_name':'原始快照名_2026/9/10 11:11','created_at':'2026-09-20T01:02:03.123456+00:00'}]
    doc=get_document(report,context);table=next(a for a in doc['appendices'] if a['id']=='other-snapshots')
    text=' '.join(table['rows'][0]['cells'])
    assert '原始快照名_2026/9/10 11:11' in text and local_time_text('2026-09-20T01:02:03.123456+00:00') in text


def test_customer_license_only_uses_entity_bound_assignment_name_version(small_report):
    report,context=small_report;host=report.asset_inventory['details']['HostSystem'][0]
    host['object_key']='host-license-case'
    host['properties']['host_license_name']='UNBOUND_STOCK_NOT_ACTUAL'
    host['properties']['host_license_evidence']={'schema':'entity-license-v1','entity_id':host['object_key'],
        'vcenter_instance_uuid':'SYN_VC','object_type':'HostSystem','assigned':True,'assignment_status':'assigned',
        'expiration_state':'unknown','name':'SYN_ASSIGNED','product_version':'8.0','days_remaining':None}
    original=deepcopy(report.model_dump());doc=get_document(report,context)
    table=next(a for a in doc['appendices'] if a['id']=='licenses')
    text=' '.join(table['headers']+table['rows'][0]['cells'])
    assert 'SYN_ASSIGNED' in text and '8.0' in text and '已分配' in text
    assert all(term not in text for term in ('有效期','到期','永久','999999','无法确认','UNBOUND_STOCK_NOT_ACTUAL'))
    assert report.model_dump()==original
    host['properties']['host_license_evidence']['entity_id']='different-host'
    assert not any(a['id']=='licenses' for a in get_document(report,context)['appendices'])


def test_no_security_results_do_not_produce_a_verified_matrix_or_empty_certificate_heading(small_report):
    report,context=small_report
    report.caliber_inputs={'object_results':[{'rule_id':'VSL-VC-006','risk_level':'P3','object_type':'vCenter',
        'result_status':'unavailable','object_name':'vc'}],'findings':[],'rule_checklist':[]}
    report.findings=[]
    for kind in ('vCenter','HostSystem'):
        for row in report.asset_inventory['details'][kind]:
            for key in list(row.get('properties') or {}):
                if 'certificate' in key:row['properties'].pop(key)
    doc=get_document(report,{})
    assert not any(a['id'] in {'security','certificates'} for a in doc['appendices'])
    blocks=[b for s in doc['sections'] if s['kind']=='environment' for b in s['blocks']]
    assert not any(b['id'] in {'security-summary','security-attention','vc-cert','cert-summary','host-cert-source'} for b in blocks)
    assert all(b.get('text','').strip() for b in blocks if b['kind']=='paragraph')


def test_word_graph_cell_widths_retain_original_numeric_proportions(small_report,tmp_path):
    from vstacklens.reports.docx_report import DocxReportEngine
    from docx import Document
    from docx.oxml.ns import qn
    report,context=small_report
    report.asset_inventory['details']['HostSystem'][0]['properties']['host_memory_allocation_ratio']=1.88
    other=deepcopy(report.asset_inventory['details']['HostSystem'][0]);other.update(object_key='host-axis-2',object_name='host-axis-2')
    other['properties']['host_memory_allocation_ratio']=1.28
    report.asset_inventory['details']['HostSystem'].append(other)
    doc=Document(DocxReportEngine().render(report,tmp_path/'bar.docx'));seen=0
    for table in doc.element.xpath('//w:tbl[w:tblPr/w:tblCaption[@w:val="图形几何"]]'):
        grid=[int(c.get(qn('w:w'))) for c in table.xpath('./w:tblGrid/w:gridCol')]
        actual=[int(c.get(qn('w:w'))) for c in table.xpath('./w:tr/w:tc/w:tcPr/w:tcW')]
        assert actual==grid
        previous=table.getprevious()
        label=''.join(previous.xpath('.//w:t/text()')) if previous is not None else ''
        filled=[i for i,cell in enumerate(table.xpath('./w:tr/w:tc')) if cell.xpath('./w:tcPr/w:shd[@w:fill="506B84" or @w:fill="8CA9BF"]')]
        if '188.0%' in label:
            assert abs(sum(grid[i] for i in filled)/sum(grid)-.94)<.003
            seen+=1
    assert seen>=1 # 188% on a 200% axis occupies 94%, not two equal cells out of three.


def test_html_escapes_embedded_data_and_uses_the_same_full_model(small_report):
    report,context=small_report;doc=get_document(report,context)
    doc['customer']='合成测试数据 </script><script>unexpected()</script>'
    rendered=html_page(doc)
    assert '\\u003c/script' in rendered
    assert '<script>unexpected()' not in rendered
    assert 'filter-search' in rendered and 'filter-cluster' in rendered and 'filter-category' in rendered
    assert 'VSTACKLENS_DOCUMENT_READY' in JS and 'overflow-x:auto' in CSS


def test_word_checklist_heading_has_three_rows_of_content(small_report,tmp_path):
    from docx import Document
    from vstacklens.reports.document_word import WordDocumentWriter
    from vstacklens.reports.docx_report import DocxReportEngine
    report,context=small_report;doc=get_document(report,context)
    doc['sections']=[{'id':'passed','title':'已通过的检查','kind':'checklist',
                      'items':[{'label':f'合成检查{i}'} for i in range(6)]}]
    doc['appendices']=[]
    path=tmp_path/'checklist.docx'
    WordDocumentWriter(DocxReportEngine(),report,doc).render(path)
    rows=[p for p in Document(path).paragraphs if p.text.startswith('✓ 合成检查')]
    assert len(rows)==3
    assert [p.paragraph_format.keep_with_next for p in rows]==[True,True,False]


def test_technical_alarm_acknowledgement_has_a_customer_state(small_report,monkeypatch):
    from dataclasses import replace
    from vstacklens.reports import document_model
    report,context=small_report;original=deepcopy(report.model_dump())
    real=document_model.get_customer_problem_view
    def captured(*args,**kwargs):
        view=real(*args,**kwargs);first=view.categories[0];display=deepcopy(first.display)
        display['technical']={'headers':['对象名称','确认状态'],
                              'rows':[[display['objects'][0]['name'],'False']]}
        return replace(view,categories=(replace(first,display=display),*view.categories[1:]))
    monkeypatch.setattr(document_model,'get_customer_problem_view',captured)
    doc=get_document(report,context)
    text=' '.join(cell for table in doc['appendices'] for row in table['rows'] for cell in row['cells'])
    assert '尚未应答' in text and 'False' not in text
    assert report.model_dump()==original
