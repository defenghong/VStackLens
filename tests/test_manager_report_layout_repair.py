"""Small offline layout counterexamples, independent of collection and UI."""
from copy import deepcopy
from pathlib import Path

import pymupdf
import pytest

from vstacklens.products.management_scope import mark_management_scope
from vstacklens.products.model import TargetRef, finish_module, make_check, make_source, new_module
from vstacklens.products.reporting import build_integrated_report, product_report_view
from vstacklens.reports.pdf_report import PdfReportEngine


@pytest.mark.parametrize("value", ["", "one short value", "中文与 English\n" * 150, "unbroken-token-" * 200])
def test_product_cell_parts_preserve_every_character(value):
    parts = PdfReportEngine._product_cell_parts(value, 140)
    assert "".join(parts) == value
    if len(value) > 1000:
        assert len(parts) > 1


def long_report():
    target = TargetRef("layout-nsx", "nsx", "offline-customer", "offline-site", "离线 Manager", "192.0.2.110")
    module = new_module(target)
    mark_management_scope(module)
    module["checks"] = [make_check("nsx.version", "版本", "normal", values={"version": "4.2.3"},
        sources=[make_source("/offline/version", version_basis="离线合成样本")], verification="offline_sample")]
    advice = "SYSLOG-BEGIN " + "Review the current Manager logging configuration and preserve the source evidence. " * 24 + " SYSLOG-END"
    module["native_alerts"] = [{"status": "OPEN", "severity": "MEDIUM", "object_name": "layout-object",
        "description": "Remote logging requires review", "recommendation": advice}]
    report = build_integrated_report("layout-only", [target], [finish_module(module)], "管理报告长建议离线核对")
    return report, advice


def test_product_pdf_long_row_has_identity_and_evidence_on_each_fragment(tmp_path):
    report, advice = long_report()
    original = deepcopy(report.model_dump())
    pure = product_report_view(report)
    derived = PdfReportEngine()._pdf_product_view(report)
    table = next(t for t in derived["sections"][0]["body"]["tables"] if t["title"] == "异常与需复核事项")
    assert len(table["rows"]) > 1
    assert all("layout-object" in row[0].replace("\u200b", "") and "E-" in row[-1] for row in table["rows"])
    restored = "".join(row[2] for row in table["rows"]).replace("\u200b", "")
    assert restored == advice
    assert "同一对象续段" in table["rows"][1][0]
    assert product_report_view(report) == pure
    pdf = PdfReportEngine().render(report, tmp_path / "long-row.pdf")
    with pymupdf.open(pdf) as document:
        text = "".join(page.get_text() for page in document).replace("\u200b", "")
        compact = "".join(text.split())
        for row in table["rows"]:
            assert "".join(row[2].replace("\u200b", "").split()) in compact
        # Geometry is essential: extracted invisible words also existed in the bug.
        for page in document:
            words = page.get_text("words")
            assert all(0 <= w[1] < w[3] <= page.rect.height for w in words)
            for end in page.search_for("SYSLOG-END"):
                assert end.y1 < page.rect.height - 30
    assert report.model_dump() == original


def test_product_layout_is_opt_in_and_vc_first_column_rule_is_unchanged():
    root = Path(__file__).resolve().parents[1] / "src/vstacklens/reports"
    template = (root / "templates/inspection-report.typ").read_text(encoding="utf-8")
    js = (root / "html_assets.py").read_text(encoding="utf-8")
    css = (root / "html_package.py").read_text(encoding="utf-8")
    assert 'split-cells: false' in template and 'split-cells: true' in template
    assert 'className = "", contextTitle = "", columnWeights = []' in js
    assert '"product-detail-table", contextTitle, weights' in js
    assert '.object-evidence-table th:first-child { width:55%; }' in css
    assert '"object-evidence-table issue-evidence-table"' in js


def test_product_word_long_row_may_continue_but_keeps_full_original(tmp_path):
    from docx import Document
    from vstacklens.reports.docx_report import DocxReportEngine
    report, advice = long_report()
    before = deepcopy(report.model_dump())
    path = DocxReportEngine().render(report, tmp_path / "long-row.docx")
    document = Document(path)
    row = next(row for table in document.tables for row in table.rows if advice in "".join(c.text for c in row.cells))
    assert not row._tr.xpath("./w:trPr/w:cantSplit")
    assert advice in row.cells[2].text
    normal_row = next(row for table in document.tables for row in table.rows
                      if any("版本" in c.text for c in row.cells) and not any("版本" in c.text for c in table.rows[0].cells))
    assert normal_row._tr.xpath("./w:trPr/w:cantSplit")
    assert report.model_dump() == before
