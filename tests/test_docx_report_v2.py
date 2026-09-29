from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT

from vstacklens.reports.docx_report_v2 import DocxReportV2Engine
from vstacklens.reports.report_model import FindingItem

from test_report_exporter import rich_report


def test_v2_report_has_staged_shell_and_customer_safe_text(tmp_path: Path) -> None:
    output = tmp_path / "v2-report.docx"
    DocxReportV2Engine().render(rich_report(), output)

    assert output.exists()
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    for heading in ("首页摘要", "1. 巡检报告概览", "2. 环境健康状态", "3. 检查结果汇总", "4. 重点问题详情与整改建议", "5. VMware 平台巡检", "8. 环境资产与对象清单", "9. 附录"):
        assert heading in text
    assert "Rule ID" not in text
    assert "Collector" not in text
    assert "INTERNAL_ONLY" not in text
    assert "子原因" not in text
    assert document.sections[0].orientation == WD_ORIENT.PORTRAIT


def test_v2_storage_policy_table_uses_landscape_section(tmp_path: Path) -> None:
    report = rich_report()
    report.vsan_summary = {
        "applicable": True,
        "health_status": "healthy",
        "host_count": 2,
        "disk_group_count": 2,
        "cache_disk_count": 2,
        "capacity_disk_count": 8,
        "capacity": {"used_percent": 62.5, "total_gb": 1000, "used_gb": 625, "free_gb": 375},
        "storage_policy_summary": {
            "checked_count": 1,
            "compliant_count": 1,
            "noncompliant_count": 0,
            "unknown_count": 0,
            "policy_categories": [{
                "policy_name": "vSAN 默认策略",
                "policy_uuid": "policy-1",
                "checked_count": 1,
                "compliant_count": 1,
                "noncompliant_count": 0,
                "unknown_count": 0,
                "not_applicable_count": 0,
            }],
        },
    }
    output = tmp_path / "v2-vsan-report.docx"
    DocxReportV2Engine().render(report, output)

    document = Document(output)
    landscape_sections = [section for section in document.sections if section.orientation == WD_ORIENT.LANDSCAPE]
    assert landscape_sections
    assert all(not section.different_first_page_header_footer for section in landscape_sections)


def test_v2_word_report_keeps_all_risk_and_unavailable_detail_rows(tmp_path: Path) -> None:
    report = rich_report().model_copy(deep=True)
    report.findings = [
        FindingItem(
            risk_level="P3",
            rule_id="VSL-VM-023",
            rule_name="VM on local datastore",
            title="虚拟机运行在本地存储上",
            object_type="VirtualMachine",
            object_name=f"vm-full-{index:02d}",
            status="open",
            current_value=True,
            expected_value=False,
            business_impact="本地存储可能限制迁移和故障恢复。",
            remediation="结合业务等级评估后迁移到共享 Datastore。",
            observed_detail={"local_datastore_names": [f"Local_{index:02d}"]},
        )
        for index in range(12)
    ]
    report.appendix["unavailable_rules"] = [
        {"title": f"未确认检查-{index:02d}", "object_name": f"对象-{index:02d}", "reason": "数据未返回"}
        for index in range(12)
    ]

    output = tmp_path / "v2-untruncated-details.docx"
    DocxReportV2Engine().render(report, output)
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    assert all(f"vm-full-{index:02d}" in text for index in range(12))
    assert all(f"未确认检查-{index:02d}" in text for index in range(12))
    assert "其余对象明细请结合 HTML 报告或附录继续查看" not in text
