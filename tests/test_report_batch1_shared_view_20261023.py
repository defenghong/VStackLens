"""Regression boundaries for saved identity and missing-vs-zero evidence."""
from copy import deepcopy

from docx import Document

from vstacklens.reports.customer_caliber import calculate
from vstacklens.reports.customer_view import get_customer_problem_view
from vstacklens.reports.docx_report import DocxReportEngine
from vstacklens.reports.report_model import CustomerInfo, EnvironmentSummary, HealthScore, ReportData, ReportInfo, RiskSummary


def same_name_report():
    rows = []
    assets = []
    for key, limit in [('vm-a', 3600), ('vm-b', 7200)]:
        rows.append(dict(rule_id='VSL-VM-011', rule_name='虚拟机资源限制', title='虚拟机资源限制',
                         risk_level='P3', object_type='VirtualMachine', object_key=key,
                         object_name='同名虚拟机', object_path=key, result_status='failed', status='open',
                         current_value=True, expected_value=False, evidence={'object_key':key}))
        assets.append(dict(object_type='VirtualMachine', object_key=key, object_name='同名虚拟机',
                           asset_location=key, properties={'cpu_limit_mhz':limit,'memory_limit_mb':-1}))
    caliber = calculate(rows, records=rows, assets={'VirtualMachine':assets}, asset_total=2, run_id='identity-test')
    return ReportData(report_info=ReportInfo(report_id='identity',run_id='identity-test',generated_at='2026-10-22T00:00:00Z'),
                      customer_info=CustomerInfo(customer_name='测试客户',site_name='测试环境',vcenter='vc.test'),
                      environment_summary=EnvironmentSummary(scope_statement='saved evidence',checked_object_total=2),
                      health_score=HealthScore(score=80,grade='关注',explanation='保存结果'),risk_summary=RiskSummary(),
                      customer_caliber=caliber,caliber_inputs={'findings':deepcopy(rows),'object_results':deepcopy(rows)},
                      asset_inventory={'details':{'VirtualMachine':assets},'summary':{'VirtualMachine':2},'total':2})


def test_shared_view_never_joins_same_name_objects_by_name():
    report = same_name_report()
    view = get_customer_problem_view(report)
    assert view.caliber['affected_total'] == 2
    objects = {item['object_key']:item for item in view.categories[0].display['objects']}
    assert set(objects) == {'vm-a','vm-b'}
    assert '3600' in objects['vm-a']['basis'] and '7200' not in objects['vm-a']['basis']
    assert '7200' in objects['vm-b']['basis'] and '3600' not in objects['vm-b']['basis']
    assert view.categories[0].display['technical'] is None  # No name-based legacy lookup.


def test_shared_view_cannot_mutate_saved_sources_or_other_consumer_projection():
    report = same_name_report()
    assets = deepcopy(report.asset_inventory)
    inputs = deepcopy(report.caliber_inputs)
    caliber = deepcopy(report.customer_caliber)
    view = get_customer_problem_view(report)
    first_consumer = view.projected_findings
    first_consumer[0]['object_name'] = 'changed by one consumer'
    assert all(row['object_name'] == '同名虚拟机' for row in view.projected_findings)
    assert report.asset_inventory == assets
    assert report.caliber_inputs == inputs
    assert report.customer_caliber == caliber


def test_resync_keeps_real_zeros_but_does_not_invent_missing_plan_or_eta():
    engine = DocxReportEngine()
    missing = Document()
    engine._resync_status_table(missing, {'resync_object_count':0,'resync_bytes':0})
    missing_text = ' '.join(cell.text for table in missing.tables for row in table.rows for cell in row.cells)
    assert '0 个' in missing_text and '0.0 GB' in missing_text
    assert '预计完成时间' not in missing_text and '计划重同步' not in missing_text
    real = Document()
    engine._resync_status_table(real, {'resync_object_count':0,'resync_bytes':0,'resync_eta_seconds':0,'planned_resync_count':0})
    real_text = ' '.join(cell.text for table in real.tables for row in table.rows for cell in row.cells)
    assert '预计完成时间' in real_text and '0 秒' in real_text
    assert '计划重同步' in real_text and '0 项' in real_text


def test_hidden_history_has_neither_an_empty_heading_nor_a_comparison_table():
    report = same_name_report()
    report.history_visible = False
    document = Document()
    engine = DocxReportEngine()
    engine._executive_summary(document, report, engine._aggregate_risks(report))
    assert '1.5 巡检结论' not in [paragraph.text for paragraph in document.paragraphs]
    assert all('上一轮' not in cell.text for table in document.tables for row in table.rows for cell in row.cells)
