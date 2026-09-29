from __future__ import annotations

import json
import time
import urllib.error
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace as NS
from zipfile import ZipFile

import pytest
from docx import Document
from pydantic import ValidationError

from vstacklens.deep.analyzer import BUILTIN_DEEP_RULES, DeepAnalyzer
from vstacklens.deep.contracts import (
    Capability,
    CapabilityMatrix,
    CapabilityStatus,
    DeepDataset,
    DeepEntity,
    DeepFinding,
    DeepRuleDefinition,
    DeepResultStatus,
    evidence_ref,
    DeepSource,
    DeepWindow,
    DatasetRecord,
    FindingScope,
    ScopeType,
    ThresholdPolicy,
    finding_id,
)
from vstacklens.deep.service import DeepInspectionService
from vstacklens.deep.interfaces import DeepCollectionPolicy
from vstacklens.deep.history import DeepHistoryStore
from vstacklens.deep.logs import classify_log_lines, correlate_log_records, log_category_for_descriptor, normalize_powercli_logs
from vstacklens.deep.report import _log_correlation_lines
from vstacklens.deep.merge import merge_records
from vstacklens.deep.resource_monitor import LatchedResourceGuard, LocalResourceMonitor
from vstacklens.deep.pyvmomi_collector import PyVmomiDeepCollector
from vstacklens.deep.report import DeepHtmlReportBuilder
from vstacklens.deep.vcenter_rest import VCenterRestReadClient
from vstacklens.deep.hardware_compatibility import esxi_version_to_hcl_release, load_bundled_vcg_index
from vstacklens.deep.dataset import DeepDatasetWriter
from vstacklens.deep.connection import EsxiHostConnectionInfo, VCenterConnectionInfo, load_esxi_host_connections_from_workbook, load_vcenter_connection_from_workbook


FIXTURE = Path(__file__).parent / "fixtures" / "deep_dataset_trend.json"


def test_workbook_host_credentials_load_without_exposing_repr(tmp_path: Path) -> None:
    values = ["User", "Address", "Username", "Password", "vCenter", "192.0.2.1", "vcenter-user", "vcenter-password", "esxi-1", "192.0.2.11", "root-readonly", "host-secret", "esxi-2", "192.0.2.12", "root-audit", "host-secret-2"]
    shared_strings = "<sst xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'>" + "".join(f"<si><t>{value}</t></si>" for value in values) + "</sst>"
    indices = {value: index for index, value in enumerate(values)}
    row_specs = [
        {"D1": "User", "E1": "Address", "F1": "Username", "G1": "Password"},
        {"D2": "vCenter", "E2": "192.0.2.1", "F2": "vcenter-user", "G2": "vcenter-password"},
        {"D3": "esxi-1", "E3": "192.0.2.11", "F3": "root-readonly", "G3": "host-secret"},
        {"D4": "esxi-2", "E4": "192.0.2.12", "F4": "root-audit", "G4": "host-secret-2"},
    ]
    rows_xml = []
    for row_spec in row_specs:
        row_id = next(int(ref[1:]) for ref in row_spec)
        cells = "".join(f"<c r='{ref}' t='s'><v>{indices[value]}</v></c>" for ref, value in row_spec.items())
        rows_xml.append(f"<row r='{row_id}'>{cells}</row>")
    workbook_path = tmp_path / "connections.xlsx"
    with ZipFile(workbook_path, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main' xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'><sheets><sheet name='Sheet1' sheetId='1' r:id='rId1'/></sheets></workbook>")
        archive.writestr("xl/_rels/workbook.xml.rels", "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'><Relationship Id='rId1' Target='worksheets/sheet1.xml'/></Relationships>")
        archive.writestr("xl/sharedStrings.xml", shared_strings)
        archive.writestr("xl/worksheets/sheet1.xml", "<worksheet xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'><sheetData>" + "".join(rows_xml) + "</sheetData></worksheet>")

    credentials = load_esxi_host_connections_from_workbook(workbook_path)

    assert [(item.host_alias, item.host) for item in credentials] == [("esxi-1", "192.0.2.11"), ("esxi-2", "192.0.2.12")]
    assert not any(secret in repr(item) for item in credentials for secret in (item.host, item.username, item.password))


def test_vcenter_workbook_loader_uses_primary_row_not_same_endpoint_readonly_row(tmp_path: Path) -> None:
    values = [
        "User", "Address", "Username", "Password",
        "vCenter", "192.0.2.1", "synthetic-primary-user", "synthetic-primary-secret",
        "vCenter Readonly", "synthetic-readonly-user", "synthetic-readonly-secret",
    ]
    shared_strings = "<sst xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'>" + "".join(f"<si><t>{value}</t></si>" for value in values) + "</sst>"
    indices = {value: index for index, value in enumerate(values)}
    row_specs = [
        {"D1": "User", "E1": "Address", "F1": "Username", "G1": "Password"},
        {"D2": "vCenter", "E2": "192.0.2.1", "F2": "synthetic-primary-user", "G2": "synthetic-primary-secret"},
        {"D3": "vCenter Readonly", "E3": "192.0.2.1", "F3": "synthetic-readonly-user", "G3": "synthetic-readonly-secret"},
    ]
    rows_xml = []
    for row_spec in row_specs:
        row_id = next(int(ref[1:]) for ref in row_spec)
        cells = "".join(f"<c r='{ref}' t='s'><v>{indices[value]}</v></c>" for ref, value in row_spec.items())
        rows_xml.append(f"<row r='{row_id}'>{cells}</row>")
    workbook_path = tmp_path / "vcenter-connections.xlsx"
    with ZipFile(workbook_path, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main' xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'><sheets><sheet name='Sheet1' sheetId='1' r:id='rId1'/></sheets></workbook>")
        archive.writestr("xl/_rels/workbook.xml.rels", "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'><Relationship Id='rId1' Target='worksheets/sheet1.xml'/></Relationships>")
        archive.writestr("xl/sharedStrings.xml", shared_strings)
        archive.writestr("xl/worksheets/sheet1.xml", "<worksheet xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'><sheetData>" + "".join(rows_xml) + "</sheetData></worksheet>")

    connection = load_vcenter_connection_from_workbook(workbook_path)

    assert connection.username == "synthetic-primary-user"
    assert connection.password == "synthetic-primary-secret"
    assert "synthetic-primary-secret" not in repr(connection)
    assert "synthetic-readonly-secret" not in repr(connection)


def test_vcenter_service_redacts_workbook_credentials_before_writing_outputs(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    from vstacklens.deep import service as service_module

    primary_user = "synthetic-primary-user"
    primary_password = "synthetic-primary-secret"
    host_user = "synthetic-esxi-user"
    host_password = "synthetic-esxi-secret"

    class FakeCollector:
        def __init__(self, *_args, **_kwargs):
            pass

        def collect(self, policy=None):
            dataset = DeepDataset.from_json(FIXTURE)
            record = dataset.records[0]
            record.summary = f"Synthetic log account {primary_user}"
            record.value = {"lines": [f"login={primary_user} password={primary_password} host={host_user} secret={host_password}"]}
            record.metadata["test_payload"] = f"{primary_user} {primary_password} {host_user} {host_password}"
            dataset.collection_log.append({"action": "synthetic.test", "detail": f"{primary_user} {primary_password} {host_user} {host_password}"})
            return dataset

    monkeypatch.setattr(service_module, "load_vcenter_connection_from_workbook", lambda _path: VCenterConnectionInfo(host="vc.synthetic", username=primary_user, password=primary_password))
    monkeypatch.setattr(service_module, "load_esxi_host_connections_from_workbook", lambda _path: [EsxiHostConnectionInfo(host_alias="esxi-synthetic", host="esxi.synthetic", username=host_user, password=host_password)])
    monkeypatch.setattr(service_module, "_workbook_connection_secrets", lambda _path: ([primary_user, host_user], [primary_password, host_password]))
    monkeypatch.setattr(service_module, "PyVmomiDeepCollector", FakeCollector)

    run = DeepInspectionService().collect_and_analyze_vcenter(
        tmp_path / "synthetic-connections.xlsx",
        tmp_path / "reports",
    )

    output_files = [path for path in tmp_path.rglob("*") if path.is_file()]
    sensitive_values = (primary_user, primary_password, host_user, host_password)
    leaks = [path.relative_to(tmp_path).as_posix() for path in output_files if any(secret.encode("utf-8") in path.read_bytes() for secret in sensitive_values)]
    assert run.html_path.exists()
    assert run.diagnostic_path.exists()
    assert leaks == []


def test_deep_redaction_masks_username_inside_upn_and_domain_account_tokens() -> None:
    from vstacklens.deep.redaction import redact_dataset_credentials

    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records[0].value = {
        "lines": [
            "authentication user=root@management.example",
            r"principal=LAB\root",
            "username=root-service-account",
            "authentication failed for root-session-id",
            "unrelated=administrator",
        ]
    }

    redacted = redact_dataset_credentials(dataset, usernames=["root"], passwords=[])
    payload = json.loads(redacted.model_dump_json())
    lines = payload["records"][0]["value"]["lines"]

    assert "root@management.example" not in lines[0]
    assert r"LAB\root" not in lines[1]
    assert "[REDACTED]@management.example" in lines[0]
    assert r"LAB\[REDACTED]" in lines[1]
    assert lines[2] == "username=[REDACTED]"
    assert lines[3] == "authentication failed for [REDACTED]-session-id"
    assert "administrator" in lines[4]


def test_workbook_host_credentials_map_to_unique_vcenter_hosts() -> None:
    host_one = NS(name="esx-one", _moId="host-1", config=NS(uuid="host-uuid-1", network=NS(vnic=[NS(spec=NS(ip=NS(ipAddress="192.0.2.11")))])))
    host_two = NS(name="esx-two", _moId="host-2", config=NS(uuid="host-uuid-2", network=NS(vnic=[NS(spec=NS(ip=NS(ipAddress="192.0.2.12")))])))
    credentials = [
        EsxiHostConnectionInfo(host_alias="esxi-1", host="192.0.2.11", username="root-a", password="secret-a"),
        EsxiHostConnectionInfo(host_alias="esxi-2", host="192.0.2.12", username="root-b", password="secret-b"),
        EsxiHostConnectionInfo(host_alias="stale-esxi", host="192.0.2.99", username="root-old", password="secret-old"),
    ]

    mapped, summary = PyVmomiDeepCollector._map_esxi_log_credentials([host_one, host_two], credentials)

    assert set(mapped) == {"host-1", "host-2"}
    assert mapped["host-1"] is credentials[0]
    assert summary["mapped_host_count"] == 2
    assert summary["unmatched_credential_row_count"] == 1
    assert summary["ambiguous_host_count"] == 0


def test_deep_fixture_runs_six_pilot_rules_and_trend_chain() -> None:
    dataset = DeepDataset.from_json(FIXTURE)
    result = DeepAnalyzer().analyze(dataset)

    assert len(result.findings) >= 7
    assert {item.category.value for item in result.findings} == {"current_risk", "historical_health", "trend"}
    assert result.report.trend_chain_verified is True
    trend = next(item for item in result.findings if item.category.value == "trend")
    assert trend.confidence_class.value == "STATISTICAL"
    assert trend.verification_guidance
    assert any(item.kind == "derived" for item in trend.evidence)


def test_degraded_dataset_cannot_report_pass_but_preserves_confirmed_findings() -> None:
    dataset = DeepDataset.from_json(FIXTURE)
    baseline = DeepAnalyzer().analyze(dataset)
    pass_id = next(item.rule_id for item in baseline.rule_results if item.status == DeepResultStatus.PASS)
    finding_id = baseline.findings[0].rule_id
    rules = tuple(rule for rule in BUILTIN_DEEP_RULES if rule.rule_id in {pass_id, finding_id})
    dataset.manifest = dataset.manifest.model_copy(update={"impact": {"degraded": True}})

    analysis = DeepAnalyzer().analyze(dataset, rules=rules)
    statuses = {item.rule_id: item for item in analysis.rule_results}

    assert statuses[pass_id].status == DeepResultStatus.INSUFFICIENT_DATA
    assert "collection was degraded" in statuses[pass_id].reason
    assert statuses[finding_id].status == DeepResultStatus.FINDING


def test_event_history_cap_does_not_suppress_complete_current_risk_passes() -> None:
    dataset = DeepDataset.from_json(FIXTURE)
    entity = DeepEntity(type="HostSystem", stable_id="event-host-fixture", display_ref="host-fixture")
    dataset.records.append(
        DatasetRecord(
            record_id="degraded-event-history-fixture",
            dataset_id=dataset.dataset_id,
            kind="event",
            entity=entity,
            collected_at_utc=dataset.manifest.created_at_utc,
            source=DeepSource(api="EventManager.QueryEvents", collector="fixture"),
            selector={"rule_id": "CL-DEEP-002"},
            window=DeepWindow(start=dataset.manifest.created_at_utc, end=dataset.manifest.created_at_utc, sample_count=0, expected_sample_count=0, completeness=1.0),
            value=0,
            unit="count",
            raw_pointer="events/event.ndjson#fixture",
            finding=False,
            summary="fixture contains no matching event",
            metadata={"rule_id": "CL-DEEP-002", "event_samples": []},
        )
    )
    license_rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "SEC-DEEP-004")
    event_rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "CL-DEEP-002")
    dataset.manifest = dataset.manifest.model_copy(update={"impact": {"degraded": True}})
    dataset.collection_log = [{"action": "event.history", "status": "degraded", "degraded": True, "reason": "collector_cap"}]

    analysis = DeepAnalyzer().analyze(dataset, rules=(license_rule, event_rule))
    statuses = {item.rule_id: item for item in analysis.rule_results}

    assert statuses["SEC-DEEP-004"].status == DeepResultStatus.PASS
    assert statuses["CL-DEEP-002"].status == DeepResultStatus.INSUFFICIENT_DATA
    assert "EventHistory was incomplete" in statuses["CL-DEEP-002"].reason


def test_historical_extension_rules_accept_task_and_event_records() -> None:
    dataset = DeepDataset.from_json(FIXTURE)
    now = "2026-09-21T03:12:00Z"
    entity = DeepEntity(type="ClusterComputeResource", stable_id="cluster-historical-fixture", display_ref="cluster-fixture")
    for rule_id, kind in (
        ("TASK-DEEP-001", "task"),
        ("CL-DEEP-002", "event"),
        ("CL-DEEP-003", "event"),
        ("STO-DEEP-002", "event"),
        ("NET-DEEP-002", "event"),
    ):
        dataset.records.append(
            DatasetRecord(
                record_id=f"historical-{rule_id}",
                dataset_id=dataset.dataset_id,
                kind=kind,
                entity=entity,
                collected_at_utc=now,
                source=DeepSource(api="fixture.history", collector="fixture", collected_at_utc=now),
                selector={"rule_id": rule_id},
                window=DeepWindow(start=now, end=now, sample_count=1, expected_sample_count=1, completeness=1.0),
                value=1,
                unit="count",
                raw_pointer=f"events/{rule_id}.ndjson#1",
                finding=True,
                summary=f"fixture historical evidence for {rule_id}",
                metadata={"rule_id": rule_id},
            )
        )
    result = DeepAnalyzer().analyze(dataset)
    statuses = {item.rule_id: item.status for item in result.rule_results}
    assert all(statuses[rule_id] == DeepResultStatus.FINDING for rule_id in ("TASK-DEEP-001", "CL-DEEP-002", "CL-DEEP-003", "STO-DEEP-002", "NET-DEEP-002"))


def test_task_history_requires_repeated_failure_or_time_cluster() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    base = datetime(2026, 9, 23, 1, 0, tzinfo=UTC)
    tasks = [
        NS(info=NS(state="error", name="RelocateVM_Task", entityName=f"vm-{index}", completeTime=base + timedelta(minutes=index)))
        for index in range(3)
    ]
    record = collector._task_history_records(
        NS(taskManager=NS(recentTask=tasks)),
        "ds-task-history",
        "2026-09-23T01:05:00Z",
    )[0]
    assert record.finding is True
    assert record.value["failed_count"] == 3
    assert record.value["repeated_task_types"] == [{"task_name": "RelocateVM_Task", "count": 3, "entities": ["vm-0", "vm-1", "vm-2"]}]
    assert record.value["time_clusters"][0]["count"] == 3

    isolated = collector._task_history_records(
        NS(taskManager=NS(recentTask=[NS(info=NS(state="error", name="RelocateVM_Task", entityName="vm-1", completeTime=base))])),
        "ds-task-history",
        "2026-09-23T01:05:00Z",
    )[0]
    assert isolated.finding is False


def test_alarm_history_ignores_normal_and_cleared_status_changes() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    event_type = type("AlarmStatusChangedEvent", (), {})
    red = event_type()
    red.fullFormattedMessage = "Alarm status changed to red and triggered"
    green = event_type()
    green.fullFormattedMessage = "Alarm status changed to green and cleared"
    records = collector._history_records([red, green], [], [], [], "ds-alarm", "2026-09-23T01:00:00Z")
    alarm = next(record for record in records if record.metadata["rule_id"] == "CL-DEEP-007")
    assert alarm.value == 1
    assert alarm.finding is True
    ha = next(record for record in records if record.metadata["rule_id"] == "CL-DEEP-002")
    assert ha.value == 0
    assert ha.finding is False
    evidence = evidence_ref(alarm)
    assert evidence.value_summary["event_samples"][0]["event_type"] == "AlarmStatusChangedEvent"


def test_history_records_cross_domain_event_correlation() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    base = datetime(2026, 9, 23, 1, 0, tzinfo=UTC)
    hosts = [NS(name="esx-01", _moId="host-1", config=NS(uuid="host-uuid-1")), NS(name="esx-02", _moId="host-2", config=NS(uuid="host-uuid-2"))]
    network = NS(fullFormattedMessage="vmnic link down", createdTime=base, entity=NS(name="esx-01", _moId="host-1"))
    storage = NS(fullFormattedMessage="APD path failure", createdTime=base + timedelta(minutes=20), entity=NS(name="esx-01", _moId="host-1"))
    late = NS(fullFormattedMessage="vmnic link down", createdTime=base + timedelta(hours=3), entity=NS(name="esx-01", _moId="host-1"))
    other_host = NS(fullFormattedMessage="vmnic link down", createdTime=base + timedelta(minutes=5), entity=NS(name="esx-02", _moId="host-2"))
    records = collector._history_records([network, storage, late, other_host], hosts, [], [], "ds-correlation", "2026-09-23T04:00:00Z")
    network_record = next(record for record in records if record.metadata["rule_id"] == "NET-DEEP-001")
    assert network_record.entity.type == "HostSystem"
    assert network_record.entity.stable_id == "host-uuid-1"
    assert network_record.metadata["event_samples"][0]["resolved_entities"][0]["stable_id"] == "host-uuid-1"
    correlation = next(record for record in records if record.metadata["rule_id"] == "HIST-DEEP-001")
    assert correlation.finding is True
    assert correlation.value["cluster_count"] == 1
    cluster = correlation.value["clusters"][0]
    assert cluster["entity_id"] == "host-uuid-1"
    assert {rule_id.split("-", 1)[0] for rule_id in cluster["rule_ids"]} == {"NET", "STO"}
    assert "NET-DEEP-001" not in cluster["rule_ids"]
    assert "NET-DEEP-002" in cluster["rule_ids"]
    narrow = collector._history_records([network, storage], hosts, [], [], "ds-correlation-narrow", "2026-09-23T04:00:00Z", correlation_window_seconds=600)
    narrow_correlation = next(record for record in narrow if record.metadata["rule_id"] == "HIST-DEEP-001")
    assert narrow_correlation.finding is False
    assert narrow_correlation.metadata["correlation_window_seconds"] == 600


def test_successful_vmotion_configuration_events_do_not_become_failures_or_cross_domain_clusters() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    base = datetime(2026, 9, 26, 1, 0, tzinfo=UTC)
    host = NS(name="esx-01", _moId="host-1", config=NS(uuid="host-uuid-1"))
    events = [
        NS(key=4001, createdTime=base, fullFormattedMessage="Firewall configuration has changed. Operation 'add' for rule set vMotion succeeded.", entity=NS(name="esx-01", _moId="host-1")),
        NS(key=4002, createdTime=base + timedelta(seconds=1), fullFormattedMessage="Firewall configuration has changed. Operation 'enable' for rule set vMotion succeeded.", entity=NS(name="esx-01", _moId="host-1")),
        NS(key=4003, createdTime=base + timedelta(seconds=2), fullFormattedMessage="APD path failure", entity=NS(name="esx-01", _moId="host-1")),
    ]
    records = collector._history_records(events, [host], [], [], "ds-vmotion-success", "2026-09-26T01:01:00Z")
    migration = next(item for item in records if item.metadata["rule_id"] == "COMPUTE-DEEP-001")
    correlation = next(item for item in records if item.metadata["rule_id"] == "HIST-DEEP-001")
    assert migration.value == 0
    assert migration.finding is False
    assert correlation.value["cluster_count"] == 0


def test_failed_vmotion_event_remains_historical_finding_evidence() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    event = NS(
        key=5001,
        createdTime=datetime(2026, 9, 26, 1, 0, tzinfo=UTC),
        fullFormattedMessage="vMotion migration failed because the target host was unavailable",
        entity=NS(name="esx-01", _moId="host-1"),
    )
    records = collector._history_records([event], [NS(name="esx-01", _moId="host-1", config=NS(uuid="host-uuid-1"))], [], [], "ds-vmotion-failed", "2026-09-26T01:01:00Z")
    migration = next(item for item in records if item.metadata["rule_id"] == "COMPUTE-DEEP-001")
    assert migration.value == 1
    assert migration.finding is True


def test_ha_history_keeps_failover_failures_but_skips_mentions_and_guest_restarts() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    base = datetime(2026, 9, 26, 1, 0, tzinfo=UTC)
    host = NS(name="esx-01", _moId="host-1", config=NS(uuid="host-uuid-1"))
    events = [
        NS(key=6001, createdTime=base, fullFormattedMessage="vSphere HA configuration was updated successfully", entity=NS(name="esx-01", _moId="host-1")),
        NS(key=6002, createdTime=base + timedelta(seconds=1), fullFormattedMessage="VM restarted after guest OS patch", entity=NS(name="esx-01", _moId="host-1")),
        NS(key=6003, createdTime=base + timedelta(seconds=2), fullFormattedMessage="vSphere HA virtual machine failover failed; alarm later changed from Gray to Green", entity=NS(name="esx-01", _moId="host-1")),
        NS(key=6004, createdTime=base + timedelta(seconds=3), fullFormattedMessage="HA restarted the protected VM after host failure", entity=NS(name="esx-01", _moId="host-1")),
    ]

    records = collector._history_records(events, [host], [], [], "ds-ha-history", "2026-09-26T01:01:00Z")
    ha = next(item for item in records if item.metadata["rule_id"] == "CL-DEEP-002")
    samples = ha.metadata["event_samples"]

    assert ha.value == 2
    assert {sample["event_key"] for sample in samples} == {6003, 6004}


def test_duplicate_link_down_events_are_not_reported_as_a_link_flap() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    base = datetime(2026, 9, 26, 1, 0, tzinfo=UTC)
    hosts = [NS(name="esx-01", _moId="host-1", config=NS(uuid="host-uuid-1"))]
    events = [
        NS(key=1001, createdTime=base, fullFormattedMessage="Physical NIC vmnic0 linkstate down", entity=NS(name="esx-01", _moId="host-1")),
        NS(key=1002, createdTime=base, fullFormattedMessage="Physical NIC vmnic0 linkstate down", entity=NS(name="esx-01", _moId="host-1")),
    ]
    records = collector._history_records(events, hosts, [], [], "ds-link-down-duplicate", "2026-09-26T01:01:00Z")
    flap_record = next(item for item in records if item.metadata["rule_id"] == "NET-DEEP-001")
    assert flap_record.value == 2
    assert flap_record.finding is False

    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-link-down-duplicate", "scope": {"hosts": 1, "vms": 0, "clusters": 0, "datastores": 0}})
    dataset.records = records
    rules = tuple(rule for rule in BUILTIN_DEEP_RULES if rule.rule_id in {"NET-DEEP-001", "NET-DEEP-002"})
    analysis = DeepAnalyzer().analyze(dataset, rules=rules)
    results = {item.rule_id: item for item in analysis.rule_results}
    assert results["NET-DEEP-001"].status == DeepResultStatus.PASS
    assert results["NET-DEEP-002"].status == DeepResultStatus.FINDING


def test_repeated_alternating_link_state_changes_meet_the_flap_threshold() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    base = datetime(2026, 9, 26, 1, 0, tzinfo=UTC)
    hosts = [NS(name="esx-01", _moId="host-1", config=NS(uuid="host-uuid-1"))]
    states = ("up", "down", "up", "down")
    events = [
        NS(
            key=2000 + index,
            createdTime=base + timedelta(seconds=index * 10),
            fullFormattedMessage=f"Physical NIC vmnic0 linkstate {state}",
            entity=NS(name="esx-01", _moId="host-1"),
        )
        for index, state in enumerate(states)
    ]
    records = collector._history_records(events, hosts, [], [], "ds-link-flap", "2026-09-26T01:01:00Z")
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-link-flap", "scope": {"hosts": 1, "vms": 0, "clusters": 0, "datastores": 0}})
    dataset.records = records
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "NET-DEEP-001")
    analysis = DeepAnalyzer().analyze(dataset, rules=(rule,))

    assert analysis.rule_results[0].status == DeepResultStatus.FINDING
    assert len(analysis.findings) == 1
    assert "vmnic0" in analysis.findings[0].fact
    assert "3 次状态转换" in analysis.findings[0].fact
    assert analysis.findings[0].threshold.id == "TH-NET-LINK-FLAP-STATE-CHANGES"


def test_unparseable_vmnic_link_state_evidence_does_not_pass() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    base = datetime(2026, 9, 26, 1, 0, tzinfo=UTC)
    host = NS(name="esx-01", _moId="host-1", config=NS(uuid="host-uuid-1"))
    event = NS(key=3001, createdTime=base, fullFormattedMessage="Physical NIC vmnic0 linkstate changed", entity=NS(name="esx-01", _moId="host-1"))
    records = collector._history_records([event], [host], [], [], "ds-link-state-unknown", "2026-09-26T01:01:00Z")
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-link-state-unknown", "scope": {"hosts": 1, "vms": 0, "clusters": 0, "datastores": 0}})
    dataset.records = records
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "NET-DEEP-001")

    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert result.status == DeepResultStatus.INSUFFICIENT_DATA


def test_cross_domain_history_without_object_identity_is_insufficient() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    event = NS(fullFormattedMessage="APD path failure", createdTime=datetime(2026, 9, 23, 1, 0, tzinfo=UTC))
    records = collector._history_records([event], [], [], [], "ds-unresolved-history", "2026-09-23T01:05:00Z")
    correlation = next(record for record in records if record.metadata["rule_id"] == "HIST-DEEP-001")
    assert correlation.finding is False
    assert correlation.window.completeness == 0.0
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records = [correlation]
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "HIST-DEEP-001")
    analysis = DeepAnalyzer().analyze(dataset, rules=(rule,))
    assert analysis.rule_results[0].status == DeepResultStatus.INSUFFICIENT_DATA


def test_deep_vcenter_correlation_window_is_configurable_from_cli() -> None:
    from vstacklens.cli import build_parser

    parser = build_parser()
    defaults = parser.parse_args(["run-deep-vcenter", "--connection-workbook", "readonly.xlsx"])
    configured = parser.parse_args(["run-deep-vcenter", "--connection-workbook", "readonly.xlsx", "--correlation-window-seconds", "600", "--event-history-days", "90", "--max-log-total-mib", "20", "--max-log-primary-mib", "15", "--max-log-fallback-mib", "5", "--log-days", "7"])
    assert defaults.correlation_window_seconds == 3600
    assert defaults.event_history_days == 30
    assert configured.correlation_window_seconds == 600
    assert configured.event_history_days == 90
    assert (defaults.max_log_total_mib, defaults.max_log_primary_mib, defaults.max_log_fallback_mib) == (16, 12, 4)
    assert (configured.max_log_total_mib, configured.max_log_primary_mib, configured.max_log_fallback_mib) == (20, 15, 5)
    assert defaults.log_days is None
    assert configured.log_days == 7


def test_deep_replay_of_local_dataset_never_opens_a_vmware_connection(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    from pyVim import connect

    def deny_connection(*_args, **_kwargs):
        raise AssertionError("offline Deep replay must not connect to vCenter or ESXi")

    monkeypatch.setattr(connect, "SmartConnect", deny_connection)
    results = DeepInspectionService().replay_datasets([FIXTURE], tmp_path / "offline-replay")

    assert len(results) == 1
    assert results[0].html_path.is_file()


def test_event_history_days_are_bounded_without_changing_performance_window() -> None:
    policy = DeepCollectionPolicy(event_history_days=90)
    assert policy.event_history_days == 90
    assert policy.history_days == 30
    with pytest.raises(ValidationError):
        DeepCollectionPolicy(event_history_days=91)


def test_event_history_query_uses_separate_configured_window() -> None:
    captured: dict[str, Any] = {}

    class EventCollector:
        def ReadNext(self, _count):
            return []

        def DestroyCollector(self):
            captured["destroyed"] = True

    def create_collector(filter_spec):
        captured["filter"] = filter_spec
        return EventCollector()

    manager = NS(CreateCollectorForEvents=create_collector)
    collector = PyVmomiDeepCollector("vc", "u", "p")
    events = collector._bounded_events(
        NS(eventManager=manager),
        DeepCollectionPolicy(history_days=30, event_history_days=90),
    )

    assert events == []
    assert captured["destroyed"] is True
    filter_window = captured["filter"].time
    assert (filter_window.endTime - filter_window.beginTime).days == 90
    assert collector._event_query_meta["window_start"] == filter_window.beginTime.replace(microsecond=0).isoformat().replace("+00:00", "Z")
    assert collector._event_query_meta["window_end"] == filter_window.endTime.replace(microsecond=0).isoformat().replace("+00:00", "Z")


def test_hardware_sensor_requires_complete_sensor_coverage() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    host = NS(
        name="esx-01",
        config=NS(instanceUuid="host-hardware-1"),
        runtime=NS(
            healthSystemRuntime=NS(
                systemHealthInfo=NS(
                    numericSensorInfo=[
                        NS(name="Power Supply 1", healthState=NS(key="red", label="Red", summary="Power supply failed")),
                        NS(name="Temperature", healthState=NS(key="green")),
                    ]
                )
            )
        ),
    )
    records = collector._hardware_sensor_records([host], "ds-hardware", "2026-09-23T01:00:00Z")
    capability = collector._hardware_sensor_capability(records)
    assert capability.status == CapabilityStatus.AVAILABLE
    assert records[0].value["issues"] == [{"sensor": "Power Supply 1", "status": "red", "summary": "Power supply failed"}]

    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-hardware"})
    dataset.records = records
    dataset.capability = CapabilityMatrix(probed_at_utc="2026-09-23T01:00:00Z", capabilities=[capability])
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "HARD-DEEP-001")
    result = DeepAnalyzer().analyze(dataset, rules=(rule,))
    assert result.rule_results[0].status == DeepResultStatus.FINDING

    unavailable = collector._hardware_sensor_records([NS(name="esx-02", config=NS(instanceUuid="host-hardware-2"), runtime=NS())], "ds-hardware", "2026-09-23T01:00:00Z")
    assert collector._hardware_sensor_capability(unavailable).status == CapabilityStatus.UNAVAILABLE


def test_bundled_vcg_index_is_integrity_checked_fresh_and_maps_supported_esxi_updates() -> None:
    index = load_bundled_vcg_index()
    try:
        assert len(index.checksum_sha256) == 64
        assert index.freshness.status == "FRESH"
        assert esxi_version_to_hcl_release("7.0.3", index.supported_releases) == "ESXi 7.0 U3"
        assert esxi_version_to_hcl_release("8.0.3", index.supported_releases) == "ESXi 8.0 U3"
        assert esxi_version_to_hcl_release("9.0.1", index.supported_releases) is None
        assert esxi_version_to_hcl_release("8.0.3", ("ESXi 8.0 U2",)) is None
    finally:
        index.close()


def test_hardware_hcl_rule_preserves_confirmed_mismatch_when_other_device_is_unknown(monkeypatch: pytest.MonkeyPatch) -> None:
    import vstacklens.deep.pyvmomi_collector as collector_module

    class FakeCatalog:
        baseline_version = "fixture-hcl"
        json_updated_time = "2026-09-26T00:00:00Z"
        freshness = NS(status="FRESH", age_days=1.0)
        supported_releases = ("ESXi 8.0 U3",)

        def __init__(self) -> None:
            self.closed = False
            self.store = object()
            self.matcher = self

        def match(self, *_args: object, **kwargs: object) -> NS:
            status = "FIRMWARE_MISMATCH" if kwargs.get("category") == "controller" else "UNKNOWN_DEVICE"
            return NS(status=status, detail=f"fixture status {status}", vcglink="https://vcg.example/device", candidate_links=[], certified_driver_versions=["1.2.3"], certified_firmware_versions=["2.3.4"])

        def close(self) -> None:
            self.closed = True

    fake_catalog = FakeCatalog()
    monkeypatch.setattr(collector_module, "load_bundled_vcg_index", lambda: fake_catalog)
    monkeypatch.setattr(
        "vstacklens.upgrade_compat.server.match_server_model",
        lambda _store, _model, _release: NS(status="SERVER_CERTIFIED", detail="fixture server certified", vcglink=None, candidate_links=()),
    )
    collector = PyVmomiDeepCollector("vc.example", "readonly", "not-a-real-password")
    host = NS(
        name="esx-fixture",
        _moId="host-fixture",
        config=NS(uuid="host-fixture-id", product=NS(version="8.0.3", build="fixture-build")),
        hardware=NS(systemInfo=NS(model="Fixture Server")),
    )
    host_entity = collector._entity(host, "HostSystem")
    devices = [
        {"object_key": "pci:0000:01:00.0", "category": "controller", "model": "Fixture HBA", "vid": "1000", "did": "005d", "svid": "1028", "ssid": "1f49", "driver_name": "lsi_mr3", "driver_version": "7.718.02.00", "firmware_version": "2.3.3", "firmware_confidence": "certain"},
        {"object_key": "pci:0000:02:00.0", "category": "nic", "model": "Fixture NIC", "vid": "8086", "did": "158b", "svid": "1028", "ssid": "0001", "driver_name": "i40en", "driver_version": "1.16.1", "firmware_version": None, "firmware_confidence": None},
    ]
    component_record = DatasetRecord(
        record_id="hardware-component-fixture",
        dataset_id="ds-hardware-compat-fixture",
        kind="hardware",
        entity=host_entity,
        collected_at_utc="2026-09-26T01:00:00Z",
        source=DeepSource(api="fixture.pci", collector="fixture", collected_at_utc="2026-09-26T01:00:00Z"),
        selector={"rule_id": "HARD-DEEP-002"},
        window=DeepWindow(start="2026-09-26T01:00:00Z", end="2026-09-26T01:00:00Z", sample_count=1, expected_sample_count=1, completeness=1.0),
        value={"device_count": 2, "devices": devices, "warnings": []},
        unit="inventory",
        raw_pointer="hardware/fixture.ndjson",
        metadata={"rule_id": "HARD-DEEP-002"},
    )

    records = collector._hardware_compatibility_records([host], [component_record], "ds-hardware-compat-fixture", "2026-09-26T01:00:00Z")
    summary = next(item for item in records if item.metadata.get("hardware_compatibility_summary"))
    phase_log = {"action": "hardware.compatibility", "status": "ok"}
    collector._update_hardware_compatibility_log([phase_log], records)
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-hardware-compat-fixture", "scope": {"hosts": 1, "vms": 0, "clusters": 0, "datastores": 0}})
    dataset.records = records
    dataset.capability.capabilities.append(
        Capability(id="hardware.vcg.compatibility", name="VCG", status=CapabilityStatus.LIMITED, profile="enhanced")
    )
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "HARD-DEEP-003")
    analysis = DeepAnalyzer().analyze(dataset, rules=(rule,))

    assert fake_catalog.closed is True
    assert summary.value["target_release_by_host"][host_entity.stable_id] == "ESXi 8.0 U3"
    assert summary.value["incompatible_count"] == 1
    assert phase_log["incompatible_count"] == 1
    assert summary.value["coverage_complete"] is False
    assert analysis.rule_results[0].status == DeepResultStatus.FINDING
    assert "FIRMWARE_MISMATCH" in analysis.findings[0].fact
    assert "未取得状态" in analysis.rule_results[0].reason


def test_missing_bundled_vcg_data_does_not_report_hardware_pass(monkeypatch: pytest.MonkeyPatch) -> None:
    import vstacklens.deep.pyvmomi_collector as collector_module

    monkeypatch.setattr(collector_module, "load_bundled_vcg_index", lambda: (_ for _ in ()).throw(FileNotFoundError("fixture missing")))
    collector = PyVmomiDeepCollector("vc.example", "readonly", "not-a-real-password")
    host = NS(name="esx-no-catalog", config=NS(uuid="host-no-catalog", product=NS(version="8.0.3", build="test")), hardware=NS(systemInfo=NS(model="Fixture Server")))
    records = collector._hardware_compatibility_records([host], [], "ds-no-vcg", "2026-09-26T01:00:00Z")
    capability = collector._hardware_compatibility_capability(records, 1, {"status": "unavailable"})
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-no-vcg", "scope": {"hosts": 1, "vms": 0, "clusters": 0, "datastores": 0}})
    dataset.records = records
    dataset.capability.capabilities.append(capability)
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "HARD-DEEP-003")
    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert capability.status == CapabilityStatus.UNAVAILABLE
    assert result.status == DeepResultStatus.NOT_EVALUATED
    assert "hardware.vcg.compatibility" in result.missing_capabilities


def test_vm_resource_review_record_captures_reservation_limit_and_shares() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    normal_vm = NS(
        name="vm-normal",
        config=NS(
            instanceUuid="vm-normal-id",
            version="vmx-21",
            cpuAllocation=NS(limit=-1, reservation=0, shares=NS(level="normal", shares=1000)),
            memoryAllocation=NS(limit=-1, reservation=0, shares=NS(level="normal", shares=81920)),
        ),
        runtime=NS(connectionState="connected", powerState="poweredOn"),
        guest=NS(toolsRunningStatus="guestToolsNotRunning", toolsVersionStatus2="guestToolsCurrent"),
        snapshot=None,
    )
    custom_vm = NS(
        name="vm-custom",
        config=NS(
            instanceUuid="vm-custom-id",
            version="vmx-21",
            cpuAllocation=NS(limit=-1, reservation=2000, shares=NS(level="custom", shares=3000)),
            memoryAllocation=NS(limit=4096, reservation=2048, shares=NS(level="normal", shares=81920)),
        ),
        runtime=NS(connectionState="connected", powerState="poweredOff"),
        guest=NS(toolsRunningStatus=""),
        snapshot=None,
    )
    records = collector._vm_hidden_risk_records([normal_vm, custom_vm], "ds-vm-resource", "2026-09-23T01:00:00Z")
    allocations = {record.entity.display_ref: record for record in records if record.metadata["rule_id"] == "VM-DEEP-005"}
    assert allocations["vm-normal"].finding is False
    assert allocations["vm-normal"].value["cpu_reservation_mhz"] == 0
    assert allocations["vm-normal"].value["cpu_shares"]["level"] == "normal"
    assert allocations["vm-custom"].finding is True
    assert allocations["vm-custom"].value["memory_limit_mb"] == 4096
    assert allocations["vm-custom"].value["memory_reservation_mb"] == 2048
    tools_record = next(record for record in records if record.entity.display_ref == "vm-normal" and record.metadata["rule_id"] == "VM-DEEP-003")
    assert tools_record.finding is True


def test_guest_os_mismatch_is_reported_only_with_running_tools_evidence() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    vm = NS(
        name="vm-os-mismatch",
        config=NS(instanceUuid="vm-os-mismatch-id", version="vmx-21", guestId="windows9Server64Guest", guestFullName="Microsoft Windows Server 2022 (64-bit)"),
        runtime=NS(connectionState="connected", powerState="poweredOn"),
        guest=NS(toolsRunningStatus="guestToolsRunning", guestFullName="Ubuntu Linux (64-bit)"),
        snapshot=None,
    )
    records = collector._vm_hidden_risk_records([vm], "ds-guest-os", "2026-09-23T01:00:00Z")
    mismatch = next(record for record in records if record.metadata["rule_id"] == "VM-DEEP-007")
    assert mismatch.finding is True
    assert mismatch.value["configured_guest_os"] == "Microsoft Windows Server 2022 (64-bit)"
    assert mismatch.value["reported_guest_os"] == "Ubuntu Linux (64-bit)"


@pytest.mark.parametrize(
    ("support_level", "expected_status"),
    [("supported", DeepResultStatus.PASS), ("deprecated", DeepResultStatus.FINDING), ("terminated", DeepResultStatus.FINDING)],
)
def test_guest_os_support_level_is_checked_against_host_environment(
    support_level: str,
    expected_status: DeepResultStatus,
) -> None:
    class Browser:
        def QueryConfigOption(self, *, key: str, host: object) -> object:
            assert key == "vmx-21"
            assert host is host_object
            return NS(guestOSDescriptor=[NS(id="ubuntu64Guest", supportLevel=support_level, fullName="Ubuntu Linux (64-bit)")])

    host_object = NS(name="esx-support", _moId="host-support", runtime=NS(connectionState="connected"))
    compute = NS(name="cluster-support", _moId="domain-support", host=[host_object])
    compute.environmentBrowser = Browser()
    host_object.parent = compute
    vm = NS(
        name="vm-support",
        config=NS(instanceUuid="vm-support-id", guestId="ubuntu64Guest", version="vmx-21"),
        runtime=NS(host=host_object),
    )
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    records = collector._guest_os_support_records([], [vm], "ds-guest-os-support", "2026-09-28T01:00:00Z")
    capability = collector._guest_os_support_capability(records, 1, {"status": "ok"})
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-guest-os-support"})
    dataset.records = [record for record in dataset.records if record.metadata.get("rule_id") != "VM-DEEP-008"] + records
    dataset.capability.capabilities = [item for item in dataset.capability.capabilities if item.id != "vm.guest_os.support"] + [capability]
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "VM-DEEP-008")

    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert result.status == expected_status
    assert records[0].value["support_status"] == ("supported" if support_level == "supported" else "needs_review")
    assert records[0].source.api == "vim.EnvironmentBrowser.QueryConfigOption"


def test_guest_os_descriptor_absence_is_insufficient_not_unsupported() -> None:
    class Browser:
        def QueryConfigOption(self, *, key: str, host: object) -> object:
            return NS(guestOSDescriptor=[NS(id="otherGuest", supportLevel="supported", fullName="Other OS")])

    host_object = NS(name="esx-support", _moId="host-support", runtime=NS(connectionState="connected"))
    compute = NS(name="cluster-support", _moId="domain-support", host=[host_object])
    compute.environmentBrowser = Browser()
    host_object.parent = compute
    vm = NS(name="vm-unlisted", config=NS(guestId="unlistedGuest", version="vmx-21"), runtime=NS(host=host_object))
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    records = collector._guest_os_support_records([], [vm], "ds-guest-os-unlisted", "2026-09-28T01:00:00Z")
    capability = collector._guest_os_support_capability(records, 1, {"status": "ok"})
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-guest-os-unlisted"})
    dataset.records = [record for record in dataset.records if record.metadata.get("rule_id") != "VM-DEEP-008"] + records
    dataset.capability.capabilities = [item for item in dataset.capability.capabilities if item.id != "vm.guest_os.support"] + [capability]
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "VM-DEEP-008")

    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert records[0].value["support_status"] == "not_listed"
    assert records[0].finding is False
    assert result.status == DeepResultStatus.INSUFFICIENT_DATA


def test_network_hidden_risks_compare_shared_network_keys_and_per_switch_uplinks() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")

    def host(name: str, moid: str, *, extra_network: bool = False, vlan: int = 10, proxy_mtu: int = 9000):
        vswitches = [
            NS(name="vSwitch0", key="key-vswitch0", mtu=1500, pnic=[f"vmnic-{moid}-0"], spec=NS(mtu=1500, policy=NS(nicTeaming=NS(policy="loadbalance_srcid", notifySwitches=True)))),
        ]
        portgroups = [NS(key=f"pg-management-{moid}", spec=NS(name="Management Network", vswitchName="vSwitch0", vlanId=vlan, policy=None))]
        vnics = [NS(device="vmk0", portgroup="Management Network", spec=NS(mtu=1500))]
        if extra_network:
            vswitches.append(NS(name="vSwitchStorage", key="key-storage", mtu=9000, pnic=[f"vmnic-{moid}-1"], spec=NS(mtu=9000, policy=NS(nicTeaming=NS(policy="loadbalance_srcid", notifySwitches=True)))))
            portgroups.append(NS(key=f"pg-storage-{moid}", spec=NS(name="Storage", vswitchName="vSwitchStorage", vlanId=20, policy=None)))
            vnics.append(NS(device="vmk1", portgroup="Storage", spec=NS(mtu=9000)))
        proxy = NS(dvsUuid="dvs-shared", dvsName="Distributed Switch", key=f"proxy-{moid}", mtu=proxy_mtu, pnic=[f"vmnic-{moid}-2"])
        return NS(name=name, _moId=moid, config=NS(instanceUuid=f"uuid-{moid}", network=NS(vswitch=vswitches, proxySwitch=[proxy], portgroup=portgroups, vnic=vnics)))

    first = host("esx-01", "host-1")
    second = host("esx-02", "host-2", extra_network=True)
    cluster = NS(name="cluster-a", _moId="domain-c1", host=[first, second])
    records = collector._network_hidden_risk_records([first, second], "ds-network", "2026-09-25T00:00:00Z", clusters=[cluster])
    uplink_records = [item for item in records if item.metadata["rule_id"] == "NET-DEEP-003"]
    drift = next(item for item in records if item.metadata["rule_id"] == "NET-DEEP-004")

    assert all(item.finding for item in uplink_records)
    assert all(item.value["host_unique_uplink_count"] >= 2 for item in uplink_records)
    assert all(item.value["single_or_zero_uplink_switches"] for item in uplink_records)
    assert drift.finding is True
    assert drift.value["comparison"]["comparable_network_property_count"] >= 4
    assert drift.value["comparison"]["same_cluster_count"] == 1
    assert any(item["property"] == "membership" and item["network_key"] == "vSwitchStorage/Storage" for item in drift.value["differences"])


def test_network_hidden_risk_reports_shared_vlan_and_vds_mtu_differences() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")

    def host(name: str, moid: str, vlan: int, proxy_mtu: int):
        network = NS(
            vswitch=[NS(name="vSwitch0", key="vSwitch0", mtu=1500, pnic=[f"vmnic-{moid}-0"], spec=NS(mtu=1500, policy=None))],
            proxySwitch=[NS(dvsUuid="dvs-shared", dvsName="Distributed Switch", key=f"proxy-{moid}", mtu=proxy_mtu, pnic=[f"vmnic-{moid}-1"])],
            portgroup=[NS(key=f"pg-{moid}", spec=NS(name="Management Network", vswitchName="vSwitch0", vlanId=vlan, policy=None))],
            vnic=[NS(device="vmk0", portgroup="Management Network", spec=NS(mtu=1500))],
        )
        return NS(name=name, _moId=moid, config=NS(instanceUuid=f"uuid-{moid}", network=network))

    first = host("esx-01", "host-1", 10, 9000)
    second = host("esx-02", "host-2", 20, 1500)
    cluster = NS(name="cluster-a", _moId="domain-c1", host=[first, second])
    records = collector._network_hidden_risk_records([first, second], "ds-network-drift", "2026-09-25T00:00:00Z", clusters=[cluster])
    drift = next(item for item in records if item.metadata["rule_id"] == "NET-DEEP-004")
    differing = {(item["network_type"], item["property"]) for item in drift.value["differences"]}

    assert drift.finding is True
    assert ("portgroups", "vlan_id") in differing
    assert ("proxy_switches", "mtu") in differing


def test_network_configuration_differences_between_clusters_are_not_reported_as_drift() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")

    def host(name: str, moid: str, mtu: int):
        network = NS(
            vswitch=[NS(name="vSwitch0", key="vSwitch0", mtu=mtu, pnic=[f"vmnic-{moid}-0"], spec=NS(mtu=mtu, policy=None))],
            proxySwitch=[],
            portgroup=[NS(key=f"pg-{moid}", spec=NS(name="Management Network", vswitchName="vSwitch0", vlanId=10, policy=None))],
            vnic=[NS(device="vmk0", portgroup="Management Network", spec=NS(mtu=1500))],
        )
        return NS(name=name, _moId=moid, config=NS(instanceUuid=f"uuid-{moid}", network=network))

    hosts = [host("esx-01", "host-1", 1500), host("esx-02", "host-2", 1500), host("esx-03", "host-3", 9000), host("esx-04", "host-4", 9000)]
    clusters = [
        NS(name="cluster-a", _moId="domain-c1", host=hosts[:2]),
        NS(name="cluster-b", _moId="domain-c2", host=hosts[2:]),
    ]
    records = collector._network_hidden_risk_records(hosts, "ds-network-clusters", "2026-09-25T00:00:00Z", clusters=clusters)
    drift = next(item for item in records if item.metadata["rule_id"] == "NET-DEEP-004")

    assert drift.finding is False
    assert drift.value["comparison"]["same_cluster_count"] == 2
    assert drift.value["comparison"]["difference_count"] == 0


def test_distributed_network_profile_capture_is_read_only_bounded_and_destroyed() -> None:
    import json
    from pyVmomi import vim

    collector = PyVmomiDeepCollector("vc", "u", "p")
    switch_type = type("VmwareDistributedVirtualSwitch", (), {})
    portgroup_type = type("DistributedVirtualPortgroup", (), {})
    vlan_type = type("VlanIdSpec", (), {})
    vlan = vlan_type()
    vlan.vlanId = 42
    team = NS(
        policy=vim.StringPolicy(value="loadbalance_loadbased"),
        notifySwitches=vim.BoolPolicy(value=True),
        rollingOrder=vim.BoolPolicy(value=False),
        uplinkPortOrder=NS(activeUplinkPort=["uplink1", "uplink2"], standbyUplinkPort=[]),
    )
    dvs = switch_type()
    dvs._moId = "dvs-1"
    dvs.uuid = "dvs-uuid-1"
    dvs.config = NS(name="dvSwitch-1", uuid="dvs-uuid-1", defaultPortConfig=NS(vlan=None, uplinkTeamingPolicy=team))
    dvpg = portgroup_type()
    dvpg._moId = "dvportgroup-1"
    dvpg.config = NS(name="Management", distributedVirtualSwitch=NS(_moId="dvs-1"), defaultPortConfig=NS(vlan=vlan, uplinkTeamingPolicy=team))

    class View:
        view = [dvs, dvpg]
        destroyed = False

        def Destroy(self):
            self.destroyed = True

    view = View()

    class ViewManager:
        calls = 0

        def CreateContainerView(self, root, object_types, recursive):
            self.calls += 1
            assert recursive is True
            return view

    view_manager = ViewManager()
    content = NS(rootFolder=NS(), viewManager=view_manager)
    profiles = collector._collect_distributed_network_profiles(content)

    assert view.destroyed is True
    assert view_manager.calls == 1
    assert profiles["status"] == "available"
    assert profiles["switch_count"] == 1
    assert profiles["portgroup_count"] == 1
    assert profiles["portgroups"][0]["vlan"]["vlanId"] == 42
    assert profiles["portgroups"][0]["teaming"]["active_uplink_count"] == 2
    assert profiles["portgroups"][0]["teaming"]["policy"] == "loadbalance_loadbased"
    json.dumps(profiles)

    guarded = collector._collect_distributed_network_profiles(content, resource_guard=lambda: "memory_limit_exceeded")
    assert guarded["status"] == "not_requested"
    assert view_manager.calls == 1


def test_enhanced_rule_is_gated_without_fake_pass() -> None:
    dataset = DeepDataset.from_json(FIXTURE)
    rule = DeepRuleDefinition(
        rule_id="ENHANCED-IO-001",
        title="Enhanced I/O latency",
        category="trend",
        profile="enhanced",
        dimension_key="io_latency",
        record_kind="perf",
        requires=["perf.level2.required"],
        verification_guidance="在 vSphere Client 中复核 I/O 延迟历史。",
    )
    result = DeepAnalyzer().analyze(dataset, rules=(rule,))
    assert result.rule_results[0].status == DeepResultStatus.NOT_EVALUATED
    assert result.findings == []


def test_enhanced_performance_analyzer_requires_samples_then_finds() -> None:
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.capability.capabilities.append(Capability(id="enhanced.perf.cpu", name="CPU Enhanced", status=CapabilityStatus.AVAILABLE, profile="enhanced"))
    for index, value in enumerate((6.0, 7.0, 6.5, 7.5, 6.2, 6.8), 1):
        dataset.records.append(
            DatasetRecord(
                record_id=f"cpu-enhanced-{index}",
                dataset_id=dataset.dataset_id,
                kind="perf",
                entity=DeepEntity(type="HostSystem", stable_id="host-enhanced", display_ref="host-enhanced"),
                collected_at_utc=f"2026-09-2{index}T03:12:00Z",
                source=DeepSource(api="fixture.QueryPerf", collector="fixture", collected_at_utc="2026-09-21T03:12:00Z"),
                selector={"family": "enhanced.perf.cpu", "counter": "cpu.ready"},
                window=DeepWindow(start=f"2026-09-2{index}T03:00:00Z", end=f"2026-09-2{index}T03:12:00Z", interval_sec=300, sample_count=2, expected_sample_count=2, completeness=1.0),
                interval_sec=300,
                value={"first": value - 1, "last": value, "max": value},
                unit="percent",
                raw_pointer=f"perf/enhanced/cpu#{index}",
                metadata={"enhanced": True, "family": "enhanced.perf.cpu"},
            )
        )
    result = DeepAnalyzer().analyze(dataset)
    enhanced = next(item for item in result.rule_results if item.rule_id == "ENHANCED-CPU-001")
    assert enhanced.status == DeepResultStatus.FINDING
    finding = next(item for item in result.findings if item.rule_id == "ENHANCED-CPU-001")
    assert "7.50%" in finding.fact


def test_enhanced_cpu_ready_uses_samples_and_requires_sustained_breach() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    start = datetime(2026, 9, 23, tzinfo=UTC)
    points = collector._enhanced_perf_points(
        "cpu.ready",
        [15000, 18000],
        [NS(timestamp=start), NS(timestamp=start + timedelta(minutes=5))],
        interval_sec=300,
    )
    assert [value for _timestamp, value in points] == [5.0, 6.0]
    usage_points = collector._enhanced_perf_points(
        "cpu.usage",
        [1362],
        [NS(timestamp=start)],
        interval_sec=300,
    )
    assert [value for _timestamp, value in usage_points] == [13.62]
    costop_points = collector._enhanced_perf_points(
        "cpu.costop",
        [9000],
        [NS(timestamp=start)],
        interval_sec=300,
    )
    assert [value for _timestamp, value in costop_points] == [3.0]

    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records = [
        DatasetRecord(
            record_id="cpu-not-sustained",
            dataset_id=dataset.dataset_id,
            kind="perf",
            entity=DeepEntity(type="HostSystem", stable_id="host-enhanced", display_ref="host-enhanced"),
            collected_at_utc="2026-09-23T01:00:00Z",
            source=DeepSource(api="fixture.QueryPerf", collector="fixture"),
            selector={"family": "enhanced.perf.cpu", "counter": "cpu.ready"},
            window=DeepWindow(start="2026-09-23T00:30:00Z", end="2026-09-23T01:00:00Z", interval_sec=300, sample_count=6, expected_sample_count=6, completeness=1.0),
            value={"max": 6.0},
            unit="percent",
            raw_pointer="perf/enhanced/cpu#1",
            metadata={"enhanced": True, "family": "enhanced.perf.cpu", "sample_values": [6.0, 6.0, 6.0, 6.0, 6.0, 4.0]},
        )
    ]
    dataset.capability.capabilities.append(Capability(id="enhanced.perf.cpu", name="CPU Enhanced", status=CapabilityStatus.AVAILABLE, profile="enhanced"))
    dataset.records.append(
        DatasetRecord(
            record_id="cpu-usage-not-ready",
            dataset_id=dataset.dataset_id,
            kind="perf",
            entity=DeepEntity(type="HostSystem", stable_id="host-enhanced", display_ref="host-enhanced"),
            collected_at_utc="2026-09-23T01:00:00Z",
            source=DeepSource(api="fixture.QueryPerf", collector="fixture"),
            selector={"family": "enhanced.perf.cpu", "counter": "cpu.usage"},
            window=DeepWindow(start="2026-09-23T00:30:00Z", end="2026-09-23T01:00:00Z", interval_sec=300, sample_count=6, expected_sample_count=6, completeness=1.0),
            value={"max": 99.0},
            unit="percent",
            raw_pointer="perf/enhanced/cpu#2",
            metadata={"enhanced": True, "family": "enhanced.perf.cpu", "sample_values": [99.0] * 6},
        )
    )
    dataset.records.append(
        DatasetRecord(
            record_id="cpu-costop-high",
            dataset_id=dataset.dataset_id,
            kind="perf",
            entity=DeepEntity(type="HostSystem", stable_id="host-enhanced", display_ref="host-enhanced"),
            collected_at_utc="2026-09-23T01:00:00Z",
            source=DeepSource(api="fixture.QueryPerf", collector="fixture"),
            selector={"family": "enhanced.perf.cpu", "counter": "cpu.costop"},
            window=DeepWindow(start="2026-09-23T00:30:00Z", end="2026-09-23T01:00:00Z", interval_sec=300, sample_count=6, expected_sample_count=6, completeness=1.0),
            value={"max": 4.0},
            unit="percent",
            raw_pointer="perf/enhanced/cpu#3",
            metadata={"enhanced": True, "family": "enhanced.perf.cpu", "sample_values": [4.0] * 6},
        )
    )
    dataset.records.extend(
        [
            DatasetRecord(
                record_id="storage-latency-low",
                dataset_id=dataset.dataset_id,
                kind="perf",
                entity=DeepEntity(type="HostSystem", stable_id="host-storage", display_ref="host-storage"),
                collected_at_utc="2026-09-23T01:00:00Z",
                source=DeepSource(api="fixture.QueryPerf", collector="fixture"),
                selector={"family": "enhanced.perf.storage", "counter": "disk.maxtotallatency"},
                window=DeepWindow(start="2026-09-23T00:30:00Z", end="2026-09-23T01:00:00Z", interval_sec=300, sample_count=6, expected_sample_count=6, completeness=1.0),
                value={"max": 1.0},
                unit="millisecond",
                raw_pointer="perf/enhanced/storage#1",
                metadata={"enhanced": True, "family": "enhanced.perf.storage", "sample_values": [1.0] * 6},
            ),
            DatasetRecord(
                record_id="storage-queue-high",
                dataset_id=dataset.dataset_id,
                kind="perf",
                entity=DeepEntity(type="HostSystem", stable_id="host-storage", display_ref="host-storage"),
                collected_at_utc="2026-09-23T01:00:00Z",
                source=DeepSource(api="fixture.QueryPerf", collector="fixture"),
                selector={"family": "enhanced.perf.storage", "counter": "disk.queueLatency"},
                window=DeepWindow(start="2026-09-23T00:30:00Z", end="2026-09-23T01:00:00Z", interval_sec=300, sample_count=6, expected_sample_count=6, completeness=1.0),
                value={"max": 99.0},
                unit="millisecond",
                raw_pointer="perf/enhanced/storage#2",
                metadata={"enhanced": True, "family": "enhanced.perf.storage", "sample_values": [99.0] * 6},
            ),
        ]
    )
    dataset.capability.capabilities.append(Capability(id="enhanced.perf.storage", name="Storage Enhanced", status=CapabilityStatus.AVAILABLE, profile="enhanced"))
    result = DeepAnalyzer().analyze(dataset)
    enhanced = next(item for item in result.rule_results if item.rule_id == "ENHANCED-CPU-001")
    assert enhanced.status == DeepResultStatus.PASS
    cpu_usage = next(item for item in result.rule_results if item.rule_id == "ENHANCED-CPU-002")
    cpu_costop = next(item for item in result.rule_results if item.rule_id == "ENHANCED-CPU-003")
    assert cpu_usage.status == DeepResultStatus.FINDING
    assert cpu_costop.status == DeepResultStatus.FINDING
    storage = next(item for item in result.rule_results if item.rule_id == "ENHANCED-STO-001")
    assert storage.status == DeepResultStatus.PASS


def test_enhanced_memory_balloon_uses_normalized_megabytes() -> None:
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records = [
        DatasetRecord(
            record_id="memory-balloon",
            dataset_id=dataset.dataset_id,
            kind="perf",
            entity=DeepEntity(type="HostSystem", stable_id="host-memory", display_ref="host-memory"),
            collected_at_utc="2026-09-23T01:00:00Z",
            source=DeepSource(api="fixture.QueryPerf", collector="fixture"),
            selector={"family": "enhanced.perf.memory", "counter": "mem.vmmemctl"},
            window=DeepWindow(start="2026-09-23T00:30:00Z", end="2026-09-23T01:00:00Z", interval_sec=300, sample_count=6, expected_sample_count=6, completeness=1.0),
            value={"max": 1024.0},
            unit="megabyte",
            raw_pointer="perf/enhanced/memory#1",
            metadata={"enhanced": True, "family": "enhanced.perf.memory", "sample_values": [1024.0] * 6},
        )
    ]
    dataset.capability.capabilities.append(Capability(id="enhanced.perf.memory", name="Memory Enhanced", status=CapabilityStatus.AVAILABLE, profile="enhanced"))
    result = DeepAnalyzer().analyze(dataset)
    enhanced = next(item for item in result.rule_results if item.rule_id == "ENHANCED-MEM-001")
    assert enhanced.status == DeepResultStatus.FINDING
    finding = next(item for item in result.findings if item.rule_id == "ENHANCED-MEM-001")
    assert "1024.00 MB" in finding.fact


def test_vsan_management_records_use_api_evidence_and_only_flag_suspended_resync() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    cluster = NS(name="vSAN-Cluster", config=NS(instanceUuid="cluster-vsan-1"))
    datastore = NS(
        name="vsanDatastore",
        summary=NS(type="vsan", capacity=1000, freeSpace=100, url="ds:///vsan-1"),
    )
    inventory = {
        "status": "collected",
        "clusters": {
            "vSAN-Cluster": {
                "source": {"managed_object": "VsanVcClusterHealthSystem/VsanObjectSystem"},
                "health_issues": [],
                "object_health_issues": [{"component": "Object compliance", "status": "yellow"}],
                "resync_object_count": 3,
                "resync_bytes": 4096,
                "resync_active_object_count": 2,
                "resync_queued_object_count": 0,
                "resync_suspended_object_count": 1,
                "disk_health_issues": [{"host": "esx-01", "disk": "naa.1", "status": "red"}],
                "disk_group_count": 1,
                "cache_disk_count": 1,
                "capacity_disk_count": 3,
            }
        },
    }
    capabilities = collector._probe_capabilities(
        NS(taskManager=None, alarmManager=None, about=None, licenseManager=None, perfManager=None),
        [],
        [],
        [],
        vsan_inventory=inventory,
    )
    records = collector._vsan_records(
        [cluster],
        [datastore],
        "ds-vsan-fixture",
        "2026-09-23T00:00:00Z",
        inventory,
    )

    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-vsan-fixture"})
    dataset.capability = capabilities
    dataset.records = records
    rules = tuple(rule for rule in BUILTIN_DEEP_RULES if rule.rule_id.startswith("VSAN-"))
    analysis = DeepAnalyzer().analyze(dataset, rules=rules)
    statuses = {result.rule_id: result.status for result in analysis.rule_results}

    assert {rule_id: statuses[rule_id] for rule_id in statuses if rule_id in {"VSAN-DEEP-001", "VSAN-DEEP-002", "VSAN-DEEP-003", "VSAN-DEEP-004"}} == {
        "VSAN-DEEP-001": DeepResultStatus.FINDING,
        "VSAN-DEEP-002": DeepResultStatus.FINDING,
        "VSAN-DEEP-003": DeepResultStatus.FINDING,
        "VSAN-DEEP-004": DeepResultStatus.FINDING,
    }
    assert statuses["VSAN-TREND-001"] == DeepResultStatus.INSUFFICIENT_DATA
    resync_record = next(record for record in records if record.metadata["rule_id"] == "VSAN-DEEP-003")
    assert resync_record.source.api == "VsanVcClusterHealthSystem/VsanObjectSystem"
    assert resync_record.value["suspended_object_count"] == 1


def test_vsan_enhanced_rules_are_not_evaluated_when_management_evidence_is_missing() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    inventory = {
        "status": "collected",
        "clusters": {
            "vSAN-Cluster": {
                "health_issues": None,
                "object_health_issues": None,
                "resync_suspended_object_count": None,
                "disk_health_issues": None,
            }
        },
    }
    capabilities = collector._vsan_capabilities(inventory)
    assert {item.id: item.status for item in capabilities} == {
        "vsan": CapabilityStatus.AVAILABLE,
        "vsan.health.api": CapabilityStatus.UNAVAILABLE,
        "vsan.object.health": CapabilityStatus.UNAVAILABLE,
        "vsan.resync": CapabilityStatus.UNAVAILABLE,
        "vsan.disk_group": CapabilityStatus.UNAVAILABLE,
        "vsan.perf.service": CapabilityStatus.UNAVAILABLE,
        "vsan.congestion": CapabilityStatus.UNAVAILABLE,
        "vsan.latency": CapabilityStatus.UNAVAILABLE,
    }


def test_vsan_capabilities_remain_not_requested_when_resource_guard_skips_collection() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    capabilities = collector._vsan_capabilities(
        {
            "status": "not_requested",
            "collection_error": "memory_limit_exceeded",
            "clusters": {},
            "performance": {"status": "not_requested", "reason": "memory_limit_exceeded"},
        }
    )
    assert all(item.status == CapabilityStatus.NOT_REQUESTED for item in capabilities)


def test_performance_capability_probe_honors_guard_between_objects() -> None:
    class PerfManager:
        historicalInterval = []

        def __init__(self):
            self.entities = []

        def QueryAvailablePerfMetric(self, entity, **_kwargs):
            self.entities.append(entity.name)
            return []

    guard_calls = {"count": 0}

    def resource_guard() -> str | None:
        guard_calls["count"] += 1
        return "memory_limit_exceeded" if guard_calls["count"] >= 3 else None

    manager = PerfManager()
    collector = PyVmomiDeepCollector("vc", "u", "p")
    matrix = collector._probe_capabilities(
        NS(taskManager=None, alarmManager=None, about=None, licenseManager=None, perfManager=manager),
        [NS(name="esx-1", _moId="host-1"), NS(name="esx-2", _moId="host-2")],
        [],
        [],
        resource_guard=resource_guard,
    )

    capability = next(item for item in matrix.capabilities if item.id == "perf.query_available_metric")
    assert manager.entities == ["esx-1"]
    assert capability.status == CapabilityStatus.NOT_REQUESTED
    assert capability.detail["reason"] == "memory_limit_exceeded"
    assert capability.detail["remaining_object_count"] == 1


def test_scope_aware_finding_id_changes_with_scope() -> None:
    entity = FindingScope(scope_type=ScopeType.ENTITY, scope_id="host-1", dimension_key="cert")
    cluster = FindingScope(scope_type=ScopeType.CLUSTER, scope_id="cluster-1", dimension_key="cert")
    aggregate = FindingScope(scope_type=ScopeType.AGGREGATE, scope_id="aggregate", dimension_key="cert", member_ids=["b", "a"])
    assert finding_id("R-1", entity) != finding_id("R-1", cluster)
    assert finding_id("R-1", aggregate) == finding_id("R-1", aggregate.model_copy(update={"member_ids": ["a", "b"]}))


def test_deep_history_accumulates_and_resolves_findings(tmp_path: Path) -> None:
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"created_at_utc": "2026-01-01T03:10:00Z"})
    analysis = DeepAnalyzer().analyze(dataset)
    store = DeepHistoryStore(tmp_path / "deep-history.db")
    first = store.record(dataset, analysis)
    assert first

    second_manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-deep-fixture-002", "created_at_utc": "2026-01-05T03:10:00Z"})
    second = dataset.model_copy(
        update={
            "manifest": second_manifest,
            "records": [record.model_copy(update={"dataset_id": "ds-deep-fixture-002"}) for record in dataset.records],
        }
    )
    second_analysis = DeepAnalyzer().analyze(second)
    lifecycles = store.record(second, second_analysis)
    finding = next(iter(lifecycles.values()))
    assert finding["occurrence_count"] == 2
    summary = store.trend_summary("vc-fixture", window_days=30)
    assert summary["state"] == "INSUFFICIENT_DATA"
    assert summary["dataset_count"] == 2

    first_finding = analysis.findings[0]
    insufficient_dataset = second.model_copy(
        update={
            "manifest": second_manifest.model_copy(update={"dataset_id": "ds-deep-fixture-003", "created_at_utc": "2026-04-05T03:10:00Z"}),
            "records": [],
        }
    )
    insufficient_analysis = DeepAnalyzer().analyze(insufficient_dataset)
    store.record(insufficient_dataset, insufficient_analysis)
    summary = store.trend_summary("vc-fixture", window_days=90)
    assert summary["state"] == "READY"
    assert summary["resolved_finding_count"] == 0
    assert summary["unverified_finding_count"] > 0
    with __import__("sqlite3").connect(tmp_path / "deep-history.db") as conn:
        row = conn.execute("SELECT status, lifecycle_json FROM deep_finding_history WHERE finding_id = ?", (first_finding.finding_id,)).fetchone()
        assert row[0] == "OPEN"
        assert json.loads(row[1])["verification_status"] == "UNVERIFIED"

    resolved_dataset = second.model_copy(
        update={
            "manifest": second_manifest.model_copy(update={"dataset_id": "ds-deep-fixture-004", "created_at_utc": "2026-04-06T03:10:00Z"}),
            "records": [record.model_copy(update={"dataset_id": "ds-deep-fixture-004"}) for record in dataset.records],
        }
    )
    resolved_analysis = DeepAnalyzer().analyze(resolved_dataset)
    pass_result = next(result for result in resolved_analysis.rule_results if result.rule_id == first_finding.rule_id).model_copy(
        update={"status": DeepResultStatus.PASS, "finding_ids": [], "reason": ""}
    )
    resolved_analysis.rule_results = [
        pass_result if result.rule_id == first_finding.rule_id else result
        for result in resolved_analysis.rule_results
    ]
    resolved_analysis.findings = [
        finding for finding in resolved_analysis.findings if finding.rule_id != first_finding.rule_id
    ]
    store.record(resolved_dataset, resolved_analysis)
    with __import__("sqlite3").connect(tmp_path / "deep-history.db") as conn:
        row = conn.execute("SELECT status, lifecycle_json FROM deep_finding_history WHERE finding_id = ?", (first_finding.finding_id,)).fetchone()
        assert row[0] == "RESOLVED"
        assert json.loads(row[1])["verification_status"] == "CONFIRMED_RESOLVED"


def test_deep_replay_orders_datasets_and_builds_30_90_day_history(tmp_path: Path) -> None:
    source = DeepDataset.from_json(FIXTURE)
    paths = []
    for index, created_at in enumerate(("2026-01-01T00:00:00Z", "2026-01-05T00:00:00Z", "2026-03-06T00:00:00Z", "2026-03-07T00:00:00Z", "2026-04-05T00:00:00Z"), start=1):
        dataset = source.model_copy(deep=True)
        dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": f"replay-{index}", "created_at_utc": created_at, "target": {"vcenter_ref": "vc-replay"}})
        path = tmp_path / f"input-{index}.json"
        dataset.to_json(path)
        paths.append(path)
    results = DeepInspectionService().replay_datasets(paths, tmp_path / "replay-output")
    assert [item.dataset_id for item in results] == [f"replay-{index}" for index in range(1, 6)]
    trend_log = next(item for item in results[-1].analysis.diagnostic.collection_log if item["action"] == "history.trend")
    assert trend_log["windows"]["30"]["state"] == "READY"
    assert trend_log["windows"]["90"]["state"] == "READY"
    store = DeepHistoryStore(tmp_path / "replay-output" / "deep-history.db")
    assert store.trend_summary("vc-replay", window_days=30)["state"] == "READY"
    assert store.trend_summary("vc-replay", window_days=90)["state"] == "READY"


def test_capacity_history_emits_per_datastore_trend_evidence(tmp_path: Path) -> None:
    def dataset_with_capacity(dataset_id: str, collected_at: str, values: dict[str, float]) -> DeepDataset:
        dataset = DeepDataset.from_json(FIXTURE)
        dataset.manifest = dataset.manifest.model_copy(
            update={"dataset_id": dataset_id, "created_at_utc": collected_at, "target": {"vcenter_ref": "vc-history"}}
        )
        dataset.records = [
            DatasetRecord(
                record_id=f"capacity-{dataset_id}-{entity_id}",
                dataset_id=dataset_id,
                kind="config",
                entity=DeepEntity(type="Datastore", stable_id=entity_id, display_ref=entity_id),
                collected_at_utc=collected_at,
                source=DeepSource(api="Datastore.summary", collector="fixture", collected_at_utc=collected_at),
                selector={"rule_id": "STO-DEEP-004"},
                window=DeepWindow(start=collected_at, end=collected_at, sample_count=1, expected_sample_count=1, completeness=1.0),
                value=value,
                unit="percent",
                raw_pointer=f"config/{entity_id}.ndjson",
                metadata={"rule_id": "STO-DEEP-004"},
            )
            for entity_id, value in values.items()
        ]
        return dataset

    store = DeepHistoryStore(tmp_path / "deep-history.db")
    for day in range(3):
        store.record_metric_observations(
            dataset_with_capacity(
                f"ds-cap-00{day + 1}",
                f"2026-09-{20 + day:02d}T00:00:00Z",
                {"ds-a": 60.0 + day, "ds-b": 10.0 + day * 0.1},
            )
        )
    current = dataset_with_capacity("ds-cap-003", "2026-09-22T00:00:00Z", {"ds-a": 62.0, "ds-b": 10.2})
    records = store.metric_records("vc-history", metric_id="datastore.used_percent", rule_id="CAP-DEEP-001")
    assert len(records) == 6
    assert all(record.metadata["history"] is True for record in records)
    dataset_dir = DeepDatasetWriter().write(current, tmp_path / "dataset", supplemental_records=records)
    metric_file = dataset_dir / "history" / "metrics" / "datastore.used_percent.ndjson"
    assert metric_file.exists()
    assert len(metric_file.read_text(encoding="utf-8").splitlines()) == 6
    manifest = json.loads((dataset_dir / "manifest.json").read_text(encoding="utf-8"))
    assert "history/metrics/datastore.used_percent.ndjson" in manifest["integrity"]["files"]

    current.records = records
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "CAP-DEEP-001")
    analysis = DeepAnalyzer().analyze(current, rules=(rule,))
    assert analysis.rule_results[0].status == DeepResultStatus.FINDING
    assert [item.scope.scope_id for item in analysis.findings] == ["ds-a"]
    assert "约 18 天" in analysis.findings[0].fact


def test_vm_count_and_thin_provision_history_trends(tmp_path: Path) -> None:
    store = DeepHistoryStore(tmp_path / "deep-history.db")
    for day, count, ratio in ((20, 10, 1.0), (21, 12, 1.08), (22, 14, 1.16)):
        dataset = DeepDataset.from_json(FIXTURE)
        dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": f"ds-growth-{day}", "created_at_utc": f"2026-09-{day}T00:00:00Z", "target": {"vcenter_ref": "vc-growth"}})
        dataset.records = [
            DatasetRecord(
                record_id=f"growth-{rule_id}-{day}",
                dataset_id=dataset.dataset_id,
                kind="perf",
                entity=DeepEntity(type="Environment", stable_id="environment", display_ref="vc-growth"),
                collected_at_utc=f"2026-09-{day}T00:00:00Z",
                source=DeepSource(api="fixture.inventory", collector="fixture"),
                selector={"rule_id": rule_id},
                window=DeepWindow(start=f"2026-09-{day}T00:00:00Z", end=f"2026-09-{day}T00:00:00Z", sample_count=1, expected_sample_count=1, completeness=1.0),
                value=value,
                unit=unit,
                metadata={"rule_id": rule_id},
            )
            for rule_id, value, unit in (("CAP-DEEP-002", count, "count"), ("CAP-DEEP-003", ratio, "ratio"))
        ]
        store.record_metric_observations(dataset)
    records = [
        *store.metric_records("vc-growth", metric_id="vm.count", rule_id="CAP-DEEP-002"),
        *store.metric_records("vc-growth", metric_id="thin_provision.ratio", rule_id="CAP-DEEP-003"),
    ]
    assert len(records) == 6
    dataset.records = records
    rules = tuple(rule for rule in BUILTIN_DEEP_RULES if rule.rule_id in {"CAP-DEEP-002", "CAP-DEEP-003"})
    analysis = DeepAnalyzer().analyze(dataset, rules=rules)
    assert {result.rule_id: result.status for result in analysis.rule_results} == {"CAP-DEEP-002": DeepResultStatus.FINDING, "CAP-DEEP-003": DeepResultStatus.FINDING}


def test_trend_report_labels_unverified_findings_separately_from_resolved() -> None:
    dataset = DeepDataset.from_json(FIXTURE)
    report = DeepAnalyzer().analyze(dataset).report
    report.trend_summaries = {
        "30": {
            "state": "READY",
            "window_days": 30,
            "dataset_count": 4,
            "finding_count_baseline": 5,
            "finding_count_current": 3,
            "new_finding_count": 1,
            "resolved_finding_count": 1,
            "persistent_finding_count": 2,
            "unverified_finding_count": 2,
            "comparison_complete": False,
        }
    }

    html = DeepHtmlReportBuilder()._html(report, dataset)

    assert "未核实 2" in html
    assert "未核实不计作已解决" in html


def test_cluster_evc_resource_pool_summary_only_flags_inconsistent_evc() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    clusters = [
        NS(name="cluster-a", summary=NS(currentEVCModeKey="intel-broadwell"), configurationEx=NS(dasConfig=None, drsConfig=None), resourcePool=NS(resourcePool=[NS(name="rp-a")])),
        NS(name="cluster-b", summary=NS(currentEVCModeKey="intel-skylake"), configurationEx=NS(dasConfig=None, drsConfig=None), resourcePool=NS(resourcePool=[])),
    ]
    records = collector._cluster_hidden_risk_records(clusters, "ds-evc", "2026-09-23T01:00:00Z")
    record = next(item for item in records if item.metadata["rule_id"] == "CL-DEEP-008")
    assert record.finding is True
    assert record.value["evc_modes"] == ["intel-broadwell", "intel-skylake"]
    assert record.value["resource_pool_counts"] == [1, 0]


def test_cluster_host_drift_compares_ntp_dns_and_syslog_state_without_serializing_targets() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    hosts = [
        NS(
            name="esx-a",
            _moId="host-a",
            config=NS(
                dateTimeInfo=NS(ntpConfig=NS(server=["ntp-a.example"])),
                network=NS(dnsConfig=NS(domainName="lab.example", address=["192.0.2.10"])),
                option=[NS(key="Syslog.global.logHost", value="tcp://syslog-a.example:514")],
            ),
        ),
        NS(
            name="esx-b",
            _moId="host-b",
            config=NS(
                dateTimeInfo=NS(ntpConfig=NS(server=["ntp-b.example"])),
                network=NS(dnsConfig=NS(domainName="lab.example", address=["192.0.2.10"])),
                option=[NS(key="Syslog.global.logHost", value="tcp://syslog-b.example:514")],
            ),
        ),
    ]
    cluster = NS(name="cluster-drift", _moId="domain-drift", host=hosts)

    record = next(
        item
        for item in collector._cluster_drift_records([cluster], "ds-host-drift", "2026-09-28T01:00:00Z")
        if item.metadata.get("rule_id") == "CL-DEEP-001"
    )

    assert record.finding is True
    assert {item["property"] for item in record.value["differences"]} == {"ntp_servers", "syslog_target_set"}
    assert "syslog-a.example" not in record.model_dump_json()
    assert "syslog-b.example" not in record.model_dump_json()
    assert record.value["syslog_target_content_included"] is False


def test_cluster_host_drift_missing_fields_is_insufficient_not_pass() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    cluster = NS(name="cluster-drift-unknown", _moId="domain-drift-unknown", host=[NS(name="esx-a", config=NS()), NS(name="esx-b", config=NS())])
    record = next(
        item
        for item in collector._cluster_drift_records([cluster], "ds-host-drift-unknown", "2026-09-28T01:00:00Z")
        if item.metadata.get("rule_id") == "CL-DEEP-001"
    )
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-host-drift-unknown"})
    dataset.records = [item for item in dataset.records if item.metadata.get("rule_id") != "CL-DEEP-001"] + [record]
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "CL-DEEP-001")

    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert record.window.completeness == 0.0
    assert result.status == DeepResultStatus.INSUFFICIENT_DATA


def test_resource_pool_tree_status_reports_red_and_yellow_with_config_evidence() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    child = NS(
        name="production",
        _moId="rp-2",
        overallStatus="yellow",
        runtime=NS(overallStatus="yellow"),
        config=NS(
            cpuAllocation=NS(reservation=18000, limit=-1, expandableReservation=True, shares=NS(level="normal", shares=4000)),
            memoryAllocation=NS(reservation=32768, limit=65536, expandableReservation=False, shares=NS(level="custom", shares=8192)),
        ),
        resourcePool=[],
    )
    root = NS(name="Resources", _moId="rp-1", overallStatus="red", runtime=NS(overallStatus="red"), resourcePool=[child])
    cluster = NS(name="cluster-a", resourcePool=root, summary=NS(currentEVCModeKey="intel-broadwell"), configurationEx=NS(dasConfig=None, drsConfig=None))

    records = collector._cluster_hidden_risk_records([cluster], "ds-resource-pool", "2026-09-26T01:00:00Z")
    record = next(item for item in records if item.metadata["rule_id"] == "CL-DEEP-009")
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records = [record]
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "CL-DEEP-009")
    analysis = DeepAnalyzer().analyze(dataset, rules=(rule,))

    assert record.finding is True
    assert record.value["status_counts"] == {"green": 0, "yellow": 1, "red": 1, "unknown": 0}
    affected_child = next(item for item in record.value["affected_pools"] if item["path"].endswith("/production"))
    assert affected_child["memory"]["reservation"] == 32768
    assert affected_child["memory"]["shares_level"] == "custom"
    assert analysis.rule_results[0].status == DeepResultStatus.FINDING
    assert "红色节点状态" in analysis.findings[0].fact


def test_resource_pool_tree_unknown_or_truncated_status_never_passes(monkeypatch: pytest.MonkeyPatch) -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "CL-DEEP-009")

    unknown_cluster = NS(
        name="cluster-unknown",
        resourcePool=NS(name="Resources", overallStatus="gray", runtime=NS(overallStatus="gray"), resourcePool=[]),
        summary=NS(currentEVCModeKey=""),
        configurationEx=NS(dasConfig=None, drsConfig=None),
    )
    unknown_record = next(
        item for item in collector._cluster_hidden_risk_records([unknown_cluster], "ds-resource-pool-unknown", "2026-09-26T01:00:00Z")
        if item.metadata["rule_id"] == "CL-DEEP-009"
    )
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records = [unknown_record]
    unknown_result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]
    assert unknown_result.status == DeepResultStatus.INSUFFICIENT_DATA

    hidden_child = NS(name="nested", overallStatus="red", runtime=NS(overallStatus="red"), resourcePool=[])
    root = NS(name="Resources", overallStatus="green", runtime=NS(overallStatus="green"), resourcePool=[hidden_child])
    truncated_cluster = NS(name="cluster-truncated", resourcePool=root, summary=NS(currentEVCModeKey=""), configurationEx=NS(dasConfig=None, drsConfig=None))
    monkeypatch.setattr("vstacklens.deep.pyvmomi_collector.MAX_RESOURCE_POOLS_PER_CLUSTER", 1)
    truncated_record = next(
        item for item in collector._cluster_hidden_risk_records([truncated_cluster], "ds-resource-pool-truncated", "2026-09-26T01:00:00Z")
        if item.metadata["rule_id"] == "CL-DEEP-009"
    )
    dataset.records = [truncated_record]
    truncated_result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]
    assert truncated_record.value["coverage_complete"] is False
    assert truncated_result.status == DeepResultStatus.INSUFFICIENT_DATA

    monkeypatch.setattr("vstacklens.deep.pyvmomi_collector.MAX_RESOURCE_POOLS_PER_SESSION", 1)
    green_cluster = NS(
        name="cluster-first",
        resourcePool=NS(name="Resources", overallStatus="green", runtime=NS(overallStatus="green"), resourcePool=[]),
        summary=NS(currentEVCModeKey=""),
        configurationEx=NS(dasConfig=None, drsConfig=None),
    )
    red_cluster = NS(
        name="cluster-skipped",
        resourcePool=NS(name="Resources", overallStatus="red", runtime=NS(overallStatus="red"), resourcePool=[]),
        summary=NS(currentEVCModeKey=""),
        configurationEx=NS(dasConfig=None, drsConfig=None),
    )
    capped_record = next(
        item for item in collector._cluster_hidden_risk_records([green_cluster, red_cluster], "ds-resource-pool-session-cap", "2026-09-26T01:00:00Z")
        if item.metadata["rule_id"] == "CL-DEEP-009"
    )
    dataset.records = [capped_record]
    capped_result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]
    assert capped_record.value["clusters"][1]["error_type"] == "session_resource_pool_limit"
    assert capped_result.status == DeepResultStatus.INSUFFICIENT_DATA


def test_resource_pool_partial_tree_keeps_confirmed_finding_and_reports_gap() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    root = NS(name="Resources", overallStatus="red", runtime=NS(overallStatus="red"), resourcePool=None)
    cluster = NS(name="cluster-partial", resourcePool=root, summary=NS(currentEVCModeKey=""), configurationEx=NS(dasConfig=None, drsConfig=None))
    record = next(
        item for item in collector._cluster_hidden_risk_records([cluster], "ds-resource-pool-partial", "2026-09-26T01:00:00Z")
        if item.metadata["rule_id"] == "CL-DEEP-009"
    )
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records = [record]
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "CL-DEEP-009")
    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert record.value["coverage_complete"] is False
    assert result.status == DeepResultStatus.FINDING
    assert "Resource Pool 资源树状态异常其他范围仍有未取得状态" in result.reason


def test_resource_pool_status_falls_back_to_legacy_runtime_and_rejects_alias_conflict() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "CL-DEEP-009")
    legacy_cluster = NS(
        name="cluster-legacy",
        resourcePool=NS(name="Resources", runtime=NS(overallStatus="yellow"), resourcePool=[]),
        summary=NS(currentEVCModeKey=""),
        configurationEx=NS(dasConfig=None, drsConfig=None),
    )
    legacy_record = next(
        item for item in collector._cluster_hidden_risk_records([legacy_cluster], "ds-resource-pool-legacy", "2026-09-26T01:00:00Z")
        if item.metadata["rule_id"] == "CL-DEEP-009"
    )
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records = [legacy_record]
    legacy_result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]
    assert legacy_record.value["affected_pools"][0]["status_source"] == "ResourcePool.runtime.overallStatus"
    assert legacy_result.status == DeepResultStatus.FINDING

    conflict_cluster = NS(
        name="cluster-conflict",
        resourcePool=NS(name="Resources", overallStatus="green", runtime=NS(overallStatus="red"), resourcePool=[]),
        summary=NS(currentEVCModeKey=""),
        configurationEx=NS(dasConfig=None, drsConfig=None),
    )
    conflict_record = next(
        item for item in collector._cluster_hidden_risk_records([conflict_cluster], "ds-resource-pool-conflict", "2026-09-26T01:00:00Z")
        if item.metadata["rule_id"] == "CL-DEEP-009"
    )
    dataset.records = [conflict_record]
    conflict_result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]
    assert conflict_record.value["coverage_complete"] is False
    assert conflict_result.status == DeepResultStatus.INSUFFICIENT_DATA


def test_vcenter_rest_reader_uses_cached_get_and_closes_only_its_session(monkeypatch: pytest.MonkeyPatch) -> None:
    from vstacklens.deep import vcenter_rest

    calls: list[tuple[str, str, str | None]] = []

    class Response:
        def __init__(self, value: bytes) -> None:
            self.value = value

        def __enter__(self) -> "Response":
            return self

        def __exit__(self, *_args: object) -> None:
            return None

        def read(self) -> bytes:
            return self.value

    def fake_urlopen(request: object, *, context: object, timeout: float) -> Response:
        method = request.get_method()
        url = request.full_url
        calls.append((method, url, request.get_header("Authorization")))
        if method == "POST":
            return Response(b'"temporary-session"')
        if method == "GET":
            return Response(b'{"overall_compliance":"COMPLIANT"}')
        return Response(b"")

    monkeypatch.setattr(vcenter_rest.urllib.request, "urlopen", fake_urlopen)
    reader = VCenterRestReadClient("vc.example", "readonly", "secret", ssl_verify=False)
    reader.open()
    payload = reader.get_json("/api/vcenter/vm/vm-1/storage/policy/compliance")
    reader.close()

    assert payload == {"overall_compliance": "COMPLIANT"}
    assert [(method, url.rsplit("/", 1)[-1]) for method, url, _auth in calls] == [
        ("POST", "session"),
        ("GET", "compliance"),
        ("DELETE", "session"),
    ]
    assert calls[1][1].endswith("/api/vcenter/vm/vm-1/storage/policy/compliance")
    assert calls[0][2] is not None
    assert calls[1][2] is None


def test_storage_policy_rule_reports_confirmed_noncompliance_and_out_of_date_with_partial_scope(monkeypatch: pytest.MonkeyPatch) -> None:
    import vstacklens.deep.pyvmomi_collector as collector_module

    responses: dict[str, object] = {
        "vm-1": {
            "overall_compliance": "NON_COMPLIANT",
            "vm_home": {"status": "COMPLIANT", "policy": "profile-a", "check_time": "2026-09-26T00:00:00Z", "failure_cause": []},
            "disks": {"disk-1000": {"status": "NON_COMPLIANT", "policy": "profile-b", "check_time": "2026-09-26T00:00:00Z", "failure_cause": [{"id": "provider.failure"}]}},
        },
        "vm-2": None,
        "vm-3": {"overall_compliance": "OUT_OF_DATE", "vm_home": {"status": "OUT_OF_DATE", "policy": "profile-a", "check_time": "2026-09-26T00:00:00Z", "failure_cause": []}, "disks": {}},
        "vm-4": urllib.error.HTTPError("https://vc.example/api", 403, "Forbidden", {}, None),
    }

    class FakeReader:
        def __init__(self, *_args: object, **_kwargs: object) -> None:
            self.closed = False

        def open(self) -> None:
            return None

        def get_json(self, path: str) -> object:
            vm_id = path.split("/vm/", 1)[1].split("/", 1)[0]
            value = responses[vm_id]
            if isinstance(value, Exception):
                raise value
            return value

        def close(self) -> None:
            self.closed = True

    fake_readers: list[FakeReader] = []

    def make_reader(*args: object, **kwargs: object) -> FakeReader:
        item = FakeReader(*args, **kwargs)
        fake_readers.append(item)
        return item

    monkeypatch.setattr(collector_module, "VCenterRestReadClient", make_reader)
    collector = PyVmomiDeepCollector("vc.example", "readonly", "not-a-real-password")
    vms = [NS(name=f"vm-{index}", _moId=f"vm-{index}", config=NS(instanceUuid=f"uuid-{index}")) for index in range(1, 5)]
    records = collector._storage_policy_compliance_records(vms, "ds-storage-policy", "2026-09-26T01:00:00Z")
    summary = next(item for item in records if item.metadata.get("storage_policy_summary"))
    vm_records = [item for item in records if not item.metadata.get("storage_policy_summary")]
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "STO-DEEP-006")
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-storage-policy"})
    dataset.records = records
    dataset.capability.capabilities.append(
        Capability(id="storage.policy.compliance", name="policy", status=CapabilityStatus.LIMITED, profile="core")
    )
    analysis = DeepAnalyzer().analyze(dataset, rules=(rule,))
    result = analysis.rule_results[0]

    assert fake_readers[0].closed is True
    assert summary.value["collection_status"] == "limited"
    assert summary.value["compliance_status_counts"]["non_compliant"] == 1
    assert summary.value["compliance_status_counts"]["out_of_date"] == 1
    assert summary.value["no_policy_association_vm_count"] == 1
    assert next(item for item in vm_records if item.entity.display_ref == "vm-1").finding is True
    assert next(item for item in vm_records if item.entity.display_ref == "vm-2").metadata["not_applicable"] is True
    assert result.status == DeepResultStatus.FINDING
    assert "STO-DEEP-006" == analysis.findings[0].rule_id
    assert {entity.display_ref for entity in analysis.findings[0].entities} == {"vm-1", "vm-3"}
    assert "其他范围仍有未取得状态" in result.reason


def test_storage_policy_unknown_and_conflicting_rollup_never_report_pass(monkeypatch: pytest.MonkeyPatch) -> None:
    import vstacklens.deep.pyvmomi_collector as collector_module

    responses = {
        "vm-unknown": {"overall_compliance": "UNRECOGNIZED", "vm_home": {"status": "UNKNOWN", "policy": "profile-a", "failure_cause": []}, "disks": {}},
        "vm-conflict": {"overall_compliance": "COMPLIANT", "vm_home": {"status": "NON_COMPLIANT", "policy": "profile-a", "failure_cause": []}, "disks": {}},
        "vm-missing-disks": {"overall_compliance": "COMPLIANT", "vm_home": {"status": "COMPLIANT", "policy": "profile-a", "failure_cause": []}},
    }

    class FakeReader:
        def __init__(self, *_args: object, **_kwargs: object) -> None:
            return None

        def open(self) -> None:
            return None

        def get_json(self, path: str) -> object:
            vm_id = path.split("/vm/", 1)[1].split("/", 1)[0]
            return responses[vm_id]

        def close(self) -> None:
            return None

    monkeypatch.setattr(collector_module, "VCenterRestReadClient", FakeReader)
    collector = PyVmomiDeepCollector("vc.example", "readonly", "not-a-real-password")
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "STO-DEEP-006")
    for vm_id in ("vm-unknown", "vm-conflict", "vm-missing-disks"):
        vm = NS(name=vm_id, _moId=vm_id, config=NS(instanceUuid=vm_id))
        records = collector._storage_policy_compliance_records([vm], f"ds-{vm_id}", "2026-09-26T01:00:00Z")
        dataset = DeepDataset.from_json(FIXTURE)
        dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": f"ds-{vm_id}"})
        dataset.records = records
        dataset.capability.capabilities.append(
            Capability(id="storage.policy.compliance", name="policy", status=CapabilityStatus.AVAILABLE, profile="core")
        )
        result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]
        assert result.status == DeepResultStatus.INSUFFICIENT_DATA
        if vm_id == "vm-missing-disks":
            record = next(item for item in records if not item.metadata.get("storage_policy_summary"))
            assert record.metadata["coverage_complete"] is False


def test_storage_policy_not_applicable_is_distinct_from_missing_policy_association(monkeypatch: pytest.MonkeyPatch) -> None:
    import vstacklens.deep.pyvmomi_collector as collector_module

    responses = {
        "vm-none": None,
        "vm-explicit-na": {"overall_compliance": "NOT_APPLICABLE", "vm_home": None, "disks": {}},
    }

    class FakeReader:
        def __init__(self, *_args: object, **_kwargs: object) -> None:
            return None

        def open(self) -> None:
            return None

        def get_json(self, path: str) -> object:
            vm_id = path.split("/vm/", 1)[1].split("/", 1)[0]
            return responses[vm_id]

        def close(self) -> None:
            return None

    monkeypatch.setattr(collector_module, "VCenterRestReadClient", FakeReader)
    collector = PyVmomiDeepCollector("vc.example", "readonly", "not-a-real-password")
    vms = [NS(name=vm_id, _moId=vm_id, config=NS(instanceUuid=vm_id)) for vm_id in responses]
    records = collector._storage_policy_compliance_records(vms, "ds-storage-policy-na", "2026-09-26T01:00:00Z")
    summary = next(item for item in records if item.metadata.get("storage_policy_summary"))
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-storage-policy-na"})
    dataset.records = records
    dataset.capability.capabilities.append(
        Capability(id="storage.policy.compliance", name="policy", status=CapabilityStatus.AVAILABLE, profile="core")
    )
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "STO-DEEP-006")
    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert summary.value["compliance_status_counts"]["not_applicable"] == 2
    assert summary.value["no_policy_association_vm_count"] == 1
    assert "1 个 GET 未返回策略关联信息" in summary.summary
    assert result.status == DeepResultStatus.NOT_APPLICABLE

    empty_dataset = DeepDataset.from_json(FIXTURE)
    empty_dataset.manifest = empty_dataset.manifest.model_copy(update={"scope": {"hosts": 0, "vms": 0, "clusters": 0, "datastores": 0}})
    empty_dataset.records = []
    empty_dataset.capability.capabilities.append(
        Capability(id="storage.policy.compliance", name="policy", status=CapabilityStatus.NOT_APPLICABLE, profile="core")
    )
    empty_result = DeepAnalyzer().analyze(empty_dataset, rules=(rule,)).rule_results[0]
    assert empty_result.status == DeepResultStatus.NOT_APPLICABLE


def test_storage_policy_permission_denial_marks_capability_unavailable(monkeypatch: pytest.MonkeyPatch) -> None:
    import vstacklens.deep.pyvmomi_collector as collector_module

    class DeniedReader:
        def __init__(self, *_args: object, **_kwargs: object) -> None:
            return None

        def open(self) -> None:
            raise urllib.error.HTTPError("https://vc.example/api/session", 403, "Forbidden", {}, None)

        def close(self) -> None:
            return None

    monkeypatch.setattr(collector_module, "VCenterRestReadClient", DeniedReader)
    collector = PyVmomiDeepCollector("vc.example", "readonly", "not-a-real-password")
    vm = NS(name="vm-1", _moId="vm-1", config=NS(instanceUuid="uuid-1"))
    records = collector._storage_policy_compliance_records([vm], "ds-storage-policy-denied", "2026-09-26T01:00:00Z")
    capability = collector._storage_policy_capability(records, 1, {"status": "unavailable"})
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-storage-policy-denied"})
    dataset.records = records
    dataset.capability.capabilities.append(capability)
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "STO-DEEP-006")
    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert capability.status == CapabilityStatus.UNAVAILABLE
    assert result.status == DeepResultStatus.NOT_EVALUATED
    assert "storage.policy.compliance" in result.missing_capabilities


def test_datastore_access_and_vmfs_mount_evidence_preserve_partial_risk_without_capacity_false_pass() -> None:
    collector = PyVmomiDeepCollector("vc.example", "readonly", "not-a-real-password")
    def datastore(name: str, accessible: bool | None, mount_access: list[bool | None], *, datastore_type: str = "VMFS", multiple: bool = True) -> NS:
        mounts = [
            NS(
                key=NS(name=f"esx-{index + 1}"),
                mountInfo=NS(mounted=True, accessible=value, inaccessibleReason="AllPathsDown_Timeout" if value is False else None),
            )
            for index, value in enumerate(mount_access)
        ]
        vmfs = NS(version="VMFS-6", majorVersion=6, extent=[NS(name="naa.device")], vmfsUpgradable=True) if datastore_type.casefold() == "vmfs" else None
        return NS(
            name=name,
            _moId=f"datastore-{name}",
            summary=NS(type=datastore_type, accessible=accessible, multipleHostAccess=multiple, capacity=1000, freeSpace=100),
            host=mounts,
            info=NS(vmfs=vmfs) if vmfs is not None else NS(),
        )

    inaccessible = datastore("ds-inaccessible", False, [False])
    access_records = collector._datastore_access_records([inaccessible], "ds-access", "2026-09-26T02:00:00Z")
    access_record = next(item for item in access_records if not item.metadata.get("storage_access_summary"))
    access_dataset = DeepDataset.from_json(FIXTURE)
    access_dataset.manifest = access_dataset.manifest.model_copy(update={"dataset_id": "ds-access"})
    access_dataset.records = access_records
    access_rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "STO-DEEP-007")
    access_result = DeepAnalyzer().analyze(access_dataset, rules=(access_rule,)).rule_results[0]
    assert access_record.finding is True
    assert access_record.value["vmfs"]["version"] == "VMFS-6"
    assert access_record.value["mounts"][0]["inaccessible_reason"] == "allpathsdowntimeout"
    assert access_result.status == DeepResultStatus.FINDING

    capacity_records = collector._storage_hidden_risk_records(
        [], [datastore("ds-full", True, [True]), inaccessible], "ds-capacity-partial", "2026-09-26T02:00:00Z"
    )
    inaccessible_capacity = next(item for item in capacity_records if item.entity.display_ref == "ds-inaccessible")
    capacity_dataset = DeepDataset.from_json(FIXTURE)
    capacity_dataset.manifest = capacity_dataset.manifest.model_copy(update={"dataset_id": "ds-capacity-partial"})
    capacity_dataset.records = capacity_records
    capacity_rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "STO-DEEP-004")
    capacity_result = DeepAnalyzer().analyze(capacity_dataset, rules=(capacity_rule,)).rule_results[0]
    assert inaccessible_capacity.value is None
    assert inaccessible_capacity.window.completeness == 0.0
    assert capacity_result.status == DeepResultStatus.FINDING
    assert "其他范围仍有未取得状态" in capacity_result.reason

    only_inaccessible_capacity = collector._storage_hidden_risk_records([], [inaccessible], "ds-capacity-unavailable", "2026-09-26T02:00:00Z")
    capacity_dataset.manifest = capacity_dataset.manifest.model_copy(update={"dataset_id": "ds-capacity-unavailable"})
    capacity_dataset.records = only_inaccessible_capacity
    inaccessible_only_result = DeepAnalyzer().analyze(capacity_dataset, rules=(capacity_rule,)).rule_results[0]
    assert inaccessible_only_result.status == DeepResultStatus.INSUFFICIENT_DATA


def test_datastore_summary_true_does_not_hide_host_mount_failure_and_unknown_never_passes() -> None:
    collector = PyVmomiDeepCollector("vc.example", "readonly", "not-a-real-password")
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "STO-DEEP-007")
    mixed = NS(
        name="ds-mixed",
        _moId="datastore-mixed",
        summary=NS(type="NFS", accessible=True, multipleHostAccess=True, capacity=1000, freeSpace=200),
        host=[
            NS(key=NS(name="esx-a"), mountInfo=NS(mounted=True, accessible=False, inaccessibleReason="PermanentDeviceLoss")),
            NS(key=NS(name="esx-b"), mountInfo=NS(mounted=True, accessible=True, inaccessibleReason=None)),
        ],
        info=NS(),
    )
    unknown = NS(
        name="ds-unknown",
        _moId="datastore-unknown",
        summary=NS(type="NFS", accessible=None, multipleHostAccess=True, capacity=1000, freeSpace=200),
        host=[NS(key=NS(name="esx-c"), mountInfo=NS(mounted=True, accessible=None))],
        info=NS(),
    )
    records = collector._datastore_access_records([mixed], "ds-mixed", "2026-09-26T02:00:00Z")
    record = next(item for item in records if not item.metadata.get("storage_access_summary"))
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-mixed"})
    dataset.records = records
    mixed_result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]
    assert record.value["state"] == "host_mount_issue"
    assert record.finding is True
    assert mixed_result.status == DeepResultStatus.FINDING

    unknown_records = collector._datastore_access_records([unknown], "ds-unknown-access", "2026-09-26T02:00:00Z")
    unknown_dataset = DeepDataset.from_json(FIXTURE)
    unknown_dataset.manifest = unknown_dataset.manifest.model_copy(update={"dataset_id": "ds-unknown-access"})
    unknown_dataset.records = unknown_records
    unknown_result = DeepAnalyzer().analyze(unknown_dataset, rules=(rule,)).rule_results[0]
    assert unknown_result.status == DeepResultStatus.INSUFFICIENT_DATA


def test_storage_path_risk_includes_shared_lun_redundancy_and_excludes_local_disks() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    host = NS(
        name="esx-paths",
        _moId="host-paths",
        config=NS(
            instanceUuid="host-paths-id",
            storageDevice=NS(
                scsiLun=[NS(key="local-lun", localDisk=True)],
                multipathInfo=NS(
                    lun=[
                        NS(key="naa-redundant", lun=NS(key="naa-redundant", localDisk=False), path=[NS(pathState="active"), NS(pathState="standby")]),
                        NS(key="naa-single", lun=NS(key="naa-single", localDisk=False), path=[NS(pathState="active")]),
                        NS(key="local-lun", lun=NS(key="local-lun", localDisk=True), path=[NS(pathState="active")]),
                    ]
                ),
            ),
        ),
    )
    records = collector._storage_hidden_risk_records([host], [], "ds-path-redundancy", "2026-09-28T01:00:00Z")
    record = next(item for item in records if item.metadata.get("rule_id") == "STO-DEEP-003")
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-path-redundancy"})
    dataset.records = [item for item in dataset.records if item.metadata.get("rule_id") != "STO-DEEP-003"] + [record]
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "STO-DEEP-003")

    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert record.finding is True
    assert record.value["shared_block_lun_count"] == 2
    assert record.value["local_lun_excluded_count"] == 1
    assert record.value["single_or_zero_usable_path_count"] == 1
    assert record.value["failed_path_count"] == 0
    assert result.status == DeepResultStatus.FINDING


def test_storage_path_missing_multipath_inventory_is_insufficient_not_pass() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    host = NS(name="esx-paths-unknown", _moId="host-paths-unknown", config=NS(instanceUuid="host-paths-unknown-id", storageDevice=NS()))
    record = next(
        item
        for item in collector._storage_hidden_risk_records([host], [], "ds-path-unavailable", "2026-09-28T01:00:00Z")
        if item.metadata.get("rule_id") == "STO-DEEP-003"
    )
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-path-unavailable"})
    dataset.records = [item for item in dataset.records if item.metadata.get("rule_id") != "STO-DEEP-003"] + [record]
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "STO-DEEP-003")

    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert record.window.completeness == 0.0
    assert result.status == DeepResultStatus.INSUFFICIENT_DATA


def test_storage_path_multipath_link_value_maps_local_scsi_device() -> None:
    collector = PyVmomiDeepCollector("vcenter.example", "readonly", "not-a-real-password")
    host = NS(
        name="esx-local-link",
        _moId="host-local-link",
        config=NS(
            instanceUuid="host-local-link-id",
            storageDevice=NS(
                scsiLun=[NS(key="scsi-disk-local", uuid="naa-local", localDisk=True)],
                multipathInfo=NS(
                    lun=[NS(key="mp-local", id="naa-local", lun=NS(value="scsi-disk-local"), path=[NS(pathState="active")])]
                ),
            ),
        ),
    )

    record = next(
        item
        for item in collector._storage_hidden_risk_records([host], [], "ds-local-link", "2026-09-28T01:00:00Z")
        if item.metadata.get("rule_id") == "STO-DEEP-003"
    )

    assert record.metadata["not_applicable"] is True
    assert record.value["local_lun_excluded_count"] == 1
    assert record.value["shared_block_lun_count"] == 0
    assert record.finding is False


def test_ha_default_heartbeat_selection_and_isolation_do_not_create_false_finding() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    cluster = NS(
        name="cluster-auto-heartbeat",
        _moId="domain-ha-auto",
        configurationEx=NS(
            dasConfig=NS(
                enabled=True,
                admissionControlEnabled=True,
                heartbeatDatastore=[],
                hBDatastoreCandidatePolicy=None,
                defaultVmSettings=NS(isolationResponse=None),
            ),
            drsConfig=NS(enabled=True),
            rule=[],
        ),
        summary=NS(currentEVCModeKey="intel-skylake"),
        resourcePool=NS(resourcePool=[]),
        host=[],
    )
    record = next(
        item
        for item in collector._cluster_hidden_risk_records([cluster], "ds-ha-default", "2026-09-25T00:00:00Z")
        if item.metadata["rule_id"] == "CL-DEEP-004"
    )
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records = [item for item in dataset.records if item.metadata.get("rule_id") != "CL-DEEP-004"] + [record]
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "CL-DEEP-004")
    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert record.finding is False
    assert record.value["heartbeat_selection_mode"] == "automatic"
    assert record.value["effective_isolation_response"] == "powerOff"
    assert record.value["isolation_response_source"] == "vSphere_default"
    assert result.status == DeepResultStatus.PASS


def test_ha_explicit_no_heartbeat_or_no_isolation_is_flagged() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    cluster = NS(
        name="cluster-user-heartbeat",
        _moId="domain-ha-user",
        configurationEx=NS(
            dasConfig=NS(
                enabled=True,
                admissionControlEnabled=True,
                heartbeatDatastore=[],
                hBDatastoreCandidatePolicy="userSelectedDs",
                defaultVmSettings=NS(isolationResponse="none"),
                dasVmConfig=[],
            ),
            drsConfig=NS(enabled=True),
            rule=[],
        ),
        summary=NS(currentEVCModeKey="intel-skylake"),
        resourcePool=NS(resourcePool=[]),
        host=[],
    )
    record = next(
        item
        for item in collector._cluster_hidden_risk_records([cluster], "ds-ha-risk", "2026-09-25T00:00:00Z")
        if item.metadata["rule_id"] == "CL-DEEP-004"
    )
    assert record.finding is True
    assert "user_selected_heartbeat_datastores_below_two" in record.value["risk_reasons"]
    assert "isolation_response_none" in record.value["risk_reasons"]


def test_ha_missing_candidate_policy_does_not_return_pass() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    cluster = NS(
        name="cluster-ha-unknown-policy",
        _moId="domain-ha-unknown",
        configurationEx=NS(
            dasConfig=NS(enabled=True, admissionControlEnabled=True, heartbeatDatastore=None, hBDatastoreCandidatePolicy="futurePolicy"),
            drsConfig=NS(enabled=True),
            rule=[],
        ),
        summary=NS(currentEVCModeKey="intel-skylake"),
        resourcePool=NS(resourcePool=[]),
        host=[],
    )
    record = next(
        item
        for item in collector._cluster_hidden_risk_records([cluster], "ds-ha-unknown", "2026-09-25T00:00:00Z")
        if item.metadata["rule_id"] == "CL-DEEP-004"
    )
    assert record.window.completeness == 0.0


def test_log_fixture_states_keep_status_limits_and_markers() -> None:
    payload = json.loads((Path(__file__).parent / "fixtures" / "deep_logs_cases.json").read_text(encoding="utf-8"))
    entity = DeepEntity(type="HostSystem", stable_id="host-log", display_ref="esx-log")
    records = normalize_powercli_logs(dataset_id="ds-logs", entity=entity, collected_at="2026-09-24T00:00:00Z", payloads=list(payload.values()), max_lines=200)
    by_category = {record.metadata["log_category"]: record for record in records}
    assert by_category["vmkernel"].value["marker_counts"]["storage"] == 1
    assert by_category["hostd"].metadata["log_status"] == "permission_denied"
    assert by_category["vpxa"].window.completeness == 1.0
    assert by_category["vobd"].value["truncated"] is True
    assert by_category["syslog"].window.completeness == 0.0


def test_fdm_descriptor_is_classified_as_ha_log() -> None:
    descriptor = NS(key="Fdm", fileName="/var/run/log/fdm.log", creator="Fdm", mimeType="text/plain", info=NS(label="Fault Domain Manager", summary=""))
    assert log_category_for_descriptor(descriptor, scope="host") == "ha"
    markers = classify_log_lines(["2026-09-26T00:00:00Z Fdm failover failed on isolated host"])
    assert markers["marker_counts"]["ha"] == 1


def test_log_normalization_enforces_per_file_and_collection_byte_limits() -> None:
    entity = DeepEntity(type="HostSystem", stable_id="host-log", display_ref="esx-log")
    budget = {"max_bytes": 4, "used_bytes": 0}
    records = normalize_powercli_logs(
        dataset_id="ds-log-limits",
        entity=entity,
        collected_at="2026-09-24T00:00:00Z",
        payloads=[
            {"key": "first", "status": "ok", "lines": ["éé"]},
            {"key": "second", "status": "ok", "lines": ["abc"]},
        ],
        max_file_bytes=3,
        byte_budget=budget,
    )
    assert records[0].value["lines"] == ["é"]
    assert records[0].metadata["truncation_reasons"] == ["file_size_limit"]
    assert records[1].value["lines"] == ["ab"]
    assert records[1].metadata["truncation_reasons"] == ["total_size_limit"]
    assert budget["used_bytes"] == 4


@pytest.mark.parametrize(
    ("field", "limit"),
    (("max_log_files", 64), ("max_log_file_bytes", 1024 * 1024), ("max_log_total_bytes", 20 * 1024 * 1024), ("max_log_primary_bytes", 15 * 1024 * 1024), ("max_log_fallback_bytes", 5 * 1024 * 1024)),
)
def test_log_collection_policy_enforces_absolute_budget_caps(field: str, limit: int) -> None:
    assert getattr(DeepCollectionPolicy(**{field: limit}), field) == limit
    with pytest.raises(ValidationError):
        DeepCollectionPolicy(**{field: limit + 1})
    with pytest.raises(ValidationError):
        DeepCollectionPolicy(**{field: -1})


def test_log_collection_policy_defaults_to_12_mib_primary_and_4_mib_fallback() -> None:
    policy = DeepCollectionPolicy()
    assert policy.max_log_total_bytes == 16 * 1024 * 1024
    assert policy.max_log_primary_bytes == 12 * 1024 * 1024
    assert policy.max_log_fallback_bytes == 4 * 1024 * 1024
    assert policy.log_collection_days is None


def test_log_collection_days_do_not_change_event_or_performance_history_windows() -> None:
    policy = DeepCollectionPolicy(log_collection_days=7)
    assert policy.log_collection_days == 7
    assert policy.event_history_days == 30
    assert policy.history_days == 30
    with pytest.raises(ValidationError):
        DeepCollectionPolicy(log_collection_days=0)


def test_log_normalizer_deduplicates_proven_overlapping_source_line_ranges() -> None:
    entity = DeepEntity(type="HostSystem", stable_id="host-log-dedupe", display_ref="esx-dedupe")
    dedupe_records = []
    primary_budget = {"max_bytes": 1024, "used_bytes": 0, "dedupe_records": dedupe_records}
    fallback_budget = {"max_bytes": 1024, "used_bytes": 0, "dedupe_records": dedupe_records}
    primary = normalize_powercli_logs(
        dataset_id="ds-log-dedupe",
        entity=entity,
        collected_at="2026-09-27T00:00:00Z",
        payloads=[{"key": "vmkernel", "status": "ok", "lines": ["line-1", "line-2"], "line_start": 1, "source": "vim.DiagnosticManager.BrowseDiagnosticLog"}],
        byte_budget=primary_budget,
    )
    fallback = normalize_powercli_logs(
        dataset_id="ds-log-dedupe",
        entity=entity,
        collected_at="2026-09-27T00:00:00Z",
        payloads=[{"key": "vmkernel", "status": "ok", "lines": ["line-1", "line-2", "line-3"], "line_start": 1, "source": "vim.DiagnosticManager.DirectESXi.BrowseDiagnosticLog"}],
        byte_budget=fallback_budget,
    )
    assert primary[0].value["lines"] == ["line-1", "line-2"]
    assert fallback[0].value["lines"] == ["line-3"]
    assert fallback[0].metadata["duplicate_line_count"] == 2
    assert primary_budget["used_bytes"] + fallback_budget["used_bytes"] == len("line-1line-2line-3")

    changed_source = normalize_powercli_logs(
        dataset_id="ds-log-dedupe",
        entity=entity,
        collected_at="2026-09-27T00:00:00Z",
        payloads=[{"key": "vmkernel", "status": "ok", "lines": ["different-content"], "line_start": 1, "source": "vim.DiagnosticManager.DirectESXi.BrowseDiagnosticLog"}],
        byte_budget=fallback_budget,
    )
    assert changed_source[0].value["lines"] == ["different-content"]
    assert changed_source[0].metadata["duplicate_line_count"] == 0


def test_diagnostic_manager_file_budget_stops_browse_calls() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    host = NS(name="esx-log", _moId="host-1")
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-log")
    descriptors = [NS(key=key, fileName=f"/var/log/{key}.log", creator=key, info=NS(label=key, summary="")) for key in ("vmkernel", "hostd", "vpxa")]
    calls: list[str] = []

    class Manager:
        def QueryDescriptions(self, host=None):
            return descriptors

        def BrowseDiagnosticLog(self, host=None, key="", start=0, lines=0):
            calls.append((key, start))
            if start == 0:
                return NS(lineStart=1, lineEnd=1, lineText=[f"2026-09-24T00:00:00Z {key} event"])
            return NS(lineStart=2, lineEnd=1, lineText=[])

    records, attempts, fallback = collector._read_diagnostic_log_scope(
        Manager(),
        host,
        entity,
        ("vmkernel", "hostd", "vpxa"),
        "ds-log-budget",
        "2026-09-24T00:01:00Z",
        file_budget={"max_files": 1, "attempted": 0},
        byte_budget={"max_bytes": 1024, "used_bytes": 0},
        max_file_bytes=1024,
    )
    statuses = {record.metadata["log_category"]: record.metadata["log_status"] for record in records}
    assert calls == [("vmkernel", 0), ("vmkernel", 2)]
    assert statuses == {"vmkernel": "ok", "hostd": "file_limit_reached", "vpxa": "file_limit_reached"}
    assert fallback == []
    assert sum(item["status"] == "file_limit_reached" for item in attempts) == 2


def test_diagnostic_manager_api_reads_bounded_log_and_falls_back_missing_category() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    host = NS(name="esx-log", _moId="host-1")
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-log")
    descriptors = [NS(key="vmkernel", fileName="/var/log/vmkernel.log", creator="vmkernel", info=NS(label="VMkernel", summary="")), NS(key="hostd", fileName="/var/log/hostd.log", creator="hostd", info=NS(label="Host Agent", summary=""))]

    class Manager:
        def QueryDescriptions(self, host=None):
            return descriptors

        def BrowseDiagnosticLog(self, host=None, key="", start=0, lines=0):
            assert lines == 200
            if key != "vmkernel" or start != 0:
                return NS(lineStart=start + 1, lineEnd=start, lineText=[])
            return NS(lineStart=1, lineEnd=1, lineText=["2026-09-24T00:00:00Z error vmnic0 link down"])

    records, attempts, fallback = collector._read_diagnostic_log_scope(Manager(), host, entity, ("vmkernel", "hostd", "vpxa"), "ds", "2026-09-24T00:01:00Z")
    by_category = {record.metadata["log_category"]: record for record in records}
    assert by_category["vmkernel"].metadata["log_status"] == "ok"
    assert by_category["vmkernel"].window.start == "2026-09-24T00:00:00Z"
    assert by_category["hostd"].metadata["log_status"] == "empty"
    assert by_category["vpxa"].metadata["log_status"] == "no_logs"
    assert fallback == ["hostd", "vpxa"]
    assert attempts[0]["status"] == "ok"


def test_diagnostic_manager_window_reads_recent_tail_and_preserves_unknown_timestamp_lines() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    collector._collection_started_at = time.perf_counter()
    host = NS(name="esx-window", _moId="host-window")
    entity = DeepEntity(type="HostSystem", stable_id="host-window", display_ref="esx-window")
    descriptor = NS(key="hostd", fileName="/var/log/hostd.log", creator="hostd", info=NS(label="Host Agent", summary=""))
    cutoff = datetime(2026, 9, 20, tzinfo=UTC)
    collected_at = "2026-09-27T00:00:00Z"
    all_lines = [
        f"{(datetime(2026, 9, 19, 23, 0, tzinfo=UTC) + timedelta(seconds=index)).isoformat().replace('+00:00', 'Z')} hostd old-{index}"
        for index in range(2850)
    ] + [
        f"{(datetime(2026, 9, 26, 0, 0, tzinfo=UTC) + timedelta(seconds=index)).isoformat().replace('+00:00', 'Z')} hostd recent-{index}"
        for index in range(150)
    ]
    all_lines[2874] = "hostd timestamp unavailable sample"
    starts: list[int] = []

    class Manager:
        def QueryDescriptions(self, host=None):
            return [descriptor]

        def BrowseDiagnosticLog(self, host=None, key="", start=0, lines=0):
            starts.append(start)
            if start == 2_147_483_647:
                return NS(lineStart=len(all_lines) + 1, lineEnd=len(all_lines), lineText=[])
            offset = max(0, start - 1)
            page_lines = all_lines[offset : offset + lines]
            return NS(lineStart=offset + 1, lineEnd=offset + len(page_lines), lineText=page_lines)

    records, attempts, fallback = collector._read_diagnostic_log_scope(
        Manager(), host, entity, ("hostd",), "ds-log-window", collected_at,
        file_budget={"max_files": 1, "attempted": 0},
        byte_budget={"max_bytes": 1024 * 1024, "used_bytes": 0, "probe_bytes": 0, "log_window_cutoff_utc": cutoff, "log_window_days": 7},
        max_file_bytes=1024 * 1024,
    )

    assert starts[0] == 2_147_483_647
    assert len(starts) <= 6
    assert all(later < earlier for earlier, later in zip(starts[1:], starts[2:]))
    assert len(records[0].value["lines"]) == 150
    assert records[0].metadata["line_start"] == 2851
    assert records[0].metadata["line_end"] == 3000
    assert records[0].metadata["log_status"] == "partial"
    assert records[0].metadata["time_window_status"] == "complete_with_unknown_time"
    assert records[0].metadata["time_unknown_line_count"] == 1
    assert 0 < records[0].metadata["time_window_excluded_line_count"] < 2850
    assert attempts[-1]["time_window_complete"] is True
    assert attempts[-1]["total_line_count"] == 3000
    assert attempts[-1]["scanned_line_count"] < 2850
    assert fallback == []


def test_diagnostic_manager_window_marks_timestamp_order_violation_as_unverified() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    collector._collection_started_at = time.perf_counter()
    host = NS(name="esx-order", _moId="host-order")
    entity = DeepEntity(type="HostSystem", stable_id="host-order", display_ref="esx-order")
    descriptor = NS(key="hostd", fileName="/var/log/hostd.log", creator="hostd", info=NS(label="Host Agent", summary=""))
    all_lines = ["2026-09-26T00:00:10Z hostd newer", "2026-09-26T00:00:00Z hostd older"]

    class Manager:
        def QueryDescriptions(self, host=None):
            return [descriptor]

        def BrowseDiagnosticLog(self, host=None, key="", start=0, lines=0):
            if start == 2_147_483_647:
                return NS(lineStart=3, lineEnd=2, lineText=[])
            offset = max(0, start - 1)
            page_lines = all_lines[offset : offset + lines]
            return NS(lineStart=offset + 1, lineEnd=offset + len(page_lines), lineText=page_lines)

    records, attempts, fallback = collector._read_diagnostic_log_scope(
        Manager(), host, entity, ("hostd",), "ds-log-order", "2026-09-27T00:00:00Z",
        file_budget={"max_files": 1, "attempted": 0},
        byte_budget={"max_bytes": 1024 * 1024, "used_bytes": 0, "probe_bytes": 0, "log_window_cutoff_utc": datetime(2026, 9, 20, tzinfo=UTC), "log_window_days": 7},
        max_file_bytes=1024 * 1024,
    )

    assert records[0].metadata["time_window_status"] == "order_unreliable"
    assert records[0].metadata["truncated"] is True
    assert attempts[-1]["time_window_complete"] is False
    assert fallback == ["hostd"]


def test_diagnostic_manager_reads_all_log_pages_and_retains_source_line_ranges() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    host = NS(name="esx-log", _moId="host-1")
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-log")
    descriptor = NS(key="hostd", fileName="/var/log/hostd.log", creator="hostd", info=NS(label="Host Agent", summary=""))
    base = datetime(2026, 9, 26, 0, 0, tzinfo=UTC)
    all_lines = [
        f"{(base + timedelta(seconds=index)).isoformat().replace('+00:00', 'Z')} hostd event-{index}"
        for index in range(450)
    ]
    starts: list[int] = []

    class Manager:
        def QueryDescriptions(self, host=None):
            return [descriptor]

        def BrowseDiagnosticLog(self, host=None, key="", start=0, lines=0):
            starts.append(start)
            assert lines == 200
            offset = max(0, start - 1)
            page_lines = all_lines[offset : offset + lines]
            return NS(
                lineStart=offset + 1,
                lineEnd=offset + len(page_lines),
                lineText=page_lines,
            )

    records, attempts, fallback = collector._read_diagnostic_log_scope(
        Manager(),
        host,
        entity,
        ("hostd",),
        "ds-log-paged",
        "2026-09-26T01:00:00Z",
        file_budget={"max_files": 1, "attempted": 0},
        byte_budget={"max_bytes": 1024 * 1024, "used_bytes": 0},
        max_file_bytes=1024 * 1024,
    )

    record = records[0]
    assert starts == [0, 201, 401, 451]
    assert record.metadata["log_status"] == "ok"
    assert record.value["line_count"] == 450
    assert record.value["truncated"] is False
    assert record.window.start == "2026-09-26T00:00:00Z"
    assert record.window.end == "2026-09-26T00:07:29Z"
    assert record.metadata["line_start"] == 1
    assert record.metadata["line_end"] == 450
    assert record.metadata["source_line_ranges"] == [
        {"record_line_start": 1, "record_line_end": 450, "source_line_start": 1, "source_line_end": 450}
    ]
    browse_attempt = next(item for item in attempts if item["source"].endswith("BrowseDiagnosticLog"))
    assert browse_attempt["page_count"] == 4
    assert browse_attempt["pagination_complete"] is True
    assert fallback == []


def test_diagnostic_manager_pagination_stops_at_existing_file_byte_limit() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    host = NS(name="esx-log", _moId="host-1")
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-log")
    descriptor = NS(key="hostd", fileName="/var/log/hostd.log", creator="hostd", info=NS(label="Host Agent", summary=""))
    starts: list[int] = []
    long_line = "2026-09-26T00:00:00Z " + ("x" * 2000)

    class Manager:
        def QueryDescriptions(self, host=None):
            return [descriptor]

        def BrowseDiagnosticLog(self, host=None, key="", start=0, lines=0):
            starts.append(start)
            return NS(lineStart=1, lineEnd=200, lineText=[long_line] * 200)

    budget = {"max_bytes": 2048, "used_bytes": 0}
    records, attempts, _fallback = collector._read_diagnostic_log_scope(
        Manager(),
        host,
        entity,
        ("hostd",),
        "ds-log-page-byte-budget",
        "2026-09-26T01:00:00Z",
        file_budget={"max_files": 1, "attempted": 0},
        byte_budget=budget,
        max_file_bytes=1024,
    )

    assert starts == [0]
    assert len(records[0].value["lines"]) == 1
    assert records[0].metadata["retained_bytes"] == 1024
    assert records[0].metadata["truncated"] is True
    assert records[0].metadata["truncation_reasons"] == ["file_size_limit"]
    assert budget["used_bytes"] == 1024
    assert next(item for item in attempts if item["source"].endswith("BrowseDiagnosticLog"))["page_count"] == 1


def test_diagnostic_manager_permission_failure_is_classified() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-log")

    class NoPermission(Exception):
        privilegeId = "Global.Diagnostics"

    class Manager:
        def QueryDescriptions(self, host=None):
            raise NoPermission("no privilege")

    records, attempts, fallback = collector._read_diagnostic_log_scope(Manager(), NS(name="esx-log"), entity, ("vmkernel", "hostd"), "ds", "2026-09-24T00:01:00Z")
    assert {record.metadata["log_status"] for record in records} == {"permission_denied"}
    assert {record.metadata["missing_privilege_id"] for record in records} == {"Global.Diagnostics"}
    assert fallback == []
    assert attempts[0]["status"] == "permission_denied"
    assert attempts[0]["missing_privilege_id"] == "Global.Diagnostics"
    powercli_attempt = next(item for item in attempts if item["source"] == "PowerCLI.Get-Log")
    assert powercli_attempt["status"] == "not_attempted"
    assert powercli_attempt["missing_privilege_id"] == "Global.Diagnostics"
    capability = collector._log_capability(records)
    assert capability.detail["missing_privilege_ids"] == ["Global.Diagnostics"]


def test_diagnostic_privilege_preflight_checks_vcenter_and_each_host() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    root = NS(_moId="group-d1")
    hosts = [NS(_moId="host-1"), NS(_moId="host-2")]

    class AuthorizationManager:
        def HasUserPrivilegeOnEntities(self, entities, userName, privId):
            assert entities == [root, *hosts]
            assert userName == "readonly-user"
            assert privId == ["Global.Diagnostics"]
            return [NS(privAvailability=[NS(privId=privId[0], isGranted=value)]) for value in (True, False, False)]

    attempts: list[dict[str, object]] = []
    content = NS(
        rootFolder=root,
        authorizationManager=AuthorizationManager(),
        sessionManager=NS(currentSession=NS(userName="readonly-user")),
    )
    privileges = collector._diagnostic_privilege_state(content, hosts, attempts)

    assert privileges == {"group-d1": True, "host-1": False, "host-2": False}
    assert attempts[0]["status"] == "ok"
    assert attempts[0]["denied_count"] == 2


def test_partial_log_capability_does_not_globally_degrade_unrelated_complete_rules() -> None:
    from vstacklens.deep.pyvmomi_collector import _session_impact_degraded

    assert _session_impact_degraded(
        [{"action": "logs.history", "status": "partial"}],
        {"degraded": False},
        {"budget_exceeded": False},
    ) is False
    assert _session_impact_degraded(
        [{"action": "vm.hidden_risk", "status": "not_requested", "reason": "memory_limit_exceeded"}],
        {"degraded": False},
        {"budget_exceeded": False},
    ) is True
    assert _session_impact_degraded([], {"degraded": True}, {"budget_exceeded": False}) is True
    assert _session_impact_degraded([], {"degraded": False}, {"budget_exceeded": True}) is True


def test_diagnostic_privilege_preflight_skips_api_and_powercli_reads() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-log")

    class Manager:
        def QueryDescriptions(self, host=None):
            raise AssertionError("preflight-denied log API must not be called")

    records, attempts, fallback = collector._read_diagnostic_log_scope(
        Manager(),
        NS(name="esx-log"),
        entity,
        ("vmkernel", "hostd"),
        "ds-preflight",
        "2026-09-24T00:01:00Z",
        preflight_missing_privilege_id="Global.Diagnostics",
    )

    assert {record.metadata["log_status"] for record in records} == {"permission_denied"}
    assert {record.metadata["missing_privilege_id"] for record in records} == {"Global.Diagnostics"}
    assert [item["status"] for item in attempts if item["source"] == "vim.DiagnosticManager.QueryDescriptions"] == ["not_attempted"]
    assert [item["status"] for item in attempts if item["source"] == "PowerCLI.Get-Log"] == ["not_attempted"]
    assert fallback == []


def test_direct_esxi_log_fallback_succeeds_after_vcenter_diagnostics_denial(monkeypatch: pytest.MonkeyPatch) -> None:
    from pyVim import connect
    from vstacklens.collection import powercli_backend

    calls = {"connect": 0, "disconnect": 0, "query": 0, "browse": 0}
    descriptor = NS(key="vmkernel", fileName="/var/run/log/vmkernel.log", creator="vmkernel", mimeType="text/plain", info=NS(label="VMkernel", summary=""))
    fdm_descriptor = NS(key="fdm", fileName="/var/run/log/fdm.log", creator="Fdm", mimeType="text/plain", info=NS(label="Fault Domain Manager", summary=""))

    class VCenterManager:
        def QueryDescriptions(self, host=None):
            raise AssertionError("the preflight-denied vCenter DiagnosticManager must not run")

    class DirectManager:
        def QueryDescriptions(self):
            calls["query"] += 1
            return [descriptor, fdm_descriptor]

        def BrowseDiagnosticLog(self, key, start, lines):
            calls["browse"] += 1
            assert key in {"vmkernel", "fdm"}
            assert lines <= 200
            if start != 0:
                return NS(lineStart=start + 1, lineEnd=start, lineText=[])
            text = "2026-09-26T00:00:00Z vmkernel: APD path failure" if key == "vmkernel" else "2026-09-26T00:00:01Z Fdm failover failed on isolated host"
            return NS(lineStart=1, lineEnd=1, lineText=[text])

    class DirectServiceInstance:
        def RetrieveContent(self):
            return NS(diagnosticManager=DirectManager())

    class FakeAuthorizationManager:
        def HasUserPrivilegeOnEntities(self, entities, userName, privId):
            assert userName == "readonly"
            assert privId == ["Global.Diagnostics"]
            return [NS(privAvailability=[NS(privId="Global.Diagnostics", isGranted=False)]) for _entity in entities]

    class UnexpectedPowerCliBackend:
        def __init__(self, *args, **kwargs):
            raise AssertionError("PowerCLI must remain skipped when vCenter reports the same Global.Diagnostics denial")

    def smart_connect(**kwargs):
        calls["connect"] += 1
        assert kwargs["host"] == "192.0.2.11"
        assert kwargs["user"] == "root-readonly"
        assert kwargs["pwd"] == "host-secret"
        return DirectServiceInstance()

    monkeypatch.setattr(connect, "SmartConnect", smart_connect)
    monkeypatch.setattr(connect, "Disconnect", lambda _si: calls.__setitem__("disconnect", calls["disconnect"] + 1))
    monkeypatch.setattr(powercli_backend, "PowerCliBackend", UnexpectedPowerCliBackend)

    host = NS(
        name="esx-01",
        _moId="host-1",
        config=NS(uuid="host-uuid-1", network=NS(vnic=[NS(spec=NS(ip=NS(ipAddress="192.0.2.11")))])),
    )
    host_two = NS(
        name="esx-02",
        _moId="host-2",
        config=NS(uuid="host-uuid-2", network=NS(vnic=[NS(spec=NS(ip=NS(ipAddress="192.0.2.12")))])),
    )
    content = NS(
        diagnosticManager=VCenterManager(),
        rootFolder=NS(_moId="group-d1"),
        authorizationManager=FakeAuthorizationManager(),
        sessionManager=NS(currentSession=NS(userName="readonly")),
    )
    credentials = [
        EsxiHostConnectionInfo(host_alias="esxi-1", host="192.0.2.11", username="root-readonly", password="host-secret"),
        EsxiHostConnectionInfo(host_alias="esxi-2", host="192.0.2.12", username="root-readonly-2", password="host-secret-2"),
    ]
    collector = PyVmomiDeepCollector("vc.example", "readonly", "vcenter-password", host_log_credentials=credentials)
    collector._collection_started_at = time.perf_counter()

    records, collection = collector._collect_historical_logs(content, [host, host_two], "ds-direct-esxi-log", "2026-09-26T00:01:00Z", DeepCollectionPolicy(max_log_files=2), [])

    direct_log = next(record for record in records if record.metadata.get("log_category") == "vmkernel" and record.metadata.get("log_status") == "ok")
    assert direct_log.entity.stable_id == "host-uuid-1"
    assert direct_log.source.api == "vim.DiagnosticManager.DirectESXi.BrowseDiagnosticLog"
    assert direct_log.value["line_count"] == 1
    ha_log = next(record for record in records if record.metadata.get("log_category") == "ha" and record.metadata.get("log_status") == "ok")
    assert ha_log.source.api == "vim.DiagnosticManager.DirectESXi.BrowseDiagnosticLog"
    assert ha_log.value["marker_counts"]["ha"] == 1
    assert collection["status"] == "partial"
    assert collection["file_read_attempt_count"] == 2
    assert collection["missing_privilege_ids"] == ["Global.Diagnostics"]
    assert any(item.get("source") == "vim.DiagnosticManager.DirectESXi" and item.get("status") == "partial" for item in collection["attempts"])
    assert any(item.get("source") == "vim.DiagnosticManager.DirectESXi" and item.get("scope") == "direct_esxi:host-uuid-2" and item.get("status") == "file_limit_reached" and item.get("connection_status") == "not_attempted" for item in collection["attempts"])
    assert sum(1 for record in records if record.entity.stable_id == "host-uuid-2" and record.metadata.get("log_status") == "file_limit_reached") == 7
    assert calls == {"connect": 1, "disconnect": 1, "query": 1, "browse": 4}
    descriptor_attempt = next(item for item in collection["attempts"] if item.get("source") == "vim.DiagnosticManager.DirectESXi.QueryDescriptions")
    assert descriptor_attempt["descriptor_category_counts"] == {"vmkernel": 1, "ha": 1}
    assert descriptor_attempt["unclassified_descriptor_count"] == 0
    serialized_records = " ".join(record.model_dump_json() for record in records)
    assert "host-secret" not in serialized_records
    assert "host-secret-2" not in serialized_records
    assert "root-readonly" not in serialized_records
    assert "vcenter-password" not in serialized_records


def test_diagnostic_manager_permission_without_privilege_id_still_uses_powercli_fallback() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-log")

    class NoPermission(Exception):
        pass

    class Manager:
        def QueryDescriptions(self, host=None):
            raise NoPermission("no privilege detail returned")

    records, attempts, fallback = collector._read_diagnostic_log_scope(
        Manager(), NS(name="esx-log"), entity, ("vmkernel", "hostd"), "ds", "2026-09-24T00:01:00Z"
    )
    assert {record.metadata["log_status"] for record in records} == {"permission_denied"}
    assert fallback == ["vmkernel", "hostd"]
    assert not any(item["source"] == "PowerCLI.Get-Log" for item in attempts)


def test_diagnostic_manager_browse_permission_preserves_missing_privilege() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-log")
    descriptor = NS(key="vmkernel", fileName="/var/log/vmkernel.log", creator="vmkernel", info=NS(label="VMkernel", summary=""))

    class NoPermission(Exception):
        privilegeId = "Global.Diagnostics"

    class Manager:
        def QueryDescriptions(self, host=None):
            return [descriptor]

        def BrowseDiagnosticLog(self, host=None, key="", start=0, lines=0):
            raise NoPermission("no privilege")

    records, attempts, fallback = collector._read_diagnostic_log_scope(
        Manager(), NS(name="esx-log"), entity, ("vmkernel",), "ds", "2026-09-24T00:01:00Z"
    )
    browse_attempt = next(item for item in attempts if item["source"] == "vim.DiagnosticManager.BrowseDiagnosticLog")
    assert records[0].metadata["log_status"] == "permission_denied"
    assert records[0].metadata["missing_privilege_id"] == "Global.Diagnostics"
    assert browse_attempt["missing_privilege_id"] == "Global.Diagnostics"
    assert fallback == []
    assert any(item["source"] == "PowerCLI.Get-Log" and item["status"] == "not_attempted" for item in attempts)


def test_diagnostic_manager_resource_guard_stops_before_next_browse() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    host = NS(name="esx-log", _moId="host-1")
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-log")
    descriptors = [NS(key=key, fileName=f"/var/log/{key}.log", creator=key, info=NS(label=key, summary="")) for key in ("vmkernel", "hostd", "vpxa")]
    calls = []
    guard_reasons = iter((None, None, None, "memory_limit_exceeded", "memory_limit_exceeded", "memory_limit_exceeded"))

    class Manager:
        def QueryDescriptions(self, host=None):
            return descriptors

        def BrowseDiagnosticLog(self, host=None, key="", start=0, lines=0):
            calls.append(key)
            if start != 0:
                return NS(lineStart=start + 1, lineEnd=start, lineText=[])
            return NS(lineStart=1, lineEnd=1, lineText=[f"2026-09-24T00:00:00Z {key} sample"])

    records, attempts, fallback = collector._read_diagnostic_log_scope(
        Manager(), host, entity, ("vmkernel", "hostd", "vpxa"), "ds-log-resource", "2026-09-24T00:01:00Z",
        file_budget={"max_files": 8, "attempted": 0}, byte_budget={"max_bytes": 1024, "used_bytes": 0},
        resource_guard=lambda: next(guard_reasons),
    )

    statuses = {record.metadata["log_category"]: record.metadata["log_status"] for record in records}
    assert calls == ["vmkernel"]
    assert statuses == {"vmkernel": "partial", "hostd": "memory_limit_exceeded", "vpxa": "memory_limit_exceeded"}
    assert fallback == []
    assert sum(item["source"] == "resource_guard" for item in attempts) == 2


def test_historical_log_collection_falls_back_from_api_permission(monkeypatch: pytest.MonkeyPatch) -> None:
    from vstacklens.collection import powercli_backend

    class NoPermission(Exception):
        pass

    class Manager:
        def QueryDescriptions(self, host=None):
            raise NoPermission("denied")

    class Cached:
        def __init__(self, keys):
            self.keys = keys

        def get_logs(self):
            return [{"key": key, "status": "ok" if key == "vmkernel" else "permission_denied", "lines": ["2026-09-24T00:00:00Z error vmkernel APD"] if key == "vmkernel" else [], "source": "PowerCLI.Get-Log"} for key in self.keys]

    class FakeBackend:
        def __init__(self, *args, **kwargs):
            self.warnings = []
            self.vcenter_logs = [{"key": "vpxd", "status": "permission_denied", "lines": [], "source": "PowerCLI.Get-Log"}]

        def prepare(self, hosts, **kwargs):
            self.keys = kwargs["log_keys"]

        def __call__(self, host):
            return Cached(self.keys)

        def get_vcenter_logs(self):
            return self.vcenter_logs

    monkeypatch.setattr(powercli_backend, "PowerCliBackend", FakeBackend)
    collector = PyVmomiDeepCollector("vc", "u", "p")
    collector._collection_started_at = time.perf_counter()
    host = NS(name="esx-log", _moId="host-1")
    records, collection = collector._collect_historical_logs(NS(diagnosticManager=Manager()), [host], "ds-fallback", "2026-09-24T00:01:00Z", DeepCollectionPolicy(), [])
    assert collection["status"] == "partial"
    assert any(attempt["source"] == "PowerCLI.Get-Log" and attempt.get("fallback_from") == "vim.DiagnosticManager" for attempt in collection["attempts"])
    assert any(record.metadata["log_category"] == "vmkernel" and record.metadata["log_status"] == "ok" for record in records)
    assert any(record.metadata["log_category"] == "vpxd" for record in records)


def test_historical_log_fallback_records_powercli_not_attempted_when_byte_budget_exhausted(monkeypatch: pytest.MonkeyPatch) -> None:
    from vstacklens.collection import powercli_backend

    class NoDescriptors:
        def QueryDescriptions(self, host=None):
            return []

    class UnexpectedBackend:
        def __init__(self, *args, **kwargs):
            raise AssertionError("PowerCLI must not start after the log byte budget is exhausted")

    monkeypatch.setattr(powercli_backend, "PowerCliBackend", UnexpectedBackend)
    collector = PyVmomiDeepCollector("vc", "user", "synthetic-secret")
    collector._collection_started_at = time.perf_counter()

    records, collection = collector._collect_historical_logs(
        NS(diagnosticManager=NoDescriptors()),
        [],
        "ds-fallback-budget-not-attempted",
        "2026-09-27T00:00:00Z",
        DeepCollectionPolicy(max_log_total_bytes=0),
        [],
    )

    powercli_attempt = next(item for item in collection["attempts"] if item.get("source") == "PowerCLI.Get-Log")
    assert powercli_attempt["status"] == "not_attempted"
    assert powercli_attempt["reason"] == "total_size_limit_reached"
    assert powercli_attempt["category_count"] == 1
    assert any(record.metadata.get("log_status") == "total_size_limit_reached" for record in records)


def test_recent_log_window_does_not_call_powercli_without_a_time_seek(monkeypatch: pytest.MonkeyPatch) -> None:
    from vstacklens.collection import powercli_backend

    class NoDescriptors:
        def QueryDescriptions(self, host=None):
            return []

    class UnexpectedPowerCLI:
        def __init__(self, *_args, **_kwargs):
            raise AssertionError("date-limited collection must not read unbounded PowerCLI log prefixes")

    monkeypatch.setattr(powercli_backend, "PowerCliBackend", UnexpectedPowerCLI)
    collector = PyVmomiDeepCollector("vc", "user", "password")
    collector._collection_started_at = time.perf_counter()
    records, collection = collector._collect_historical_logs(
        NS(diagnosticManager=NoDescriptors()),
        [],
        "ds-log-window-powercli-skip",
        "2026-09-27T00:00:00Z",
        DeepCollectionPolicy(log_collection_days=7),
        [],
    )

    powercli = next(item for item in collection["attempts"] if item.get("source") == "PowerCLI.Get-Log")
    assert powercli["status"] == "not_attempted"
    assert powercli["reason"] == "time_window_seek_not_supported_by_backend"
    assert next(item for item in collection["attempts"] if item.get("source") == "vim.DiagnosticManager.DirectESXi")["status"] == "not_attempted"
    assert any(record.metadata.get("log_category") == "vpxd" and record.metadata["log_status"] == "time_window_seek_not_supported" for record in records)
    assert collection["requested_log_days"] == 7


def test_historical_log_guard_prevents_starting_powercli_fallback(monkeypatch: pytest.MonkeyPatch) -> None:
    from vstacklens.collection import powercli_backend

    class NoPermission(Exception):
        pass

    class Manager:
        def QueryDescriptions(self, host=None):
            raise NoPermission("denied")

    class UnexpectedBackend:
        def __init__(self, *args, **kwargs):
            raise AssertionError("PowerCLI fallback must not start after resource guard")

    monkeypatch.setattr(powercli_backend, "PowerCliBackend", UnexpectedBackend)
    stop_reasons = iter((None, None, "memory_limit_exceeded"))
    collector = PyVmomiDeepCollector("vc", "u", "p")
    collector._collection_started_at = time.perf_counter()
    host = NS(name="esx-log", _moId="host-1")
    records, collection = collector._collect_historical_logs(
        NS(diagnosticManager=Manager()),
        [host],
        "ds-log-resource-stop",
        "2026-09-24T00:01:00Z",
        DeepCollectionPolicy(),
        [],
        resource_guard=lambda: next(stop_reasons),
    )

    stopped = [record for record in records if record.metadata["log_status"] == "memory_limit_exceeded"]
    assert len(stopped) == 8
    assert all(record.source.api == "DeepCollectionPolicy.resource_guard" for record in stopped)
    assert collection["file_read_attempt_count"] == 0
    assert collection["successful_count"] == 0
    assert any(item["source"] == "resource_guard" and item["status"] == "memory_limit_exceeded" for item in collection["attempts"])


def test_historical_log_fallback_obeys_global_file_cap_and_host_specific_keys(monkeypatch: pytest.MonkeyPatch) -> None:
    from vstacklens.collection import powercli_backend

    class NoPermission(Exception):
        pass

    class Manager:
        def QueryDescriptions(self, host=None):
            raise NoPermission("denied")

    class Cached:
        def __init__(self, keys):
            self.keys = keys

        def get_logs(self):
            return [{"key": key, "status": "ok", "lines": [f"2026-09-24T00:00:00Z {key} sample"], "source": "PowerCLI.Get-Log"} for key in self.keys]

    backend_instances = []

    class FakeBackend:

        def __init__(self, *args, **kwargs):
            backend_instances.append(self)
            self.warnings = []
            self.vcenter_logs = [{"key": "vpxd", "status": "ok", "lines": ["2026-09-24T00:00:00Z vpxd sample"], "source": "PowerCLI.Get-Log"}]
            self.log_file_attempt_count = 0

        def prepare(self, hosts, **kwargs):
            self.prepared = (hosts, kwargs)
            self.log_file_attempt_count = kwargs["max_log_files"]

        def __call__(self, host):
            return Cached(self.prepared[1]["log_keys_by_host"][str(host._moId)])

        def get_vcenter_logs(self):
            return self.vcenter_logs

    monkeypatch.setattr(powercli_backend, "PowerCliBackend", FakeBackend)
    collector = PyVmomiDeepCollector("vc", "u", "p")
    collector._collection_started_at = time.perf_counter()
    hosts = [NS(name="esx-2", _moId="host-2"), NS(name="esx-1", _moId="host-1")]
    records, collection = collector._collect_historical_logs(
        NS(diagnosticManager=Manager()),
        hosts,
        "ds-log-cap",
        "2026-09-24T00:01:00Z",
        DeepCollectionPolicy(max_log_files=2),
        [],
    )
    prepared_hosts, kwargs = backend_instances[0].prepared
    assert [host.name for host in prepared_hosts] == ["esx-1"]
    assert kwargs["log_keys_by_host"] == {"host-1": ["vmkernel"]}
    assert kwargs["vcenter_log_keys"] == ["vpxd"]
    assert collection["file_read_attempt_count"] == 2
    assert collection["budget_limited_count"] == 13
    assert sum(record.metadata["log_status"] == "file_limit_reached" for record in records) == 13


def test_historical_log_fallback_isolates_vcenter_and_per_host_failures(monkeypatch: pytest.MonkeyPatch) -> None:
    from vstacklens.collection import powercli_backend

    class NoPermission(Exception):
        pass

    class Manager:
        def QueryDescriptions(self, host=None):
            raise NoPermission("denied")

    class Cached:
        def __init__(self, host_id: str, keys: list[str]):
            self.host_id = host_id
            self.keys = keys

        def get_logs(self):
            if self.host_id == "host-fail":
                raise RuntimeError("host scoped log result failed")
            return [{"key": key, "status": "ok", "lines": [f"2026-09-24T00:00:00Z {key} sample"], "source": "PowerCLI.Get-Log"} for key in self.keys]

    class FakeBackend:
        def __init__(self, *args, **kwargs):
            self.warnings = []
            self.log_file_attempt_count = 0

        def prepare(self, hosts, **kwargs):
            self.keys_by_host = kwargs["log_keys_by_host"]
            self.log_file_attempt_count = kwargs["max_log_files"]

        def get_vcenter_logs(self):
            raise RuntimeError("vCenter scoped log result failed")

        def __call__(self, host):
            host_id = str(host._moId)
            return Cached(host_id, self.keys_by_host[host_id])

    monkeypatch.setattr(powercli_backend, "PowerCliBackend", FakeBackend)
    collector = PyVmomiDeepCollector("vc", "u", "p")
    collector._collection_started_at = time.perf_counter()
    hosts = [NS(name="esx-fail", _moId="host-fail"), NS(name="esx-ok", _moId="host-ok")]
    records, collection = collector._collect_historical_logs(
        NS(diagnosticManager=Manager()),
        hosts,
        "ds-log-scope-isolation",
        "2026-09-24T00:01:00Z",
        DeepCollectionPolicy(),
        [],
    )

    fallback = [record for record in records if record.source.api == "PowerCLI.Get-Log"]
    statuses = {(record.entity.type, record.entity.stable_id, record.metadata["log_category"]): record.metadata["log_status"] for record in fallback}
    attempt = next(item for item in collection["attempts"] if item["source"] == "PowerCLI.Get-Log")
    assert statuses[("vCenter", collector._environment_entity().stable_id, "vpxd")] == "error"
    assert statuses[("HostSystem", collector._entity(hosts[0], "HostSystem").stable_id, "vmkernel")] == "error"
    assert statuses[("HostSystem", collector._entity(hosts[1], "HostSystem").stable_id, "vmkernel")] == "ok"
    assert attempt["status"] == "partial"
    assert attempt["scope_failures"] == [
        {"scope": "vcenter", "status": "error"},
        {"scope": "host", "entity_id": "host-fail", "status": "error"},
    ]
    assert collection["status"] == "partial"


def test_log_records_retain_time_correlated_event_evidence() -> None:
    entity = DeepEntity(type="HostSystem", stable_id="host-log", display_ref="esx-log")
    records = normalize_powercli_logs(dataset_id="ds-logs", entity=entity, collected_at="2026-09-24T00:00:00Z", payloads=[{"key": "vmkernel", "status": "ok", "lines": ["2026-09-24T00:00:30Z error vmnic0 link down"]}])
    event = NS(createdTime=datetime(2026, 9, 24, 0, 1, tzinfo=UTC), fullFormattedMessage="Physical NIC vmnic0 linkstate is down", entity=NS(name="esx-log"), key=487709, dataset_pointer="events/event.ndjson#NET-DEEP-001/host-log")
    correlate_log_records(records, [event])
    assert records[0].metadata["related_event_count"] == 1
    assert records[0].metadata["related_event_samples"][0]["event_type"] == "SimpleNamespace"
    assert records[0].metadata["related_event_samples"][0]["event_key"] == 487709
    assert records[0].metadata["related_event_samples"][0]["dataset_pointer"] == "events/event.ndjson#NET-DEEP-001/host-log"
    assert records[0].metadata["related_event_samples"][0]["log_timestamp"] == "2026-09-24T00:00:30Z"
    assert records[0].metadata["related_event_samples"][0]["log_line_index"] == 1
    records[0].metadata["source_line_numbers"] = [3621]
    correlate_log_records(records, [event])
    assert records[0].metadata["related_event_samples"][0]["source_line_number"] == 3621
    assert records[0].metadata["related_event_samples"][0]["matching_categories"] == ["network"]
    other = NS(createdTime=datetime(2026, 9, 24, 0, 1, tzinfo=UTC), fullFormattedMessage="Physical NIC vmnic0 linkstate is down", entity=NS(name="other-host"))
    records = normalize_powercli_logs(dataset_id="ds-logs", entity=entity, collected_at="2026-09-24T00:00:00Z", payloads=[{"key": "vmkernel", "status": "ok", "lines": ["2026-09-24T00:00:30Z error vmnic0 link down"]}])
    correlate_log_records(records, [other])
    assert records[0].metadata["related_event_count"] == 0
    unrelated = NS(createdTime=datetime(2026, 9, 24, 0, 1, tzinfo=UTC), fullFormattedMessage="Virtual machine snapshot created", entity=NS(name="esx-log"))
    correlate_log_records(records, [unrelated])
    assert records[0].metadata["related_event_count"] == 0


def test_event_log_correlation_prefers_stable_object_ids_when_names_collide() -> None:
    collected = "2026-09-24T00:00:00Z"
    first = DeepEntity(type="HostSystem", stable_id="host-uuid-1", display_ref="shared-esx")
    second = DeepEntity(type="HostSystem", stable_id="host-uuid-2", display_ref="shared-esx")
    payload = [{"key": "vmkernel", "status": "ok", "lines": ["2026-09-24T00:00:30Z error link down"]}]
    logs = [
        *normalize_powercli_logs(dataset_id="ds-collision", entity=first, collected_at=collected, payloads=payload),
        *normalize_powercli_logs(dataset_id="ds-collision", entity=second, collected_at=collected, payloads=payload),
    ]
    event = NS(
        createdTime=datetime(2026, 9, 24, 0, 1, tzinfo=UTC),
        fullFormattedMessage="Physical NIC linkstate is down",
        host=NS(name="shared-esx", host=NS(_moId="host-moid-1")),
    )
    correlate_log_records(logs, [event], object_id_aliases={"host-moid-1": "host-uuid-1"})
    assert [record.metadata["related_event_count"] for record in logs] == [1, 0]
    assert logs[0].metadata["related_event_samples"][0]["object_ids"] == ["host-uuid-1"]

    same_name_wrong_id = NS(
        createdTime=datetime(2026, 9, 24, 0, 1, tzinfo=UTC),
        fullFormattedMessage="Physical NIC linkstate is down",
        host=NS(name="shared-esx", host=NS(_moId="host-moid-unknown")),
    )
    correlate_log_records(logs, [same_name_wrong_id])
    assert [record.metadata["related_event_count"] for record in logs] == [0, 0]

    ambiguous_name = NS(
        createdTime=datetime(2026, 9, 24, 0, 1, tzinfo=UTC),
        fullFormattedMessage="Physical NIC linkstate is down",
        entity=NS(name="shared-esx"),
    )
    correlate_log_records(logs, [ambiguous_name])
    assert [record.metadata["related_event_count"] for record in logs] == [0, 0]

    candidate = DatasetRecord(
        record_id="task-same-name-other-host",
        dataset_id="ds-collision",
        kind="task",
        entity=second,
        collected_at_utc=collected,
        source=DeepSource(api="TaskManager", collector="fixture", collected_at_utc=collected),
        selector={"rule_id": "fixture"},
        window=DeepWindow(start=collected, end=collected, sample_count=1, expected_sample_count=1, completeness=1.0),
        value={"state": "error"},
        unit="task",
        raw_pointer="fixture.ndjson#task",
    )
    correlate_log_records([logs[0]], [], [candidate])
    assert logs[0].metadata["related_record_count"] == 0


def test_saved_event_sample_dicts_can_replay_log_correlation_with_raw_pointer() -> None:
    entity = DeepEntity(type="HostSystem", stable_id="host-log", display_ref="esx-log")
    records = normalize_powercli_logs(dataset_id="ds-saved-events", entity=entity, collected_at="2026-09-24T00:00:00Z", payloads=[{"key": "vmkernel", "status": "ok", "lines": ["2026-09-24T00:00:30Z vmnic0 linkstate down"]}])
    sample = {
        "timestamp": "2026-09-24T00:01:00Z",
        "event_type": "vim.event.EventEx",
        "event_key": 487709,
        "object_ids": ["host-log"],
        "object_names": ["esx-log"],
        "message": "Physical NIC vmnic0 linkstate is down",
        "dataset_pointer": "events/event.ndjson#NET-DEEP-001/host-log",
    }

    correlate_log_records(records, [sample])

    related = records[0].metadata["related_event_samples"]
    assert records[0].metadata["related_event_count"] == 1
    assert related[0]["event_type"] == "vim.event.EventEx"
    assert related[0]["event_key"] == 487709
    assert related[0]["dataset_pointer"] == "events/event.ndjson#NET-DEEP-001/host-log"
    assert related[0]["log_timestamp"] == "2026-09-24T00:00:30Z"
    assert related[0]["matching_categories"] == ["network"]


def test_log_correlation_uses_event_rule_classification_when_message_is_generic() -> None:
    entity = DeepEntity(type="HostSystem", stable_id="host-rule-hint", display_ref="esx-rule-hint")
    records = normalize_powercli_logs(
        dataset_id="ds-rule-hint-correlation",
        entity=entity,
        collected_at="2026-09-24T00:00:00Z",
        payloads=[{"key": "vmkernel", "status": "ok", "lines": ["2026-09-24T00:00:30Z vmnic0 linkstate down"]}],
    )
    event = {
        "timestamp": "2026-09-24T00:01:00Z",
        "event_type": "vim.event.EventEx",
        "event_key": 487710,
        "object_ids": ["host-rule-hint"],
        "message": "Physical adapter transition was reported",
    }

    correlate_log_records(records, [event], event_rule_ids={id(event): {"NET-DEEP-002"}})

    related = records[0].metadata["related_event_samples"][0]
    assert related["matching_categories"] == ["network"]
    assert related["rule_ids"] == ["NET-DEEP-002"]
    assert related["interfaces"] == []
    assert related["state"] == ""
    assert related["evidence_level"] == "same_object_time_category"
    assert related.get("dataset_pointer") is None
    assert related["event_source_api"] == "EventManager.QueryEvents"


def test_log_correlation_does_not_invent_event_dataset_pointer() -> None:
    entity = DeepEntity(type="HostSystem", stable_id="host-pointerless", display_ref="esx-pointerless")
    logs = normalize_powercli_logs(
        dataset_id="ds-pointerless-correlation",
        entity=entity,
        collected_at="2026-09-24T00:00:00Z",
        payloads=[{"key": "vmkernel", "status": "ok", "lines": ["2026-09-24T00:00:30Z vmnic0 linkstate up"]}],
    )
    event = {
        "timestamp": "2026-09-24T00:01:00Z",
        "event_type": "vim.event.EventEx",
        "event_key": 594284,
        "object_ids": ["host-pointerless"],
        "message": "Physical NIC vmnic0 linkstate is up",
    }

    correlate_log_records(logs, [event])

    related = logs[0].metadata["related_event_samples"][0]
    assert related["event_source_api"] == "EventManager.QueryEvents"
    assert related.get("dataset_pointer") is None
    saved = evidence_ref(logs[0]).value_summary["related_event_samples"][0]
    assert saved["event_source_api"] == "EventManager.QueryEvents"
    assert saved["event_key"] == 594284
    line = _log_correlation_lines(NS(records=logs))[0]
    assert "Event Key 594284 的关联摘要保存在日志 Evidence 中" in line
    assert "EventManager.QueryEvents" in line
    assert "events/event.ndjson" not in line


def test_network_correlation_requires_matching_interface_and_link_state_when_explicit() -> None:
    entity = DeepEntity(type="HostSystem", stable_id="host-log", display_ref="esx-log")
    collected = "2026-08-28T10:10:00Z"
    logs = normalize_powercli_logs(
        dataset_id="ds-network-specific-correlation",
        entity=entity,
        collected_at=collected,
        payloads=[
            {
                "key": "vobd",
                "status": "ok",
                "lines": [
                    "2026-08-28T10:00:00Z vmnic4 linkstate down",
                    "2026-08-28T10:09:00Z vmnic4 linkstate up",
                    "2026-08-28T10:09:30Z vmnic5 linkstate down",
                ],
            }
        ],
    )
    down_event = {
        "timestamp": "2026-08-28T10:10:00Z",
        "event_type": "vim.event.EventEx",
        "event_key": 487709,
        "object_ids": ["host-log"],
        "message": "Physical NIC vmnic4 linkstate is down",
        "dataset_pointer": "events/event.ndjson#NET-DEEP-001/host-log",
    }

    correlate_log_records(logs, [down_event])

    related = logs[0].metadata["related_event_samples"]
    assert logs[0].metadata["related_event_count"] == 1
    assert related[0]["event_key"] == 487709
    assert related[0]["log_line_index"] == 1
    assert related[0]["log_timestamp"] == "2026-08-28T10:00:00Z"
    assert related[0]["time_distance_seconds"] == 600.0

    contradictory_logs = normalize_powercli_logs(
        dataset_id="ds-network-specific-correlation",
        entity=entity,
        collected_at=collected,
        payloads=[
            {
                "key": "vobd",
                "status": "ok",
                "lines": [
                    "2026-08-28T10:09:00Z vmnic4 linkstate up",
                    "2026-08-28T10:09:30Z vmnic5 linkstate down",
                ],
            }
        ],
    )

    correlate_log_records(contradictory_logs, [down_event])

    assert contradictory_logs[0].metadata["related_event_count"] == 0


def test_network_correlation_prefers_same_interface_and_state_and_reports_basis() -> None:
    entity = DeepEntity(type="HostSystem", stable_id="host-specific", display_ref="esx-specific")
    logs = normalize_powercli_logs(
        dataset_id="ds-network-correlation-priority",
        entity=entity,
        collected_at="2026-08-28T10:10:00Z",
        payloads=[
            {
                "key": "vobd",
                "status": "ok",
                "lines": [
                    "2026-08-28T10:09:59Z network diagnostic observed on vmnic4",
                    "2026-08-28T10:00:00Z vmnic4 linkstate down",
                ],
            }
        ],
    )
    event = {
        "timestamp": "2026-08-28T10:10:00Z",
        "event_type": "vim.event.EventEx",
        "event_key": 487709,
        "object_ids": ["host-specific"],
        "message": "Physical NIC vmnic4 linkstate is down",
        "dataset_pointer": "events/event.ndjson#NET-DEEP-001/host-specific",
    }
    weak_events = [
        {
            "timestamp": f"2026-08-28T10:09:5{second}Z",
            "event_type": "vim.event.EventEx",
            "event_key": 487800 + second,
            "object_ids": ["host-specific"],
            "rule_ids": ["NET-DEEP-002"],
            "message": "Physical adapter transition was reported",
            "dataset_pointer": f"events/event.ndjson#NET-DEEP-002/host-specific/{second}",
        }
        for second in range(3)
    ]

    correlate_log_records(logs, [event, *weak_events])

    related = logs[0].metadata["related_event_samples"][0]
    assert related["log_line_index"] == 2
    assert related["time_distance_seconds"] == 600.0
    assert related["interfaces"] == ["vmnic4"]
    assert related["state"] == "down"
    assert related["evidence_level"] == "same_object_time_category_interface_state"
    assert logs[0].metadata["related_event_evidence_counts"] == {
        "same_object_time_category_interface_state": 1,
        "same_object_time_category": 3,
    }

    diagnostic = evidence_ref(logs[0]).value_summary
    sample = diagnostic["related_event_samples"][0]
    assert sample["dataset_pointer"] == event["dataset_pointer"]
    assert sample["log_line_index"] == 2
    assert sample["interfaces"] == ["vmnic4"]
    assert sample["state"] == "down"

    report_lines = _log_correlation_lines(NS(records=logs))
    assert "Event Key 487709" in report_lines[0]
    assert "vmnic4 down" in report_lines[0]
    assert "同对象、同网卡、同方向" in report_lines[0]
    assert "共现不代表因果" not in report_lines[0]


def test_sanitized_real_support_bundle_event_correlation_replays_expected_links() -> None:
    import json

    fixture_path = Path(__file__).parent / "fixtures" / "deep_support_bundle_correlation_sanitized.json"
    fixture = json.loads(fixture_path.read_text(encoding="utf-8"))
    entity = DeepEntity(type="HostSystem", stable_id=fixture["stable_id"], display_ref="matched-r4-host")
    normalized_lines = [
        f"{item['timestamp']} {item['interface']} linkstate {item['state']}"
        for item in fixture["log_samples"]
    ]
    records = normalize_powercli_logs(
        dataset_id="offline-sanitized-support-bundle-correlation",
        entity=entity,
        collected_at="2026-08-28T10:29:59Z",
        payloads=[
            {
                "key": "vobd",
                "status": "ok",
                "source": "ESXiSupportBundle.vobd",
                "lines": normalized_lines,
            }
        ],
    )
    records[0].metadata["source_line_numbers"] = [item["source_line_number"] for item in fixture["log_samples"]]
    records[0].raw_pointer = "provided-support-bundle:var/run/log/vobd.0.gz"

    correlate_log_records(records, fixture["event_samples"], window_seconds=fixture["window_seconds"])

    assert records[0].metadata["related_event_count"] == 6
    actual = {
        item["event_key"]: {
            "source_line_number": item["source_line_number"],
            "time_distance_seconds": item["time_distance_seconds"],
            "matching_categories": item["matching_categories"],
        }
        for item in records[0].metadata["related_event_samples"]
    }
    expected = {
        item["event_key"]: {
            "source_line_number": item["source_line_number"],
            "time_distance_seconds": item["time_distance_seconds"],
            "matching_categories": ["network"],
        }
        for item in fixture["expected_links"]
    }
    assert actual == expected


def test_log_event_task_performance_chain_keeps_all_sources() -> None:
    entity = DeepEntity(type="HostSystem", stable_id="host-chain", display_ref="esx-chain")
    collected = "2026-09-24T00:00:00Z"
    logs = normalize_powercli_logs(dataset_id="ds-chain", entity=entity, collected_at=collected, payloads=[{"key": "vmkernel", "status": "ok", "lines": ["2026-09-24T00:00:30Z error vmnic0 link down"]}])
    event = NS(createdTime=datetime(2026, 9, 24, 0, 1, tzinfo=UTC), fullFormattedMessage="Physical NIC vmnic0 linkstate is down", entity=NS(name="esx-chain"))
    common = dict(dataset_id="ds-chain", entity=entity, collected_at_utc=collected, source=DeepSource(api="fixture", collector="fixture", collected_at_utc=collected), selector={"rule_id": "fixture"}, window=DeepWindow(start=collected, end=collected, sample_count=1, expected_sample_count=1, completeness=1.0), value=1, unit="count", raw_pointer="fixture.ndjson")
    task = DatasetRecord(record_id="task-chain", kind="task", summary="network adapter reconfiguration task failed", metadata={"rule_id": "TASK-DEEP-NET-001"}, **common)
    perf = DatasetRecord(record_id="perf-chain", kind="perf", summary="network drop sample", metadata={"rule_id": "ENHANCED-NET-001"}, **common)
    unrelated_task = DatasetRecord(record_id="task-unrelated", kind="task", summary="vMotion migration failed", metadata={"rule_id": "COMPUTE-DEEP-001"}, **common)
    correlate_log_records(logs, [event], [task, perf, unrelated_task])
    assert logs[0].metadata["related_event_count"] == 1
    assert {item["record_id"] for item in logs[0].metadata["related_record_samples"]} == {"task-chain", "perf-chain"}


def test_deleted_event_reference_name_failure_preserves_managed_object_id() -> None:
    from vstacklens.deep.logs import event_object_references

    class DeletedManagedObject:
        _moId = "vm-37045"

        def __init__(self):
            self.name_reads = 0

        @property
        def name(self):
            self.name_reads += 1
            raise RuntimeError("object was deleted during event processing")

    deleted = DeletedManagedObject()
    event = NS(
        createdTime=datetime(2026, 9, 25, 1, 0, tzinfo=UTC),
        fullFormattedMessage="VM reset after guest failure",
        vm=NS(name="deleted-vm", vm=deleted),
    )
    object_ids, object_names = event_object_references(event)
    assert object_ids == {"vm-37045"}
    assert "deleted-vm" in object_names
    assert deleted.name_reads == 0

    collector = PyVmomiDeepCollector("vc", "u", "p")
    records = collector._history_records([event], [], [], [], "ds-deleted-event-ref", "2026-09-25T01:01:00Z")
    reset = next(item for item in records if item.metadata["rule_id"] == "COMPUTE-DEEP-002")
    sample = reset.metadata["event_samples"][0]
    assert sample["object_ids"] == ["vm-37045"]
    assert any(item["stable_id"] == "moid:vm-37045" for item in sample["resolved_entities"])
    assert deleted.name_reads == 0


def test_log_record_correlation_uses_full_evidence_time_window() -> None:
    entity = DeepEntity(type="HostSystem", stable_id="host-window", display_ref="esx-window")
    logs = normalize_powercli_logs(
        dataset_id="ds-window-correlation",
        entity=entity,
        collected_at="2026-09-24T04:00:00Z",
        payloads=[{"key": "vmkernel", "status": "ok", "lines": ["2026-09-24T00:00:30 error link down"]}],
    )

    def record(record_id: str, start: str, end: str) -> DatasetRecord:
        return DatasetRecord(
            record_id=record_id,
            dataset_id="ds-window-correlation",
            kind="perf",
            entity=entity,
            collected_at_utc="2026-09-24T04:00:00Z",
            source=DeepSource(api="fixture.performance", collector="fixture", collected_at_utc="2026-09-24T04:00:00Z"),
            selector={"rule_id": "fixture"},
            window=DeepWindow(start=start, end=end, sample_count=2, expected_sample_count=2, completeness=1.0),
            value=1,
            unit="count",
            raw_pointer=f"fixture.ndjson#{record_id}",
            summary="network packet drop performance sample",
        )

    inside = record("perf-window-overlap", "2026-09-24T00:00:00Z", "2026-09-24T03:00:00Z")
    outside = record("perf-window-outside", "2026-09-24T02:00:00Z", "2026-09-24T03:00:00Z")
    correlate_log_records(logs, [], [inside, outside])

    related = logs[0].metadata["related_record_samples"]
    assert [item["record_id"] for item in related] == ["perf-window-overlap"]
    assert related[0]["time_distance_seconds"] == 0.0
    assert related[0]["window_start"] == inside.window.start
    assert related[0]["window_end"] == inside.window.end


def test_evidence_merge_deduplicates_and_preserves_conflicts() -> None:
    dataset = DeepDataset.from_json(FIXTURE)
    record = dataset.records[0]
    duplicate = record.model_copy(deep=True)
    conflict = record.model_copy(update={"value": "different", "source": record.source.model_copy(update={"api": "alternate.readonly.source"})})
    merged, log = merge_records([record, duplicate, conflict])
    assert len(merged) == 2
    assert log["duplicate_count"] == 1
    assert log["conflict_count"] == 1
    assert all(item.metadata.get("conflict") is True for item in merged)
    duplicate_evidence = next(item for item in merged if item.metadata.get("duplicate_count"))
    source_ref = evidence_ref(duplicate_evidence)
    assert len(source_ref.value_summary["duplicate_evidence"]) == 2
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records = merged
    rule_id = str(record.metadata["rule_id"])
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == rule_id)
    analysis = DeepAnalyzer().analyze(dataset, rules=(rule,))
    assert analysis.rule_results[0].status == DeepResultStatus.INSUFFICIENT_DATA
    assert "conflicting evidence" in analysis.rule_results[0].reason


def test_resource_collection_log_reports_budget_and_deltas() -> None:
    policy = DeepCollectionPolicy(budget_seconds=10)
    before = {"status": "ok", "cpu_user_sec": 1.0, "cpu_system_sec": 0.5, "rss_bytes": 100, "read_bytes": 10, "write_bytes": 20, "network_sent_bytes": 30, "network_recv_bytes": 40}
    after = {"status": "ok", "cpu_user_sec": 2.0, "cpu_system_sec": 1.0, "rss_bytes": 140, "read_bytes": 15, "write_bytes": 25, "network_sent_bytes": 45, "network_recv_bytes": 60}
    result = PyVmomiDeepCollector._resource_collection_log(before, after, 2.0, policy)
    assert result["status"] == "ok"
    assert result["cpu_sec"] == 1.5
    assert result["rss_delta_bytes"] == 40
    assert result["budget_exceeded"] is False


@pytest.mark.parametrize(
    ("elapsed", "rss", "expected"),
    [(1, 1024, None), (11, 1024, "budget_exceeded"), (1, 2048 * 1024 * 1024, "memory_limit_exceeded")],
)
def test_resource_guard_stops_optional_work(elapsed: float, rss: int, expected: str | None) -> None:
    policy = DeepCollectionPolicy(budget_seconds=10, resource_memory_limit_mb=2048)
    sample = {"status": "ok", "rss_bytes": rss}
    assert PyVmomiDeepCollector._resource_guard_status(elapsed, sample, policy) == expected
    if expected is None:
        assert PyVmomiDeepCollector._resource_guard_status(elapsed, sample, policy, cancelled=True) == "cancelled"
    cpu_sample = {"status": "ok", "rss_bytes": 1024, "cpu_percent_of_machine": 96.0}
    assert PyVmomiDeepCollector._resource_guard_status(1, cpu_sample, policy) == "cpu_limit_exceeded"


def test_certificate_collection_guard_stops_before_next_host() -> None:
    expiry = datetime.now(UTC) + timedelta(days=45)
    hosts = [
        NS(name="esx-1", _moId="host-1", config=NS(certificateInfo=NS(notAfter=expiry))),
        NS(name="esx-2", _moId="host-2", config=NS(certificateInfo=NS(notAfter=expiry))),
    ]
    guard_reasons = iter((None, None, None, "memory_limit_exceeded"))
    guard = LatchedResourceGuard(lambda: next(guard_reasons))
    collector = PyVmomiDeepCollector("vc", "u", "p")
    collector._probe_host_tls_not_after = lambda *_args, **_kwargs: None
    records = collector._certificate_records(
        hosts,
        "ds-resource-cert",
        "2026-09-25T00:00:00Z",
        resource_guard=guard,
    )

    assert [record.entity.display_ref for record in records] == ["vc", "esx-1", "esx-2"]
    assert [record.value["certificate_status"] for record in records] == ["unavailable", "available", "not_requested"]
    assert all(record.metadata["rule_id"] == "SEC-DEEP-003" for record in records)


def test_certificate_collection_covers_vcenter_hosts_and_expired_certificates() -> None:
    soon = datetime.now(UTC) + timedelta(days=30)
    expired = datetime.now(UTC) - timedelta(days=2)
    collector = PyVmomiDeepCollector("vc", "u", "p")
    collector._probe_host_tls_not_after = lambda endpoint, **_kwargs: soon if endpoint == "vc" else None
    hosts = [NS(name="esx-expired", _moId="host-expired", config=NS(certificateInfo=NS(notAfter=expired)))]

    records = collector._certificate_records(hosts, "ds-certificate-lifecycle", "2026-09-28T01:00:00Z")

    by_entity = {record.entity.type: record for record in records}
    assert set(by_entity) == {"vCenter", "HostSystem"}
    assert by_entity["vCenter"].value["source"] == "vCenter TLS certificate"
    assert by_entity["vCenter"].finding is True
    assert by_entity["HostSystem"].value["days_remaining"] < 0
    assert by_entity["HostSystem"].finding is True


def test_certificate_read_failure_is_incomplete_not_a_clean_result() -> None:
    collector = PyVmomiDeepCollector("vc", "u", "p")
    collector._probe_host_tls_not_after = lambda *_args, **_kwargs: None
    hosts = [NS(name="esx-cert-missing", _moId="host-cert-missing", config=NS())]
    records = collector._certificate_records(hosts, "ds-certificate-unavailable", "2026-09-28T01:00:00Z")
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-certificate-unavailable"})
    dataset.records = [item for item in dataset.records if item.metadata.get("rule_id") != "SEC-DEEP-003"] + records
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "SEC-DEEP-003")

    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert len(records) == 2
    assert all(record.window.completeness == 0.0 for record in records)
    assert result.status == DeepResultStatus.INSUFFICIENT_DATA


def test_license_expiry_and_capacity_are_reported_without_license_keys() -> None:
    license_item = NS(
        name="vSphere Evaluation",
        editionKey="evaluation",
        expirationDate=datetime.now(UTC) + timedelta(days=20),
        used=3,
        total=2,
        costUnit="cpuCore",
        licenseKey="SENSITIVE-LICENSE-KEY",
    )
    records = PyVmomiDeepCollector("vc", "u", "p")._lifecycle_records(
        NS(licenseManager=NS(licenses=[license_item])),
        [],
        [],
        "ds-license-risk",
        "2026-09-28T01:00:00Z",
    )
    record = next(item for item in records if item.metadata.get("rule_id") == "SEC-DEEP-004")

    assert record.finding is True
    assert record.value["evaluation_license"] is True
    assert record.value["expiring_or_expired_count"] == 1
    assert record.value["used_exceeds_total_count"] == 1
    assert "SENSITIVE-LICENSE-KEY" not in record.model_dump_json()


def test_empty_license_inventory_is_insufficient_not_pass() -> None:
    records = PyVmomiDeepCollector("vc", "u", "p")._lifecycle_records(
        NS(licenseManager=NS(licenses=[])),
        [],
        [],
        "ds-license-empty",
        "2026-09-28T01:00:00Z",
    )
    record = next(item for item in records if item.metadata.get("rule_id") == "SEC-DEEP-004")
    dataset = DeepDataset.from_json(FIXTURE)
    dataset.manifest = dataset.manifest.model_copy(update={"dataset_id": "ds-license-empty"})
    dataset.records = [item for item in dataset.records if item.metadata.get("rule_id") != "SEC-DEEP-004"] + records
    rule = next(item for item in BUILTIN_DEEP_RULES if item.rule_id == "SEC-DEEP-004")

    result = DeepAnalyzer().analyze(dataset, rules=(rule,)).rule_results[0]

    assert record.value["collection_status"] == "empty"
    assert record.window.completeness == 0.0
    assert result.status == DeepResultStatus.INSUFFICIENT_DATA


def test_resource_guard_latches_first_stop_reason_for_session() -> None:
    checks = iter((None, "budget_exceeded", None))
    guard = LatchedResourceGuard(lambda: next(checks))
    assert [guard(), guard(), guard()] == [None, "budget_exceeded", "budget_exceeded"]


def test_resource_guard_stops_later_core_phases_and_degraded_analysis_has_no_pass(monkeypatch: pytest.MonkeyPatch) -> None:
    from pyVim import connect
    from vstacklens.deep import pyvmomi_collector as collector_module

    content = NS(
        about=NS(apiVersion="8.0", version="8.0", build="test"),
        taskManager=None,
        alarmManager=None,
        licenseManager=None,
        perfManager=None,
        diagnosticManager=None,
        eventManager=None,
    )
    service_instance = NS(RetrieveContent=lambda: content)

    class Monitor:
        def __init__(self, *_args, **_kwargs):
            self.samples = 0

        def start(self):
            return None

        def latest_sample(self):
            self.samples += 1
            rss = 1024 if self.samples == 1 else 20 * 1024 * 1024
            return {"status": "ok", "rss_bytes": rss, "cpu_percent_of_machine": 1.0}

        def stop(self):
            return {"status": "ok", "sample_count": self.samples}

    resource_sample = {"status": "ok", "cpu_user_sec": 0.1, "cpu_system_sec": 0.1, "rss_bytes": 1024, "read_bytes": 1, "write_bytes": 1, "network_sent_bytes": 1, "network_recv_bytes": 1}
    monkeypatch.setattr(collector_module, "LocalResourceMonitor", Monitor)
    monkeypatch.setattr(connect, "SmartConnect", lambda **_kwargs: service_instance)
    monkeypatch.setattr(connect, "Disconnect", lambda _service: None)
    collector = PyVmomiDeepCollector("vc", "u", "p")
    collector._resource_snapshot = lambda include_network=True: dict(resource_sample)
    collector._view = lambda _content, _kind: []
    snapshot_calls = []
    collector._snapshot_records = lambda *_args: snapshot_calls.append(True) or []

    dataset = collector.collect(policy=DeepCollectionPolicy(resource_memory_limit_mb=10))
    actions = {item["action"]: item for item in dataset.collection_log}
    analysis = DeepAnalyzer().analyze(dataset)

    assert snapshot_calls == []
    assert actions["lifecycle.certificate"]["status"] == "partial"
    assert actions["vm.snapshot"]["status"] == "not_requested"
    assert actions["vm.snapshot"]["reason"] == "memory_limit_exceeded"
    assert dataset.manifest.impact["degraded"] is True
    assert all(item.status != DeepResultStatus.PASS for item in analysis.rule_results)


def test_resource_monitor_starts_before_vcenter_connect_and_connection_failure_returns_unavailable_dataset(monkeypatch: pytest.MonkeyPatch) -> None:
    from pyVim import connect
    from vstacklens.deep import pyvmomi_collector as collector_module

    calls = []

    class Monitor:
        def __init__(self, *_args, **_kwargs):
            pass

        def start(self):
            calls.append("monitor_start")

        def stop(self):
            calls.append("monitor_stop")
            return {"status": "limited"}

    def smart_connect(**_kwargs):
        calls.append("connect")
        raise TimeoutError("synthetic connect timeout SECRET-CREDENTIAL")

    monkeypatch.setattr(collector_module, "LocalResourceMonitor", Monitor)
    monkeypatch.setattr(connect, "SmartConnect", smart_connect)
    collector = PyVmomiDeepCollector("vc", "u", "SECRET-CREDENTIAL")
    resource_sample = {"status": "ok", "cpu_user_sec": 0.1, "cpu_system_sec": 0.1, "rss_bytes": 4096, "read_bytes": 1, "write_bytes": 1, "network_sent_bytes": 1, "network_recv_bytes": 1}
    collector._resource_snapshot = lambda include_network=True: dict(resource_sample)

    dataset = collector.collect()

    assert calls == ["monitor_start", "connect", "monitor_stop"]
    assert dataset.manifest.impact["degraded"] is True
    assert dataset.manifest.collection_status["failed"] == 1
    assert next(item for item in dataset.collection_log if item["action"] == "connection")["status"] == "timeout"
    log_collection = next(item for item in dataset.collection_log if item["action"] == "logs.history")
    assert log_collection["status"] == "unavailable"
    categories = {item["category"]: item for item in log_collection["category_results"]}
    assert len(categories) == 8
    assert categories["vpxd"]["status"] == "timeout"
    assert all(categories[name]["reason"] == "host_inventory_unavailable" for name in ("vmkernel", "hostd", "vpxa", "vobd", "syslog", "vsan"))
    log_record = next(record for record in dataset.records if record.kind == "log")
    assert log_record.metadata["log_category"] == "vpxd"
    assert log_record.metadata["log_status"] == "timeout"
    assert "SECRET-CREDENTIAL" not in dataset.model_dump_json()
    analysis = DeepAnalyzer().analyze(dataset)
    assert all(result.status != DeepResultStatus.PASS for result in analysis.rule_results)


def test_run_deep_vcenter_cli_returns_nonzero_for_connection_failure(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    import argparse
    from vstacklens import cli

    report_dir = tmp_path / "deep-report"
    failure = NS(
        report_dir=report_dir,
        html_path=report_dir / "index.html",
        diagnostic_path=tmp_path / "deep-report.deep-diagnostic.json",
        docx_path=None,
        analysis=NS(diagnostic=NS(collection_log=[{"action": "connection", "status": "timeout"}], impact={"degraded": True})),
    )

    class Service:
        def collect_and_analyze_vcenter(self, *_args, **_kwargs):
            return failure

    monkeypatch.setattr(cli, "DeepInspectionService", Service)
    args = argparse.Namespace(
        connection_workbook=tmp_path / "readonly.xlsx",
        report_dir=report_dir,
        docx_out=None,
        resource_budget_seconds=60,
        resource_memory_limit_mb=512,
        resource_cpu_limit_percent=90.0,
        max_log_total_mib=16,
        max_log_primary_mib=12,
        max_log_fallback_mib=4,
        correlation_window_seconds=3600,
        event_history_days=30,
    )

    with pytest.raises(SystemExit) as exit_info:
        cli.cmd_run_deep_vcenter(args)

    output = capsys.readouterr().out
    assert exit_info.value.code == 2
    assert "collection unavailable (timeout)" in output
    assert "deep-report.dataset.json" in output


def test_retrieve_content_timeout_returns_unavailable_dataset_and_disconnects(monkeypatch: pytest.MonkeyPatch) -> None:
    from pyVim import connect
    from vstacklens.deep import pyvmomi_collector as collector_module

    calls = []
    stopped = {"value": False}

    class Monitor:
        def __init__(self, *_args, **_kwargs):
            pass

        def start(self):
            calls.append("monitor_start")

        def stop(self):
            if not stopped["value"]:
                calls.append("monitor_stop")
                stopped["value"] = True
            return {"status": "limited"}

    def retrieve_content():
        calls.append("retrieve_content")
        raise TimeoutError("synthetic RetrieveContent timeout SECRET-CREDENTIAL")

    service = NS(RetrieveContent=retrieve_content)
    monkeypatch.setattr(collector_module, "LocalResourceMonitor", Monitor)
    monkeypatch.setattr(connect, "SmartConnect", lambda **_kwargs: calls.append("connect") or service)
    monkeypatch.setattr(connect, "Disconnect", lambda _service: calls.append("disconnect"))
    collector = PyVmomiDeepCollector("vc", "u", "SECRET-CREDENTIAL")
    resource_sample = {"status": "ok", "cpu_user_sec": 0.1, "cpu_system_sec": 0.1, "rss_bytes": 4096, "read_bytes": 1, "write_bytes": 1, "network_sent_bytes": 1, "network_recv_bytes": 1}
    collector._resource_snapshot = lambda include_network=True: dict(resource_sample)

    dataset = collector.collect()

    assert calls == ["monitor_start", "connect", "retrieve_content", "monitor_stop", "disconnect"]
    assert next(item for item in dataset.collection_log if item["action"] == "connection")["status"] == "ok"
    assert next(item for item in dataset.collection_log if item["action"] == "session.content")["status"] == "timeout"
    assert next(item for item in dataset.collection_log if item["action"] == "logs.history")["status"] == "unavailable"
    assert "SECRET-CREDENTIAL" not in dataset.model_dump_json()


def test_daily_performance_guard_stops_before_query_after_counter_probe() -> None:
    class PerfManager:
        perfCounter = [NS(key=1, nameInfo=NS(key="used"), groupInfo=NS(key="storage"))]

        def __init__(self):
            self.available_queries = 0
            self.performance_queries = 0

        def QueryAvailablePerfMetric(self, **_kwargs):
            self.available_queries += 1
            return [NS(counterId=1, instance="")]

        def QueryPerf(self, **_kwargs):
            self.performance_queries += 1
            return []

    manager = PerfManager()
    guard_checks = iter((None, "cpu_limit_exceeded"))
    collector = PyVmomiDeepCollector("vc", "u", "p")
    records = collector._collect_daily_perf(
        NS(perfManager=manager),
        [NS(name="ds-1", _moId="datastore-1")],
        "ds-resource-daily-perf",
        "2026-09-25T00:00:00Z",
        DeepCollectionPolicy(),
        resource_guard=lambda: next(guard_checks),
    )

    assert records == []
    assert manager.available_queries == 1
    assert manager.performance_queries == 0


def test_enhanced_performance_guard_stops_before_next_object_query() -> None:
    class PerfManager:
        perfCounter = []

        def __init__(self):
            self.entities = []

        def QueryAvailablePerfMetric(self, entity, **_kwargs):
            self.entities.append(entity.name)
            return []

    checks = {"count": 0}

    def resource_guard() -> str | None:
        checks["count"] += 1
        return None if checks["count"] == 1 else "memory_limit_exceeded"

    perf_manager = PerfManager()
    collector = PyVmomiDeepCollector("vc", "u", "p")
    records, capabilities = collector._collect_enhanced_perf_records(
        NS(perfManager=perf_manager),
        [NS(name="esx-1", _moId="host-1"), NS(name="esx-2", _moId="host-2")],
        [],
        "ds-resource-guard",
        "2026-09-25T00:00:00Z",
        DeepCollectionPolicy(),
        resource_guard=resource_guard,
    )

    assert perf_manager.entities == ["esx-1"]
    assert records == []
    assert all(item.status == CapabilityStatus.NOT_REQUESTED for item in capabilities)
    assert all(item.detail["reason"] == "memory_limit_exceeded" for item in capabilities)


def test_event_collection_cancellation_marks_partial_window_without_false_pass() -> None:
    calls = {"count": 0, "destroyed": False}

    def cancel_requested() -> bool:
        calls["count"] += 1
        return calls["count"] >= 3

    class EventCollector:
        def ReadNext(self, _count):
            event = NS(createdTime=datetime(2026, 9, 24, 0, 0, tzinfo=UTC), fullFormattedMessage="Physical NIC vmnic0 linkstate is down")
            return [event] * _count

        def DestroyCollector(self):
            calls["destroyed"] = True

    manager = NS(CreateCollectorForEvents=lambda _spec: EventCollector())
    collector = PyVmomiDeepCollector("vc", "u", "p", cancel_requested=cancel_requested)
    events = collector._bounded_events(NS(eventManager=manager), DeepCollectionPolicy(max_history_records=1001))
    assert len(events) == 1000
    assert calls["destroyed"] is True
    assert collector._event_query_meta["reason"] == "cancelled_during_event_query"
    history = collector._history_records(events, [], [], [], "ds-cancel", "2026-09-24T00:01:00Z", collector._event_query_meta)
    link = next(item for item in history if item.metadata["rule_id"] == "NET-DEEP-001")
    network_error = next(item for item in history if item.metadata["rule_id"] == "NET-DEEP-002")
    assert link.finding is False
    assert network_error.finding is True
    assert link.window.completeness == 0.0


def test_event_collection_resource_guard_stops_between_batches_and_marks_degraded() -> None:
    state = {"guard_calls": 0, "read_calls": 0, "destroyed": False}

    def should_stop() -> str | None:
        state["guard_calls"] += 1
        return "budget_exceeded" if state["guard_calls"] >= 3 else None

    class EventCollector:
        def ReadNext(self, count):
            state["read_calls"] += 1
            return [NS(createdTime=datetime(2026, 9, 24, 0, 0, tzinfo=UTC), fullFormattedMessage="event") for _ in range(count)]

        def DestroyCollector(self):
            state["destroyed"] = True

    manager = NS(CreateCollectorForEvents=lambda _spec: EventCollector())
    collector = PyVmomiDeepCollector("vc", "u", "p")
    events = collector._bounded_events(
        NS(eventManager=manager),
        DeepCollectionPolicy(max_history_records=2000),
        should_stop=should_stop,
    )

    assert len(events) == 1000
    assert state["read_calls"] == 1
    assert state["destroyed"] is True
    assert collector._event_query_meta["degraded"] is True
    assert collector._event_query_meta["reason"] == "budget_exceeded"
    assert collector._event_query_meta["batches"] == 1


def test_local_resource_monitor_captures_peak_cpu_and_rss() -> None:
    calls = {"count": 0}

    def sample() -> dict[str, object]:
        calls["count"] += 1
        value = calls["count"]
        return {"status": "ok", "cpu_user_sec": value * 0.01, "cpu_system_sec": value * 0.005, "rss_bytes": value * 1024, "read_bytes": value * 10, "write_bytes": value * 20}

    monitor = LocalResourceMonitor(sample, interval_seconds=0.1)
    monitor.start()
    time.sleep(0.25)
    result = monitor.stop()
    assert result["status"] == "ok"
    assert result["sample_count"] >= 3
    assert result["peak_rss_bytes"] >= result["sample_count"] * 1024
    assert result["peak_cpu_percent_of_one_core"] > 0


def test_customer_visible_finding_requires_verification_guidance() -> None:
    # The contract must fail before a customer-visible finding can reach a
    # renderer without a manual verification path.
    with pytest.raises(ValueError, match="verification_guidance"):
        DeepFinding(
            finding_id="f:bad",
            rule_id="BAD-001",
            rule_version="1.0.0",
            title="缺少人工复核指引",
            category="current_risk",
            confidence_class="FACT",
            scope=FindingScope(scope_type=ScopeType.ENTITY, scope_id="host-1", dimension_key="bad"),
            entities=[DeepEntity(type="HostSystem", stable_id="host-1")],
            verification_guidance="",
        )


def test_empty_connection_failure_reports_explain_no_analysis_without_false_pass(tmp_path: Path) -> None:
    from vstacklens.deep.contracts import DeepReport
    from vstacklens.deep.report import DeepHtmlReportBuilder, DeepWordReportBuilder

    dataset = DeepDataset.from_json(FIXTURE)
    dataset.records = []
    dataset.manifest = dataset.manifest.model_copy(update={"scope": {"hosts": 0, "vms": 0, "clusters": 0, "datastores": 0}, "impact": {"degraded": True}})
    dataset.collection_log = [{"action": "connection", "status": "timeout"}]
    report = DeepReport(analysis_scope=[], pass_summary=[], findings=[])
    html_path = DeepHtmlReportBuilder().render(report, tmp_path / "failure.html", dataset)
    docx_path = DeepWordReportBuilder().render(report, tmp_path / "failure.docx", dataset)
    html = html_path.read_text(encoding="utf-8")
    docx = Document(docx_path)
    word_text = "\n".join([p.text for p in docx.paragraphs] + [c.text for table in docx.tables for row in table.rows for c in row.cells])

    for text in (html, word_text):
        assert "本次未能连接 vCenter 并取得环境清单" in text
        assert "本报告不表示环境正常" in text
        assert "timeout" not in text.casefold()
        assert "UNAVAILABLE" not in text
    assert "无可展示的通过项聚合" in word_text
    assert "Evidence 附录" not in word_text


@pytest.mark.parametrize(
    ("log_status", "successful_count", "expected_phrase"),
    [
        ("partial", 20, "日志仅部分采集成功（20/36 个对象与日志类别组合成功读取或确认无内容）"),
        ("unavailable", 0, "本次未取得可用的历史日志（0/36 个对象与日志类别组合成功读取或确认无内容）"),
    ],
)
def test_customer_report_discloses_log_coverage_without_claiming_normal(
    tmp_path: Path,
    log_status: str,
    successful_count: int,
    expected_phrase: str,
) -> None:
    from vstacklens.deep.contracts import DeepReport
    from vstacklens.deep.report import DeepHtmlReportBuilder, DeepWordReportBuilder

    dataset = DeepDataset.from_json(FIXTURE)
    dataset.collection_log = [
        {
            "action": "logs.history",
            "status": log_status,
            "category_count": 36,
            "successful_count": successful_count,
            "missing_privilege_ids": ["Global.Diagnostics"],
            "truncated_count": 20 if log_status == "partial" else 0,
        }
    ]
    report = DeepReport(analysis_scope=[], pass_summary=[], findings=[])
    html_path = DeepHtmlReportBuilder().render(report, tmp_path / f"logs-{log_status}.html", dataset)
    docx_path = DeepWordReportBuilder().render(report, tmp_path / f"logs-{log_status}.docx", dataset)
    html = html_path.read_text(encoding="utf-8")
    word = Document(docx_path)
    word_text = "\n".join([paragraph.text for paragraph in word.paragraphs] + [cell.text for table in word.tables for row in table.rows for cell in row.cells])

    for rendered in (html, word_text):
        assert expected_phrase in rendered
        assert "部分来源读取权限受限" in rendered
        assert "本次未形成日志与事件/任务/性能的可验证关联" in rendered
        assert "未读取、未保留或未覆盖的历史范围不作正常性判断" in rendered
        assert "Global.Diagnostics" not in rendered
        assert "远端 Syslog" not in rendered


def test_customer_report_log_coverage_describes_verified_correlation_without_claiming_causality() -> None:
    from vstacklens.deep.report import _log_coverage_note

    dataset = DeepDataset.from_json(FIXTURE)
    dataset.collection_log = [{"action": "logs.history", "status": "ok", "category_count": 1, "successful_count": 1}]
    dataset.records = [
        DatasetRecord(
            record_id="log-with-correlation",
            dataset_id=dataset.dataset_id,
            kind="log",
            entity=DeepEntity(type="HostSystem", stable_id="host-1"),
            collected_at_utc="2026-09-25T00:00:00Z",
            source=DeepSource(api="vim.DiagnosticManager.BrowseDiagnosticLog"),
            window=DeepWindow(start="2026-09-24T00:00:00Z", end="2026-09-25T00:00:00Z"),
            value={"lines": []},
            metadata={"rule_id": "LOG-DEEP-001", "related_event_count": 1, "related_record_count": 0},
        )
    ]

    note = _log_coverage_note(dataset)

    assert note is not None
    assert "1 条日志记录存在同对象、时间窗且类型相关的证据关联" in note
    assert "共现不代表因果" in note
    assert "未形成日志与事件/任务/性能的可验证关联" not in note


def test_customer_report_log_coverage_shows_actual_retained_time_ranges_and_pages() -> None:
    from vstacklens.deep.report import _log_coverage_note

    dataset = DeepDataset.from_json(FIXTURE)
    dataset.collection_log = [
        {
            "action": "logs.history",
            "status": "partial",
            "category_count": 3,
            "successful_count": 1,
            "partially_successful_count": 1,
            "segment_read_count": 8,
            "deduplicated_source_line_count": 2,
            "retained_log_bytes": 1024 * 1024,
            "retained_primary_log_bytes": 768 * 1024,
            "retained_fallback_log_bytes": 256 * 1024,
            "max_log_total_bytes": 16 * 1024 * 1024,
            "max_log_primary_bytes": 12 * 1024 * 1024,
            "max_log_fallback_bytes": 4 * 1024 * 1024,
            "requested_log_days": 7,
            "log_window_cutoff_utc": "2026-09-20T12:00:00Z",
            "time_window_status": "complete_with_unknown_time",
            "source_log_bytes_read": 1200 * 1024,
            "time_window_excluded_line_count": 250,
            "time_unknown_line_count": 1,
            "attempts": [
                {"source": "PowerCLI.Get-Log", "status": "partial", "result_status_counts": {"interface_unavailable": 9}},
                {"source": "vim.DiagnosticManager.DirectESXi", "status": "not_attempted", "result_status_counts": {"no_logs": 2, "total_size_limit_reached": 1}},
            ],
            "truncated_count": 2,
        }
    ]
    entity = DeepEntity(type="HostSystem", stable_id="host-log-window", display_ref="esx-log-window")
    dataset.records = [
        DatasetRecord(
            record_id="log-vpxd-window",
            dataset_id=dataset.dataset_id,
            kind="log",
            entity=DeepEntity(type="vCenter", stable_id="vc-log-window", display_ref="vcenter-log-window"),
            collected_at_utc="2026-09-25T12:00:00Z",
            source=DeepSource(api="vim.DiagnosticManager.BrowseDiagnosticLog"),
            window=DeepWindow(start="2026-09-24T00:00:00Z", end="2026-09-25T12:00:00Z", sample_count=2),
            value={"lines": ["line-one", "line-two"], "line_count": 2, "marker_counts": {"error": 2, "network": 1}},
            metadata={"rule_id": "LOG-DEEP-001", "log_category": "vpxd", "log_status": "partial", "time_source": "line_timestamp"},
        ),
        DatasetRecord(
            record_id="log-hostd-window",
            dataset_id=dataset.dataset_id,
            kind="log",
            entity=entity,
            collected_at_utc="2026-09-25T12:00:00Z",
            source=DeepSource(api="vim.DiagnosticManager.BrowseDiagnosticLog"),
            window=DeepWindow(start="2026-09-24T12:00:00Z", end="2026-09-25T12:00:00Z", sample_count=1),
            value={"lines": ["line-three"], "line_count": 1},
            metadata={"rule_id": "LOG-DEEP-001", "log_category": "hostd", "log_status": "partial", "time_source": "line_timestamp"},
        ),
    ]

    note = _log_coverage_note(dataset)

    assert note is not None
    assert "另有 1 个组合部分读取" in note
    assert "实际保留 3 行日志，分段读取 8 次" in note
    assert "1.00/16.00 MiB" in note
    assert "vCenter 主来源 0.75/12.00 MiB" in note
    assert "备用来源合计 0.25/4.00 MiB" in note
    assert "跨来源按同对象、类别、源行号及相同行内容去重 2 行" in note
    assert "最近 7 天" in note
    assert "窗口定位完成但含时间无法判定行" in note
    assert "跳过窗口外 250 行" in note
    assert "保留 1.00 MiB" in note
    assert "保留并标记时间无法判定 1 行" in note
    assert "更早的未读取行不计入本次日志分析" in note
    assert "PowerCLI partial（结果：接口不可用或未返回可用日志 9）" in note
    assert "直连 ESXi not_attempted（结果：未发现匹配的日志描述符 2、日志总预算耗尽 1）" in note
    assert "日志行时间戳跨度 2026-09-24 00:00:00 UTC 至 2026-09-25 12:00:00 UTC（约 1.5 天" in note
    assert "类别时间窗：hostd" in note and "vCenter" in note
    assert "不代表连续或全量覆盖" in note
    assert "日志关键词标记累计命中：错误 2、网络 1" in note
    assert "同一行或不同来源可能重复计数" in note
    assert "line-one" not in note


def test_customer_reports_show_matched_log_event_task_and_performance_pointers(tmp_path: Path) -> None:
    from vstacklens.deep.contracts import DeepReport
    from vstacklens.deep.report import DeepHtmlReportBuilder, DeepWordReportBuilder

    dataset = DeepDataset.from_json(FIXTURE)
    dataset.collection_log = [{"action": "logs.history", "status": "ok", "category_count": 1, "successful_count": 1}]
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-1")
    log = DatasetRecord(
        record_id="log-host-1",
        dataset_id=dataset.dataset_id,
        kind="log",
        entity=entity,
        collected_at_utc="2026-09-25T00:00:00Z",
        source=DeepSource(api="vim.DiagnosticManager.BrowseDiagnosticLog"),
        window=DeepWindow(start="2026-09-24T00:00:00Z", end="2026-09-25T00:00:00Z"),
        value={"lines": ["PRIVATE LOG BODY"]},
        raw_pointer="logs/log.ndjson#host-1/ha/direct-esxi",
        summary="HA log sample",
        metadata={
            "rule_id": "LOG-DEEP-001",
            "log_category": "ha",
            "correlation_summary": "vmnic4 link down",
            "related_event_count": 1,
            "related_event_samples": [
                {
                    "event_type": "vim.event.HostFailedEvent",
                    "event_key": 487709,
                    "dataset_pointer": "events/event.ndjson#NET-DEEP-001/host-1",
                    "timestamp": "2026-09-24T23:59:48Z",
                    "log_timestamp": "2026-09-24T23:59:36Z",
                    "log_line_index": 1,
                    "source_line_number": 3621,
                    "object": "esx-1",
                    "interfaces": ["vmnic4"],
                    "state": "down",
                    "host": "esx-1",
                    "message": "PRIVATE EVENT MESSAGE",
                    "object_ids": ["host-1"],
                    "time_distance_seconds": 12,
                    "matching_categories": ["ha"],
                    "parent_dataset_pointer": "parent:r4/events/event.ndjson#event-key-487709",
                }
            ],
            "related_record_count": 2,
            "related_record_samples": [
                {
                    "record_id": "task-1",
                    "kind": "task",
                    "source": "TaskManager.recentTask",
                    "timestamp": "2026-09-24T23:59:45Z",
                    "window_start": "2026-09-24T23:59:45Z",
                    "window_end": "2026-09-24T23:59:45Z",
                    "time_distance_seconds": 15,
                    "matching_categories": ["ha"],
                    "summary": "PRIVATE TASK SUMMARY",
                },
                {
                    "record_id": "perf-1",
                    "kind": "perf",
                    "source": "PerformanceManager.QueryPerf",
                    "timestamp": "2026-09-25T00:00:00Z",
                    "window_start": "2026-09-24T23:55:00Z",
                    "window_end": "2026-09-25T00:00:00Z",
                    "time_distance_seconds": 0,
                    "matching_categories": ["performance"],
                    "summary": "PRIVATE PERFORMANCE SUMMARY",
                },
            ],
        },
    )
    task = DatasetRecord(
        record_id="task-1",
        dataset_id=dataset.dataset_id,
        kind="task",
        entity=entity,
        collected_at_utc="2026-09-25T00:00:00Z",
        source=DeepSource(api="TaskManager.recentTask"),
        window=DeepWindow(start="2026-09-24T23:59:45Z", end="2026-09-24T23:59:45Z"),
        value=1,
        raw_pointer="events/task.ndjson#task-1",
    )
    perf = DatasetRecord(
        record_id="perf-1",
        dataset_id=dataset.dataset_id,
        kind="perf",
        entity=entity,
        collected_at_utc="2026-09-25T00:00:00Z",
        source=DeepSource(api="PerformanceManager.QueryPerf"),
        window=DeepWindow(start="2026-09-24T23:55:00Z", end="2026-09-25T00:00:00Z"),
        value=95.0,
        raw_pointer="perf/perf.ndjson#cpu-ready",
    )
    dataset.records = [log, task, perf]
    report = DeepReport(analysis_scope=[], pass_summary=[], findings=[])
    html_path = DeepHtmlReportBuilder().render(report, tmp_path / "correlated.html", dataset)
    docx_path = DeepWordReportBuilder().render(report, tmp_path / "correlated.docx", dataset)
    html = html_path.read_text(encoding="utf-8")
    word_doc = Document(docx_path)
    word = "\n".join([paragraph.text for paragraph in word_doc.paragraphs] + [cell.text for table in word_doc.tables for row in table.rows for cell in row.cells])

    for rendered in (html, word):
        assert "日志时间关联（共现不代表因果）" in rendered
        assert "vim.event.HostFailedEvent" in rendered
        assert "Event Key 487709" in rendered
        assert "日志 2026-09-24T23:59:36Z（源日志第 3621 行），摘要 vmnic4 link down" in rendered
        assert "vmnic4 down" in rendered
        assert "events/event.ndjson#NET-DEEP-001/host-1" in rendered
        assert "parent:r4/events/event.ndjson#event-key-487709" in rendered
        assert "TaskManager.recentTask" in rendered
        assert "PerformanceManager.QueryPerf" in rendered
        assert "logs/log.ndjson#host-1/ha/direct-esxi" in rendered
        assert "events/task.ndjson#task-1" in rendered
        assert "perf/perf.ndjson#cpu-ready" in rendered
        assert "PRIVATE LOG BODY" not in rendered
        assert "PRIVATE EVENT MESSAGE" not in rendered
        assert "PRIVATE TASK SUMMARY" not in rendered
        assert "PRIVATE PERFORMANCE SUMMARY" not in rendered


def test_threshold_policy_requires_source_metadata() -> None:
    with pytest.raises(ValueError, match="requires rationale"):
        ThresholdPolicy(id="bad", value=0.8, unit="ratio", source="vstacklens_engineering")


def test_deep_service_emits_customer_report_and_internal_diagnostic(tmp_path: Path) -> None:
    report_dir = tmp_path / "deep-report"
    docx_path = tmp_path / "deep-report.docx"
    result = DeepInspectionService().analyze_dataset(FIXTURE, report_dir, docx_out=docx_path)

    assert result.html_path.exists()
    assert result.docx_path == docx_path
    assert docx_path.exists()
    assert result.diagnostic_path.exists()
    assert (tmp_path / "deep-history.db").exists()
    html = result.html_path.read_text(encoding="utf-8")
    assert "L0 摘要" in html
    assert "本次已完成分析范围" in html
    assert "趋势" in html
    assert "verification_guidance" not in html
    assert "NOT_EVALUATED" not in html
    assert "TrendAnalyzer" in html

    text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    assert "人工复核指引" in text
    assert "Evidence 附录" in text
    assert "NOT_EVALUATED" not in text

    diagnostic = json.loads(result.diagnostic_path.read_text(encoding="utf-8"))
    assert {item["status"] for item in diagnostic["rule_results"]} >= {"FINDING", "PASS"}


def test_customer_reports_summarize_event_evidence_and_keep_full_dataset_pointer(tmp_path: Path) -> None:
    from vstacklens.deep.contracts import DeepReport, DeepSource, DeepWindow
    from vstacklens.deep.report import DeepHtmlReportBuilder, DeepWordReportBuilder, _evidence_summary_lines

    dataset = DeepDataset.from_json(FIXTURE)
    entity = DeepEntity(type="HostSystem", stable_id="host-1", display_ref="esx-1")
    record = DatasetRecord(
        record_id="event-history-host-1",
        kind="event",
        dataset_id=dataset.dataset_id,
        collected_at_utc="2026-09-25T00:00:00Z",
        source=DeepSource(api="EventManager.QueryEvents"),
        entity=entity,
        selector={"rule_id": "HIST-DEEP-001"},
        window=DeepWindow(start="2026-09-24T00:00:00Z", end="2026-09-25T00:00:00Z", sample_count=5),
        value=5,
        unit="count",
        summary="五条历史事件命中该规则。",
        raw_pointer="events/event.ndjson#host-1/HIST-DEEP-001",
        metadata={
            "event_filter": ["link down"],
            "event_samples": [
                {
                    "timestamp": f"2026-09-24T00:0{index}:00Z",
                    "event_type": "HostConnectionLostEvent",
                    "event_key": index,
                    "host": "esx-1",
                    "message": f"sample-event-{index}",
                    "object_ids": [f"private-moid-{index}"],
                    "resolved_entities": [{"stable_id": f"private-stable-id-{index}"}],
                }
                for index in range(1, 6)
            ],
        },
    )
    evidence = evidence_ref(record)
    assert len(evidence.value_summary["event_samples"]) == 3
    assert evidence.value_summary["event_samples_count"] == 5
    assert len(record.metadata["event_samples"]) == 5
    config_evidence = evidence.model_copy(
        update={
            "kind": "config",
            "value_summary": {
                "value": {
                    "configured_guest": "RHEL 9",
                    "guest_reported": "RHEL 8",
                    "tools_running": True,
                    "internal_payload": [1, 2, 3],
                    "distributed_network_profiles": {
                        "status": "available",
                        "switch_count": 1,
                        "portgroup_count": 1,
                        "portgroups": [
                            {"dvs_name": "dvSwitch-1", "name": "Management", "vlan": {"spec_type": "VlanIdSpec", "vlanId": 42}, "teaming": {"policy": "loadbalance_loadbased"}}
                        ],
                    },
                },
                "unit": "state",
                "summary": "Configured and reported operating systems do not match.",
            },
        }
    )
    config_lines = _evidence_summary_lines(config_evidence)
    assert "摘要：Configured and reported operating systems do not match." in config_lines
    assert "关键数据：configured_guest=RHEL 9；guest_reported=RHEL 8；tools_running=True" in config_lines
    assert any("dvSwitch-1/Management" in line and "VLAN ID 42" in line and "loadbalance_loadbased" in line for line in config_lines)
    assert all(not line.startswith("{") for line in config_lines)
    finding = DeepFinding(
        finding_id="f:event-summary",
        rule_id="HIST-DEEP-001",
        rule_version="1.0.0",
        title="历史链路事件",
        category="historical_health",
        confidence_class="FACT",
        scope=FindingScope(scope_type=ScopeType.ENTITY, scope_id="host-1", dimension_key="history"),
        entities=[entity],
        evidence=[evidence],
        fact="历史链路事件达到规则条件。",
        finding="发现重复链路事件。",
        verification_guidance="在 vCenter Events 中按时间与主机复核。",
    )
    report = DeepReport(findings=[finding])
    html_path = DeepHtmlReportBuilder().render(report, tmp_path / "event-summary.html", dataset)
    docx_path = DeepWordReportBuilder().render(report, tmp_path / "event-summary.docx", dataset)
    html_text = html_path.read_text(encoding="utf-8")
    payload_text = (html_path.parent / "deep_report_payload.json").read_text(encoding="utf-8")
    event_doc = Document(docx_path)
    docx_text = "\n".join(paragraph.text for paragraph in event_doc.paragraphs)
    evidence_index = next(index for index, paragraph in enumerate(event_doc.paragraphs) if paragraph.text.startswith("Evidence Ref"))

    for rendered in (html_text, docx_text):
        assert "五条历史事件命中该规则" in rendered
        assert "sample-event-1" in rendered
        assert "sample-event-3" in rendered
        assert "其余 2 条样本保留在 Dataset 原始证据中" in rendered
        assert "events/event.ndjson#host-1/HIST-DEEP-001" in rendered
        assert "sample-event-4" not in rendered
        assert "private-moid-1" not in rendered
        assert "resolved_entities" not in rendered
    assert event_doc.paragraphs[evidence_index].paragraph_format.keep_with_next
    assert event_doc.paragraphs[evidence_index + 1].text.startswith("来源：")
    assert event_doc.paragraphs[evidence_index + 1].paragraph_format.keep_with_next
    assert "sample-event-4" not in payload_text
    assert "private-moid-1" not in payload_text
    assert event_doc.paragraphs[evidence_index].paragraph_format.keep_with_next
    assert event_doc.paragraphs[evidence_index + 1].text.startswith("来源：")
    assert event_doc.paragraphs[evidence_index + 1].paragraph_format.keep_with_next

    many_evidence = [
        evidence.model_copy(
            update={
                "ref": f"ev:event/host-{index}/sample",
                "entity": DeepEntity(type="HostSystem", stable_id=f"host-{index}", display_ref=f"esx-{index}"),
            }
        )
        for index in range(12)
    ]
    many_finding = finding.model_copy(update={"finding_id": "f:event-summary-many", "evidence": many_evidence})
    many_report = DeepReport(findings=[many_finding])
    many_html_path = DeepHtmlReportBuilder().render(many_report, tmp_path / "event-summary-many.html", dataset)
    many_docx_path = DeepWordReportBuilder().render(many_report, tmp_path / "event-summary-many.docx", dataset)
    many_html = many_html_path.read_text(encoding="utf-8")
    many_tables = Document(many_docx_path).tables

    assert 'class="evidence-table"' in many_html
    assert "esx-11" in many_html
    assert "ev:event/host-11/sample" in many_html
    evidence_table = next(table for table in many_tables if len(table.rows) == 13)
    assert evidence_table.rows[0].cells[0].text == "对象"
    assert evidence_table.rows[-1].cells[0].text == "esx-11"

    uplink_finding = finding.model_copy(update={"finding_id": "f:uplink-table", "rule_id": "NET-DEEP-003", "evidence": many_evidence[:5]})
    uplink_report = DeepReport(findings=[uplink_finding])
    uplink_html_path = DeepHtmlReportBuilder().render(uplink_report, tmp_path / "uplink-table.html", dataset)
    uplink_docx_path = DeepWordReportBuilder().render(uplink_report, tmp_path / "uplink-table.docx", dataset)
    assert 'class="evidence-table"' in uplink_html_path.read_text(encoding="utf-8")
    assert any(len(table.rows) == 6 for table in Document(uplink_docx_path).tables)
