"""Word supplement: retain one evidence table without losing unique fields."""
from types import SimpleNamespace

import pytest
from docx import Document

from vstacklens.reports.docx_report import DocxReportEngine, RiskGroup


@pytest.mark.parametrize(
    ("headers", "values"),
    [
        (["所在虚拟机", "快照名称", "创建时间/存在时长", "快照链深度"],
         ["vm-a", "backup-before-upgrade", "2026-10-01T08:00:00+08:00", "4"]),
        (["虚拟机名称", "ISO镜像"],
         ["vm-a", "[Shared-NAS] ISO/ubuntu-24.04.4-live-server-amd64.iso"]),
        (["虚拟机", "Datastore", "所属主机或集群", "影响说明"],
         ["vm-a", "Local-62", "cluster-a", "迁移前先确认应用依赖，保留备份。"]),
    ],
)
def test_original_detail_table_is_rendered_once_with_its_unique_evidence(monkeypatch, headers, values):
    shared = {
        "current": "保存的巡检结论。", "impact": "", "remediation": "复核对象后处理。",
        "steps": ["确认保存的对象信息后处理，再验证结果。"],
        "objects": [{"name": "vm-a", "level": "P3", "location": "cluster-a", "basis": "新增重复表标识"}],
        "technical": {"headers": headers, "rows": [values]},
        "table_notes": ["vm-a：先确认设备要求，保留必要的配置。"],
    }
    engine = DocxReportEngine()
    monkeypatch.setattr(engine, "_shared_problem", lambda *_: shared)
    finding = SimpleNamespace(object_name="vm-a", risk_level="P3")
    group = RiskGroup("P3", "虚拟机", "保存的问题", [finding], "VSL-VM-001")
    document = Document()
    engine._risk_recommendations(document, None, [group])
    assert len(document.tables) == 1
    text = "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert all(value in text for value in values)
    assert "新增重复表标识" not in text
    assert [paragraph.text for paragraph in document.paragraphs].count(shared["table_notes"][0]) == 1


def test_category_without_legacy_table_keeps_the_only_object_evidence(monkeypatch):
    shared = {
        "current": "2 台虚拟机存在资源配置问题。", "impact": "", "remediation": "按对象用途确认。",
        "steps": ["调整资源限制，可在线修改，不用关机。"],
        "technical": None,
        "objects": [
            {"name": "vm-a", "level": "P3", "location": "cluster-a", "basis": "CPU 限制：3600 MHz。"},
            {"name": "vm-b", "level": "P3", "location": "cluster-b", "basis": "内存预留：16384 MB。对象处理要求：先核对设备要求。"},
        ],
    }
    engine = DocxReportEngine()
    monkeypatch.setattr(engine, "_shared_problem", lambda *_: shared)
    findings = [SimpleNamespace(object_name=name, risk_level="P3") for name in ("vm-a", "vm-b")]
    group = RiskGroup("P3", "虚拟机", "资源配置", findings, "VSL-VM-010")
    document = Document()
    engine._risk_recommendations(document, None, [group])
    assert len(document.tables) == 1
    text = "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert "CPU 限制：3600 MHz。" in text
    assert "内存预留：16384 MB。对象处理要求：先核对设备要求。" in text


def test_cover_omits_delimiters_for_empty_environment_or_priority_fields(monkeypatch):
    monkeypatch.setattr("vstacklens.reports.docx_report.delivery_overview", lambda _: {
        "environment": "", "scope": "192.0.2.111", "conclusion": "需关注", "explanation": "保存的检查结果。",
        "priorities": [{"level": "P2", "title": "", "objects": 1}], "limitations": [],
        "recorded_results": {"available": True, "failed_levels": {"P2": 1}},
    })
    report = SimpleNamespace(
        customer_info=SimpleNamespace(customer_name="测试客户", vcenter="192.0.2.111"),
        report_info=SimpleNamespace(generated_at="2026-10-22T00:00:00+08:00"), product_modules=[],
    )
    document = Document()
    DocxReportEngine()._cover(document, report)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert "环境与范围：192.0.2.111" in paragraphs
    assert "P2 · 1 个对象" in paragraphs
    assert all("： ·" not in text and "·  ·" not in text for text in paragraphs)


def test_old_customer_recommendation_uses_shared_business_time_wording():
    assert "维护窗口" not in DocxReportEngine()._clean("建议在维护窗口完成处理并验证结果。")
