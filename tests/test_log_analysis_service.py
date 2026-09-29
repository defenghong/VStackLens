from __future__ import annotations

import io
import json
import logging
import sqlite3
import tarfile
import threading
import time
import urllib.error
import zipfile
from pathlib import Path

import pytest
from docx import Document

import vstacklens.application.cloud_log_diagnosis as cloud_module
from vstacklens.application import InspectionService, LogAnalysisConfig
from vstacklens.application.cloud_log_diagnosis import (
    DEFAULT_CLOUD_MODEL,
    DEFAULT_CLOUD_PROVIDER,
    MODEL_PROTOCOL,
    CloudModelClient,
    CloudModelConfig,
    CloudModelResult,
)
from vstacklens.application.logging_config import LOGGER_NAME, configure_runtime_logging
import vstacklens.application.log_analysis_service as log_analysis_module
from vstacklens.application.log_analysis_service import LogAnalysisCloudRequiredError, LogAnalysisService, log_analysis_display_stats
import vstacklens.db.connection as db_connection
from vstacklens.db.connection import connect, init_db


def _support_bundle_zip(path: Path) -> Path:
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr(
            "var/log/vmware/vpxd/vpxd.log",
            "\n".join(
                [
                    "2026-06-10T10:00:00Z warning connection refused while connecting hostd",
                    "2026-06-10T10:01:00Z error authentication failed for session token",
                    "2026-06-10T10:02:00Z warning certificate handshake failed",
                ]
            ),
        )
        bundle.writestr(
            "var/run/log/vmkernel.log",
            "\n".join(
                [
                    "2026-06-10T10:03:00Z cpu0:2097152 storage latency detected on datastore naa.6000",
                    "2026-06-10T10:04:00Z cpu0:2097152 vmnic2 link down packet loss detected",
                ]
            ),
        )
        bundle.writestr("manifest.txt", "support bundle manifest")
    return path


def _nested_support_bundle_zip(path: Path) -> Path:
    tar_buffer = io.BytesIO()
    with tarfile.open(fileobj=tar_buffer, mode="w:gz") as tar:
        data = b"2026-06-10T10:05:00Z warning vmnic3 link down packet loss detected\n"
        info = tarfile.TarInfo("var/run/log/vmkernel.log")
        info.size = len(data)
        tar.addfile(info, io.BytesIO(data))
        manifest = b"nested support bundle manifest"
        manifest_info = tarfile.TarInfo("manifest.txt")
        manifest_info.size = len(manifest)
        tar.addfile(manifest_info, io.BytesIO(manifest))
    tar_buffer.seek(0)
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("host-support.tgz", tar_buffer.getvalue())
    return path


def _deep_nested_support_bundle_zip(path: Path, depth: int = 6) -> Path:
    payload = b"2026-06-10T10:05:00Z warning vmnic3 link down packet loss detected\n"
    for level in range(depth):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as nested:
            name = "var/run/log/vmkernel.log" if level == 0 else f"nested-{level}.zip"
            nested.writestr(name, payload)
        payload = buffer.getvalue()
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("host-support.zip", payload)
    return path


def _corrupt_nested_support_bundle_zip(path: Path) -> Path:
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("host-support.tgz", b"not a valid tar payload")
        bundle.writestr("var/run/log/vmkernel.log", "2026-06-10T10:05:00Z warning vmnic3 link down packet loss detected\n")
    return path


def _sensitive_support_bundle_zip(path: Path) -> Path:
    sensitive_line = (
        "2026-06-10T10:01:00Z error authentication failed "
        "password=secret-password passwd=secret-passwd pwd=secret-pwd "
        "token=token-123 access_token=access-123 refresh_token=refresh-123 "
        "api_key=key-123 secret=secret-value sessionId=session-123 cookie=session-cookie "
        "Authorization: Bearer bearer-token Bearer bare-token Basic basic-token"
    )
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("var/log/vmware/vpxd/vpxd.log", sensitive_line)
    return path


def _tar_add_text(tar: tarfile.TarFile, path: str, text: str) -> None:
    data = text.encode("utf-8")
    info = tarfile.TarInfo(path)
    info.size = len(data)
    tar.addfile(info, io.BytesIO(data))


def _snapshot_support_bundle_zip(path: Path) -> Path:
    vm_dir = "esx-host/vmfs/volumes/630c602f-6216a9aa-3c03-78ac447b3b2e/B-longxiagm-133.242"
    vsan_dir = "esx-host/vmfs/volumes/vsan:525731262becee94-187a12b3670f1951/3e6aee69-1a70-96ba-d4e6-3868dd1ad2e0"
    tar_buffer = io.BytesIO()
    with tarfile.open(fileobj=tar_buffer, mode="w:gz") as tar:
        _tar_add_text(
            tar,
            f"{vm_dir}/vmware.log",
            "\n".join(
                [
                    '2026-06-04T02:13:53.271Z No(00) vcpu-0 - ConfigDB: Setting softPowerOff = "TRUE"',
                    "2026-06-04T02:13:53.334Z In(05) vcpu-0 - VMX: Issuing power-off request...",
                    "2026-06-04T02:13:53.334Z In(05) vcpu-0 - SnapshotVMX_ConsolidateCancel: Requesting snapshot consolidate cancel.",
                    "2026-06-04T02:29:13.924Z In(05) worker-94624005 - DISKLIB-VMFS_SPARSE : Canceling vmfs combine.",
                    "2026-06-04T02:29:13.935Z In(05) worker-94624005 - DISKLIB-LIB_CHAINMODIFY : Failed to combine : Operation was cancelled (33).",
                    "2026-06-04T02:29:15.214Z In(05) vcpu-0 - ConsolidateItemComplete: Cancelling in-progress online consolidate.",
                    "2026-06-04T02:29:15.464Z In(05) vcpu-0 - ConsolidateEnd: Snapshot consolidate complete: Cancelled (34).",
                    "2026-06-04T02:29:26.001Z In(05) vmx - Transitioned vmx/execState/val to poweredOff",
                ]
            ),
        )
        _tar_add_text(
            tar,
            f"{vm_dir}/B-longxiagm-133.242.vmsd",
            "\n".join(['snapshot.lastUID = "106"', 'snapshot.needConsolidate = "TRUE"']),
        )
        _tar_add_text(
            tar,
            f"{vm_dir}/B-longxiagm-133.242.vmx",
            "\n".join(
                [
                    'displayName = "B-longxiagm-10.240.7.242"',
                    'scsi0:0.fileName = "B-longxiagm-133.242.vmdk"',
                    f'scsi0:1.fileName = "/vmfs/volumes/vsan:525731262becee94-187a12b3670f1951/3e6aee69-1a70-96ba-d4e6-3868dd1ad2e0/B-longxiagm-10.240.7.242-000003.vmdk"',
                ]
            ),
        )
        _tar_add_text(
            tar,
            "esx-host/var/run/log/vpxa.log",
            "\n".join(
                [
                    "2026-06-06T22:19:28.638Z info vpxa[123] Timed out waiting for task vim.Task:haTask-120-vim.VirtualMachine.removeAllSnapshots-11703068",
                    "2026-06-06T22:19:28.638Z warning vpxa[123] No taskinfo property updates for haTask-120-vim.VirtualMachine.removeAllSnapshots-11703068 in 300000 ms",
                ]
            ),
        )
        descriptors = {
            "B-longxiagm-172.16.133.242.vmdk": ('CID=8521875c', 'parentCID=ffffffff', ""),
            "B-longxiagm-10.240.7.242-000001.vmdk": ('CID=242aa430', 'parentCID=8521875c', 'parentFileNameHint="B-longxiagm-172.16.133.242.vmdk"'),
            "B-longxiagm-10.240.7.242-000002.vmdk": ('CID=6bb648ab', 'parentCID=242aa430', 'parentFileNameHint="B-longxiagm-10.240.7.242-000001.vmdk"'),
            "B-longxiagm-10.240.7.242-000003.vmdk": ('CID=754b316c', 'parentCID=6bb648ab', 'parentFileNameHint="B-longxiagm-10.240.7.242-000002.vmdk"'),
        }
        for name, lines in descriptors.items():
            _tar_add_text(tar, f"{vsan_dir}/{name}", "\n".join(line for line in lines if line))
    tar_buffer.seek(0)
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("host-support.tgz", tar_buffer.getvalue())
    return path


def _migration_network_support_bundle_zip(path: Path) -> Path:
    tar_buffer = io.BytesIO()
    with tarfile.open(fileobj=tar_buffer, mode="w:gz") as tar:
        _tar_add_text(
            tar,
            "vc/var/log/vmware/vpxd/vpxd.log",
            "\n".join(
                [
                    "2026-06-02T18:33:10Z info vpxd[1220] Task started: MigrateVM_Task for vm esx-host-01 from esx01 to esx02",
                    "2026-06-02T18:35:43Z warning vpxd[1220] Relocate completed for esx-host-01, waiting for network backing update on dvPort 318",
                ]
            ),
        )
        _tar_add_text(
            tar,
            "esx02/var/run/log/hostd.log",
            "\n".join(
                [
                    "2026-06-02T18:36:01Z info hostd[2100] ReconfigureVM_Task esx-host-01 ethernet0 connectable.connected = false",
                    "2026-06-02T18:36:08Z info hostd[2100] esx-host-01 Vigor transport disconnected during post-migration state transition",
                    "2026-06-02T18:36:20Z info hostd[2100] esx-host-01 Vigor transport reconnected after vNIC connect refresh",
                    "2026-06-02T18:36:31Z info hostd[2100] ReconfigureVM_Task esx-host-01 ethernet0 connectable.connected = true",
                ]
            ),
        )
        _tar_add_text(
            tar,
            "esx02/var/run/log/vmkernel.log",
            "\n".join(
                [
                    "2026-06-02T18:36:05Z cpu7:2097152 net: dvPort 318 portgroup PG-App VLAN 133 uplink vmnic2 dropped packets after migration",
                    "2026-06-02T18:37:05Z cpu7:2097152 net: vmnic2 link up for portgroup PG-App",
                ]
            ),
        )
        _tar_add_text(
            tar,
            "esx02/vmfs/volumes/datastore/esx-host-01/esx-host-01.vmx",
            "\n".join(
                [
                    'displayName = "esx-host-01"',
                    'ethernet0.virtualDev = "vmxnet3"',
                    'ethernet0.networkName = "PG-App"',
                    'ethernet0.connectable.connected = "TRUE"',
                    'ethernet0.connectable.startConnected = "TRUE"',
                ]
            ),
        )
        _tar_add_text(
            tar,
            "esx02/commands/esxcli_network.txt",
            "Portgroup PG-App VLAN 133 uplink vmnic2 active\n",
        )
        _tar_add_text(
            tar,
            "esx02/commands/vmkerrcode_-l.txt",
            "\n".join(
                [
                    "VMK_ENETUNREACH Network unreachable",
                    "VMK_ENETRESET Network dropped connection on reset",
                    "VMK_VLAN_FILTERED Pkts dropped because of VLAN mismatch",
                ]
            ),
        )
        _tar_add_text(tar, "esx02/etc/vmware/usb.ids", "ffff Hardware vendor dictionary\n")
        _tar_add_text(tar, "esx02/etc/vmware/pci.ids", "ffff PCI vendor dictionary\n")
    tar_buffer.seek(0)
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("vmware-support-hosts.tgz", tar_buffer.getvalue())
    return path


def _arp_conflict_support_bundle_zip(path: Path) -> Path:
    conflict_line = (
        "2026-06-14T09:32:12Z cpu12:2097152)WARNING: arp: "
        "00:50:56:6d:f7:1b is using my IP address 10.240.6.187 on vmk2!"
    )
    tar_buffer = io.BytesIO()
    with tarfile.open(fileobj=tar_buffer, mode="w:gz") as tar:
        _tar_add_text(
            tar,
            "esx-atcdi-node-0115/var/run/log/vmkernel-172.16.102.187.log",
            "\n".join([conflict_line for _ in range(586)]),
        )
        _tar_add_text(
            tar,
            "esx-atcdi-node-0115/var/run/log/hostd.log",
            "2026-06-14T09:31:45Z info hostd[2100] Network health check requested for vmk2\n",
        )
        _tar_add_text(
            tar,
            "esx-atcdi-node-0115/commands/esxcli_network.txt",
            "vmk2 10.240.6.187 VLAN 103 uplink vmnic4\n",
        )
    tar_buffer.seek(0)
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("vmware-support-hosts.tgz", tar_buffer.getvalue())
    return path


def _arp_conflict_variant_support_bundle_zip(path: Path) -> Path:
    lines = [
        "2026-06-14T09:32:12Z cpu12:2097152)WARNING: arp: 00:50:56:6d:f7:1b is using my IP address 10.240.6.187 on vmk2",
        "2026-06-14T09:32:13Z cpu12:2097152)WARNING: arp: 00:50:56:6d:f7:1b is using my IP address 10.240.6.187 on vmk2.",
    ]
    tar_buffer = io.BytesIO()
    with tarfile.open(fileobj=tar_buffer, mode="w:gz") as tar:
        _tar_add_text(
            tar,
            "esx-atcdi-node-0115/var/run/log/vmkernel-172.16.102.187.log",
            "\n".join(lines),
        )
    tar_buffer.seek(0)
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("vmware-support-hosts.tgz", tar_buffer.getvalue())
    return path


def _cp932_support_bundle_zip(path: Path) -> Path:
    message = "2026-06-10T10:03:00Z warning ホスト接続エラー\n".encode("cp932")
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("var/log/vmware/vpxd/vpxd.log", message)
    return path


def _snapshot_root_name_edge_bundle_zip(path: Path) -> Path:
    vm_dir = "esx-host/vmfs/volumes/datastore1/SQLVM"
    tar_buffer = io.BytesIO()
    with tarfile.open(fileobj=tar_buffer, mode="w:gz") as tar:
        _tar_add_text(
            tar,
            f"{vm_dir}/SQLVM.vmx",
            "\n".join(
                [
                    'displayName = "SQLVM"',
                    'scsi0:0.fileName = "SQL-0001.vmdk"',
                    'scsi0:1.fileName = "SQL-0001-000001.vmdk"',
                ]
            ),
        )
        descriptors = {
            "SQL-0001.vmdk": ('CID=11111111', 'parentCID=ffffffff', ""),
            "SQL-0001-000001.vmdk": ('CID=22222222', 'parentCID=11111111', 'parentFileNameHint="SQL-0001.vmdk"'),
        }
        for name, lines in descriptors.items():
            _tar_add_text(tar, f"{vm_dir}/{name}", "\n".join(line for line in lines if line))
    tar_buffer.seek(0)
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr("host-support.tgz", tar_buffer.getvalue())
    return path

def _static_vsan_support_bundle_zip(path: Path) -> Path:
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.writestr(
            "esx01/commands/vmkfstools_-P.txt",
            "\n".join(
                [
                    "vSAN datastore static inventory output",
                    "Object UUID: 1234-5678",
                    "Policy: hostFailuresToTolerate = 1",
                ]
            ),
        )
        bundle.writestr(
            "esx01/commands/dump-vmdk-rdm-info.txt",
            "\n".join(
                [
                    "Static VMDK mapping output",
                    "esx-host-01.vmdk -> naa.6000",
                ]
            ),
        )
        bundle.writestr("manifest.txt", "support bundle manifest")
    return path


def _docx_text(path: Path) -> str:
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            text += "\n" + "\t".join(cell.text for cell in row.cells)
    return text


def _flush_runtime_log_handlers() -> None:
    for handler in logging.getLogger(LOGGER_NAME).handlers:
        handler.flush()


def _close_runtime_log_handler(log_path: Path) -> None:
    resolved = log_path.resolve()
    logger = logging.getLogger(LOGGER_NAME)
    for handler in list(logger.handlers):
        if Path(getattr(handler, "baseFilename", "")).resolve() == resolved:
            logger.removeHandler(handler)
            handler.close()


NO_MODEL_FORBIDDEN_TERMS = (
    "当前判断",
    "诊断结论",
    "是否可以完全定位",
    "置信度与边界",
    "当前能够确认",
    "当前不能确认",
    "最终根因",
    "根因判断",
    "证据推理",
    "模型辅助分析",
    "AI 根因分析",
    "大模型判断",
)


def _assert_no_model_log_check_report(combined: str) -> None:
    assert "VStackLens 日志粗排查报告" in combined
    assert "1. 排查概览" in combined
    assert "1.1 客户问题描述" in combined
    assert "1.2 日志包处理结果" in combined
    assert "1.3 当前排查范围" in combined
    assert "2. 相关日志线索" in combined
    assert "3. 初步排查方向" in combined
    assert "4. 附录" in combined
    for term in NO_MODEL_FORBIDDEN_TERMS:
        assert term not in combined


def _assert_rule_only_engine(
    engine: dict,
    *,
    provider: str = DEFAULT_CLOUD_PROVIDER,
    model: str = DEFAULT_CLOUD_MODEL,
    fallback_reason: str = "",
) -> None:
    expected = {
        "name": "rule_based_log_diagnosis_with_optional_cloud",
        "mode": "rule_based",
        "model_used": False,
        "model_source": "rule_only",
        "model_protocol": MODEL_PROTOCOL,
        "fallback_reason": fallback_reason,
        "cloud_provider": provider,
        "model_name": model,
    }
    for key, value in expected.items():
        assert engine.get(key) == value
    assert engine.get("model_quality", "") in {"", "failed"}
    assert "model_quality_reason" in engine
    if not fallback_reason:
        assert engine.get("cloud_assist_enabled") is False
        assert engine.get("api_call_attempted") is False
        assert engine.get("api_response_received") is False
        assert engine.get("cloud_status") == "not_enabled"


def _assert_cloud_engine(engine: dict, *, provider: str = DEFAULT_CLOUD_PROVIDER, model: str = DEFAULT_CLOUD_MODEL) -> None:
    expected = {
        "name": "rule_based_log_diagnosis_with_optional_cloud",
        "mode": "cloud_assisted",
        "model_used": True,
        "model_source": "cloud",
        "model_protocol": MODEL_PROTOCOL,
        "fallback_reason": "",
        "cloud_provider": provider,
        "model_name": model,
        "customer_report_mode": "cloud_assisted",
    }
    for key, value in expected.items():
        assert engine.get(key) == value
    assert engine.get("model_quality") == "accepted"
    assert engine.get("model_quality_reason")
    assert engine.get("api_call_attempted") is True
    assert engine.get("api_response_received") is True
    assert engine.get("cloud_status") == "accepted"


def _assert_cloud_triage_engine(engine: dict, *, provider: str = DEFAULT_CLOUD_PROVIDER, model: str = DEFAULT_CLOUD_MODEL) -> None:
    expected = {
        "name": "rule_based_log_diagnosis_with_optional_cloud",
        "mode": "cloud_triage",
        "model_used": True,
        "model_source": "cloud",
        "model_protocol": MODEL_PROTOCOL,
        "fallback_reason": "",
        "cloud_provider": provider,
        "model_name": model,
        "customer_report_mode": "rule_based",
    }
    for key, value in expected.items():
        assert engine.get(key) == value
    assert engine.get("model_quality") == "downgraded"
    assert engine.get("model_quality_reason")
    assert engine.get("api_call_attempted") is True
    assert engine.get("api_response_received") is True
    assert engine.get("cloud_status") == "responded_downgraded"


def _assert_cloud_rejected_engine(engine: dict, *, provider: str = DEFAULT_CLOUD_PROVIDER, model: str = DEFAULT_CLOUD_MODEL) -> None:
    expected = {
        "name": "rule_based_log_diagnosis_with_optional_cloud",
        "mode": "cloud_rejected",
        "model_used": True,
        "model_source": "cloud",
        "model_protocol": MODEL_PROTOCOL,
        "cloud_provider": provider,
        "model_name": model,
        "customer_report_mode": "rule_based",
    }
    for key, value in expected.items():
        assert engine.get(key) == value
    assert engine.get("model_quality") == "rejected"
    assert engine.get("model_quality_reason")
    assert engine.get("fallback_reason")
    assert engine.get("api_call_attempted") is True
    assert engine.get("api_response_received") is True
    assert engine.get("cloud_status") == "responded_rejected"


def _cloud_failure_payload(tmp_path: Path) -> tuple[Path, dict, str]:
    report_dirs = sorted((tmp_path / "reports").glob("log-analysis-*"))
    assert len(report_dirs) == 1
    report_dir = report_dirs[0]
    assert not (report_dir / "index.html").exists()
    assert not (report_dir / "VStackLens-Log-Analysis-Report.docx").exists()
    assert (report_dir / "log_analysis_payload.json").exists()
    assert (report_dir / "diagnosis.json").exists()
    debug_path = report_dir / "model_analysis_failure_debug.json"
    assert debug_path.exists()
    payload_text = (report_dir / "log_analysis_payload.json").read_text(encoding="utf-8")
    debug = json.loads(debug_path.read_text(encoding="utf-8"))
    payload = json.loads(payload_text)
    assert payload["metadata"]["report_title"] == "模型分析失败诊断信息"
    assert payload["metadata"]["analysis_status"] == "model_analysis_failed"
    assert payload["summary"]["model_analysis_failed"] is True
    assert payload["diagnosis"]["model_analysis_failed"] is True
    assert debug["title"] == "模型分析失败诊断信息"
    assert debug["failure_reason"] == payload["diagnosis"]["model_analysis_failure_reason"]
    return report_dir, payload, payload_text + "\n" + json.dumps(debug, ensure_ascii=False)


def _cloud_fallback_result_payload(result) -> tuple[dict, str]:
    payload_text = (result.report_dir / "log_analysis_payload.json").read_text(encoding="utf-8")
    diagnosis_text = (result.report_dir / "diagnosis.json").read_text(encoding="utf-8")
    html_text = result.html_path.read_text(encoding="utf-8")
    doc_text = _docx_text(result.docx_path)
    payload = json.loads(payload_text)
    combined = payload_text + "\n" + diagnosis_text + "\n" + html_text + "\n" + doc_text
    return payload, combined


def _cloud_content(current_judgement: str | None = None, evidence_chain: list[dict] | None = None) -> dict:
    return {
        "mode": "diagnosis",
        "can_determine_root_cause": False,
        "current_judgement": current_judgement
        or "当前日志显示 esx-host-01 在相关时间段存在 MigrateVM_Task / Relocate 迁移记录，hostd 日志中出现 Vigor transport disconnected / reconnected 与 ethernet0 connectable 状态切换，vmkernel 日志中也出现 dvPort 318、PG-App、VLAN 133、uplink vmnic2 dropped packets 相关线索。结合客户描述“在线迁移后网络不通，重连虚拟网卡恢复”，当前更倾向于优先排查迁移后虚拟网卡连接状态刷新、目标主机网络路径、端口组 / DVS 绑定或 Guest OS 网卡状态未及时恢复方向，但当前 support bundle 尚不能单独证明最终根因。",
        "likely_causes": ["连接状态刷新异常", "目标侧网络路径或端口绑定需要复核"],
        "evidence_chain": evidence_chain
        if evidence_chain is not None
        else [
            {
                "strength": "中",
                "file": "host-support.tgz!vc/var/log/vmware/vpxd/vpxd.log",
                "location": "lines 1-2",
                "message": "2026-06-02T18:33:10Z Task started: MigrateVM_Task for vm esx-host-01 from esx01 to esx02\n2026-06-02T18:35:43Z Relocate completed for esx-host-01, waiting for network backing update on dvPort 318",
                "relevance": "该证据说明客户问题对象在相关时间段存在迁移和网络 backing 更新线索。",
            },
            {
                "strength": "强",
                "file": "host-support.tgz!esx02/var/run/log/hostd.log",
                "location": "lines 1-4",
                "message": "2026-06-02T18:36:08Z esx-host-01 Vigor transport disconnected during post-migration state transition\n2026-06-02T18:36:20Z esx-host-01 Vigor transport reconnected after vNIC connect refresh",
                "relevance": "该证据与客户描述的迁移后网络不通、重连虚拟网卡后恢复现象直接相关。",
            },
        ],
        "evidence_reasoning": ["vCenter 迁移记录、hostd 中 Vigor transport 断开/重连及虚拟网卡连接状态变化可以共同支撑优先排查迁移后虚拟网卡连接状态和目标侧网络路径，但仍缺少故障时 Guest OS 与网络设备侧闭环证据。"],
        "can_confirm": ["当前日志包能够确认 esx-host-01 存在与客户问题相关的迁移和虚拟网卡/网络状态变化线索。"],
        "cannot_confirm": ["当前 support bundle 尚不能单独证明最终根因。"],
        "missing_materials": ["故障发生时间点", "vCenter Task/Event 记录", "现场网络状态截图"],
        "recommended_actions": ["结合故障时间复核任务事件", "核对相关对象的配置和运行状态"],
        "evidence_quality": "medium",
        "overclaiming_risk": "当前证据仍缺少现场状态闭环，存在过度判断风险。",
        "confidence": "中",
    }


def _arp_cloud_content() -> dict:
    return {
        "mode": "diagnosis",
        "can_determine_root_cause": False,
        "confidence": "中",
        "current_judgement": (
            "当前日志可以确认主机 esx-atcdi-node-0115 的 vmk2 接口反复检测到 IP 地址 "
            "10.240.6.187 被 MAC 00:50:56:6d:f7:1b 使用，属于 VMkernel 层面的 ARP/IP 冲突告警。"
            "该问题可能导致 vmk2 所在网络通信异常或间歇性不稳定。结合客户网络不通类问题，"
            "当前更倾向于优先排查重复 IP、端口组 / VLAN 或目标网络路径方向；但仅凭当前 support bundle "
            "不能确认冲突 MAC 对应的具体设备身份。"
        ),
        "most_likely_cause": "vmk2 所在网络存在重复 IP 或异常 ARP 冲突，需要结合交换机和 vCenter 配置确认冲突 MAC 对应设备。",
        "likely_causes": ["VMkernel 适配器重复 IP", "虚拟机或物理设备配置了相同 IP", "相关 VLAN / 端口组配置需要复核"],
        "evidence_chain": [
            {
                "strength": "强",
                "file": "vmware-support-hosts.tgz!esx-atcdi-node-0115/var/run/log/vmkernel-172.16.102.187.log",
                "location": "lines 1-586",
                "message": "arp: 00:50:56:6d:f7:1b is using my IP address 10.240.6.187 on vmk2!",
                "relevance": "该日志反复出现 586 次，直接说明 vmk2 检测到自身 IP 10.240.6.187 被 00:50:56:6d:f7:1b 使用。",
            }
        ],
        "evidence_reasoning": [
            "vmkernel 日志中的 ARP/IP 冲突告警属于现场运行日志，不是静态字典或配置清单；vmk2、10.240.6.187 和 00:50:56:6d:f7:1b 与客户网络不通类现象存在直接排查关系。",
            "该证据能够证明主机侧检测到重复 IP/ARP 冲突，但不能单独识别冲突 MAC 属于哪台虚拟机、物理服务器或网络设备。",
        ],
        "can_confirm": [
            "主机 esx-atcdi-node-0115 的 vmk2 接口检测到 10.240.6.187 与 00:50:56:6d:f7:1b 存在 ARP/IP 冲突。",
            "该告警在当前日志片段中重复出现，具备优先排查价值。",
        ],
        "cannot_confirm": [
            "当前 support bundle 不能确认冲突 MAC 对应的具体设备身份。",
            "当前日志不能单独确认业务网络不通的最终根因，只能确认 VMkernel 层面存在 ARP/IP 冲突强线索。",
        ],
        "missing_materials": [
            "vCenter 中 VMkernel 适配器配置截图",
            "10.240.6.187 的规划归属",
            "交换机 MAC 地址表 / ARP 表",
            "相关 VLAN / 端口组配置",
            "是否有虚拟机或物理设备配置了相同 IP",
        ],
        "recommended_actions": [
            "优先核对 10.240.6.187 是否被其他 VMkernel、虚拟机、物理服务器或网络设备重复配置。",
            "结合交换机 MAC 地址表和 vCenter 网络配置定位 00:50:56:6d:f7:1b 的归属。",
        ],
        "evidence_quality": "strong",
        "overclaiming_risk": "当前证据可以确认 ARP/IP 冲突告警，但不能直接确认冲突设备身份或最终业务影响范围。",
    }


def _patch_cloud_success(monkeypatch: pytest.MonkeyPatch, content: dict | None = None, plan: dict | None = None) -> list[dict]:
    calls: list[dict] = []

    def fake_complete_json(self: CloudModelClient, prompt: str) -> CloudModelResult:
        calls.append({"stage": "planning", "config": self.config, "prompt": prompt})
        return CloudModelResult(
            ok=True,
            content=plan
            or {
                "requests": [
                    {
                        "file": "host-support.tgz!esx02/var/run/log/hostd.log",
                        "component": "hostd",
                        "keywords": ["esx-host-01", "Vigor transport", "connectable", "disconnected"],
                        "context_lines": 20,
                        "reason": "查看迁移后虚拟网卡连接状态和 Vigor transport 状态变化。",
                    },
                    {
                        "file": "host-support.tgz!esx02/var/run/log/vmkernel.log",
                        "component": "vmkernel",
                        "keywords": ["dvPort", "portgroup", "VLAN", "uplink", "dropped"],
                        "context_lines": 20,
                        "reason": "查看目标主机网络路径、端口组和上行链路线索。",
                    },
                ],
                "stop_reason": "已请求与客户问题对象、迁移和网络状态相关的日志片段。",
            },
            elapsed_ms=45,
            response_received=True,
        )

    def fake_diagnose(self: CloudModelClient, prompt: str) -> CloudModelResult:
        calls.append({"stage": "diagnosis", "config": self.config, "prompt": prompt})
        return CloudModelResult(ok=True, content=content or _cloud_content(), elapsed_ms=123, response_received=True)

    monkeypatch.setattr(CloudModelClient, "complete_json", fake_complete_json)
    monkeypatch.setattr(CloudModelClient, "diagnose", fake_diagnose)
    return calls


def test_log_analysis_requires_support_bundle_and_problem_only(tmp_path: Path) -> None:
    service = LogAnalysisService()
    empty_errors = service.validate_config(LogAnalysisConfig(problem_description="主机连接偶发超时"))

    assert "请导入 VMware support bundle 日志包。" in empty_errors

    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")
    errors = service.validate_config(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="客户反馈管理界面偶发超时",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )

    assert errors == []


def test_log_analysis_reads_nested_vmware_support_bundle_archives(tmp_path: Path) -> None:
    bundle = _nested_support_bundle_zip(tmp_path / "vmware-support.zip")
    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="网络链路异常，需要分析 support bundle。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )

    assert result.summary["total_files"] == 2
    assert result.summary["log_files"] == 2
    assert result.summary["evidence_count"] >= 1
    assert any("host-support.tgz!var/run/log/vmkernel.log" in evidence["file"] for finding in result.findings for evidence in finding["evidence"])


def test_log_analysis_limits_deep_nested_zip_scan_without_crashing(tmp_path: Path) -> None:
    bundle = _deep_nested_support_bundle_zip(tmp_path / "vmware-support.zip")

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="网络链路异常，需要分析 support bundle。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )

    assert result.summary["total_files"] >= 1
    assert result.summary["log_files"] == 0
    assert result.summary["evidence_count"] == 0


def test_log_analysis_skips_corrupt_nested_tar_without_crashing(tmp_path: Path) -> None:
    bundle = _corrupt_nested_support_bundle_zip(tmp_path / "vmware-support.zip")

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="网络链路异常，需要分析 support bundle。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )

    assert result.summary["total_files"] >= 2
    assert result.summary["log_files"] >= 1
    assert result.summary["evidence_count"] >= 1


def test_log_analysis_decodes_cp932_logs() -> None:
    service = LogAnalysisService()

    assert "ホスト接続エラー" == service._decode_text("ホスト接続エラー".encode("cp932"))


def test_log_analysis_generates_independent_html_docx_and_database_rows(tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")
    db_path = tmp_path / "log-analysis.db"
    report_root = tmp_path / "reports"
    progress_messages: list[str] = []

    result = InspectionService(config_path=tmp_path / "desktop_config.json").run_log_analysis(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="客户反馈 vCenter 管理界面偶发超时，部分主机日志存在链路告警。",
            customer_name="测试客户",
            report_title="VMware 日志分析报告",
            db_path=db_path,
            report_output_dir=report_root,
        ),
        progress_callback=lambda progress: progress_messages.append(progress.message),
    )

    assert result.html_path.exists()
    assert result.docx_path.exists()
    assert result.summary["log_files"] == 3
    assert result.summary["evidence_count"] >= 4
    assert {"读取 support bundle 文件清单", "分析日志关键特征", "生成 HTML 日志分析报告", "生成 Word 日志分析报告", "日志分析完成"} <= set(progress_messages)

    html_text = result.html_path.read_text(encoding="utf-8")
    doc_text = "\n".join(paragraph.text for paragraph in Document(result.docx_path).paragraphs)
    combined = html_text + "\n" + doc_text
    _assert_no_model_log_check_report(combined)
    assert "不连接 vCenter" in combined
    assert "不写入虚拟化健康巡检结果" in combined
    assert "secret-password" not in combined

    with connect(db_path) as conn:
        assert conn.execute("SELECT COUNT(*) AS c FROM log_analysis_runs").fetchone()["c"] == 1
        report_types = {row["report_type"] for row in conn.execute("SELECT report_type FROM log_analysis_reports").fetchall()}
        assert report_types == {"log_analysis_html", "log_analysis_docx"}
        assert conn.execute("SELECT COUNT(*) AS c FROM inspection_runs").fetchone()["c"] == 0
        assert conn.execute("SELECT COUNT(*) AS c FROM reports").fetchone()["c"] == 0


def test_log_analysis_preserves_chinese_metadata_without_local_paths(tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")
    db_path = tmp_path / "log-analysis.db"
    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="客户反馈管理界面登录超时，需分析发生时间窗口内的关键日志。",
            customer_name="测试客户",
            report_title="日志分析交付报告",
            db_path=db_path,
            report_output_dir=tmp_path / "reports",
        )
    )

    html_text = result.html_path.read_text(encoding="utf-8")
    document = Document(result.docx_path)
    doc_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )
    for table in document.tables:
        for row in table.rows:
            doc_text += "\n" + "\t".join(cell.text for cell in row.cells)
    combined = html_text + "\n" + doc_text

    assert "测试客户" in combined
    assert "客户反馈管理界面登录超时" in combined
    assert "????" not in combined
    assert str(tmp_path) not in combined
    assert "sqlite" not in combined.lower()
    assert "src/vstacklens" not in combined.replace("\\", "/").lower()


def test_log_analysis_redacts_sensitive_evidence_from_payload_html_and_docx(tmp_path: Path) -> None:
    bundle = _sensitive_support_bundle_zip(tmp_path / "vmware-support.zip")
    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="authentication failed after token refresh",
            customer_name="Security Customer",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )

    payload_text = (result.report_dir / "log_analysis_payload.json").read_text(encoding="utf-8")
    html_text = result.html_path.read_text(encoding="utf-8")
    doc_text = _docx_text(result.docx_path)
    combined = payload_text + "\n" + html_text + "\n" + doc_text

    for raw_value in (
        "secret-password",
        "secret-passwd",
        "secret-pwd",
        "token-123",
        "access-123",
        "refresh-123",
        "key-123",
        "secret-value",
        "session-123",
        "session-cookie",
        "bearer-token",
        "bare-token",
        "basic-token",
    ):
        assert raw_value not in combined

    for redacted in (
        "password=***",
        "passwd=***",
        "pwd=***",
        "token=***",
        "access_token=***",
        "refresh_token=***",
        "api_key=***",
        "secret=***",
        "sessionId=***",
        "cookie=***",
        "Authorization: Bearer ***",
        "Bearer ***",
        "Basic ***",
    ):
        assert redacted in combined


def test_generic_problem_fallback_generates_problem_driven_report(tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")
    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="客户反馈 vCenter 管理界面登录超时，想确认是否和认证或证书有关。",
            customer_name="测试客户",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )

    payload = json.loads((result.report_dir / "log_analysis_payload.json").read_text(encoding="utf-8"))
    diagnosis = payload["diagnosis"]
    combined = result.html_path.read_text(encoding="utf-8") + "\n" + _docx_text(result.docx_path)

    assert diagnosis["scenario"] == "generic_problem_driven_diagnosis"
    _assert_rule_only_engine(diagnosis["diagnosis_engine"])
    assert diagnosis["confidence"] in {"低", "中", "高"}
    assert diagnosis["what_cannot_be_confirmed"]
    assert diagnosis["missing_materials"]
    _assert_no_model_log_check_report(combined)
    assert "3.1" in combined
    assert "3.2" in combined
    assert "2. 关键发现" not in combined
    assert "3. 日志证据摘要" not in combined
    assert "rule_based_log_diagnosis_with_optional_cloud" not in combined


def test_snapshot_consolidation_problem_generates_problem_driven_diagnosis(tmp_path: Path) -> None:
    bundle = _snapshot_support_bundle_zip(tmp_path / "vmware-support.zip")
    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="快照整合失败，想知道什么原因，怎么解决。",
            customer_name="测试客户",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )

    payload = json.loads((result.report_dir / "log_analysis_payload.json").read_text(encoding="utf-8"))
    diagnosis = payload["diagnosis"]
    html_text = result.html_path.read_text(encoding="utf-8")
    doc_text = _docx_text(result.docx_path)
    combined = html_text + "\n" + doc_text
    diagnosis_text = json.dumps(diagnosis, ensure_ascii=False)

    assert diagnosis["scenario"] == "snapshot_consolidation_failure"
    _assert_rule_only_engine(diagnosis["diagnosis_engine"])
    assert diagnosis["can_fully_determine"] is False
    assert diagnosis["affected_objects"]["vm_name"] == "B-longxiagm-10.240.7.242"
    timeline_text = json.dumps(diagnosis.get("timeline", []), ensure_ascii=False)
    _assert_no_model_log_check_report(combined)
    assert "SnapshotVMX_ConsolidateCancel" in timeline_text
    assert "Operation was cancelled (33)" in timeline_text
    assert "Cancelled (34)" in timeline_text
    assert "快照整合失败诊断报告" not in combined
    assert "SnapshotVMX_ConsolidateCancel" in combined
    assert "Operation was cancelled (33)" in timeline_text
    assert "Cancelled (34)" in timeline_text
    assert "-000003.vmdk" in combined
    assert "removeAllSnapshots" in combined
    assert "2. 关键发现" not in combined
    assert "问题结论" not in combined
    assert "日志证据摘要" not in combined
    assert "AI 根因分析" not in combined
    assert str(tmp_path) not in combined
    assert "sqlite" not in combined.lower()
    assert "src/vstacklens" not in combined.replace("\\", "/").lower()


def test_snapshot_vmdk_root_name_boundary_is_preserved(tmp_path: Path) -> None:
    bundle = _snapshot_root_name_edge_bundle_zip(tmp_path / "vmware-support.zip")

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="快照整合失败，想知道什么原因，怎么解决。",
            customer_name="测试客户",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )

    payload = json.loads((result.report_dir / "log_analysis_payload.json").read_text(encoding="utf-8"))
    diagnosis = payload["diagnosis"]
    affected = diagnosis["affected_objects"]

    assert affected["snapshot_disks"]
    assert any(item["file"] == "SQL-0001-000001.vmdk" for item in affected["snapshot_disks"])
    assert "SQL-0001.vmdk" in json.dumps(affected.get("vmdk_chain_files", []), ensure_ascii=False)
    assert "SQL.vmdk" not in json.dumps(affected.get("vmdk_chain_files", []), ensure_ascii=False)


def test_problem_driven_diagnosis_for_migration_network_issue(tmp_path: Path) -> None:
    bundle = _migration_network_support_bundle_zip(tmp_path / "vmware-support.zip")
    problem = "esx-host-01在线迁移后网络不通，重连虚拟网卡又能通"
    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description=problem,
            customer_name="测试客户",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )

    payload = json.loads((result.report_dir / "log_analysis_payload.json").read_text(encoding="utf-8"))
    diagnosis = payload["diagnosis"]
    html_text = result.html_path.read_text(encoding="utf-8")
    doc_text = _docx_text(result.docx_path)
    combined = html_text + "\n" + doc_text
    diagnosis_text = json.dumps(diagnosis, ensure_ascii=False)
    evidence_text = json.dumps(diagnosis.get("evidence_chain", []), ensure_ascii=False)

    assert diagnosis["scenario"] == "vm_migration_network_loss"
    _assert_rule_only_engine(diagnosis["diagnosis_engine"])
    assert diagnosis["confidence"] == "中"
    assert diagnosis["affected_objects"]["primary_object"] == "esx-host-01"
    assert {"迁移", "网络", "虚拟网卡"} <= set(diagnosis["affected_objects"]["domains"])
    assert "性能" not in set(diagnosis["affected_objects"]["domains"])
    assert problem in combined
    assert "esx-host-01" in combined
    _assert_no_model_log_check_report(combined)
    assert "当前证据指向" not in combined
    assert "故障发生时间点" in combined
    assert "迁移任务截图或 vCenter Task/Event 记录" in combined
    assert "迁移前源主机和迁移后目标主机" in combined
    assert "分布式端口" in combined
    assert "vNIC connect 状态截图" in combined
    assert "ARP" in combined
    assert "端口组" in combined
    assert "VLAN" in combined
    assert "上行链路" in combined
    assert "dvPort、vmkernel、uplink 状态" in combined
    assert "Guest OS 内网卡状态、IP、路由、ARP、网关连通性" in combined
    assert "Vigor transport" in combined
    assert "重连虚拟网卡恢复" in diagnosis_text
    assert "判断依据及日志证据" not in combined
    assert "问题时间线" not in combined
    assert "2. 关键发现" not in combined
    assert "3. 日志证据摘要" not in combined
    assert "VMK_ENETUNREACH" not in combined
    assert "VMK_ENETRESET" not in combined
    assert "VMK_VLAN_FILTERED" not in combined
    assert "vmkerrcode_-l.txt" not in evidence_text
    assert "usb.ids" not in evidence_text
    assert "pci.ids" not in evidence_text
    timeline_text = json.dumps(diagnosis.get("timeline", []), ensure_ascii=False)
    assert "MigrateVM_Task" in timeline_text
    assert "connectable.connected" in timeline_text
    assert "vmkerrcode_-l.txt" not in timeline_text
    assert "usb.ids" not in timeline_text
    assert "pci.ids" not in timeline_text
    assert "AI 根因分析" not in combined
    assert "rule_based_log_diagnosis_with_optional_cloud" not in combined
    assert str(tmp_path) not in combined
    assert "sqlite" not in combined.lower()
    assert "src/vstacklens" not in combined.replace("\\", "/").lower()


def test_problem_driven_diagnosis_recognizes_arp_conflict_variants(tmp_path: Path) -> None:
    service = LogAnalysisService()
    sources = [
        {
            "path": "/var/run/log/vmkernel.log",
            "archive_path": "vmware-support-hosts.tgz!esx-atcdi-node-0115/var/run/log/vmkernel-172.16.102.187.log",
            "host": "esx-atcdi-node-0115",
            "host_management_ip": "10.240.5.187",
            "text": "\n".join(
                [
                    "2026-06-14T09:32:12Z cpu12:2097152)WARNING: arp: 00:50:56:6d:f7:1b is using my IP address 10.240.6.187 on vmk2",
                    "2026-06-14T09:32:13Z cpu12:2097152)WARNING: arp: 00:50:56:6d:f7:1b is using my IP address 10.240.6.187 on vmk2.",
                ]
            ),
        }
    ]

    candidates = service._cloud_candidate_strong_evidence(sources)

    assert len(candidates) == 1
    assert candidates[0]["type"] == "ARP/IP 冲突"
    assert candidates[0]["vmkernel_interface"] == "vmk2"
    assert candidates[0]["conflict_ip"] == "10.240.6.187"
    assert candidates[0]["conflict_mac"] == "00:50:56:6d:f7:1b"
    assert candidates[0]["count"] == 2


def test_default_rule_mode_does_not_call_cloud(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")

    def fail_if_called(self: CloudModelClient, prompt: str) -> CloudModelResult:
        raise AssertionError("cloud model should not be called by default")

    monkeypatch.setattr(CloudModelClient, "complete_json", fail_if_called)
    monkeypatch.setattr(CloudModelClient, "diagnose", fail_if_called)

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="客户反馈管理界面偶发超时，需要分析日志。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )

    payload = json.loads((result.report_dir / "log_analysis_payload.json").read_text(encoding="utf-8"))
    _assert_rule_only_engine(payload["diagnosis"]["diagnosis_engine"])
    _assert_no_model_log_check_report(result.html_path.read_text(encoding="utf-8") + "\n" + _docx_text(result.docx_path))


def test_cloud_enabled_requires_api_key(tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")

    errors = LogAnalysisService().validate_config(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="客户反馈管理界面偶发超时，需要分析日志。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key="",
        )
    )

    assert any("API Key" in error for error in errors)


def test_formal_cloud_analysis_passes_config_and_emits_cloud_progress(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = _migration_network_support_bundle_zip(tmp_path / "vmware-support.zip")
    calls = _patch_cloud_success(monkeypatch)
    progress: list[tuple[str, int, str]] = []

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="esx-host-01 在线迁移后网络不通，重连虚拟网卡又能通。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_provider="deepseek",
            cloud_api_url="https://api.deepseek.example/chat/completions",
            cloud_model_name="deepseek-formal-model",
            cloud_api_key="sk-formal-secret",
            cloud_timeout_seconds=37,
        ),
        progress_callback=lambda item: progress.append((item.stage, item.percent, item.message)),
    )

    assert [call["stage"] for call in calls] == ["planning", "diagnosis"]
    for call in calls:
        config = call["config"]
        assert config.provider == "deepseek"
        assert config.api_url == "https://api.deepseek.example/chat/completions"
        assert config.model_name == "deepseek-formal-model"
        assert config.api_key == "sk-formal-secret"
        assert config.timeout_seconds == 37
    assert ("cloud_planning", 60, "正在调用云端模型规划日志取证范围") in progress
    assert ("cloud_diagnosing", 68, "正在调用云端模型进行日志判断") in progress

    payload_text = (result.report_dir / "log_analysis_payload.json").read_text(encoding="utf-8")
    engine = json.loads(payload_text)["diagnosis"]["diagnosis_engine"]
    _assert_cloud_engine(engine, provider="deepseek", model="deepseek-formal-model")
    assert "sk-formal-secret" not in payload_text


def test_cloud_success_generates_model_assisted_diagnosis_report(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = _migration_network_support_bundle_zip(tmp_path / "vmware-support.zip")
    api_key = "sk-secret-api-key"
    calls = _patch_cloud_success(monkeypatch)

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="esx-host-01 在线迁移后网络不通，重连虚拟网卡又能通。",
            customer_name="测试客户",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key=api_key,
        )
    )

    assert len(calls) == 2
    assert [call["stage"] for call in calls] == ["planning", "diagnosis"]
    planning_prompt = calls[0]["prompt"]
    diagnosis_prompt = calls[1]["prompt"]
    assert "模型驱动取证" in planning_prompt
    assert "日志目录索引" in planning_prompt
    assert "log_directory_index" in planning_prompt
    assert "提取日志片段" in diagnosis_prompt
    assert "extracted_log_snippets" in diagnosis_prompt
    assert "不是报告润色器" in diagnosis_prompt
    assert "强证据" in diagnosis_prompt
    assert "中证据" in diagnosis_prompt
    assert "弱证据" in diagnosis_prompt
    assert "mode=triage" in diagnosis_prompt
    assert "不能编造" in diagnosis_prompt
    assert api_key not in planning_prompt + diagnosis_prompt
    assert str(tmp_path) not in planning_prompt + diagnosis_prompt
    payload_text = (result.report_dir / "log_analysis_payload.json").read_text(encoding="utf-8")
    diagnosis_text = (result.report_dir / "diagnosis.json").read_text(encoding="utf-8")
    html_text = result.html_path.read_text(encoding="utf-8")
    doc_text = _docx_text(result.docx_path)
    combined = payload_text + "\n" + diagnosis_text + "\n" + html_text + "\n" + doc_text
    payload = json.loads(payload_text)
    diagnosis = payload["diagnosis"]

    _assert_cloud_engine(diagnosis["diagnosis_engine"])
    assert diagnosis["diagnosis_engine"]["model_quality"] == "accepted"
    assert diagnosis["key_evidence_count"] >= 2
    assert diagnosis["current_judgement"].startswith("当前日志显示")
    assert "1. 诊断结论" in combined
    assert "1.1 当前判断" in combined
    assert "当前能够确认" in combined
    assert "当前不能确认" in combined
    assert "需要客户补充的信息" in combined
    assert "VStackLens 日志排查报告" not in combined
    assert "鎶ュ憡" not in combined
    assert "鍒ゆ柇" not in combined
    assert "�" not in combined
    assert api_key not in combined
    assert "sk-secret" not in combined
    for local_marker in ("D:\\日志", "D:\\LOG", "C:\\Users", "D:\\软件开发", "src\\vstacklens", "site-packages", "sqlite"):
        assert local_marker.lower() not in combined.lower()

    with connect(tmp_path / "log-analysis.db") as conn:
        names = {row["report_name"] for row in conn.execute("SELECT report_name FROM log_analysis_reports").fetchall()}
    assert names == {"VStackLens 模型辅助日志诊断报告"}


def test_cloud_arp_conflict_candidate_can_drive_model_diagnosis(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = _arp_conflict_support_bundle_zip(tmp_path / "vmware-support.zip")
    plan = {
        "requests": [
            {
                "file": "vmware-support-hosts.tgz!esx-atcdi-node-0115/var/run/log/vmkernel-172.16.102.187.log",
                "component": "vmkernel",
                "keywords": ["arp", "using my IP address", "vmk2", "10.240.6.187", "00:50:56:6d:f7:1b"],
                "context_lines": 20,
                "reason": "候选强证据显示 vmk2 存在 ARP/IP 冲突，需要提取上下文。",
            }
        ],
        "stop_reason": "优先验证 ARP/IP 冲突强候选。",
    }
    calls = _patch_cloud_success(monkeypatch, _arp_cloud_content(), plan=plan)

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="业务虚拟机网络异常，怀疑 VMkernel 网络存在 ARP/IP 冲突。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key="sk-secret-api-key",
        )
    )

    payload_text = (result.report_dir / "log_analysis_payload.json").read_text(encoding="utf-8")
    diagnosis_text = (result.report_dir / "diagnosis.json").read_text(encoding="utf-8")
    report_text = result.html_path.read_text(encoding="utf-8") + "\n" + _docx_text(result.docx_path)
    combined = payload_text + "\n" + diagnosis_text + "\n" + report_text
    payload = json.loads(payload_text)
    diagnosis = payload["diagnosis"]
    strong = diagnosis.get("candidate_strong_evidence") or []

    assert [call["stage"] for call in calls] == ["planning", "diagnosis"]
    assert "vmk2" in calls[0]["prompt"]
    assert "10.240.6.187" in calls[0]["prompt"]
    assert "00:50:56:6d:f7:1b" in calls[0]["prompt"]
    _assert_cloud_engine(diagnosis["diagnosis_engine"])
    assert strong
    assert strong[0]["type"] == "ARP/IP 冲突"
    assert strong[0]["host"] == "esx-atcdi-node-0115"
    assert strong[0]["host_management_ip"] == "10.240.5.187"
    assert strong[0]["vmkernel_interface"] == "vmk2"
    assert strong[0]["conflict_ip"] == "10.240.6.187"
    assert strong[0]["conflict_mac"] == "00:50:56:6d:f7:1b"
    assert strong[0]["count"] == 586
    assert "VStackLens 模型辅助日志诊断报告" in report_text
    assert "vmk2" in combined
    assert "10.240.6.187" in combined
    assert "00:50:56:6d:f7:1b" in combined
    assert "ARP/IP 冲突" in combined
    assert "当前能够确认" in combined
    assert "当前不能确认" in combined
    assert "net-dvs_-l.txt" not in combined
    assert "sk-secret" not in combined
    assert str(tmp_path) not in combined


def test_cloud_network_output_ignoring_arp_candidate_is_downgraded(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = _arp_conflict_support_bundle_zip(tmp_path / "vmware-support.zip")
    _patch_cloud_success(monkeypatch, _cloud_content())

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="业务虚拟机网络异常，怀疑 VMkernel 网络存在 ARP/IP 冲突。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key="sk-secret-api-key",
        )
    )

    payload, combined = _cloud_fallback_result_payload(result)
    diagnosis = payload["diagnosis"]

    _assert_cloud_triage_engine(diagnosis["diagnosis_engine"])
    assert payload["summary"]["cloud_assist_failed"] is True
    assert diagnosis["cloud_assist_status"] == "responded_downgraded"
    assert diagnosis["cloud_assist_quality"] == "downgraded"
    assert payload["metadata"]["cloud_assist_failure_reason"]
    assert "网络问题存在 ARP/IP 冲突强候选但模型输出未引用" in diagnosis["diagnosis_engine"]["model_quality_reason"]
    assert "VStackLens 模型辅助粗排查报告" not in combined
    assert "sk-secret" not in combined


def test_cloud_generic_judgement_downgrade_falls_back_to_rule_report(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = _migration_network_support_bundle_zip(tmp_path / "vmware-support.zip")
    content = _cloud_content(
        current_judgement="当前日志线索与客户问题存在关联，建议优先围绕时间窗口、问题对象和关键组件状态进一步核对。"
    )
    _patch_cloud_success(monkeypatch, content)

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="esx-host-01 在线迁移后网络不通，重连虚拟网卡又能通。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key="sk-secret-api-key",
        )
    )

    payload, combined = _cloud_fallback_result_payload(result)
    diagnosis = payload["diagnosis"]

    _assert_cloud_triage_engine(diagnosis["diagnosis_engine"])
    assert diagnosis["model_output_mode"] == "triage"
    assert diagnosis["cloud_assist_failed"] is True
    assert payload["metadata"]["cloud_assist_failure_reason"]
    assert "current_judgement 过于泛化" in diagnosis["diagnosis_engine"]["model_quality_reason"]
    assert "VStackLens 模型辅助粗排查报告" not in combined
    assert "sk-secret" not in combined


def test_cloud_diagnosis_with_weak_or_empty_evidence_is_not_full_diagnosis(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = _static_vsan_support_bundle_zip(tmp_path / "vmware-support.zip")
    content = _cloud_content(
        current_judgement="当前日志显示静态 vSAN 与 VMDK 清单中存在客户描述方向的有限线索，结合客户描述只能作为排查入口，当前 support bundle 尚不能单独证明最终根因。"
    )
    content.update(
        {
            "mode": "diagnosis",
            "can_determine_root_cause": True,
            "evidence_quality": "weak",
            "confidence": "高",
        }
    )
    _patch_cloud_success(monkeypatch, content)

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="vSAN / IO 拥堵，想确认是不是存储路径或 vSAN 异常导致。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key="sk-secret-api-key",
        )
    )

    payload, combined = _cloud_fallback_result_payload(result)
    diagnosis = payload["diagnosis"]

    _assert_cloud_triage_engine(diagnosis["diagnosis_engine"])
    assert diagnosis["model_output_mode"] == "triage"
    assert diagnosis["model_evidence_quality"] == "weak"
    assert "模型声称可确认根因但证据链不足" in diagnosis["diagnosis_engine"]["model_quality_reason"]
    assert "VStackLens 模型辅助粗排查报告" not in combined
    assert "VStackLens 模型辅助诊断报告" not in combined
    assert "vmkfstools" not in json.dumps(diagnosis.get("evidence_chain", []), ensure_ascii=False)
    assert "dump-vmdk" not in json.dumps(diagnosis.get("evidence_chain", []), ensure_ascii=False)


def test_cloud_empty_evidence_diagnosis_is_downgraded(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = tmp_path / "vmware-support.zip"
    with zipfile.ZipFile(bundle, "w") as archive:
        archive.writestr("manifest.txt", "support bundle manifest")
    _patch_cloud_success(monkeypatch, _cloud_content(evidence_chain=[]))

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="客户反馈业务虚拟机偶发网络不通，需要分析日志。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key="sk-secret-api-key",
        )
    )

    payload, combined = _cloud_fallback_result_payload(result)
    diagnosis = payload["diagnosis"]

    _assert_cloud_triage_engine(diagnosis["diagnosis_engine"])
    assert diagnosis["key_evidence_count"] == 0
    assert "当前没有可用证据链" in diagnosis["diagnosis_engine"]["model_quality_reason"]
    assert "VStackLens 模型辅助粗排查报告" not in combined


def test_cloud_rejected_output_falls_back_to_rule_report_without_leaks(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = _migration_network_support_bundle_zip(tmp_path / "vmware-support.zip")
    api_key = "sk-secret-api-key"
    content = _cloud_content(
        current_judgement=f"当前日志显示 esx-host-01 存在线索，但模型错误输出了敏感内容 {api_key} 和本机路径 D:\\软件开发\\secret\\trace.py。"
    )
    _patch_cloud_success(monkeypatch, content)

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="esx-host-01 在线迁移后网络不通，重连虚拟网卡又能通。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key=api_key,
        )
    )

    payload, combined = _cloud_fallback_result_payload(result)
    diagnosis = payload["diagnosis"]

    _assert_cloud_rejected_engine(diagnosis["diagnosis_engine"])
    assert "VStackLens 日志粗排查报告" in combined
    assert "VStackLens 模型辅助诊断报告" not in combined
    assert "VStackLens 模型辅助粗排查报告" not in combined
    assert api_key not in combined
    assert "sk-secret" not in combined
    assert "D:\\软件开发" not in combined
    assert "trace.py" not in combined


def test_cloud_http_failure_blocks_customer_report_without_key(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")
    api_key = "sk-secret-api-key"

    def fake_complete_json(self: CloudModelClient, prompt: str) -> CloudModelResult:
        return CloudModelResult(False, fallback_reason=f"HTTP 401: bad key {api_key}")

    monkeypatch.setattr(CloudModelClient, "complete_json", fake_complete_json)

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="客户反馈管理界面偶发超时，需要分析日志。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key=api_key,
        )
    )

    payload, combined = _cloud_fallback_result_payload(result)
    engine = payload["diagnosis"]["diagnosis_engine"]

    assert engine["model_used"] is False
    assert engine["model_source"] == "rule_only"
    assert engine["api_call_attempted"] is True
    assert engine["api_response_received"] is False
    assert engine["cloud_status"] == "request_failed"
    assert engine["fallback_reason"] == "HTTP 401: bad key ***"
    assert "VStackLens 日志粗排查报告" in combined
    assert api_key not in combined
    assert "sk-secret" not in combined


def test_connection_success_does_not_hide_formal_analysis_failure(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    class SuccessResponse:
        def __enter__(self) -> "SuccessResponse":
            return self

        def __exit__(self, exc_type, exc, traceback) -> None:
            return None

        def read(self) -> bytes:
            body = {"choices": [{"message": {"content": '{"ok": true, "message": "连接成功"}'}}]}
            return json.dumps(body, ensure_ascii=False).encode("utf-8")

    monkeypatch.setattr(cloud_module.urllib.request, "urlopen", lambda request, timeout: SuccessResponse())
    connection_result = CloudModelClient(CloudModelConfig(api_key="sk-secret-formal")).test_connection()
    assert connection_result.ok is True

    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")

    def fake_complete_json(self: CloudModelClient, prompt: str) -> CloudModelResult:
        return CloudModelResult(False, fallback_reason="HTTP 401: bad key sk-secret-formal")

    monkeypatch.setattr(CloudModelClient, "complete_json", fake_complete_json)

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="客户反馈管理界面偶发超时，需要分析日志。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key="sk-secret-formal",
        )
    )

    payload, combined = _cloud_fallback_result_payload(result)
    engine = payload["diagnosis"]["diagnosis_engine"]

    assert engine["api_call_attempted"] is True
    assert engine["api_response_received"] is False
    assert engine["model_used"] is False
    assert engine["model_source"] == "rule_only"
    assert engine["model_quality"] == "failed"
    assert engine["fallback_reason"] == "HTTP 401: bad key ***"
    assert "VStackLens 日志粗排查报告" in combined
    assert "sk-secret-formal" not in combined


def test_cloud_invalid_model_response_is_reported_as_cloud_rejected(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")

    def fake_complete_json(self: CloudModelClient, prompt: str) -> CloudModelResult:
        return CloudModelResult(True, content={"requests": []}, response_received=True)

    def fake_diagnose(self: CloudModelClient, prompt: str) -> CloudModelResult:
        return CloudModelResult(False, fallback_reason="cloud JSON schema rejected", response_received=True)

    monkeypatch.setattr(CloudModelClient, "complete_json", fake_complete_json)
    monkeypatch.setattr(CloudModelClient, "diagnose", fake_diagnose)

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="客户反馈管理界面偶发超时，需要分析日志。",
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
            cloud_assist_enabled=True,
            cloud_api_key="sk-secret-api-key",
        )
    )

    payload, combined = _cloud_fallback_result_payload(result)
    engine = payload["diagnosis"]["diagnosis_engine"]

    assert engine["model_used"] is True
    assert engine["model_source"] == "cloud"
    assert engine["mode"] == "cloud_rejected"
    assert engine["model_quality"] == "rejected"
    assert engine["api_call_attempted"] is True
    assert engine["api_response_received"] is True
    assert engine["cloud_status"] == "responded_rejected"
    assert engine["fallback_reason"] == "cloud JSON schema rejected"
    assert "VStackLens 日志粗排查报告" in combined
    assert "sk-secret" not in combined


def test_cloud_client_uses_chat_completions_messages_protocol(monkeypatch: pytest.MonkeyPatch) -> None:
    requests: list[dict] = []

    class FakeResponse:
        def __enter__(self) -> "FakeResponse":
            return self

        def __exit__(self, exc_type, exc, traceback) -> None:
            return None

        def read(self) -> bytes:
            body = {"choices": [{"message": {"content": json.dumps(_cloud_content(), ensure_ascii=False)}}]}
            return json.dumps(body, ensure_ascii=False).encode("utf-8")

    def fake_urlopen(request, timeout: int):
        body = json.loads(request.data.decode("utf-8"))
        requests.append({"url": request.full_url, "headers": dict(request.header_items()), "body": body, "timeout": timeout})
        return FakeResponse()

    monkeypatch.setattr(cloud_module.urllib.request, "urlopen", fake_urlopen)

    result = CloudModelClient(CloudModelConfig(api_url="https://api.example.com/chat/completions", api_key="sk-secret")).diagnose("客户问题和证据摘要")

    assert result.ok is True
    assert requests[0]["url"] == "https://api.example.com/chat/completions"
    assert requests[0]["body"]["model"] == DEFAULT_CLOUD_MODEL
    assert "messages" in requests[0]["body"]
    assert requests[0]["body"]["messages"][0]["role"] == "system"
    assert requests[0]["body"]["messages"][1] == {"role": "user", "content": "客户问题和证据摘要"}
    assert requests[0]["body"]["response_format"] == {"type": "json_object"}
    assert "prompt" not in requests[0]["body"]


def test_cloud_client_retries_without_response_format_on_unsupported_http(monkeypatch: pytest.MonkeyPatch) -> None:
    requests: list[dict] = []

    class FakeResponse:
        def __enter__(self) -> "FakeResponse":
            return self

        def __exit__(self, exc_type, exc, traceback) -> None:
            return None

        def read(self) -> bytes:
            content = f"模型说明前缀 {json.dumps(_cloud_content(), ensure_ascii=False)} 模型说明后缀"
            body = {"choices": [{"message": {"content": content}}]}
            return json.dumps(body, ensure_ascii=False).encode("utf-8")

    def fake_urlopen(request, timeout: int):
        body = json.loads(request.data.decode("utf-8"))
        requests.append(body)
        if len(requests) == 1:
            raise urllib.error.HTTPError(request.full_url, 400, "response_format unsupported", hdrs=None, fp=None)
        return FakeResponse()

    monkeypatch.setattr(cloud_module.urllib.request, "urlopen", fake_urlopen)

    result = CloudModelClient(CloudModelConfig(api_key="sk-secret")).diagnose("证据摘要")

    assert result.ok is True
    assert len(requests) == 2
    assert "response_format" in requests[0]
    assert "response_format" not in requests[1]
    assert result.content and result.content["confidence"] == "中"


def test_cloud_client_test_connection_success_and_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    class SuccessResponse:
        def __enter__(self) -> "SuccessResponse":
            return self

        def __exit__(self, exc_type, exc, traceback) -> None:
            return None

        def read(self) -> bytes:
            body = {"choices": [{"message": {"content": '{"ok": true, "message": "连接成功"}'}}]}
            return json.dumps(body).encode("utf-8")

    monkeypatch.setattr(cloud_module.urllib.request, "urlopen", lambda request, timeout: SuccessResponse())
    success = CloudModelClient(CloudModelConfig(api_key="sk-secret")).test_connection()
    assert success.ok is True
    assert success.content == {"ok": True, "message": "连接成功"}

    def fail_urlopen(request, timeout: int):
        raise urllib.error.HTTPError(request.full_url, 401, "bad key sk-secret", hdrs=None, fp=None)

    monkeypatch.setattr(cloud_module.urllib.request, "urlopen", fail_urlopen)
    failure = CloudModelClient(CloudModelConfig(api_key="sk-secret")).test_connection()
    assert failure.ok is False
    assert "认证失败" in failure.fallback_reason
    assert "HTTP 401" in failure.fallback_reason
    assert "sk-secret" not in failure.fallback_reason


def test_cloud_client_marks_invalid_success_response_as_response_received(monkeypatch: pytest.MonkeyPatch) -> None:
    class InvalidResponse:
        def __enter__(self) -> "InvalidResponse":
            return self

        def __exit__(self, exc_type, exc, traceback) -> None:
            return None

        def read(self) -> bytes:
            return b"not json"

    monkeypatch.setattr(cloud_module.urllib.request, "urlopen", lambda request, timeout: InvalidResponse())

    result = CloudModelClient(CloudModelConfig(api_key="sk-secret")).diagnose("证据摘要")

    assert result.ok is False
    assert result.response_received is True
    assert "云端接口返回不是有效 JSON" in result.fallback_reason
    assert "sk-secret" not in result.fallback_reason


def test_cloud_client_returns_readable_sanitized_connection_failures(monkeypatch: pytest.MonkeyPatch) -> None:
    cases = [
        (
            urllib.error.HTTPError(
                "https://api.example.com/chat/completions",
                403,
                "Forbidden",
                hdrs=None,
                fp=io.BytesIO(b'{"error":{"message":"model permission denied for sk-secret"}}'),
            ),
            "权限不足",
        ),
        (
            urllib.error.HTTPError(
                "https://api.example.com/chat/completions",
                429,
                "Too Many Requests",
                hdrs=None,
                fp=io.BytesIO(b'{"error":{"message":"insufficient balance"}}'),
            ),
            "限流",
        ),
        (urllib.error.URLError("timed out"), "网络连接超时"),
        (ValueError("model_not_found: deepseek-reasoner is unavailable"), "模型不可用"),
    ]

    for exc, expected in cases:
        def fail_urlopen(request, timeout: int, failure: Exception = exc):
            raise failure

        monkeypatch.setattr(cloud_module.urllib.request, "urlopen", fail_urlopen)
        result = CloudModelClient(CloudModelConfig(api_key="sk-secret")).test_connection()

        assert result.ok is False
        assert expected in result.fallback_reason
        assert "sk-secret" not in result.fallback_reason


def test_log_analysis_display_stats_prefers_v3_diagnosis_fields() -> None:
    stats = log_analysis_display_stats(
        {"log_files": 7, "evidence_count": 0},
        {
            "key_evidence_count": 4,
            "missing_materials": ["故障发生时间点", "迁移任务截图"],
            "recommended_next_steps": ["核对端口组", "检查上行链路", "复核 Guest OS 网卡"],
        },
        [],
    )

    assert stats == {
        "log_files": 7,
        "key_evidence": 4,
        "missing_materials": 2,
        "recommendations": 3,
    }


def test_log_analysis_display_stats_falls_back_for_legacy_payloads() -> None:
    stats = log_analysis_display_stats(
        {"LogFile": 5, "Evidence": 6},
        {},
        [{"title": "网络异常"}, {"title": "认证失败"}],
    )

    assert stats == {
        "log_files": 5,
        "key_evidence": 6,
        "missing_materials": 0,
        "recommendations": 2,
    }


def test_log_analysis_display_stats_falls_back_to_diagnosis_log_files() -> None:
    stats = log_analysis_display_stats(
        {"log_files": 0, "evidence_count": 0},
        {
            "key_evidence_count": 2,
            "evidence_chain": [
                {"file": "hostd.log", "message": "MigrateVM_Task"},
                {"file": "vmkernel.log", "message": "Vigor transport disconnected"},
            ],
            "missing_materials": ["迁移任务截图"],
            "handling_recommendations": ["核对目标主机网络路径"],
        },
        [],
    )

    assert stats["log_files"] == 2
    assert stats["key_evidence"] == 2
    assert stats["missing_materials"] == 1
    assert stats["recommendations"] == 1


def test_service_snapshot_vmdk_name_helper_distinguishes_base_and_snapshot() -> None:
    service = LogAnalysisService()

    assert service._is_snapshot_vmdk_name("SQL-0001.vmdk") is False
    assert service._is_snapshot_vmdk_name("SQL-0001-000001.vmdk") is True
    assert service._vmdk_chain_root("SQL-0001.vmdk") == "SQL-0001"
    assert service._vmdk_chain_root("SQL-0001-000001.vmdk") == "SQL-0001"


def test_service_read_tar_member_text_tolerates_extractfile_failure() -> None:
    service = LogAnalysisService()
    member = tarfile.TarInfo("broken.log")
    member.size = 128

    class BrokenTar:
        def extractfile(self, tarinfo):
            raise tarfile.TarError("broken member")

    assert service._read_tar_member_text(BrokenTar(), member, 1024) == ""


def test_report_center_lists_log_analysis_reports_without_health_metrics(tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")
    db_path = tmp_path / "log-analysis.db"
    report_root = tmp_path / "reports"
    service = InspectionService(config_path=tmp_path / "desktop_config.json")
    service.run_log_analysis(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description="网络链路异常，需要分析 support bundle。",
            customer_name="测试客户",
            db_path=db_path,
            report_output_dir=report_root,
        )
    )

    items = service.list_report_packages(db_path, report_root=report_root)
    types = {item.report_type for item in items}
    log_items = [item for item in items if item.report_type.startswith(("日志分析", "日志排查", "日志粗排查"))]

    assert {"日志粗排查 HTML 报告", "日志粗排查 Word 报告"} <= types
    assert "HTML 报告包" not in types
    assert "Word 报告" not in types
    assert len(log_items) == 2
    assert {item.score for item in log_items} == {"不适用"}
    assert all(item.history_summary == "独立日志分析，不参与健康巡检历史对比。" for item in log_items)
    assert all(item.asset_summary["LogFile"] > 0 for item in log_items)
    assert all(item.asset_summary["Evidence"] > 0 for item in log_items)
    assert all("Supplemental" in item.asset_summary for item in log_items)
    assert all("Recommendation" in item.asset_summary for item in log_items)


def test_report_center_uses_diagnosis_stats_from_log_payload(tmp_path: Path) -> None:
    db_path = tmp_path / "log-analysis.db"
    report_dir = tmp_path / "reports" / "log-analysis-run"
    report_dir.mkdir(parents=True)
    html_path = report_dir / "index.html"
    docx_path = report_dir / "VStackLens-Log-Analysis-Report.docx"
    html_path.write_text("<html>log report</html>", encoding="utf-8")
    docx_path.write_bytes(b"docx placeholder")
    payload = {
        "summary": {"log_files": 7, "evidence_count": 0},
        "diagnosis": {
            "key_evidence_count": 4,
            "missing_materials": ["故障发生时间点", "迁移任务截图"],
            "recommended_next_steps": ["核对端口组", "检查上行链路", "复核 Guest OS 网卡"],
        },
        "findings": [],
    }
    (report_dir / "log_analysis_payload.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    init_db(db_path)
    with connect(db_path) as conn:
        conn.execute(
            """
            INSERT INTO log_analysis_runs (
              log_run_id, customer_name, report_title, support_bundle_path,
              support_bundle_name, problem_description, run_status,
              current_stage, progress_percent, summary_json,
              started_at, finished_at, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                "log-run-1",
                "测试客户",
                "VMware 日志分析报告",
                "vmware-support.zip",
                "vmware-support.zip",
                "esx-host-01 在线迁移后网络不通，重连虚拟网卡又能通。",
                "success",
                "success",
                100,
                json.dumps({"log_files": 1, "evidence_count": 1}, ensure_ascii=False),
                "2026-06-13T10:00:00Z",
                "2026-06-13T10:01:00Z",
                "2026-06-13T10:00:00Z",
                "2026-06-13T10:01:00Z",
            ),
        )
        for report_type, path in (("log_analysis_html", html_path), ("log_analysis_docx", docx_path)):
            conn.execute(
                """
                INSERT INTO log_analysis_reports (
                  report_id, log_run_id, report_name, report_type, report_status,
                  file_path, generated_at, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    f"{report_type}-1",
                    "log-run-1",
                    "VMware 日志分析报告",
                    report_type,
                    "success",
                    str(path),
                    "2026-06-13T10:01:00Z",
                    "2026-06-13T10:01:00Z",
                    "2026-06-13T10:01:00Z",
                ),
            )

    items = InspectionService(config_path=tmp_path / "desktop_config.json").list_report_packages(db_path, report_root=tmp_path / "reports")
    log_items = [item for item in items if item.report_type.startswith(("日志分析", "日志排查", "日志粗排查"))]

    assert len(log_items) == 2
    for item in log_items:
        assert item.asset_summary["LogFile"] == 7
        assert item.asset_summary["Evidence"] == 4
        assert item.asset_summary["Supplemental"] == 2
        assert item.asset_summary["Recommendation"] == 3


def test_log_analysis_report_center_is_tolerant_of_old_databases(tmp_path: Path) -> None:
    db_path = tmp_path / "old.db"
    conn = sqlite3.connect(db_path)
    try:
        conn.execute("CREATE TABLE inspection_runs (run_id TEXT PRIMARY KEY)")
        conn.commit()
    finally:
        conn.close()

    items = InspectionService(config_path=tmp_path / "desktop_config.json").list_report_packages(db_path, report_root=tmp_path / "reports")

    assert items == []


def test_log_analysis_retries_when_database_is_locked(monkeypatch, tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")
    db_path = tmp_path / "log-analysis.db"
    init_db(db_path)
    monkeypatch.setattr(db_connection, "CONNECT_TIMEOUT_SECONDS", 0.05)
    monkeypatch.setattr(db_connection, "BUSY_TIMEOUT_MS", 50)

    locker = sqlite3.connect(db_path, timeout=0.05, isolation_level=None, check_same_thread=False)
    locker.execute("PRAGMA busy_timeout = 50")
    locker.execute("BEGIN IMMEDIATE")

    def release_lock() -> None:
        locker.rollback()
        locker.close()

    timer = threading.Timer(0.4, release_lock)
    timer.start()
    try:
        started = time.monotonic()
        result = LogAnalysisService().run(
            LogAnalysisConfig(
                support_bundle_path=bundle,
                problem_description="客户反馈管理界面偶发超时，需要分析日志。",
                db_path=db_path,
                report_output_dir=tmp_path / "reports",
            )
        )
    finally:
        timer.cancel()
        try:
            locker.close()
        except sqlite3.Error:
            pass

    assert time.monotonic() - started >= 0.2
    assert result.html_path.exists()
    assert result.docx_path.exists()
    with connect(db_path) as conn:
        run = conn.execute("SELECT run_status FROM log_analysis_runs WHERE log_run_id = ?", (result.log_run_id,)).fetchone()
        assert run["run_status"] == "success"
        assert conn.execute("SELECT COUNT(*) AS c FROM log_analysis_reports WHERE log_run_id = ?", (result.log_run_id,)).fetchone()["c"] == 2


def test_report_center_read_transaction_does_not_block_log_analysis_write(tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")
    db_path = tmp_path / "log-analysis.db"
    init_db(db_path)
    reader = connect(db_path)
    try:
        reader.execute("BEGIN")
        reader.execute("SELECT COUNT(*) FROM log_analysis_runs").fetchone()
        service = InspectionService(config_path=tmp_path / "desktop_config.json")
        result = service.run_log_analysis(
            LogAnalysisConfig(
                support_bundle_path=bundle,
                problem_description="网络链路异常，需要分析 support bundle。",
                db_path=db_path,
                report_output_dir=tmp_path / "reports",
            )
        )
        items = service.list_report_packages(db_path, report_root=tmp_path / "reports")
    finally:
        reader.rollback()
        reader.close()

    assert result.html_path.exists()
    assert result.docx_path.exists()
    assert {"日志粗排查 HTML 报告", "日志粗排查 Word 报告"} <= {item.report_type for item in items}


def test_log_analysis_failure_leaves_failed_run_without_report_rows(monkeypatch, tmp_path: Path) -> None:
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")
    db_path = tmp_path / "log-analysis.db"

    def fail_render_html(payload: dict, output_dir: Path) -> Path:
        raise RuntimeError("render exploded")

    monkeypatch.setattr(log_analysis_module, "render_log_analysis_html", fail_render_html)

    with pytest.raises(RuntimeError, match="render exploded"):
        LogAnalysisService().run(
            LogAnalysisConfig(
                support_bundle_path=bundle,
                problem_description="客户反馈管理界面偶发超时，需要分析日志。",
                db_path=db_path,
                report_output_dir=tmp_path / "reports",
            )
        )

    with connect(db_path) as conn:
        rows = conn.execute("SELECT run_status, error_message FROM log_analysis_runs").fetchall()
        assert len(rows) == 1
        assert rows[0]["run_status"] == "failed"
        assert "render exploded" in rows[0]["error_message"]
        assert conn.execute("SELECT COUNT(*) AS c FROM log_analysis_reports").fetchone()["c"] == 0


def test_runtime_logging_creates_file_without_sensitive_problem_text(tmp_path: Path) -> None:
    log_path = configure_runtime_logging(tmp_path / "logs")
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")
    sensitive_problem = "SECRET-PASSWORD-12345 客户反馈管理界面偶发超时"

    result = LogAnalysisService().run(
        LogAnalysisConfig(
            support_bundle_path=bundle,
            problem_description=sensitive_problem,
            db_path=tmp_path / "log-analysis.db",
            report_output_dir=tmp_path / "reports",
        )
    )
    _flush_runtime_log_handlers()
    text = log_path.read_text(encoding="utf-8")
    _close_runtime_log_handler(log_path)

    assert log_path.exists()
    assert "Log analysis start" in text
    assert result.log_run_id in text
    assert "length=" in text
    assert "sha256_12=" in text
    assert "SECRET-PASSWORD-12345" not in text
    assert sensitive_problem not in text


def test_runtime_logging_records_traceback_on_log_analysis_failure(monkeypatch, tmp_path: Path) -> None:
    log_path = configure_runtime_logging(tmp_path / "logs")
    bundle = _support_bundle_zip(tmp_path / "vmware-support.zip")

    def fail_render_html(payload: dict, output_dir: Path) -> Path:
        raise RuntimeError("render exploded")

    monkeypatch.setattr(log_analysis_module, "render_log_analysis_html", fail_render_html)

    with pytest.raises(RuntimeError, match="render exploded"):
        LogAnalysisService().run(
            LogAnalysisConfig(
                support_bundle_path=bundle,
                problem_description="客户反馈管理界面偶发超时，需要分析日志。",
                db_path=tmp_path / "log-analysis.db",
                report_output_dir=tmp_path / "reports",
            )
        )
    _flush_runtime_log_handlers()
    text = log_path.read_text(encoding="utf-8")
    _close_runtime_log_handler(log_path)

    assert "Log analysis failed" in text
    assert "Traceback" in text
    assert "RuntimeError: render exploded" in text
    assert "log_run_id=logrun" in text


def test_desktop_database_lock_message_is_customer_friendly() -> None:
    pytest.importorskip("PySide6")
    from vstacklens.desktop.app import FRIENDLY_DB_LOCK_ERROR, friendly_error_message

    message = friendly_error_message("sqlite3.OperationalError: database is locked")

    assert message == FRIENDLY_DB_LOCK_ERROR
    assert "database is locked" not in message
