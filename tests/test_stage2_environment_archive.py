from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import pytest

from vstacklens.application.inspection_service import InspectionService, DesktopInspectionConfig
from vstacklens.application.inspection_runner import InspectionRunner, MockRunRequest
from vstacklens.db.connection import connect

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def archives(tmp_path):
    service = InspectionService(config_path=tmp_path / "config.json")
    db = tmp_path / "isolated.db"
    fields = dict(address="192.0.2.111", site_name="实验站点")
    a = service.save_environment_archive(db, name="客户甲生产环境", customer_name="客户甲", **fields)
    b = service.save_environment_archive(db, name="客户乙生产环境", customer_name="客户乙", **fields)
    empty = service.save_environment_archive(db, name="尚未巡检", customer_name="客户丙", address="192.0.2.112")
    return service, db, a, b, empty


def mock_run(db, env, output, *, failing=False, all_formats=True):
    return InspectionRunner().run_mock(MockRunRequest(db_path=db, environment_id=env, vcenter="192.0.2.111",
        rulepack_path=ROOT / "rulepacks/builtin-vsphere-v1", report_dir=output,
        fixture_path=ROOT / f"tests/fixtures/inventory_{'fail' if failing else 'pass'}.json",
        docx_out=output.with_suffix(".docx") if all_formats else None,
        pdf_out=output.with_suffix(".pdf") if all_formats else None))


def historical_rows(db):
    with connect(db) as conn:
        return {t: [tuple(r) for r in conn.execute(f"SELECT * FROM {t} ORDER BY rowid")]
                for t in ("inspection_runs", "inventory_snapshots", "inventory_objects", "rule_results", "findings", "reports")}


def test_same_address_two_customers_keep_data_and_three_reports_separate(archives, tmp_path):
    service, db, a, b, _ = archives
    assert a != b
    first = mock_run(db, a, tmp_path / "customer-a", failing=True)
    second = mock_run(db, b, tmp_path / "customer-b")
    for env, run, customer in ((a, first, "客户甲"), (b, second, "客户乙")):
        data = service.get_latest_valid_desktop_data(db, environment_id=env)
        assert data.run_id == run.run_id
        assert data.customer_name == customer
        assert all(item.run_id == run.run_id for item in data.recent_reports)
        assert {x.run_id for x in service.list_platform_assets(db, data.run_id)} == {run.run_id}
        assert {x.run_id for x in service.list_report_history(db, environment_id=env)} == {run.run_id}
        reports = service.list_report_packages(db, tmp_path, environment_id=env)
        assert reports and all(item.run_id == run.run_id for item in reports)
        assert set(service.desktop_report_formats(db, run.run_id)) == {"HTML", "Word", "PDF"}
        payload = json.loads((run.report_dir / "data/customer_report_payload.json").read_text(encoding="utf-8"))
        assert payload["report_context"]["customer_info"]["customer_name"] == customer
        from docx import Document
        word = Document(run.docx_path)
        visible = "\n".join([p.text for p in word.paragraphs] + [p.text for section in word.sections for p in section.header.paragraphs])
        assert customer in visible
        assert ("客户乙" if customer=="客户甲" else "客户甲") not in visible
        import pymupdf
        with pymupdf.open(run.pdf_path) as pdf:
            assert customer in "".join(p.get_text() for p in pdf)


def test_rename_and_disable_preserve_every_history_record_and_file(archives, tmp_path):
    service, db, a, _, _ = archives
    run = mock_run(db, a, tmp_path / "before-name")
    before = historical_rows(db)
    returned = service.save_environment_archive(db, environment_id=a, name="更名后的环境", customer_name="客户甲", address="192.0.2.111", site_name="实验站点")
    assert returned == a
    service.set_environment_enabled(db, a, False)
    assert historical_rows(db) == before
    assert service.get_latest_valid_desktop_data(db, environment_id=a).run_id == run.run_id
    assert service.list_report_packages(db, tmp_path, environment_id=a)
    assert all(p.is_file() for p in service.desktop_report_formats(db, run.run_id).values())
    with pytest.raises(ValueError, match="停用"):
        mock_run(db, a, tmp_path / "disabled")
    assert historical_rows(db) == before
    service.set_environment_enabled(db, a, True)
    next_run = mock_run(db, a, tmp_path / "after-name", all_formats=False)
    assert service.environment_run_ids(db, a) == {run.run_id, next_run.run_id}
    comparison = service.get_history_comparison(db, environment_id=a)
    assert {item.run_id for item in comparison.run_options} == {run.run_id, next_run.run_id}


def test_empty_environment_never_falls_back_to_another_environment(archives, tmp_path):
    service, db, a, b, empty = archives
    run = mock_run(db, a, tmp_path / "only-a", all_formats=False)
    for env in (b, empty, "does-not-exist"):
        assert not service.get_latest_valid_desktop_data(db, environment_id=env).has_data
        assert service.list_report_history(db, environment_id=env) == []
        assert service.list_report_packages(db, tmp_path, environment_id=env) == []
        comparison = service.get_history_comparison(db, environment_id=env)
        assert comparison.state == "empty"
        assert not comparison.run_options
    denied = service.get_history_comparison(db, comparison_run_id=run.run_id, environment_id=b)
    assert denied.state == "run_unavailable"


def test_archive_distinguishes_partial_collection_missing_reports_and_live_status(archives, tmp_path):
    service, db, a, b, _ = archives
    run = mock_run(db, a, tmp_path / "partial", all_formats=False)
    for path in service.desktop_report_formats(db, run.run_id).values():
        # Move only this isolated test artifact, preserving it for inspection.
        path.rename(path.with_name(path.name + ".retained"))
    rows = {e.environment_id: e for e in service.list_environment_archives(db)}
    assert rows[a].last_collected_at
    assert rows[a].collection_result.startswith("部分采集")
    assert "缺失" in rows[a].report_result
    assert not rows[a].available_formats
    assert rows[b].collection_result == "尚无巡检记录"
    assert rows[b].report_result == "尚无报告"


def test_profile_validation_and_failed_save_do_not_leave_partial_business_rows(archives):
    service, db, a, _, _ = archives
    with connect(db) as conn:
        before = {t: [tuple(r) for r in conn.execute(f"SELECT * FROM {t}")] for t in ("customers", "sites", "vcenters")}
    with pytest.raises(ValueError, match="已有相同地址"):
        service.save_environment_archive(db, name="重复档案", customer_name="客户甲", site_name="实验站点", address="192.0.2.111")
    with pytest.raises(ValueError, match="密码"):
        service.save_environment_archive(db, name="不合法", customer_name="新客户", address="https://user:private@192.0.2.111")
    with connect(db) as conn:
        assert before == {t: [tuple(r) for r in conn.execute(f"SELECT * FROM {t}")] for t in before}


def test_selected_environment_reaches_formal_service_request_without_saved_password(archives, tmp_path, monkeypatch):
    service, db, a, _, _ = archives
    captured = []
    def cancelled(request, **_):
        captured.append(request)
        from vstacklens.application.inspection_runner import RunnerCancelledError
        raise RunnerCancelledError("offline request capture")
    monkeypatch.setattr(service.runner, "run_vcenter", cancelled)
    from vstacklens.application.inspection_runner import RunnerCancelledError
    config = DesktopInspectionConfig(db_path=db, rulepack_path=ROOT / "rulepacks/builtin-vsphere-v1",
        report_output_dir=tmp_path / "unused", vcenter="192.0.2.111", username="offline", password="session-private", environment_id=a)
    with pytest.raises(RunnerCancelledError):
        service.run_vcenter_inspection(config)
    assert captured[0].environment_id == a
    assert "session-private" not in service.config_path.read_text(encoding="utf-8")
    assert "environment_id" not in service.config_path.read_text(encoding="utf-8")


def test_actual_widgets_route_all_entries_with_scope_and_clear_password(archives, tmp_path, qt_application_lifetime):
    service, db, a, b, empty = archives
    first = mock_run(db, a, tmp_path / "a", failing=True, all_formats=False)
    second = mock_run(db, b, tmp_path / "b", all_formats=False)
    service.save_config(DesktopInspectionConfig(db_path=db, rulepack_path=ROOT / "rulepacks/builtin-vsphere-v1", report_output_dir=tmp_path))
    from vstacklens.desktop.app import MainWindow
    from PySide6.QtWidgets import QPushButton
    window = MainWindow(service=service)
    window.show()
    for env, run in ((a, first), (b, second), (empty, None)):
        row = next(i for i,e in enumerate(window._environment_archives) if e.environment_id == env)
        for target in ("asset_center", "risk_center", "history_compare", "report_center"):
            window.environment_table.setCurrentCell(row, 0)
            window.environment_links[target].click()
            assert window._environment_id == env
            assert window._caliber_run_id == (run.run_id if run else None)
            assert all(item.run_id == run.run_id for item in window._asset_items) if run else not window._asset_items
            assert all(item.run_id == run.run_id for item in window._report_items) if run else not window._report_items
        if not run:
            window.risk_non_reported.setChecked(True)
            assert not window._risk_items
    window._select_inspection_environment(a)
    window.password_input.setText("previous-session-secret")
    window._select_inspection_environment(b)
    assert not window.password_input.text()
    assert window._current_config().environment_id == b
    window.password_input.setText("session-for-old-address")
    selected = next(e for e in service.list_environment_archives(db) if e.environment_id == b)
    service.save_environment_archive(db, environment_id=b, name=selected.name, customer_name=selected.customer_name,
        customer_id=selected.customer_id, site_name=selected.site_name, address="192.0.2.119")
    window._refresh_all_data()
    assert window.vcenter_input.text() == "192.0.2.119"
    assert not window.password_input.text()
    window.password_input.setText("same-target-session")
    service.save_environment_archive(db, environment_id=b, name=selected.name, customer_name=selected.customer_name,
        customer_id=selected.customer_id, site_name="更名后的站点", address="192.0.2.119")
    window._refresh_all_data()
    assert window.site_input.text() == "更名后的站点"
    assert window.password_input.text() == "same-target-session"
    window.go_to_module("inspection_center")
    window._select_inspection_environment(a)
    assert window._environment_id == a
    assert window._caliber_run_id == first.run_id
    assert "客户甲" in window.current_object_label.text()
    assert window._connection_status_text("success") == "未实时检测"
    assert not any("PPT" in button.text() for button in window.findChildren(QPushButton))
    window.go_to_module("dashboard")
    assert window._environment_id is None
    window.close()


def test_inline_editor_registers_and_renames_without_ppt_or_password_persistence(archives, tmp_path, qt_application_lifetime):
    service, db, _, _, _ = archives
    legacy_ppt = tmp_path / "customer-owned.pptx"
    legacy_ppt.write_bytes(b"retained customer PowerPoint")
    service.save_config(DesktopInspectionConfig(db_path=db, report_output_dir=tmp_path))
    from vstacklens.desktop.app import MainWindow
    from PySide6.QtWidgets import QPushButton
    window = MainWindow(service=service)
    window.go_to_module("vcenter_management"); window.show()
    window._edit_environment(None)
    window.environment_name_input.setText("新增环境")
    window.environment_customer_input.setEditText("离线新客户")
    window.environment_address_input.setText("192.0.2.118")
    save = next(b for b in window.findChildren(QPushButton) if b.text() == "保存资料")
    save.click()
    created = next(e for e in service.list_environment_archives(db) if e.name == "新增环境")
    assert not window.environment_editor.isVisible()
    window.environment_edit_button.click()
    window.environment_name_input.setText("改名且保留标识")
    save.click()
    renamed = next(e for e in service.list_environment_archives(db) if e.name == "改名且保留标识")
    assert renamed.environment_id == created.environment_id
    window.environment_toggle_button.click()
    assert not next(e for e in service.list_environment_archives(db) if e.environment_id == created.environment_id).enabled
    assert legacy_ppt.read_bytes() == b"retained customer PowerPoint"
    window.close()


@pytest.mark.parametrize("failure", [ConnectionError("offline simulated connection failure"), PermissionError("NoPermission: offline access denied")])
def test_connection_diagnostic_is_not_reported_as_successful_inventory_collection(archives, tmp_path, monkeypatch, failure):
    service, db, a, _, _ = archives
    from vstacklens.application.inspection_runner import VCenterRunRequest
    from vstacklens.application.connection_probe import ProbeResult
    class OfflineFailureCollector:
        def __init__(self, *_args, **_kwargs):
            pass
        def precheck(self):
            raise failure
    monkeypatch.setattr("vstacklens.application.inspection_runner.PyVmomiCollector", OfflineFailureCollector)
    monkeypatch.setattr("vstacklens.application.inspection_runner.socket_probe", lambda *_: ProbeResult("failed", "离线模拟不可达"))
    InspectionRunner().run_vcenter(VCenterRunRequest(db_path=db, environment_id=a,
        vcenter="192.0.2.111", username="offline", password="offline-fake",
        rulepack_path=ROOT / "rulepacks/builtin-vsphere-v1", report_dir=tmp_path / "diagnostic"))
    env = next(e for e in service.list_environment_archives(db) if e.environment_id == a)
    assert "未采集到清单" in env.collection_result
    from vstacklens.application.connection_probe import friendly_connection_error
    assert friendly_connection_error(failure) in env.collection_result
    assert env.last_collected_at == ""
    assert "诊断报告" in env.report_result
