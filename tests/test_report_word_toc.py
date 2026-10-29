from pathlib import Path
from types import SimpleNamespace
import json
import os
import subprocess

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from vstacklens.reports import word_toc


def document(tmp_path):
    path=tmp_path/'generated.docx';doc=Document();doc.add_heading('合成测试数据',1)
    setting=OxmlElement('w:updateFields');setting.set(qn('w:val'),'true');doc.settings.element.append(setting)
    doc.add_paragraph('检查统计保持不变。');doc.save(path)
    return path


def test_ordinary_tests_never_launch_word(tmp_path,monkeypatch,caplog):
    path=document(tmp_path);original=path.read_bytes()
    monkeypatch.setattr(word_toc,'word_executable',lambda:pytest.fail('Word detection should be bypassed'))
    with caplog.at_level('INFO',logger='vstacklens.report_generation'):
        result=word_toc.update_word_toc(path)
    assert result['reason']=='disabled_for_automatic_tests' and path.read_bytes()==original
    assert 'open_on_update' in caplog.text


def test_no_word_fallback_keeps_standard_document_bytes(tmp_path,monkeypatch):
    path=document(tmp_path);original=path.read_bytes();monkeypatch.setenv('VSTACKLENS_WORD_TOC_MODE','auto')
    monkeypatch.setattr(word_toc,'word_executable',lambda:None)
    result=word_toc.update_word_toc(path)
    assert result['mode']=='open_on_update' and result['reason']=='microsoft_word_not_installed'
    assert path.read_bytes()==original


@pytest.mark.parametrize('failure',['timeout','worker','invalid','unsafe'])
def test_optional_word_failure_is_transactional_and_cleans_only_owned_process(tmp_path,monkeypatch,failure):
    path=document(tmp_path);original=path.read_bytes();monkeypatch.setenv('VSTACKLENS_WORD_TOC_MODE','auto')
    monkeypatch.setattr(word_toc,'word_executable',lambda:Path('WINWORD.EXE'))
    cleaned=[]
    monkeypatch.setattr(word_toc,'_cleanup_owned',lambda state:cleaned.append(state.copy()) or True)
    def run(command,**kwargs):
        state_path=Path(command[command.index('-StatePath')+1]);candidate=Path(command[command.index('-DocumentPath')+1])
        assert candidate!=path and candidate.read_bytes()==original
        state_path.write_text(json.dumps({'owned':True,'pid':8181,'creation_filetime':99,'status':'updated' if failure in {'invalid','unsafe'} else 'fallback'}),'utf-8')
        if failure=='timeout':raise subprocess.TimeoutExpired(command,kwargs['timeout'])
        if failure=='invalid':candidate.write_bytes(b'invalid archive')
        if failure=='unsafe':
            d=Document(candidate);d.add_paragraph('password=not-a-real-secret');d.save(candidate)
        return SimpleNamespace(returncode=1 if failure=='worker' else 0)
    monkeypatch.setattr(word_toc.subprocess,'run',run)
    result=word_toc.update_word_toc(path,timeout=.1)
    assert result['mode']=='open_on_update' and path.read_bytes()==original
    assert cleaned and cleaned[0]['pid']==8181 and result['owned_process_exited']


def test_pid_reuse_does_not_terminate_an_unrelated_word(monkeypatch):
    monkeypatch.setattr(word_toc,'_creation_time',lambda pid:100)
    monkeypatch.setattr(word_toc.subprocess,'run',lambda *a,**k:pytest.fail('Must not kill a reused PID'))
    assert word_toc._cleanup_owned({'owned':True,'pid':8181,'creation_filetime':99})


def test_identity_read_failure_is_not_reported_as_an_exited_word(monkeypatch):
    monkeypatch.setattr(word_toc,'_creation_time',lambda pid:-1)
    monkeypatch.setattr(word_toc.subprocess,'run',lambda *a,**k:pytest.fail('Unknown identity must not be killed'))
    assert not word_toc._cleanup_owned({'owned':True,'pid':8181,'creation_filetime':99})


@pytest.mark.skipif(os.environ.get('VSTACKLENS_RUN_WORD_NATIVE')!='1',reason='Native Word is an explicit separate integration test')
def test_native_word_toc_updates_in_private_process_and_exits(tmp_path,monkeypatch):
    if word_toc.word_executable() is None:pytest.skip('Microsoft Word not installed')
    monkeypatch.setenv('VSTACKLENS_WORD_TOC_MODE','auto')
    from vstacklens.reports.document_word import _field
    p=tmp_path/'native.docx';d=Document();settings=OxmlElement('w:updateFields');settings.set(qn('w:val'),'true');d.settings.element.append(settings)
    toc=d.add_paragraph();_field(toc,' TOC \\o "1-3" \\h \\z \\u ')
    d.add_heading('合成测试数据第一章',1);d.add_paragraph('正文。');d.add_page_break();d.add_heading('合成测试数据第二章',1);d.save(p)
    result=word_toc.update_word_toc(p)
    assert result['mode']=='native_word' and result['owned_process_exited'],result
    from zipfile import ZipFile
    with ZipFile(p) as z:
        text=z.read('word/document.xml').decode('utf-8')
        assert 'PAGEREF' in text and '合成测试数据第二章' in text
