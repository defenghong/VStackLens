"""Management report counterexamples; synthetic evidence, not live faults."""
from copy import deepcopy
import pytest
from vstacklens.products.model import TargetRef, new_module, make_check, make_source, finish_module
from vstacklens.products.management_scope import mark_management_scope
from vstacklens.products.reporting import build_integrated_report, product_report_view, _capacity_display


def fixture(product="avi", checks=None, alerts=(), management=True):
    target=TargetRef("report-offline-"+product,product,"offline-customer","offline-site","离线管理设备长名称","192.0.2.109")
    module=new_module(target)
    if management: mark_management_scope(module)
    module["checks"]=checks or [make_check("AVI-A1" if product=="avi" else "nsx.version","管理版本","normal",values={"version":"offline"},sources=[make_source("/offline",version_basis="离线合成")],verification="offline_sample")]
    module["native_alerts"]=list(alerts)
    module=finish_module(module)
    return build_integrated_report("offline-report-manager",[target],[module],"管理端离线验证")


@pytest.mark.parametrize("product,key,cleared,active",[("avi","state","ALERT_STATE_OFF","ALERT_STATE_ON"),("nsx","status","RESOLVED","OPEN")])
def test_c1_lifecycle_separates_cleared_active_unknown_and_preserves_severity(product,key,cleared,active):
    alerts=[{key:state,"severity":"MEDIUM","object_name":name,"description":name+"原生说明"} for state,name in [(cleared,"已清除对象"),(active,"活动对象"),("UNRECOGNIZED","未知对象")]]
    report=fixture(product,alerts=alerts);before=report.model_dump()
    view=product_report_view(report);body=view["sections"][0]["body"]
    assert [c["object"] for c in body["concerns"]]==["活动对象","未知对象"]
    assert "MEDIUM" in body["concerns"][0]["status"] and "活动" in body["concerns"][0]["status"]
    assert "状态待确认" in body["concerns"][1]["status"]
    history=next(t for t in body["tables"] if t["title"]=="已清除告警历史")
    assert history["rows"][0][0]=="已清除对象" and cleared in history["rows"][0][2]
    assert report.model_dump()==before and report.risk_summary.total==0


def resource_check(metrics):
    return make_check("AVI-A3","设备内部资源","insufficient_data",values={"metrics":metrics,"partitions":None},sources=[make_source("/offline/controller-metrics",version_basis="离线合成")],verification="offline_sample")


def test_c3_nonempty_controller_sample_is_in_shared_body_with_window_and_missing_count():
    metric={"object_id":"node-1","metric_id":"controller_stats.avg_cpu_usage","unit":"PERCENT","requested_window":{"start":"2026-10-18T00:00:00+00:00","stop":"2026-10-18T00:14:00+00:00","step":60},"returned_window":{},"returned_sample_count":1,"samples":[{"timestamp":"2026-10-18T00:14:00+00:00","value":12.345678}],"missing_samples":14,"complete":False}
    base=fixture();valid=deepcopy(base.product_modules[0]["checks"][0])
    report=fixture(checks=[valid,resource_check([metric])]);original=deepcopy(report.product_modules)
    body=product_report_view(report)["sections"][0]["body"]
    text=str(body["tables"])
    for expected in ("controller_stats.avg_cpu_usage","12.345678 %","缺样 14","00:14:00","间隔 60 秒","窗口不完整"):
        assert expected in text
    assert report.product_modules==original


def test_missing_samples_are_not_zero_and_management_scope_leaks_reject():
    valid=deepcopy(fixture().product_modules[0]["checks"][0])
    body=product_report_view(fixture(checks=[valid,resource_check([])]))["sections"][0]["body"]
    assert "未返回资源样本" in str(body) and "0 %" not in str(body)
    report=fixture();report.product_modules[0]["checks"].append(make_check("AVI-C5","后端池","not_deployed",values={"count":0}))
    with pytest.raises(ValueError,match="范围外"):
        product_report_view(report)


def test_legacy_scope_is_not_silently_marked_or_filtered():
    report=fixture(management=False);report.product_modules[0]["checks"].append(make_check("AVI-C5","旧业务范围","not_deployed",values={"count":0}))
    before=report.model_dump();view=product_report_view(report)
    assert any(c["title"]=="旧业务范围" for c in view["sections"][0]["checks"])
    assert report.model_dump()==before


def test_capacity_conversion_has_explicit_si_and_iec_units_and_no_fake_zero():
    assert _capacity_display(1048576,"bytes")=="1.00 MiB"
    assert _capacity_display(1048576,"KiB")=="1.00 GiB"
    assert _capacity_display(1048576,"kilobytes")=="1000.00 MiB"
    assert _capacity_display(None,"kilobytes")=="未采集"
    assert "未换算" in _capacity_display(1,"unknown")


def test_real_three_formats_share_sample_and_alarm_lifecycle(tmp_path):
    import json
    import pymupdf
    from docx import Document
    from vstacklens.reports.html_package import HtmlReportPackageBuilder
    from vstacklens.reports.exporter import ReportExportEngine
    from vstacklens.reports.pdf_report import PdfReportEngine
    metric={"metric_id":"controller_stats.avg_cpu_usage","unit":"PERCENT","requested_window":{"start":"2026-10-18T00:00:00Z","stop":"2026-10-18T00:14:00Z","step":60},"returned_sample_count":1,"samples":[{"timestamp":"2026-10-18T00:14:00Z","value":12.345678}],"missing_samples":14,"complete":False}
    valid=deepcopy(fixture().product_modules[0]["checks"][0])
    report=fixture(checks=[valid,resource_check([metric])],alerts=[{"state":"ALERT_STATE_OFF","severity":"MEDIUM","description":"离线已清除事件"},{"state":"ALERT_STATE_ON","severity":"MEDIUM","description":"离线活动事件"}])
    before=report.model_dump()
    html,_=HtmlReportPackageBuilder().render(report,{},tmp_path/"html")
    word,_,_=ReportExportEngine().render(report,tmp_path/"report.docx")
    pdf=PdfReportEngine().render(report,tmp_path/"report.pdf")
    payload=json.loads((html.parent/"data/customer_report_payload.json").read_text(encoding="utf-8"))
    assert payload["productReport"]==product_report_view(report)
    doc=Document(word)
    word_text="\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    with pymupdf.open(pdf) as rendered: pdf_text="\n".join(p.get_text() for p in rendered)
    for text in (word_text,pdf_text,json.dumps(payload,ensure_ascii=False)):
        normalized="".join(text.replace("\u200b", "").split())
        for expected in ("12.345678","controller_stats.avg_cpu_usage","缺样14","MEDIUM","已清除告警历史","ALERT_STATE_OFF","活动"):
            assert expected in normalized
        assert "ServiceEngine" not in normalized and "后端池" not in normalized
    assert report.model_dump()==before
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if "E-" in cell.text: assert cell._tc.xpath("./w:tcPr/w:noWrap")


def test_wholly_unavailable_still_blocks_customer_health_report():
    from vstacklens.products.reporting import NoValidInspectionData
    with pytest.raises(NoValidInspectionData): fixture(checks=[resource_check([])])


def test_partial_alarm_read_failure_remains_visible_even_if_returned_history_is_cleared():
    valid=deepcopy(fixture().product_modules[0]["checks"][0])
    failed=make_check("AVI-A4","原生告警读取","collection_failed",values={"partial":True},summary="后续告警页读取失败，结果不完整",sources=[make_source("/offline/alerts")])
    body=product_report_view(fixture(checks=[valid,failed],alerts=[{"state":"ALERT_STATE_OFF","severity":"MEDIUM","description":"历史已清除"}]))["sections"][0]["body"]
    assert any(c["status"]=="采集失败" and "后续告警页" in c["observation"] for c in body["concerns"])
    assert not any(c["title"]=="设备原生告警" for c in body["concerns"])


def test_failed_vcenter_target_is_not_removed_from_a_mixed_manager_scope():
    t=TargetRef("offline-vc","vcenter","offline-customer","offline-site","失败VC","192.0.2.111")
    vc=new_module(t);vc["checks"]=[make_check("execution.failure","读取失败","collection_failed",values={"reason":"offline"})]
    avi=fixture().product_modules[0]
    a=TargetRef(avi["target_id"],"avi","offline-customer","offline-site",avi["target_name"],"192.0.2.109")
    report=build_integrated_report("offline-mixed",[t,a],[finish_module(vc),avi],"混合范围")
    view=product_report_view(report)
    assert "VC" in view["scope"] and any(s["product"]=="vcenter" for s in view["sections"])


def test_management_cloud_is_configuration_not_runtime_health():
    report=fixture();report.product_modules[0]["inventory"]=[{"object_type":"cloud","object_id":"cloud-registration","object_name":"配置登记Cloud","values":{"name":"配置登记Cloud","vtype":"CLOUD_VCENTER"}}]
    text=str(product_report_view(report)["sections"][0]["body"]["tables"])
    assert "Cloud 登记配置" in text and "不据此判断运行连接状态" in text
    assert "运行详情见证据" not in text
