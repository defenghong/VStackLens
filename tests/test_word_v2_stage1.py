from __future__ import annotations

from pathlib import Path

from docx import Document

from test_report_exporter import rich_report
from vstacklens.reports.word_v2_renderer import WordV2Renderer
from vstacklens.reports.word_v2_view_model import STAGE1_CATEGORY_IDS, WordV2ViewModelBuilder


def _category(category_id: str, level: str, title: str, section: str, objects: list[dict], evidence: list[dict]) -> dict:
    return {
        "category_id": category_id,
        "priority": level,
        "risk_level": level,
        "title": title,
        "description": f"经检查，{title}。",
        "remediation": f"建议结合业务和维护窗口处理{title}。",
        "section": section,
        "affected_object_count": len(objects),
        "affected_objects": objects,
        "evidence_items": evidence,
    }


def stage1_report():
    report = rich_report()
    categories = [
        _category(
            "CLUSTER-DRS", "P1", "集群 DRS 配置问题", "Cluster",
            [{"object_key": "cluster-1", "object_name": "Cluster01", "object_type": "ClusterComputeResource", "object_path": "VC01 / DC01 / Cluster01"}],
            [{"object_key": "cluster-1", "object_name": "Cluster01", "observed_value": "未启用", "reference_value": "启用", "judgement_text": "Cluster01 当前 DRS 未启用。"}],
        ),
        _category(
            "HOST-PNIC", "P1", "ESXi 主机物理网卡异常", "Host",
            [{"object_key": "host-1", "object_name": "esxi01", "object_type": "HostSystem", "object_path": "VC01 / DC01 / Cluster01 / esxi01"}],
            [{"object_key": "host-1", "object_name": "esxi01", "observed_value": 1, "judgement_text": "esxi01 存在物理网卡链路异常。"}],
        ),
        _category(
            "VM-CDROM", "P2", "虚拟机挂载 ISO 镜像", "Virtual Machine",
            [{"object_key": "vm-1", "object_name": "VM01", "object_type": "VirtualMachine", "object_path": "VC01 / DC01 / VM01"}],
            [{"object_key": "vm-1", "object_name": "VM01", "judgement_text": "VM01 当前存在 ISO 挂载。"}],
        ),
        _category(
            "VM-SNAPSHOT", "P2", "虚拟机存在快照", "Virtual Machine",
            [{"object_key": "vm-2", "object_name": "VM02", "object_type": "VirtualMachine", "object_path": "VC01 / DC01 / VM02"}],
            [{"object_key": "vm-2", "object_name": "VM02", "observed_value": 12, "judgement_text": "VM02 快照存在 12 天。"}],
        ),
    ]
    report.problem_categories = categories
    report.report_presentation["problem_categories"] = categories
    report.report_presentation["risk_summary"] = {"P1": 2, "P2": 2, "P3": 0}
    report.report_presentation["environment_scale"] = {"vcenter": 1, "datacenter": 1, "cluster": 1, "host": 2, "vm": 2, "datastore": 2, "vsan_cluster": 1}
    report.vsan_summary = {
        "applicable": True,
        "health_status": "healthy",
        "architecture": "OSA",
        "host_count": 2,
        "disk_group_count": 2,
        "cache_disk_count": 2,
        "capacity_disk_count": 6,
        "object_count": 12,
        "vmdk_count": 8,
        "resync_object_count": 0,
        "storage_policy_summary": {"checked_count": 8, "compliant_count": 8},
        "capacity": {"total_gb": 1000, "used_gb": 300, "used_percent": 30},
    }
    return report


def test_stage1_view_model_sorts_engineering_order_and_keeps_customer_titles() -> None:
    model = WordV2ViewModelBuilder(category_ids=STAGE1_CATEGORY_IDS).build(stage1_report())

    assert [item.title for item in model.risk_items] == [
        "集群 DRS 配置问题",
        "ESXi 主机物理网卡异常",
        "虚拟机挂载 ISO 镜像",
        "虚拟机存在快照",
    ]
    assert all(item.source_category_id for item in model.risk_items)
    assert model.vsan.applicable is True
    assert not model.validation_errors


def test_stage1_renderer_has_only_confirmed_chapters_and_no_internal_text(tmp_path: Path) -> None:
    output = tmp_path / "word-v2-stage1.docx"
    model = WordV2ViewModelBuilder(category_ids=STAGE1_CATEGORY_IDS).build(stage1_report())
    WordV2Renderer().render(model, output)

    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    for heading in ("1. 巡检概况", "2. 巡检结果总体分析", "3. 风险项详细说明及优化建议", "4. vSAN 专项巡检", "6. 巡检结论与建议"):
        assert heading in text
    for forbidden in ("首页摘要", "VMware 平台巡检", "环境资产与对象清单", "source_category_id", "CLUSTER-DRS", "VSL-", "子原因"):
        assert forbidden not in text
    assert "虚拟机挂载 ISO 镜像" in text
    assert "虚拟机存在快照" in text
    assert "集群 DRS 配置问题" in text
    assert "ESXi 主机物理网卡异常" in text
