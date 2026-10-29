"""Third-batch Word presentation consumes common words and display formatting."""
from copy import deepcopy
from datetime import datetime

from docx import Document

from vstacklens.reports.customer_format import CAPACITY_NOTE
from vstacklens.reports.docx_report import (
    CUSTOMER_BASIS_WIDTHS_DXA, ISO_DETAIL_WIDTHS_DXA, DocxReportEngine, RiskGroup,
)
from vstacklens.reports.report_model import (
    CustomerInfo, EnvironmentSummary, FindingItem, HealthScore, ReportData, ReportInfo, RiskSummary,
)


def _report(details=None):
    return ReportData(
        report_info=ReportInfo(report_id="batch3-word", run_id="batch3-word", generated_at="2026-10-22T06:57:12Z"),
        customer_info=CustomerInfo(customer_name="客户", site_name="实验环境", vcenter="vc.test"),
        environment_summary=EnvironmentSummary(scope_statement="既有数据", checked_object_total=1),
        health_score=HealthScore(score=80, grade="关注", explanation="既有结论"),
        risk_summary=RiskSummary(), asset_inventory={"details": details or {}, "summary": {}},
    )


def _text(document):
    return "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )


def test_resource_steps_consume_shared_groups_and_restart_numbering(monkeypatch):
    display = {
        "current": "3 台配置预留，5 台配置资源限制。", "expected": "", "impact": "共同影响。",
        "remediation": "NSX Manager 的资源预留由产品部署设定，不建议降低，保持现状即可。",
        "steps": ["不得读取已废弃的平铺步骤"],
        "step_groups": [
            {"title": "预留", "steps": ["识别预留。", "确认需求。", "调整预留。", "观察资源。"]},
            {"title": "资源限制", "steps": ["检查限制。", "确认用途。", "调整限制。", "复核性能。"]},
        ],
    }
    before = deepcopy(display)
    engine = DocxReportEngine()
    monkeypatch.setattr(engine, "_shared_problem", lambda *_: display)
    document = Document()
    engine._add_risk_sentence(document, RiskGroup("P3", "虚拟机", "资源配置", []), _report())
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert paragraphs[1:6] == ["预留", "1. 识别预留。", "2. 确认需求。", "3. 调整预留。", "4. 观察资源。"]
    assert paragraphs[6:11] == ["资源限制", "1. 检查限制。", "2. 确认用途。", "3. 调整限制。", "4. 复核性能。"]
    assert "5." not in _text(document)
    assert "不得读取" not in _text(document)
    assert display["remediation"] in _text(document)
    assert display == before


def test_alarm_current_and_standard_have_distinct_shared_labels(monkeypatch):
    display = {
        "current_label": "当前值", "expected_label": "判定标准",
        "current": "当前有 1 条红色活动告警：Datacenters 的“Backup job status”告警。",
        "expected": "红色活动告警应为 0 条。", "impact": "共同影响。", "remediation": "共同建议。",
        "steps": ["查看告警。"],
    }
    before = deepcopy(display)
    engine = DocxReportEngine()
    monkeypatch.setattr(engine, "_shared_problem", lambda *_: display)
    document = Document()
    engine._add_risk_sentence(document, RiskGroup("P2", "vCenter", "告警", []), _report())
    text = _text(document)
    assert f"当前值：{display['current']}" in text
    assert f"判定标准：{display['expected']}" in text
    assert "当前状态" not in text and "参考 0" not in text
    assert display == before


def test_storage_table_keeps_mixed_units_in_gb_and_half_up_without_mutating_data(monkeypatch):
    report = _report({"Datastore": [
        {"object_name": "Data_66", "properties": {"datastore_capacity_gb": 931.25, "used_percent": 11.45,
         "free_percent": 88.55, "datastore_filesystem_type": "VMFS", "accessible": True}},
        {"object_name": "Name-With-Datastore", "properties": {"datastore_capacity_gb": 2048,
         "used_percent": 10, "free_percent": 90, "datastore_filesystem_type": "VMFS", "accessible": True}},
    ]})
    before = deepcopy(report.model_dump())
    engine = DocxReportEngine()
    engine._customer_mode = True
    monkeypatch.setattr(engine, "_version_certificate_license_rows", lambda _: [])
    document = Document()
    engine._asset_overview(document, report)
    stores = next(table for table in document.tables if table.cell(0, 0).text == "Datastore")
    rows = [[cell.text for cell in row.cells] for row in stores.rows]
    assert rows[0] == ["Datastore", "类型", "总容量", "已用", "剩余", "状态"]
    assert rows[1][0:4] == ["Data_66", "VMFS", "931.3 GB", "11.5%"]
    assert rows[2][0] == "Name-With-Datastore"
    assert rows[2][2] == "2,048.0 GB"
    assert _text(document).count(CAPACITY_NOTE) == 1
    assert report.model_dump() == before


def test_vsan_capacity_labels_use_one_table_unit_and_note_is_not_repeated():
    summary = {"capacity": {"total_gb": 32192.75, "used_gb": 3682.5, "free_gb": 28510.25, "used_percent": 11.44}}
    before = deepcopy(summary)
    engine = DocxReportEngine()
    document = Document()
    engine._add_capacity_note(document)
    engine._add_vsan_capacity_chart(document, summary)
    text = _text(document)
    assert "总容量 31.4 TB；已用 3.6 TB；可用 27.8 TB；使用率 11.4%" in text
    assert "已用 3.6 TB" in text and "可用 27.8 TB" in text
    assert text.count(CAPACITY_NOTE) == 1
    assert summary == before


def test_snapshot_time_is_local_but_snapshot_name_and_pure_dates_are_unchanged():
    engine = DocxReportEngine()
    original = {"name": "虚拟机快照 2026/9/10 11:11:10", "create_time": "2026-09-10T03:11:22.427854+00:00"}
    expected = datetime.fromisoformat(original["create_time"]).astimezone().strftime("%Y-%m-%d %H:%M")
    before = deepcopy(original)
    assert engine._snapshot_time_display(original) == expected
    assert engine._snapshot_name_label(original) == original["name"]
    assert engine._time_label("2031-05-25") == "2031-05-25"
    assert engine._date_label(original["create_time"]) == expected[:10]
    assert original == before


def test_missing_capacity_does_not_turn_into_zero_or_create_snapshot_capacity():
    engine = DocxReportEngine()
    assert engine._snapshot_size_display({"name": "真实快照", "create_time": "2026-10-22"}) == ""
    assert engine._gb_size_label(None) == "未采集"
    assert engine._resync_bytes_label(None) == "未采集"
    assert engine._resync_bytes_label(0) == "0.0 GB"
    assert engine._gb_size_label(931.25) == "931.3 GB"
    assert engine._fmt_percent(11.45) == "11.5%"


def test_updated_labels_keep_existing_table_widths_and_data():
    engine = DocxReportEngine()
    assert engine._detail_widths(["虚拟机名称", "ISO 镜像"]) == ISO_DETAIL_WIDTHS_DXA
    assert engine._detail_widths(["对象", "所在范围", "判断依据"]) == CUSTOMER_BASIS_WIDTHS_DXA
    shared = {"current": "", "expected": "", "impact": "", "remediation": "", "steps": [],
              "basis_label": "判断依据", "objects": [{"name": "原始名称", "level": "P3", "location": "原始位置", "basis": "共同事实。"}]}
    engine._shared_problem = lambda *_: shared
    assert engine._detail_table(_report(), RiskGroup("P3", "虚拟机", "标题", [])) == (
        ["对象", "所在范围", "判断依据"], [["原始名称（P3）", "原始位置", "共同事实。"]]
    )


def test_feature_and_ssh_adapters_describe_different_kinds_of_state():
    engine = DocxReportEngine()
    report = _report({
        "ClusterComputeResource": [{"object_name": "Cluster-A", "properties": {"ha_enabled": False, "drs_enabled": True}}],
        "HostSystem": [{"object_name": "host-a", "properties": {"ssh_running": True}}],
    })
    cluster_finding = FindingItem(risk_level="P2", rule_id="VSL-CL-001", rule_name="HA", title="HA",
                                 object_type="ClusterComputeResource", object_name="Cluster-A", status="open", current_value=False)
    host_finding = FindingItem(risk_level="P3", rule_id="VSL-HOST-001", rule_name="SSH", title="SSH",
                              object_type="HostSystem", object_name="host-a", status="open", current_value=True)
    before = deepcopy(report.model_dump())
    assert engine._cluster_feature_table(report, RiskGroup("P2", "集群", "HA", [cluster_finding]), "ha_enabled", "HA 当前状态")[1][0][1] == "vSphere HA 未启用"
    assert engine._cluster_feature_table(report, RiskGroup("P3", "集群", "DRS", [cluster_finding]), "drs_enabled", "DRS 当前状态")[1][0][1] == "DRS 已启用"
    headers, rows = engine._host_ssh_table(report, RiskGroup("P3", "主机", "SSH", [host_finding]))
    assert headers[-1] == "SSH 服务状态" and rows[0][-1] == "正在运行"
    assert report.model_dump() == before


def test_missing_total_capacity_hides_whole_column_but_known_zero_keeps_it(monkeypatch):
    engine = DocxReportEngine()
    engine._customer_mode = True
    monkeypatch.setattr(engine, "_version_certificate_license_rows", lambda _: [])
    missing = _report({"Datastore": [{"object_name": "DS-no-size", "properties": {
        "used_percent": 25, "free_percent": 75, "accessible": True,
    }}]})
    missing_before = deepcopy(missing.model_dump())
    document = Document()
    engine._asset_overview(document, missing)
    store = next(table for table in document.tables if table.cell(0, 0).text == "Datastore")
    assert [cell.text for cell in store.rows[0].cells] == ["Datastore", "类型", "已用", "剩余", "状态"]
    assert "总容量" not in _text(document) and CAPACITY_NOTE not in _text(document)
    assert "0.0 GB" not in _text(document)
    assert missing.model_dump() == missing_before

    zero = missing.model_copy(deep=True)
    zero.asset_inventory["details"]["Datastore"][0]["properties"]["datastore_capacity_gb"] = 0
    zero_before = deepcopy(zero.model_dump())
    zero_doc = Document()
    engine._asset_overview(zero_doc, zero)
    store = next(table for table in zero_doc.tables if table.cell(0, 0).text == "Datastore")
    assert [cell.text for cell in store.rows[0].cells] == ["Datastore", "类型", "总容量", "已用", "剩余", "状态"]
    assert store.rows[1].cells[2].text == "0.0 GB"
    assert _text(zero_doc).count(CAPACITY_NOTE) == 1
    assert zero.model_dump() == zero_before


def test_summary_recommendation_reuses_the_same_shared_sentence_as_problem_body(monkeypatch):
    engine = DocxReportEngine()
    finding = FindingItem(risk_level="P3", rule_id="VSL-VM-023", rule_name="本地存储", title="本地存储",
                          object_type="VirtualMachine", object_name="Datastore-in-real-name", status="open",
                          remediation="旧建议：必要时迁移到共享 Datastore。")
    shared = {"title": "共同标题", "remediation": "结合业务用途、备份和故障恢复要求人工确认，不直接判定必须迁移。"}
    before = deepcopy(finding.model_dump())
    monkeypatch.setattr(engine, "_shared_problem", lambda *_: shared)
    rows = engine._summary_rows([RiskGroup("P3", "虚拟机", "本地存储", [finding])], report=_report())
    assert rows == [["P3", "虚拟机", shared["title"], "1", shared["remediation"]]]
    assert "Datastore" not in rows[0][-1]
    assert finding.model_dump() == before
