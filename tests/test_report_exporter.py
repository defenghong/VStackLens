from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZipFile
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from vstacklens.application.inspection_runner import InspectionRunner, MockRunRequest, VCenterRunRequest
from vstacklens.cli import cmd_run_mock
from vstacklens.collection.mock_collector import MockCollector
from vstacklens.db.connection import connect
from vstacklens.reports.docx_report import DocxReportEngine
from vstacklens.reports.exporter import ReportExportEngine
from vstacklens.reports.html_report import HtmlReportEngine
from vstacklens.reports.html_package import HtmlReportPackageBuilder
from vstacklens.reports.pdf_report import PdfReportEngine
from vstacklens.findings.evidence_builder import EvidenceBuilder
from vstacklens.reports.report_model import (
    CustomerInfo,
    EnvironmentSummary,
    FindingItem,
    HealthScore,
    ReportDataFactory,
    RemediationItem,
    ReportData,
    ReportInfo,
    RiskSummary,
)


ROOT = Path(__file__).resolve().parents[1]
RULEPACK = ROOT / "rulepacks" / "builtin-vsphere-v1"
FIXTURE = ROOT / "tests" / "fixtures" / "inventory_pass.json"
FORBIDDEN_WORD_TERMS = [
    "立即整改",
    "紧急整改",
    "必须整改",
    "严重威胁",
    "授权服务",
    "不可判定项",
    "无不可判定项",
    "更新域",
    "更新目录",
    "报告生成说明",
    "规则覆盖清单",
    "巡检范围",
    "风险统计",
    "主要关注方向",
    "失败",
]


def rich_report() -> ReportData:
    return ReportData(
        report_info=ReportInfo(report_id="report-test", generated_at="2026-06-03T00:00:00+08:00", run_id="run-test"),
        customer_info=CustomerInfo(customer_name="交付客户", site_name="生产站点", vcenter="vcsa.test.local"),
        environment_summary=EnvironmentSummary(scope_statement="本报告基于本次 vSphere SDK 巡检结果生成。", checked_object_total=8),
        health_score=HealthScore(score=82, grade="B", explanation="环境存在少量需要优先处理的风险。"),
        risk_summary=RiskSummary(P1=1, P2=1, P3=1, P4=1, total=4),
        risk_distribution={"by_rule": {"VSL-VC-001": 1}},
        asset_inventory={
            "summary": {"vCenter": 1, "ClusterComputeResource": 1, "HostSystem": 2, "VirtualMachine": 3, "Datastore": 2},
            "details": {
                "vCenter": [
                    {
                        "object_name": "vcsa.test.local",
                        "properties": {"connected": True, "version": "8.0.3", "build": "24853646", "certificate_days_remaining": 716, "license_days_remaining": 999999},
                        "asset_location": "vCenter 根对象",
                    }
                ],
                "ClusterComputeResource": [
                    {
                        "object_name": "Cluster-A",
                        "properties": {"host_count": 2, "ha_enabled": True, "drs_enabled": True},
                        "asset_location": "vcsa.test.local / Datacenter / Cluster-A",
                    }
                ],
                "HostSystem": [
                    {
                        "object_name": "esxi-01",
                        "properties": {
                            "asset_location": "vcsa.test.local / Datacenter / Cluster-A",
                            "connection_state": "connected",
                            "cpu_usage_avg": 20.5,
                            "memory_usage_avg": 48.2,
                            "ssh_running": True,
                            "syslog_configured": False,
                            "host_license_name": "vSphere 7 Enterprise Plus",
                            "host_certificate_days_remaining": 1811,
                            "host_license_expiration_days": 999999,
                        },
                        "asset_location": "vcsa.test.local / Datacenter / Cluster-A",
                    },
                    {
                        "object_name": "esxi-02",
                        "properties": {
                            "asset_location": "vcsa.test.local / Datacenter / Cluster-A",
                            "connection_state": "connected",
                            "cpu_usage_avg": 31.2,
                            "memory_usage_avg": 52.4,
                            "host_license_name": "vSphere 7 Enterprise Plus",
                            "host_certificate_days_remaining": 1810,
                            "host_license_expiration_days": 999999,
                        },
                        "asset_location": "vcsa.test.local / Datacenter / Cluster-A",
                    },
                ],
                "Datastore": [
                    {
                        "object_name": "DS01",
                        "properties": {"datastore_filesystem_type": "VMFS", "used_percent": 72.5, "free_percent": 27.5, "accessible": True},
                        "asset_location": "vcsa.test.local / Datacenter / 挂载集群 Cluster-A",
                    },
                    {
                        "object_name": "DS02",
                        "properties": {"datastore_filesystem_type": "VMFS", "used_percent": 44.0, "free_percent": 56.0, "accessible": True},
                        "asset_location": "vcsa.test.local / Datacenter / 挂载集群 Cluster-A",
                    },
                ],
                "VirtualMachine": [
                    {
                        "object_name": "app-01",
                        "properties": {"asset_location": "vcsa.test.local / Datacenter / Cluster-A", "power_state": "poweredOn", "tools_running": False, "vcpu_count": 4},
                        "asset_location": "vcsa.test.local / Datacenter / Cluster-A",
                    }
                ],
            },
            "total": 9,
        },
        findings=[
            FindingItem(
                risk_level="P1",
                rule_id="VSL-VC-004",
                rule_name="vCenter 告警状态",
                title="vCenter 存在红色活动告警",
                object_type="vCenter",
                object_name="vcsa.test.local",
                status="open",
                current_value=1,
                expected_value=0,
                business_impact="重要故障可能仍处于活动状态，需要优先确认影响范围。",
                remediation="逐项检查告警根因，确认恢复后清理已解决告警。",
                observed_detail={"active_red_alarms": [{"alarm_name": "Host memory status", "entity_name": "esxi-01", "status": "红色"}]},
                recommended_action_zh="逐项打开红色告警，确认对象状态和根因。",
            ),
            FindingItem(
                risk_level="P2",
                rule_id="VSL-HOST-001",
                rule_name="ESXi SSH 服务",
                title="ESXi 主机 SSH 服务已开启",
                object_type="HostSystem",
                object_name="esxi-01",
                status="open",
                current_value=True,
                expected_value=False,
                business_impact="长期开启远程维护入口会增加误操作和安全管理风险。",
                remediation="结合运维需要评估关闭或设置超时策略。",
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-VM-004",
                rule_name="VMware Tools 状态",
                title="VMware Tools 未运行",
                object_type="VirtualMachine",
                object_name="app-01",
                status="open",
                current_value=False,
                expected_value=True,
                business_impact="虚拟机运行状态和驱动集成信息可能无法及时反馈。",
                remediation="确认系统状态后启动或重新安装 VMware Tools。",
            ),
            FindingItem(
                risk_level="P4",
                rule_id="VSL-CL-016",
                rule_name="集群 vMotion 网络覆盖",
                title="集群 vMotion 网络覆盖不完整",
                object_type="ClusterComputeResource",
                object_name="Cluster-A",
                status="open",
                current_value=1,
                expected_value=0,
                business_impact="集群内部分主机缺少 vMotion 网络会限制在线迁移能力。",
                remediation="为缺少 vMotion 的主机补齐 VMkernel 网络配置。",
                observed_detail={"missing_vmotion_hosts": ["esxi-02"]},
            ),
        ],
        recommendations=["优先处理 P1/P2 风险，并在处理后复跑巡检确认。"],
        remediation_plan=[
            RemediationItem(
                risk_level="P1",
                rule_id="VSL-VC-001",
                title="恢复 vCenter API 访问",
                object_count=1,
                owner_role="virtualization_admin",
                effort="medium",
                verification_method="复跑巡检确认",
            )
        ],
        appendix={
            "certificate_license_evidence": [
                {
                    "rule_id": "VSL-HOST-023",
                    "result_status": "passed",
                    "result_status_label": "通过",
                    "object_type": "HostSystem",
                    "object_type_label": "ESXi",
                    "object_name": "esxi-01",
                    "current_value_zh": "正式授权",
                    "expected_value_zh": "已分配有效授权",
                    "source_path": "host product information",
                    "collected_at": "2026-06-03T00:00:00+08:00",
                },
                {
                    "rule_id": "VSL-VC-006",
                    "result_status": "passed",
                    "result_status_label": "通过",
                    "object_type": "vCenter",
                    "object_type_label": "vCenter",
                    "object_name": "vcsa.test.local",
                    "current_value_zh": "正式授权",
                    "expected_value_zh": "vCenter 授权有效",
                    "source_path": "licenseAssignmentManager",
                    "collected_at": "2026-06-03T00:00:00+08:00",
                },
            ],
            "unavailable_rules": [
                {
                    "rule_id": "VSL-TEST-UNAVAILABLE",
                    "title": "测试不可判定规则",
                    "object_name": "test-object",
                    "data_quality": "missing",
                    "reason": "测试数据缺失",
                }
            ],
            "rule_checklist": [],
        },
    )


def minimal_report() -> ReportData:
    return rich_report()


def table_layout_report() -> ReportData:
    report = rich_report().model_copy(deep=True)
    cluster_props = report.asset_inventory["details"]["ClusterComputeResource"][0]["properties"]
    cluster_props["ha_enabled"] = False
    cluster_props["drs_enabled"] = False
    report.findings.extend(
        [
            FindingItem(
                risk_level="P2",
                rule_id="VSL-CL-001",
                rule_name="Cluster HA disabled",
                title="集群 HA 未启用",
                object_type="ClusterComputeResource",
                object_name="Cluster-A",
                status="open",
                current_value=False,
                expected_value=True,
                business_impact="集群级高可用保护不可用。",
                remediation="评估集群容量后启用 HA。",
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-CL-003",
                rule_name="Cluster DRS disabled",
                title="集群 DRS 未启用",
                object_type="ClusterComputeResource",
                object_name="Cluster-A",
                status="open",
                current_value=False,
                expected_value=True,
                business_impact="集群无法自动均衡工作负载。",
                remediation="检查虚拟机规则后启用 DRS。",
            ),
            FindingItem(
                risk_level="P2",
                rule_id="VSL-HOST-019",
                rule_name="ESXi host physical NIC link issue",
                title="ESXi 主机物理网卡链路异常",
                object_type="HostSystem",
                object_name="esxi-01",
                status="open",
                current_value=1,
                expected_value=0,
                business_impact="主机网络链路异常可能影响业务或迁移网络。",
                remediation="定位异常 vmnic 和对应物理交换机端口。",
                observed_detail={"host_name": "esxi-01", "link_speed_detail": [{"device": "vmnic1", "status": "down", "speed_mb": 0}]},
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-HOST-014",
                rule_name="ESXi host syslog configuration",
                title="ESXi 主机未配置远程 Syslog",
                object_type="HostSystem",
                object_name="esxi-01",
                status="open",
                current_value=False,
                expected_value=True,
                business_impact="故障追溯和审计证据可能不足。",
                remediation="为 ESXi 主机配置企业统一 Syslog 服务。",
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-HOST-CORE",
                rule_name="ESXi host log and core dump configuration",
                title="ESXi 主机日志与核心转储配置不完整",
                object_type="HostSystem",
                object_name="esxi-01",
                status="open",
                current_value="远程日志未配置；核心转储未配置",
                expected_value="远程日志和核心转储配置完整",
                business_impact="日志和核心转储缺失可能影响故障后分析。",
                remediation="配置远程 Syslog 或持久化日志目录，并确认核心转储目标可用。",
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-VM-001",
                rule_name="VM old snapshot",
                title="虚拟机存在超期快照",
                object_type="VirtualMachine",
                object_name="app-01",
                status="open",
                current_value=33,
                expected_value=7,
                business_impact="长期快照会占用存储并影响虚拟机性能。",
                remediation="确认快照用途后删除或合并快照。",
                observed_detail={"snapshots": [{"name": "before-upgrade", "create_time": "2026-05-01T00:00:00+08:00", "size_mb": 2048}]},
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-VM-TOOLS-001",
                rule_name="VMware Tools install state",
                title="虚拟机未安装 VMware Tools",
                object_type="VirtualMachine",
                object_name="app-01",
                status="open",
                current_value=False,
                expected_value=True,
                business_impact="客户机状态采集和优雅关机能力受限。",
                remediation="为支持的客户机安装 VMware Tools 或 open-vm-tools，并复核运行状态。",
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-VM-TOOLS-002",
                rule_name="VMware Tools outdated",
                title="VMware Tools 版本过旧",
                object_type="VirtualMachine",
                object_name="app-01",
                status="open",
                current_value=True,
                expected_value=False,
                business_impact="驱动和管理能力可能不是最新状态。",
                remediation="在维护窗口升级 VMware Tools。",
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-VM-099",
                rule_name="VM resource limit configuration",
                title="虚拟机配置了资源 Limit",
                object_type="VirtualMachine",
                object_name="app-resource-config-1234567890abcdef1234567890abcdef",
                status="open",
                current_value="CPU limit: 2000MHz; memory limit: 4096MB; inherited from very-long-resource-pool-name-for-validation",
                expected_value="CPU/Memory limit 未启用",
                business_impact="资源限制可能导致虚拟机在高峰期无法使用完整资源。",
                remediation="建议结合业务窗口核对该虚拟机资源限制配置，确认是否为业务要求。",
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-VM-100",
                rule_name="VM resource limit configuration",
                title="vCLS 默认虚拟机资源 Limit 需确认",
                object_type="VirtualMachine",
                object_name="vCLS-1234567890abcdef1234567890abcdef",
                status="open",
                current_value="CPU limit: 2000MHz; memory limit: 4096MB",
                expected_value="CPU/Memory limit 未启用",
                business_impact="系统默认虚拟机不应进入客户风险正文。",
                remediation="系统默认虚拟机由平台自动管理。",
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-HOST-099",
                rule_name="ESXi version mismatch",
                title="ESXi 主机版本不一致",
                object_type="HostSystem",
                object_name="esxi-01.lab.local-172.16.10.15",
                status="open",
                current_value="ESXi-8.0.3-build-24585383-with-extra-long-build-description",
                expected_value="与集群内主机版本保持一致",
                business_impact="版本差异可能影响集群维护和兼容性判断。",
                remediation="建议在维护窗口统一核对 ESXi 版本与补丁基线。",
            ),
            FindingItem(
                risk_level="P2",
                rule_id="VSL-NET-015",
                rule_name="vSwitch and portgroup security policy",
                title="端口组安全策略配置不符合建议",
                object_type="HostSystem",
                object_name="esxi-01",
                status="open",
                current_value=2,
                expected_value=0,
                business_impact="端口组安全策略放宽可能影响网络隔离边界。",
                remediation="核对端口组用途，非必要场景关闭相关策略。",
                observed_detail={
                    "host_name": "esxi-01",
                    "portgroup_security_issues": [
                        {
                            "switch": "vSwitch0",
                            "portgroup": "VM Network",
                            "policy": "混杂模式",
                            "current_value": "已启用",
                            "recommended_value": "禁用",
                        }
                    ],
                },
            ),
            FindingItem(
                risk_level="P2",
                rule_id="VSL-HOST-024",
                rule_name="ESXi host hardware health",
                title="ESXi 主机硬件健康状态异常",
                object_type="HostSystem",
                object_name="esxi-01",
                status="open",
                current_value=1,
                expected_value=0,
                business_impact="主机硬件状态异常可能影响承载能力。",
                remediation="结合硬件管理平台复核传感器状态。",
                observed_detail={"host_name": "esxi-01", "host_hardware_health_issues": [{"component": "Power Supply 1", "status": "red"}]},
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-VM-023",
                rule_name="VM running on local datastore",
                title="虚拟机运行在本地存储上",
                object_type="VirtualMachine",
                object_name="app-01",
                status="open",
                current_value=True,
                expected_value=False,
                business_impact="本地存储可能限制迁移和故障恢复能力。",
                remediation="确认业务等级，必要时迁移到共享 Datastore。",
                observed_detail={"local_datastore_names": ["local-esxi-01"], "datastore_names": ["local-esxi-01"]},
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-VM-024",
                rule_name="VM configured OS differs from guest OS",
                title="虚拟机操作系统配置与实际不一致",
                object_type="VirtualMachine",
                object_name="app-01",
                status="open",
                current_value=True,
                expected_value=False,
                business_impact="OS 信息不准确可能影响资产台账和兼容性评估。",
                remediation="核对实际操作系统，必要时修正虚拟机客户机操作系统配置。",
                observed_detail={"guest_os_actual": "Ubuntu Linux (64-bit)", "guest_os_configured": "Microsoft Windows Server 2019"},
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-HOST-025",
                rule_name="ESXi host power policy",
                title="ESXi 主机电源策略非高性能",
                object_type="HostSystem",
                object_name="esxi-01",
                status="open",
                current_value=False,
                expected_value=True,
                business_impact="非高性能电源策略可能造成性能波动。",
                remediation="结合业务负载确认是否调整为高性能策略。",
                observed_detail={"host_power_policy": "Balanced"},
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-DS-019",
                rule_name="Datastore shared across clusters",
                title="Datastore 跨集群共享需确认规划",
                object_type="Datastore",
                object_name="DS02",
                status="open",
                current_value=True,
                expected_value=False,
                business_impact="跨集群共享 Datastore 需要确认规划边界。",
                remediation="结合存储规划确认共享范围。",
                observed_detail={"datastore_cluster_count": 2, "datastore_cluster_names": ["Cluster-A", "Cluster-B"], "datastore_host_count": 6},
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-DS-020",
                rule_name="vSAN datastore capacity health",
                title="vSAN 存储健康存在异常",
                object_type="Datastore",
                object_name="vsanDatastore",
                status="open",
                current_value=85.5,
                expected_value="< 80%",
                business_impact="vSAN 容量水位偏高会降低对象修复和重同步余量。",
                remediation="清理无效数据、迁移负载或规划扩容。",
                observed_detail={"datastore_name": "vsanDatastore", "vsan_used_percent": 85.5, "vsan_free_gb": 512.0},
            ),
            FindingItem(
                risk_level="P3",
                rule_id="VSL-HOST-026",
                rule_name="Host resource allocation overcommit ratio",
                title="主机资源分配比例偏高",
                object_type="HostSystem",
                object_name="esxi-01",
                status="open",
                current_value=True,
                expected_value=False,
                business_impact="资源分配比例偏高会增加容量余量压力。",
                remediation="结合性能指标和容量增长计划评估是否回收或扩容。",
                observed_detail={"host_vcpu_to_pcpu_ratio": 5.0, "host_memory_allocation_ratio": 1.6},
            ),
        ]
    )
    return report


def docx_text(path: Path) -> str:
    doc = Document(path)
    parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells if cell.text.strip())
    return "\n".join(parts)


def docx_headings(path: Path) -> list[str]:
    doc = Document(path)
    return [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.style.name.startswith("Heading") and paragraph.text.strip()]


def table_headers(table) -> list[str]:
    return [cell.text.strip() for cell in table.rows[0].cells]


def find_table_with_headers(path: Path, headers: list[str]):
    doc = Document(path)
    for table in doc.tables:
        if table_headers(table) == headers:
            return table
    raise AssertionError(f"table with headers not found: {headers}")


def find_tables_with_headers(path: Path, headers: list[str]):
    doc = Document(path)
    return [table for table in doc.tables if table_headers(table) == headers]


def cell_fill(cell) -> str | None:
    shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


def cell_width(cell) -> int:
    tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
    return int(tc_w.get(qn("w:w"))) if tc_w is not None else 0


def table_width(table) -> int:
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    return int(tbl_w.get(qn("w:w"))) if tbl_w is not None else 0


def table_layout(table) -> str | None:
    layout = table._tbl.tblPr.find(qn("w:tblLayout"))
    return layout.get(qn("w:type")) if layout is not None else None


def table_grid_widths(table) -> list[int]:
    return [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid.findall(qn("w:gridCol"))]


def cell_alignment(cell):
    return cell.paragraphs[0].alignment if cell.paragraphs else None


def row_cant_split(row) -> bool:
    return row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None


def first_run_font_size(cell) -> float | None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.text:
                return run.font.size.pt if run.font.size is not None else None
    return None


def risk_category_numbers(headings: list[str]) -> list[int]:
    numbers: list[int] = []
    for heading in headings:
        match = re.match(r"2\.(\d+) (?:P1|P2|P3)", heading)
        if match:
            numbers.append(int(match.group(1)))
    return numbers


def docx_xml(path: Path, member: str) -> str:
    with ZipFile(path) as zf:
        return zf.read(member).decode("utf-8")


def test_exporter_renders_docx_by_default_suffix(tmp_path: Path) -> None:
    output = tmp_path / "report"
    rendered_path, report_type, report_name = ReportExportEngine().render(minimal_report(), output)

    assert rendered_path == tmp_path / "report.docx"
    assert report_type == "docx"
    assert report_name == "VStackLens Word Report"
    assert rendered_path.exists()
    text = docx_text(rendered_path)
    headings = docx_headings(rendered_path)
    assert "vSphere 环境健康巡检报告" in text
    assert "交付客户" in text
    assert "vcsa.test.local" in text
    assert "1. 执行摘要" in headings
    assert "1.4 重点问题" in headings
    assert "1.5 巡检结论" in headings
    assert "2. 问题与整改建议" in headings
    assert "2.1 P1" in headings
    assert "2.2 P2" in headings
    assert "2.3 P3" in headings
    assert "3. VMware 巡检" in headings
    assert "5. 环境资产与对象清单" in headings
    assert not any("附录" in heading for heading in headings)
    assert "完整的对象明细、证据和采集来源保留在同一次巡检生成的 HTML 报告和数据包中。" not in text
    assert "P1" in text
    assert "对象\n告警内容\n确认状态\n时间\n状态" in text
    assert "健康评分" not in text
    for term in FORBIDDEN_WORD_TERMS:
        assert term not in text


def test_exporter_generates_idle_vsan_resync_panel(tmp_path: Path) -> None:
    report = minimal_report().model_copy(update={
        "vsan_summary": {
            "status": "collected",
            "resync_object_count": 0,
            "resync_bytes": 0,
        }
    })

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "vsan-report.docx")
    text = docx_text(rendered_path)

    assert "当前无待同步任务" in text
    assert "0 个" in text
    assert "0 B" in text
    assert "0 秒" in text
    assert "0 项" in text


def test_vsan_charts_do_not_render_zero_abnormal_or_empty_segment(tmp_path: Path) -> None:
    report = minimal_report().model_copy(update={
        "vsan_summary": {
            "status": "collected",
            "object_count": 165,
            "vmdk_count": 86,
            "object_issue_count": 0,
            "capacity": {"status": "normal", "used_percent": 10.1, "total_gb": 100.0, "used_gb": 10.1, "free_gb": 89.9},
        }
    })
    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "vsan-charts.docx")
    doc = Document(rendered_path)
    object_bar = next(table for table in doc.tables if table.rows[0].cells[0].text == "vSAN 对象")
    vmdk_bar = next(table for table in doc.tables if table.rows[0].cells[0].text == "VMDK")
    capacity_bar = next(table for table in doc.tables if table.rows[0].cells[0].text == "容量")
    assert [cell.text for cell in object_bar.rows[0].cells] == ["vSAN 对象", "正常 165"]
    assert [cell.text for cell in vmdk_bar.rows[0].cells] == ["VMDK", "已识别 86"]
    assert len(capacity_bar.rows[0].cells) == 3
    assert "异常 0" in docx_text(rendered_path)


def test_docx_feedback_revision_uses_native_toc_field(tmp_path: Path) -> None:
    rendered_path, _, _ = ReportExportEngine().render(minimal_report(), tmp_path / "report.docx")

    document_xml = docx_xml(rendered_path, "word/document.xml")
    settings_xml = docx_xml(rendered_path, "word/settings.xml")

    assert 'TOC \\o "1-3" \\h \\z \\u' in document_xml
    assert 'w:fldCharType="begin"' in document_xml
    assert 'w:fldCharType="separate"' in document_xml
    assert 'w:fldCharType="end"' in document_xml
    assert "w:updateFields" in settings_xml


def test_html_report_engine_autoescapes_jinja_template_values(tmp_path: Path) -> None:
    attack = "<script>alert(1)</script>"
    report = rich_report().model_copy(deep=True)
    report.customer_info.customer_name = attack
    report.findings[0].title = attack
    report.findings[0].object_name = f"vm-{attack}"

    output = tmp_path / "report.html"
    HtmlReportEngine().render(report, output)
    html = output.read_text(encoding="utf-8")

    assert attack not in html
    assert "vm-<script>alert(1)</script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "vm-&lt;script&gt;alert(1)&lt;/script&gt;" in html


def test_docx_v2_issue_summary_table_has_expected_headers_colors_and_widths(tmp_path: Path) -> None:
    rendered_path, _, _ = ReportExportEngine().render(table_layout_report(), tmp_path / "report.docx")
    table = find_table_with_headers(rendered_path, ["等级", "分类", "问题描述", "影响数量", "建议动作"])

    widths = [cell_width(cell) for cell in table.rows[0].cells]
    assert widths == [950, 1400, 2400, 950, 3825]

    expected_fills = {"P1": "C00000", "P2": "7030A0", "P3": "FFC000"}
    seen: dict[str, str] = {}
    for row in table.rows[1:]:
        level = row.cells[0].text.strip()
        if level in expected_fills:
            seen[level] = cell_fill(row.cells[0])
    assert seen == expected_fills


def test_docx_summary_counts_grouped_issues_separately_from_object_details(tmp_path: Path) -> None:
    report = table_layout_report()
    duplicate = report.findings[0].model_copy(deep=True)
    duplicate.object_name = "vcsa-standby.test.local"
    report.findings.append(duplicate)
    engine = DocxReportEngine()
    groups = engine._customer_visible_risk_groups(engine._aggregate_risks(report))
    group_counts = engine._risk_counts(report, groups)
    detail_counts = engine._finding_counts(groups)

    assert sum(group_counts.values()) == len(groups)
    assert sum(detail_counts.values()) == sum(len(group.findings) for group in groups)
    assert sum(detail_counts.values()) > sum(group_counts.values())

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "report.docx")
    text = docx_text(rendered_path)

    assert f"{sum(group_counts.values())}" in text
    assert f"{sum(detail_counts.values())}" in text
    assert "类建议关注问题" in text
    assert "条对象明细" in text
    assert "按问题类型统计" in text
    assert "按影响对象明细统计" in text
    assert "Word 正文仅保留代表性对象明细" in text
    assert "完整对象清单建议在 HTML 报告中查看" in text


def test_docx_merges_expired_and_deep_snapshots_into_one_p1_problem(tmp_path: Path) -> None:
    report = rich_report().model_copy(deep=True)
    report.findings = []
    for index in range(14):
        report.findings.append(FindingItem(
            risk_level="P1",
            rule_id="VSL-VM-001",
            rule_name="虚拟机存在超期快照",
            title="虚拟机存在超期快照",
            object_type="VirtualMachine",
            object_name=f"vm-{index:02d}",
            object_path=f"vcsa.test.local / Cluster-A / vm-{index:02d}",
            status="failed",
            business_impact="长期快照可能影响性能。",
            observed_detail={"snapshots": [{"name": f"snapshot-{index:02d}", "age_days": 120}]},
        ))
    for index in range(3):
        report.findings.append(FindingItem(
            risk_level="P3",
            rule_id="VSL-VM-007",
            rule_name="虚拟机快照链过深",
            title="虚拟机快照链过深",
            object_type="VirtualMachine",
            object_name=f"vm-{index:02d}",
            object_path=f"vcsa.test.local / Cluster-A / vm-{index:02d}",
            status="failed",
            business_impact="过深快照链可能影响备份。",
            observed_detail={"snapshots": [{"name": f"snapshot-{index:02d}", "chain_depth": 4}]},
        ))

    engine = DocxReportEngine()
    groups = engine._customer_visible_risk_groups(engine._aggregate_risks(report))
    snapshots = [group for group in groups if "快照" in group.title]

    assert len(snapshots) == 1
    assert snapshots[0].title == "虚拟机存在快照"
    assert snapshots[0].level == "P1"
    assert len(snapshots[0].findings) == 14
    assert engine._risk_counts(report, groups) == {"P1": 1, "P2": 0, "P3": 0}
    output, _, _ = ReportExportEngine().render(report, tmp_path / "snapshot-merged.docx")
    text = docx_text(output)
    assert "虚拟机存在快照" in text
    assert "虚拟机快照链过深" not in text
    assert "超期快照或过深快照链" in text


def test_docx_powered_off_problem_explains_excluded_templates(tmp_path: Path) -> None:
    report = rich_report().model_copy(deep=True)
    report.findings = [FindingItem(
        risk_level="P3",
        rule_id="VSL-VM-018",
        rule_name="存在关机的虚拟机",
        title="存在关机的虚拟机",
        object_type="VirtualMachine",
        object_name="vm-06",
        status="failed",
        business_impact="环境中存在关机的虚拟机。",
    )]
    report.asset_inventory["details"]["VirtualMachine"] = [
        {
            "object_name": f"vm-{index:02d}",
            "properties": {
                "power_state": "poweredOff",
                "is_template": index < 6,
                "is_system_vm": False,
            },
        }
        for index in range(15)
    ]

    output, _, _ = ReportExportEngine().render(report, tmp_path / "powered-off-scope.docx")
    text = docx_text(output)
    detail = find_table_with_headers(output, ["虚拟机名称", "当前电源状态", "建议处理"])

    assert "另有 6 台关机的模板或系统虚拟机，不计入。" in text
    assert detail.rows[1].cells[1].text.strip() == "关机"
    assert "poweredOff" not in " ".join(cell.text for row in detail.rows for cell in row.cells)
    assert "未采集天" not in " ".join(cell.text for row in detail.rows for cell in row.cells)


def test_docx_customer_text_replaces_internal_tsm_ssh_service_name(tmp_path: Path) -> None:
    report = rich_report().model_copy(deep=True)
    report.findings = [FindingItem(
        risk_level="P3",
        rule_id="VSL-HOST-019",
        rule_name="ESXi 主机 SSH 服务已开启",
        title="ESXi 主机 SSH 服务已开启",
        object_type="HostSystem",
        object_name="esxi-01",
        status="failed",
        remediation="停止 TSM-SSH 服务，并复核主机服务状态。",
    )]

    output, _, _ = ReportExportEngine().render(report, tmp_path / "ssh-service.docx")
    text = docx_text(output)

    assert "TSM-SSH" not in text
    assert "停止 SSH 服务" in text


def test_docx_customer_report_hides_p4_items(tmp_path: Path) -> None:
    rendered_path, _, _ = ReportExportEngine().render(rich_report(), tmp_path / "report.docx")
    text = docx_text(rendered_path)

    assert "P4" not in text
    assert "P4级" not in text
    assert "集群 vMotion 网络覆盖不完整" not in text


def test_docx_tables_use_fixed_total_width_and_semantic_alignment(tmp_path: Path) -> None:
    rendered_path, _, _ = ReportExportEngine().render(table_layout_report(), tmp_path / "report.docx")
    doc = Document(rendered_path)

    assert doc.tables
    for table in doc.tables:
        label = table.rows[0].cells[0].text.strip()
        expected_layout = "fixed" if label in {"vSAN 对象", "VMDK", "容量"} else "autofit"
        assert table_layout(table) == expected_layout
        assert table.autofit is (expected_layout == "autofit")
        assert table_width(table) == 9525
        assert sum(table_grid_widths(table)) == 9525
        assert [cell_width(cell) for cell in table.rows[0].cells] == table_grid_widths(table)
        assert row_cant_split(table.rows[0])
        for cell in table.rows[0].cells:
            assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
            assert cell_alignment(cell) == WD_ALIGN_PARAGRAPH.CENTER
        for row in table.rows[1:]:
            assert row_cant_split(row)
            for cell in row.cells:
                assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER

    summary_table = find_table_with_headers(rendered_path, ["等级", "分类", "问题描述", "影响数量", "建议动作"])
    assert cell_alignment(summary_table.rows[1].cells[0]) == WD_ALIGN_PARAGRAPH.CENTER
    assert cell_alignment(summary_table.rows[1].cells[1]) == WD_ALIGN_PARAGRAPH.CENTER
    assert cell_alignment(summary_table.rows[1].cells[2]) == WD_ALIGN_PARAGRAPH.LEFT


def test_docx_known_tables_match_feedback_revision_columns(tmp_path: Path) -> None:
    rendered_path, _, _ = ReportExportEngine().render(table_layout_report(), tmp_path / "report.docx")

    expected_tables = {
        ("vCenter", "数据中心", "集群", "ESXi 主机", "共享/业务 Datastore", "虚拟机"): [1587, 1587, 1587, 1588, 1588, 1588],
        ("等级", "分类", "问题描述", "影响数量", "建议动作"): [950, 1400, 2400, 950, 3825],
        ("名称", "版本", "Build", "连接状态"): [2199, 2200, 2563, 2563],
        ("主机", "所属集群", "连接状态", "CPU 使用率", "内存使用率"): [1905, 1905, 1905, 1905, 1905],
        ("Datastore", "类型", "已用", "剩余", "状态"): [1905, 1905, 1905, 1905, 1905],
        ("集群", "HA 当前状态"): [4762, 4763],
        ("集群", "DRS 当前状态"): [4762, 4763],
        ("ESXi 主机", "网卡", "绑定交换机", "链路状态", "实际/配置速率", "判断与建议"): [1500, 900, 1400, 1150, 1850, 2725],
        ("ESXi 主机", "所属集群", "当前状态"): [3175, 3175, 3175],
        ("所在虚拟机", "快照名称", "创建时间/存在时长"): [3175, 3175, 3175],
        ("ESXi 主机", "vSwitch或端口组", "策略项", "当前状态", "建议"): [1531, 2268, 1531, 1531, 2664],
        ("ESXi 主机", "组件", "健康状态"): [3175, 3175, 3175],
        ("虚拟机", "Datastore", "所属主机或集群", "影响说明"): [2500, 2200, 2200, 2625],
        ("虚拟机", "实际操作系统", "配置操作系统"): [2500, 3512, 3513],
        ("Datastore", "关联集群数", "关联集群", "主机数"): [2600, 1600, 3725, 1600],
        ("Datastore", "已用率", "剩余容量", "状态"): [2800, 1900, 2200, 2625],
        ("ESXi 主机", "vCPU:pCPU", "内存分配比例", "当前状态"): [2600, 2300, 2300, 2325],
        ("对象", "当前状态", "建议"): [2500, 3300, 3725],
    }

    for headers, widths in expected_tables.items():
        table = find_table_with_headers(rendered_path, list(headers))
        assert [cell_width(cell) for cell in table.rows[0].cells] == widths

    snapshot_table = find_table_with_headers(rendered_path, ["所在虚拟机", "快照名称", "创建时间/存在时长"])
    assert any(row.cells[2].text.strip() for row in snapshot_table.rows[1:])

    text = docx_text(rendered_path)
    assert "ESXi主机" not in text
    assert "序号" not in table_headers(find_table_with_headers(rendered_path, ["等级", "分类", "问题描述", "影响数量", "建议动作"]))


def test_docx_snapshot_rows_omit_uncollected_field_stack(tmp_path: Path) -> None:
    report = rich_report().model_copy(deep=True)
    report.findings = [
        FindingItem(
            risk_level="P3",
            rule_id="VSL-VM-001",
            rule_name="VM snapshot",
            title="虚拟机存在快照",
            object_type="VirtualMachine",
            object_name="app-01",
            status="open",
            current_value=True,
            expected_value=False,
            business_impact="长期保留快照会占用存储，可能影响虚拟机性能，并增加备份、迁移和快照合并风险。",
            remediation="确认快照用途后删除或合并快照。",
            observed_detail={"snapshots": [{}]},
        )
    ]

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "report.docx")
    snapshot_table = next(table for table in Document(rendered_path).tables if "app-01" in "\n".join(cell.text for row in table.rows for cell in row.cells))
    row_text = " ".join(cell.text for cell in snapshot_table.rows[1].cells)

    assert "快照名称：未采集" not in row_text
    assert "创建时间：未采集" not in row_text
    assert "快照链深度：未采集" not in row_text
    assert row_text.endswith("未采集")
    assert "虚拟机存在快照" in row_text


def test_vsl_vm_007_collected_snapshot_names_and_times_reach_docx(tmp_path: Path) -> None:
    vm_snapshots = {
        "WF-alan软件（ubuntu）": 3,
        "WF-知识库业务（centos9）": 4,
        "WF-知识库业务（win）": 3,
        "centos9-备份-勿开": 4,
    }
    report = rich_report().model_copy(deep=True)
    report.findings = []
    expected_names = []
    for vm_name, snapshot_count in vm_snapshots.items():
        snapshots = [
            {
                "name": f"{vm_name}-snapshot-{snapshot_index}",
                "create_time": f"2026-09-{10 + snapshot_index:02d}T03:11:22+00:00",
            }
            for snapshot_index in range(1, snapshot_count + 1)
        ]
        evidence = EvidenceBuilder()._evidence_vsl_vm_007(
            SimpleNamespace(object_name=vm_name),
            {"snapshot_chain_depth": snapshot_count, "max_snapshot_chain_depth": 2, "snapshots": snapshots},
        )
        expected_names.extend(item["name"] for item in snapshots)
        report.findings.append(
            FindingItem(
                risk_level="P3",
                rule_id="VSL-VM-007",
                rule_name="VM snapshot chain too deep",
                title="虚拟机快照链过深",
                object_type="VirtualMachine",
                object_name=vm_name,
                status="open",
                current_value=snapshot_count,
                expected_value=2,
                business_impact="快照链过深可能影响虚拟机性能。",
                remediation="确认用途后删除或合并无用快照。",
                observed_detail=evidence["observed_detail"],
            )
        )

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "snapshot-chain-report.docx")
    table = find_table_with_headers(rendered_path, ["所在虚拟机", "快照名称", "创建时间/存在时长"])
    rows = [" ".join(cell.text for cell in row.cells) for row in table.rows[1:]]
    visible = "\n".join(rows)
    assert len(rows) == 14
    assert all(name in visible for name in expected_names)
    assert "centos9-备份-勿开" in visible
    assert "未采集" not in visible


def test_docx_snapshot_table_falls_back_to_collected_inventory_when_evidence_detail_is_missing(tmp_path: Path) -> None:
    report = rich_report().model_copy(deep=True)
    vm = next(item for item in report.asset_inventory["details"]["VirtualMachine"] if item["object_name"] == "app-01")
    vm["properties"]["snapshots"] = [
        {"name": "inventory-only-snapshot", "create_time": "2026-09-12T08:00:00+00:00"}
    ]
    report.findings = [
        FindingItem(
            risk_level="P3",
            rule_id="VSL-VM-007",
            rule_name="VM snapshot chain too deep",
            title="虚拟机快照链过深",
            object_type="VirtualMachine",
            object_name="app-01",
            status="open",
            current_value=3,
            expected_value=2,
            business_impact="快照链过深可能影响虚拟机性能。",
            remediation="确认用途后删除或合并无用快照。",
            observed_detail={"snapshot_chain_depth": 3},
        )
    ]

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "snapshot-inventory-fallback.docx")
    table = find_table_with_headers(rendered_path, ["所在虚拟机", "快照名称", "创建时间/存在时长"])
    visible = "\n".join(" ".join(cell.text for cell in row.cells) for row in table.rows[1:])
    assert "inventory-only-snapshot" in visible
    assert "2026-09-12T08:00:00+00:00" in visible
    assert "未采集" not in visible


def test_docx_pnic_details_show_binding_link_and_rate_basis(tmp_path: Path) -> None:
    report = rich_report().model_copy(deep=True)
    report.findings = [
        FindingItem(
            risk_level="P2",
            rule_id="VSL-HOST-019",
            rule_name="ESXi host physical NIC link issue",
            title="ESXi 主机物理网卡链路异常",
            object_type="HostSystem",
            object_name="esxi-01",
            status="open",
            current_value=2,
            expected_value=0,
            business_impact="网络链路异常可能影响业务。",
            remediation="检查链路和交换机端口。",
            observed_detail={
                "host_name": "esxi-01",
                "physical_nic_details": [
                    {
                        "device": "vmnic0", "is_uplink": True, "assigned_switches": ["vSwitch0"],
                        "link_state": "up", "actual_speed_mb": 1000, "configured_speed_mb": None,
                        "autonegotiation": True, "assessment": "up", "issue_codes": [],
                    },
                    {
                        "device": "vmnic1", "is_uplink": False, "assigned_switches": [],
                        "link_state": "down", "actual_speed_mb": None, "configured_speed_mb": None,
                        "autonegotiation": True, "assessment": "unused", "issue_codes": [],
                    },
                    {
                        "device": "vmnic2", "is_uplink": True, "assigned_switches": ["vSwitch0"],
                        "link_state": "down", "actual_speed_mb": None, "configured_speed_mb": None,
                        "autonegotiation": True, "assessment": "down", "issue_codes": ["link_down"],
                    },
                    {
                        "device": "vmnic3", "is_uplink": True, "assigned_switches": ["vSAN-vDS"],
                        "link_state": "up", "actual_speed_mb": 1000, "configured_speed_mb": 10000,
                        "autonegotiation": False, "assessment": "speed_mismatch", "issue_codes": ["speed_mismatch"],
                    },
                ],
            },
        )
    ]

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "pnic-evidence-report.docx")
    table = find_table_with_headers(
        rendered_path,
        ["ESXi 主机", "网卡", "绑定交换机", "链路状态", "实际/配置速率", "判断与建议"],
    )
    rows = {row.cells[1].text.strip(): " ".join(cell.text for cell in row.cells) for row in table.rows[1:]}
    assert "未绑定" in rows["vmnic1"]
    assert "不按上行链路判故障" in rows["vmnic1"]
    assert "vSwitch0" in rows["vmnic2"] and "down" in rows["vmnic2"]
    assert "vSAN-vDS" in rows["vmnic3"]
    assert "实际 1000 Mbps / 配置 10000 Mbps" in rows["vmnic3"]
    assert "固定配置速率与实际速率不一致" in rows["vmnic3"]


def test_docx_pnic_table_does_not_truncate_collected_adapter_rows(tmp_path: Path) -> None:
    physical_nics = [
        {
            "device": f"vmnic{index}",
            "is_uplink": True,
            "assigned_switches": ["vSwitch0"],
            "link_state": "down" if index == 0 else "up",
            "actual_speed_mb": None if index == 0 else 10000,
            "configured_speed_mb": None,
            "autonegotiation": True,
            "assessment": "down" if index == 0 else "up",
            "issue_codes": ["link_down"] if index == 0 else [],
        }
        for index in range(12)
    ]
    report = rich_report().model_copy(deep=True)
    report.findings = [
        FindingItem(
            risk_level="P2",
            rule_id="VSL-HOST-019",
            rule_name="ESXi host physical NIC link issue",
            title="ESXi 主机物理网卡链路异常",
            object_type="HostSystem",
            object_name="esxi-01",
            status="open",
            current_value=1,
            expected_value=0,
            business_impact="网络链路异常可能影响业务。",
            remediation="检查链路和交换机端口。",
            observed_detail={"host_name": "esxi-01", "physical_nic_details": physical_nics},
        )
    ]

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "pnic-all-rows.docx")
    table = find_table_with_headers(
        rendered_path,
        ["ESXi 主机", "网卡", "绑定交换机", "链路状态", "实际/配置速率", "判断与建议"],
    )
    assert len(table.rows) - 1 == 12
    assert any(row.cells[1].text.strip() == "vmnic11" for row in table.rows[1:])
    assert "其余对象建议在 HTML 报告中进一步查看" not in docx_text(rendered_path)


def test_docx_general_risk_details_are_not_truncated_after_ten_rows(tmp_path: Path) -> None:
    report = rich_report().model_copy(deep=True)
    report.findings = [
        FindingItem(
            risk_level="P3",
            rule_id="TEST-VM-DETAILS",
            rule_name="Test virtual machine details",
            title="虚拟机巡检明细测试",
            object_type="VirtualMachine",
            object_name=f"vm-detail-{index:02d}",
            status="open",
            current_value=f"value-{index:02d}",
            expected_value="expected",
            business_impact="需要查看该对象的明细。",
            remediation="逐项核查。",
        )
        for index in range(12)
    ]

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "all-risk-details.docx")
    visible = docx_text(rendered_path)
    assert all(f"vm-detail-{index:02d}" in visible for index in range(12))
    assert "其余对象建议在 HTML 报告中进一步查看" not in visible


def test_docx_vsan_object_and_disk_issue_details_are_not_truncated(tmp_path: Path) -> None:
    report = rich_report().model_copy(deep=True)
    object_issues = [
        {"component": f"vsan-object-{index:02d}", "status": "red", "summary": f"object-detail-{index:02d}"}
        for index in range(12)
    ]
    disk_issues = [
        {"host": "esxi-01", "disk": f"naa-issue-{index:02d}", "status": "error", "summary": f"disk-detail-{index:02d}"}
        for index in range(12)
    ]
    report.vsan_summary = {
        "status": "collected",
        "architecture": "osa",
        "clusters": ["vSAN"],
        "health_issue_count": 0,
        "object_issue_count": len(object_issues),
        "object_count": 100,
        "vmdk_count": 0,
        "host_count": 1,
        "disk_group_count": 1,
        "cache_disk_count": 1,
        "capacity_disk_count": 1,
        "disk_topology": [
            {"host": "esxi-01", "disk_groups": [{"cache_disk": "naa-cache", "capacity_disks": ["naa-capacity"]}]}
        ],
        "physical_disks": [
            {"host": "esxi-01", "name": "naa-cache", "summary_health": "green"},
            {"host": "esxi-01", "name": "naa-capacity", "summary_health": "green"},
        ],
        "disk_issues": disk_issues,
        "object_issues": object_issues,
        "resync_object_count": 0,
        "resync_bytes": 0,
        "capacity": {"total_gb": 100, "used_gb": 25, "free_gb": 75, "used_percent": 25, "status": "normal"},
        "network": {"vmkernels": []},
    }

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "all-vsan-details.docx")
    object_table = find_table_with_headers(rendered_path, ["对象/组件", "Health", "说明"])
    disk_table = find_table_with_headers(rendered_path, ["主机", "设备", "状态", "说明"])
    assert len(object_table.rows) - 1 == 12
    assert len(disk_table.rows) - 1 == 12
    visible = docx_text(rendered_path)
    assert "vsan-object-11" in visible and "object-detail-11" in visible
    assert "naa-issue-11" in visible and "disk-detail-11" in visible
    assert "未取得或未确认的数据不按正常处理" in visible


def test_docx_legacy_pnic_evidence_labels_host_level_switch_mapping(tmp_path: Path) -> None:
    report = table_layout_report()
    host = next(item for item in report.asset_inventory["details"]["HostSystem"] if item["object_name"] == "esxi-01")
    host["properties"]["affected_switches"] = ["vSwitch0"]
    finding = next(item for item in report.findings if item.rule_id == "VSL-HOST-019")
    finding.observed_detail = {
        "host_name": "esxi-01",
        "affected_nics": ["vmnic1"],
        "link_speed_detail": [{"device": "vmnic1", "status": "down", "speed_mb": None}],
    }

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "legacy-pnic-evidence.docx")
    table = find_table_with_headers(
        rendered_path,
        ["ESXi 主机", "网卡", "绑定交换机", "链路状态", "实际/配置速率", "判断与建议"],
    )
    row = next(row for row in table.rows[1:] if row.cells[1].text.strip() == "vmnic1")
    text = " ".join(cell.text for cell in row.cells)
    assert "vSwitch0（主机级映射，未逐卡保存）" in text
    assert "down" in text
    assert "旧版证据未记录" in text


def test_docx_fallback_detail_tables_wrap_long_tokens_and_left_align_text(tmp_path: Path) -> None:
    rendered_path, _, _ = ReportExportEngine().render(table_layout_report(), tmp_path / "report.docx")
    fallback_tables = find_tables_with_headers(rendered_path, ["对象", "当前状态", "建议"])

    assert fallback_tables
    for table in fallback_tables:
        assert [cell_width(cell) for cell in table.rows[0].cells] == [2500, 3300, 3725]
        for row in table.rows[1:]:
            assert cell_alignment(row.cells[0]) == WD_ALIGN_PARAGRAPH.LEFT
            assert cell_alignment(row.cells[1]) == WD_ALIGN_PARAGRAPH.LEFT
            assert cell_alignment(row.cells[2]) == WD_ALIGN_PARAGRAPH.LEFT
            for cell in row.cells:
                if cell.text.strip():
                    assert first_run_font_size(cell) == 8.5

    long_name_row = next(
        row
        for table in fallback_tables
        for row in table.rows[1:]
        if "app-resource-config" in row.cells[0].text.replace("\u200b", "")
    )
    # Object identity must remain copyable/searchable; the renderer no longer
    # injects zero-width break characters into names.
    assert "\u200b" not in long_name_row.cells[0].text
    assert long_name_row.cells[1].text.startswith("当前：")
    assert "very-long-resource-pool-name-for-validation" not in long_name_row.cells[1].text
    assert len(long_name_row.cells[1].text) < 80


def test_docx_customer_word_refines_finding_display_terms_and_duplicates(tmp_path: Path) -> None:
    rendered_path, _, _ = ReportExportEngine().render(table_layout_report(), tmp_path / "report.docx")
    text = docx_text(rendered_path)
    headings = docx_headings(rendered_path)

    assert "P4" not in text
    assert "P4级" not in text
    assert "期望值" not in text
    assert "Limit" not in text
    assert "CPU limit" not in text
    assert "memory limit" not in text
    assert "CPU/Memory limit" not in text
    assert "vCLS-" not in text
    assert "VMware vCLS" not in text

    assert "创建时间/存在时长" in text
    assert "快照保留时间" not in text
    assert "Syslog" not in text
    assert "syslog" not in text.lower()
    assert "VMware Tools" not in text
    assert "ESXi 主机日志与核心转储配置不完整" not in text
    assert "core dump" not in text.lower()

    assert not any("Tools" in heading for heading in headings)
    assert "vCenter 可以更准确获取客户机状态、IP、心跳和运行信息" not in text
    assert "支持优雅关机/重启" not in text


def test_docx_customer_report_hides_tools_and_syslog_findings_without_mutating_collection(tmp_path: Path) -> None:
    report = table_layout_report()
    original_findings = list(report.findings)
    engine = DocxReportEngine()
    collected_groups = engine._aggregate_risks(report)
    visible_groups = engine._customer_visible_risk_groups(collected_groups)

    assert report.findings == original_findings
    assert any(engine._is_tools_group(group) for group in collected_groups)
    assert any(engine._is_remote_syslog_group(group) for group in collected_groups)
    assert not any(engine._is_tools_group(group) for group in visible_groups)
    assert not any(engine._is_syslog_group(group) for group in visible_groups)

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "report.docx")
    text = docx_text(rendered_path)
    visible_detail_count = sum(len(group.findings) for group in visible_groups)

    assert "VMware Tools" not in text
    assert "Syslog" not in text
    assert f"{len(visible_groups)} 类建议关注问题" in text
    assert f"{visible_detail_count} 条对象明细" in text


def test_docx_cpu_ready_finding_uses_metric_units(tmp_path: Path) -> None:
    report = rich_report().model_copy(deep=True)
    report.risk_summary = RiskSummary(P1=0, P2=1, P3=0, P4=0, total=1)
    report.findings = [
        FindingItem(
            risk_level="P2",
            rule_id="VSL-VM-004",
            rule_name="虚拟机 CPU Ready 状态",
            title="虚拟机 CPU Ready 过高",
            object_type="VirtualMachine",
            object_name="WF-知识库业务（win）",
            status="open",
            current_value=13.19,
            expected_value="< 5",
            current_value_zh="13.19% CPU Ready",
            expected_value_zh="< 5%",
            evidence_summary_zh="WF-知识库业务（win）存在 CPU 调度等待偏高，CPU Ready 当前值为 13.19%，建议低于 5%。",
            observed_detail={"current_value": 13.19, "cpu_ready_percent": 13.19},
            expected_detail={"expected_value": "< 5"},
            observed_detail_zh="当前值：13.19% CPU Ready；CPU Ready：13.19%",
            expected_detail_zh="建议状态：< 5%",
            business_impact="应用响应可能变慢，用户体验下降。",
            remediation="复核虚拟机 vCPU 配置和主机 CPU 调度压力。",
            recommended_action_zh="复核虚拟机 vCPU 配置和主机 CPU 调度压力。",
        )
    ]

    rendered_path, _, _ = ReportExportEngine().render(report, tmp_path / "report.docx")
    text = docx_text(rendered_path)

    assert "CPU Ready" in text
    assert "CPU 调度等待" in text
    assert "13.19%" in text
    assert "< 5%" in text
    assert "当前值为 13.19，建议状态为 < 5" not in text
    assert re.search(r"当前值[:：]\s*13\.19(?!%)", text) is None
    assert "CPU 使用率 13.19%" not in text


def test_docx_v2_risk_category_numbering_is_continuous(tmp_path: Path) -> None:
    rendered_path, _, _ = ReportExportEngine().render(minimal_report(), tmp_path / "report.docx")
    numbers = risk_category_numbers(docx_headings(rendered_path))

    assert numbers == list(range(1, len(numbers) + 1))
    assert numbers == [1, 2, 3]


def test_report_recommendations_do_not_use_unavailable_item_wording() -> None:
    recommendations = ReportDataFactory()._recommendations({}, [{"rule_id": "VSL-TEST"}])
    text = "\n".join(recommendations)

    assert "\u4e0d\u53ef\u5224\u5b9a\u9879" not in text
    assert "\u6570\u636e\u7f3a\u53e3" in text


def test_exporter_writes_report_json(tmp_path: Path) -> None:
    output = tmp_path / "inspection_result.json"
    rendered_path, report_type, _ = ReportExportEngine().render(minimal_report(), output)

    assert rendered_path == output
    assert report_type == "json"
    assert '"schema_version": "1.0"' in output.read_text(encoding="utf-8")


@pytest.mark.parametrize("suffix", [".pdf", ".pptx"])
def test_exporter_does_not_claim_pdf_or_ppt_delivery(tmp_path: Path, suffix: str) -> None:
    with pytest.raises(ValueError, match="HTML .* Word"):
        ReportExportEngine().render(minimal_report(), tmp_path / f"report{suffix}")


def test_docx_safe_float_keeps_uncollected_values_unavailable(tmp_path: Path) -> None:
    engine = DocxReportEngine()
    assert engine._safe_float(None) is None
    assert engine._safe_float("") is None
    assert engine._safe_float("not-a-number") is None
    assert engine._storage_attention({"used_percent": None}) == "未采集"

    report = minimal_report().model_copy(deep=True)
    datastore = report.asset_inventory["details"]["Datastore"][0]["properties"]
    datastore["used_percent"] = None
    output = tmp_path / "report.docx"

    rendered_path, _, _ = ReportExportEngine().render(report, output)
    text = docx_text(rendered_path)

    assert "-1.0" not in text


def test_run_mock_docx_out_generates_docx_and_report_row(tmp_path: Path) -> None:
    docx_out = tmp_path / "mock-report.docx"
    args = SimpleNamespace(
        db=str(tmp_path / "mock.db"),
        rulepack=str(RULEPACK),
        fixture=str(FIXTURE),
        report_dir=str(tmp_path / "report-package"),
        zip_report=False,
        report_title="Word Smoke",
        customer_name="Word 客户",
        site_name="生产站点",
        previous_run_id=None,
        html_out=None,
        out=None,
        docx_out=str(docx_out),
    )

    cmd_run_mock(args)

    assert docx_out.exists()
    with connect(Path(args.db)) as conn:
        rows = conn.execute("SELECT report_type, report_status, file_path FROM reports").fetchall()
    docx_rows = [row for row in rows if row["report_type"] == "docx"]
    assert docx_rows
    assert docx_rows[0]["report_status"] == "success"
    assert Path(docx_rows[0]["file_path"]) == docx_out


@pytest.mark.parametrize("failed_format", ["html", "word", "pdf"])
def test_runner_keeps_other_report_formats_when_one_export_fails(monkeypatch, tmp_path: Path, failed_format: str) -> None:
    error_text = f"{failed_format} simulated failure"
    if failed_format == "html":
        monkeypatch.setattr(HtmlReportPackageBuilder, "render", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError(error_text)))
    elif failed_format == "word":
        monkeypatch.setattr(ReportExportEngine, "render", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError(error_text)))
    else:
        monkeypatch.setattr(PdfReportEngine, "render", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError(error_text)))

    report_dir = tmp_path / "report-package"
    request = MockRunRequest(
        db_path=tmp_path / "run.db",
        rulepack_path=RULEPACK,
        report_dir=report_dir,
        fixture_path=FIXTURE,
        customer_name="测试客户",
        site_name="测试站点",
        docx_out=report_dir / "VStackLens-Word-Report.docx",
        pdf_out=report_dir / "VStackLens-PDF-Report.pdf",
    )

    result = InspectionRunner().run_mock(request)

    assert (result.html_error is not None) is (failed_format == "html")
    assert (result.docx_error is not None) is (failed_format == "word")
    assert (result.pdf_error is not None) is (failed_format == "pdf")
    assert (result.report_path.is_file()) is (failed_format != "html")
    assert (result.docx_path is not None and result.docx_path.is_file()) is (failed_format != "word")
    assert (result.pdf_path is not None and result.pdf_path.is_file()) is (failed_format != "pdf")
    with connect(request.db_path) as conn:
        statuses = {row["report_type"]: row["report_status"] for row in conn.execute(
            "SELECT report_type, report_status FROM reports WHERE run_id = ?",
            (result.run_id,),
        )}
    assert statuses["html_package"] == ("failed" if failed_format == "html" else "success")
    assert statuses["docx"] == ("failed" if failed_format == "word" else "success")
    assert statuses["pdf"] == ("failed" if failed_format == "pdf" else "success")


def test_run_vcenter_sdk_path_docx_out_generates_docx(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    class FakeCollector:
        def __init__(self, *args, **kwargs):
            pass

        def precheck(self) -> None:
            return None

        def collect(self, context, plan, progress=None):
            return MockCollector(FIXTURE).collect(context, plan)

    monkeypatch.setattr("vstacklens.application.inspection_runner.PyVmomiCollector", FakeCollector)
    docx_out = tmp_path / "vcenter-report.docx"

    result = InspectionRunner().run_vcenter(
        VCenterRunRequest(
            db_path=tmp_path / "vcenter.db",
            rulepack_path=RULEPACK,
            report_dir=tmp_path / "vcenter-report",
            report_title="SDK Word Smoke",
            customer_name="SDK 客户",
            site_name="生产站点",
            docx_out=docx_out,
            vcenter="vcsa.test.local",
            username="administrator@vsphere.local",
            password="secret",
        )
    )

    assert result.docx_path == docx_out
    assert docx_out.exists()
    with connect(tmp_path / "vcenter.db") as conn:
        report_types = {row["report_type"] for row in conn.execute("SELECT report_type FROM reports").fetchall()}
    assert "docx" in report_types
