"""Presentation checks: preserve stored judgments and independent session safety."""
from pathlib import Path

from docx import Document
from PySide6.QtWidgets import QLabel

from test_pdf_report import _report
from test_stage4_desktop_runtime import window  # isolated official Qt setup
from test_stage3_batch_inspection import setup
from vstacklens.reports.customer_caliber import ensure_report
from vstacklens.reports.docx_report import DocxReportEngine
from vstacklens.reports.html_package import HtmlReportPackageBuilder
from vstacklens.reports.html_assets import SINGLE_PAGE_REPORT_JS
from vstacklens.reports.pdf_report import PdfReportEngine
from vstacklens.reports.presentation import delivery_overview
from vstacklens.desktop.widgets import NumberSummary


def test_saved_grades_are_not_overwritten_by_catalog_defaults(tmp_path):
    from vstacklens.reports.presentation import customer_check_statistics

    report, context = _report()
    context["object_results"] = [dict(rule_id="VSL-VM-004", risk_level=level, result_status="failed",
        object_name=f"vm-{i}", object_type="VirtualMachine", category="vm", check_name="CPU Ready", rule_name="CPU Ready", status="failed")
        for i,level in enumerate(("P1","P2","P2","P3"))]
    context["rule_catalog"] = [dict(rule_id="VSL-VM-004",risk_level="P3",category="vm")]
    report.caliber_inputs["object_results"]=context["object_results"]
    ensure_report(report, context)
    saved=report.model_dump()
    engine=PdfReportEngine()
    first=engine._build_payload(report,context)
    assert {r["level"]:r["count"] for r in first["scope"]["level_distribution"]} == dict(P1=1,P2=2,P3=1)
    assert first["recorded_results"]["failed_levels"] == dict(P1=1,P2=2,P3=1)
    context["rule_catalog"][0]["risk_level"]="P1"
    second=engine._build_payload(report,context)
    assert first["scope"]["level_distribution"]==second["scope"]["level_distribution"]
    html=HtmlReportPackageBuilder()._single_page_payload(report,context)
    assert html["summary"]["delivery"]["recorded_results"]["failed_levels"]==dict(P1=1,P2=2,P3=1)
    statistics=customer_check_statistics(report,context)
    assert {key:statistics[key] for key in ("total","passed","failed")}==dict(total=4,passed=0,failed=4)
    assert statistics["levels"]==dict(P1=1,P2=2,P3=1)
    assert {key:html["summary"]["checkStatistics"][key] for key in ("total","passed","failed")}==dict(total=4,passed=0,failed=4)
    doc=Document(DocxReportEngine().render(report,tmp_path/"grades.docx"))
    paragraphs="\n".join(doc.element.xpath('//w:t/text()'))
    assert statistics["text"] in paragraphs
    assert "P1 1 条、P2 2 条、P3 1 条" not in paragraphs  # Approved concise Word summary.
    assert report.model_dump()==saved


def test_incomplete_capacity_and_unavailable_checks_keep_known_usage():
    report,context=_report()
    store=report.asset_inventory["details"]["Datastore"][0]
    store["properties"]={"used_percent":96}
    report.environment_summary.result_status_summary={"passed":5,"failed":0,"unavailable":1}
    source_status=dict(report.environment_summary.result_status_summary)
    payload=PdfReportEngine()._build_payload(report,context)
    assert payload["datastores"][0]["usage"]==96
    assert not payload["datastores"][0]["total_text"]
    assert payload["overcommit_display"]==""  # A32a omits missing customer fields; it does not invent zero.
    assert not payload["scope"]["all_passed"]
    assert report.environment_summary.result_status_summary==source_status
    assert store["properties"]=={"used_percent":96}


def test_certificate_conflict_keeps_original_days_and_judgment():
    report,_=_report()
    report.asset_inventory["details"]["HostSystem"][0]["properties"]["host_certificate_days_remaining"]=20
    saved=report.model_dump()
    assert any("原始证书数据不一致" in x and "20 天" in x for x in delivery_overview(report)["limitations"])
    assert report.model_dump()==saved


def test_word_a4_and_parent_issue_repeated_header(tmp_path):
    report,_=_report()
    doc=Document(DocxReportEngine().render(report,tmp_path/"a4.docx"))
    assert round(doc.sections[0].page_width.cm,1)==21
    assert round(doc.sections[0].page_height.cm,1)==29.7
    from docx.table import Table
    heading=next(p for p in doc.paragraphs if p.text.startswith('附录 B '))
    assert 'CPU Ready' in heading.text and heading.style.name=='Heading 1'
    sibling=heading._p.getnext()
    while sibling is not None and not sibling.tag.endswith('}tbl'):sibling=sibling.getnext()
    assert sibling is not None
    evidence=Table(sibling,doc)
    assert 'tblHeader' in evidence.rows[0]._tr.xml
    assert all('tblHeader' not in r._tr.xml for r in evidence.rows[1:])
    assert all(run.font.size is None or run.font.size.pt>=10 for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs for run in p.runs)


def test_missing_saved_object_results_are_not_reported_as_verified_zero(tmp_path):
    from vstacklens.reports.presentation import customer_check_statistics

    report,_=_report()
    ensure_report(report)
    before=report.model_dump()
    opening=delivery_overview(report)
    assert not opening["recorded_results"]["available"]
    assert opening["check_counts"]=={}
    statistics=customer_check_statistics(report)
    assert not statistics["available"] and statistics["text"]==""
    word=Document(DocxReportEngine().render(report,tmp_path/"missing-results.docx"))
    paragraphs="\n".join(p.text for p in word.paragraphs)
    assert "未保存当次逐对象结果" not in paragraphs
    assert "本次共检查 0 条" not in paragraphs and "通过 0 条" not in paragraphs
    assert opening["conclusion"] not in {"正常","健康","通过"}
    html=HtmlReportPackageBuilder()._single_page_payload(report,{})
    assert not html["summary"]["delivery"]["recorded_results"]["available"]
    assert html["summary"]["checkStatistics"]["text"]==""
    assert 's.checkStatistics?.text ?' in SINGLE_PAGE_REPORT_JS
    assert report.model_dump()==before


def test_complete_target_information_is_readable_without_credentials(window):
    widget,(_,_,_,_,_,controls),app=window
    widget._new_plan();widget.plan_editor_tabs.setCurrentIndex(1)
    widget.plan_targets.setCurrentCell(0,1);app.processEvents()
    labels=widget.plan_target_information.findChildren(QLabel)
    text="\n".join(label.text() for label in labels)
    env=widget._plan_profiles[0]
    assert env.customer_name in text and env.name in text and env.address in text
    assert "密码" not in text
    assert not widget.plan_consent.isChecked() and not controls["calls"]


def test_report_buttons_fit_small_view(window,tmp_path,monkeypatch):
    from vstacklens.application.inspection_service import ReportCenterItem
    widget,_,app=window
    widget.resize(1093,614);widget.go_to_module("report_center");app.processEvents()
    file=tmp_path/"offline-ui.txt";file.write_text("offline fixture",encoding="utf-8")
    monkeypatch.setattr(widget.service,"desktop_report_formats",lambda *a:{"HTML":file,"Word":file,"PDF":file})
    item=ReportCenterItem(run_id="run-0123456789abcdef0123456789abcdef",report_path=file,report_dir=tmp_path,payload_path=None,
        title="离线客户 长名称虚拟化巡检报告",customer_name="离线客户 生产保障与项目交付部门",site_name="离线长名称环境档案",
        assessment_time="2026-10-14T09:00:00+08:00",score=95,risk_summary={},asset_summary={},top_risks=[],history_summary="",status="success",report_type="html_package",path_hint=str(file))
    widget._task_numbers[item.run_id]="S-20261014-01";widget._populate_reports([item]);app.processEvents()
    widget._toggle_report_details();app.processEvents()
    area=widget.report_splitter.widget(1)
    assert area.isVisible()
    assert area.widget().width()<=area.viewport().width()
    for control in (widget.report_word_button,widget.report_pdf_button):
        assert control.isVisible() and control.isEnabled()
        right=control.mapTo(widget,control.rect().topRight()).x()
        assert right<widget.width()


def test_word_host_counter_is_not_a_nic_link_state():
    from vstacklens.reports.docx_report import RiskGroup
    from vstacklens.reports.report_model import FindingItem
    report,_=_report()
    finding=FindingItem(rule_id="VSL-NET-003",rule_name="物理网卡错误计数",title="物理网卡错误计数异常",risk_level="P3",object_type="HostSystem",object_name="esxi-01",status="failed",current_value_zh="12 次",affected_components=["net.errorsRx.summation","net.droppedTx.summation"])
    headers,rows=DocxReportEngine()._detail_table(report,RiskGroup("P3","网络",finding.title,[finding],""))
    assert headers==["ESXi 主机","指标 / 范围","当次记录","判断限制与建议"]
    assert len(rows)==1 and rows[0][2]=="12 次"
    assert all(m in rows[0][1] for m in finding.affected_components)
    assert "不将主机级计数拆分" in rows[0][3]
    finding.observed_detail={"pnic_error_count":12,"pnic_error_detail":[{"counter":"net.errorsRx.summation","value":7},{"counter":"net.droppedTx.summation","value":5}]}
    _,rows=DocxReportEngine()._detail_table(report,RiskGroup("P3","网络",finding.title,[finding],""))
    assert [r[2] for r in rows]==["7 次","5 次"] and all("12 次" in r[3] for r in rows)


def test_html_certificate_conflict_preserves_saved_days():
    certificates=HtmlReportPackageBuilder()._single_page_certificates({"run_timing":{"started_at":"2026-10-13"}},
        {"host":[{"name":"esxi-01","cluster":"C","certificateDate":"2026-06-24","certificateDays":20}]},{"name":"vc"},[])
    entry=certificates["clusters"][0]["hosts"][0]
    assert entry["remainingDays"]==20 and entry["remainingLabel"]=="20 天" and not entry["expired"]
    assert entry["dataConflict"] and entry["imminent"] and entry["expiresOn"]=="2026-06-24"


def test_opening_preserves_recorded_conclusion_advice_and_limits():
    report, _ = _report()
    ensure_report(report)
    report.environment_summary.result_status_summary = {"passed": 3, "unavailable": 2, "error": 1}
    report.appendix["unavailable_rules"] = [{"reason": "权限不足：网络清单未取得"}]
    before = report.model_dump()
    opening = delivery_overview(report)
    assert opening["conclusion"] == report.overall_status
    assert opening["check_counts"] == {}  # Only saved, explicitly graded records may form customer totals.
    assert not opening["recorded_results"]["available"]
    assert report.environment_summary.result_status_summary == {"passed":3,"unavailable":2,"error":1}
    assert not any("检查记录不可用" in x or "检查记录出错" in x for x in opening["limitations"])
    assert "权限不足：网络清单未取得" not in opening["limitations"]
    assert report.appendix["unavailable_rules"] == before["appendix"]["unavailable_rules"]
    assert opening["priorities"][0]["suggestion"] == report.customer_caliber["categories"][0]["remediation_zh"]
    assert report.model_dump() == before


def test_three_formats_use_identical_opening_and_keep_source(tmp_path):
    report, context = _report()
    ensure_report(report, context)
    before = report.model_dump()
    expected = delivery_overview(report)
    html = HtmlReportPackageBuilder()._single_page_payload(report, context)
    pdf = PdfReportEngine()._build_payload(report, context)
    word = DocxReportEngine().render(report, tmp_path / "report.docx")
    assert html["summary"]["delivery"] == pdf["delivery"] == expected
    word_document=Document(word)
    paragraphs = "\n".join(word_document.element.xpath('//w:t/text()'))+''.join(p.text for p in word_document.sections[0].header.paragraphs)
    assert expected["customer"] in paragraphs and expected["scope"] in paragraphs
    from vstacklens.reports.document_model import get_document
    assert get_document(report,context)['opening']['conclusion'] in paragraphs
    assert get_document(report,context)['opening']['conclusion'] in {'危险','关注','存在整改事项','健康','评估受限'}
    # The batch-four cover only contains the four approved customer summaries;
    # internal legacy-caliber explanations remain in the preserved source.
    from vstacklens.reports.document_model import get_document
    opening=get_document(report,context)['opening']
    assert set(opening)=={'conclusion','problem_text','priorities','more','statistics'}
    assert report.model_dump() == before
    assert {p["level"] for p in pdf["problems"]} <= {"P1", "P2", "P3"}
    assert all("basis" in p for p in pdf["problems"])


def test_failed_connection_cannot_be_described_as_verified_health():
    report, _ = _report()
    report.asset_inventory["details"]["vCenter"][0]["properties"]["collection_mode"] = "connection_failed"
    original = report.model_dump()
    opening = delivery_overview(report)
    assert opening["conclusion"] not in {"正常", "健康", "通过"}
    assert not any("连接未成功" in x for x in opening["limitations"])
    assert report.model_dump() == original


def test_word_missing_flags_remain_uncollected_and_zero_is_not_missing():
    engine = DocxReportEngine()
    assert engine._boolean_label(None, "可访问", "不可访问") == "未采集"
    assert engine._boolean_label(True, "可访问", "不可访问") == "可访问"
    assert engine._boolean_label(False, "可访问", "不可访问") == "不可访问"
    assert engine._fmt(0) == "0"


def test_number_summary_keeps_missing_distinct_from_zero(qt_application_lifetime):
    widget = NumberSummary()
    widget.set_counts(None, "")
    assert "None" not in widget.text() and "—" in widget.text()
    widget.set_counts(0, 0)
    assert widget.text() == "0 类问题 · 0 个涉及对象"
    widget.deleteLater()


def test_word_chapter_order_native_toc_and_page_numbers(tmp_path):
    report, _ = _report()
    target = DocxReportEngine().render(report, tmp_path / "report.docx")
    doc = Document(target)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    numbered = [int(h.split(".", 1)[0].split(" ", 1)[0]) for h in headings if h[0].isdigit()]
    assert numbered == sorted(numbered)
    assert 'TOC \\o "1-3"' in doc.element.xml
    assert "PAGE" in doc.sections[0].footer._element.xml
    assert doc.styles["Heading 1"].paragraph_format.keep_with_next
    assert '<w:br w:type="page"' not in doc.element.xml
    assert next(p for p in doc.paragraphs if p.text == "目录").paragraph_format.page_break_before
    import re
    assert [re.sub(r'^\d{2}\s+','',h) for h in headings if re.match(r'^\d{2}\s+',h)]==['问题总览','问题详情','环境详情']
    assert any(h.startswith('附录 A ') for h in headings)
    for table in doc.tables:
        # Single-row bars are charts, not tabular data headers.
        if len(table.rows) > 1 and table.cell(0, 0).text != "Resync 状态":
            assert "tblHeader" in table.rows[0]._tr.xml
            assert all(p.paragraph_format.keep_with_next for c in table.rows[0].cells for p in c.paragraphs)


def test_pdf_empty_usage_section_is_conditional_and_missing_not_zero(tmp_path):
    import pymupdf

    template = (Path(__file__).resolve().parents[1] / "src/vstacklens/reports/templates/inspection-report.typ").read_text(encoding="utf-8")
    assert '#if data.top_cpu.len() > 0 or data.top_memory.len() > 0 [' in template
    assert '虚拟机 CPU / 内存实际用量未采集' not in template
    # Display formatting is shared in Python rather than independently rounded in Typst.
    assert PdfReportEngine()._percent_text(None) == ""
    assert PdfReportEngine()._percent_text(0) == "0.0%"
    assert '#let num(value) = if value == none { "" }' in template
    report,context=_report()
    vm=report.asset_inventory["details"]["VirtualMachine"][0]["properties"]
    vm.pop("cpu_usage_mhz"); vm.pop("memory_usage_mb")
    source_missing=dict(vm)
    engine=PdfReportEngine()
    missing=engine._build_payload(report,context)
    assert missing["top_cpu"]==[] and missing["top_memory"]==[]
    missing_path=engine.render(report,tmp_path/"missing-usage.pdf",report_context=context)
    with pymupdf.open(missing_path) as document:
        missing_text="".join("".join(page.get_text().split()) for page in document)
        missing_headings=["".join(item[1].split()) for item in document.get_toc()]
    assert not any("虚拟机实际资源用量Top5" in heading for heading in missing_headings)
    assert "虚拟机CPU/内存实际用量未采集" not in missing_text
    assert vm==source_missing
    vm.update(cpu_usage_mhz=0,memory_usage_mb=0)
    zero=engine._build_payload(report,context)
    assert len(zero["top_cpu"])==len(zero["top_memory"])==1
    assert zero["top_cpu"][0]["value"]==0 and zero["top_memory"][0]["value"]==0
    assert zero["top_cpu"][0]["display"]=="0.00 GHz"
    assert zero["top_memory"][0]["display"]=="0.0 GB"
    zero_path=engine.render(report,tmp_path/"zero-usage.pdf",report_context=context)
    with pymupdf.open(zero_path) as document:
        zero_text="".join("".join(page.get_text().replace('\u200b','').split()) for page in document)
        zero_headings=["".join(item[1].split()) for item in document.get_toc()]
    assert 'CPU实际用量Top5' in zero_text and '内存实际用量Top5' in zero_text
    assert "0.00GHz" in zero_text and "0.0GB" in zero_text
    assert vm["cpu_usage_mhz"]==0 and vm["memory_usage_mb"]==0


def test_pdf_zero_problem_explanations_stay_compact_and_small_bars_do_not_split(tmp_path):
    import pymupdf

    def recorded_report(status):
        report,context=_report()
        row=dict(rule_id="VSL-VM-004",risk_level="P2",result_status=status,status=status,
                 object_name="app-01",object_type="VirtualMachine",rule_name="CPU Ready",check_name="CPU Ready")
        context["object_results"]=[row]
        report.caliber_inputs["object_results"]=[dict(row)]
        if status=="passed":
            report.findings=[]
        ensure_report(report,context)
        return report,context

    # Missing records are omitted rather than displayed as verified zero checks.
    missing,missing_context=_report()
    missing.findings=[]
    missing_engine=PdfReportEngine()
    missing_payload=missing_engine._build_payload(missing,missing_context)
    assert missing_payload["scope"]["no_results"] and not missing_payload["scope"]["all_passed"]
    missing_path=missing_engine.render(missing,tmp_path/"no-records.pdf",report_context=missing_context)
    with pymupdf.open(missing_path) as document:
        missing_headings=["".join(item[1].split()) for item in document.get_toc()]
        missing_text="".join("".join(page.get_text().split()) for page in document)
    assert not any("巡检范围结果" in heading for heading in missing_headings)
    assert "本次共检查0条" not in missing_text and "全部通过" not in missing_text

    # A real passed record retains its failed=0 value and concise zero-problem presentation.
    zero,zero_context=recorded_report("passed")
    zero_engine=PdfReportEngine()
    zero_payload=zero_engine._build_payload(zero,zero_context)
    assert (zero_payload["scope"]["total_checks"],zero_payload["scope"]["passed_checks"],zero_payload["scope"]["problem_checks"])==(1,1,0)
    assert zero_payload["problem_count"]==0 and zero_payload["problems"]==[]
    zero_path=zero_engine.render(zero,tmp_path/"zero-problems.pdf",report_context=zero_context)
    with pymupdf.open(zero_path) as document:
        zero_pages=len(document)
        zero_headings=["".join(item[1].split()) for item in document.get_toc()]
        zero_text="".join("".join(page.get_text().split()) for page in document)
        assert "本次共检查1条，通过1条，未通过0条。" in zero_text
        assert "未通过检查结果按等级分布" not in zero_text and "全部通过" not in zero_text
        # V-01 restores honest category counts, including known zero values;
        # the zero-problem report still has no failed-result distribution.
        cover=document[0]
        # The requested cover de-duplication leaves totals in the metric
        # cards, rather than repeating the combined sentence elsewhere.
        assert '0类问题、0个受影响对象' not in ''.join(cover.get_text().split())
        assert cover.search_for('问题类别') and cover.search_for('受影响对象')
        cover_text=''.join(cover.get_text().split())
        assert all(level+'0' in cover_text for level in ('P1','P2','P3'))
    failed,failed_context=recorded_report("failed")
    failed_path=PdfReportEngine().render(failed,tmp_path/"one-problem.pdf",report_context=failed_context)
    with pymupdf.open(failed_path) as document:
        assert zero_pages<=len(document)  # Empty issue sections must not expand beyond a real issue card.


def test_html_print_expands_all_evidence_and_restores_reading_state():
    assert 'addEventListener("beforeprint"' in SINGLE_PAGE_REPORT_JS
    assert 'addEventListener("afterprint"' in SINGLE_PAGE_REPORT_JS
    assert "printDetailsState.set(detail, detail.open)" in SINGLE_PAGE_REPORT_JS
    assert "detail.open = true" in SINGLE_PAGE_REPORT_JS
    assert "detail.open = open" in SINGLE_PAGE_REPORT_JS
    assert "min-width:0!important" in HtmlReportPackageBuilder()._single_page_css()


def test_pdf_missing_network_evidence_is_not_a_verified_zero_and_maintenance_keeps_connection():
    from copy import deepcopy

    report, context = _report()
    host = report.asset_inventory["details"]["HostSystem"][0]["properties"]
    host["in_maintenance_mode"] = True
    original_host=deepcopy(host)
    payload = PdfReportEngine()._build_payload(report, context)
    assert payload["connected_host_count"] == 1
    assert payload["network"]["vmkernel_count"] == 0  # original normalized data stays intact
    assert payload["network"]["vmkernel_display"] == ""
    assert payload["network"]["pnic_anomaly_display"] == ""
    assert payload["network"]["pnic_anomaly_count"] is None
    assert payload["network"]["error_count"] is None
    assert payload["narratives"]["network"] == ""
    assert host==original_host and host["connection_state"]=="connected" and host["in_maintenance_mode"] is True
    assert "端口组策略异常 0 项" not in payload["narratives"]["security"]
    host.update(vmkernel_adapters=[], physical_nic_details=[], portgroup_security_issue_count=0)
    original_empty_host=deepcopy(host)
    payload = PdfReportEngine()._build_payload(report, context)
    assert payload["network"]["vmkernel_display"] == "0"
    assert payload["network"]["pnic_anomaly_display"] == "0 条"
    assert payload["network"]["pnic_anomaly_count"] == 0
    assert payload["network"]["error_count"] is None
    assert payload["narratives"]["network"] == "在用物理链路未发现异常。"
    assert "错误或丢包计数 0" not in payload["narratives"]["network"]
    assert payload["connected_host_count"] == 1
    assert host==original_empty_host and host["in_maintenance_mode"] is True
    assert "policy_display" not in payload["network"]
    assert "portgroup_details" not in payload["security"]


def test_actual_plan_editor_compacts_redundant_headers_without_bypassing_consent(window):
    widget, (_, _, _, _, _, controls), app = window
    widget.resize(1093, 614)  # logical size for 1366 x 768 at 125 percent
    widget._new_plan(); widget.plan_editor_tabs.setCurrentIndex(1)
    app.processEvents()
    assert not widget.plan_page_header.isVisible()
    assert not widget.plan_status.isVisible()
    assert widget.plan_targets.viewport().height() >= 125
    assert not widget.plan_consent.isChecked() and not widget.plan_authorize_button.isEnabled()
    assert widget.plan_save_button.isVisible()
    assert widget.plan_save_button.mapTo(widget, widget.plan_save_button.rect().bottomRight()).y() < widget.height()
    assert not controls["calls"]
    assert all(widget.plan_targets.item(r, 1).toolTip() for r in range(widget.plan_targets.rowCount()))
    widget.plan_tabs.setCurrentIndex(0); app.processEvents()
    assert widget.plan_page_header.isVisible()
