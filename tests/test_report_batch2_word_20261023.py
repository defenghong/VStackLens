"""Batch two A: customer Word uses shared facts without rejudging sources."""
from copy import deepcopy

from docx import Document

from vstacklens.reports.docx_report import DocxReportEngine
from vstacklens.reports.report_model import (
    CustomerInfo, EnvironmentSummary, FindingItem, HealthScore, ReportData, ReportInfo, RiskSummary,
)


def report_with_assets(details=None):
    return ReportData(
        report_info=ReportInfo(report_id="b2-word", run_id="b2-word", generated_at="2026-10-22T00:00:00Z"),
        customer_info=CustomerInfo(customer_name="测试客户", site_name="实验环境", vcenter="vc.test"),
        environment_summary=EnvironmentSummary(scope_statement="保存的巡检数据", checked_object_total=116),
        health_score=HealthScore(score=80, grade="关注", explanation="保存的结论"),
        risk_summary=RiskSummary(), asset_inventory={"details": details or {}, "summary": {"VirtualMachine": 116}},
    )


def visible_text(document):
    return "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )


def test_summary_displays_only_shared_problem_and_object_counts(monkeypatch):
    engine = DocxReportEngine()
    facts = {
        "counts": {"problem_count": 10, "affected_total": 72, "problem_label": "问题类别", "affected_label": "受影响对象"},
        "snapshots": {"overview_label": "有快照的虚拟机 20 台"},
    }
    monkeypatch.setattr(engine, "_display_facts", lambda _: facts)
    report = report_with_assets()
    before = deepcopy(report.model_dump())
    document = Document()
    engine._executive_summary(document, report, [])
    text = visible_text(document)
    assert "问题类别：10 类；受影响对象：72 个。" in text
    assert "有快照的虚拟机 20 台" in text
    assert all(value not in text for value in ("91", "94", "28 条", "22 条", "17 台", "跨等级对象", "等级内去重"))
    assert report.model_dump() == before


def test_summary_displays_shared_execution_statistics_without_counting_again(monkeypatch):
    engine = DocxReportEngine()
    facts = {
        "counts": {"problem_count": 10, "affected_total": 71, "problem_label": "问题类别", "affected_label": "受影响对象"},
        "snapshots": {"overview_label": "有快照的虚拟机 20 台"},
    }
    monkeypatch.setattr(engine, "_display_facts", lambda _: facts)
    statistics = {"total": 1378, "passed": 1285, "failed": 93,
                  "text": "本次共检查1378项，通过1285项，未通过93项。", "by_object_type_rows": []}
    received_reports = []

    def saved_statistics(report, context=None):
        received_reports.append(report)
        return statistics

    monkeypatch.setattr("vstacklens.reports.presentation.customer_check_statistics", saved_statistics)
    report = report_with_assets()
    before = deepcopy(report.model_dump())
    document = Document()
    engine._executive_summary(document, report, [])
    assert received_reports == [report]
    assert [paragraph.text for paragraph in document.paragraphs].count(statistics["text"]) == 1
    assert "受影响对象：71 个。" in visible_text(document)
    assert report.model_dump() == before


def test_cover_does_not_show_redundant_record_level_counts(monkeypatch):
    monkeypatch.setattr("vstacklens.reports.docx_report.delivery_overview", lambda _: {
        "environment": "", "scope": "vc.test", "conclusion": "关注", "explanation": "保存的结论",
        "priorities": [], "limitations": [],
        "recorded_results": {"available": True, "failed_levels": {"P1": 0, "P2": 19, "P3": 75}},
    })
    document = Document()
    DocxReportEngine()._cover(document, report_with_assets())
    text = visible_text(document)
    assert "总体结论：关注" in text
    assert "当次未通过结果条数" not in text
    assert "P2 19 条" not in text and "P3 75 条" not in text


def test_certificate_dates_use_shared_object_identity_and_source(monkeypatch):
    report = report_with_assets({
        "vCenter": [{"object_key": "vc-a", "object_name": "vc.test", "properties": {"certificate_days_remaining": 581}}],
        "HostSystem": [
            {"object_key": "host-a", "object_name": "same-name", "properties": {}},
            {"object_key": "host-b", "object_name": "same-name", "properties": {}},
            {"object_key": "host-c", "object_name": "no-date", "properties": {}},
        ],
    })
    certificates = {
        "vcenter": {"object_key": "vc-a", "name": "vc.test", "days": 581, "expiry_date": "2028-05-25",
                    "date_source": "按采集日期和保存的剩余天数推算", "derived": True},
        "hosts": [
            {"object_key": "host-a", "name": "same-name", "days": 600, "expiry_date": "2028-06-13",
             "date_source": "按采集日期和保存的剩余天数推算", "derived": True},
            {"object_key": "host-b", "name": "same-name", "expiry_date": "2031-06-01",
             "date_source": "ESXi 443 TLS 证书直接读取", "derived": False},
        ],
    }
    engine = DocxReportEngine()
    monkeypatch.setattr(engine, "_display_facts", lambda _: {"certificates": certificates})
    before = deepcopy(report.model_dump())
    rows = engine._version_certificate_license_rows(report)
    assert "2028-05-25" in rows[0][2] and "推算" in rows[0][2]
    assert "600 天" in rows[1][2] and "2031" not in rows[1][2]
    assert "2031-06-01" in rows[2][2] and "443 TLS" in rows[2][2]
    assert "剩余" not in rows[2][2]  # A date must not invent remaining days.
    assert rows[3][2] == ""  # No date or day placeholder for a missing certificate fact.
    assert report.model_dump() == before


def test_missing_chart_values_are_hidden_but_known_zero_is_retained():
    engine = DocxReportEngine()
    missing = Document()
    engine._add_vsan_object_chart(missing, {})
    engine._add_vsan_capacity_chart(missing, {})
    assert len(missing.tables) == 0
    assert not visible_text(missing)
    zero = Document()
    engine._add_vsan_object_chart(zero, {"object_count": 0, "object_issue_count": 0})
    assert "vSAN 对象 0；异常 0" in visible_text(zero)
    assert all(value not in visible_text(zero) for value in ("未确认", "待确认", "不可用"))


def test_vsan_section_uses_measured_shared_facts_without_overall_normal(monkeypatch):
    report = report_with_assets()
    report.vsan_summary = {
        "status": "collected", "health_issue_count": 0, "object_issue_count": 0,
        "object_count": 175, "vmdk_count": 89, "resync_object_count": 0, "resync_bytes": 0,
        "capacity": {"used_percent": 11.4},
        "storage_policy_summary": {"status": "unknown"},
    }
    facts = {"items": [
        {"key": "objects", "text": "已识别 175 个 vSAN 对象，健康异常 0 项"},
        {"key": "capacity", "text": "容量使用率 11.4%"},
        {"key": "resync", "text": "当前无待重同步对象或数据"},
    ], "summary": "已识别 175 个 vSAN 对象，健康异常 0 项；容量使用率 11.4%；当前无待重同步对象或数据。"}
    engine = DocxReportEngine()
    engine._customer_mode = True
    monkeypatch.setattr(engine, "_display_facts", lambda _: {"vsan": facts})
    before = deepcopy(report.model_dump())
    document = Document()
    engine._vsan_section(document, report)
    text = visible_text(document)
    assert facts["summary"] in text
    assert all(item["text"] in text for item in facts["items"])
    assert all(value not in text for value in ("整体状态正常", "总体 Health 正常", "未确认", "待确认", "不可用", "Storage Policy"))
    assert report.model_dump() == before


def test_local_storage_customer_titles_use_shared_display_without_changing_source():
    original_title = "虚拟机运行在本地存储上"
    approved_title = "虚拟机磁盘存放在本地存储"
    report = report_with_assets({"VirtualMachine": [{
        "object_type": "VirtualMachine", "object_key": "vm-local", "object_name": "app-local",
        "asset_location": "vc.test / Datacenter / Cluster-A / 主机 host-a",
        "properties": {"power_state": "poweredOn", "is_template": False, "datastore_names": ["Local-A"]},
    }]})
    report.findings = [FindingItem(
        risk_level="P3", rule_id="VSL-VM-023", rule_name=original_title, title=original_title,
        object_type="VirtualMachine", object_name="app-local", status="open",
        current_value=True, expected_value=False, evidence={"object_key": "vm-local"},
        observed_detail={"local_datastore_names": ["Local-A"]},
        business_impact="可能影响主机维护时的迁移安排。", remediation="先核对用途、备份与恢复要求。",
        remediation_steps=["确认用途及Datastore。", "完成迁移或确认例外。", "复跑巡检。"],
    )]
    original_assets = deepcopy(report.asset_inventory)
    original_findings = [finding.model_dump() for finding in report.findings]
    engine = DocxReportEngine()
    groups = engine._aggregate_risks(report)
    assert len(groups) == 1 and len(groups[0].findings) == 1
    assert groups[0].title == original_title  # Internal source identity is retained.
    summary = Document()
    engine._executive_summary(summary, report, groups)
    summary_text = visible_text(summary)
    assert approved_title in summary_text and original_title not in summary_text
    details = Document()
    engine._risk_recommendations(details, report, groups)
    detail_text = visible_text(details)
    assert any(paragraph.text.endswith(approved_title) for paragraph in details.paragraphs)
    assert approved_title in detail_text and original_title not in detail_text
    follow_up = Document()
    engine._follow_up(follow_up, groups, report=report)
    assert approved_title in visible_text(follow_up) and original_title not in visible_text(follow_up)
    cover = Document()
    engine._cover(cover, report)
    assert approved_title in visible_text(cover) and original_title not in visible_text(cover)
    assert [finding.model_dump() for finding in report.findings] == original_findings
    assert report.asset_inventory == original_assets
    assert len(groups) == 1 and len(groups[0].findings) == 1
