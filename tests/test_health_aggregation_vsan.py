from __future__ import annotations

from datetime import UTC, datetime
import json
from pathlib import Path
from types import SimpleNamespace as NS

from docx import Document
from pyVmomi import vim

from vstacklens.collection.pyvmomi_collector import PyVmomiCollector
from vstacklens.core.context import RunContext
from vstacklens.db.connection import connect, init_db
from vstacklens.db.repositories import ensure_default_scope, finalize_run_summary, insert_rule, insert_rule_result, insert_run
from vstacklens.findings.deduplication import FindingDeduplicator
from vstacklens.reports.health_status import assess_health
from vstacklens.reports.context_builder import ReportContextBuilder
from vstacklens.reports.docx_report import DocxReportEngine
from vstacklens.reports.html_package import HtmlReportPackageBuilder
from vstacklens.reports.report_model import ReportDataFactory
from vstacklens.inventory.normalizer import InventoryNormalizer
from vstacklens.rules.executor import RuleExecutor
from vstacklens.rules.rulepack_loader import RulePackLoader
from vstacklens.rules.schema_validator import SchemaValidator
from vstacklens.reports.presentation import (
    build_health_impact_summary,
    build_risk_category_summary,
    build_risk_group_summary,
    build_risk_object_summary,
)


def test_same_rule_hits_many_vms_counts_one_category_and_many_objects() -> None:
    findings = [
        {"rule_id": "VSL-VM-001", "risk_level": "P1", "title": "虚拟机存在超期快照", "object_name": f"vm-{index}", "health_impact": "attention"}
        for index in range(15)
    ]
    assert build_risk_category_summary(findings) == {"P1": 1, "P2": 0, "P3": 0, "P4": 0}
    assert build_risk_object_summary(findings) == {"P1": 15, "P2": 0, "P3": 0, "P4": 0}
    assert build_risk_group_summary(findings)[0]["object_count"] == 15


def test_snapshot_priority_does_not_make_environment_critical() -> None:
    health = assess_health({"P1": 1, "P2": 0, "P3": 0}, {"passed": 10}, 2, health_impact={"none": 0, "attention": 1, "critical": 0})
    assert health["label"] == "关注"
    assert health["label"] != "危险"


def test_vsan_object_inaccessible_can_make_environment_critical() -> None:
    health = assess_health({"P1": 1, "P2": 0, "P3": 0}, {"passed": 10}, 2, health_impact={"none": 0, "attention": 0, "critical": 1})
    assert health["label"] == "危险"


def test_resync_only_does_not_create_vsan_issue() -> None:
    collector = object.__new__(PyVmomiCollector)
    result = collector._vsan_datastore_info(
        True,
        ["cluster-a"],
        76.0,
        24.0,
        100.0,
        False,
        {"status": "collected", "clusters": {"cluster-a": {
            "api_status": "collected",
            "cluster_enabled": True,
            "health_issues": [],
            "disk_health_issues": [],
            "object_health_issues": [],
            "resync_object_count": 3,
            "resync_bytes": 1024,
        }}},
    )
    assert result["vsan_resync_object_count"] == 3
    assert result["vsan_issue_count"] == 0
    assert result["vsan_capacity_status"] == "attention"


def test_html_report_has_five_linear_chapters_and_print_expansion() -> None:
    html = HtmlReportPackageBuilder()._html({"report_context": {"report_info": {"report_title": "test"}, "vsan_summary": {"status": "collected"}}})
    for section in ("summary", "environment", "issues", "passed", "inventory"):
        assert f'href="#{section}"' in html
        assert f'id="{section}" class="report-section"' in html
    assert 'href="#vcenter"' not in html
    assert 'href="#vsan"' not in html
    assert 'href="#history"' not in html
    assert 'class="report-module"' not in html
    css = HtmlReportPackageBuilder()._css()
    one_page_css = HtmlReportPackageBuilder()._single_page_css()
    assert ".report-nav { position:sticky" in one_page_css
    assert "@media print" in one_page_css
    assert "@media print" in css
    html_without_vsan = HtmlReportPackageBuilder()._html({"report_context": {"report_info": {"report_title": "test"}, "vsan_summary": {"status": "not_applicable"}}})
    assert 'id="vsan"' not in html_without_vsan
    assert [html.find(f'href="#{section}"') for section in ("summary", "environment", "issues", "passed", "inventory")] == sorted(html.find(f'href="#{section}"') for section in ("summary", "environment", "issues", "passed", "inventory"))


def test_html_storage_policy_is_conditional_and_categorized() -> None:
    builder = HtmlReportPackageBuilder()
    base = {"report_info": {"report_title": "test"}, "vsan_summary": {"applicable": True, "status": "collected"}}
    without_policy = builder._html({"report_context": base})
    assert "存储策略分类" not in without_policy
    with_policy = builder._html({"report_context": {**base, "vsan_summary": {
        **base["vsan_summary"],
        "storage_policy_summary": {
            "source": "vsan_object_api", "scope": "vsan_object", "coverage": "object_only", "pbm_confirmed": False,
            "checked_count": 1, "compliant_count": 1, "noncompliant_count": 0, "unknown_count": 0,
            "policy_categories": [{"policy_name": "Policy", "policy_uuid": "profile-1", "checked_count": 1, "compliant_count": 1, "noncompliant_count": 0, "unknown_count": 0}],
        },
    }}})
    script = builder._js()
    assert "存储策略分类" in script
    assert '["总体健康"' in script
    assert '["重同步对象数"' in script
    assert '["容量使用率"' in script
    assert '["主机", "设备", "IP", "子网", "网络标签", "MTU"]' in script
    assert '<details><summary>查看 VMkernel 明细' in script
    css = builder._css()
    assert ".grid.vsan-summary-grid" in css
    assert ".vsan-checks-table th:nth-child(4)" in css
    assert "#vsan .vsan-checks-table { min-width: 640px; }" in css
    assert "policyAvailable" in script
    assert "profile-1" in with_policy
    assert "vsan_object_api" in with_policy


def test_vsan_unknown_and_zero_are_distinct() -> None:
    collector = object.__new__(PyVmomiCollector)
    unavailable = collector._vsan_datastore_info(True, ["cluster-a"], 85.0, 15.0, 100.0, True, {"status": "unavailable", "clusters": {}})
    collected = collector._vsan_datastore_info(True, ["cluster-a"], 85.0, 15.0, 100.0, True, {"status": "collected", "clusters": {"cluster-a": {
        "api_status": "collected", "cluster_enabled": True, "health_issues": [], "disk_health_issues": [], "object_health_issues": [], "resync_object_count": 0, "resync_bytes": 0,
    }}})
    assert unavailable["vsan_resync_object_count"] is None
    assert unavailable["vsan_issue_count"] is None
    assert collected["vsan_resync_object_count"] == 0
    assert collected["vsan_resync_bytes"] == 0


def test_vsan_report_context_preserves_missing_resync_values() -> None:
    row = {
        "object_name": "vsanDatastore",
        "object_type": "Datastore",
        "properties_json": json.dumps({
            "datastore_is_vsan": True,
            "vsan_cluster_names": ["cluster-a"],
            "vsan_api_status": "unsupported",
            "vsan_resync_object_count": None,
            "vsan_resync_bytes": None,
        }),
    }

    summary = ReportContextBuilder()._vsan_summary([row])

    assert summary["resync_object_count"] is None
    assert summary["resync_bytes"] is None
    document = Document()
    DocxReportEngine()._resync_status_table(document, summary)
    resync_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "状态未确认" in resync_text
    assert resync_text.count("未采集") == 4


def test_vsan_network_cross_checks_selected_adapter_and_recovers_missing_label() -> None:
    rows = [
        {
            "object_name": "vSAN",
            "object_type": "Datastore",
            "properties_json": json.dumps({
                "datastore_is_vsan": True,
                "datastore_host_names": ["esxi-01"],
                "vsan_api_status": "collected",
            }),
        },
        {
            "object_name": "esxi-01",
            "object_type": "HostSystem",
            "properties_json": json.dumps({
                "vsan_enabled": True,
                "vmkernel_adapters": [
                    {"device": "vmk0", "network_label": "Management Network"},
                    {"device": "vmk1", "network_label": "", "ip_address": "10.250.20.69"},
                ],
                "vsan_vmk_adapters": [
                    {"device": "vmk1", "network_label": "vsan", "ip_address": "10.250.20.69"},
                ],
            }),
        },
    ]
    summary = ReportContextBuilder()._vsan_summary(rows)
    network = summary["network"]
    assert len(network["vmkernels"]) == 2
    assert network["vmkernels"][1]["network_label"] == "vsan"
    assert network["cross_validation"] == {"path_count": 2, "checked": 1, "recovered_labels": 1, "mismatches": 0}


def test_vsan_report_semantics_promote_only_confirmed_abnormal_categories() -> None:
    builder = ReportContextBuilder()
    categories, findings, _ = builder._vsan_report_semantics({
        "status": "collected",
        "clusters": ["vSAN"],
        "health_issue_count": 1,
        "disk_issue_count": 0,
        "object_issue_count": 0,
        "object_count": 100,
        "vmdk_count": 50,
        "resync_object_count": 2,
        "resync_bytes": 1024,
        "physical_disks": [{"name": "naa.1", "host": "esxi-01", "summary_health": "green"}],
        "disk_topology": [{"host": "esxi-01", "disk_groups": [{"cache_disk": "naa.1", "capacity_disks": []}]}],
        "network": {"status": "collected", "vmkernels": [{"ip_address": "192.168.1.10", "network_label": "vsan"}]},
        "capacity": {"status": "normal", "used_percent": 10},
    })
    status = {item["category_id"]: item["status"] for item in categories}
    assert status["VSAN-CLUSTER-HEALTH"] == "需关注"
    assert status["VSAN-RESYNC"] == "正常"
    assert [item["rule_id"] for item in findings] == ["VSAN-01"]
    assert findings[0]["risk_level"] == "P1"
    assert builder._remediation_plan(findings)[0]["rule_id"] == "VSAN-01"


def test_vsan_report_semantics_keeps_missing_disk_health_and_network_label_unconfirmed() -> None:
    categories, findings, _ = ReportContextBuilder()._vsan_report_semantics({
        "status": "collected",
        "clusters": ["vSAN"],
        "health_issue_count": 0,
        "disk_issue_count": 0,
        "object_issue_count": 0,
        "object_count": 100,
        "vmdk_count": 50,
        "resync_object_count": 0,
        "resync_bytes": 0,
        "disk_details": [{"health": "未确认"}],
        "network": {"vmkernels": [{"device": "vmk1", "ip_address": "10.250.20.69", "network_label": "未记录"}]},
        "capacity": {"status": "normal", "used_percent": 10},
        "policy_noncompliant_count": 0,
        "storage_policy_summary": {"checked_count": 0, "policy_categories": []},
    })
    statuses = {item["category_id"]: item["status"] for item in categories}
    assert statuses["VSAN-DISK"] == "未确认"
    assert statuses["VSAN-NETWORK"] == "未确认"
    assert statuses["VSAN-POLICY"] == "未确认"
    assert findings == []


def test_v2_vsan_risk_and_visibility_flow_from_saved_inventory(tmp_path) -> None:
    db_path = tmp_path / "vsan-report.db"
    init_db(db_path)
    with connect(db_path) as conn:
        customer_id, site_id, vcenter_id = ensure_default_scope(conn, "vc.local")
        run_id = insert_run(conn, customer_id, site_id, vcenter_id, "manual")
        conn.execute(
            "INSERT INTO inventory_snapshots (snapshot_id, run_id, customer_id, site_id, vcenter_id, snapshot_type, collected_at, object_count, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            ("snapshot-vsan", run_id, customer_id, site_id, vcenter_id, "vcenter", "2026-09-24T00:00:00Z", 2, "2026-09-24T00:00:00Z"),
        )
        datastore = {
            "datastore_is_vsan": True,
            "datastore_host_names": ["esxi-01"],
            "vsan_cluster_names": ["vSAN"],
            "vsan_api_status": "collected",
            "vsan_architecture": "osa",
            "vsan_health_issues": [],
            "vsan_disk_health_issues": [],
            "vsan_object_health_issues": [],
            "vsan_disk_group_count": 1,
            "vsan_cache_disk_count": 1,
            "vsan_capacity_disk_count": 0,
            "vsan_disk_topology": [{"host": "esxi-01", "disk_groups": [{"cache_disk": "naa.cache", "capacity_disks": []}]}],
            "vsan_physical_disks": [{"host": "esxi-01", "name": "naa.cache", "summary_health": "green"}],
            "vsan_object_count": 10,
            "vsan_vmdk_count": 5,
            "vsan_resync_object_count": 2,
            "vsan_resync_bytes": 1024,
            "vsan_used_percent": 10.0,
            "datastore_capacity_gb": 100.0,
            "datastore_used_gb": 10.0,
            "datastore_free_gb": 90.0,
        }
        host = {
            "vsan_enabled": True,
            "vmkernel_adapters": [{"device": "vmk1", "ip_address": "10.250.20.69", "network_label": "vsan", "mtu": 1500}],
        }
        for object_id, object_type, object_name, properties in (
            ("datastore-1", "Datastore", "vSAN", datastore),
            ("host-1", "HostSystem", "esxi-01", host),
        ):
            conn.execute(
                "INSERT INTO inventory_objects (object_id, snapshot_id, run_id, customer_id, vcenter_id, object_type, object_key, object_name, properties_json, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (object_id, "snapshot-vsan", run_id, customer_id, vcenter_id, object_type, object_id, object_name, json.dumps(properties), "2026-09-24T00:00:00Z"),
            )

        builder = ReportContextBuilder()
        html_builder = HtmlReportPackageBuilder()

        def check(expected_p1: int, expected_status: str) -> None:
            context = builder.build(conn, run_id)
            report = ReportDataFactory().from_context(context)
            category = next(item for item in context["vsan_summary"]["report_categories"] if item["category_id"] == "VSAN-DISK")
            assert category["status"] == expected_status
            assert context["risk_summary"]["P1"] == expected_p1
            word_engine = DocxReportEngine()
            visible_groups = word_engine._customer_visible_risk_groups(word_engine._aggregate_risks(report))
            assert word_engine._risk_counts(report, visible_groups)["P1"] == expected_p1
            page = html_builder._single_page_payload(report, context)
            assert sum(1 for item in page["problems"] if item["level"] == "P1") == expected_p1
            assert [item["name"] for item in page["environment"]["clusters"]] == ["vSAN"]
            assert 'href="#environment"' in html_builder._html(page)
            assert 'id="vsan"' not in html_builder._html(page)

        check(0, "正常")
        datastore["vsan_disk_health_issues"] = [{"host": "esxi-01", "disk": "naa.cache", "status": "red", "summary": "磁盘健康异常"}]
        datastore["vsan_physical_disks"][0]["summary_health"] = "red"
        conn.execute("UPDATE inventory_objects SET properties_json = ? WHERE object_id = 'datastore-1'", (json.dumps(datastore),))
        check(1, "需关注")

        datastore["datastore_is_vsan"] = False
        datastore["vsan_api_status"] = None
        host["vsan_enabled"] = False
        conn.execute("UPDATE inventory_objects SET properties_json = ? WHERE object_id = 'datastore-1'", (json.dumps(datastore),))
        conn.execute("UPDATE inventory_objects SET properties_json = ? WHERE object_id = 'host-1'", (json.dumps(host),))
        context = builder.build(conn, run_id)
        report = ReportDataFactory().from_context(context)
        assert context["vsan_summary"]["status"] == "not_applicable"
        assert context["risk_summary"]["P1"] == 0
        page = html_builder._single_page_payload(report, context)
        html = html_builder._html(page)
        assert len(page["environment"]["clusters"]) == 0
        assert 'href="#vsan"' not in html
        assert 'id="vsan"' not in html
        doc = Document(DocxReportEngine().render(report, tmp_path / "no-vsan.docx"))
        headings = [paragraph.text for paragraph in doc.paragraphs if paragraph.style.name.startswith("Heading")]
        assert not any("vSAN 专项巡检" in heading for heading in headings)
        assert "4. 后续处理建议" in headings


def test_vsan_vmkernel_selection_excludes_management_and_vmotion() -> None:
    collector = object.__new__(PyVmomiCollector)
    host = NS(
        name="esxi-01",
        config=NS(
            virtualNicManagerInfo=NS(netConfig=[
                NS(nicType="management", selectedVnic=["vmk0"]),
                NS(nicType="vmotion", selectedVnic=["vmk1"]),
                NS(nicType="vsan", selectedVnic=["vmk2"]),
            ]),
            network=NS(vnic=[
                NS(device="vmk0", spec=NS(portgroup="Management", mtu=1500, ip=NS(ipAddress="10.240.0.1", subnetMask="255.255.255.0"))),
                NS(device="vmk1", spec=NS(portgroup="vMotion", mtu=1500, ip=NS(ipAddress="10.240.1.1", subnetMask="255.255.255.0"))),
                NS(device="vmk2", spec=NS(portgroup="vSAN", mtu=9000, ip=NS(ipAddress="10.240.2.1", subnetMask="255.255.255.0"))),
            ]),
        ),
    )
    result = collector._vsan_vmk_adapters(host)
    assert [item["device"] for item in result] == ["vmk2"]
    assert result[0]["ip_address"] == "10.240.2.1"
    assert result[0]["mtu"] == 9000


def test_vsan_vmkernel_selection_resolves_real_vsan_key_reference() -> None:
    collector = object.__new__(PyVmomiCollector)
    host = NS(
        name="esxi-01",
        config=NS(
            virtualNicManagerInfo=NS(netConfig=[NS(nicType="vsan", selectedVnic=["vsan.key-vim.host.VirtualNic-vmk1"])]),
            network=NS(vnic=[NS(key="key-vim.host.VirtualNic-vmk1", device="vmk1", spec=NS(portgroup="vSAN", mtu=1500, ip=NS(ipAddress="10.250.20.69", subnetMask="255.255.255.0")))]),
        ),
    )
    result = collector._vsan_vmk_adapters(host)
    assert result and result[0]["device"] == "vmk1"
    assert result[0]["ip_address"] == "10.250.20.69"


def test_vsan_disk_mappings_are_grouped_by_host() -> None:
    collector = object.__new__(PyVmomiCollector)
    host = NS(name="esxi-01")
    disk_system = NS(QueryDiskMappings=lambda host=None: [
        NS(diskMapping=NS(uuid="group-1", ssd=NS(canonicalName="naa.cache"), nonSsd=[NS(canonicalName="naa.capacity-1"), NS(canonicalName="naa.capacity-2")]))
    ])
    result = collector._query_vsan_disk_inventory(disk_system, NS(host=[host]))
    assert result["disk_group_count"] == 1
    assert result["cache_disk_count"] == 1
    assert result["capacity_disk_count"] == 2
    assert result["disk_topology"][0]["host"] == "esxi-01"
    assert result["disk_topology"][0]["disk_groups"][0]["capacity_disks"] == ["naa.capacity-1", "naa.capacity-2"]


def test_vsan_disk_mappings_accept_real_disk_map_info_ex_mapping_field() -> None:
    collector = object.__new__(PyVmomiCollector)
    host = NS(name="esxi-01")
    disk_system = NS(QueryDiskMappings=lambda host=None: [
        NS(mapping=NS(ssd=NS(canonicalName="naa.cache"), nonSsd=[NS(canonicalName="naa.capacity-1")]))
    ])
    result = collector._query_vsan_disk_inventory(disk_system, NS(host=[host]))
    assert result["disk_group_count"] == 1
    assert result["cache_disk_count"] == 1
    assert result["capacity_disk_count"] == 1
    assert result["disk_topology"][0]["disk_groups"][0]["cache_disk"] == "naa.cache"
    assert result["disk_topology"][0]["disk_groups"][0]["capacity_disks"] == ["naa.capacity-1"]


def test_vm_kernel_adapters_resolve_dvs_network_label_and_switch() -> None:
    collector = object.__new__(PyVmomiCollector)
    dpg = NS(
        key="dvportgroup-1080",
        name="vSAN-PG",
        config=NS(
            name="vsan",
            distributedVirtualSwitch=NS(uuid="dvs-uuid", name="vSAN"),
        ),
    )
    index = collector._distributed_portgroup_index([dpg])
    host = NS(
        name="esxi-01",
        config=NS(
            network=NS(
                portgroup=[],
                proxySwitch=[NS(dvsUuid="dvs-uuid", dvsName="vSAN")],
                vnic=[NS(
                    device="vmk1",
                    spec=NS(
                        portgroup="",
                        distributedVirtualPort=NS(switchUuid="dvs-uuid", portgroupKey="dvportgroup-1080"),
                        mtu=1500,
                        ip=NS(ipAddress="10.250.20.69", subnetMask="255.255.255.0"),
                    ),
                )],
            ),
            virtualNicManagerInfo=NS(netConfig=[NS(nicType="vsan", selectedVnic=["vmk1"])]),
        ),
    )

    all_adapters = collector._vmkernel_adapters(host, index)
    vsan_adapters = collector._vsan_vmk_adapters(host, index)

    assert all_adapters[0]["network_label"] == "vsan"
    assert all_adapters[0]["switch"] == "vSAN"
    assert vsan_adapters[0]["host_name"] == "esxi-01"
    assert vsan_adapters[0]["ip_address"] == "10.250.20.69"
    assert vsan_adapters[0]["network_label"] == "vsan"
    assert vsan_adapters[0]["switch"] == "vSAN"


def test_vsan_physical_disk_health_preserves_health_capacity_and_membership() -> None:
    collector = object.__new__(PyVmomiCollector)
    summary = NS(
        physicalDisksHealth=[NS(
            hostname="esxi-01",
            disks=[NS(
                name="naa.disk-1",
                uuid="disk-uuid-1",
                summaryHealth="green",
                operationalHealth="green",
                operationalHealthDescription="OK",
                capacityHealth="green",
                capacity=1024,
                usedCapacity=512,
                inCmmds=True,
                inVsi=True,
                vsanDiskGroupUuid="group-1",
                scsiDisk=NS(canonicalName="naa.disk-1"),
            )],
        )],
    )

    disks = collector._vsan_physical_disk_details(summary)

    assert disks == [{
        "host": "esxi-01",
        "name": "naa.disk-1",
        "uuid": "disk-uuid-1",
        "summary_health": "green",
        "operational_health": "green",
        "operational_health_description": "OK",
        "capacity_health": "green",
        "capacity_bytes": 1024,
        "used_capacity_bytes": 512,
        "in_cmmds": True,
        "in_vsi": True,
        "disk_group_uuid": "group-1",
        "scsi_device": "naa.disk-1",
    }]
    assert collector._vsan_physical_disk_details(NS()) is None


def test_vsan_api_version_negotiation_does_not_use_sdk_v3_default() -> None:
    collector = PyVmomiCollector("vc.local", "user", "secret", ssl_verify=False)
    fake = NS(GetLatestVmodlVersion=lambda host, port: "vsan.version.version22")
    assert collector._resolve_vsan_api_version(fake) == "vsan.version.version22"


def test_vsan_zero_object_health_is_not_reported_as_an_issue() -> None:
    collector = object.__new__(PyVmomiCollector)
    summary = NS(results=[
        NS(name="inaccessible", health="inaccessible", numObjects=0),
        NS(name="inaccessible", health="inaccessible", numObjects=2),
    ])
    issues = collector._extract_vsan_health_issues(summary)
    assert len(issues) == 1
    assert issues[0]["status"] == "inaccessible"


def test_vsan_resync_parser_accepts_current_api_field_names() -> None:
    collector = object.__new__(PyVmomiCollector)
    object_system = NS(QuerySyncingVsanObjectsSummary=lambda cluster: NS(totalObjectsToSync=3, totalBytesToSync=4096))
    result = collector._query_vsan_resync_summary(object_system, NS(name="cluster-a"))
    assert result["object_count"] == 3
    assert result["bytes"] == 4096
    assert result["object_count_source"] == "totalObjectsToSync"
    assert result["bytes_source"] == "totalBytesToSync"


def test_vsan_health_state_details_preserve_native_fields_and_object_counts() -> None:
    collector = object.__new__(PyVmomiCollector)
    summary = NS(
        overallHealth="info",
        overallHealthDescription="Online vSAN health issue",
        timestamp="2026-09-20T00:00:00Z",
        objectHealth=NS(objectHealthDetail=[
            NS(health="healthy", numObjects=156, objUuids=["a"]),
            NS(health="inaccessible", numObjects=0, objUuids=[]),
        ]),
    )
    details = collector._vsan_health_state_details(summary)
    assert details == [
        {"health": "healthy", "num_objects": 156, "object_uuids": ["a"]},
        {"health": "inaccessible", "num_objects": 0, "object_uuids": []},
    ]


def test_vsan_object_inventory_preserves_identity_and_vdisk_fields() -> None:
    collector = object.__new__(PyVmomiCollector)
    object_system = NS(
        QueryObjectIdentities=lambda **kwargs: NS(identities=[
            NS(uuid="obj-1", type="vdisk", description="disk.vmdk", vmInstanceUuid="vm-1", spbmProfileName="Policy", spbmProfileUuid="profile-1"),
            NS(uuid="obj-2", type="namespace", description="vm-home", vmInstanceUuid="vm-1"),
        ]),
        QueryVsanObjectInformation=lambda **kwargs: [
            NS(vsanObjectUuid="obj-1", vsanHealth="healthy", spbmProfileUuid="profile-1", spbmComplianceResult=NS(complianceStatus="compliant")),
            NS(vsanObjectUuid="obj-2", vsanHealth="healthy", spbmComplianceResult=NS(complianceStatus="compliant")),
        ],
    )
    result = collector._query_vsan_object_inventory(object_system, NS(name="cluster-a"))
    assert result["collection_status"] == "collected"
    assert result["object_count"] == 2
    assert result["vmdk_count"] == 1
    assert result["vmdks"][0]["object_uuid"] == "obj-1"
    assert result["objects"][0]["health"] == "healthy"
    assert result["storage_policy_summary"]["source"] == "vsan_object_api"
    assert result["storage_policy_summary"]["checked_count"] == 2
    assert result["storage_policy_summary"]["policy_categories"][0]["policy_name"] == "Policy"


def test_vsan_object_inventory_requests_identity_without_space_summary_and_supports_vos_method() -> None:
    collector = object.__new__(PyVmomiCollector)
    calls = []

    def query_identities(**kwargs):
        calls.append(kwargs)
        return NS(identities=[NS(uuid="obj-1", type="vdisk", description="disk.vmdk", spbmProfileName="Policy", spbmProfileUuid="profile-1")])

    object_system = NS(
        QueryObjectIdentities=query_identities,
        VosQueryVsanObjectInformation=lambda **kwargs: [
            NS(vsanObjectUuid="obj-1", vsanHealth="healthy", spbmComplianceResult=NS(complianceStatus="nonCompliant"))
        ],
    )
    result = collector._query_vsan_object_inventory(object_system, NS(name="cluster-a"))
    assert calls[0]["includeObjIdentity"] is True
    assert calls[0]["includeSpaceSummary"] is False
    assert result["policy_noncompliant_count"] == 1
    assert result["storage_policy_summary"]["noncompliant_count"] == 1
    assert result["storage_policy_summary"]["compliance_collection_status"] == "collected"


def test_vsan_object_policy_without_compliance_is_not_reported_as_zero() -> None:
    collector = object.__new__(PyVmomiCollector)
    object_system = NS(
        QueryObjectIdentities=lambda **kwargs: NS(
            identities=[NS(uuid="obj-1", type="vdisk", description="disk.vmdk", spbmProfileName="Policy", spbmProfileUuid="profile-1")]
        )
    )
    result = collector._query_vsan_object_inventory(object_system, NS(name="cluster-a"))
    assert result["policy_noncompliant_count"] is None
    assert result["storage_policy_summary"]["noncompliant_count"] == 0
    assert result["storage_policy_summary"]["unknown_count"] == 1


def test_vmdk_inventory_uses_vsan_object_policy_when_rest_is_unavailable() -> None:
    collector = object.__new__(PyVmomiCollector)
    collector._service_instance = None
    collector._rest_policy_cache = {}
    collector._vsan_policy_by_object_uuid = {
        "object-uuid": {
            "status": "compliant",
            "policy": "profile-1",
            "policy_uuid": "profile-1",
            "policy_name": "Policy",
            "source": "vsan_object_api",
            "scope": "vsan_object",
            "coverage": "object_only",
            "pbm_confirmed": False,
        }
    }
    backing = vim.vm.device.VirtualDisk.FlatVer2BackingInfo(
        fileName="[vSAN] SQL01/SQL01.vmdk",
        uuid="disk-uuid",
        backingObjectId="object-uuid",
        thinProvisioned=True,
        eagerlyScrub=False,
        diskMode="persistent",
    )
    device = vim.vm.device.VirtualDisk(key=2000, unitNumber=0, capacityInKB=2048, backing=backing)
    result = collector._vm_vmdk_inventory(NS(config=NS(hardware=NS(device=[device]))), vim)
    assert result[0]["storage_policy"]["policy_name"] == "Policy"
    assert result[0]["storage_policy"]["source"] == "vsan_object_api"


def test_vmdk_inventory_preserves_backing_and_read_only_policy_shape() -> None:
    collector = object.__new__(PyVmomiCollector)
    collector._service_instance = None
    collector._rest_policy_cache = {}
    backing = vim.vm.device.VirtualDisk.FlatVer2BackingInfo(
        fileName="[vSAN] SQL01/SQL01.vmdk",
        uuid="disk-uuid",
        backingObjectId="object-uuid",
        thinProvisioned=True,
        eagerlyScrub=False,
        diskMode="persistent",
    )
    device = vim.vm.device.VirtualDisk(key=2000, unitNumber=0, capacityInKB=2048, backing=backing)
    result = collector._vm_vmdk_inventory(NS(config=NS(hardware=NS(device=[device]))), vim)
    assert result[0]["capacity_bytes"] == 2048 * 1024
    assert result[0]["file_name"] == "[vSAN] SQL01/SQL01.vmdk"
    assert result[0]["uuid"] == "disk-uuid"
    assert result[0]["backing_object_id"] == "object-uuid"
    assert result[0]["storage_policy"] is None


def test_vsan_native_space_usage_preserves_capacity_and_thresholds() -> None:
    collector = object.__new__(PyVmomiCollector)
    usage_system = NS(
        QuerySpaceUsage=lambda cluster=None, whatifCapacityOnly=False: NS(
            totalCapacityB=1000,
            freeCapacityB=400,
            uncommittedB=50,
            capacityHealthThreshold=NS(yellowValue=750, redValue=800, enabled=True),
            spaceOverview=NS(physicalUsedB=600, usedB=600, overheadB=25, primaryCapacityB=575, reservedCapacityB=10),
            spaceDetail=NS(spaceUsageByObjectType=[NS(objType="vdisk", physicalUsedB=500, usedB=500, overheadB=20, primaryCapacityB=480)]),
        )
    )
    result = collector._query_vsan_space_usage(usage_system, NS(name="cluster-a"))
    assert result["collection_status"] == "collected"
    assert result["coverage"] == "native_vsan"
    assert result["total_capacity_bytes"] == 1000
    assert result["used_capacity_bytes"] == 600
    assert result["capacity_health_threshold"] == {"warning_bytes": 750, "error_bytes": 800, "enabled": True}
    assert result["space_by_object_type"][0]["object_type"] == "vdisk"


def test_word_vsan_heading_levels_and_non_applicable_suppression(tmp_path) -> None:
    context = {
        "run": {"run_id": "run-test", "created_at": "2026-09-19T00:00:00Z"},
        "scope": {"customer_name": "客户", "site_name": "站点", "vcenter_host": "vc.local"},
        "environment_info": {}, "asset_summary": {}, "asset_total": 0, "findings": [], "exception_findings": [],
        "risk_summary": {"P1": 0, "P2": 0, "P3": 0, "P4": 0}, "risk_object_summary": {"P1": 0, "P2": 0, "P3": 0, "P4": 0},
        "result_status_summary": {"passed": 1, "failed": 0, "unavailable": 0}, "health_impact_summary": {"none": 0, "attention": 0, "critical": 0},
        "vsan_summary": {"applicable": True, "status": "collected", "health_status": "healthy", "clusters": ["cluster-a"], "health_issue_count": 0, "object_issue_count": 0, "disk_issue_count": 0, "resync_object_count": 0, "resync_bytes": 0, "disk_topology": []},
        "executive_summary": {"score": 100}, "score_breakdown": {}, "rule_catalog": [], "unavailable_results": [], "not_applicable_summary": [], "certificate_license_evidence": [], "rule_checklist": [], "remediation_plan": [], "object_results": [],
    }
    report = ReportDataFactory().from_context(context)
    path = DocxReportEngine().render(report, tmp_path / "vsan.docx")
    doc = Document(path)
    heading_styles = {paragraph.text: paragraph.style.name for paragraph in doc.paragraphs if paragraph.text.startswith("4.")}
    assert heading_styles["4.1 集群健康"] == "Heading 2"
    assert heading_styles["4.2 磁盘与磁盘组"] == "Heading 2"
    assert not any("Storage Policy" in paragraph.text for paragraph in doc.paragraphs)
    assert any("判定说明" in paragraph.text for paragraph in doc.paragraphs)
    resync_table = next(table for table in doc.tables if table.rows[0].cells[0].text == "Resync 状态")
    resync_text = "\n".join(cell.text for row in resync_table.rows for cell in row.cells)
    assert "当前无待同步任务" in resync_text
    assert "0 个" in resync_text
    assert "0 B" in resync_text
    assert "0 秒" in resync_text
    assert "0 项" in resync_text

    active_context = {
        **context,
        "vsan_summary": {**context["vsan_summary"], "resync_object_count": 2, "resync_bytes": 1024},
    }
    report = ReportDataFactory().from_context(active_context)
    path = DocxReportEngine().render(report, tmp_path / "vsan-resync-active.docx")
    active_table = next(table for table in Document(path).tables if table.rows[0].cells[0].text == "Resync 状态")
    active_text = "\n".join(cell.text for row in active_table.rows for cell in row.cells)
    assert "当前有待同步任务" in active_text
    assert "2 个" in active_text
    assert "1.0 KB" in active_text
    assert active_text.count("未采集") == 2

    context["vsan_summary"] = {"applicable": False, "status": "not_applicable"}
    report = ReportDataFactory().from_context(context)
    path = DocxReportEngine().render(report, tmp_path / "no-vsan.docx")
    headings = [paragraph.text for paragraph in Document(path).paragraphs if paragraph.style.name.startswith("Heading")]
    assert not any("vSAN 专项巡检" in heading for heading in headings)
    assert "4. 后续处理建议" in headings
    assert "5. 环境资产与对象清单" in headings


def test_word_vsan_details_and_judgements_are_rendered(tmp_path) -> None:
    context = {
        "run": {"run_id": "run-vsan-details", "created_at": "2026-09-19T00:00:00Z"},
        "scope": {"customer_name": "客户", "site_name": "站点", "vcenter_host": "vc.local"},
        "environment_info": {}, "asset_summary": {}, "asset_total": 0, "findings": [], "exception_findings": [],
        "risk_summary": {"P1": 0, "P2": 0, "P3": 0, "P4": 0}, "risk_object_summary": {"P1": 0, "P2": 0, "P3": 0, "P4": 0},
        "result_status_summary": {"passed": 1, "failed": 0, "unavailable": 0}, "health_impact_summary": {"none": 0, "attention": 0, "critical": 0},
        "vsan_summary": {
            "applicable": True, "status": "collected", "architecture": "osa", "clusters": ["vSAN"],
            "host_count": 1, "disk_group_count": 1, "cache_disk_count": 1, "capacity_disk_count": 2,
            "object_count": 165, "vmdk_count": 86, "health_issue_count": 0, "disk_issue_count": 0,
            "object_issue_count": 0, "resync_object_count": 0, "resync_bytes": 0,
            "disk_topology": [{"host": "192.168.10.64", "disk_groups": [{"id": "未记录", "cache_disk": "naa.cache", "capacity_disks": ["naa.cap1", "naa.cap2"]}]}],
            "physical_disks": [
                {"host": "192.168.10.64", "name": "naa.cache", "summary_health": "green"},
                {"host": "192.168.10.64", "name": "naa.cap1", "summary_health": "green"},
                {"host": "192.168.10.64", "name": "naa.cap2", "summary_health": "green"},
            ],
            "network": {"status": "collected", "vmkernels": [
                {"host_name": "192.168.10.64", "device": "vmk0", "ip_address": "192.168.10.64", "subnet_mask": "255.255.255.0", "network_label": "Management Network", "mtu": 1500},
                {"host_name": "192.168.10.64", "device": "vmk1", "ip_address": "10.250.20.71", "subnet_mask": "255.255.255.0", "network_label": "vsan", "mtu": 1500},
                {"host_name": "192.168.10.64", "device": "vmk2", "ip_address": "192.168.30.12", "subnet_mask": "255.255.255.0", "network_label": "Storage", "mtu": 1500},
            ]},
            "capacity": {"status": "normal", "used_percent": 10.1, "total_gb": 32192.8, "used_gb": 3247.7, "free_gb": 28945.1},
        },
        "executive_summary": {"score": 100}, "score_breakdown": {}, "rule_catalog": [], "unavailable_results": [], "not_applicable_summary": [], "certificate_license_evidence": [], "rule_checklist": [], "remediation_plan": [], "object_results": [],
    }
    report = ReportDataFactory().from_context(context)
    path = DocxReportEngine().render(report, tmp_path / "vsan-details.docx")
    doc = Document(path)
    headers = [{cell.text for cell in table.rows[0].cells} for table in doc.tables]
    disk_index = next(index for index, header in enumerate(headers) if {"主机", "磁盘组", "磁盘角色", "设备名称", "健康状态"}.issubset(header))
    vmk_index = next(index for index, header in enumerate(headers) if {"主机", "设备", "IP", "子网", "Network Label", "MTU"}.issubset(header))
    assert len(doc.tables[disk_index].rows) == 4
    assert len(doc.tables[vmk_index].rows) == 4
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "未发现对象不可访问或异常健康状态" in text
    assert "低于 75% 关注阈值" in text
    assert "当前没有待重同步对象" in text
    assert "Storage Policy" not in text


def test_snapshot_name_and_create_time_survive_full_report_pipeline(tmp_path) -> None:
    root = Path(__file__).resolve().parents[1]
    raw_rules = RulePackLoader().load_raw(root / "rulepacks" / "builtin-vsphere-v1")
    rule = next(SchemaValidator().validate_rule(item) for item in raw_rules if item.get("rule_id") == "VSL-VM-001")
    collector = object.__new__(PyVmomiCollector)
    vm = NS(snapshot=NS(rootSnapshotList=[NS(name="before-upgrade", createTime=datetime(2026, 5, 1, tzinfo=UTC), childSnapshotList=[])]))
    properties = {"snapshot_age_days_max": collector._snapshot_age_days_max(vm), "snapshots": collector._snapshot_details(vm)}
    inventory = InventoryNormalizer().normalize({"objects": [{"object_type": "VirtualMachine", "object_key": "vm-1", "object_name": "app-01", "object_path": "app-01", "properties": properties}]})
    db_path = tmp_path / "snapshot-pipeline.db"
    init_db(db_path)
    with connect(db_path) as conn:
        customer_id, site_id, vcenter_id = ensure_default_scope(conn, "vc.local")
        run_id = insert_run(conn, customer_id, site_id, vcenter_id, "manual")
        run = RunContext(run_id=run_id, customer_id=customer_id, site_id=site_id, vcenter_id=vcenter_id, db_path=db_path)
        insert_rule(conn, rule)
        results = RuleExecutor().execute(run, inventory, [rule])
        for result in results:
            insert_rule_result(conn, result)
        FindingDeduplicator().upsert_findings(conn, site_id, results)
        finalize_run_summary(conn, run_id)
        context = ReportContextBuilder().build(conn, run_id)
    assert context["findings"][0]["observed_detail"]["snapshots"][0]["name"] == "before-upgrade"
    report = ReportDataFactory().from_context(context)
    html_dir = tmp_path / "html"
    HtmlReportPackageBuilder().render(report, context, html_dir)
    html_payload = (html_dir / "data" / "customer_report_payload.json").read_text(encoding="utf-8")
    assert "before-upgrade" in html_payload
    assert "2026-05-01" in html_payload
    docx_path = DocxReportEngine().render(report, tmp_path / "snapshot.docx")
    doc = Document(docx_path)
    docx_text = "\n".join([paragraph.text for paragraph in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    assert "before-upgrade" in docx_text
    assert "2026-05-01" in docx_text
