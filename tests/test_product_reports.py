from copy import deepcopy
import json

from docx import Document
import pymupdf
import pytest

from vstacklens.products.model import TargetRef, finish_module, make_check, new_module
from vstacklens.products.reporting import (
    NoValidInspectionData, build_integrated_report, deduplicate_observations, product_report_view,
    build_product_evidence, product_evidence_bytes, product_evidence_filename, validate_product_coverage,
)
from vstacklens.reports.exporter import ReportExportEngine
from vstacklens.reports.html_package import HtmlReportPackageBuilder
from vstacklens.reports.pdf_report import PdfReportEngine
from vstacklens.reports.report_model import (
    CustomerInfo, EnvironmentSummary, HealthScore, ReportData, ReportInfo, RiskSummary,
)


def target(product="nsx", target_id="target-one", customer="customer-one"):
    return TargetRef(target_id, product, customer, "site-one", "长名称管理目标_" + "ABCDEFGHIJK_" * 4,
                     "192.0.2.44", customer_name="隔离离线验收客户", site_name="实验站点")


def source(path="/api/v1/cluster/status"):
    return {"method": "GET", "path": path, "collected_at": "2026-10-17T01:00:00Z",
            "version_basis": "4.2.3 实例接口说明"}


def module(t, *, status="normal", count=1):
    value = new_module(t)
    value["version"] = "4.2.3.0.0" if t.product == "nsx" else "30.2.5"
    value["transport_security"] = {"identity_verified": False, "allow_untrusted_lab": True}
    value["limitations"] = ["离线人工样本；不能作为现场验收。"]
    value["checks"] = [make_check(f"check-{i}", f"管理状态核验 {i}", status,
                                 values={"state": "STABLE", "field_missing": None},
                                 sources=[source()], basis="设备原生状态，未设资源阈值。",
                                 object_id=f"node-{i}", object_name=t.name,
                                 verification="live_observed") for i in range(count)]
    for c in value["checks"]:
        c["collected_at"] = "2026-10-17T01:00:00Z"
    return finish_module(value)


def legacy_report():
    return ReportData(report_info=ReportInfo(report_id="old-report", run_id="old-run", generated_at="2026-10-16T01:00:00Z"),
                      customer_info=CustomerInfo(customer_name="原客户", site_name="原站点", vcenter="vc.example"),
                      environment_summary=EnvironmentSummary(scope_statement="原始 VC 范围", checked_object_total=1),
                      health_score=HealthScore(score=77, grade="B", explanation="原始历史判断", label="关注"),
                      risk_summary=RiskSummary(P2=3, total=3),
                      asset_inventory={"total": 1, "details": {"vCenter": [{"object_name": "vc.example", "properties": {"connected": True}}]}})


def test_legacy_defaults_and_original_grade_are_preserved():
    legacy = legacy_report()
    before = legacy.model_dump()
    t = target()
    report = build_integrated_report("new-run", [target("vcenter", "old-vc"), t], [module(t)], "综合巡检", legacy_report=legacy)
    assert legacy.model_dump() == before
    assert report.health_score.model_dump() == legacy.health_score.model_dump()
    assert report.risk_summary.model_dump() == legacy.risk_summary.model_dump()
    assert report.selected_products == ["vcenter", "nsx"]
    assert report.integration_summary["score_scope"]["target_id"] == "old-vc"
    assert legacy.selected_products == ["vcenter"] and legacy.product_modules == []


def test_identical_fact_deduplicates_display_and_preserves_evidence():
    t = target()
    m = module(t)
    m["checks"].append(deepcopy(m["checks"][0]))
    before = deepcopy(m)
    report = build_integrated_report("run", [t], [m], "NSX 巡检")
    view = product_report_view(report)
    assert len(view["sections"][0]["checks"]) == 1
    assert view["sections"][0]["checks"][0]["evidence_count"] == 2
    assert len(report.product_modules[0]["checks"]) == 2
    assert report.integration_summary["verified_pass"] == 0
    assert "不计入" in view["sections"][0]["checks"][0]["limits"]
    assert m == before


def test_conflicting_records_and_different_measurement_windows_survive():
    t = target()
    m = module(t)
    changed = deepcopy(m["checks"][0]); changed["values"] = {"state": "UNSTABLE"}; changed["status"] = "abnormal"
    m["checks"].append(changed)
    rows, conflicts = deduplicate_observations([m])
    assert len(rows) == 2 and len(conflicts) == 1
    assert all("conflict" in row for row in rows)
    changed["measurement_window"] = {"seconds": 60}
    rows, conflicts = deduplicate_observations([m])
    assert len(rows) == 2 and not conflicts


def test_same_address_cross_customer_and_untrusted_identity_never_merge():
    a, b = target(), target(target_id="target-two", customer="customer-two")
    ma, mb = module(a), module(b)
    ma["target_identity"] = mb["target_identity"] = {"canonical_object_id": "same-name-not-trusted"}
    report = build_integrated_report("run", [a, b], [ma, mb], "独立范围")
    assert report.integration_summary["observation_count"] == 2
    assert len(product_report_view(report)["sections"]) == 2


@pytest.mark.parametrize("status", ["permission_denied", "unsupported", "collection_failed", "insufficient_data", "cancelled"])
def test_all_modules_without_effective_data_cannot_publish(status):
    t = target()
    with pytest.raises(NoValidInspectionData):
        build_integrated_report("run", [t], [module(t, status=status)], "无数据")


def test_native_medium_remains_native_and_failures_do_not_hide_other_target():
    a, b = target(), target("avi", "avi-one")
    ma, mb = module(a), module(b, status="permission_denied")
    ma["native_alerts"] = [{"severity": "MEDIUM", "object_name": "Manager", "status": "OPEN", "description": "原生告警"}]
    report = build_integrated_report("run", [a, b], [ma, mb], "组合巡检")
    view = product_report_view(report)
    assert view["sections"][0]["native_alerts"][0]["severity"] == "MEDIUM"
    assert view["sections"][1]["checks"][0]["status"] == "权限不足"
    assert report.risk_summary.total == report.risk_summary.P2 == 0


@pytest.mark.parametrize("product", ["nsx", "avi"])
def test_single_product_real_three_engines_same_values_no_empty_vcenter(tmp_path, product):
    t = target(product)
    m = module(t, count=2)
    m["checks"][1]["status"] = "attention"
    m["checks"][1]["summary"] = "备份未配置，需要人工确认保护范围。"
    m["native_alerts"] = [{"severity": "MEDIUM", "object_name": "管理节点", "status": "OPEN", "description": "离线原生告警样本"}]
    report = build_integrated_report("offline-run", [t], [m], f"{product.upper()} 只读巡检")
    assert report.health_score.score is None and report.health_score.grade == "未评分"
    before = report.model_dump()
    html, _ = HtmlReportPackageBuilder().render(report, {}, tmp_path / "report")
    word, _, _ = ReportExportEngine().render(report, tmp_path / "report.docx")
    pdf = PdfReportEngine().render(report, tmp_path / "report.pdf")
    payload = json.loads((html.parent / "data/customer_report_payload.json").read_text(encoding="utf-8"))
    assert payload["productReport"] == product_report_view(report)
    html_text = html.read_text(encoding="utf-8")
    assert 'id="environment"' not in html_text and 'id="inventory"' not in html_text
    word_doc = Document(word)
    word_text = "\n".join([p.text for p in word_doc.paragraphs] + [c.text for table in word_doc.tables for row in table.rows for c in row.cells])
    with pymupdf.open(pdf) as document:
        pdf_text = "\n".join(page.get_text() for page in document)
    for text in (word_text, pdf_text):
        assert "STABLE" in text and "MEDIUM" in text and "备份未配置" in text
        assert "未验证对端身份" in text and "现场读取" in text
        assert "VMkernel" not in text and "vCenter 数量" not in text
    assert report.model_dump() == before
    assert all(table._tbl.xpath(".//w:tblHeader") for table in word_doc.tables)
    view = product_report_view(report)
    assert view["coverage"]["classified_product_checks"] == 2
    assert view["evidence_filename"] in word_text and view["evidence_filename"] in pdf_text
    expected = product_evidence_bytes(report)
    assert (html.parent / view["evidence_filename"]).read_bytes() == expected
    assert (pdf.parent / view["evidence_filename"]).read_bytes() == expected


def test_product_projection_rejects_unsafe_input_without_silently_dropping_evidence():
    t = target()
    m = module(t)
    m["checks"][0]["values"].update(password="sensitive-hidden", license_key="hidden-key", certificate_pem="full-certificate")
    before = deepcopy(m)
    with pytest.raises(ValueError) as error:
        build_integrated_report("run", [t], [m], "脱敏")
    assert all(secret not in str(error.value) for secret in ("sensitive-hidden", "hidden-key", "full-certificate"))
    assert m == before


def test_direct_engine_rejects_manually_constructed_empty_product_report(tmp_path):
    report = legacy_report()
    report.selected_products = ["avi"]
    report.product_modules = [{"product": "avi", "checks": [], "target_id": "t"}]
    for generate in (lambda: HtmlReportPackageBuilder().render(report, {}, tmp_path / "html"),
                     lambda: ReportExportEngine().render(report, tmp_path / "word.docx"),
                     lambda: PdfReportEngine().render(report, tmp_path / "pdf.pdf")):
        with pytest.raises(NoValidInspectionData):
            generate()
    assert not list(tmp_path.glob("*.docx")) and not list(tmp_path.glob("*.pdf"))


def test_multiple_vcenters_keep_individual_original_models_in_three_engines(tmp_path):
    a, b, nsx = target("vcenter", "vc-one"), target("vcenter", "vc-two"), target()
    first, second = legacy_report(), legacy_report()
    first.customer_info.vcenter = "vc-first.example"
    second.customer_info.vcenter = "vc-second.example"
    second.report_info.run_id = "second-old-run"
    second.health_score.score = 41
    second.health_score.grade = "D"
    first.integration_summary = {"legacy_target_id": a.target_id}
    second.integration_summary = {"legacy_target_id": b.target_id}
    before = deepcopy([first.model_dump(), second.model_dump()])
    report = build_integrated_report("combined-run", [a, b, nsx], [module(nsx)], "多目标综合报告", legacy_report=[first, second])
    saved = report.integration_summary["legacy_assessments"]
    assert report.health_score.score is None and report.health_score.grade == "分别评估"
    assert report.integration_summary["score_scope"]["kind"] == "per_vcenter"
    assert [(r["target_id"], r["score"], r["grade"]) for r in saved] == [("vc-one", 77, "B"), ("vc-two", 41, "D")]
    html, _ = HtmlReportPackageBuilder().render(report, {}, tmp_path / "html")
    word, _, _ = ReportExportEngine().render(report, tmp_path / "multi.docx")
    pdf = PdfReportEngine().render(report, tmp_path / "multi.pdf")
    payload = json.loads((html.parent / "data/customer_report_payload.json").read_text(encoding="utf-8"))
    assert [r["target_id"] for r in payload["legacyReports"]] == ["vc-one", "vc-two"]
    word_text = "\n".join(p.text for p in Document(word).paragraphs)
    with pymupdf.open(pdf) as document:
        pdf_text = "\n".join(p.get_text() for p in document)
    assert all(marker in text for marker in ("vc-one", "vc-two", "NSX") for text in (word_text, pdf_text))
    assert [first.model_dump(), second.model_dump()] == before


def test_long_inventory_retains_every_object_and_sensitive_publish_fails_closed(tmp_path):
    t = target()
    m = module(t)
    names = [f"long-manager-object-{i:03d}-" + "A_long_continuous_name_" * 4 for i in range(24)]
    m["inventory"] = [{"name": name, "state": "UP", "address": None} for name in names]
    report = build_integrated_report("long-run", [t], [m], "长表巡检")
    html, _ = HtmlReportPackageBuilder().render(report, {}, tmp_path / "html")
    word, _, _ = ReportExportEngine().render(report, tmp_path / "long.docx")
    pdf = PdfReportEngine().render(report, tmp_path / "long.pdf")
    word_doc = Document(word)
    word_text = "\n".join(c.text for table in word_doc.tables for row in table.rows for c in row.cells)
    with pymupdf.open(pdf) as document:
        pdf_text = "\n".join(p.get_text() for p in document)
    html_data = (html.parent / "data/customer_report_payload.json").read_text(encoding="utf-8")
    # Assert the full name (not just its ordinal prefix) in all formats. PDF
    # line wrapping may add whitespace/soft breaks inside a continuous token;
    # every original non-layout character must still be present for all 24.
    visible = ["".join(text.replace("\u200b", "").split()) for text in (word_text, pdf_text, html_data)]
    for name in names:
        assert all(name in text for text in visible)
    from vstacklens.application.artifact_safety import UnsafeArtifactError
    secret = "deliberately-sensitive-runtime-value"
    report.product_modules[0]["checks"][0]["summary"] = secret
    with pytest.raises(UnsafeArtifactError):
        HtmlReportPackageBuilder().render(report, {}, tmp_path / "blocked", secrets=(secret,))
    assert not (tmp_path / "blocked").exists()
    with pytest.raises(UnsafeArtifactError):
        ReportExportEngine().render(report, tmp_path / "blocked.docx", secrets=(secret,))
    with pytest.raises(UnsafeArtifactError):
        PdfReportEngine().render(report, tmp_path / "blocked.pdf", secrets=(secret,))
    assert not (tmp_path / "blocked.docx").exists() and not (tmp_path / "blocked.pdf").exists()


def test_successful_vcenter_matches_explicit_id_after_another_vcenter_failed():
    failed, successful = target("vcenter", "failed-vc"), target("vcenter", "success-vc")
    original = legacy_report()
    original.integration_summary = {"legacy_target_id": successful.target_id}
    report = build_integrated_report("run", [failed, successful], [], "VC 部分失败", legacy_report=[original])
    assert report.integration_summary["legacy_assessments"][0]["target_id"] == "success-vc"


def test_long_check_evidence_splits_bounded_cells_and_keeps_every_service(tmp_path):
    t = target()
    m = module(t)
    m["checks"][0]["values"] = {"services": [{"name": f"manager-service-{i:03d}-" + "LongName" * 12, "state": "UP"} for i in range(24)]}
    report = build_integrated_report("long-value-run", [t], [m], "完整服务证据")
    rows = product_report_view(report)["sections"][0]["checks"][0]["details"]
    assert any(row["label"] == "实际状态或值（续）" for row in rows)
    assert all(len(row["text"]) <= 850 for row in rows)
    pdf = PdfReportEngine().render(report, tmp_path / "long-value.pdf")
    with pymupdf.open(pdf) as document:
        text = "\n".join(p.get_text() for p in document)
    assert all(f"manager-service-{i:03d}" in text for i in range(24))


def test_failed_vcenter_with_effective_nsx_reports_failed_scope_in_all_three_formats(tmp_path):
    vc, nsx = target("vcenter", "failed-vc"), target()
    failed = module(vc, status="collection_failed")
    failed["checks"][0]["title"] = "vCenter 连接失败诊断"
    failed["checks"][0]["summary"] = "未取得 vCenter 清单，保留具体失败范围。"
    failed["errors"] = [{"stage": "collection", "type": "authentication", "reason": "账号认证失败"}]
    diagnostic = legacy_report()
    diagnostic.asset_inventory["details"]["vCenter"][0]["properties"]["connected"] = False
    report = build_integrated_report("partial-run", [vc, nsx], [failed, module(nsx)], "部分目标失败巡检", legacy_report=diagnostic)
    view = product_report_view(report)
    assert not view["has_vcenter"] and not report.integration_summary["legacy_assessments"]
    assert [m["product"] for m in view["sections"]] == ["vcenter", "nsx"]
    assert report.integration_summary["status_counts"]["collection_failed"] == 1
    html, _ = HtmlReportPackageBuilder().render(report, {}, tmp_path / "html")
    word, _, _ = ReportExportEngine().render(report, tmp_path / "partial.docx")
    pdf = PdfReportEngine().render(report, tmp_path / "partial.pdf")
    html_text = html.read_text(encoding="utf-8")
    assert 'id="environment"' not in html_text
    doc = Document(word)
    word_text = "\n".join([p.text for p in doc.paragraphs] + [c.text for table in doc.tables for row in table.rows for c in row.cells])
    with pymupdf.open(pdf) as document:
        pdf_text = "\n".join(p.get_text() for p in document)
    for text in (word_text, pdf_text):
        normalized = "".join(text.replace("\u200b", "").split())
        assert "vCenter连接失败诊断" in normalized and "采集失败" in normalized and "STABLE" in normalized
        assert "VMkernel" not in text and "vCenter 数量" not in text


def test_all_failed_vcenter_diagnostic_inventory_is_not_effective_inspection_data():
    vc = target("vcenter", "failed-vc")
    diagnostic = legacy_report()
    diagnostic.asset_inventory["details"]["vCenter"][0]["properties"]["connected"] = False
    with pytest.raises(NoValidInspectionData):
        build_integrated_report("failed-run", [vc], [module(vc, status="collection_failed")], "失败诊断", legacy_report=diagnostic)


def test_avi_native_alert_reason_is_preserved_when_summary_is_absent():
    avi = target("avi", "avi-native-reason")
    data = module(avi)
    data["native_alerts"] = [{"obj_name": "Controller", "level": "ALERT_MEDIUM", "state": "ACTIVE", "reason": "原生 Controller 告警原因", "recommended_action": "原生设备建议"}]
    report = build_integrated_report("native-reason", [avi], [data], "AVI 原生告警")
    alert = product_report_view(report)["sections"][0]["native_alerts"][0]
    assert alert["description"] == "原生 Controller 告警原因"
    assert alert["recommendation"] == "原生设备建议"
    assert alert["severity"] == "ALERT_MEDIUM"


def test_native_alert_sources_only_use_the_owning_alarm_check():
    avi = target("avi", "avi-native-source")
    data = module(avi)
    data["checks"] = [
        make_check("AVI-A4", "原生告警", "attention", judgement="NATIVE", sources=[source("/api/alert")]),
        make_check("AVI-A6", "授权", "normal", judgement="NATIVE", sources=[source("/api/licensing/status")]),
        make_check("AVI-C1:cloud", "Cloud", "normal", judgement="NATIVE", sources=[source("/api/cloud")]),
    ]
    data["native_alerts"] = [{"level": "MEDIUM", "reason": "原始原因", "timestamp": 123456789}]
    alert = product_report_view(build_integrated_report("source", [avi], [data], "来源"))["sections"][0]["native_alerts"][0]
    assert "/api/alert" in alert["time"]
    assert "/api/cloud" not in alert["time"] and "/api/licensing" not in alert["time"]
    assert "设备原始发生时间：123456789" in alert["time"]


def test_combined_opening_and_navigation_cover_all_selected_targets_in_three_formats(tmp_path):
    vc, nsx, avi = target("vcenter", "overview-vc"), target("nsx", "overview-nsx"), target("avi", "overview-avi")
    vc.name, nsx.name, avi.name = "VC-unique-target", "NSX-unique-target", "AVI-unique-target"
    legacy = legacy_report()
    report = build_integrated_report("scope", [vc, nsx, avi], [module(vc), module(nsx), module(avi)], "三产品巡检", legacy_report=legacy)
    view = product_report_view(report)
    assert [t["target_id"] for t in view["targets"]] == [vc.target_id, nsx.target_id, avi.target_id]
    html, _ = HtmlReportPackageBuilder().render(report, {}, tmp_path / "html")
    word, _, _ = ReportExportEngine().render(report, tmp_path / "scope.docx")
    pdf = PdfReportEngine().render(report, tmp_path / "scope.pdf")
    doc = Document(word)
    texts = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    toc = next(p.text for p in doc.paragraphs if "NSX-unique-target" in p.text and "AVI-unique-target" in p.text)
    assert "VC-unique-target" in toc
    with pymupdf.open(pdf) as document:
        opening = document[0].get_text() + document[1].get_text()
    for marker in (vc.name, nsx.name, avi.name):
        assert marker in texts and marker in opening
    payload = json.loads((html.parent / "data/customer_report_payload.json").read_text(encoding="utf-8"))
    assert payload["productReport"]["targets"] == view["targets"]
    assert report.health_score.score == legacy.health_score.score


def test_long_iso_display_wrap_preserves_every_source_character_and_model(tmp_path):
    report = legacy_report()
    original_path = "[datastore-prod] ISO/Windows_Server_" + "ExtremelyLongName" * 14 + ".iso"
    report.asset_inventory["details"]["VirtualMachine"] = [{"object_name": "long-iso-vm", "properties": {"iso_paths": [original_path], "iso_mounted": True}}]
    before = report.model_dump()
    engine = PdfReportEngine()
    payload = engine._build_payload(report, {})
    projected = payload["appendix"]["iso_vms"][0]["iso_paths"]
    assert "\u200b" in projected and projected.replace("\u200b", "") == original_path
    pdf = engine.render(report, tmp_path / "iso.pdf")
    with pymupdf.open(pdf) as document:
        joined = "".join(p.get_text() for p in document).replace("\u200b", "")
    # The concise PDF refers readers to full Word/HTML object lists.
    from vstacklens.reports.docx_report import DocxReportEngine
    from vstacklens.reports.html_package import HtmlReportPackageBuilder
    from vstacklens.reports.document_model import get_document
    word=DocxReportEngine().render(report,tmp_path/'iso.docx')
    text=''.join(Document(word).element.xpath('//w:t/text()')).replace('\u200b','')
    assert ''.join(original_path.split()) in ''.join(text.split())
    html,_=HtmlReportPackageBuilder().render(report,{},tmp_path/'full')
    doc=get_document(report)
    assert any(original_path in '\n'.join(r['cells']) for a in doc['appendices'] for r in a['rows'])
    assert html.is_file()
    assert report.asset_inventory == before["asset_inventory"]


def test_native_alert_pdf_time_wrap_preserves_exact_time_and_public_model(tmp_path):
    avi = target("avi", "time-wrap")
    data = module(avi)
    alarm_source = source("/api/alert")
    alarm_source["collected_at"] = "2026-10-17T12:29:29.483447+00:00"
    data["checks"] = [make_check("AVI-A4", "原生告警", "attention", sources=[alarm_source], judgement="NATIVE")]
    data["native_alerts"] = [{"level": "MEDIUM", "state": "ACTIVE", "reason": "原生原因", "timestamp": 123456789}]
    report = build_integrated_report("time", [avi], [data], "时间换行")
    original = product_report_view(report)
    engine = PdfReportEngine()
    projected = engine._build_payload(report, {})["product_report"]
    original_time = original["sections"][0]["native_alerts"][0]["time"]
    displayed_time = projected["sections"][0]["native_alerts"][0]["time"]
    assert "\u200b" in displayed_time and displayed_time.replace("\u200b", "") == original_time
    assert product_report_view(report) == original
    pdf = engine.render(report, tmp_path / "time.pdf")
    with pymupdf.open(pdf) as document:
        text = "".join(p.get_text() for p in document).replace("\u200b", "")
    filename = product_evidence_filename(report)
    evidence = json.loads((tmp_path / filename).read_text(encoding="utf-8"))
    assert evidence["modules"] == report.product_modules
    assert evidence["modules"][0]["checks"][0]["sources"][0]["collected_at"] == "2026-10-17T12:29:29.483447+00:00"
    assert filename in "".join(text.split())


def test_engineering_evidence_is_complete_stable_and_references_fail_closed():
    t = target()
    data = module(t, count=2)
    data["checks"][1]["values"]["long_original_value"] = "完整原始字段" * 1200
    data["inventory"] = [{"object_id": "uuid-kept-in-evidence", "name": "Manager", "values": {"native": {"anything": [1, 2, None]}}}]
    report = build_integrated_report("evidence", [t], [data], "完整工程证据")
    before = deepcopy(report.product_modules)
    evidence = build_product_evidence(report)
    view = product_report_view(report)
    assert evidence["modules"] == before
    assert len(evidence["modules"][0]["checks"][1]["values"]["long_original_value"]) == 7200
    assert product_evidence_bytes(report) == product_evidence_bytes(report)
    assert product_evidence_filename(report) == view["evidence_filename"]
    assert validate_product_coverage(report, view, evidence)["original_modules_equal"]
    changed = deepcopy(evidence); changed["modules"][0]["checks"].pop()
    with pytest.raises(ValueError, match="完整模块不一致"):
        validate_product_coverage(report, view, changed)
    missing = deepcopy(view); missing["sections"][0]["body"]["check_evidence_ids"].pop()
    with pytest.raises(ValueError, match="没有全部归入"):
        validate_product_coverage(report, missing, evidence)
    invalid = deepcopy(view); invalid["sections"][0]["body"]["evidence_ids"].append("E-NOT-REAL")
    with pytest.raises(ValueError, match="不存在"):
        validate_product_coverage(report, invalid, evidence)
    report.product_modules[0]["checks"][0]["values"]["new_fact"] = 1
    assert product_evidence_filename(report) != view["evidence_filename"]
    assert before == evidence["modules"]


def test_aggregate_abnormal_services_and_backends_cannot_be_hidden_in_attachment(tmp_path):
    t = target("avi")
    data = module(t)
    data["checks"] = [make_check("aggregate", "聚合业务状态", "attention", values={
        "services": [{"name": "stopped-service", "state": "STOPPED"}, {"name": "role-disabled-service", "state": "STOPPED", "enabled": False},
                     {"name": "unhealthy-service", "health": "UNHEALTHY"}, {"name": "failed-oper-service", "oper_status": {"state": "OPER_FAILED"}}],
        "nodes": [{"name": "unreachable-controller", "state": "CLUSTER_UNREACHABLE"}],
        "interfaces": [{"interface_id": "eth-test", "link_status": "DOWN"}],
        "pools": [{"name": "pool-one", "servers": [{"name": "backend-a", "state": "DOWN"},
                                                       {"server": {"ip": {"addr": "192.0.2.91"}, "port": 443}, "is_server_up": False}]}],
        "monitor_config": {"failed_checks": 3, "successful_checks": 2, "origin_status": "unknown", "reference_state": "unknown_partial_scope"},
    }, sources=[source("/api/pool")], object_id="")]
    report = build_integrated_report("abnormal-children", [t], [data], "异常子对象")
    view = product_report_view(report)
    table = next(x for x in view["sections"][0]["body"]["tables"] if x["title"] == "需复核状态的具体对象")
    assert len(table["rows"]) == 8
    assert {r[0] for r in table["rows"]} == {"stopped-service", "role-disabled-service", "unhealthy-service", "failed-oper-service", "unreachable-controller", "eth-test", "backend-a", "192.0.2.91:443"}
    assert any("配置已停用" in row[2] for row in table["rows"] if row[0] == "role-disabled-service")
    missing = deepcopy(view)
    next(x for x in missing["sections"][0]["body"]["tables"] if x["title"] == table["title"])["rows"].pop()
    with pytest.raises(ValueError, match="聚合记录"):
        validate_product_coverage(report, missing, build_product_evidence(report))
    assert data["checks"][0]["object_id"] == ""
    html, _ = HtmlReportPackageBuilder().render(report, {}, tmp_path / "html")
    word, _, _ = ReportExportEngine().render(report, tmp_path / "abnormal.docx")
    pdf = PdfReportEngine().render(report, tmp_path / "abnormal.pdf")
    doc = Document(word)
    word_text = "\n".join(c.text for table in doc.tables for row in table.rows for c in row.cells)
    with pymupdf.open(pdf) as document:
        pdf_text = "\n".join(page.get_text() for page in document)
    html_body = json.loads((html.parent / "data/customer_report_payload.json").read_text(encoding="utf-8"))["productReport"]["sections"][0]["body"]
    for marker in ("stopped-service", "role-disabled-service", "unhealthy-service", "failed-oper-service", "unreachable-controller", "eth-test", "backend-a", "192.0.2.91:443"):
        assert all(marker in "".join(text.replace("\u200b", "").split()) for text in (word_text, pdf_text, json.dumps(html_body)))


def test_unreferenced_configuration_is_scoped_catalog_not_system_name_guess():
    t = target("avi")
    data = module(t)
    data["checks"] = [make_check("AVI-C5", "Pool", "not_deployed", values={"count": 0}, sources=[source("/api/pool")], completeness="complete")]
    for i in range(11):
        data["checks"].append(make_check(f"AVI-C7:hm{i}", f"健康检查 System-{i}", "insufficient_data", object_id=f"hm{i}",
                                         values={"name": f"System-{i}", "type": "HEALTH_MONITOR_TCP", "monitor_port": 80, "origin_status": "unknown"}, sources=[source("/api/healthmonitor")]))
    report = build_integrated_report("catalog", [t], [data], "配置目录")
    body = product_report_view(report)["sections"][0]["body"]
    assert len(body["configuration_catalogs"]) == 1 and body["configuration_catalogs"][0][1] == "11"
    assert "当前完整所选范围未发现引用" in body["configuration_catalogs"][0][2]
    assert not body["concerns"] and len(body["empty_business"]) == 1
    assert not any(table["title"] == "需复核状态的具体对象" for table in body["tables"])
    assert all(check["status"] == "insufficient_data" for check in report.product_modules[0]["checks"][1:])
    data["checks"][0]["completeness"] = "partial"
    partial = build_integrated_report("partial-catalog", [t], [data], "配置目录")
    scope = product_report_view(partial)["sections"][0]["body"]["configuration_catalogs"][0][2]
    assert "引用范围未完整核实" in scope and "系统默认属性未确认" in scope


def test_certificate_identity_merges_alias_uses_but_preserves_distinct_certificates():
    t = target("avi")
    data = module(t)
    cert = {"uuid": "cert-a", "name": "Portal cert", "identity": {"certificate_sha256": "a" * 64},
            "x509_validity": {"valid_from": "2026-01-01Z", "expires_at": "2027-01-01Z"},
            "usages": [{"usage": "management_portal"}]}
    data["checks"] = [make_check("AVI-A5:cert-a", "证书", "normal", object_id="cert-a", values=cert, sources=[source("/api/sslkeyandcertificate")])]
    alias = deepcopy(cert); alias.update(uuid="cert-alias", name="Alias cert", usages=[{"usage": "management_secure_channel"}])
    distinct = deepcopy(cert); distinct.update(uuid="cert-b", name="Other cert", identity={"certificate_sha256": "b" * 64})
    data["inventory"] = [{"object_type": "certificate", "object_id": "cert-alias", "values": alias},
                         {"object_type": "certificate", "object_id": "cert-b", "values": distinct}]
    report = build_integrated_report("cert", [t], [data], "证书用途")
    rows = product_report_view(report)["sections"][0]["body"]["certificate_rows"]
    assert len(rows) == 2
    merged = next(row for row in rows if "Alias cert" in row[0])
    assert "Portal cert" in merged[0] and "管理门户" in merged[1] and "管理安全通道" in merged[1]
    assert "2027-01-01Z" in merged[2]
    assert report.product_modules[0]["checks"] == data["checks"]


def test_empty_avi_resource_window_does_not_repeat_old_sample_summary():
    t = target("avi")
    data = module(t)
    data["checks"] = [make_check("AVI-A3", "设备内部资源", "insufficient_data", values={"metrics": [], "partitions": None},
                                 sources=[source("/api/analytics/metrics/controller/sample")], summary="展示已有 Controller 资源样本")]
    # Another real read makes a partial report publishable; empty metrics alone
    # cannot establish a healthy resource reading.
    data["checks"].append(make_check("AVI-A1", "版本", "normal", values={"version": "30.2.5"}, sources=[source("/api/initial-data")]))
    report = build_integrated_report("empty-window", [t], [data], "资源窗口")
    text = json.dumps(product_report_view(report)["sections"][0]["body"], ensure_ascii=False)
    assert "所请求窗口未返回资源样本" in text and "原因未验证" in text
    assert "内部磁盘分区数据未采集" in text and "展示已有 Controller 资源样本" not in text
    assert build_product_evidence(report)["modules"][0]["checks"][0]["summary"] == "展示已有 Controller 资源样本"


def test_tls_catalog_names_configured_usage_and_separates_unreferenced_scope():
    t = target("avi")
    data = module(t)
    data["inventory"] = [
        {"object_type": "ssl_profile", "object_name": "System-Portal", "values": {"name": "System-Portal", "reference_state": "referenced", "usages": [{"usage": "management_portal"}]}},
        {"object_type": "ssl_profile", "object_name": "System-Unused", "values": {"name": "System-Unused", "reference_state": "unreferenced_in_complete_selected_scope", "usages": []}},
    ]
    report = build_integrated_report("tls-scope", [t], [data], "配置用途")
    body = product_report_view(report)["sections"][0]["body"]
    labels = [row[2] for row in body["configuration_catalogs"]]
    assert any("已被当前所选配置引用" in label for label in labels)
    assert any("当前完整所选范围未发现引用" in label for label in labels)
    rows = [row for table in body["tables"] for row in table["rows"]]
    assert any(row[0] == "System-Portal" and "管理门户（配置引用）" in row[2] for row in rows)
    assert not any(row[0] == "System-Unused" for row in rows)


def test_original_run_finish_and_later_source_times_are_separate_in_three_formats(tmp_path):
    t = target("avi")
    data = module(t)
    data["completed_at"] = "2026-10-17T01:00:00Z"
    data["checks"][0]["sources"].append({**source(), "collected_at": "2026-10-17T01:00:00.999999+00:00"})
    data.setdefault("coverage", {})["targeted_overlay"] = {"base_completed_at": "2026-10-17T01:00:00Z", "full_recollection": False,
                                           "supplement_sources": [{"collected_at": "2026-10-17T02:00:00Z"}]}
    data["checks"][0]["values"]["targeted_metric_evidence"] = [{"source": {"collected_at": "2026-10-17T02:00:03Z"}}]
    report = build_integrated_report("source-times", [t], [data], "已有数据重放与定点补读")
    view = product_report_view(report)
    assert view["sections"][0]["time"] == "原始整体采集完成 2026-10-17T01:00:00Z；补充来源取数 2026-10-17T02:00:00Z 至 2026-10-17T02:00:03Z"
    html, _ = HtmlReportPackageBuilder().render(report, {}, tmp_path / "html")
    word, _, _ = ReportExportEngine().render(report, tmp_path / "times.docx")
    pdf = PdfReportEngine().render(report, tmp_path / "times.pdf")
    word_text = "\n".join(p.text for p in Document(word).paragraphs)
    with pymupdf.open(pdf) as document:
        pdf_text = "\n".join(page.get_text() for page in document)
    for text in (word_text, pdf_text):
        normalized = "".join(text.replace("\u200b", "").split())
        assert "原始整体采集完成2026-10-17T01:00:00Z" in normalized
        assert "补充来源取数2026-10-17T02:00:00Z至2026-10-17T02:00:03Z" in normalized
        assert "采集覆盖时间" not in text
    html_view = json.loads((html.parent / "data/customer_report_payload.json").read_text(encoding="utf-8"))["productReport"]
    assert html_view == view
    assert build_product_evidence(report)["modules"][0] == data
    ordinary = deepcopy(data)
    ordinary["coverage"].pop("targeted_overlay")
    ordinary["checks"][0]["values"].pop("targeted_metric_evidence")
    plain = build_integrated_report("ordinary-time", [t], [ordinary], "没有补读标记")
    assert product_report_view(plain)["sections"][0]["time"] == "原始整体采集完成 2026-10-17T01:00:00Z"


def test_cluster_body_uses_native_state_without_old_service_summary_contradiction():
    t = target("avi")
    data = module(t)
    data["checks"] = [make_check("AVI-A2", "集群运行状态", "normal", values={"cluster": {"state": "CLUSTER_UP_NO_HA"}, "nodes_count": 1},
                                 summary="关键服务没有取得已核实的逐服务数据", sources=[source("/api/cluster/runtime")]),
                      make_check("AVI-A2:services", "关键管理服务", "insufficient_data", values={"services": [{"name": "service", "state": "CLUSTER_ACTIVE"}], "reported_count": 1,
                                 "native_state_counts": {"CLUSTER_ACTIVE": 1}}, sources=[source("/api/cluster/runtime")])]
    report = build_integrated_report("cluster-services", [t], [data], "聚合与服务分别展示")
    body_text = json.dumps(product_report_view(report)["sections"][0]["body"], ensure_ascii=False)
    assert "CLUSTER_UP_NO_HA" in body_text and "原生节点数：1" in body_text
    assert "关键服务没有取得已核实的逐服务数据" not in body_text
    assert build_product_evidence(report)["modules"][0]["checks"][0]["summary"] == data["checks"][0]["summary"]


def test_same_source_object_lifecycle_fields_share_one_row_but_distinct_nodes_do_not_merge():
    t = target()
    data = module(t)
    services = [{"node_id": "node-a", "service": "message-bus", "runtime_state": "stopped", "monitor_runtime_state": "stopped"},
                {"node_id": "node-b", "service": "message-bus", "runtime_state": "stopped", "monitor_runtime_state": "stopped"}]
    data["checks"][0]["values"] = {"services": services}
    report = build_integrated_report("lifecycle-facts", [t], [data], "同一对象不同生命周期事实")
    body = product_report_view(report)["sections"][0]["body"]
    rows = next(table["rows"] for table in body["tables"] if table["title"] == "需复核状态的具体对象")
    assert len(rows) == 2  # two exact source objects, never merged by same service name
    assert all("服务生命周期：stopped" in row[1] and "监控生命周期：stopped" in row[1] for row in rows)
    assert build_product_evidence(report)["modules"][0]["checks"][0]["values"]["services"] == services


def test_integrated_word_vcenter_chapter_starts_with_executive_text_and_pure_vcenter_keeps_break(tmp_path):
    vc, nsx = target("vcenter", "vc-word-chapter"), target()
    legacy = legacy_report()
    report = build_integrated_report("word-chapter", [vc, nsx], [module(nsx)], "综合章分页", legacy_report=legacy)
    # Stored legacy reports already carry their canonical customer caliber;
    # prepare this minimal fixture before testing a presentation-only render.
    from vstacklens.reports.customer_caliber import ensure_report
    ensure_report(report)
    before = report.model_dump()
    word, _, _ = ReportExportEngine().render(report, tmp_path / "integrated.docx")
    paragraphs = Document(word).paragraphs
    chapter = next(p for p in paragraphs if p.text.startswith("VC／vSAN 专项"))
    summary = next(p for p in paragraphs if p.text == "1. 执行摘要")
    assert chapter.paragraph_format.page_break_before is True
    assert chapter.paragraph_format.keep_with_next is True
    assert summary.paragraph_format.page_break_before is False
    assert report.model_dump() == before
    pure = legacy_report()
    path, _, _ = ReportExportEngine().render(pure, tmp_path / "pure-vcenter.docx")
    pure_summary = next(p for p in Document(path).paragraphs if p.text == "目录")
    assert pure_summary.paragraph_format.page_break_before is True


def test_vcenter_cover_summary_heading_stays_with_summary_text(tmp_path):
    vc, nsx, avi = target("vcenter", "cover-vc"), target(), target("avi", "cover-avi")
    report = build_integrated_report("cover", [vc, nsx, avi], [module(vc), module(nsx), module(avi)], "综合封面", legacy_report=legacy_report())
    engine = PdfReportEngine()
    expected = "".join(engine._build_payload(report, {})["narratives"]["executive"].split())[:24]
    pdf = engine.render(report, tmp_path / "cover.pdf")
    with pymupdf.open(pdf) as document:
        summary_pages = ["".join(page.get_text().split()) for page in document if "执行摘要" in page.get_text()]
    assert summary_pages and all(expected in text for text in summary_pages)
